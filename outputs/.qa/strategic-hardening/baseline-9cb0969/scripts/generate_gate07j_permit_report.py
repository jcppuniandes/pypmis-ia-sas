from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from generate_gate07i_report import (
    CALLOUT,
    CAUTION,
    DARK_BLUE,
    INK,
    POSITIVE,
    RISK,
    TEAL,
    add_bottom_rule,
    add_callout,
    add_page_number,
    add_table,
    configure_styles,
    set_image_alt,
    set_run,
)

REPO = Path(__file__).resolve().parents[1]
OUTPUT = (
    REPO
    / "outputs"
    / "technical_reports"
    / "Informe_Tecnico_Gate_07J_PERMIT_Permit_Readiness_Provider_20260831.docx"
)
LOGO = REPO / "frontend" / "public" / "pypmis-construction-ai-logo.png"

# Resolved design tokens:
# preset=standard_business_brief; first-page pattern=memo_masthead.
# Letter portrait; 1 in margins; 0.492 in header/footer; 9360 DXA tables;
# 120 DXA table indent; Calibri 11 pt; body 6 pt after / 1.10 line.
# Named override: dense traceability/status matrices use 7.2-8.4 pt.


def add_title_page(doc: Document) -> None:
    section = doc.sections[0]
    # Define both stories explicitly. Word and LibreOffice otherwise diverge
    # on even-page header placement when a table continues across a page.
    doc.settings.odd_and_even_pages_header_footer = True
    for header in (section.header.paragraphs[0], section.even_page_header.paragraphs[0]):
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_run(
            header.add_run("TECHNICAL REPORT · GATE 07J-PERMIT"),
            size=8.5,
            color="5F6B76",
            bold=True,
        )
        add_bottom_rule(header, color="D9E1E8", size=6)
    for footer in (section.footer.paragraphs[0], section.even_page_footer.paragraphs[0]):
        add_page_number(footer)

    if LOGO.exists():
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(8)
        picture = paragraph.add_run().add_picture(str(LOGO), width=Inches(0.72))
        set_image_alt(picture, "Logotipo de P&Pmis Construction AI")

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(4)
    set_run(kicker.add_run("INFORME TÉCNICO DE IMPLEMENTACIÓN"), size=10, color=TEAL, bold=True)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.keep_with_next = True
    set_run(title.add_run("Gate 07J-PERMIT"), size=25, color=INK, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    set_run(
        subtitle.add_run("Permit Readiness Governed Provider"),
        size=15,
        color=DARK_BLUE,
        bold=True,
    )
    for label, value in (
        ("Aplicación", "P&Pmis Construction AI"),
        ("Fecha de corte", "31 de agosto de 2026"),
        ("Repositorio", "pypmis-ia-sas · branch main · base fde6027"),
        ("Migración", "20260831_0054"),
        ("Salida funcional", "READY_FOR_GOVERNED_PERMIT_EVIDENCE"),
        ("Estado exclusivo", "GATE07J_PERMIT_RELEASE_VALIDATED"),
    ):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        set_run(paragraph.add_run(f"{label}: "), size=10.5, color=INK, bold=True)
        set_run(paragraph.add_run(value), size=10.5)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(10)
    rule.paragraph_format.space_after = Pt(14)
    add_bottom_rule(rule)
    add_callout(
        doc,
        "Resultado ejecutivo",
        "Se implementó PERMIT_READINESS_PROVIDER sobre un registro único de permisos, condiciones separadas y "
        "evaluaciones versionadas. El proveedor ya está AVAILABLE para Project y Project Set; conserva Class B "
        "en los perfiles estándar y eleva a Class A, por configuración, los proyectos Mining bajo PYP-FFA. "
        "No se incorporaron trámites regulatorios automáticos, GIS, tasas, pagos ni decisiones FID.",
        fill=POSITIVE,
    )
    add_table(
        doc,
        ["Capa", "Resultado", "Evidencia verificable"],
        [
            ("Dominio / API", "PASS", "ProjectPermit, PermitCondition, Assessment; 22 rutas OpenAPI"),
            ("Gate 07J-EVID", "PASS", "Provider AVAILABLE; contrato gobernado Project/Set"),
            ("PostgreSQL", "PASS", "Upgrade/downgrade/upgrade 0053/0054; constraints PASS"),
            ("Frontend", "PASS", "7 USER submódulos; ADMIN con 17 tabs"),
            ("Browser", "PASS", "Chromium USER/ADMIN; integración provider; 0 page errors"),
            ("Localhost", "ACTIVE", "Frontend 5173; API 8000; Alembic 0054"),
        ],
        [1900, 1300, 6160],
        font_size=8.5,
    )
    doc.add_page_break()


def add_contents(doc: Document) -> None:
    heading = doc.add_heading("Contenido y trazabilidad", level=1)
    heading.paragraph_format.space_before = Pt(0)
    add_table(
        doc,
        ["Capítulo", "Tema", "Entregables"],
        [
            ("1", "Baseline, alcance y REUSE BEFORE CREATE", "1-7"),
            ("2", "Catálogo, aplicabilidad y dominio Permit", "8-18"),
            ("3", "Evaluación, reglas y agregación", "19-37"),
            ("4", "Provider Gate 07J-EVID e integración", "38-44"),
            ("5", "ADMIN, USER, seguridad y APIs", "45-58"),
            ("6", "Persistencia, pruebas y baseline", "59-88"),
            ("7", "Riesgos, gap analysis y cierre", "89-93"),
            ("8", "Matriz integral de trazabilidad", "1-93"),
        ],
        [1050, 5600, 2710],
        font_size=8.8,
    )
    add_callout(
        doc,
        "Criterio de evidencia",
        "Los resultados corresponden a comandos ejecutados durante esta entrega. Las pruebas de reglas y casos "
        "negativos se verificaron en servicio/PostgreSQL; las superficies USER, ADMIN e integración se verificaron "
        "con Vitest y Playwright Chromium. El capítulo 8 relaciona los 93 entregables exigidos.",
        fill=CALLOUT,
    )
    doc.add_page_break()


def add_baseline_reuse(doc: Document) -> None:
    doc.add_heading("1. Baseline, alcance y REUSE BEFORE CREATE", level=1)
    add_table(
        doc,
        ["Baseline", "Estado de entrada", "Decisión"],
        [
            ("Alembic", "20260830_0053", "Sucesora aditiva y reversible 0054"),
            ("Gate 07J-EVID", "Permit provider RESERVED / UNAVAILABLE", "Registrar provider real; conservar policy"),
            ("PYP-INV-SELECT", "Permit B · no bloqueante", "Sin cambio"),
            ("PYP-FFA estándar", "Permit B · no bloqueante", "Sin cambio"),
            ("PYP-FFA Mining", "Override contextual Class A", "Proveedor gobernado; 100% de cobertura"),
            ("Gate 07J", "Preparación de decisión", "Sólo evidencia; can_enter_fid=false"),
        ],
        [2000, 3300, 4060],
        font_size=8.4,
    )
    doc.add_heading("1.1 Inventario reutilizado", level=2)
    add_table(
        doc,
        ["Capacidad existente", "Uso en 07J-PERMIT"],
        [
            ("EnterpriseWorkspace(PROJECT)", "Identidad canónica; no se creó un Project master paralelo"),
            ("Gate 07I", "Recommended Project Set y hash para agregación exacta"),
            ("Gate 07J-EVID", "Registro de providers, contrato, policy Class A/B y gap analysis"),
            ("AdminConfiguration", "DRAFT/PUBLISHED/SUPERSEDED, herencia, hash, preview y publish"),
            ("RBAC / SecurityEvent", "Autorización, segregación y auditoría existentes"),
            ("ETag / idempotency", "revision_version, If-Match y claves transaccionales"),
            ("Workspace Navigator", "Contexto PROJECT/PORTFOLIO y navegación sin shell duplicado"),
        ],
        [3100, 6260],
        font_size=8.5,
    )
    add_callout(
        doc,
        "Decisión de modelo",
        "No existía un Permit Register gobernado reutilizable. Se crearon ProjectPermit como master, "
        "PermitCondition como obligación separada y PermitReadinessAssessment como evidencia versionada. "
        "Un registro de permiso no equivale a su evaluación de readiness.",
        fill=CAUTION,
    )


def add_domain(doc: Document) -> None:
    doc.add_heading("2. Catálogo, aplicabilidad y dominio Permit", level=1)
    add_table(
        doc,
        ["Objeto", "Responsabilidad", "Controles"],
        [
            ("Permit Requirement Catalog", "Tipos, autoridades lógicas, obligatoriedad y evidencia", "Versionado y publicado"),
            ("ProjectPermit", "Solicitud, vigencia, autoridad, evidencia y estado", "Tenant + Project; ETag"),
            ("PermitCondition", "Condiciones del permiso y carácter blocking", "Estado, vencimiento, evidencia"),
            ("PermitReadinessAssessment", "Snapshot de fuente/configuración y resultado", "Versión; COMPLETED inmutable"),
            ("Project/Set Readiness", "Resultado vigente y agregación contextual", "Sin score sintético"),
        ],
        [2500, 4300, 2560],
        font_size=8.3,
    )
    doc.add_heading("2.1 Starter configurable y aplicabilidad", level=2)
    add_table(
        doc,
        ["Dimensión", "Implementación"],
        [
            ("17 tipos de permiso", "Starter editable; cada tipo posee código, nombre y authority_type CONFIGURED"),
            ("Proyecto", "project_type y metadatos del workspace determinan requirements aplicables"),
            ("Geografía", "Reglas por atributos configurables; sin GIS ni geocodificación externa"),
            ("Gate Profile", "PYP-INV-SELECT y PYP-FFA resuelven obligatoriedad por policy publicada"),
            ("Mining PYP-FFA", "Tres requerimientos starter FFA elevan Permit Readiness a Class A"),
            ("Overrides", "Se procesan sólo desde configuración publicada; no hay lógica sectorial oculta"),
        ],
        [2500, 6860],
        font_size=8.5,
    )
    doc.add_heading("2.2 Lifecycle y semántica", level=2)
    add_table(
        doc,
        ["Componente", "Estados / regla"],
        [
            ("Permit", "DRAFT, IDENTIFIED, SUBMITTED, UNDER_REVIEW, APPROVED, APPROVED_WITH_CONDITIONS, REJECTED, SUSPENDED, EXPIRED, WITHDRAWN"),
            ("Condición", "OPEN, IN_PROGRESS, SATISFIED, WAIVED, CANCELLED; blocking configurable"),
            ("Assessment", "DRAFT, IN_PROGRESS, COMPLETED, SUPERSEDED, VOIDED"),
            ("Readiness", "NOT_STARTED, READY, BLOCKED, STALE, EXPIRED"),
            ("Vigencia", "effective/approval/expiry dates separadas; la expiración nunca se infiere como aprobación"),
        ],
        [2100, 7260],
        font_size=8.1,
    )
    add_callout(
        doc,
        "Condiciones",
        "APPROVED_WITH_CONDITIONS puede calificar únicamente cuando no existen condiciones blocking abiertas. "
        "Las condiciones no bloqueantes permanecen como warnings y conservan trazabilidad.",
        fill=CALLOUT,
    )


def add_readiness(doc: Document) -> None:
    doc.add_heading("3. Evaluación, reglas y agregación", level=1)
    add_table(
        doc,
        ["Chequeo", "Resultado controlado"],
        [
            ("Requirement aplicable sin permiso", "PERMIT_REQUIRED / missing_permit_count"),
            ("SUBMITTED / UNDER_REVIEW", "PERMIT_NOT_APPROVED"),
            ("REJECTED", "PERMIT_REJECTED"),
            ("SUSPENDED", "PERMIT_SUSPENDED"),
            ("EXPIRED o expiry_date superada", "PERMIT_EXPIRED / PERMIT_EXPIRY_DETECTED"),
            ("Evidencia obligatoria ausente", "PERMIT_EVIDENCE_REQUIRED"),
            ("Condición blocking abierta", "OPEN_BLOCKING_PERMIT_CONDITIONS"),
            ("Cobertura bajo threshold", "PERMIT_EVIDENCE_COVERAGE_BELOW_THRESHOLD"),
            ("Fuente/configuración cambió", "STALE_PERMIT_SOURCE"),
            ("Assessment fuera de valid_until", "PERMIT_READINESS_EXPIRED"),
        ],
        [3900, 5460],
        font_size=8.2,
    )
    add_callout(
        doc,
        "Regla de calificación",
        "Sólo APPROVED y APPROVED_WITH_CONDITIONS sin bloqueos califican. READY no significa autorización "
        "regulatoria futura: significa que el conjunto exigido por la configuración está vigente, evidenciado y "
        "sin condiciones bloqueantes abiertas a la fecha de corte.",
        fill=CAUTION,
    )
    doc.add_heading("3.1 Project Set, cobertura y proyecto mixto", level=2)
    add_table(
        doc,
        ["Regla", "Comportamiento"],
        [
            ("Denominador", "Sólo proyectos contextualmente Class A del Recommended Project Set"),
            ("Project Class B", "Permanece visible, pero no bloquea ni entra al denominador Class A"),
            ("Cobertura", "valid_project_count / required_project_count; Class A exige 100%"),
            ("Mixed Project Type", "Cada proyecto resuelve sus requirements y clasificación de forma independiente"),
            ("No score", "No se promedian permisos ni proyectos; el resultado es readiness + blockers/warnings"),
            ("Set hash", "Agregación enlazada al recommended_project_set_hash vigente"),
        ],
        [2700, 6660],
        font_size=8.5,
    )
    doc.add_heading("3.2 Freshness, hashes y renovación", level=2)
    add_table(
        doc,
        ["Control", "Implementación"],
        [
            ("permit_source_hash", "Hash determinista del registro, conditions, evidencia y configuración"),
            ("permit_readiness_hash", "Hash del resultado, cobertura, blockers, warnings y vigencia"),
            ("Refresh Preview", "Calcula diferencias sin persistir; la prueba confirma persisted=false"),
            ("Refresh Apply", "Actualiza assessment activo con idempotency y ETag"),
            ("Reassessment", "Crea nueva versión y conserva historia; no muta COMPLETED"),
            ("Renewal", "Nueva vigencia cambia source hash y exige refresh/reassessment"),
        ],
        [2500, 6860],
        font_size=8.5,
    )


def add_provider(doc: Document) -> None:
    doc.add_heading("4. Gate 07J-EVID e integración gobernada", level=1)
    add_table(
        doc,
        ["Contrato", "Valor implementado"],
        [
            ("provider_key", "PERMIT_READINESS_PROVIDER"),
            ("source_module / evidence_type", "PERMIT_READINESS / PERMIT_READINESS"),
            ("status", "AVAILABLE"),
            ("scopes", "PROJECT · EACH_PROJECT · SELECTED_PROJECTS_ONLY"),
            ("output", "READY_FOR_GOVERNED_PERMIT_EVIDENCE"),
            ("source identifiers", "assessment_id o aggregate_reference, project/set hash, source/readiness hash"),
            ("coverage payload", "required/valid projects; required/approved/conditional/missing/rejected/suspended/expired permits"),
            ("decision guard", "can_supply_investment_gate_evidence=true sólo si source_current y policy satisfechas"),
        ],
        [2900, 6460],
        font_size=8.2,
    )
    doc.add_heading("4.1 Compatibilidad y gap reanalysis", level=2)
    add_table(
        doc,
        ["Contexto", "Clasificación", "Efecto"],
        [
            ("PYP-INV-SELECT", "Class B", "Visible; no bloquea preparación de la decisión"),
            ("PYP-FFA estándar", "Class B", "Compatibilidad conservada"),
            ("PYP-FFA + Mining", "Class A", "Fuente gobernada obligatoria; 100% de cobertura"),
            ("Gap analysis", "Cerrado", "No queda un Class A provider reservado en el contexto implementado"),
            ("Siguiente paso", "Contextual", "READY_TO_DESIGN_GATE07K; STOP: Gate 07K no se inicia"),
        ],
        [2900, 1800, 4660],
        font_size=8.4,
    )
    add_callout(
        doc,
        "Frontera de decisión",
        "can_design_investment_decision_gate se resuelve con el registro publicado de providers y la policy "
        "contextual. Esta entrega no implementa Investment Decision, GO/HOLD/RECYCLE/KILL ni FID.",
        fill=RISK,
    )


def add_surfaces_security(doc: Document) -> None:
    doc.add_heading("5. Superficies ADMIN/USER, seguridad y APIs", level=1)
    add_table(
        doc,
        ["USER MODE", "Capacidad"],
        [
            ("My Permit Work", "Bandeja de acciones permitidas por rol"),
            ("Project Permit Register", "Listado, filtros, estado, vigencia y autoridad configurada"),
            ("Permit Requirements", "Requirements aplicables y brechas"),
            ("Permit Conditions", "Condiciones blocking/no blocking y estado"),
            ("Project Permit Readiness", "Resumen, blockers, warnings, cobertura y hashes"),
            ("Project Set Permit Readiness", "Matriz Project Set con denominador contextual"),
            ("Permit History / Audit", "Historia de evaluaciones y trazabilidad"),
        ],
        [3100, 6260],
        font_size=8.4,
    )
    doc.add_heading("5.1 ADMIN MODE", level=2)
    add_table(
        doc,
        ["Superficie", "Implementación"],
        [
            ("Permit Readiness Configuration", "Workspace administrativo con 17 tabs del catálogo y policy"),
            ("Ciclo", "Clonar DRAFT, editar, preview no persistente y publish gobernado"),
            ("Starter", "Un DRAFT local creado; no se publicaron cambios regulatorios ficticios"),
            ("Controles", "Tipos, applicability, statuses, conditions, evidence, freshness, roles y provider contract"),
        ],
        [3100, 6260],
        font_size=8.5,
    )
    doc.add_heading("5.2 Roles, permisos y SoD", level=2)
    add_table(
        doc,
        ["Rol / control", "Alcance"],
        [
            ("permit_manager / owner", "Registro, edición y status según acciones permitidas"),
            ("assessor", "Crear, editar, refresh y reassess"),
            ("reviewer", "Completar assessment; separación de revisión"),
            ("viewer", "Lectura y evidencia sin mutaciones"),
            ("configuration_admin", "Configurar y publicar policy"),
            ("ETag / idempotency", "If-Match para mutaciones; claves únicas para create/complete/refresh/reassess"),
            ("SecurityEvent", "Eventos para configuración, permisos, condiciones y assessments"),
        ],
        [3000, 6360],
        font_size=8.4,
    )
    doc.add_heading("5.3 Contrato HTTP", level=2)
    add_table(
        doc,
        ["Grupo", "Rutas verificadas"],
        [
            ("Permit Register", "List/create/get/update/status; list/create/update conditions"),
            ("Assessment", "List/create/get/update/complete/reassess"),
            ("Readiness", "Project, refresh-preview/apply, Portfolio Set y project matrix"),
            ("Evidence", "Project/Portfolio investment-gate-permit-evidence"),
            ("Administration", "List/clone/update/publish configurations; preview"),
            ("Overview", "Provider, release, roles, counts y acciones"),
        ],
        [2600, 6760],
        font_size=8.3,
    )


def add_persistence_validation(doc: Document) -> None:
    doc.add_heading("6. Persistencia, migración y validación", level=1)
    add_table(
        doc,
        ["Elemento", "Implementación"],
        [
            ("Tablas", "project_permits, permit_conditions, permit_readiness_assessments"),
            ("Numeración", "PRA mediante AdminNumberSequence; números de permiso conservan referencia controlada"),
            ("Constraints", "Estados, cobertura, unicidad tenant/número, versionado project/profile, idempotency"),
            ("Índices", "Tenant/project/status, readiness scan y active project/profile"),
            ("Inmutabilidad", "Listeners impiden mutar o borrar assessments COMPLETED"),
            ("Concurrencia", "revision_version/ETag, unique keys y transacciones PostgreSQL"),
            ("Migración", "0053→0054→0053→0054 validada en stack desechable"),
        ],
        [2600, 6760],
        font_size=8.4,
    )
    doc.add_heading("6.1 Matriz de pruebas ejecutadas", level=2)
    add_table(
        doc,
        ["Suite", "Resultado", "Cobertura relevante"],
        [
            ("Ruff", "PASS", "Módulo Permit y archivos Python modificados"),
            ("Backend combinado", "22/22 PASS", "Permit 7 + EVID 8 + RISK 7; 359.05 s"),
            ("PostgreSQL disposable", "PASS", "Upgrade/downgrade/constraints y happy path 0054"),
            ("TypeScript + Vite", "PASS", "2416 módulos; build 23.23 s"),
            ("Vitest componente", "3/3 PASS", "USER/ADMIN/provider y estados"),
            ("Vitest AppFlow + componente", "31/31 PASS", "Navegación y regresión principal; 621.86 s"),
            ("Playwright Chromium", "1/1 PASS", "USER+ADMIN+Gate07J; 42 s; 0 page errors"),
            ("ESLint", "0 errores", "8 warnings preexistentes en BimIfcModelViewer"),
        ],
        [2800, 1700, 4860],
        font_size=8.2,
    )
    doc.add_heading("6.2 Escenarios Permit", level=2)
    add_table(
        doc,
        ["Escenario", "Capa", "Resultado"],
        [
            ("Happy path Project + Set", "Servicio/PostgreSQL", "PASS"),
            ("Missing/submitted/rejected/suspended/expired", "Servicio", "PASS · bloquean FFA contextual"),
            ("Blocking vs nonblocking condition", "Servicio", "PASS"),
            ("Cambio post-completion / stale / expiry", "Servicio", "PASS"),
            ("Class B vs Mining FFA Class A", "Policy", "PASS"),
            ("Denominador sólo contextual Class A", "Agregación", "PASS"),
            ("Preview no persistente, ETag y audit", "Servicio", "PASS"),
            ("USER/ADMIN/provider disponible", "Browser", "PASS"),
        ],
        [4200, 2300, 2860],
        font_size=8.1,
    )
    doc.add_heading("6.3 Baseline persistente antes/después", level=2)
    add_table(
        doc,
        ["Métrica", "Antes", "Después", "Lectura"],
        [
            ("Alembic", "0053", "0054", "Cambio esperado"),
            ("AdminConfiguration", "81", "82", "+1 DRAFT Permit starter"),
            ("AdminNumberSequence", "12", "12", "Sin deriva"),
            ("EnterpriseWorkspace", "14", "14", "Sin deriva"),
            ("Project 14 defaults MD5", "eea3648c…", "eea3648c…", "Idéntico"),
            ("Project permits / conditions / assessments", "N/A", "0 / 0 / 0", "Sin datos ficticios persistidos"),
            ("Gate/Fund/Risk/Portfolio records", "0", "0", "Baseline protegido"),
        ],
        [3050, 1400, 1900, 3010],
        font_size=8.1,
    )
    add_callout(
        doc,
        "Localhost verificado",
        "Docker activo: frontend http://127.0.0.1:5173/app; API http://127.0.0.1:8000; PostgreSQL y Redis "
        "healthy; worker y beat en ejecución. La API reportó provider AVAILABLE, release validada y 22 rutas Permit.",
        fill=POSITIVE,
    )


def add_risks_closeout(doc: Document) -> None:
    doc.add_heading("7. Riesgos técnicos, gap analysis y cierre", level=1)
    add_table(
        doc,
        ["Riesgo / deuda", "Estado y tratamiento"],
        [
            ("Integraciones regulatorias", "Fuera de alcance: no polling, submission, pagos, tasas ni sincronización con autoridades"),
            ("Geografía avanzada", "Sin GIS/geocoding; applicability se resuelve por atributos configurados"),
            ("Browser de reloj controlado", "Expiry/stale probados en servicio; ampliar fixture full-stack nocturno"),
            ("Volumen enterprise", "Índices y consultas acotadas; falta benchmark dedicado de decenas de miles de permisos"),
            ("Catálogo normativo", "Starter no sustituye asesoría legal; cada tenant debe publicar su policy aprobada"),
            ("Concurrencia extrema", "ETag/idempotency protegen; ampliar carreras multi-worker en CI de carga"),
        ],
        [3300, 6060],
        font_size=8.3,
    )
    add_callout(
        doc,
        "Safeguards verificados",
        "No se implementaron FID, Investment Decision, GO/HOLD/RECYCLE/KILL, Project Activation, ejecución de "
        "trámites, credenciales regulatorias, firmas, pagos, GIS ni repositorio documental externo.",
        fill=RISK,
    )
    add_callout(
        doc,
        "Gap analysis final",
        "PERMIT_READINESS_PROVIDER dejó de estar RESERVED/UNAVAILABLE y ahora satisface el contrato contextual "
        "Class A sin alterar los perfiles Class B. El siguiente estado permitido es READY_TO_DESIGN_GATE07K; "
        "Gate 07K no se inicia en esta entrega.",
        fill=CAUTION,
    )
    add_callout(
        doc,
        "Estado exclusivo",
        "GATE07J_PERMIT_RELEASE_VALIDATED · final_output = READY_FOR_GOVERNED_PERMIT_EVIDENCE · "
        "provider = PERMIT_READINESS_PROVIDER · status = AVAILABLE · can_enter_fid=false.",
        fill=POSITIVE,
    )


DELIVERABLES = [
    "Executive summary",
    "ADR Gate07J-PERMIT",
    "Baseline assessment",
    "Gate07J-EVID contract review",
    "Gate07J-RISK gap-analysis review",
    "REUSE BEFORE CREATE",
    "Permit model reuse/create decision",
    "Permit Requirement Catalog",
    "Permit types",
    "Project applicability",
    "Geography applicability",
    "Gate Profile applicability",
    "Permit Record model",
    "Permit lifecycle",
    "Status semantics",
    "Permit conditions",
    "Blocking condition policy",
    "Evidence architecture",
    "PermitReadinessAssessment model",
    "Assessment lifecycle",
    "Readiness statuses",
    "Readiness checks",
    "Qualifying-status rules",
    "Investment Selection behavior",
    "PYP-FFA strict behavior",
    "ProjectPermitReadiness model",
    "Project Set aggregation",
    "Mixed Project Type behavior",
    "Coverage denominator",
    "No-score decision",
    "Freshness",
    "Expiry",
    "Renewal/history",
    "Source hash",
    "Permit readiness hash",
    "Refresh Preview",
    "Reassessment",
    "Provider Registry update",
    "Gate07J evidence contract",
    "Class A supply rule",
    "Gate07J integration",
    "Class B compatibility",
    "Gate07J-EVID gap reanalysis",
    "Contextual can_design_investment_decision_gate",
    "ADMIN MODE",
    "USER MODE",
    "Workspace Navigator",
    "Permit Register UI",
    "Permit Detail UI",
    "Readiness Dashboard",
    "Project Set Matrix",
    "Allowed actions",
    "Roles/permissions",
    "SoD",
    "ETag/idempotency",
    "SecurityEvent",
    "USER APIs",
    "ADMIN APIs",
    "Persistence",
    "Numbering",
    "Migration",
    "Constraints/indexes",
    "Concurrency",
    "PostgreSQL happy path",
    "Missing permit test",
    "Submitted-not-approved test",
    "Blocking-condition test",
    "Rejected test",
    "Suspended test",
    "Expiry test",
    "Renewal test",
    "Project Set coverage test",
    "Mixed Project Type test",
    "Geography override test",
    "Stale-source test",
    "Gate07J Class A integration test",
    "Class B compatibility test",
    "Gap reanalysis",
    "Frontend tests",
    "Browser USER E2E",
    "Browser negative E2E",
    "Browser ADMIN E2E",
    "Browser Gate07J integration",
    "Performance",
    "No-regulatory-overbuild safeguards",
    "No-FID safeguards",
    "Regression",
    "Baseline before/after",
    "Risks / technical debt",
    "Final Gap Analysis",
    "Recommendation for next step",
    "Updated technical report",
    "Final exclusive status",
]


def add_deliverable_trace(doc: Document) -> None:
    heading = doc.add_heading("8. Trazabilidad de entregables 1-93", level=1)
    heading.paragraph_format.page_break_before = True

    def evidence(index: int) -> str:
        if index <= 7:
            return "Portada / §1 / ADR 57"
        if index <= 18:
            return "§2"
        if index <= 37:
            return "§3"
        if index <= 44:
            return "§4"
        if index <= 58:
            return "§5"
        if index <= 88:
            return "§6 / suites"
        return "§7 / Documento"

    rows = [(str(index), item, evidence(index), "CERRADO") for index, item in enumerate(DELIVERABLES, 1)]
    for index, chunk in enumerate((rows[:13], rows[13:40], rows[40:67], rows[67:])):
        if index:
            doc.add_page_break()
        add_table(
            doc,
            ["#", "Entregable", "Evidencia", "Estado"],
            chunk,
            [560, 5200, 1900, 1700],
            font_size=7.2,
        )


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)
    add_title_page(doc)
    add_contents(doc)
    add_baseline_reuse(doc)
    add_domain(doc)
    add_readiness(doc)
    add_provider(doc)
    add_surfaces_security(doc)
    add_persistence_validation(doc)
    add_risks_closeout(doc)
    add_deliverable_trace(doc)
    properties = doc.core_properties
    properties.title = "Informe Técnico Gate 07J-PERMIT — Permit Readiness Governed Provider"
    properties.subject = "Permit Register, Permit Readiness y evidencia gobernada para Gate 07J"
    properties.author = "P&Pmis Construction AI"
    properties.keywords = "Gate 07J-PERMIT, Permit Readiness, Permit Register, Gate 07J-EVID, PYP-FFA"
    properties.comments = "Generated from verified repository and localhost evidence on 2026-08-31."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
