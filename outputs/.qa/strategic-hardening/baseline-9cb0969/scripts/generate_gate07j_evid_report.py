from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from generate_gate07i_report import (
    CALLOUT,
    CAUTION,
    DARK_BLUE,
    INK,
    POSITIVE,
    RISK,
    TEAL,
    add_bottom_rule,
    add_callout,
    add_page_number,
    add_table,
    configure_styles,
    set_image_alt,
    set_run,
)

REPO = Path(__file__).resolve().parents[1]
OUTPUT = (
    REPO
    / "outputs"
    / "technical_reports"
    / "Informe_Tecnico_Gate_07J_EVID_Investment_Evidence_Policy_20260830.docx"
)
LOGO = REPO / "frontend" / "public" / "pypmis-construction-ai-logo.png"

# Resolved token map:
# preset=standard_business_brief; first-page pattern=memo_masthead.
# Named overrides: compact matrix tables use 7.3–8.3 pt body text while
# retaining 9360 DXA width, 120 DXA indent and preset cell margins.


def add_title_page(doc: Document) -> None:
    section = doc.sections[0]
    doc.settings.odd_and_even_pages_header_footer = True
    for header in (section.header.paragraphs[0], section.even_page_header.paragraphs[0]):
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_run(header.add_run("TECHNICAL REPORT · GATE 07J-EVID"), size=8.5, color="5F6B76", bold=True)
        add_bottom_rule(header, color="D9E1E8", size=6)
    for footer in (section.footer.paragraphs[0], section.even_page_footer.paragraphs[0]):
        add_page_number(footer)

    if LOGO.exists():
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(8)
        picture = paragraph.add_run().add_picture(str(LOGO), width=Inches(0.72))
        set_image_alt(picture, "Logotipo de P&Pmis Construction AI")

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(4)
    set_run(kicker.add_run("INFORME TÉCNICO DE IMPLEMENTACIÓN"), size=10, color=TEAL, bold=True)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    set_run(title.add_run("Gate 07J-EVID"), size=25, color=INK, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    set_run(
        subtitle.add_run("Investment Decision Evidence Matrix & Readiness Policy"),
        size=14.5,
        color=DARK_BLUE,
        bold=True,
    )
    for label, value in (
        ("Aplicación", "P&Pmis Construction AI"),
        ("Fecha de corte", "30 de agosto de 2026"),
        ("Repositorio", "pypmis-ia-sas · branch main"),
        ("Alembic", "20260830_0052 (head, sin migración nueva)"),
        ("Estado exclusivo", "GATE07J_EVID_POLICY_VALIDATED"),
    ):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        set_run(paragraph.add_run(f"{label}: "), size=10.5, color=INK, bold=True)
        set_run(paragraph.add_run(value), size=10.5)
    rule = doc.add_paragraph()
    add_bottom_rule(rule)
    add_callout(
        doc,
        "Resultado ejecutivo",
        "Se implementó una política contextual A/B/C sobre AdminConfiguration, un registro de proveedores "
        "code-backed y la integración aditiva con Gate 07J. La solución impide degradar Class A a evidencia manual, "
        "resuelve overrides por Project Type, conserva Preparations históricas y no crea Investment Decision ni FID.",
        fill=POSITIVE,
    )
    add_table(
        doc,
        ["Capa", "Resultado", "Evidencia"],
        [
            ("Política y Gate 07J", "PASS", "A/B/C, snapshot, no fallback, historial inmutable"),
            ("Proveedores", "PASS", "7 entradas; 5 fuentes gobernadas reutilizadas; 2 reservadas"),
            ("PostgreSQL", "PASS", "Head 0052; 69/69 + 68/68; downgrade/upgrade"),
            ("Frontend", "PASS", "12 tabs ADMIN, TypeScript, Vitest 30/30"),
            ("Browser", "PASS", "Clone, Preview, 412, Save, Publish; consola/red controladas"),
            ("Localhost", "ACTIVE", "Frontend 5173; API 8000; draft inicial disponible"),
        ],
        [1900, 1300, 6160],
        font_size=8.8,
    )
    doc.add_page_break()


def add_baseline_and_reuse(doc: Document) -> None:
    doc.add_heading("1. Baseline, alcance y REUSE BEFORE CREATE", level=1)
    add_callout(
        doc,
        "Alcance aplicado",
        "Gate 07J-EVID es configuración + registro + validación. Gate 07J conserva ownership de Preparation, "
        "Evidence Items, readiness y handoff. La implementación no introduce un nuevo motor de readiness.",
    )
    add_table(
        doc,
        ["Capacidad inspeccionada", "Decisión de reutilización"],
        [
            ("Gate 07J profiles / requirements", "Extensión aditiva mediante snapshot de política efectiva"),
            ("Gate 07I", "Fuente gobernada de Portfolio Recommendation"),
            ("Gate 07G / 07G-A", "Budget Feasibility, dependencies y CAPEX/EXPENSE como detalle"),
            ("Gate 07H", "Fuente gobernada de Resource Feasibility"),
            ("Gate 07J-FEL", "Definition, FEL, PDRI y Cost Estimate Maturity"),
            ("Gate 07J-FUND", "Fuente gobernada de Funding Readiness"),
            ("AdminConfiguration", "DRAFT/PUBLISHED/SUPERSEDED, herencia, hash y revisión"),
            ("RBAC / SecurityEvent / ETag", "Patrones existentes; sin autoridad de inversión"),
        ],
        [3100, 6260],
        font_size=8.8,
    )
    doc.add_heading("1.1 Baseline funcional", level=2)
    add_table(
        doc,
        ["Gate", "Contrato vigente"],
        [
            ("07G", "READY_FOR_PORTFOLIO_BUDGET_RECOMMENDATION"),
            ("07G-A", "GATE07G_A_OPC_ALIGNMENT_VALIDATED"),
            ("07H", "READY_FOR_PORTFOLIO_RESOURCE_RECOMMENDATION"),
            ("07I", "READY_FOR_INVESTMENT_GATE_PREPARATION"),
            ("07J", "READY_FOR_INVESTMENT_DECISION; can_enter_fid=false"),
            ("07J-FEL", "GATE07J_FEL_RELEASE_VALIDATED"),
            ("07J-FUND", "GATE07J_FUND_RELEASE_VALIDATED"),
        ],
        [1600, 7760],
        font_size=8.9,
    )


def add_architecture(doc: Document) -> None:
    doc.add_heading("2. Arquitectura A/B/C y regla de no fallback", level=1)
    add_table(
        doc,
        ["Clase", "Semántica", "Validación", "Fallo"],
        [
            (
                "A",
                "GOVERNED_SOURCE_REQUIRED",
                "Provider AVAILABLE, contrato, scope, current y readiness",
                "Bloquea; manual prohibido",
            ),
            (
                "B",
                "GOVERNED_EVIDENCE_SUFFICIENT",
                "Referencia controlada, owner/issuer, fechas, hash y validación",
                "SUBMITTED no es VALID",
            ),
            ("C", "NOT_REQUIRED", "No crea requisito efectivo", "Sin blocker ni warning"),
        ],
        [800, 2600, 4000, 1960],
        font_size=8.4,
    )
    add_callout(
        doc,
        "Control crítico",
        "Class A nunca se convierte silenciosamente en Class B. Provider ausente produce SOURCE_PROVIDER_MISSING; "
        "contrato incompleto produce PROVIDER_CONTRACT_INVALID; scope incompatible produce SCOPE_UNSUPPORTED; "
        "source_current=false produce SOURCE_STALE.",
        fill=RISK,
    )
    doc.add_heading("2.1 Modelo de política", level=2)
    add_table(
        doc,
        ["Grupo", "Campos gobernados"],
        [
            ("Identidad", "requirement_key, evidence_type, classification, required, blocking"),
            ("Fuente", "provider_key, provider_required, minimum readiness/stage/maturity"),
            ("Contexto", "Gate Profile, Governance Model, Project Type, scope"),
            ("Freshness", "validity_period_days y current-with source flags"),
            ("Coverage", "minimum coverage; EACH_PROJECT no promedia Projects faltantes"),
            ("Manual", "allow_manual_reference y manual_validation_required sólo para B"),
        ],
        [2200, 7160],
        font_size=8.8,
    )


def add_providers(doc: Document) -> None:
    doc.add_heading("3. Provider Registry y validación de contrato", level=1)
    add_table(
        doc,
        ["Evidence", "Provider", "Source", "Status", "Scopes"],
        [
            ("Portfolio Recommendation", "GATE07I_PORTFOLIO_RECOMMENDATION", "Gate 07I", "AVAILABLE", "Portfolio / Set"),
            ("Budget Feasibility", "GATE07G_BUDGET_FEASIBILITY", "Gate 07G/A", "AVAILABLE", "Portfolio / Set"),
            ("Resource Feasibility", "GATE07H_RESOURCE_FEASIBILITY", "Gate 07H", "AVAILABLE", "Portfolio / Set"),
            ("Definition/FEL/PDRI/Estimate", "GATE07J_FEL_PROJECT_DEFINITION", "Gate 07J-FEL", "AVAILABLE", "Each / Selected"),
            ("Funding Readiness", "GATE07J_FUND_FUNDING_READINESS", "Gate 07J-FUND", "AVAILABLE", "Portfolio / Set"),
            ("Risk Readiness", "RISK_READINESS_PROVIDER", "Reserved", "UNAVAILABLE", "Portfolio / Set / Each"),
            ("Permit Readiness", "PERMIT_READINESS_PROVIDER", "Reserved", "UNAVAILABLE", "Each / Selected"),
        ],
        [1900, 2850, 1450, 1250, 1910],
        font_size=7.8,
    )
    add_callout(
        doc,
        "Contrato mínimo",
        "evidence_type · scope · readiness_status · source_current · evidence_hash · as_of_date · warnings · blockers. "
        "El registro es code-backed y la UI no acepta endpoints ejecutables arbitrarios.",
        fill=CALLOUT,
    )


def profile_rows(ffa: bool) -> list[tuple[str, ...]]:
    return [
        ("Portfolio Recommendation", "A", "Sí", "Gate07I", "Portfolio", "Contextual", "—"),
        ("Budget Feasibility", "A", "Sí", "Gate07G/A", "Portfolio", "Contextual", "—"),
        ("Resource Feasibility", "A" if ffa else "B", "Sí" if ffa else "No", "Gate07H" if ffa else "Gate07H / manual", "Portfolio", "Contextual", "—"),
        ("Project Definition", "A", "Sí", "Gate07J-FEL", "Each", "365 d", "100%"),
        ("FEL Stage", "A", "Sí", "Gate07J-FEL", "Each", "365 d", "100%"),
        ("PDRI / Definition", "A", "Sí", "Gate07J-FEL", "Each", "365 d", "100%"),
        ("Funding Readiness", "A" if ffa else "B", "Sí" if ffa else "No", "Gate07J-FUND" if ffa else "Gate07J-FUND / manual", "Portfolio", "180 d", "—"),
        ("Risk Readiness", "A" if ffa else "B", "Sí" if ffa else "No", "Risk provider" if ffa else "Manual validated", "Project Set", "90 d", "100%" if ffa else "—"),
        ("Permit Readiness", "B", "Sí" if ffa else "No", "Manual validated", "Each", "180 d", "100%" if ffa else "80%"),
        ("Contracting Strategy", "B", "Sí" if ffa else "No", "Manual validated", "Each", "365 d", "100%" if ffa else "80%"),
        ("Execution Plan", "B", "Sí" if ffa else "No", "Manual validated", "Each", "365 d", "100%" if ffa else "80%"),
        ("Schedule Maturity", "B", "Sí" if ffa else "No", "Manual validated", "Each", "180 d", "100%" if ffa else "80%"),
        ("Cost Estimate Maturity", "A" if ffa else "B", "Sí" if ffa else "No", "Gate07J-FEL" if ffa else "Manual validated", "Each", "180 d", "100%" if ffa else "80%"),
        ("Environmental / Social", "B" if ffa else "C", "No", "Manual / N/A", "Each", "180 d" if ffa else "—", "100%" if ffa else "—"),
        ("Land Access", "B" if ffa else "C", "No", "Manual / N/A", "Each", "180 d" if ffa else "—", "100%" if ffa else "—"),
        ("Stakeholder Readiness", "B" if ffa else "C", "No", "Manual / N/A", "Each", "180 d" if ffa else "—", "100%" if ffa else "—"),
        ("Technical Assurance", "B" if ffa else "C", "No", "Manual / N/A", "Each", "365 d" if ffa else "—", "100%" if ffa else "—"),
        ("Other Governed Evidence", "C", "No", "N/A", "Portfolio", "—", "—"),
    ]


def add_matrices(doc: Document) -> None:
    doc.add_heading("4. Matrices efectivas por Investment Gate Profile", level=1)
    doc.add_heading("4.1 PYP-INV-SELECT — Investment Selection Preparation", level=2)
    add_table(
        doc,
        ["Evidence", "Class", "Block", "Provider", "Scope", "Validity", "Coverage"],
        profile_rows(False),
        [2500, 600, 650, 1900, 1150, 1100, 1460],
        font_size=7.3,
    )
    doc.add_heading("4.2 PYP-FFA — Full Funds Authorization Preparation", level=2)
    add_table(
        doc,
        ["Evidence", "Class", "Block", "Provider", "Scope", "Validity", "Coverage"],
        profile_rows(True),
        [2500, 600, 650, 1900, 1150, 1100, 1460],
        font_size=7.3,
    )


def add_context_and_adrs(doc: Document) -> None:
    doc.add_heading("5. Contexto, overrides y decisiones ADR", level=1)
    add_table(
        doc,
        ["Resolución", "Comportamiento"],
        [
            ("PYP-FFA + Mining + Permit", "Base B → A; provider reservado; manual insuficiente; blocking 100%"),
            ("PYP-FFA + Building + Permit", "Conserva B; evidencia controlada requiere validación"),
            ("Mixed Mining + Building", "EACH_PROJECT resuelve ambos Projects sin colapsar tipo"),
            ("DIRECT_INTERNAL / CONTRACTOR_DELIVERY", "C por defecto; no vuelve obligatorio el gate"),
            ("CAPITAL_OWNER", "Aplica matriz publicada e inheritance contextual"),
        ],
        [3200, 6160],
        font_size=8.6,
    )
    doc.add_heading("5.1 Ocho decisiones obligatorias", level=2)
    add_table(
        doc,
        ["Domain", "INV-SELECT", "FFA / override", "Decision y acción"],
        [
            ("Risk", "B · 90 d", "A · block · 100%", "Provider missing; construir primero"),
            ("Permit", "B · 180 d · 80%", "B; Mining=A · block · 100%", "Manual base; provider después de Risk"),
            ("Contracting", "B · no block", "B · block · 100%", "Documento gobernado suficiente; no engine"),
            ("Execution Plan", "B · no block", "B · block · 100%", "PEP validado suficiente; no engine"),
            ("Schedule Maturity", "B · no block", "B · block · 100%", "Assessment suficiente; no CPM"),
            ("Cost Estimate", "B · manual", "A · FEL · block · 100%", "Reusar FEL; no duplicar fuente"),
            ("Environmental/Social", "C", "B · 180 d · 100%", "Validación manual; no nuevo módulo"),
            ("Land Access", "C", "B · 180 d · 100%", "Validación manual; no nuevo módulo"),
        ],
        [1900, 1900, 2600, 2960],
        font_size=8.0,
    )
    add_callout(
        doc,
        "Decisiones de detalle",
        "Project Dependencies permanece en el resumen de Budget Feasibility. CAPEX/EXPENSE/UNCLASSIFIED permanece "
        "como detalle presupuestal y de restricciones de funding; ninguno se convierte en evidence type independiente.",
        fill=CAUTION,
    )


def add_gate07j_integration(doc: Document) -> None:
    doc.add_heading("6. Integración con Gate 07J y compatibilidad histórica", level=1)
    add_table(
        doc,
        ["Momento", "Comportamiento"],
        [
            ("Nueva Preparation", "Obtiene política PUBLISHED efectiva por portfolio/profile"),
            ("Snapshot", "policy id/revision/hash, matrix, providers y Project contexts"),
            ("Class A", "Resuelve adapter gobernado; error permanece blocking; manual prohibido"),
            ("Class B", "Governed source opcional o SUBMITTED → VALIDATED manual"),
            ("Class C", "NOT_REQUIRED; sin blocker/warning"),
            ("Refresh", "Conserva snapshot; no auto-upgrade silencioso"),
            ("FROZEN / legacy", "Hash, profile, requirements y evidence no se reinterpretan"),
            ("Handoff", "Gate 07J sigue siendo owner; can_enter_fid=false"),
        ],
        [2600, 6760],
        font_size=8.7,
    )
    doc.add_heading("6.1 Validity, coverage, blocking y conditions", level=2)
    add_table(
        doc,
        ["Control", "Regla"],
        [
            ("Validity", "Por requirement; no existe expiración global única"),
            ("Freshness", "Current-with Project Set / Recommendation / Budget / FEL según evidence"),
            ("Coverage", "EACH_PROJECT valida todos; 2/3 = 66,67%; threshold 100% bloquea"),
            ("Expired B", "VALID vencida se vuelve EXPIRED y genera EVIDENCE_EXPIRED"),
            ("Conditions", "VALID_WITH_CONDITIONS sólo si Gate Profile lo permite; usa patrón Gate 07J"),
            ("Blocking", "A siempre bloquea si fuente inválida; B sólo cuando requirement.blocking"),
        ],
        [2200, 7160],
        font_size=8.8,
    )


def add_surfaces_and_api(doc: Document) -> None:
    doc.add_heading("7. Superficies ADMIN/USER, APIs y seguridad", level=1)
    add_callout(
        doc,
        "ADMIN MODE",
        "Enterprise Strategy Manager → Investment Governance → Investment Decision Evidence Policy. Se usa "
        "CompactModuleHeader y 12 tabs: Overview, Investment Gate Profiles, Evidence Matrix, Project Type Overrides, "
        "Provider Registry, Readiness Rules, Validity, Coverage, Blocking Policy, Missing Providers, Preview e History.",
        fill=POSITIVE,
    )
    add_table(
        doc,
        ["API", "Método", "Control"],
        [
            ("/investment-evidence-policy/overview", "GET", "Policy, providers, gaps, ADRs, boundaries"),
            ("/configurations", "GET", "Historia por revisión"),
            ("/configurations/{id}/clone", "POST", "If-Match; DRAFT"),
            ("/configurations/{id}", "PUT", "If-Match; 412 ETAG_MISMATCH"),
            ("/configurations/{id}/preview", "POST", "persisted=false"),
            ("/configurations/{id}/publish", "POST", "Supersede + immutable publish"),
            ("/providers[/{key}]", "GET", "Registry read-only"),
            ("/effective/preview", "POST", "Contextual dry-run"),
            ("/missing-providers", "GET", "Deterministic gap analysis"),
        ],
        [4650, 1000, 3710],
        font_size=8.2,
    )
    doc.add_heading("7.1 RBAC, ETag y auditoría", level=2)
    add_table(
        doc,
        ["Área", "Implementación"],
        [
            ("Permissions", "read · configure · publish · provider.read"),
            ("Roles", "investment_evidence_policy_admin / viewer; organization_admin heredado"),
            ("Actions", "view, clone, edit, preview, publish, view registry, view gaps"),
            ("Concurrency", "DRAFT writes usan ETag/If-Match; stale = 412"),
            ("SecurityEvent", "cloned, updated, published, superseded; sin ruido por reads"),
            ("USER", "Consume resultado dentro de Gate 07J; no proceso transaccional independiente"),
        ],
        [2500, 6860],
        font_size=8.7,
    )


def add_validation(doc: Document) -> None:
    doc.add_heading("8. Pruebas, PostgreSQL, browser y regresión", level=1)
    add_table(
        doc,
        ["Validación", "Resultado", "Cobertura"],
        [
            ("Backend policy + Gate 07J", "11/11 PASS", "A/B/C, mixed types, stale, 412, audit, historical"),
            ("PostgreSQL baseline suite", "69/69 PASS", "07D/07E y gobernanza transversal"),
            ("PostgreSQL portfolio suite", "68/68 PASS", "07E–07J-EVID + provider/service tests"),
            ("Migration", "PASS", "0052 → 0051 → 0052; EVID no añade schema"),
            ("Frontend targeted", "2/2 PASS", "Matrix, registry, dry-run, next provider"),
            ("Frontend AppFlow", "28/28 PASS", "USER/ADMIN navigation and existing modules"),
            ("Browser 07J-EVID", "1/1 PASS", "Clone, override, Preview, 412, Save, Publish"),
            ("Browser console/network", "PASS", "0 uncaught; 0 unexpected 4xx/5xx; controlled 412"),
            ("Ruff / ESLint / TypeScript", "PASS", "Backend and frontend static validation"),
        ],
        [3100, 1500, 4760],
        font_size=8.3,
    )
    doc.add_heading("8.1 Casos funcionales críticos", level=2)
    add_table(
        doc,
        ["Caso", "Resultado"],
        [
            ("Risk A + provider unavailable", "SOURCE_PROVIDER_MISSING; blocker; no manual fallback"),
            ("Permit B manual", "SUBMITTED hasta validación independiente; luego VALID"),
            ("Land Access C", "NOT_REQUIRED; sin blocker/warning"),
            ("Mining override", "PYP-FFA Permit B → A"),
            ("Mixed Set", "Mining y Building resueltos por Project"),
            ("Stale A", "SOURCE_STALE; no degradación"),
            ("Expired B", "EXPIRED por reglas Gate 07J"),
            ("Coverage", "No se promedian Projects faltantes"),
            ("Invalid provider contract", "PROVIDER_CONTRACT_INVALID"),
            ("Historical", "Preparation pre-policy queda intacta"),
        ],
        [3000, 6360],
        font_size=8.7,
    )


def add_baseline_after(doc: Document) -> None:
    doc.add_heading("9. Protección de baseline persistente", level=1)
    add_table(
        doc,
        ["Medida", "Antes", "Después", "Conclusión"],
        [
            ("Alembic head", "20260830_0052", "20260830_0052", "Sin migración"),
            ("AdminConfiguration", "79", "80", "+1 starter DRAFT esperado"),
            ("Gate Profiles", "2", "2", "Sin reemplazo de lifecycle"),
            ("EVID policy revisions", "0", "1", "Starter DRAFT, no auto-publish"),
            ("Provider registry", "0 tablas", "7 code-backed", "Sin endpoints ejecutables"),
            ("Gate07J Preparations / Items", "0 / 0", "0 / 0", "Preview no persistió"),
            ("Enterprise Workspaces", "14", "14", "Sin cambio"),
            ("Project ID14 hash", "1ba97904…", "1ba97904…", "Inalterado"),
            ("FEL / Funds / Funding", "0 / 0 / 0", "0 / 0 / 0", "Inalterado"),
            ("Budget / Resource Scenarios", "0 / 0", "0 / 0", "Inalterado"),
        ],
        [2500, 1900, 1900, 3060],
        font_size=8.0,
    )
    add_callout(
        doc,
        "Entorno efímero",
        "Las suites PostgreSQL se ejecutaron en PostgreSQL 16 con tmpfs y proyecto Compose aislado. Al terminar se "
        "eliminaron contenedores y red. El baseline local sólo recibió el único starter DRAFT solicitado.",
        fill=CALLOUT,
    )


def add_gap_and_closeout(doc: Document) -> None:
    doc.add_heading("10. Final Gap Analysis", level=1)
    add_table(
        doc,
        ["Grupo exacto", "Salida validada"],
        [
            (
                "AVAILABLE GOVERNED SOURCES",
                "Portfolio, Budget, Resource, Definition, FEL, PDRI, Funding y Cost Estimate (FFA vía FEL)",
            ),
            (
                "MISSING GOVERNED SOURCES",
                "Risk para PYP-FFA; Permit para PYP-FFA + Mining",
            ),
            (
                "GOVERNED EVIDENCE SUFFICIENT",
                "Risk Selection; Permit base; Contracting; Execution; Schedule; Cost Selection; FFA E&S/Land/Stakeholder/Technical",
            ),
            (
                "NOT REQUIRED",
                "Selection E&S/Land/Stakeholder/Technical; Other; governance fuera de CAPITAL_OWNER",
            ),
        ],
        [3000, 6360],
        font_size=8.4,
    )
    add_callout(
        doc,
        "Decisión única siguiente",
        "NEXT = BUILD_MISSING_GOVERNED_PROVIDER · provider = RISK_READINESS_PROVIDER. Es el primer gap blocking "
        "Class A para PYP-FFA. No se inicia su implementación en este gate.",
        fill=CAUTION,
    )
    add_callout(
        doc,
        "Estado exclusivo final",
        "GATE07J_EVID_POLICY_VALIDATED. final_output = READY_FOR_INVESTMENT_DECISION_EVIDENCE_POLICY. "
        "can_design_investment_decision_gate=false mientras exista el gap Risk; can_enter_fid=false siempre.",
        fill=POSITIVE,
    )
    add_callout(
        doc,
        "STOP boundary",
        "No Risk Manager, Permit Manager, Contracting Manager, Execution Plan Manager, Schedule Manager, Cost "
        "Estimate Manager, Investment Decision, FID, GO/HOLD/RECYCLE/KILL, Project Initialization, Activation o Execution.",
        fill=RISK,
    )


def add_deliverable_trace(doc: Document) -> None:
    doc.add_heading("11. Trazabilidad de entregables 1–41", level=1)
    items = [
        "Executive summary",
        "ADR Gate07J-EVID",
        "Baseline assessment",
        "REUSE BEFORE CREATE",
        "A/B/C architecture",
        "No-silent-fallback rule",
        "Gate Profile integration",
        "Evidence policy model",
        "Provider Registry",
        "Existing provider inventory",
        "Provider contract validation",
        "PYP-INV-SELECT matrix",
        "PYP-FFA matrix",
        "Project Type override matrix",
        "Mixed Project Set behavior",
        "Risk ADR",
        "Permit ADR",
        "Contracting ADR",
        "Execution Plan ADR",
        "Schedule Maturity ADR",
        "Cost Estimate Maturity ADR",
        "Environmental/Social ADR",
        "Land Access ADR",
        "Dependency reuse decision",
        "CAPEX/EXPENSE evidence decision",
        "Validity/coverage/blocking rules",
        "Missing Provider Analysis",
        "Gate07J integration",
        "Historical compatibility",
        "ADMIN/USER surfaces",
        "APIs",
        "Persistence/migration decision",
        "Tests",
        "Browser ADMIN E2E",
        "Browser Gate07J integration",
        "Regression",
        "Baseline before/after",
        "Final Gap Analysis",
        "Next-step decision",
        "Updated technical report",
        "Final exclusive status",
    ]
    section_map = (
        ["Portada"]
        + ["ADR 55"]
        + ["§1"] * 2
        + ["§2"] * 2
        + ["§4 / §6"]
        + ["§2"]
        + ["§3"] * 3
        + ["§4.1", "§4.2", "§5", "§5"]
        + ["§5.1"] * 8
        + ["§5.1", "§5.1", "§6.1", "§10", "§6", "§6", "§7", "§7", "§2 / ADR 55", "§8", "§8", "§8", "§8", "§9", "§10", "§10", "Documento", "§10"]
    )
    rows = [(str(index), item, section_map[index - 1], "CERRADO") for index, item in enumerate(items, 1)]
    add_table(
        doc,
        ["#", "Entregable", "Evidencia", "Estado"],
        rows,
        [600, 5000, 1760, 2000],
        font_size=7.8,
    )


def normalize_single_running_header(section, settings) -> None:
    """Make one default header/footer authoritative in Word and LibreOffice."""

    # LibreOffice's DOCX importer treats a lone default reference as an
    # odd-page header/footer, so the generator creates distinct but identical
    # even-page parts. Remove only first-page variants inherited from templates.
    settings.odd_and_even_pages_header_footer = True
    section.different_first_page_header_footer = False
    sect_pr = section._sectPr
    for child in list(sect_pr):
        if child.tag not in {qn("w:headerReference"), qn("w:footerReference")}:
            continue
        if child.get(qn("w:type")) == "first":
            sect_pr.remove(child)


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)
    add_title_page(doc)
    add_baseline_and_reuse(doc)
    add_architecture(doc)
    add_providers(doc)
    add_matrices(doc)
    add_context_and_adrs(doc)
    add_gate07j_integration(doc)
    add_surfaces_and_api(doc)
    add_validation(doc)
    add_baseline_after(doc)
    add_gap_and_closeout(doc)
    add_deliverable_trace(doc)
    normalize_single_running_header(section, doc.settings)

    properties = doc.core_properties
    properties.title = "Informe Técnico Gate 07J-EVID — Investment Decision Evidence Policy"
    properties.subject = "Política contextual A/B/C, proveedores gobernados e integración con Gate 07J"
    properties.author = "P&Pmis Construction AI"
    properties.keywords = "Gate 07J-EVID, Evidence Policy, A/B/C, Provider Registry, Investment Gate Preparation"
    properties.comments = "Generated from verified repository evidence on 2026-08-30."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
