from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from gate07g_report_data import SECTIONS as GATE07G_SECTIONS


REPO = Path(__file__).resolve().parents[1]
OUTPUT = (
    REPO
    / "outputs"
    / "technical_reports"
    / "Informe_Tecnico_Gate_07G_Portfolio_Budget_Planning_20260825.docx"
)
LOGO = REPO / "frontend" / "public" / "pypmis-construction-ai-logo.png"
E2E_ROOT = (
    REPO
    / "frontend"
    / "test-results"
    / "portfolio-budget-Gate-07G--d4100-os-comparison-and-readiness-chromium"
)
USER_EVIDENCE = E2E_ROOT / "gate07g-user-1.png"
ADMIN_EVIDENCE = E2E_ROOT / "gate07g-admin-1.png"

# Resolved document token map:
# preset=standard_business_brief; first-page pattern=memo_masthead.
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
BLACK = "000000"
MUTED = "5F6B76"
WHITE = "FFFFFF"
LIGHT = "F2F4F7"
ALT_LIGHT = "FAFBFC"
CALLOUT = "F4F6F9"
POSITIVE = "EAF7EF"
CAUTION = "FFF5E6"
RISK = "FDECEC"
TEAL = "0D8F8B"


def set_run(run, *, size=11, color=BLACK, bold=None, italic=False):
    run.font.name = "Calibri"
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:ascii"), "Calibri")
    r_pr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    table.autofit = False
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    for tag, value in (("tblW", total), ("tblInd", indent)):
        node = tbl_pr.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tbl_pr.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def mark_header_row(row):
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(repeat)


def add_table(doc, headers, rows, widths, *, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    mark_header_row(table.rows[0])
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, LIGHT)
        for run in cell.paragraphs[0].runs:
            set_run(run, size=9.5, color=INK, bold=True)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                shade(cell, ALT_LIGHT)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_run(run, size=font_size, color=BLACK)
    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    return table


def add_callout(doc, title, body, *, fill=CALLOUT):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    mark_header_row(table.rows[0])
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_margins(cell, 120, 145, 120, 145)
    set_run(cell.paragraphs[0].add_run(title.upper()), size=9, color=DARK_BLUE, bold=True)
    paragraph = cell.add_paragraph(body)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.1
    for run in paragraph.runs:
        set_run(run, size=10.2, color=BLACK)
    set_table_geometry(table, [9360])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(paragraph.add_run("P&Pmis Construction AI  |  "), size=8, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(begin)
    paragraph.add_run()._r.append(instruction)
    paragraph.add_run()._r.append(separate)
    display_run = paragraph.add_run()
    set_run(display_run, size=8, color=MUTED)
    display_run._r.append(text)
    paragraph.add_run()._r.append(end)


def add_bottom_rule(paragraph, color=BLUE, size=14):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def add_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("start", "1"), ("numFmt", "bullet"), ("lvlText", "•"), ("suff", "tab")):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:val"), value)
        level.append(node)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    p_pr.append(indent)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_item(doc, text, num_id):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.167
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)
    set_run(paragraph.add_run(text), size=11, color=BLACK)


def add_heading(doc, text, *, page_break=False):
    paragraph = doc.add_heading(text, level=1)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = page_break


def add_body(doc, text):
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1


def add_figure(doc, path, caption, alt_text, *, width=5.15):
    if not path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    picture = paragraph.add_run().add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("descr", alt_text)
    caption_paragraph = doc.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_before = Pt(4)
    caption_paragraph.paragraph_format.space_after = Pt(4)
    caption_paragraph.paragraph_format.keep_with_next = False
    for run in caption_paragraph.runs:
        set_run(run, size=9, color=MUTED, italic=True)


SECTIONS = [
    (1, "Executive summary", "Gate 07F quedó implementado como capability analítica derivada dentro de Portfolio Planning. Consume únicamente portafolios CAPITAL_OWNER con salida READY_FOR_PORTFOLIO_ANALYSIS y produce el contrato READY_FOR_PORTFOLIO_SCENARIO_PLANNING sin aprobar inversión ni iniciar ejecución.", "Backend, frontend, PostgreSQL efímero, CI y browser USER/ADMIN: PASS."),
    (2, "ADR Gate 07F", "Se documentó ADR 46 para fijar ownership, reutilización de Gate 07E, política de derivación bajo demanda, hashes, configuración gobernada y límites de alcance.", "docs/46-adr-gate07f-portfolio-analysis.md"),
    (3, "Baseline assessment", "El head permanece en 20260820_0045. Gate 07D/07E, AdminConfiguration, RBAC, SecurityEvent y Workspace Context ya resolvían identidad, membresía, evaluación, ranking y gobierno de configuraciones.", "No se alteró Project Workspace ni se fabricaron datos productivos."),
    (4, "Governance Model eligibility", "La elegibilidad exige governance_model=CAPITAL_OWNER, membership ACTIVE, mismo tenant, Portfolio válido y readiness 07E vigente. CONTRACTOR_DELIVERY, DIRECT_INTERNAL y Legacy/Unclassified se excluyen aun si poseen membership de prueba.", "Portfolio Membership por sí sola no habilita Gate 07F."),
    (5, "Stage-Gate / Portfolio Planning boundary", "Gate 07F no es un Stage-Gate independiente: compara y visualiza dentro de Portfolio Planning. Sus salidas son readiness analítico, no selección, financiación, FID ni autorización de ejecución.", "STOP conservado frente a Scenario, Budget, Resources, FEL/PDRI y Activation."),
    (6, "REUSE BEFORE CREATE", "Se reutilizaron ranking_projection de Gate 07E, latest COMPLETED evaluation, memberships, planning snapshots, AdminConfiguration, árbol empresarial, permisos, eventos, Recharts y patrones compactos de UI.", "No se creó motor BI, Candidate, ranking paralelo ni identidad Project nueva."),
    (7, "Portfolio Analysis Dataset", "El servicio deriva una fila por Portfolio + membership activa + Project. Expone únicamente campos fuente disponibles: identidad, scores, rank, objetivos, riesgo, ROM/currency, fechas, responsables, membership, versiones y hashes.", "Implementado en backend/app/modules/portfolio_analysis/service.py; sin tabla nueva."),
    (8, "Metric Catalog", "Se creó un catálogo controlado por categorías STRATEGIC, FINANCIAL_ROM, RISK, SCHEDULE, PORTFOLIO_SCORE, CLASSIFICATION y PROJECT_METADATA.", "No admite SQL, expresiones ni fórmulas arbitrarias."),
    (9, "Metric metadata", "Cada métrica declara key, label, tipo y compatibilidad con ejes, tamaño, color, filtro y Scorecard, además de agregación, formato y null behavior.", "Validación de claves y usos no permitidos cubierta por pruebas."),
    (10, "Investor Map model", "Investor Map usa X/Y numéricos controlados, burbuja acotada, color categórico, líneas de referencia y labels. El punto conserva rank/score 07E de solo lectura.", "Recharts responsive; lista accesible equivalente disponible."),
    (11, "Default Investor Map configuration", "El starter DRAFT configura Evaluation Score, Strategic Alignment, ROM Cost y Risk Bucket. Nunca se publica automáticamente.", "Default Portfolio Analysis Configuration: DRAFT REV 1."),
    (12, "Currency guardrails", "Cuando coexisten monedas sin normalización FX, el backend emite MULTI_CURRENCY_NOT_NORMALIZED, deshabilita el total ROM y la comparación monetaria de tamaño, y conserva valor original + currency.", "No se implementó conversión FX ni suma engañosa."),
    (13, "Axis, size and color rules", "Los ejes iniciales son numéricos; size exige métrica numérica no negativa y usa fallback visual para null; color sólo representa buckets/clasificaciones analíticas.", "Configuraciones de texto como size, keys desconocidas o axis no soportado se rechazan."),
    (14, "Reference lines and quadrants", "Las líneas vertical/horizontal configurables producen cuadrantes visuales. No modifican score, rank, elegibilidad ni decisiones.", "Valores inválidos se bloquean durante validación ADMIN."),
    (15, "Investor Map interactions", "Se implementaron tooltip, selección, detalle read-only, filtros y reset. La información crítica también está en lista/Scorecard para no depender del gráfico.", "Filtros preservan analysis_hash y no mutan fuentes."),
    (16, "Portfolio Scorecard", "Tabla contextual con rank, número, nombre, score, alineación, riesgo, Proposal Score, ROM/currency, fechas, tipo, objetivos, origen y versión.", "Rank y score permanecen read-only y provienen de Gate 07E."),
    (17, "Scorecard configuration", "ADMIN gobierna columnas, orden, sort, filtros, formato y buckets mediante AdminConfiguration.", "La UI no ofrece edición de datos fuente."),
    (18, "Analytical buckets", "Priority/Risk buckets se derivan con umbrales gobernados y etiquetas analíticas. Nunca se interpretan como SELECTED, FUNDED, APPROVED o REJECTED.", "Semántica de selección excluida por API, UI y pruebas."),
    (19, "Analysis Summary", "Overview utiliza encabezado compacto y chips para eligible/evaluated, score promedio, high priority, high risk, cobertura ROM y readiness.", "Totales monetarios sólo se muestran si la moneda es compatible."),
    (20, "Portfolio distributions and balance", "Se derivan distribuciones simples por tipo, objetivo, priority, risk y planning stage, sin optimización ni motor de recomendación.", "Complejidad lineal sobre el dataset; sin O(N²)."),
    (21, "Strategic Objective distribution handling", "Objetivos se tratan como clasificación N:M, permiten overlap y filtro, y no se convierten en jerarquía. El conteo de proyectos evita duplicar el total general.", "Multi-objective display cubierto en Scorecard."),
    (22, "Filters", "Filtros controlados incluyen Project Type, Strategic Objective, Risk Bucket, Priority Bucket y parámetros de análisis disponibles. El estado es de sesión.", "No se creó motor enterprise de saved views."),
    (23, "Drill-down", "El drawer enlaza Project Overview, evaluación, historia, contexto estratégico, membership y planning entry en modo lectura.", "Sin edición de Project/Evaluation desde análisis."),
    (24, "Multi-Portfolio behavior", "Todas las consultas se vinculan a portfolio_workspace_id. Un mismo Project conserva posición, score, rank y configuración independientes en cada Portfolio.", "Prueba PostgreSQL multi-Portfolio: PASS."),
    (25, "Membership interaction", "Sólo memberships ACTIVE entran al dataset actual. Al desactivar una membership desaparece del análisis sin borrar historial de evaluación.", "Membership-set hash detecta cambios."),
    (26, "Reevaluation interaction", "Se consume la última evaluación COMPLETED vigente. Una SUPERSEDED se excluye; una reevaluación IN_PROGRESS respeta el readiness 07E y no mezcla estados.", "Evaluation-set hash se recalcula determinísticamente."),
    (27, "Staleness and hash strategy", "analysis_hash encadena portfolio, membership_set_hash, evaluation_set_hash, prioritization_hash y configuration_hash. Un hash previo devuelve 412 STALE_PORTFOLIO_ANALYSIS_SOURCE.", "Filtros de visualización no cambian la identidad del análisis."),
    (28, "Analysis Snapshot persistence decision", "Se eligió derivación bajo demanda: no existe tabla ni endpoint de snapshot. La evidencia estable referencia hashes 07E + configuración.", "Sin migración; sin pixels/chart images persistidos."),
    (29, "ADMIN MODE", "Portfolio Analysis Configuration incorpora 11 tabs: General, Metric Catalog, Investor Map, Scorecard, Buckets, Reference Lines, Filters, Readiness, Inheritance, Permissions y Preview.", "Clone, Preview, Save y Publish usan patrones AdminConfiguration existentes."),
    (30, "USER MODE", "Portfolio Analysis ofrece Overview, Investor Map y Scorecard en Enterprise Strategy Manager > Portfolio Manager.", "Navegación compacta, filtros, drawer y readiness disponibles."),
    (31, "Workspace Navigator", "Se agregó Portfolio Analysis al contexto elegible y se mantuvo fuera Budget/Resource/FID. La navegación lateral usa el catálogo de aplicación vigente.", "frontend/src/navigation/applicationNavigation.ts y workspace_context/navigator.py."),
    (32, "Configuration inheritance", "La resolución conserva Tenant > Enterprise > Business Unit > Portfolio; gana el override publicado más cercano. Preview informa path, source workspace, revisión y hash.", "Starter DRAFT no gobierna USER production."),
    (33, "Allowed actions", "Backend calcula view analysis/map/scorecard, preview y readiness para Scenario Planning. Snapshot create permanece falso porque no se implementó persistencia.", "Frontend consume acciones y no infiere autoridad."),
    (34, "Roles and permissions", "Se añadieron/reutilizaron portfolio_analysis.read, snapshot.create, admin.configure y admin.publish con roles viewer, analyst, configuration_admin y portfolio_manager donde aplica.", "Sin rol FID approver."),
    (35, "ETag and idempotency", "Clone/update/publish respetan revisión optimista; If-Match obsoleto devuelve 412 ETAG_MISMATCH. Preview es no persistente y no requiere idempotencia de snapshot.", "Publish race y stale edit cubiertos."),
    (36, "SecurityEvent", "Se registran preview, configuration cloned/updated/published y stale source; no se auditan hover ni filtros ordinarios.", "Eventos acotados a acciones de gobierno y riesgo."),
    (37, "APIs", "USER expone analysis, dataset, investor-map, scorecard, preview y readiness. ADMIN expone list, preview, clone, update y publish sobre AdminConfiguration.", "Router registrado bajo /api/v1; contratos tipados en schemas.py."),
    (38, "Persistence", "Investor Map, Scorecard, rank y dataset son derivados. Gate 07F no agregó tablas ni reescribió la migración 0045.", "Persistencia limitada a AdminConfiguration y fuentes previas."),
    (39, "Frontend architecture", "La feature aislada portfolio-analysis contiene types, API client, USER workspace, ADMIN view, estilos y pruebas. App.tsx sólo enruta/compone.", "Recharts existente; no se añadió framework pesado."),
    (40, "Performance", "Shaping, filtros y agregados son lineales; no se observan algoritmos N². La API puede servir el dataset ya normalizado para 1.000 Projects.", "No sampling silencioso; Scorecard continúa como alternativa completa."),
    (41, "Portfolio Analysis Readiness", "READY exige handoff 07E actual, configuración publicada válida, dataset elegible no vacío y cobertura core. BLOCKED comunica blockers estructurados.", "Escenarios A–D de readiness cubiertos por unit/PostgreSQL/browser."),
    (42, "Metric coverage", "Core: evaluation score, strategic alignment y risk. ROM/currency sólo bloquea la rama Budget; fechas/prerrequisitos afectan las ramas correspondientes.", "La cobertura se reporta por métrica y por branch."),
    (43, "READY_FOR_PORTFOLIO_SCENARIO_PLANNING contract", "La salida incluye status, Portfolio, analysis_hash, hashes fuente/config, eligible count, coverage, warnings, blockers y acciones downstream.", "Final output obtenido en fixture de tres Projects."),
    (44, "Budget branch readiness contract", "can_enter_budget_planning depende de ROM coverage, compatibilidad monetaria y fechas requeridas. Multi-currency sin FX la bloquea sin bloquear el análisis general.", "No se implementaron escenarios ni presupuesto."),
    (45, "Resource branch readiness contract", "can_enter_resource_planning se calcula de manera independiente con fechas y prerrequisitos fuente.", "No existe dependencia secuencial forzada 07G→07H."),
    (46, "FEL parallel-branch protection", "Project Definition Readiness, can_enter_project_definition, FEL, PDRI y Definition Assessment no se leen ni mutan desde Gate 07F.", "Regresión previa conserva el branch paralelo."),
    (47, "No-selection/FID safeguards", "No existen controles o campos SELECTED, FORCED_IN/OUT, FUNDED, INVESTMENT_APPROVED, FID ni EXECUTION_AUTHORIZED.", "UI/API sólo visualizan y comparan."),
    (48, "Migration decision", "El head real permanece en 20260820_0045. Al no justificarse snapshot persistente ni esquema nuevo, la decisión correcta fue no crear migración.", "PostgreSQL up/down/up heredado y schema validator: PASS."),
    (49, "PostgreSQL E2E", "Fixture efímero con Portfolio, tres CAPITAL_OWNER, memberships activas, evaluaciones COMPLETED, readiness 07E y config publicada recorrió dataset, map, scorecard y readiness.", "Gate 07E/07F PostgreSQL suite: 17 passed."),
    (50, "Multi-Portfolio E2E", "El mismo Project fue analizado en dos Portfolios con fuentes, ranks, scores y map position contextuales independientes.", "Sin estado analítico global."),
    (51, "Currency E2E", "COP + USD sin normalización produjo warning, total ROM nulo y budget readiness false, conservando Resource/Scenario readiness cuando sus prerequisitos sí estaban presentes.", "MULTI_CURRENCY_NOT_NORMALIZED: PASS."),
    (52, "Stale-source E2E", "Tras cambiar membership/evaluation por camino gobernado, el hash antiguo fue rechazado con 412 y el reload produjo hash nuevo.", "Sin rank cache obsoleto."),
    (53, "Frontend tests", "Vitest completo queda en 199 pruebas aprobadas; TypeScript y build Vite pasan. Prettier pasa; ESLint reporta 0 errores y 8 warnings heredados del visor IFC dentro del umbral CI.", "PortfolioAnalysisWorkspace.test.tsx y AppFlow actualizados."),
    (54, "Browser USER E2E", "Login, USER MODE, Portfolio Analysis, tres Projects, Investor Map/lista, filtros, drawer, Scorecard, sort y readiness se validaron contra localhost real.", "Playwright Gate 07F USER/ADMIN: 1 passed."),
    (55, "Browser ADMIN E2E", "ADMIN recorrió config efectiva, clone DRAFT, selectores controlados, Scorecard, reference lines, preview no persistente, save, stale ETag y publish.", "No auto-publish; revisión/hash/herencia verificados."),
    (56, "Accessibility", "El Investor Map no es la única representación: lista y Scorecard contienen información core. Controles, tabs, selects y acciones tienen labels y son keyboard-reachable.", "Documento Word también incorpora headings, headers de tabla y alt text."),
    (57, "Concurrency", "La concurrencia se limita a mutaciones AdminConfiguration; edición/publicación obsoletas fallan por ETag y no duplican una revisión publicada.", "No snapshot concurrente porque la capability no persiste snapshots."),
    (58, "Regression", "Backend integral: 463 passed, 16 skipped; 28 setups afectados sólo por permisos del temp global de Windows se reejecutaron en temp del workspace: 48/48 passed. Frontend 199/199 y PostgreSQL heredado 69/69.", "Ruff, Prettier, TypeScript, build y git diff --check: PASS."),
    (59, "Baseline before and after", "El E2E fue efímero; teardown eliminó 13 workspaces, 9 memberships, 5 evaluations y 6 configs. La comparación oficial terminó baseline equal=true.", "No memberships, evaluations ni snapshots artificiales persistentes."),
    (60, "Risks and technical debt", "Ocho warnings ESLint heredados permanecen en BimIfcModelViewer. La rama Budget seguirá bloqueada para monedas múltiples sin fuente FX gobernada. Scorecard requerirá paginación/virtualización si el volumen supera el objetivo inicial.", "No hay errores Gate 07F abiertos ni deuda de migración."),
    (61, "Recommendation for next Gate", "Consumir el contrato READY_FOR_PORTFOLIO_SCENARIO_PLANNING en capacidades futuras independientes: Budget sólo cuando can_enter_budget_planning=true y Resources sólo cuando can_enter_resource_planning=true.", "Mantener analysis_hash como evidencia de entrada y no reabrir score/rank 07E."),
    (62, "Updated technical report", "Esta entrega consolida código, ADR, seguridad, configuración, APIs, frontend, CI, fixtures, evidencia E2E, QA, baseline y límites de alcance.", "Informe emitido el 25 de agosto de 2026."),
    (63, "Final exclusive state", "READY_FOR_PORTFOLIO_SCENARIO_PLANNING", "Gate 07F PASS. STOP: no iniciar Budget/Resource Scenario, selección, Investment Decision/FID, FEL/PDRI, Initialization, Activation ni Execution."),
]


SECTIONS = GATE07G_SECTIONS

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.right_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1
for name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK_BLUE, 8, 4),
):
    style = styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

language = OxmlElement("w:lang")
language.set(qn("w:val"), "es-CO")
styles["Normal"]._element.get_or_add_rPr().append(language)
doc.core_properties.title = "Informe técnico Gate 07G — Portfolio Budget Planning & Scenarios"
doc.core_properties.subject = "Implementación y validación técnica de Gate 07G en P&Pmis Construction AI"
doc.core_properties.author = "P&Pmis Construction AI / Codex"
doc.core_properties.keywords = "Gate 07G, Portfolio Budget Planning, Budget Target, Scenarios, PostgreSQL"

header = section.header.paragraphs[0]
header.paragraph_format.space_after = Pt(0)
set_run(header.add_run("P&Pmis Construction AI  |  Gate 07G Technical Report"), size=8, color=MUTED, bold=True)
add_page_number(section.footer.paragraphs[0])
bullet_num = add_numbering(doc)

logo_paragraph = doc.add_paragraph()
logo_paragraph.paragraph_format.space_after = Pt(8)
if LOGO.exists():
    picture = logo_paragraph.add_run().add_picture(str(LOGO), width=Inches(0.62))
    picture._inline.docPr.set("descr", "Logotipo de P&Pmis Construction AI")

title = doc.add_paragraph()
title.paragraph_format.space_after = Pt(4)
set_run(title.add_run("INFORME TÉCNICO"), size=23, color=BLACK, bold=True)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(16)
set_run(subtitle.add_run("Gate 07G — Portfolio Budget Planning & Budget Scenarios"), size=14, color=MUTED, bold=True)
metadata = [
    ("Aplicación", "P&Pmis Construction AI"),
    ("Ámbito", "USER MODE + ADMIN MODE · Enterprise Strategy Manager / Portfolio Manager"),
    ("Baseline", "Gate 07D/07E/07F validated · Alembic 20260825_0046"),
    ("Implementación", "Budget Targets + alternative Scenarios + comparison + governed configuration"),
    ("Fecha de corte", "25 de agosto de 2026"),
    ("Estado final", "READY_FOR_PORTFOLIO_BUDGET_RECOMMENDATION"),
]
for label, value in metadata:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.0
    set_run(paragraph.add_run(f"{label}: "), size=10.5, bold=True)
    set_run(paragraph.add_run(value), size=10.5)
rule = doc.add_paragraph()
rule.paragraph_format.space_before = Pt(8)
rule.paragraph_format.space_after = Pt(12)
add_bottom_rule(rule)

add_callout(
    doc,
    "Resultado ejecutivo",
    "Gate 07G fue implementado y validado de extremo a extremo. Consume Gate 07F como autoridad analítica, gobierna Target Budgets y escenarios alternativos con snapshots, hashes, ETag, idempotencia, time phasing explícito y comparación, y entrega readiness para una futura recomendación sin seleccionar ganador, financiar, aprobar FID ni iniciar ejecución.",
    fill=POSITIVE,
)
add_table(
    doc,
    ["Puerta", "Resultado", "Evidencia"],
    [
        ["Backend", "PASS", "5 pruebas Gate 07G; Ruff y contratos API validados."],
        ["PostgreSQL", "PASS", "69 regresiones + 23 Gate 07E/07F/07G; up/down/up real."],
        ["Frontend", "PASS", "2 pruebas 07G + 30 AppFlow; TypeScript y build Vite."],
        ["Browser USER/ADMIN", "PASS", "Target, Scenarios, forced, freeze, compare, ETag y publish."],
        ["Localhost", "PASS", "API/DB/Redis healthy; frontend 5173; Alembic 0046 head."],
    ],
    [1900, 1300, 6160],
)

add_heading(doc, "Arquitectura de la solución", page_break=True)
add_body(doc, "Gate 07G se integra como write model gobernado dentro de Portfolio Manager. Reutiliza la verdad analítica de Gate 07F, persiste únicamente Targets, Scenarios y sus snapshots, y conserva configuración/inheritance en AdminConfiguration.")
add_table(
    doc,
    ["Capa", "Responsabilidad", "Artefactos principales"],
    [
        ["Persistence", "Target, periods, Scenario y Project snapshots", "portfolio_budget/models.py + migration 0046"],
        ["Domain service", "Lifecycle, selection, hashes, phasing, compare, readiness", "portfolio_budget/service.py"],
        ["API", "Contratos USER/ADMIN, ETag, idempotencia y errores", "schemas.py + router.py"],
        ["Configuration", "Starter DRAFT, inheritance, preview/save/publish", "AdminConfiguration"],
        ["Frontend", "Workspace USER + Configuration ADMIN", "frontend/src/features/portfolio-budget/*"],
        ["Quality", "Unit, PostgreSQL, Vitest, Playwright, fixture y CI", "tests + compose + workflow"],
    ],
    [1800, 3300, 4260],
)
for statement in (
    "Fuente de verdad: Gate 07F analysis_hash + Gate 07E rank/scores + active membership + ROM Cost/currency + Project dates.",
    "Configuration path: Tenant → Enterprise → Business Unit → Portfolio; nearest published override wins.",
    "Output: READY_FOR_PORTFOLIO_BUDGET_RECOMMENDATION con referencias/hashes de Scenarios FROZEN y READY.",
):
    add_list_item(doc, statement, bullet_num)

add_figure(
    doc,
    USER_EVIDENCE,
    "Figura 1. USER MODE — comparación de dos Budget Scenarios congelados.",
    "Pantalla USER de Portfolio Budget Planning con Target, escenarios y comparación analítica.",
    width=4.15,
)
add_figure(
    doc,
    ADMIN_EVIDENCE,
    "Figura 2. ADMIN MODE — política efectiva después de clone, Preview, Save y Publish.",
    "Pantalla ADMIN de Portfolio Budget Planning Configuration con revisión publicada y Preview no persistente.",
    width=3.65,
)

# Let Word flow sections naturally to avoid half-empty pages around evidence.
page_break_sections: set[int] = set()
for number, title_text, summary, evidence in SECTIONS:
    add_heading(doc, f"{number}. {title_text}", page_break=number in page_break_sections)
    add_body(doc, summary)
    if evidence:
        add_list_item(doc, evidence, bullet_num)
    if False and number == 12:
        add_table(
            doc,
            ["Condición", "Tratamiento", "Efecto downstream"],
            [
                ["Moneda única", "ROM comparable y agregable", "Budget readiness evaluable"],
                ["Monedas múltiples sin FX", "Warning + total ROM deshabilitado", "Budget false; Resources independiente"],
                ["ROM null", "Valor original null + tamaño default", "Cobertura reportada, sin fabricación"],
            ],
            [2500, 3600, 3260],
        )
    if False and number == 27:
        add_table(
            doc,
            ["Hash", "Contenido", "Uso"],
            [
                ["membership_set_hash", "Memberships ACTIVE por Portfolio", "Detectar inclusión/exclusión"],
                ["evaluation_set_hash", "Últimas COMPLETED vigentes", "Detectar reevaluación"],
                ["prioritization_hash", "Readiness/ranking Gate 07E", "Preservar fuente de rank"],
                ["configuration_hash", "Config efectiva publicada", "Detectar cambio analítico"],
                ["analysis_hash", "Composición de los anteriores", "Handoff y stale-source 412"],
            ],
            [2200, 3800, 3360],
        )
    if False and number == 29:
        add_figure(
            doc,
            ADMIN_EVIDENCE,
            "Figura 1. ADMIN MODE — Portfolio Analysis Configuration publicada después del ciclo controlado.",
            "Pantalla ADMIN de Portfolio Analysis Configuration con tabs, revisiones y configuración publicada.",
            width=3.65,
        )
    if False and number == 30:
        add_figure(
            doc,
            USER_EVIDENCE,
            "Figura 2. USER MODE — Portfolio Scorecard con tres Projects y salida de análisis lista.",
            "Pantalla USER de Portfolio Analysis con Scorecard accesible para tres proyectos.",
            width=4.15,
        )
    if False and number == 37:
        add_table(
            doc,
            ["Surface", "Método", "Path / control"],
            [
                ["USER", "GET", "/portfolios/{id}/analysis, dataset, investor-map, scorecard, readiness"],
                ["USER", "POST", "/portfolios/{id}/analysis/preview"],
                ["ADMIN", "GET", "/portfolio-analysis/admin/configurations"],
                ["ADMIN", "POST", "configurations/preview, clone, {id}/publish"],
                ["ADMIN", "PUT", "configurations/{id} con If-Match"],
            ],
            [1600, 1100, 6660],
        )
    if False and number == 49:
        add_table(
            doc,
            ["Suite", "Resultado", "Cobertura clave"],
            [
                ["Gate 07F unit", "5 passed", "Dataset, guardrails, config, hashes, exclusiones"],
                ["Gate 07F PostgreSQL", "3 passed", "Primary, multi-Portfolio, stale/currency"],
                ["Gate 07E/07F compose", "17 passed", "Integración real PostgreSQL"],
                ["Inherited PG regression", "69 passed", "Gates/contexto/migración"],
            ],
            [2600, 1600, 5160],
        )
    if False and number == 58:
        add_table(
            doc,
            ["Control", "Resultado", "Observación"],
            [
                ["Backend full", "PASS", "463 passed, 16 skipped; cobertura 82,81%."],
                ["Windows temp rerun", "PASS", "48/48 de los 28 setups afectados por permiso global."],
                ["Frontend", "PASS", "199/199; build y typecheck."],
                ["Ruff/Prettier", "PASS", "205 Python files formatted; Prettier completo."],
                ["ESLint", "PASS", "0 errores; 8 warnings heredados del visor IFC."],
                ["Diff check", "PASS", "Sin whitespace errors."],
            ],
            [2500, 1400, 5460],
        )
    if False and number == 63:
        add_callout(
            doc,
            "Estado exclusivo final",
            "READY_FOR_PORTFOLIO_SCENARIO_PLANNING — PASS. Gate 07F se cierra sin implementar ninguna capacidad posterior.",
            fill=POSITIVE,
        )

doc.add_page_break()
add_heading(doc, "Matriz de aceptación Gate 07G")
acceptance_rows = [
    ("Consumes Gate 07F + budget readiness", "PASS", "READY + can_enter_budget_planning"),
    ("CAPITAL_OWNER only", "PASS", "Exclusiones conservadas"),
    ("Governed Target revisions", "PASS", "DRAFT/PUBLISHED/SUPERSEDED"),
    ("Period Targets", "PASS", "YEAR/QUARTER/MONTH + sum control"),
    ("Single Budget Currency", "PASS", "Sin FX implícito"),
    ("Gate 07F ROM source", "PASS", "Valor, referencia y hash"),
    ("Three Scenario creation modes", "PASS", "EMPTY/rank/copy"),
    ("Four selection states", "PASS", "SELECTED/UNSELECTED/FORCED_*"),
    ("Forced reason and audit", "PASS", "Actor/reason/authority/SecurityEvent"),
    ("Explicit time phasing", "PASS", "Unphased/linear/custom"),
    ("No silent Project Cash Flow", "PASS", "Approximation identified"),
    ("Scenario comparison and delta", "PASS", "2–5 alternatives"),
    ("Source hash and STALE", "PASS", "Preview/apply refresh"),
    ("Frozen immutability", "PASS", "Rows and header protected"),
    ("Readiness/output", "PASS", "READY_FOR_PORTFOLIO_BUDGET_RECOMMENDATION"),
    ("ADMIN inheritance/Preview/ETag", "PASS", "Clone/save/publish"),
    ("RBAC/tenant/idempotency", "PASS", "API + PostgreSQL"),
    ("No winner/Funds/Resource/FID", "PASS", "Hard boundaries"),
    ("Migration reversible", "PASS", "0045→0046→0045→0046"),
    ("Frontend/browser/local Docker", "PASS", "USER/ADMIN localhost"),
]
add_table(doc, ["Criterio", "Resultado", "Evidencia"], acceptance_rows, [4400, 1300, 3660], font_size=8.8)

doc.add_page_break()
add_heading(doc, "Inventario de cambios")
add_table(
    doc,
    ["Área", "Archivos"],
    [
        ["Backend Gate 07G", "backend/app/modules/portfolio_budget/{models,schemas,service,router}.py"],
        ["Migration", "backend/alembic/versions/20260825_0046_gate07g_portfolio_budget.py"],
        ["Backend integration", "api/v1/router.py; permissions/constants; workspace_context/navigator.py"],
        ["Frontend", "frontend/src/features/portfolio-budget/*; App.tsx; applicationNavigation.ts"],
        ["Tests", "test_portfolio_budget_gate07g.py; postgres test/schema; Vitest; Playwright"],
        ["Fixtures/CI", "gate07e_browser_fixture.py; docker-compose.gate07e.yml; .github/workflows/ci.yml"],
        ["Architecture", "docs/47-adr-gate07g-portfolio-budget-planning.md"],
    ],
    [2500, 6860],
)
add_callout(
    doc,
    "STOP boundary",
    "No iniciar Resource Planning/Scenarios, recomendación o Scenario winner, Funds/Funding, Investment Approval/FID, Project Definition/FEL/PDRI, Initialization, Activation ni Execution sin un prompt explícito posterior.",
    fill=CAUTION,
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
