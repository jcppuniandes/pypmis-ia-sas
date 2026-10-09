"""Preset, structure, accessibility and content audit for Gate 07M-EVID."""

from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

REPO = Path(__file__).resolve().parents[1]
DOCX = REPO / "outputs" / "technical_reports" / "Informe_Tecnico_Gate_07M_EVID_Policy_Update_20260903.docx"
OUTPUT_DIR = REPO / "outputs" / "technical_reports"


def close(actual: int | None, expected: int, tolerance: int = 1000) -> bool:
    return actual is not None and abs(actual - expected) <= tolerance


def write_json(name: str, payload: dict[str, object]) -> None:
    (OUTPUT_DIR / name).write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")


def run() -> None:
    doc = Document(DOCX)
    section = doc.sections[0]
    normal = doc.styles["Normal"]
    heading1 = doc.styles["Heading 1"]
    heading2 = doc.styles["Heading 2"]
    body = " ".join(paragraph.text for paragraph in doc.paragraphs)
    body += " " + " ".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    preset_checks = {
        "single_portrait_section": len(doc.sections) == 1 and section.page_height > section.page_width,
        "letter_page_size": close(section.page_width, Inches(8.5)) and close(section.page_height, Inches(11)),
        "one_inch_margins": all(close(value, Inches(1)) for value in (section.left_margin, section.right_margin, section.top_margin, section.bottom_margin)),
        "header_footer_distance": close(section.header_distance, Inches(0.492)) and close(section.footer_distance, Inches(0.492)),
        "normal_calibri_11pt": normal.font.name == "Calibri" and close(normal.font.size, Pt(11), 100),
        "normal_spacing": close(normal.paragraph_format.space_after, Pt(6), 100) and abs(float(normal.paragraph_format.line_spacing) - 1.10) < 0.001,
        "heading1_tokens": close(heading1.font.size, Pt(16), 100) and str(heading1.font.color.rgb) == "2E74B5",
        "heading2_tokens": close(heading2.font.size, Pt(13), 100) and str(heading2.font.color.rgb) == "2E74B5",
        "heading_hierarchy": sum(p.style.name == "Heading 1" for p in doc.paragraphs) == 9 and sum(p.style.name == "Heading 2" for p in doc.paragraphs) >= 9,
        "header_present": "TECHNICAL REPORT · GATE 07M-EVID" in section.header.paragraphs[0].text,
        "footer_present": "P&Pmis Construction AI" in section.footer.paragraphs[0].text,
    }
    table_geometry = row_headers = cell_margins = no_split = True
    for table in doc.tables:
        properties = table._tbl.tblPr
        table_geometry &= properties.find(qn("w:tblW")) is not None and properties.find(qn("w:tblInd")) is not None
        row_headers &= table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None
        for row in table.rows:
            no_split &= row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
            for cell in row.cells:
                cell_margins &= cell._tc.get_or_add_tcPr().find(qn("w:tcMar")) is not None
    preset_checks.update({
        "table_geometry_complete": table_geometry,
        "table_header_flags_complete": row_headers,
        "cell_margins_complete": cell_margins,
        "rows_prevent_split": no_split,
    })
    content_checks = {
        "final_output_present": "READY_FOR_BUSINESS_CASE_AND_ECONOMIC_FINANCIAL_GATE_EVIDENCE" in body,
        "exclusive_state_present": "GATE07M_EVID_POLICY_UPDATE_VALIDATED" in body,
        "next_present": "READY_TO_REVALIDATE_NEW_FID_PATH" in body,
        "stop_present": "STOP" in body,
        "both_providers_present": all(key in body for key in ("BUSINESS_CASE_READINESS_PROVIDER", "ECONOMIC_FINANCIAL_READINESS_PROVIDER")),
        "policy_identity_present": all(value in body for value in ("787", "a70075f0b392ab5d17a21353d7434ca73b51d11f57ccbd14774aa1013e795392")),
        "traceability_101_present": any(cell.text.strip() == "101" for table in doc.tables for row in table.rows for cell in row.cells),
        "no_fake_bullets": not any(p.text.strip().startswith(("•", "- ")) for p in doc.paragraphs),
        "no_placeholder_text": not re.search(r"\b(?:todo|lorem ipsum|placeholder|codex-file-citation)\b", body, flags=re.IGNORECASE),
    }
    a11y_checks = {
        "all_images_have_alt": all(shape._inline.docPr.get("descr") for shape in doc.inline_shapes),
        "heading_levels_present": all(any(p.style.name == level for p in doc.paragraphs) for level in ("Heading 1", "Heading 2")),
        "tables_have_header_rows": row_headers,
        "document_title_present": bool(doc.core_properties.title),
        "document_language_content": "Informe Técnico" in doc.core_properties.title,
    }
    preset = {
        "preset": "standard_business_brief",
        "first_page_pattern": "memo_masthead",
        "file": str(DOCX),
        "tables": len(doc.tables),
        "inline_images": len(doc.inline_shapes),
        "checks": preset_checks,
        "result": "PASS" if all(preset_checks.values()) else "FAIL",
        "failed": [key for key, value in preset_checks.items() if not value],
    }
    style = {
        "file": str(DOCX),
        "checks": content_checks,
        "result": "PASS" if all(content_checks.values()) else "FAIL",
        "failed": [key for key, value in content_checks.items() if not value],
    }
    a11y = {
        "file": str(DOCX),
        "checks": a11y_checks,
        "result": "PASS" if all(a11y_checks.values()) else "FAIL",
        "failed": [key for key, value in a11y_checks.items() if not value],
    }
    write_json("gate07m_evid_preset_audit.json", preset)
    write_json("gate07m_evid_style_lint.json", style)
    write_json("gate07m_evid_a11y.json", a11y)
    result = {"preset": preset["result"], "style": style["result"], "a11y": a11y["result"]}
    print(json.dumps(result, indent=2, ensure_ascii=False))
    if not all(value == "PASS" for value in result.values()):
        raise SystemExit(1)


if __name__ == "__main__":
    run()
