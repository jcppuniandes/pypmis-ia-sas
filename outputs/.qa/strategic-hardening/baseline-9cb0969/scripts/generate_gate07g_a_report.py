from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from generate_gate07i_report import (
    CALLOUT,
    CAUTION,
    DARK_BLUE,
    INK,
    POSITIVE,
    TEAL,
    add_bottom_rule,
    add_callout,
    add_page_number,
    add_table,
    configure_styles,
    set_image_alt,
    set_run,
)

REPO = Path(__file__).resolve().parents[1]
OUTPUT = (
    REPO
    / "outputs"
    / "technical_reports"
    / "Informe_Tecnico_Gate_07G_A_OPC_Budget_Alignment_20260830.docx"
)
LOGO = REPO / "frontend" / "public" / "pypmis-construction-ai-logo.png"
E2E_ROOT = (
    REPO
    / "frontend"
    / "test-results"
    / "portfolio-budget-Gate-07G--d4100-os-comparison-and-readiness-chromium"
)
USER_EVIDENCE = E2E_ROOT / "gate07g-user-1.png"
ADMIN_EVIDENCE = E2E_ROOT / "gate07g-admin-1.png"

# Resolved document token map:
# preset=standard_business_brief; first-page pattern=memo_masthead.


def add_title_page(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(header.add_run("TECHNICAL REPORT · GATE 07G-A"), size=8.5, color="5F6B76", bold=True)
    add_bottom_rule(header, color="D9E1E8", size=6)
    add_page_number(section.footer.paragraphs[0])

    if LOGO.exists():
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(8)
        picture = paragraph.add_run().add_picture(str(LOGO), width=Inches(0.72))
        set_image_alt(picture, "Logotipo de P&Pmis Construction AI")

    kicker = doc.add_paragraph()
    set_run(kicker.add_run("INFORME TÉCNICO DE IMPLEMENTACIÓN"), size=10, color=TEAL, bold=True)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    set_run(title.add_run("Gate 07G-A"), size=25, color=INK, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    set_run(
        subtitle.add_run("OPC Budget Alignment Enhancements"),
        size=15,
        color=DARK_BLUE,
        bold=True,
    )
    for label, value in (
        ("Aplicación", "P&Pmis Construction AI"),
        ("Fecha de corte", "30 de agosto de 2026"),
        ("Repositorio", "pypmis-ia-sas · branch main"),
        ("Migración", "20260830_0052 (head)"),
        ("Estado", "GATE07G_A_OPC_ALIGNMENT_VALIDATED"),
    ):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        set_run(paragraph.add_run(f"{label}: "), size=10.5, color=INK, bold=True)
        set_run(paragraph.add_run(value), size=10.5)
    rule = doc.add_paragraph()
    add_bottom_rule(rule)
    add_callout(
        doc,
        "Resultado ejecutivo",
        "Gate 07G-A quedó implementado como una extensión aditiva de Portfolio Budget Planning. La solución "
        "deriva Waterline sin crear un ranking paralelo, incorpora Project Dependencies transversales y agrega "
        "desglose CAPEX/EXPENSE/UNCLASSIFIED sin inventar clasificación. El recorrido PostgreSQL, USER y ADMIN "
        "se validó en localhost y finalizó sin datos temporales residuales.",
        fill=POSITIVE,
    )
    add_table(
        doc,
        ["Capa", "Resultado", "Evidencia principal"],
        [
            ("Backend / API", "PASS", "Waterline, dependencies, breakdown, ETag, RBAC, idempotencia y eventos"),
            ("PostgreSQL", "PASS", "0051 → 0052 → 0051 → 0052; suites 69/69 y 60/60"),
            ("Frontend", "PASS", "ESLint, TypeScript, Vitest 2/2 y Vite production build"),
            ("Browser E2E", "PASS", "07E → 07F → 07G-A; USER + ADMIN; create/edit/deactivate"),
            ("Localhost", "ACTIVE", "Frontend 5173; API 8000; PostgreSQL head 0052"),
        ],
        [1900, 1300, 6160],
        font_size=8.9,
    )
    doc.add_page_break()


def add_scope(doc: Document) -> None:
    doc.add_heading("1. Alcance, baseline y decisiones", level=1)
    add_callout(
        doc,
        "Alcance aplicado",
        "Se reutilizó Gate 07G existente y se añadieron exclusivamente capacidades de alineación presupuestal OPC. "
        "No se creó una nueva Investment Gate, un optimizador, un ganador de Scenario, un Funds allocation, una "
        "aprobación FID ni activación de Project.",
    )
    add_table(
        doc,
        ["Decisión", "Implementación", "Consecuencia"],
        [
            ("Waterline derivada", "Gate 07E rank + Gate 07G selection/demand/target", "No persiste rank ni cumulative demand"),
            ("Dependency transversal", "ProjectDependency Project→Project, tenant-scoped", "Reutilizable entre Portfolios y gates"),
            ("Budget class explícita", "CAPEX, EXPENSE y UNCLASSIFIED", "Sin split porcentual fabricado"),
            ("Historical compatibility", "Legacy rows migradas como LEGACY_UNCLASSIFIED", "Scenarios previos conservan hashes y lectura"),
            ("Freeze policy", "BLOCK_FREEZE para mandatory unmet", "Optional dependencies sólo generan warning"),
        ],
        [2200, 3600, 3560],
        font_size=8.7,
    )
    doc.add_heading("1.1 ADR y continuidad arquitectónica", level=2)
    add_table(
        doc,
        ["Artefacto", "Contenido"],
        [
            ("docs/54-adr-gate07g-a-opc-budget-alignment.md", "Ownership, derivación, persistencia mínima, compatibilidad y límites"),
            ("AdminConfiguration", "Política heredable y Preview no persistente"),
            ("SecurityEvent", "Auditoría de altas, cambios y desactivaciones; sin ruido de lectura"),
            ("Workspace Context / Navigation", "Reutilización de Portfolio Manager USER y configuración ADMIN"),
        ],
        [3300, 6060],
        font_size=8.9,
    )


def add_domain(doc: Document) -> None:
    doc.add_heading("2. Dominio, persistencia y migración", level=1)
    add_table(
        doc,
        ["Componente", "Cambio aditivo", "Control"],
        [
            ("PortfolioBudgetTarget", "capex_budget_target, expense_budget_target", "Nullable; no obliga clasificación inexistente"),
            ("PortfolioBudgetTargetPeriod", "capex_amount, expense_amount", "Desglose opcional por periodo"),
            ("PortfolioBudgetScenarioProject", "capex, expense, unclassified, source", "Suma exacta = cost_value"),
            ("ProjectDependency", "Pred→Succ, type, strength, dates, source, evidence", "Tenant, ETag, idempotencia y status"),
        ],
        [3100, 3100, 3160],
        font_size=8.4,
    )
    doc.add_heading("2.1 Integridad de ProjectDependency", level=2)
    add_table(
        doc,
        ["Regla", "Comportamiento"],
        [
            ("Dirección", "Predecessor → Successor; nunca se infiere la relación inversa."),
            ("Strength", "MANDATORY puede bloquear Freeze; OPTIONAL sólo advierte."),
            ("Status", "ACTIVE participa en el grafo; INACTIVE conserva historia."),
            ("Safeguards", "Rechazo de self-reference, cross-tenant, duplicado activo y ciclo mandatory."),
            ("Concurrencia", "PUT y deactivate requieren If-Match; stale revision responde 412."),
            ("Idempotencia", "Create y deactivate usan claves tenant-scoped."),
        ],
        [2600, 6760],
        font_size=8.9,
    )
    doc.add_heading("2.2 Alembic 0052", level=2)
    add_table(
        doc,
        ["Paso", "Resultado", "Observación"],
        [
            ("Upgrade", "PASS", "20260827_0051 → 20260830_0052"),
            ("Schema validation", "PASS", "Tabla, FKs, checks, índices y columnas verificadas en PostgreSQL 16"),
            ("Downgrade", "PASS", "0052 → 0051 sin afectar objetos previos"),
            ("Re-upgrade", "PASS", "0051 → 0052 reproducible"),
            ("Legacy data", "PASS", "unclassified=cost_value; source=LEGACY_UNCLASSIFIED"),
        ],
        [1800, 1300, 6260],
        font_size=8.8,
    )


def add_waterline(doc: Document) -> None:
    doc.add_heading("3. Waterline derivada", level=1)
    add_callout(
        doc,
        "Principio de cálculo",
        "Waterline conserva el portfolio_rank contextual de Gate 07E y acumula exclusivamente la demanda seleccionada "
        "del Scenario frente al Target total o al periodo solicitado. El resultado se calcula en memoria y nunca "
        "reordena Projects ni persiste una frontera paralela.",
        fill=CALLOUT,
    )
    add_table(
        doc,
        ["Entrada", "Fuente", "Salida derivada"],
        [
            ("Rank", "Gate 07E contextual portfolio rank", "Orden estable y read-only"),
            ("Selection", "Scenario SELECTED/FORCED_IN/FORCED_OUT", "Included, excluded y forced state"),
            ("Demand", "Cost snapshot + explicit time phasing", "Total o periodo; sin mutar Scenario"),
            ("Target", "Published Budget Target / period", "Variance y waterline position"),
            ("Dependencies", "Active ProjectDependency graph", "MET, UNMET, OPTIONAL_WARNING, NOT_APPLICABLE"),
        ],
        [2200, 3600, 3560],
        font_size=8.6,
    )
    doc.add_heading("3.1 Filtros y accesibilidad", level=2)
    add_table(
        doc,
        ["Filtro", "Semántica"],
        [
            ("Period", "TOTAL o period_key; recalcula demanda y Target correspondiente."),
            ("Project Type", "Filtro visual; no cambia selección ni ranking."),
            ("Strategic Objective", "Clasificación N:M proveniente de Project."),
            ("Risk Bucket", "Dato de análisis existente; no se recalcula."),
            ("Selection State", "Visibilidad de Selected/Forced/Unselected."),
            ("Budget Class", "CAPEX, EXPENSE o UNCLASSIFIED según datos gobernados."),
        ],
        [2700, 6660],
        font_size=8.9,
    )
    add_callout(
        doc,
        "Representación accesible",
        "La Waterline dispone de tabla con nombre accesible, columnas de demanda acumulada, frontera, breakdown y "
        "dependency status, además de filtros etiquetados y detalle inspeccionable. El color no es el único canal.",
        fill=POSITIVE,
    )
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    set_run(paragraph.add_run("Identificador técnico: "), size=9.5, color=INK, bold=True)
    set_run(paragraph.add_run("GATE07E_CONTEXTUAL_PORTFOLIO_RANK"), size=9.5)


def add_dependencies(doc: Document) -> None:
    doc.add_heading("4. Evaluación de dependencias", level=1)
    add_table(
        doc,
        ["Caso", "Resultado derivado", "Efecto"],
        [
            ("Successor incluido; predecessor incluido", "MET", "Sin blocker"),
            ("Mandatory predecessor excluido", "UNMET", "BLOCK_FREEZE por política default"),
            ("Optional predecessor excluido", "OPTIONAL_WARNING", "Warning; Freeze permitido"),
            ("Successor excluido", "NOT_APPLICABLE", "Relación no evalúa selección"),
            ("Forced successor + missing predecessor", "FORCED conflict", "Warning explícito, sin autocorrección"),
            ("Cadena A→B→C", "Transitive unmet", "Detalle de predecessor faltante aguas arriba"),
        ],
        [2850, 2700, 3810],
        font_size=8.6,
    )
    doc.add_heading("4.1 Registro USER y política ADMIN", level=2)
    add_table(
        doc,
        ["USER", "ADMIN"],
        [
            ("Crear relación con type, strength, fechas, fuente y descripción", "Enable/disable dependency evaluation"),
            ("Editar strength/description bajo ETag", "Catálogo de tipos y strengths permitidos"),
            ("Desactivar sin borrar historial", "Mandatory unmet policy: BLOCK_FREEZE"),
            ("Ver estado por Scenario sin autocorrección", "Graph limit y Preview no persistente"),
        ],
        [4680, 4680],
        font_size=8.8,
    )


def add_breakdown(doc: Document) -> None:
    doc.add_heading("5. CAPEX / EXPENSE / UNCLASSIFIED", level=1)
    add_callout(
        doc,
        "No fabricated split",
        "El desglose sólo se informa cuando Project contiene planning_budget_breakdown gobernado. Si la fuente no "
        "clasifica el costo, capex y expense permanecen null y el valor completo se registra como UNCLASSIFIED.",
        fill=CAUTION,
    )
    add_table(
        doc,
        ["Nivel", "Campos", "Regla"],
        [
            ("Target", "total, capex target, expense target, unallocated", "Opcional; total continúa siendo autoridad"),
            ("Target period", "amount, capex, expense", "Permite control temporal explícito"),
            ("Scenario Project", "capex, expense, unclassified, source", "COALESCE parts + unclassified = cost"),
            ("Scenario summary", "selected y variance por clase", "Agregación sólo de filas seleccionadas"),
            ("Comparison", "class deltas + unmet dependency delta", "No elige alternativa"),
        ],
        [2000, 3500, 3860],
        font_size=8.7,
    )


def add_ui(doc: Document) -> None:
    doc.add_heading("6. Experiencia USER y ADMIN", level=1)
    add_table(
        doc,
        ["Superficie", "Capacidad añadida", "Protección"],
        [
            ("USER · Waterline", "Period/filter/detail + accessible table", "No nueva priorización"),
            ("USER · Dependencies", "Create/edit/deactivate register", "No auto-selection"),
            ("USER · Scenario", "CAPEX/EXPENSE/Unclassified y dependency status", "Frozen history intacta"),
            ("USER · Comparison", "Breakdown y dependency deltas", "No winner"),
            ("ADMIN", "Waterline, dependency, breakdown y warning policies", "Preview persisted=false"),
        ],
        [2200, 3800, 3360],
        font_size=8.5,
    )
    doc.add_heading("6.1 Evidencia Browser E2E", level=2)
    add_callout(
        doc,
        "Captura reproducible",
        "Las figuras siguientes provienen del recorrido real 07E → 07F → 07G-A contra API, PostgreSQL y frontend "
        "locales. El fixture fue eliminado al terminar, incluyendo la nueva ProjectDependency.",
        fill=CALLOUT,
    )
    add_table(
        doc,
        ["Figura", "Validación"],
        [
            ("USER", "Target, dos Scenarios FROZEN, comparación y tabs Waterline/Dependencies."),
            ("ADMIN", "Clone, Preview persisted=false, Save, stale ETag y Publish."),
        ],
        [1900, 7460],
        font_size=8.7,
    )
    add_figure(
        doc,
        USER_EVIDENCE,
        "Figura 1. Evidencia USER: escenarios y comparación Gate 07G-A en localhost.",
        "Pantalla USER de Portfolio Budget Planning con tabs Waterline, Dependencies y Comparison.",
        width=Inches(4.5),
    )
    add_figure(
        doc,
        ADMIN_EVIDENCE,
        "Figura 2. Evidencia ADMIN: política publicada y Preview no persistente.",
        "Pantalla ADMIN de Portfolio Budget Planning Configuration con políticas Waterline, Dependencies y CAPEX EXPENSE.",
        width=Inches(3.15),
    )


def add_figure(doc: Document, path: Path, caption: str, alt_text: str, *, width: Inches) -> None:
    if not path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    picture = paragraph.add_run().add_picture(str(path), width=width)
    set_image_alt(picture, alt_text)
    label = doc.add_paragraph(caption)
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.space_after = Pt(6)
    for run in label.runs:
        set_run(run, size=8.7, color="5F6B76", italic=True)


def add_security(doc: Document) -> None:
    doc.add_heading("7. API, seguridad y configuración", level=1)
    add_table(
        doc,
        ["Contrato", "Control"],
        [
            ("GET Scenario Waterline", "Scenario scope + optional period_key; derived response"),
            ("GET Scenario Dependencies", "Derived summary, per-Project status, blockers y warnings"),
            ("GET/POST Portfolio Dependencies", "RBAC read/manage + tenant/project validation"),
            ("GET/PUT/Deactivate Dependency", "ETag, idempotency, history y SecurityEvent"),
            ("ADMIN Preview", "Policy override evaluado con persisted=false"),
        ],
        [3900, 5460],
        font_size=8.9,
    )
    doc.add_heading("7.1 Permisos y acciones", level=2)
    add_table(
        doc,
        ["Permiso / acción", "Asignación y uso"],
        [
            ("portfolio_dependency.read", "Viewer/planner/manager: registro y estado derivado"),
            ("portfolio_dependency.manage", "Planner/manager/admin: create, edit y deactivate"),
            ("can_view_waterline", "Lectura Waterline según autoridad de Portfolio Budget"),
            ("can_view_capex_expense_breakdown", "Visualización de datos disponibles; no edición de fuente"),
            ("can_freeze_budget_scenario", "Se retira cuando existe mandatory unmet bajo BLOCK_FREEZE"),
        ],
        [3800, 5560],
        font_size=8.8,
    )
    add_callout(
        doc,
        "Eventos de seguridad",
        "project_dependency.created, project_dependency.updated y project_dependency.deactivated registran actor, "
        "tenant, target y detalle. Waterline, filtros, hover y lecturas ordinarias no producen eventos ruidosos.",
    )


def add_validation(doc: Document) -> None:
    doc.add_heading("8. Validación técnica y evidencia", level=1)
    add_table(
        doc,
        ["Control", "Resultado", "Evidencia"],
        [
            ("Ruff + py_compile", "PASS", "portfolio_budget, migration, tests y fixture"),
            ("Backend focalizado", "8/8 funcional", "Waterline, breakdown, cycles, block, forced y transitive"),
            ("Coverage aislada", "44,43% global", "8 tests PASS; exit no-cero sólo por umbral monolítico 75%"),
            ("PostgreSQL migration suite", "69/69", "Schema y regresión sobre base efímera"),
            ("PostgreSQL focused suite", "60/60", "Upgrade/downgrade/re-upgrade y acceptance"),
            ("ESLint", "PASS", "Feature portfolio-budget sin errores"),
            ("TypeScript", "PASS", "tsc --noEmit"),
            ("Vitest", "2/2", "Workspace USER + Configuration ADMIN"),
            ("Vite", "PASS", "Production build"),
            ("Playwright prerequisites", "2/2", "07E FLOW_PASS y 07F readiness"),
            ("Playwright Gate 07G-A", "1/1", "USER/ADMIN + dependency create/edit/deactivate"),
            ("Fixture teardown", "PASS", "13 workspaces, 9 memberships y 1 dependency eliminados"),
        ],
        [2800, 1600, 4960],
        font_size=8.2,
    )
    add_callout(
        doc,
        "Interpretación de cobertura",
        "Los ocho casos focalizados pasaron. El comando local que carga toda la aplicación aplica un umbral de "
        "cobertura global de 75% y reportó 44,43% al ejecutar sólo dos archivos; esto no corresponde a un fallo "
        "funcional. La batería PostgreSQL completa sí finalizó con código 0.",
        fill=CAUTION,
    )


def add_acceptance(doc: Document) -> None:
    doc.add_heading("9. Matriz de aceptación Gate 07G-A", level=1)
    rows = [
        ("Waterline reuses Gate 07E rank", "PASS", "rank_source explícito; sin rank persistido"),
        ("Total and period Waterline", "PASS", "TOTAL/period_key + phasing"),
        ("Selection and forced state visible", "PASS", "SELECTED/FORCED_IN/FORCED_OUT"),
        ("Project filters", "PASS", "Type/objective/risk/selection/budget class"),
        ("Accessible representation", "PASS", "Named table, labels and text statuses"),
        ("ProjectDependency transverse model", "PASS", "Tenant-scoped Project→Project"),
        ("Mandatory/Optional", "PASS", "Blocker vs warning"),
        ("Active/Inactive", "PASS", "Deactivation preserves history"),
        ("Self/cross-tenant/duplicate safeguards", "PASS", "Service + DB constraints"),
        ("Cycle safeguard", "PASS", "Mandatory directed graph"),
        ("Transitive chains", "PASS", "Derived transitive unmet detail"),
        ("Forced conflicts", "PASS", "Explicit warning; no autocorrection"),
        ("CAPEX/EXPENSE", "PASS", "Only governed source data"),
        ("UNCLASSIFIED", "PASS", "Default for unclassified cost"),
        ("Target and period breakdown", "PASS", "Nullable additive fields"),
        ("Scenario and comparison enrichment", "PASS", "Summaries, period and deltas"),
        ("ADMIN sections", "PASS", "Structured policy + Preview"),
        ("Preview non-persistent", "PASS", "persisted=false"),
        ("RBAC/tenant isolation", "PASS", "Permissions and same-tenant assertions"),
        ("ETag/idempotency", "PASS", "Mutation contracts"),
        ("SecurityEvent", "PASS", "Governed changes only"),
        ("Historical compatibility", "PASS", "LEGACY_UNCLASSIFIED + legacy hash path"),
        ("Migration reversible", "PASS", "0051 to 0052 to 0051"),
        ("USER browser", "PASS", "Waterline/dependencies/scenarios/comparison"),
        ("ADMIN browser", "PASS", "Clone/preview/save/stale ETag/publish"),
        ("No optimizer/winner/FID", "PASS", "Explicit downstream boundaries"),
    ]
    add_table(doc, ["Criterio", "Estado", "Evidencia"], rows, [4100, 1200, 4060], font_size=7.9)


def add_inventory(doc: Document) -> None:
    doc.add_heading("10. Inventario de cambios y cierre", level=1)
    add_table(
        doc,
        ["Área", "Archivos principales"],
        [
            ("ADR", "docs/54-adr-gate07g-a-opc-budget-alignment.md"),
            ("Backend", "backend/app/modules/portfolio_budget/{models,schemas,service,router}.py"),
            ("Migration", "backend/alembic/versions/20260830_0052_gate07g_a_opc_alignment.py"),
            ("Security", "enterprise_structure/{constants,permissions}.py"),
            ("Frontend", "frontend/src/features/portfolio-budget/*"),
            ("Testing", "test_portfolio_budget_gate07g_a.py; postgres validator/tests; Playwright"),
            ("CI/fixture", ".github/workflows/ci.yml; docker-compose.gate07e.yml; gate07e_browser_fixture.py"),
        ],
        [2200, 7160],
        font_size=8.7,
    )
    doc.add_heading("10.1 Riesgos y recomendaciones", level=2)
    add_table(
        doc,
        ["Tema", "Tratamiento / recomendación"],
        [
            ("Datos de clasificación", "Mantener UNCLASSIFIED hasta recibir fuente financiera gobernada."),
            ("Graph size", "Usar max_graph_depth de configuración y monitorear Portfolios de alto volumen."),
            ("Period targets", "No interpretar cero como asignación aprobada cuando la fuente es desconocida."),
            ("UX comparison", "Evolucionar JSON técnico a tabla dedicada sin cambiar contratos ni autoridad."),
            ("CI", "Conservar secuencia 07E→07F→07G-A y teardown con ProjectDependency."),
        ],
        [2600, 6760],
        font_size=8.8,
    )
    doc.add_page_break()
    doc.add_heading("11. Conclusión", level=1)
    add_table(
        doc,
        ["Cierre", "Resultado"],
        [
            ("Funcional", "Waterline, Project Dependencies y budget breakdown disponibles."),
            ("Técnico", "API, migración 0052, frontend y CI integrados."),
            ("Operacional", "Localhost activo en 5173/8000; fixture sin residuos."),
            ("Release", "GATE07G_A_OPC_ALIGNMENT_VALIDATED"),
        ],
        [2500, 6860],
        font_size=9.0,
    )
    add_callout(
        doc,
        "Estado exclusivo final",
        "GATE07G_A_OPC_ALIGNMENT_VALIDATED. El módulo queda preparado para análisis presupuestal gobernado. "
        "STOP: no implementar optimizer, Scenario winner, Investment Decision/FID, financiación, Project "
        "Initialization, Activation o Execution sin una instrucción posterior explícita.",
        fill=POSITIVE,
    )


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)
    add_title_page(doc)
    add_scope(doc)
    add_domain(doc)
    add_waterline(doc)
    add_dependencies(doc)
    add_breakdown(doc)
    add_ui(doc)
    add_security(doc)
    add_validation(doc)
    add_acceptance(doc)
    add_inventory(doc)

    properties = doc.core_properties
    properties.title = "Informe Técnico Gate 07G-A — OPC Budget Alignment Enhancements"
    properties.subject = "Implementación y validación técnica de Waterline, Dependencies y CAPEX/EXPENSE"
    properties.author = "P&Pmis Construction AI"
    properties.keywords = "Gate 07G-A, Waterline, Project Dependencies, CAPEX, EXPENSE, Portfolio Budget"
    properties.comments = "Generated from verified repository evidence on 2026-08-30."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
