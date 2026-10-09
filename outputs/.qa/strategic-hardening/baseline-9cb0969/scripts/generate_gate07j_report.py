from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

from gate07j_report_data import SECTIONS
from generate_gate07i_report import (
    ALT_LIGHT,
    BLUE,
    CALLOUT,
    DARK_BLUE,
    INK,
    MUTED,
    POSITIVE,
    RISK,
    TEAL,
    add_bottom_rule,
    add_callout,
    add_page_number,
    add_table,
    configure_styles,
    set_image_alt,
    set_run,
)

REPO = Path(__file__).resolve().parents[1]
OUTPUT = REPO / "outputs" / "technical_reports" / "Informe_Tecnico_Gate_07J_Investment_Gate_Preparation_20260826.docx"
LOGO = REPO / "frontend" / "public" / "pypmis-construction-ai-logo.png"


def add_title_page(doc: Document) -> None:
    doc.settings.odd_and_even_pages_header_footer = True
    section = doc.sections[0]
    for header in (section.header.paragraphs[0], section.even_page_header.paragraphs[0]):
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_run(header.add_run("TECHNICAL REPORT · GATE 07J"), size=8.5, color=MUTED, bold=True)
        add_bottom_rule(header, color="D9E1E8", size=6)
    for footer in (section.footer.paragraphs[0], section.even_page_footer.paragraphs[0]):
        add_page_number(footer)

    if LOGO.exists():
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = 8
        picture = paragraph.add_run().add_picture(str(LOGO), width=Inches(0.72))
        set_image_alt(picture, "Logotipo de P&Pmis Construction AI")
    kicker = doc.add_paragraph()
    set_run(kicker.add_run("INFORME TÉCNICO DE IMPLEMENTACIÓN"), size=10, color=TEAL, bold=True)
    title = doc.add_paragraph()
    set_run(title.add_run("Gate 07J"), size=25, color=INK, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = 16
    set_run(subtitle.add_run("Investment Gate Preparation"), size=15, color=DARK_BLUE, bold=True)
    for label, value in (
        ("Aplicación", "P&Pmis Construction AI"),
        ("Fecha", "26 de agosto de 2026"),
        ("Repositorio", "pypmis-ia-sas · branch main"),
        ("Migración", "20260826_0049 (head)"),
        ("Estado", "IMPLEMENTADO, DESPLEGADO Y VALIDADO"),
    ):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = 2
        set_run(paragraph.add_run(f"{label}: "), size=10.5, color=INK, bold=True)
        set_run(paragraph.add_run(value), size=10.5)
    rule = doc.add_paragraph()
    add_bottom_rule(rule)
    add_callout(
        doc,
        "Resultado ejecutivo",
        "Gate 07J PASS · READY_FOR_INVESTMENT_DECISION · can_enter_fid=false. "
        "La salida es un paquete de evidencia congelado para una decisión posterior; no es aprobación de inversión.",
        fill=POSITIVE,
    )
    add_table(
        doc,
        ["Capa", "Resultado", "Evidencia"],
        [
            ("Backend / API", "PASS", "Módulo Gate 07J; ETag; idempotencia; SOD; SecurityEvent"),
            ("PostgreSQL", "PASS", "110 pruebas; upgrade/downgrade/upgrade"),
            ("Frontend", "PASS", "TypeScript; build Vite; Vitest; navegación USER/ADMIN"),
            ("Browser", "PASS", "USER/ADMIN reales; 0 errores de consola"),
            ("Localhost", "ACTIVE", "http://localhost:5173/app · schema head 0049"),
        ],
        [1900, 1300, 6160],
        font_size=9.1,
    )
    doc.add_page_break()


def add_contents(doc: Document) -> None:
    doc.add_heading("Contenido y trazabilidad", level=1)
    add_table(
        doc,
        ["Capítulo", "Alcance", "Cobertura"],
        [
            ("A", "Gobierno, fuente y perfiles", "Entregables 1–11"),
            ("B", "Dominio y evidencia", "Entregables 12–25"),
            ("C", "Readiness, lifecycle y handoff", "Entregables 26–34"),
            ("D", "Seguridad, API y experiencia", "Entregables 35–43"),
            ("E", "Validación, despliegue y cierre", "Entregables 44–48"),
        ],
        [1100, 5200, 3060],
    )
    add_callout(doc, "Lectura del informe", "Los 48 entregables están numerados y respaldados por evidencia ejecutada en esta entrega.")
    doc.add_page_break()


def add_architecture(doc: Document) -> None:
    doc.add_heading("Arquitectura de preparación implementada", level=1)
    add_table(
        doc,
        ["Entrada gobernada", "Gate 07J", "Salida controlada"],
        [
            ("Gate 07I RECOMMENDED/current", "Profiles + Evidence Package + SOD + Coverage", "FROZEN preparation hash"),
            ("Gate 07G/07H refs + Gate 07F hashes", "DRAFT → IN_PREPARATION → READY_FOR_REVIEW → FROZEN", "READY_FOR_INVESTMENT_DECISION"),
            ("Selected CAPITAL_OWNER Projects", "No scoring; no approval; explicit refresh", "can_enter_fid=false"),
        ],
        [2850, 3580, 2930],
        font_size=8.8,
    )
    add_callout(doc, "Boundary", "Gate 07J no implementa Investment Decision, FID, Funds, FEL/PDRI scoring, Initialization, Activation ni Execution.", fill=RISK)


def add_deliverables(doc: Document) -> None:
    chapters = {1: "A. Gobierno, fuente y perfiles", 12: "B. Dominio y evidencia", 26: "C. Readiness, lifecycle y handoff", 35: "D. Seguridad, API y experiencia", 44: "E. Validación, despliegue y cierre"}
    for number, title, body, evidence in SECTIONS:
        if number in chapters:
            doc.add_heading(chapters[number], level=1)
        heading = doc.add_heading(f"{number}. {title}", level=2)
        heading.paragraph_format.keep_with_next = True
        paragraph = doc.add_paragraph(body)
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.keep_with_next = True
        fill = POSITIVE if "PASS" in evidence or number == 48 else RISK if "risk" in title.lower() or "STOP" in evidence else CALLOUT
        add_callout(doc, "Evidencia / decisión", evidence, fill=fill)


def add_appendix(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("Anexo A · Inventario técnico", level=1)
    add_table(
        doc,
        ["Componente", "Implementación"],
        [
            ("Backend", "app/modules/investment_gate_preparation/{models,schemas,service,router}.py"),
            ("Migración", "alembic/versions/20260826_0049_gate07j_investment_gate_preparation.py"),
            ("Frontend", "src/features/investment-gate-preparation/*"),
            ("Navegación", "applicationNavigation.ts + workspace_context/navigator.py"),
            ("Seguridad", "constants.py + permissions.py + SecurityEvent"),
            ("Tests", "unit, PostgreSQL, Vitest y Playwright/browser"),
            ("ADR", "docs/50-adr-gate07j-investment-gate-preparation.md"),
        ],
        [2000, 7360],
    )
    doc.add_heading("Anexo B · Matriz de validación", level=1)
    add_table(
        doc,
        ["Validación", "Resultado", "Conteo / nota"],
        [
            ("Gate 07J unit", "PASS", "4/4"),
            ("PostgreSQL + regression", "PASS", "110/110"),
            ("Frontend build + Vitest", "PASS", "TypeScript + 1/1"),
            ("Browser USER/ADMIN", "PASS", "0 errores de consola"),
            ("Migration cycle", "PASS", "0048 → 0049 → 0048 → 0049"),
            ("Ruff / Prettier / ESLint", "PASS", "0 errores Gate 07J"),
        ],
        [3400, 1400, 4560],
    )
    add_callout(doc, "STOP", "No iniciar Investment Decision/FID, GO/HOLD/RECYCLE/KILL, Funds/Funding, FEL/PDRI scoring, Initialization, Activation o Execution sin autorización explícita.", fill=RISK)


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)
    add_title_page(doc)
    add_contents(doc)
    add_architecture(doc)
    add_deliverables(doc)
    add_appendix(doc)
    properties = doc.core_properties
    properties.title = "Informe Técnico Gate 07J — Investment Gate Preparation"
    properties.subject = "Implementación, validación y despliegue técnico Gate 07J"
    properties.author = "P&Pmis Construction AI"
    properties.keywords = "Gate 07J, Investment Gate Preparation, evidence, readiness, FID boundary"
    properties.comments = "Generated from verified repository evidence on 2026-08-26."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
