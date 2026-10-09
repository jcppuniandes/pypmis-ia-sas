from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

from gate07j_fel_report_data import BASELINE_FACTS, CLOSEOUT_ACCEPTANCE, CLOSEOUT_RUNS, REMEDIATIONS, SECTIONS
from generate_gate07i_report import (
    CALLOUT,
    DARK_BLUE,
    INK,
    MUTED,
    POSITIVE,
    RISK,
    TEAL,
    add_bottom_rule,
    add_callout,
    add_page_number,
    add_table,
    configure_styles,
    set_image_alt,
    set_run,
)

REPO = Path(__file__).resolve().parents[1]
OUTPUT = REPO / "outputs" / "technical_reports" / "Informe_Tecnico_Gate_07J_FEL_Project_Definition_20260826.docx"
LOGO = REPO / "frontend" / "public" / "pypmis-construction-ai-logo.png"


def add_title_page(doc: Document) -> None:
    doc.settings.odd_and_even_pages_header_footer = True
    section = doc.sections[0]
    for header in (section.header.paragraphs[0], section.even_page_header.paragraphs[0]):
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_run(header.add_run("TECHNICAL REPORT - GATE 07J-FEL"), size=8.5, color=MUTED, bold=True)
        add_bottom_rule(header, color="D9E1E8", size=6)
    add_page_number(section.footer.paragraphs[0])
    add_page_number(section.even_page_footer.paragraphs[0])

    if LOGO.exists():
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = 8
        picture = paragraph.add_run().add_picture(str(LOGO), width=Inches(0.72))
        set_image_alt(picture, "Logotipo de P&Pmis Construction AI")
    kicker = doc.add_paragraph()
    set_run(kicker.add_run("INFORME TÉCNICO DE IMPLEMENTACIÓN"), size=10, color=TEAL, bold=True)
    title = doc.add_paragraph()
    set_run(title.add_run("Gate 07J-FEL"), size=25, color=INK, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = 16
    set_run(subtitle.add_run("Project Definition / FEL Manager"), size=15, color=DARK_BLUE, bold=True)
    for label, value in (
        ("Aplicación", "P&Pmis Construction AI"),
        ("Fecha", "27 de agosto de 2026"),
        ("Repositorio", "pypmis-ia-sas - branch main"),
        ("Migración", "20260826_0050 (head)"),
        ("Estado", "IMPLEMENTADO, DESPLEGADO Y VALIDADO"),
    ):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = 2
        set_run(paragraph.add_run(f"{label}: "), size=10.5, color=INK, bold=True)
        set_run(paragraph.add_run(value), size=10.5)
    rule = doc.add_paragraph()
    add_bottom_rule(rule)
    add_callout(
        doc,
        "Resultado ejecutivo",
        "Gate 07J-FEL-H PASS - GATE07J_FEL_RELEASE_VALIDATED. Tres corridas E2E consecutivas validaron los "
        "lifecycles USER y ADMIN contra API y PostgreSQL reales, sin contaminación del baseline y sin ampliar el alcance funcional.",
        fill=POSITIVE,
    )
    add_table(
        doc,
        ["Capa", "Resultado", "Evidencia"],
        [
            ("Backend / API", "PASS", "FEL, assessments, maturity, actions, readiness y source adapter"),
            ("PostgreSQL", "PASS", "116/116 pruebas; upgrade/downgrade/upgrade"),
            ("Frontend", "PASS", "206 suite + 1 isolated; TypeScript, Prettier, ESLint y Vite"),
            ("Browser E2E", "PASS", "3/3 USER+ADMIN; consola, página y red limpias"),
            ("Localhost", "ACTIVE", "http://localhost:5173/app - schema head 0050"),
        ],
        [1900, 1300, 6160],
        font_size=9.1,
    )
    doc.add_page_break()


def add_contents(doc: Document) -> None:
    doc.add_heading("Contenido y trazabilidad", level=1)
    add_table(
        doc,
        ["Capítulo", "Alcance", "Cobertura"],
        [
            ("A", "Arquitectura, gobierno y límites", "Entregables 1-10"),
            ("B", "Frameworks, scoring y configuración", "Entregables 11-19"),
            ("C", "Dominio, lifecycle y readiness", "Entregables 20-36"),
            ("D", "Seguridad, API y experiencia", "Entregables 37-45"),
            ("E", "Persistencia, validación y cierre", "Entregables 46-54"),
            ("F", "Browser E2E Release Closeout", "3 runs, matriz, remediaciones y baseline"),
        ],
        [1100, 5200, 3060],
    )
    add_callout(
        doc,
        "Lectura del informe",
        "Los 54 entregables se conservan y el capítulo F añade el cierre funcional E2E, las remediaciones y la prueba de no contaminación.",
    )
    doc.add_page_break()


def add_architecture(doc: Document) -> None:
    doc.add_heading("Arquitectura de evidencia de definición", level=1)
    add_table(
        doc,
        ["Configuración publicada", "Project Definition / FEL", "Consumidor"],
        [
            ("Framework + stages + mappings", "FEL status + assessments + maturity", "Investment Gate Preparation"),
            ("Elements + scoring direction + thresholds", "Evidence + gaps + definition actions", "PROJECT_DEFINITION_FEL contract"),
            ("Review + SOD + inheritance", "READY / BLOCKED / STALE / NOT_STARTED", "Gate requirement source adapter"),
        ],
        [3000, 3360, 3000],
        font_size=8.8,
    )
    add_callout(
        doc,
        "Separación de autoridad",
        "Project Definition/FEL informa madurez y evidencia. Gate 07J prepara un paquete. Ninguna de las dos capacidades decide inversión, entra a FID o activa ejecución.",
        fill=RISK,
    )


def add_deliverables(doc: Document) -> None:
    chapters = {
        1: "A. Arquitectura, gobierno y límites",
        11: "B. Frameworks, scoring y configuración",
        20: "C. Dominio, lifecycle y readiness",
        37: "D. Seguridad, API y experiencia",
        46: "E. Persistencia, validación y cierre",
    }
    for number, title, body, evidence in SECTIONS:
        if number in chapters:
            doc.add_heading(chapters[number], level=1)
        heading = doc.add_heading(f"{number}. {title}", level=2)
        heading.paragraph_format.keep_with_next = True
        paragraph = doc.add_paragraph(body)
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.keep_with_next = True
        fill = POSITIVE if "PASS" in evidence or number == 54 else RISK if number in {10, 53} else CALLOUT
        add_callout(doc, "Evidencia / decisión", evidence, fill=fill)


def add_browser_closeout(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("F. Browser E2E Release Closeout", level=1)
    add_callout(
        doc,
        "Resultado de release",
        "GATE07J_FEL_RELEASE_VALIDATED. Las tres corridas consecutivas usaron fixtures independientes, backend real, "
        "PostgreSQL real y la interfaz servida en localhost. La navegación manual posterior confirmó USER MODE y ADMIN MODE; "
        "la certificación funcional proviene del E2E automatizado, no del smoke visual.",
        fill=POSITIVE,
    )
    add_table(
        doc,
        ["Ejecución", "Resultado", "Cobertura", "Duración", "Aislamiento"],
        CLOSEOUT_RUNS,
        [1550, 1500, 1900, 1350, 3060],
        font_size=8.8,
    )
    paragraph = doc.add_paragraph(
        "Cada corrida inició FEL, selló una evaluación, cambió readiness mediante una acción, aplicó SoD, creó una "
        "reevaluación y cerró el lifecycle de configuración ADMIN con Preview, ETag y publicación."
    )
    paragraph.paragraph_format.keep_with_next = True
    add_table(
        doc,
        ["ID", "Criterio", "Resultado", "Evidencia observada"],
        CLOSEOUT_ACCEPTANCE,
        [1100, 3550, 1200, 3510],
        font_size=8.25,
    )
    add_callout(
        doc,
        "Navegación vs. validación funcional",
        "El smoke del navegador integrado verificó disponibilidad, marca, cambio USER/ADMIN, navegación y carga del starter DRAFT. "
        "Las operaciones mutables y los casos negativos fueron ejecutados sólo sobre fixtures desechables por el E2E.",
    )
    add_table(
        doc,
        ["ID", "Hallazgo", "Corrección aplicada"],
        REMEDIATIONS,
        [1100, 3380, 4880],
        font_size=8.45,
    )
    add_callout(
        doc,
        "Protección de datos persistentes",
        "El teardown eliminó 39 workspaces, 24 configuraciones, 9 assessments, 6 FEL statuses, 3 actions y 3 reviewers "
        "del lote E2E. Los contadores posteriores y el hash del Project Workspace ID 14 coincidieron exactamente con el baseline previo.",
        fill=POSITIVE,
    )
    add_table(
        doc,
        ["Control", "Antes", "Después", "Resultado"],
        BASELINE_FACTS,
        [2600, 2700, 2700, 1360],
        font_size=8.25,
    )


def add_appendix(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("Anexo A - Inventario técnico", level=1)
    add_table(
        doc,
        ["Componente", "Implementación"],
        [
            ("Backend", "app/modules/project_definition/{models,schemas,service,router}.py"),
            ("Migración", "alembic/versions/20260826_0050_gate07j_fel_project_definition.py"),
            ("Integración", "investment_gate_preparation/service.py - governed source adapter"),
            ("Frontend", "src/features/project-definition/*"),
            ("Navegación", "applicationNavigation.ts + workspace_context/navigator.py"),
            ("Seguridad", "constants.py + permissions.py + SecurityEvent"),
            ("Tests", "unit, PostgreSQL, Vitest y browser/E2E"),
            ("ADR", "docs/51-adr-gate07j-fel-project-definition.md"),
        ],
        [2000, 7360],
    )
    doc.add_heading("Anexo B - Matriz de validación", level=1)
    add_table(
        doc,
        ["Validación", "Resultado", "Conteo / nota"],
        [
            ("PostgreSQL + regression", "PASS", "116/116"),
            ("Migration cycle", "PASS", "0049 -> 0050 -> 0049 -> 0050"),
            ("Ruff", "PASS", "0 errores en alcance Gate 07J-FEL"),
            ("Frontend lint/types/build", "PASS", "Prettier, ESLint, tsc y Vite"),
            ("Vitest", "PASS", "206 suite + 1 isolated; 207 casos cubiertos"),
            ("Browser E2E USER/ADMIN", "PASS", "3/3; persisted=false; 0 errores"),
            ("Baseline antes/después", "PASS", "Contadores y Project 14 MD5 idénticos"),
            ("Localhost", "ACTIVE", "frontend 5173; API 8000; schema 0050"),
        ],
        [3400, 1400, 4560],
    )
    add_callout(
        doc,
        "STOP",
        "No iniciar Funds/Funding, Investment Decision/FID, GO/HOLD/RECYCLE/KILL, Project Initialization, Activation, Execution, CPM o P6 sin autorización explícita.",
        fill=RISK,
    )


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)
    add_title_page(doc)
    add_contents(doc)
    add_architecture(doc)
    add_deliverables(doc)
    add_browser_closeout(doc)
    add_appendix(doc)
    properties = doc.core_properties
    properties.title = "Informe Técnico Gate 07J-FEL - Project Definition / FEL Manager"
    properties.subject = "Implementación, validación y despliegue técnico Gate 07J-FEL"
    properties.author = "P&Pmis Construction AI"
    properties.keywords = "Gate 07J, FEL, Project Definition, maturity, evidence, readiness"
    properties.comments = "Updated from verified repository evidence on 2026-08-27."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
