from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

REPO = Path(__file__).resolve().parents[1]
DOCX = (
    REPO
    / "outputs"
    / "technical_reports"
    / "Informe_Tecnico_Gate_07G_A_OPC_Budget_Alignment_20260830.docx"
)
OUTPUT = REPO / "outputs" / "technical_reports" / "gate07g_a_preset_audit.json"


def close(actual: int | None, expected: int, tolerance: int = 1000) -> bool:
    return actual is not None and abs(actual - expected) <= tolerance


def run() -> None:
    doc = Document(DOCX)
    section = doc.sections[0]
    paragraphs = list(doc.paragraphs)
    cell_paragraphs = [
        paragraph
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    ]
    text = " ".join(paragraph.text for paragraph in paragraphs + cell_paragraphs)
    checks = {
        "single_portrait_section": len(doc.sections) == 1 and section.page_height > section.page_width,
        "letter_page_size": close(section.page_width, Inches(8.5)) and close(section.page_height, Inches(11)),
        "one_inch_margins": all(
            close(value, Inches(1))
            for value in (section.left_margin, section.right_margin, section.top_margin, section.bottom_margin)
        ),
        "header_footer_distance": close(section.header_distance, Inches(0.492))
        and close(section.footer_distance, Inches(0.492)),
        "normal_11pt": close(doc.styles["Normal"].font.size, Pt(11), 100),
        "heading1_16pt": close(doc.styles["Heading 1"].font.size, Pt(16), 100),
        "heading2_13pt": close(doc.styles["Heading 2"].font.size, Pt(13), 100),
        "heading_hierarchy": sum(p.style.name == "Heading 1" for p in paragraphs) == 11
        and sum(p.style.name == "Heading 2" for p in paragraphs) == 8,
        "page_number_footer": " PAGE " in section.footer.paragraphs[0]._p.xml,
        "image_alt_present": len(doc.inline_shapes) >= 3
        and all(shape._inline.docPr.get("descr") for shape in doc.inline_shapes),
        "no_fake_bullets": not any(p.text.strip().startswith(("•", "- ")) for p in paragraphs),
        "no_placeholders": not re.search(r"\b(?:todo|lorem ipsum|placeholder|codex-file-citation)\b", text, re.IGNORECASE),
        "release_status": "GATE07G_A_OPC_ALIGNMENT_VALIDATED" in text,
        "waterline_contract": all(
            token in text
            for token in ("GATE07E_CONTEXTUAL_PORTFOLIO_RANK", "period_key", "No nueva priorización")
        ),
        "dependency_contract": all(
            token in text
            for token in ("ProjectDependency", "BLOCK_FREEZE", "OPTIONAL_WARNING", "Transitive unmet")
        ),
        "breakdown_contract": all(
            token in text for token in ("CAPEX", "EXPENSE", "UNCLASSIFIED", "LEGACY_UNCLASSIFIED")
        ),
        "validation_evidence": all(
            token in text for token in ("69/69", "60/60", "2/2", "1/1", "0051 → 0052")
        ),
        "stop_boundary": all(
            token in text
            for token in ("optimizer", "Scenario winner", "Investment Decision/FID", "Activation", "Execution")
        ),
    }
    table_geometry = True
    header_flags = True
    margins = True
    rows_no_split = True
    for table in doc.tables:
        table_geometry &= table._tbl.tblPr.find(qn("w:tblW")) is not None
        table_geometry &= table._tbl.tblPr.find(qn("w:tblInd")) is not None
        header_flags &= table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None
        for row in table.rows:
            rows_no_split &= row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
            for cell in row.cells:
                margins &= cell._tc.get_or_add_tcPr().find(qn("w:tcMar")) is not None
    checks.update(
        {
            "table_geometry": table_geometry,
            "table_header_flags": header_flags,
            "table_cell_margins": margins,
            "table_rows_no_split": rows_no_split,
        }
    )
    report = {
        "preset": "standard_business_brief",
        "first_page_pattern": "memo_masthead",
        "file": str(DOCX),
        "tables": len(doc.tables),
        "inline_images": len(doc.inline_shapes),
        "checks": checks,
        "result": "PASS" if all(checks.values()) else "FAIL",
        "failed": [name for name, passed in checks.items() if not passed],
    }
    OUTPUT.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(report, indent=2, ensure_ascii=False))
    if report["result"] != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    run()
