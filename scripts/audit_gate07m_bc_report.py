"""Preset, structure and content audit for the Gate 07M-BC Word report."""

from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

REPO = Path(__file__).resolve().parents[1]
DOCX = REPO / "outputs" / "technical_reports" / "Informe_Tecnico_Gate_07M_BC_Business_Case_Manager_20260902.docx"
OUTPUT = REPO / "outputs" / "technical_reports" / "gate07m_bc_preset_audit.json"


def close(actual: int | None, expected: int, tolerance: int = 1000) -> bool:
    return actual is not None and abs(actual - expected) <= tolerance


def run() -> None:
    doc = Document(DOCX)
    section = doc.sections[0]
    normal = doc.styles["Normal"]
    heading1 = doc.styles["Heading 1"]
    heading2 = doc.styles["Heading 2"]
    body_text = " ".join(paragraph.text for paragraph in doc.paragraphs)
    body_text += " " + " ".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    checks: dict[str, bool] = {
        "single_portrait_section": len(doc.sections) == 1 and section.page_height > section.page_width,
        "letter_page_size": close(section.page_width, Inches(8.5)) and close(section.page_height, Inches(11)),
        "one_inch_margins": all(
            close(value, Inches(1))
            for value in (section.left_margin, section.right_margin, section.top_margin, section.bottom_margin)
        ),
        "header_footer_distance": close(section.header_distance, Inches(0.492))
        and close(section.footer_distance, Inches(0.492)),
        "normal_calibri_11pt": normal.font.name == "Calibri" and close(normal.font.size, Pt(11), 100),
        "normal_spacing": close(normal.paragraph_format.space_after, Pt(6), 100)
        and abs(float(normal.paragraph_format.line_spacing) - 1.10) < 0.001,
        "heading1_tokens": close(heading1.font.size, Pt(16), 100)
        and str(heading1.font.color.rgb) == "2E74B5"
        and close(heading1.paragraph_format.space_before, Pt(16), 100)
        and close(heading1.paragraph_format.space_after, Pt(8), 100),
        "heading2_tokens": close(heading2.font.size, Pt(13), 100)
        and str(heading2.font.color.rgb) == "2E74B5"
        and close(heading2.paragraph_format.space_before, Pt(12), 100)
        and close(heading2.paragraph_format.space_after, Pt(6), 100),
        "heading_hierarchy_present": sum(p.style.name == "Heading 1" for p in doc.paragraphs) == 9
        and sum(p.style.name == "Heading 2" for p in doc.paragraphs) >= 7,
        "header_present": "TECHNICAL REPORT · GATE 07M-BC" in section.header.paragraphs[0].text,
        "footer_present": "P&Pmis Construction AI" in section.footer.paragraphs[0].text,
        "image_alt_present": all(shape._inline.docPr.get("descr") for shape in doc.inline_shapes),
        "final_output_present": "READY_FOR_GOVERNED_BUSINESS_CASE_EVIDENCE" in body_text,
        "exclusive_state_present": "GATE07M_BC_RELEASE_VALIDATED" in body_text,
        "next_present": "BUILD_GATE07M_FIN" in body_text,
        "not_available_present": "NOT_AVAILABLE" in body_text,
        "traceability_74_present": any(
            cell.text.strip() == "74" for table in doc.tables for row in table.rows for cell in row.cells
        ),
        "no_fake_bullets": not any(p.text.strip().startswith(("•", "- ")) for p in doc.paragraphs),
        "no_placeholder_text": not re.search(
            r"\b(?:todo|lorem ipsum|placeholder|codex-file-citation)\b", body_text, flags=re.IGNORECASE
        ),
    }
    table_geometry = True
    row_header_flags = True
    cell_margins = True
    no_split_rows = True
    for table in doc.tables:
        properties = table._tbl.tblPr
        table_geometry &= properties.find(qn("w:tblW")) is not None
        table_geometry &= properties.find(qn("w:tblInd")) is not None
        row_header_flags &= table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None
        for row in table.rows:
            no_split_rows &= row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
            for cell in row.cells:
                cell_margins &= cell._tc.get_or_add_tcPr().find(qn("w:tcMar")) is not None
    checks.update(
        {
            "table_geometry_complete": table_geometry,
            "table_header_flags_complete": row_header_flags,
            "cell_margins_complete": cell_margins,
            "rows_prevent_split": no_split_rows,
        }
    )
    report = {
        "preset": "standard_business_brief",
        "first_page_pattern": "memo_masthead",
        "file": str(DOCX),
        "tables": len(doc.tables),
        "inline_images": len(doc.inline_shapes),
        "checks": checks,
        "result": "PASS" if all(checks.values()) else "FAIL",
        "failed": [key for key, value in checks.items() if not value],
    }
    OUTPUT.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(report, indent=2, ensure_ascii=False))
    if report["result"] != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    run()
