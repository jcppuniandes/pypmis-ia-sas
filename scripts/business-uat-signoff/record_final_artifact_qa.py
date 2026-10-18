"""Record actual complete-page review after root has inspected final renders."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "docs/validation/business-uat-signoff-20261018"
OUT = ROOT / "outputs/business-uat-signoff-20261018"
QA = ROOT / "outputs/.qa/business-uat-signoff-20261018"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--all-pages-reviewed", action="store_true", required=True)
    parser.add_argument("--draft", required=True)
    args = parser.parse_args()
    assert args.draft.isalnum()
    build = read_json(QA / "documents-build-manifest.json")
    workbook_build = read_json(QA / "workbooks/build-manifest.json")
    for manifest in [build, workbook_build]:
        assert manifest["data_sha256"] == digest(DATA / "review-pack-data.json")
        assert manifest["baseline_sha256"] == digest(DATA / "baseline.json")
    workbook = read_json(QA / "workbooks/workbook-verification.json")
    scenarios = read_json(QA / "workbooks/formula-scenarios.json")
    assert workbook["result"] == "PASS" and len(workbook["checks"]) == 21
    assert len(scenarios["scenarios"]) == 20 and scenarios["all_pass"]
    outputs = {}
    for item, label, audit_name in zip(build["outputs"], ["signoff", "report"], ["signoff", "technical"]):
        path = ROOT / item["path"]
        assert digest(path) == item["sha256"]
        assert item["source_parts_preserved"]
        a11y = read_json(QA / f"{audit_name}-a11y.json")
        assert not a11y["findings"] and not any(a11y["counts"].values())
        pdf = QA / f"word-{label}/{args.draft}/{path.stem}.pdf"
        image_dir = QA / f"word-{label}/pages-{args.draft}"
        exported = read_json(Path(str(pdf) + ".export.json"))
        provenance = read_json(image_dir / "render-provenance.json")
        assert exported["input_docx_sha256"] == provenance["input_docx_sha256"] == digest(path)
        assert exported["output_pdf_sha256"] == provenance["input_pdf_sha256"] == digest(pdf)
        reader = PdfReader(pdf)
        count = len(reader.pages)
        pages = [image_dir / f"page-{n}.png" for n in range(1, count + 1)]
        assert all(p.is_file() for p in pages)
        assert len(list(image_dir.glob("page-*.png"))) == count
        assert all(provenance["pages"][page.name] == digest(page) for page in pages)
        assert all(len(page.extract_text().strip()) > 80 for page in reader.pages)
        if label == "signoff":
            forms = [table for table in Document(path).tables if table.cell(0, 0).text == "Campo de registro autorizado"]
            assert len(forms) == 15
            assert all(len(table.rows) == 13 for table in forms)
            assert all(not row.cells[1].text.strip() for table in forms for row in table.rows[1:])
            form_pages = [page for page in reader.pages if "Campo de registro autorizado" in page.extract_text()]
            assert len(form_pages) == 15
            assert all("Fecha objetivo acordada" in page.extract_text() for page in form_pages)
        outputs[path.name] = {
            "sha256": digest(path),
            "bytes": path.stat().st_size,
            "page_count": count,
            "native_word_read_only_pdf": pdf.relative_to(ROOT).as_posix(),
            "pdf_sha256": digest(pdf),
            "visual_review": "ROOT_INSPECTED_EVERY_FINAL_PAGE",
            "rendered_pages": [{"page": n, "path": page.relative_to(ROOT).as_posix(), "sha256": digest(page)} for n, page in enumerate(pages, 1)],
            "a11y_findings": a11y["counts"],
            "source_preserved_parts": 17,
            "source_total_parts": 19,
        }
    result = {
        "schema_version": 1,
        "date": "2026-10-18",
        "result": "PASS",
        "activity": "DOCUMENT_AND_WORKBOOK_QA_NOT_BUSINESS_ACCEPTANCE",
        "data_sha256": digest(DATA / "review-pack-data.json"),
        "baseline_sha256": digest(DATA / "baseline.json"),
        "outputs": outputs,
        "word_human_forms": "15 forms, 12 human fields each blank; one page per area visually confirmed",
        "workbook_native_and_visual_checks": {"passed": 21, "total": 21, "source": "outputs/.qa/business-uat-signoff-20261018/workbooks/workbook-verification.json", "sha256": digest(QA / "workbooks/workbook-verification.json")},
        "workbook_formula_scenarios": {"passed": 20, "total": 20, "simulation_only": True, "restored_human_approvals": 0, "engine": "Artifact Tool, not native Microsoft Excel"},
        "business_approvals": 0,
        "business_release_authorized": False,
        "document_qa_is_not_application_uat": True,
    }
    (DATA / "final-artifact-qa.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"result": result["result"], "word_pages": {key: value["page_count"] for key, value in outputs.items()}, "workbook_checks": 21, "formula_scenarios": 20, "business_approvals": 0}, ensure_ascii=False))


if __name__ == "__main__":
    main()
