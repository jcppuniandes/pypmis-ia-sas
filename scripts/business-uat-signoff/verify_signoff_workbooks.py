"""Read-only semantic/native verification; never authors an Office workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "outputs/business-uat-signoff-20261018"
QA = ROOT / "outputs/.qa/business-uat-signoff-20261018/workbooks"
NAME = "P&Pmis_Strategic_Macroprocess_UAT_Traceability.xlsx"
TRACKER = "P&Pmis_Strategic_Macroprocess_Business_UAT_Signoff_Tracker.xlsx"
SOURCE = ROOT / "outputs/gate07j-controlled-class-b-20261016" / NAME
EXPECTED_SOURCE = "baad276e386ea9550b5bac0eafae8bf6e9910361ea29047da58d952ab5dd757a"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def xml(obj) -> str:
    return ET.tostring(obj.to_tree(), encoding="unicode")


def signature(cell):
    return {
        "value": cell.value,
        "data_type": cell.data_type,
        "font": str(cell.font),
        "fill": str(cell.fill),
        "border": str(cell.border),
        "alignment": str(cell.alignment),
        "protection": str(cell.protection),
        "number_format": cell.number_format,
        "hyperlink": str(cell.hyperlink),
        "comment": str(cell.comment),
    }


def blank_cells(sheet, ranges):
    nonblank = []
    for ref in ranges:
        for row in sheet[ref]:
            for cell in row:
                if cell.value is not None and cell.value != "":
                    nonblank.append(cell.coordinate)
    return nonblank


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--visual-reviewed", action="store_true")
    args = parser.parse_args()
    source = load_workbook(SOURCE, data_only=False)
    successor = load_workbook(OUT / NAME, data_only=False)
    tracker = load_workbook(OUT / TRACKER, data_only=False)
    cached = load_workbook(OUT / TRACKER, data_only=True)
    checks = []

    def check(name, result, detail=None):
        checks.append({"check": name, "pass": bool(result), "detail": detail})

    check("Immutable source SHA256", digest(SOURCE) == EXPECTED_SOURCE)
    check("Eight original trace tabs in original order", successor.sheetnames == source.sheetnames and len(source.sheetnames) == 8)
    diffs = []
    feature_diffs = []
    for original in source:
        current = successor[original.title]
        for row in original:
            for cell in row:
                if signature(cell) != signature(current[cell.coordinate]):
                    diffs.append(f"{original.title}!{cell.coordinate}")
        for label, a, b in [
            ("tables", [xml(t) for t in original.tables.values()], [xml(t) for t in current.tables.values()]),
            ("merges", str(original.merged_cells), str(current.merged_cells)),
            ("freeze", original.freeze_panes, current.freeze_panes),
            ("gridlines", original.sheet_view.showGridLines, current.sheet_view.showGridLines),
            ("conditional_formats", [(str(k), [xml(r) for r in v]) for k, v in original.conditional_formatting._cf_rules.items()], [(str(k), [xml(r) for r in v]) for k, v in current.conditional_formatting._cf_rules.items()]),
        ]:
            if a != b:
                feature_diffs.append(f"{original.title}: {label}")
        source_dv = [xml(v) for v in original.data_validations.dataValidation]
        current_dv = [xml(v) for v in current.data_validations.dataValidation]
        if not all(v in current_dv for v in source_dv):
            feature_diffs.append(f"{original.title}: source validation changed/lost")
        added_dv = [v for v in current.data_validations.dataValidation if xml(v) not in source_dv]
        if added_dv and (original.title != "02_UAT_Cases" or any(str(v.sqref) != "AB8:AB58" for v in added_dv)):
            feature_diffs.append(f"{original.title}: unexpected validation addition")
        for row, dim in original.row_dimensions.items():
            new = current.row_dimensions[row]
            if (dim.height, dim.hidden, dim.outlineLevel) != (new.height, new.hidden, new.outlineLevel):
                feature_diffs.append(f"{original.title}: row dimension {row}")
        for col, dim in original.column_dimensions.items():
            new = current.column_dimensions[col]
            if (dim.width, dim.hidden, dim.outlineLevel) != (new.width, new.hidden, new.outlineLevel):
                feature_diffs.append(f"{original.title}: column dimension {col}")
    check("All existing trace cells preserve value formula and semantic style", not diffs, diffs[:30])
    check("Native tables merges panes validations CF and dimensions retained", not feature_diffs, feature_diffs[:30])
    check("Eight original native trace tables", sum(len(s.tables) for s in successor) == 8)
    check("Defined names retained", [xml(v) for v in source.defined_names.values()] == [xml(v) for v in successor.defined_names.values()])
    original_ids = [source["02_UAT_Cases"].cell(r, 1).value for r in range(8, 59)]
    current_ids = [successor["02_UAT_Cases"].cell(r, 1).value for r in range(8, 59)]
    check("51 original case IDs unchanged", original_ids == current_ids and len(set(current_ids)) == 51)
    case_sheet = successor["02_UAT_Cases"]
    check("Only five prepared signoff fields added", all(case_sheet.cell(7, c).value for c in range(27, 32)) and case_sheet.max_column == 31)
    check("Trace business decisions approval evidence conditions blank", not blank_cells(case_sheet, ["AB8:AD58"]))
    check("Trace human signatures preserve blank source", not blank_cells(successor["08_Release_Criteria"], ["E23:I37"]))
    expected_tabs = ["01_Signoff_Matrix", "02_Acceptance_Conditions", "03_Known_Findings", "04_Evidence_Register", "05_Document_Acceptance", "06_Final_Release_Decision"]
    check("Six exact tracker tabs", tracker.sheetnames == expected_tabs)
    blank_scope = {
        expected_tabs[0]: ["C8:N22", "D27:H29"],
        expected_tabs[1]: ["F8:F15", "H8:I15", "K8:O15"],
        expected_tabs[2]: ["G8:O15"],
        expected_tabs[4]: ["F8:K97"],
        expected_tabs[5]: ["B30:C31", "A41:J55"],
    }
    nonblank = {sheet: blank_cells(tracker[sheet], ranges) for sheet, ranges in blank_scope.items()}
    check("Every actual human identity decision date review signature blank", not any(nonblank.values()), nonblank)
    check("8 candidate conditions not accepted", [tracker[expected_tabs[1]].cell(r, 10).value for r in range(8, 16)] == ["PROPOSED_NOT_ACCEPTED"] * 8)
    check("90 document acceptance rows", len([r for r in range(8, 98) if tracker[expected_tabs[4]].cell(r, 1).value]) == 90)
    values = cached[expected_tabs[5]]
    check("Cached authorized completions 0 of 15", values["B5"].value == 0 and values["B6"].value == 15 and values["B7"].value == 0)
    check("Cached decision remains UAT_SIGNOFF_PENDING", values["B34"].value == "UAT_SIGNOFF_PENDING")
    check("Cached counts ACCEPT AWC REJECT conditions rejections all zero", all(values[c].value == 0 for c in ["B8", "B9", "B10", "B12", "B19"]))
    check("Actual current runtime continuity proof not inferred", values["B27"].value == "UNCHANGED_OR_REVALIDATED")
    formula_errors = []
    for sheet in cached:
        for row in sheet:
            for cell in row:
                if cell.data_type == "e":
                    formula_errors.append(f"{sheet.title}!{cell.coordinate}={cell.value}")
    check("Saved formula cached values free of errors", not formula_errors, formula_errors)
    scenarios = json.loads((QA / "formula-scenarios.json").read_text())
    required_scenarios = {
        "All fields blank", "ACCEPT text only", "14 of 15 complete",
        "15 ACCEPT but one missing approval", "15 ACCEPT but findings unreviewed",
        "Authorized material REJECT", "Accepted release-blocking condition",
        "Nonblocking condition missing owner/date",
        "All ACCEPT and all review obligations satisfied",
        "Documented nonblocking accepted condition in authorized area",
        "All approvals but missing four-MISSING acknowledgment",
        "All approvals but missing document approval",
        "REJECT text without material rejection register",
        "All 15 ACCEPT cannot silently become WITH_CONDITIONS",
        "AWC unknown condition ID", "AWC condition belongs to a different area",
        "AWC without an accepted open condition",
        "RESOLVED condition with contradictory blocking YES stays pending",
        "Open S2 finding cannot be accepted away",
        "Accepted S2 condition cannot be treated as nonblocking release",
    }
    check("Actual formula boundary scenarios passed then restored", scenarios["all_pass"] and scenarios["restored_human_approvals"] == 0 and required_scenarios == {s["name"] for s in scenarios["scenarios"]})
    check("Visual QA of every tracker sheet and affected trace fields", args.visual_reviewed)
    result = {"result": "PASS" if all(c["pass"] for c in checks) else "NOT_VALIDATED", "checks": checks, "formula_engine": "Artifact Tool. Native Microsoft Excel was not executed.", "source_sha256": digest(SOURCE), "outputs": {NAME: digest(OUT / NAME), TRACKER: digest(OUT / TRACKER)}}
    (QA / "workbook-verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result["result"] == "PASS" else 1)


if __name__ == "__main__":
    main()
