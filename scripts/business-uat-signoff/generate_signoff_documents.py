"""Build two source-derived Word documents without recording business approval.

Run only after the root agent has completed the two-DOCX authoring marker.
No application command, runtime startup, database write or test run is performed.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "docs/validation/business-uat-signoff-20261018"
OUT = ROOT / "outputs/business-uat-signoff-20261018"
QA = ROOT / "outputs/.qa/business-uat-signoff-20261018"
REFERENCE = (
    ROOT
    / "outputs/gate07j-controlled-class-b-20261016"
    / "P&Pmis_Strategic_Macroprocess_Final_Release_Closeout.docx"
)
REFERENCE_SHA = "2a0c470c57b833b89c85758863653afa04a03b941481e18ee37aabcdfc76d031"
NAMES = (
    "P&Pmis_Strategic_Macroprocess_Business_UAT_Signoff.docx",
    "P&Pmis_Informe_Tecnico_Business_UAT_Signoff_20261018.docx",
)
LABELS = [
    "Representante autorizado",
    "Cargo o rol designado",
    "Organización",
    "Base de autoridad",
    "Fecha de decisión",
    "Decisión",
    "Condiciones acordadas",
    "Comentarios",
    "Evidencia efectivamente revisada",
    "Firma o referencia de aprobación",
    "Propietario de condición",
    "Fecha objetivo acordada",
]
ALLOWED_CHANGED_PARTS = {"word/document.xml", "word/footer1.xml"}


def load(path):
    return json.loads(path.read_text(encoding="utf-8-sig"))


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def prose(text):
    """Separate compressed prose counts without altering IDs or cryptographic hashes."""
    text = str(text)
    paths = []

    def protect_path(match):
        paths.append(match.group(0))
        return "PATHPROTECT" + chr(64 + len(paths)) + "X"

    text = re.sub(
        r"(?:outputs|docs|scripts|backend|frontend)/[^\s]+", protect_path, text
    )
    for before, after in [
        ("ACCEPT_WITH_CONDITIONS0", "ACCEPT_WITH_CONDITIONS 0"),
        ("conserva20261011_0068", "conserva 20261011_0068"),
        ("HEAD6672fd5", "HEAD 6672fd5"),
        ("MD5ef", "MD5 ef"),
        ("cadenaA0", "cadena A0"),
        ("susSHA", "sus SHA"),
        ("geometríaLetter", "geometría Letter"),
        ("independienteClass", "independiente Class"),
        ("ceroQRA", "cero QRA"),
    ]:
        text = text.replace(before, after)
    text = re.sub(r"(\d+)(S[12])\b", r"\1 \2", text)

    def display_token(match):
        token = match.group(0)
        if (
            "_" in token
            or re.fullmatch(r"[a-fA-F0-9]{7,}", token)
            or re.fullmatch(
                r"(?:BA|DOC|E|R|PROC|P|A|B|C|N|S)\d{1,3}|Gate\d+[A-Z]|E2E|MD5|SHA256",
                token,
            )
        ):
            return token
        return re.sub(
            r"(?<=[A-Za-zÁÉÍÓÚáéíóúÑñ])(?=\d)|(?<=\d)(?=[A-Za-zÁÉÍÓÚáéíóúÑñ])",
            " ",
            token,
        )

    text = re.sub(r"[A-Za-zÁÉÍÓÚáéíóúÑñ0-9_]+", display_token, text)
    text = re.sub(
        r"(?<=\d)(?=(?:áreas|condiciones|rechazos|días|campos|pasos|filas|workspaces|configuraciones|hojas|páginas|documentos|referencias|partes|decisiones|pendientes|entradas|casos|de)\b)",
        " ",
        text,
    )
    text = re.sub(
        r"\b(actor|revisor|Project|Tenant|Portfolio|HEAD|Class|publicador)(\d+)\b",
        r"\1 \2",
        text,
    )
    text = re.sub(r"\b(\d+)(PASS|SKIP|S1|S2|XML)\b", r"\1 \2", text)
    text = re.sub(r"\b(ACCEPT|REJECT|CONDITIONS)(\d+)\b", r"\1 \2", text)
    for before, after in [
        ("corte16", "corte del 16"),
        ("corte15", "corte del 15"),
        ("guías14", "guías del 14"),
        ("Closeout16", "Closeout del 16"),
        ("Guide14", "Guide del 14"),
        ("Guide15", "Guide del 15"),
        ("técnico16", "técnico del 16"),
        ("auditoría18", "auditoría del 18"),
        ("documentales18", "documentales del 18"),
        ("Word16", "Word del 16"),
        ("cierre16", "cierre del 16"),
        ("histórica16", "histórica del 16"),
        ("histórico14", "histórico del 14"),
        ("observación16", "observación del 16"),
        ("consulta18", "consulta del 18"),
        ("0%", "0 %"),
        ("100%", "100 %"),
        ("ClassB", "Class B"),
        ("lectorQRA", "lector QRA"),
        ("aceptorQRA", "aceptor QRA"),
        ("revisorGate", "revisor Gate"),
        ("autoridadFID", "autoridad FID"),
        ("activadorProject", "activador Project"),
        ("31workspaces", "31 workspaces"),
        ("external_keyPRJ", "external_key PRJ"),
        ("MD5ef", "MD5 ef"),
        ("1de1", "1 de 1"),
        ("19de19", "19 de 19"),
        ("A0-A31", "A0-A31"),
        ("Corte 14 Oct 2026", "Corte del 14 de octubre de 2026"),
        ("Remediación 15 Oct 2026", "Remediación del 15 de octubre de 2026"),
        ("Cierre 16 Oct 2026", "Cierre del 16 de octubre de 2026"),
        ("vigencia 18 Oct", "vigencia del 18 de octubre"),
    ]:
        text = text.replace(before, after)
    text = re.sub(r"([:;,])(?=\S)", r"\1 ", text)
    for index, path in enumerate(paths, 1):
        text = text.replace("PATHPROTECT" + chr(64 + index) + "X", path)
    return text


def paragraph(document, text="", style=None, bold=False):
    result = document.add_paragraph(style=style)
    run = result.add_run(prose(text))
    run.bold = bold
    result.paragraph_format.widow_control = True
    if bold:
        result.paragraph_format.keep_with_next = True
    return result


def heading(document, text, level=1):
    result = document.add_paragraph(text, style=f"Heading {level}")
    result.paragraph_format.keep_with_next = True
    result.paragraph_format.keep_together = True
    return result


def table(document, headers, rows, widths=None, blank_rows=False, trailing_space=True):
    """Reuse the reference table component with an intentional column grid."""
    reference = Document(REFERENCE).tables[0]
    result = document.add_table(rows=1, cols=len(headers))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    result._tbl.remove(result._tbl.tblPr)
    result._tbl.insert(0, deepcopy(reference._tbl.tblPr))
    usable = (7772400 - 2 * 658495) / 914400
    widths = widths or [usable / len(headers)] * len(headers)
    for index, width in enumerate(widths):
        result.columns[index].width = Inches(width)
    for row_index, row_values in enumerate([headers, *rows]):
        row = result.rows[0] if row_index == 0 else result.add_row()
        tr_pr = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        tr_pr.append(no_split)
        if row_index == 0:
            repeated = OxmlElement("w:tblHeader")
            repeated.set(qn("w:val"), "true")
            tr_pr.append(repeated)
        for column_index, value in enumerate(row_values):
            cell = row.cells[column_index]
            if cell._tc.tcPr is not None:
                cell._tc.remove(cell._tc.tcPr)
            properties = deepcopy(
                reference.cell(0 if row_index == 0 else 1, 0)._tc.tcPr
            )
            properties.find(qn("w:tcW")).set(
                qn("w:w"), str(round(widths[column_index] * 1440))
            )
            shade = properties.find(qn("w:shd"))
            if shade is None:
                shade = OxmlElement("w:shd")
                properties.append(shade)
            shade.set(
                qn("w:fill"),
                "17324D"
                if row_index == 0
                else "F3F6F8"
                if row_index % 2 == 0
                else "FFFFFF",
            )
            cell._tc.insert(0, properties)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.06
            p.paragraph_format.keep_with_next = row_index == 0 or (
                len(rows) <= 6 and row_index < len(rows)
            )
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(prose(value))
            run.font.name = "Aptos"
            run.font.size = Pt(9.5)
            run.bold = row_index == 0
            run.font.color.rgb = RGBColor.from_string(
                "FFFFFF" if row_index == 0 else "202A33"
            )
            if blank_rows and row_index > 0:
                p.paragraph_format.space_after = Pt(2)
    if trailing_space:
        paragraph(document)
    return result


def start_document(title, intro):
    source = Document(REFERENCE)
    logo = deepcopy(source.paragraphs[0]._p)
    section = deepcopy(source._element.body.sectPr)
    body = source._element.body
    for element in list(body):
        body.remove(element)
    body.append(logo)
    body.append(section)
    paragraph(source, title, "Title")
    paragraph(
        source,
        "P&Pmis Construction AI. Capital Portfolio & Investment Governance. "
        "Emisión del paquete 18 de octubre de 2026.",
        "Subtitle",
    )
    paragraph(source, intro)
    return source


def page(document, title):
    document.add_page_break()
    heading(document, title)


def evidence_names(data, ids):
    register = {row["id"]: row for row in data["evidence_register"]}
    return "; ".join(f"{key} {register[key]['title']}" for key in ids)


def baseline_paragraph(document, baseline):
    paragraph(
        document,
        "Código guardado: main / "
        + baseline["frozen_commit"]
        + ". La referencia original de la validación16 fue "
        "10a2442bfb69b97dc72f9c78e2419f157b400973. El commit actual incorpora "
        "el desarrollo guardado; sus cinco hashes de producto y nueve de fixtures "
        "coinciden con el cierre técnico. Alembic conserva20261011_0068, sin nueva migración.",
    )


def decisions_paragraph(document):
    paragraph(
        document,
        "Cada representante debe registrar exactamente una decisión: ACCEPT, "
        "ACCEPT_WITH_CONDITIONS o REJECT. Una condición requiere ID, descripción, "
        "severidad, impacto de liberación, workaround, propietario y fecha objetivo "
        "acordados. Un rechazo requiere motivo, proceso afectado, severidad, "
        "evidencia y remediación. No se infiere aceptación de silencio, asistencia, "
        "comentario no registrado o prueba técnica aprobada.",
    )


def findings_section(document, data):
    paragraph(
        document,
        "Los seis hallazgos siguientes permanecen OPEN. Ninguno está aceptado "
        "como condición de negocio. Propietarios, fechas y decisiones requieren "
        "acuerdo autorizado y permanecen vacíos en los formularios y el tracker.",
    )
    for item in data["findings"]:
        heading(document, item["id"].replace("-", " ") + " " + item["severity"], 2)
        paragraph(document, item["title"], bold=True)
        paragraph(document, item["impact"]).paragraph_format.keep_with_next = True
        paragraph(
            document, "Workaround: " + item["workaround"]
        ).paragraph_format.keep_with_next = True
        paragraph(
            document, "Evidencia propuesta: " + ", ".join(item["evidence_ids"]) + "."
        )
    heading(document, "Observaciones de consistencia visual", 2)
    for item in data["ui_observations"]:
        paragraph(document, item["id"] + ". " + item["title"], bold=True)
        paragraph(document, item["impact"] + " " + item["workaround"])
    paragraph(
        document,
        "UAT-S3-006 exige reconocimiento expreso por la posible omisión de Project "
        "al filtrar por responsable. El revisor elegirá aceptar una condición, "
        "rechazar la liberación o proponer reclasificación con justificación. "
        "No se cambia el backend para ajustar estados a su presentación.",
    )


def gate_section(document, data):
    paragraph(
        document,
        "El cierre técnico16 documenta19 requisitos required:15 VALID y4 MISSING "
        "no bloqueantes por la política publicada. La fila opcional Class C es "
        "NOT_REQUIRED. La cobertura Project es100%, pero no se declara19de19 "
        "requisitos satisfechos.",
    )
    table(
        document,
        ["Requisito que permanece MISSING", "Situación publicada"],
        [
            [key, "MISSING no bloqueante; no satisfecho"]
            for key in data["required_missing"]
        ],
        [3.6, 3.46],
    )
    paragraph(
        document,
        "Los cuatro Class B controlados fueron SUBMITTED por actor4 y VALID "
        "por revisor5, con EACH_PROJECT y cobertura1de1. Son evidencia "
        "representativa UAT, no licencias reales, contratos firmados, NTP ni "
        "autorización para ejecutar Delivery.",
    )
    table(
        document,
        ["Class B", "Control y ventana de validez capturada"],
        [
            ["permit-readiness", "Revisión independiente;180días"],
            ["contracting-strategy-readiness", "Revisión independiente;365días"],
            ["execution-plan-readiness", "Revisión independiente;365días"],
            ["schedule-maturity", "Revisión independiente;180días"],
        ],
        [3.6, 3.46],
    )
    paragraph(
        document,
        "Gate07J quedó READY y FROZEN sin auto-GO. Gate07K registró PYP-FFA / "
        "GO / DECIDED, is_fid=true, investment_authorized=true y "
        "full_funds_authorized=true, sin activación automática. Inicializó actor12 "
        "y activó actor13 mediante comandos canónicos separados. Post-FID quedó "
        "NO_REACTIVATION_REQUIRED después de la activación. Estos actores son "
        "cuentas de UAT, no representantes de aceptación empresarial.",
    )
    paragraph(
        document,
        "La aceptación debe reconocer que Recommendation no es FID; Budget "
        "FROZEN no es inversión aprobada; HOLD_FOR_CAPACITY no es Gate07K HOLD; "
        "FEL es etapa y PDRI es madurez. Funding no es Budget y FIN conserva su "
        "cálculo determinístico. QRA Cost/Schedule es contextual, no universal. "
        "Integrated QRA no está liberado. Un resultado QRA no autoriza FID, "
        "contingencia, actualización Budget ni write-back de cronograma.",
    )


def documentation_section(document, data):
    paragraph(
        document,
        "El paquete utiliza las versiones controladas E01 a E06 sin reescribir "
        "las fuentes. Cada revisor registrará para cada documento "
        "ACCEPTABLE, ACCEPTABLE_WITH_CORRECTIONS o NOT_ACCEPTABLE; la decisión "
        "documental se conserva vacía hasta revisión real.",
    )
    table(
        document,
        ["ID", "Documento y revisión", "Aceptación registrada"],
        [
            [r["id"], r["name"] + "\n" + r["revision"], ""]
            for r in data["required_documents"]
        ],
        [0.58, 4.52, 1.96],
    )
    heading(document, "Addendum de vigencia documental", 2)
    for text in data["documentation_addendum"]:
        paragraph(document, text)
    paragraph(
        document,
        "Final Release Closeout16 permanece byte-idéntico: solo se actualizará "
        "con la sección de Business Sign-Off después de decisiones empresariales "
        "reales, preservando su conclusión e historia técnica.",
    )


def area_forms(document, data):
    for area in data["areas"]:
        page(
            document,
            area["id"] + " " + area["name"].replace("/", " ").replace("&", "and"),
        )
        paragraph(document, area["purpose"])
        paragraph(
            document, "Alcance de revisión: " + ", ".join(area["review_step_ids"]) + "."
        )
        paragraph(
            document,
            "Evidencia propuesta: "
            + evidence_names(data, area["evidence_proposed"])
            + ".",
        )
        for number, question in enumerate(area["questions"], 1):
            paragraph(document, str(number) + ". " + question)
        if area["id"] == "BA14":
            paragraph(
                document,
                "Reconocimiento de los cuatro MISSING no bloqueantes revisado y "
                "registrado por el representante: ______________________________",
            )
        paragraph(
            document,
            "Opciones permitidas: ACCEPT / ACCEPT_WITH_CONDITIONS / REJECT. "
            "No hay opción seleccionada.",
        )
        table(
            document,
            ["Campo de registro autorizado", "Respuesta"],
            [[x, ""] for x in LABELS],
            [2.45, 4.61],
            blank_rows=True,
        )
        paragraph(
            document,
            "Revisar propósito, workflow, evidencias y hallazgos antes de registrar "
            "decisión y aprobación. Un campo vacío no es aceptación ni rechazo.",
        )


def final_decision(document, data, compact=False):
    heading(document, "Reglas de decisión y declaración actual", 2)
    paragraph(
        document,
        "BUSINESS_ACCEPTED exige15 ACCEPT explícitos, personas autorizadas, "
        "evidencia completa, ningún rechazo y0S1/S2 abiertos. "
        "BUSINESS_ACCEPTED_WITH_CONDITIONS exige15 decisiones completas sin "
        "REJECT, al menos un área con ACCEPT_WITH_CONDITIONS y condiciones "
        "no bloqueantes con propietarios, fechas y "
        "workarounds acordados, y0S1/S2. Un REJECT material o condición "
        "bloqueante produce BUSINESS_REJECTED hasta resolución autorizada.",
    )
    paragraph(
        document,
        "Con decisiones o firmas incompletas se conserva UAT_SIGNOFF_PENDING. "
        "Hoy0de15 áreas tienen decisión registrada:0 ACCEPT,0 "
        "ACCEPT_WITH_CONDITIONS,0 REJECT y15pendientes. Hay0condiciones "
        "aceptadas y0rechazos empresariales registrados; esos ceros no afirman "
        "aceptabilidad de los hallazgos abiertos.",
    )
    table(
        document,
        ["Registro de decisión final autorizada", "Contenido"],
        [
            [x, ""]
            for x in [
                "Decisión",
                "Representante y organización",
                "Base de autoridad",
                "Fecha",
                "Condiciones o rechazos aplicables",
                "Firma o referencia de aprobación",
            ]
        ],
        [2.8, 4.26],
        blank_rows=True,
        trailing_space=not compact,
    )
    paragraph(document, "Estado exacto conservado: " + data["release_state"] + ".")
    paragraph(
        document,
        "No se autoriza liberación empresarial ni producción. Recomendación única: "
        + data["next_recommendation"]
        + ".",
    )


def signoff_document(data, baseline):
    doc = start_document(
        "Aceptación empresarial del macroproceso estratégico",
        "El paquete está preparado para revisión de15áreas autorizadas. No "
        "contiene firmas ni decisiones. La evidencia técnica histórica permanece "
        "disponible y la aceptación empresarial continúa pendiente.",
    )
    heading(doc, "A Resumen ejecutivo de liberación")
    paragraph(
        doc,
        "Avance empresarial0de15 áreas,0%. ACCEPT0; ACCEPT_WITH_CONDITIONS0; REJECT0; pendientes15. Estado: UAT_SIGNOFF_PENDING.",
    )
    baseline_paragraph(doc, baseline)
    paragraph(
        doc,
        "La auditoría del 18 verificó integridad de código y evidencias y preservación READ ONLY de Project 14. Después de conservar la observación inicial de servicios detenidos, se inició solo la DB privada existente y se comprobó continuidad por SQL READ ONLY. Las filas completas de ocho tablas y la selección de negocio coinciden con el corte del 16. API y Redis siguen detenidos; no se reejecutaron pruebas ni se creó aprobación.",
    )
    decisions_paragraph(doc)
    page(doc, "B Alcance y propósito empresarial")
    paragraph(
        doc,
        "El alcance cubre preparación estratégica, planificación de capital, origen de inversión, Portfolio, Five Case, fuentes gobernadas, Gate07J, Gate07K/FID y gobernanza posterior hasta inicialización y activación de Project Workspace. Project es una instancia canónica PROJECT, no una identidad paralela.",
    )
    paragraph(
        doc,
        "Quedan fuera Project Delivery, compras y construcción operacionales, Field Management, Operational Readiness, Asset Operations, realización operativa de beneficios, P6 directo y QRA integrado o avanzado. No se comienza ninguno automáticamente después de preparar este paquete.",
    )
    heading(doc, "C Mapa de revisión del proceso", 2)
    paragraph(doc, data["review_steps_note"])
    table(
        doc,
        ["ID", "Actividad bajo revisión"],
        [[r["id"], r["name"]] for r in data["review_steps"]],
        [0.65, 6.41],
    )
    page(doc, "D Responsabilidades ADMIN y USER")
    paragraph(
        doc,
        "ADMIN configura y USER opera. La matriz distingue responsabilidades propuestas de permisos efectivos y de una delegación empresarial real. Deben respetarse tenant, ámbito, vigencia y separación de funciones.",
    )
    table(
        doc,
        ["Ámbito", "ADMIN", "USER"],
        [[r["domain"], r["admin"], r["user"]] for r in data["admin_user_rows"]],
        [1.5, 2.78, 2.78],
    )
    paragraph(
        doc,
        "La revisión de seguridad debe reconocer aislamiento tenant, RBAC, autor/publicador de política, ejecutor/aceptorQRA, preparador/revisorGate, autoridadFID, inicializador/activadorProject y validador independienteClassB. Un cambio de política solicitado se registra como cambio de negocio, no como edición silenciosa del sign-off.",
    )
    heading(doc, "Autoridades propuestas para revisión", 2)
    paragraph(
        doc,
        "La RACI fuente E11 contiene 35 procesos. R ejecuta, A responde por la aprobación, C se consulta e I se informa. Son responsabilidades funcionales propuestas, no nombres ni delegaciones aceptadas. El revisor debe confirmar la asignación organizacional real y la independencia exigida por la política efectiva.",
    )
    table(
        doc,
        ["Actividad", "R propuesto", "A propuesto"],
        [
            [
                "Publicación de configuración",
                "Administrador autor",
                "Publicador independiente cuando se exige",
            ],
            [
                "Business Case",
                "Preparador Business Case",
                "Revisor Business Case independiente",
            ],
            [
                "Modelo y resultado QRA",
                "Analista QRA",
                "Revisor o aceptor independiente",
            ],
            [
                "Gate evidence y freeze",
                "Gate Preparer",
                "Investment Evidence Reviewer independiente",
            ],
            [
                "Decisión FID",
                "Investment Decision Preparer",
                "Decision Maker o comité autorizado",
            ],
            [
                "Inicialización y activación",
                "Project Workspace Initializer",
                "Project Workspace Activator independiente",
            ],
        ],
        [2.05, 2.25, 2.76],
    )
    page(doc, "E F G Rutas de negocio bajo aceptación")
    for letter, route in zip(["E", "F", "G"], data["routes"]):
        heading(doc, letter + " " + route["id"].replace("_", " "), 2)
        paragraph(doc, route["purpose"])
        paragraph(doc, route["sequence"])
        paragraph(
            doc,
            "Comprobación disponible: " + route["validated"] + " " + route["boundary"],
        )
        paragraph(doc, "Evidencia propuesta: " + ", ".join(route["evidence_ids"]) + ".")
    page(doc, "H Gate07J FID y Post FID")
    gate_section(doc, data)
    page(doc, "I Hallazgos conocidos para decisión")
    findings_section(doc, data)
    page(doc, "J Documentación de negocio y evidencia")
    documentation_section(doc, data)
    page(doc, "Registro de evidencias propuestas")
    paragraph(
        doc,
        "Las referencias siguientes identifican las fuentes de revisión. Sus rutas, clasificación y límites de prueba están en la hoja04 del tracker; los hashes constan en el manifiesto controlado del paquete. La revisión técnica histórica no constituye aceptación humana; E23 es continuidad SQL actual, no un E2E nuevo.",
    )
    for item in data["evidence_register"]:
        paragraph(
            doc,
            item["id"] + " " + item["title"] + ". Corte " + item["cut"] + ".",
        )
    page(doc, "K Matriz de áreas y formularios de aprobación")
    paragraph(
        doc,
        "La matriz resume cobertura y estado de registro; no anticipa el resultado de una revisión. Las15páginas de formulario posteriores capturan los12campos de cada área. La evidencia propuesta está separada de la evidencia efectivamente revisada.",
    )
    paragraph(
        doc,
        "El tracker contiene la matriz completa de sign-off con representante, rol, evidencia revisada, decisión, condiciones, propietario, plazo y aprobación. En este documento el índice siguiente remite a los formularios de captura, sin prellenar esos campos humanos.",
    )
    table(
        doc,
        ["Área", "Cobertura propuesta", "Estado de registro"],
        [
            [
                a["id"] + " " + a["name"],
                ", ".join(a["review_step_ids"]),
                "Pendiente sin decisión",
            ]
            for a in data["areas"]
        ],
        [3.35, 1.81, 1.9],
    )
    area_forms(doc, data)
    page(doc, "L Condiciones rechazos y decisión final")
    paragraph(
        doc,
        "Registro de condiciones: no hay condiciones empresariales acordadas. Los hallazgos son candidatos a revisión, no condiciones aceptadas. El tracker permite capturar todos los atributos sin prellenar propietarios ni fechas.",
    )
    paragraph(
        doc,
        "Cada condición se registrará con ID, área, hallazgo o descripción, severidad, bloqueo de liberación explícito, workaround, propietario, fecha objetivo y estado. Ninguno de estos acuerdos se presupone.",
    )
    table(
        doc,
        [
            "Condición y área",
            "Severidad e impacto",
            "Workaround",
            "Propietario fecha estado",
        ],
        [["", "", "", ""]],
        [1.8, 1.7, 1.65, 1.91],
        blank_rows=True,
        trailing_space=False,
    )
    paragraph(doc, "Registro de rechazos: no hay rechazos autorizados registrados.")
    table(
        doc,
        [
            "Rechazo área representante",
            "Motivo proceso evidencia",
            "Severidad remediación impacto",
        ],
        [["", "", ""]],
        [2.3, 2.38, 2.38],
        blank_rows=True,
        trailing_space=False,
    )
    final_decision(doc, data, compact=True)
    return doc


def technical_document(data, baseline):
    doc = start_document(
        "Informe técnico de preparación de aceptación empresarial",
        "Se preparó el paquete de aceptación y su seguimiento sin cambiar "
        "funciones, políticas o registros de negocio. Las15decisiones "
        "autorizadas siguen pendientes. Este informe separa la evidencia "
        "técnica histórica de la auditoría actual y de la aceptación humana.",
    )
    heading(doc, "1 Resultado del incremento")
    paragraph(
        doc,
        "Se creó un documento de aceptación, este informe y un tracker de6hojas, además de referencias de sign-off en una copia controlada de trazabilidad. El paquete facilita revisión por15áreas y35actividades sin duplicar lógica de producto ni sustituir autoridades empresariales.",
    )
    paragraph(
        doc,
        "No hay decisiones, firmas, condiciones aceptadas ni rechazos registrados. Avance0de15,0%. El estado sigue siendo "
        + data["release_state"]
        + ".",
    )
    heading(doc, "2 Código guardado y congelación")
    baseline_paragraph(doc, baseline)
    paragraph(
        doc,
        "El árbol inicial de main estaba limpio. La captura de git status, diff, diffstat y log está en baseline.json y startup-review.json. Los cambios de este incremento se restringen a documentación, datos del paquete y scripts de artefactos. No se modifican archivos backend/frontend de producto.",
    )
    heading(doc, "3 Validaciones históricas y auditoría actual")
    counts = baseline["historical_test_counts"]
    table(
        doc,
        ["Ejecución histórica", "PASS", "FAIL", "ERROR", "SKIP"],
        [
            [
                name,
                *(str(counts[name][key]) for key in ["PASS", "FAIL", "ERROR", "SKIP"]),
            ]
            for name in [
                "broad",
                "browser",
                "class_b",
                "bridge",
                "canonical",
                "postactivation",
                "frontend",
                "legacy",
                "ci_contract_guard",
            ]
        ],
        [2.42, 1.16, 1.16, 1.16, 1.16],
    )
    paragraph(
        doc,
        "Cada XML se informa por separado; los focused no se suman a898PASS de914entradas broad. Los16SKIP siguen siendo omisiones, no PASS. La auditoría18 comprobó los hashes de9XML; no ejecutó pruebas nuevas ni browser nuevo.",
    )
    paragraph(
        doc,
        "La observación inicial encontró la UAT privada detenida y su volumen PostgreSQL 16 existente. Se conserva intacta esa captura en baseline.json. Root inició después solo la DB privada existente; runtime-readonly-20261018.json acredita transaction_read_only=on, selección de negocio idéntica al cierre del 16 y conteos y hashes de filas completas idénticos en ocho tablas críticas. API y Redis permanecen detenidos. Es una observación SQL actual de continuidad, no un nuevo cálculo de readiness del servicio, recorrido E2E, browser ni aceptación humana.",
    )
    heading(doc, "4 Identidad y preservación operativa")
    paragraph(
        doc,
        "La evidencia histórica corresponde a Tenant1 fresh-copper-1765b06813, Enterprise1, BU2, Portfolio3 y Project4. No se reutiliza Tenant4/Project16 de la UAT temporal perdida. La reconstrucción persistente fue autorizada expresamente.",
    )
    paragraph(
        doc,
        "Project14 operacional permanece id14, PRJ-PPMIS, external_keyPRJ-PYP-PPMIS, active, version1 y MD5ef67b78a4d4902d2dc6db3efb017b80b. La consulta18 BEGIN READ ONLY / SELECT / ROLLBACK coincide con la observación16 en Project14,31workspaces,95configuraciones publicadas, políticas, ceroQRA y Alembic0068. No hubo DML.",
    )
    heading(doc, "5 Alcance funcional que se somete a revisión")
    paragraph(
        doc,
        "Los35pasos R01-R35 son índice de revisión del prompt. Se conservan los44Pxx del Business Process,35PROC de RACI y51casosA/B/C/N de UAT. La correspondencia de área es documental y no transforma resultados históricos en aceptación nueva.",
    )
    for route in data["routes"]:
        paragraph(
            doc, route["id"] + ". " + route["validated"] + " " + route["boundary"]
        )
    heading(doc, "6 Gate Evidence FID y Post FID")
    gate_section(doc, data)
    heading(doc, "7 Separación de responsabilidades y seguridad")
    paragraph(
        doc,
        "La matriz ADMIN/USER cubre Workspace Types, Idea/BA, Portfolio, Business Case, FEL/PDRI, Risk, Funding/FIN, QRA, Evidence Policy, Decision Criteria y autoridad Gate. Reutiliza las guías y permisos auditados. La RACI propuesta no asigna personas ni delega autoridad. La firma exige representante real, cargo, organización y base de autorización.",
    )
    heading(doc, "8 Hallazgos y límites de presentación")
    findings_section(doc, data)
    heading(doc, "9 Documentación y vigencia")
    documentation_section(doc, data)
    heading(doc, "10 Evidencia y reglas de captura")
    paragraph(
        doc,
        "review-pack-data.json centraliza 15 áreas, 35 actividades, 11 filas ADMIN/USER, 6 hallazgos, 2 observaciones y 23 referencias. Se verificó que los 23 paths existen y los 6 SHA documentales coinciden. E23 conserva el suplemento SQL actual sin reemplazar la captura inicial. El tracker incorpora 6 hojas para matriz, condiciones, hallazgos, evidencias, aceptación documental y decisión final. Sus campos humanos permanecen vacíos.",
    )
    paragraph(
        doc,
        "La verificación de los XLSX aprobó 21 controles nativos y visuales. El tracker conserva 6 hojas, 15 áreas, 90 filas para aceptación documental, 23 evidencias y 15 espacios de rechazo. Veinte escenarios de fórmula se comprobaron solo en memoria, sin exportar aceptaciones de prueba; se restauró el estado 0 de 15. La copia controlada de trazabilidad conserva 8 hojas y 8 tablas; mantiene intactos valores, estilos y fórmulas A:Z y agrega solo AA:AE para propuestas de área y campos humanos vacíos.",
    )
    paragraph(
        doc,
        "El paquete de revisión integra 11 componentes lógicos mediante cruces explícitos de hojas y secciones. Se preservan seis fuentes byte-idénticas en 01_Fuentes_Baseline16 y se separan los artefactos actuales en 02_Paquete_Revision. El manifiesto registra esa relación sin crear conclusiones técnicas competidoras ni reescribir el Closeout del 16.",
    )
    paragraph(
        doc,
        "Los formularios requieren representante, rol, organización, base de autoridad, fecha, decisión, condiciones, comentarios, evidencia realmente revisada y firma, además de propietario y fecha objetivo cuando aplique. La autorización para proseguir sin firmas habilita preparar el paquete, no aprobarlo en nombre del negocio.",
    )
    heading(doc, "11 Protección de documentos e historia técnica")
    paragraph(
        doc,
        "Los6documentos fuentes conservan susSHA. Final Release Closeout16 no recibe una sección de aceptación inexistente; el prompt exige actualizarlo después de decisiones reales. La copia de trazabilidad agrega referencias sin cambiar PASS/FAIL, IDs o historia técnica. Nuevas funciones o cambios materiales exigirían invalidar la aceptación afectada y acordar repruebas.",
    )
    paragraph(
        doc,
        "Los2Word se derivan del paquete retenido del Closeout16. Conservan geometríaLetter, logo, estilos, relaciones y mobiliario. Solo cambian body y fecha de footer para su propósito nuevo. El inventario de partes verifica preservación byte-idéntica de las demás partes. La revisión visual final se realiza sobre todas sus páginas antes de entrega.",
    )
    heading(doc, "12 Determinación y siguiente paso")
    final_decision(doc, data)
    heading(doc, "13 Registro de fuentes controladas")
    register = {x["id"]: x for x in data["evidence_register"]}
    for item in data["required_documents"]:
        heading(doc, item["id"] + " " + item["name"].replace("&", "and"), 2)
        paragraph(doc, item["revision"] + ". " + register[item["evidence_id"]]["path"])
        paragraph(doc, "SHA256 " + item["sha256"])
    return doc


def patch_output(document, output):
    stream = io.BytesIO()
    document.save(stream)
    with ZipFile(REFERENCE) as source, ZipFile(stream) as generated:
        body = generated.read("word/document.xml")
        footer = source.read("word/footer1.xml").replace(
            b"16 octubre 2026", b"18 octubre 2026"
        )
        assert footer != source.read("word/footer1.xml"), "Footer date slot not found"
        with ZipFile(output, "w") as target:
            for info in source.infolist():
                target.writestr(
                    info,
                    body
                    if info.filename == "word/document.xml"
                    else footer
                    if info.filename == "word/footer1.xml"
                    else source.read(info.filename),
                )
    with ZipFile(REFERENCE) as source, ZipFile(output) as target:
        assert source.namelist() == target.namelist()
        part_rows = [
            {
                "part": name,
                "source_sha256": hashlib.sha256(source.read(name)).hexdigest(),
                "output_sha256": hashlib.sha256(target.read(name)).hexdigest(),
                "changed": source.read(name) != target.read(name),
            }
            for name in source.namelist()
        ]
        assert {r["part"] for r in part_rows if r["changed"]} == ALLOWED_CHANGED_PARTS
        src_xml = etree.fromstring(source.read("word/document.xml"))
        out_xml = etree.fromstring(target.read("word/document.xml"))
        namespace = {
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        }
        assert etree.tostring(
            src_xml.find("w:body/w:sectPr", namespace)
        ) == etree.tostring(out_xml.find("w:body/w:sectPr", namespace))
        assert etree.tostring(src_xml.find("w:body/w:p", namespace)) == etree.tostring(
            out_xml.find("w:body/w:p", namespace)
        )
    return {
        "path": output.relative_to(ROOT).as_posix(),
        "sha256": digest(output),
        "bytes": output.stat().st_size,
        "parts": part_rows,
        "source_parts_preserved": True,
        "visual_review": "PENDING_ROOT_RENDER_AND_ALL_PAGES_REVIEW",
    }


def validate(data, baseline):
    assert digest(REFERENCE) == REFERENCE_SHA
    assert data["date"] == "2026-10-18"
    assert len(data["areas"]) == 15 and len(data["review_steps"]) == 35
    assert (
        data["counts"]["completed"]
        == data["counts"]["accept"]
        == data["counts"]["reject"]
        == 0
    )
    assert data["counts"]["pending"] == 15 and not data["business_release_authorized"]
    assert not any(data["area_blank_fields"].values())
    assert not data["conditions_register"] and not data["rejection_register"]
    assert (
        baseline["code_baseline_unchanged"]
        and baseline["captured_hash_counts"]["all_match"]
    )
    runtime = load(ROOT / data["baseline_supplement_file"])
    assert runtime["selected_matches_delivery_20261016"]
    assert runtime["all_eight_table_counts_and_row_md5_match"]
    assert runtime["transaction_proof"]["transaction_read_only"] == "on"
    refs = {row["id"]: row for row in data["evidence_register"]}
    for item in refs.values():
        assert (ROOT / item["path"]).is_file(), item["path"]
    for document in data["required_documents"]:
        assert (
            digest(ROOT / refs[document["evidence_id"]]["path"]) == document["sha256"]
        )
    return refs


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--authoring-marker-confirmed", action="store_true")
    parser.add_argument("--validate-only", action="store_true")
    args = parser.parse_args()
    data = load(DATA / "review-pack-data.json")
    baseline = load(DATA / "baseline.json")
    validate(data, baseline)
    if args.validate_only:
        print("DATA_AND_SOURCE_VALIDATION_PASS")
        return
    if not args.authoring_marker_confirmed:
        parser.error(
            "Root must confirm the two-DOCX authoring marker before creating documents"
        )
    OUT.mkdir(parents=True, exist_ok=True)
    QA.mkdir(parents=True, exist_ok=True)
    outputs = []
    for name, builder in zip(NAMES, [signoff_document, technical_document]):
        outputs.append(patch_output(builder(data, baseline), OUT / name))
    validate(data, baseline)
    manifest = {
        "schema_version": 1,
        "source": str(REFERENCE),
        "source_sha256": REFERENCE_SHA,
        "data_sha256": digest(DATA / "review-pack-data.json"),
        "baseline_sha256": digest(DATA / "baseline.json"),
        "outputs": outputs,
        "business_decisions": 0,
        "database_writes": 0,
        "source_document_changes": 0,
    }
    (QA / "documents-build-manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(
        json.dumps(
            {
                "outputs": [
                    {k: v for k, v in x.items() if k != "parts"} for x in outputs
                ]
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
