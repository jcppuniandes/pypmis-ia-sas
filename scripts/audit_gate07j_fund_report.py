from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

REPO = Path(__file__).resolve().parents[1]
DOCX = REPO / "outputs" / "technical_reports" / "Informe_Tecnico_Gate_07J_FUND_Funds_Manager_20260828.docx"
OUTPUT = REPO / "outputs" / "technical_reports" / "gate07j_fund_preset_audit.json"


def close(actual: int | None, expected: int, tolerance: int = 1000) -> bool:
    return actual is not None and abs(actual - expected) <= tolerance


def run() -> None:
    doc = Document(DOCX)
    section = doc.sections[0]
    body = [paragraph.text for paragraph in doc.paragraphs]
    cells = [
        paragraph.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    ]
    text = " ".join(body + cells)
    checks = {
        "single_portrait_section": len(doc.sections) == 1 and section.page_height > section.page_width,
        "letter_page_size": all(
            close(item.page_width, Inches(8.5)) and close(item.page_height, Inches(11))
            for item in doc.sections
        ),
        "one_inch_margins": all(
            close(value, Inches(1))
            for item in doc.sections
            for value in (item.left_margin, item.right_margin, item.top_margin, item.bottom_margin)
        ),
        "header_footer_distance": all(
            close(item.header_distance, Inches(0.492)) and close(item.footer_distance, Inches(0.492))
            for item in doc.sections
        ),
        "normal_11pt": close(doc.styles["Normal"].font.size, Pt(11), 100),
        "heading1_16pt": close(doc.styles["Heading 1"].font.size, Pt(16), 100),
        "heading2_13pt": close(doc.styles["Heading 2"].font.size, Pt(13), 100),
        "heading_hierarchy": sum(p.style.name == "Heading 1" for p in doc.paragraphs) == 9
        and sum(p.style.name == "Heading 2" for p in doc.paragraphs) == 15,
        "page_number_footer": all(
            " PAGE " in item.footer.paragraphs[0]._p.xml for item in doc.sections
        ),
        "image_alt_present": all(shape._inline.docPr.get("descr") for shape in doc.inline_shapes),
        "no_fake_bullets": not any(p.text.strip().startswith(("•", "- ")) for p in doc.paragraphs),
        "no_placeholders": not re.search(r"\b(?:todo|lorem ipsum|placeholder|codex-file-citation)\b", text, re.I),
        "final_output_present": "READY_FOR_GOVERNED_FUNDING_EVIDENCE" in text,
        "release_status_present": "GATE07J_FUND_RELEASE_VALIDATED" in text,
        "lifecycle_adr_present": all(
            token in text
            for token in ("ADR 53", "EXHAUSTED", "SUPERSEDED", "RELEASED", "EXPIRED")
        ),
        "browser_three_runs_present": all(token in text for token in ("FLOW 1", "FLOW 2", "FLOW 3", "3/3 FLOW_PASS")),
        "governed_evidence_present": all(
            token in text for token in ("FUNDS_MANAGER", "FUNDING_READINESS", "GOVERNED_SOURCE")
        ),
        "negative_cases_present": all(
            token in text
            for token in (
                "Partial funding",
                "Currency mismatch",
                "Restriction mismatch",
                "Blocking condition",
                "Stale source",
                "Over-allocation",
                "No double count",
            )
        ),
        "budget_boundary_present": "Budget != Funding" in text,
        "baseline_hash_present": "eea3648c6908a267d1e0651828ff4459" in text,
        "validation_counts_present": all(token in text for token in ("126/126", "69/69", "57/57", "3/3")),
        "baseline_match_present": "PERSISTENT_BASELINE_MATCH=True" in text,
        "stop_boundary_present": all(
            token in text
            for token in ("Investment Decision/FID", "Tesorería", "Contabilidad", "Project Initialization", "Activation", "Execution")
        ),
    }
    table_geometry = True
    header_flags = True
    margins = True
    rows_no_split = True
    for table in doc.tables:
        table_geometry &= table._tbl.tblPr.find(qn("w:tblW")) is not None
        table_geometry &= table._tbl.tblPr.find(qn("w:tblInd")) is not None
        header_flags &= table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None
        for row in table.rows:
            rows_no_split &= row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
            for cell in row.cells:
                margins &= cell._tc.get_or_add_tcPr().find(qn("w:tcMar")) is not None
    checks.update(
        {
            "table_geometry": table_geometry,
            "table_header_flags": header_flags,
            "table_cell_margins": margins,
            "table_rows_no_split": rows_no_split,
        }
    )
    report = {
        "preset": "standard_business_brief",
        "first_page_pattern": "memo_masthead",
        "file": str(DOCX),
        "tables": len(doc.tables),
        "inline_images": len(doc.inline_shapes),
        "checks": checks,
        "result": "PASS" if all(checks.values()) else "FAIL",
        "failed": [name for name, passed in checks.items() if not passed],
    }
    OUTPUT.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(report, indent=2, ensure_ascii=False))
    if report["result"] != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    run()
