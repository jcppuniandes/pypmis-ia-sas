"""Preset, structure and content audit for the Gate 07M-FIN Word report."""

from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

REPO = Path(__file__).resolve().parents[1]
DOCX = REPO / "outputs" / "technical_reports" / "Informe_Tecnico_Gate_07M_FIN_Economic_Financial_Model_20260902.docx"
OUTPUT = REPO / "outputs" / "technical_reports" / "gate07m_fin_preset_audit.json"


def close(actual: int | None, expected: int, tolerance: int = 1000) -> bool:
    return actual is not None and abs(actual - expected) <= tolerance


def run() -> None:
    doc = Document(DOCX)
    section = doc.sections[0]
    normal = doc.styles["Normal"]
    heading1 = doc.styles["Heading 1"]
    heading2 = doc.styles["Heading 2"]
    body = " ".join(paragraph.text for paragraph in doc.paragraphs)
    body += " " + " ".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    checks = {
        "single_portrait_section": len(doc.sections) == 1 and section.page_height > section.page_width,
        "letter_page_size": close(section.page_width, Inches(8.5)) and close(section.page_height, Inches(11)),
        "one_inch_margins": all(close(value, Inches(1)) for value in (section.left_margin, section.right_margin, section.top_margin, section.bottom_margin)),
        "header_footer_distance": close(section.header_distance, Inches(0.492)) and close(section.footer_distance, Inches(0.492)),
        "normal_calibri_11pt": normal.font.name == "Calibri" and close(normal.font.size, Pt(11), 100),
        "normal_spacing": close(normal.paragraph_format.space_after, Pt(6), 100) and abs(float(normal.paragraph_format.line_spacing) - 1.10) < 0.001,
        "heading1_tokens": close(heading1.font.size, Pt(16), 100) and str(heading1.font.color.rgb) == "2E74B5",
        "heading2_tokens": close(heading2.font.size, Pt(13), 100) and str(heading2.font.color.rgb) == "2E74B5",
        "heading_hierarchy": sum(p.style.name == "Heading 1" for p in doc.paragraphs) == 9 and sum(p.style.name == "Heading 2" for p in doc.paragraphs) >= 10,
        "header_present": "TECHNICAL REPORT · GATE 07M-FIN" in section.header.paragraphs[0].text,
        "footer_present": "P&Pmis Construction AI" in section.footer.paragraphs[0].text,
        "image_alt_present": all(shape._inline.docPr.get("descr") for shape in doc.inline_shapes),
        "final_output_present": "READY_FOR_GOVERNED_ECONOMIC_FINANCIAL_EVIDENCE" in body,
        "exclusive_state_present": "GATE07M_FIN_RELEASE_VALIDATED" in body,
        "next_present": "UPDATE_GATE07J_EVID_FOR_BUSINESS_CASE_AND_ECONOMIC_FINANCIAL_READINESS" in body,
        "provider_present": "ECONOMIC_FINANCIAL_READINESS_PROVIDER" in body,
        "traceability_81_present": any(cell.text.strip() == "81" for table in doc.tables for row in table.rows for cell in row.cells),
        "no_fake_bullets": not any(p.text.strip().startswith(("•", "- ")) for p in doc.paragraphs),
        "no_placeholder_text": not re.search(r"\b(?:todo|lorem ipsum|placeholder|codex-file-citation)\b", body, flags=re.IGNORECASE),
    }
    table_geometry = row_headers = cell_margins = no_split = True
    for table in doc.tables:
        properties = table._tbl.tblPr
        table_geometry &= properties.find(qn("w:tblW")) is not None and properties.find(qn("w:tblInd")) is not None
        row_headers &= table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None
        for row in table.rows:
            no_split &= row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
            for cell in row.cells:
                cell_margins &= cell._tc.get_or_add_tcPr().find(qn("w:tcMar")) is not None
    checks.update({"table_geometry_complete": table_geometry, "table_header_flags_complete": row_headers, "cell_margins_complete": cell_margins, "rows_prevent_split": no_split})
    report = {
        "preset": "standard_business_brief",
        "first_page_pattern": "memo_masthead",
        "file": str(DOCX),
        "tables": len(doc.tables),
        "inline_images": len(doc.inline_shapes),
        "checks": checks,
        "result": "PASS" if all(checks.values()) else "FAIL",
        "failed": [key for key, value in checks.items() if not value],
    }
    OUTPUT.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(report, indent=2, ensure_ascii=False))
    if report["result"] != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    run()
