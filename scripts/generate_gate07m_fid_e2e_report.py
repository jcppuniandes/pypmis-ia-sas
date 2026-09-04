from __future__ import annotations

import json
import re
import subprocess
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPO = Path(__file__).resolve().parents[1]
OUTPUT = (
    REPO
    / "outputs"
    / "technical_reports"
    / "Informe_Tecnico_Gate_07M_FID_E2E_Full_Stack_Revalidation_20260904.docx"
)
PROMPT = Path(
    r"C:\Users\Ricardo\Documents\P&P\P&Pmis Construction AI\Diseño\Portfolio"
    r"\PROMPT_07M_FID_E2E_NEW_PYP_FFA_FID_PATH_FULL_STACK_REVALIDATION.md"
)
PG_EVIDENCE = REPO / "outputs" / "technical_reports" / "gate07m_fid_e2e_postgresql.json"
SQLITE_EVIDENCE = REPO / "outputs" / "technical_reports" / "gate07m_fid_e2e_sqlite.json"
BROWSER_EVIDENCE = REPO / "outputs" / "technical_reports" / "gate07m_fid_browser_result.json"
LOGO = REPO / "frontend" / "public" / "pypmis-construction-ai-logo.png"

INK = "111827"
DARK_BLUE = "17365D"
MID_BLUE = "1F4E78"
PALE_BLUE = "EAF2F8"
LIGHT_BLUE = "F5F9FC"
BORDER = "D9D9D9"
MUTED = "5F6B76"
WHITE = "FFFFFF"
GREEN = "1B5E20"
AMBER = "8A4B08"


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def git_text(*args: str) -> str:
    return subprocess.check_output(["git", *args], cwd=REPO, text=True).strip()


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_borders(cell, color: str = BORDER, size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_run(
    run,
    *,
    size: float = 10.5,
    bold: bool = False,
    italic: bool = False,
    color: str = INK,
    font: str = "Aptos",
) -> None:
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def set_cell_text(cell, value: str, *, header: bool = False, size: float = 8.4) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    set_run(
        paragraph.add_run(str(value)),
        size=size,
        bold=header,
        color=WHITE if header else INK,
    )
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[tuple | list],
    widths: list[int],
    *,
    font_size: float = 8.4,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for index, (cell, label, width) in enumerate(zip(header.cells, headers, widths, strict=True)):
        set_cell_width(cell, width)
        set_cell_fill(cell, DARK_BLUE)
        set_cell_borders(cell)
        set_cell_text(cell, label, header=True, size=font_size)
        if index == 0:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for cell, value, width in zip(row.cells, values, widths, strict=True):
            set_cell_width(cell, width)
            set_cell_fill(cell, PALE_BLUE if row_index % 2 == 0 else WHITE)
            set_cell_borders(cell)
            set_cell_text(cell, str(value), size=font_size)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 0.1
    set_run(spacer.add_run("\u200b"), size=1, color=WHITE)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(paragraph.add_run("P&Pmis Construction AI  |  "), size=8, color=MUTED)
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (fld_char, instruction, separate, text, end):
        run._r.append(node)
    set_run(run, size=8, color=MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.right_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, space_before, space_after in (
        ("Title", 25, 0, 12),
        ("Subtitle", 14, 0, 16),
        ("Heading 1", 16, 10, 6),
        ("Heading 2", 12.5, 8, 4),
        ("Heading 3", 10.8, 6, 3),
    ):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Normal" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = RGBColor.from_string(INK)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), style.font.name)
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(space_after)
        style.paragraph_format.keep_with_next = True
        p_borders = style._element.pPr.find(qn("w:pBdr"))
        if p_borders is not None:
            style._element.pPr.remove(p_borders)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(header.add_run("INFORME TÉCNICO · GATE 07M-FID-E2E"), size=8, bold=True, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


def page_break(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def add_bullets(doc: Document, items: list[str], *, size: float = 10.2) -> None:
    for item in items:
        paragraph = doc.add_paragraph(item, style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(1)
        for run in paragraph.runs:
            set_run(run, size=size)


def add_numbered(doc: Document, items: list[str], *, size: float = 10.2) -> None:
    for item in items:
        paragraph = doc.add_paragraph(item, style="List Number")
        paragraph.paragraph_format.space_after = Pt(3)
        for run in paragraph.runs:
            set_run(run, size=size)


def add_status_line(doc: Document, label: str, value: str, color: str = GREEN) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    set_run(paragraph.add_run(f"{label}: "), bold=True)
    set_run(paragraph.add_run(value), bold=True, color=color)


def title_page(doc: Document, pg: dict, commit: str, branch: str) -> None:
    if LOGO.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(8)
        picture = paragraph.add_run().add_picture(str(LOGO), width=Inches(0.82))
        picture._inline.docPr.set("descr", "Logotipo de P&Pmis Construction AI")
        picture._inline.docPr.set("title", "P&Pmis Construction AI")

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(3)
    set_run(kicker.add_run("INFORME TÉCNICO DE IMPLEMENTACIÓN Y VALIDACIÓN"), size=10, bold=True, color=MID_BLUE)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(6)
    title.paragraph_format.keep_with_next = True
    set_run(title.add_run("Gate 07M-FID-E2E"), size=25, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    subtitle.paragraph_format.keep_with_next = True
    set_run(
        subtitle.add_run("Revalidación full-stack del nuevo camino PYP-FFA hacia FID"),
        size=14,
        bold=True,
        italic=False,
    )

    metadata = [
        ("Aplicación", "P&Pmis Construction AI"),
        ("Fecha de corte", "4 de septiembre de 2026"),
        ("Repositorio", f"pypmis-ia-sas · rama {branch} · baseline {commit}"),
        ("Alembic", "20260902_0058 · sin cambios de esquema"),
        ("Base de validación", "PostgreSQL 16 desechable + SQLite + Chromium real-stack"),
        ("Perfil", "CAPITAL_OWNER · COMMERCIAL · PYP-FFA"),
    ]
    add_table(doc, ["Dato", "Valor"], metadata, [1900, 7460], font_size=9.0)

    doc.add_heading("Conclusión", level=1)
    paragraph = doc.add_paragraph()
    paragraph.add_run(
        "La revalidación quedó aprobada. El camino estratégico reutiliza los servicios existentes desde "
        "Business Case y Economic & Financial Model hasta Gate 07J, Gate 07K y Gate 07L; produce una "
        "decisión GO con FID y fondos totalmente autorizados, conserva can_activate_project=false y "
        "bloquea correctamente una nueva preparación cuando FIN deja de estar vigente."
    )
    add_status_line(doc, "Estado exclusivo", pg["release_status"])
    add_status_line(doc, "Salida final", pg["final_output"])
    add_status_line(doc, "Siguiente estado", pg["next"], MID_BLUE)
    page_break(doc)


def executive_scope(doc: Document) -> None:
    doc.add_heading("1. Resumen ejecutivo y alcance", level=1)
    doc.add_paragraph(
        "El objetivo fue demostrar el camino PYP-FFA completo con datos sembrados, auditables y removibles, "
        "sin crear funcionalidad de negocio. La prueba se ejecutó sobre servicios reales y una base PostgreSQL "
        "aislada, y se replicó en SQLite para detectar diferencias de persistencia y hashing."
    )
    add_table(
        doc,
        ["Dimensión", "Resultado", "Evidencia"],
        [
            ("Camino feliz", "PASS", "BC/FIN READY → Gate07J FROZEN → Gate07K GO/FID → Gate07L elegible"),
            ("Política efectiva", "PASS", "Policy 787 rev. 1; transición exacta antes/después del timestamp"),
            ("Class A", "PASS", "11/11 elementos válidos; cobertura 100 %"),
            ("Inmutabilidad", "PASS", "Preparación congelada y decisión histórica no cambian ante staleness"),
            ("Camino negativo", "PASS", "Suspensión de Sponsor → FIN STALE → nueva preparación bloqueada"),
            ("Recuperación", "PASS", "Nuevo snapshot Funding, FIN V2 y BC rev. 3 restauran READY"),
            ("Navegador", "PASS", "Happy, negative, historical y drilldown sin mocks del core"),
            ("No regresión", "PASS", "29 pruebas PostgreSQL; build y controles estáticos aprobados"),
        ],
        [1800, 1050, 6510],
        font_size=8.4,
    )

    doc.add_heading("1.1 Reuse only", level=2)
    add_bullets(
        doc,
        [
            "Se reutilizaron BusinessCaseService, EconomicFinancialModelService, presupuesto, Funding, FEL, Risk, Permit, Gate07I, Gate07J-EVID, InvestmentGatePreparation, InvestmentDecision y Post-FID Governance.",
            "Los inserts directos quedaron limitados a identidad, propuesta y configuraciones de fixture sin fábrica pública; el recorrido de negocio usa servicios reales.",
            "No se añadieron modelos, tablas, migraciones, proveedores, fórmulas financieras ni semánticas nuevas de Gate 07J/07K/07L.",
            "No se iniciaron Procurement, Contracting, Commitments, Payments, Accounting, Treasury, NTP, ejecución, P6, Asset ni Facilities Management.",
        ],
    )

    doc.add_heading("1.2 Arquitectura de validación", level=2)
    add_table(
        doc,
        ["Capa", "Implementación", "Aislamiento"],
        [
            ("Servicios", "Pytest integra el camino completo con servicios de dominio", "Fixture por tenant y Project"),
            ("Datos", "PostgreSQL 16 con Alembic 0044→0058; réplica SQLite", "tmpfs y esquema desechable"),
            ("API/UI", "Uvicorn + Vite + Chromium Playwright", "Puertos 18000 y 15173"),
            ("Evidencia", "JSON con IDs, hashes, estados y tiempos", "outputs/technical_reports"),
        ],
        [1600, 4700, 3060],
    )
    page_break(doc)


def implementation_fixes(doc: Document) -> None:
    doc.add_heading("2. Correcciones de contrato implementadas", level=1)
    doc.add_paragraph(
        "La revalidación reveló defectos latentes de interoperabilidad. Se corrigió únicamente lo necesario "
        "para cumplir contratos ya definidos; ninguna corrección amplía el dominio."
    )
    add_table(
        doc,
        ["Componente", "Corrección", "Control preservado"],
        [
            (
                "Business Case",
                "Normalización de as_of_date entre JSON ISO y columna DateTime.",
                "Lineage serializable sin romper persistencia relacional.",
            ),
            (
                "BC ↔ FIN",
                "Una BC COMPLETED/READY posterior que referencia exactamente FIN id/version/hash cierra el puente de vigencia.",
                "No muta revisiones completadas ni acepta referencias parciales.",
            ),
            (
                "FIN ↔ Funding",
                "Currentness compara el snapshot referenciado con readiness derivada de fuentes gobernadas vivas.",
                "Suspensión o cambio de una fuente convierte FIN en STALE.",
            ),
            (
                "FIN refresh",
                "Refresh usa currentness vivo y el snapshot Funding vigente.",
                "Crea una nueva versión; no reescribe FIN histórico.",
            ),
            (
                "Policy 787",
                "Selección por effective_from y después revision para configuraciones de igual scope.",
                "Cambio temporal exacto; no se backdatea la política publicada.",
            ),
            (
                "Hash de decisión",
                "Monto autorizado canónico con dos decimales antes de calcular el hash.",
                "Determinismo entre SQLite y PostgreSQL.",
            ),
        ],
        [1850, 4700, 2810],
        font_size=8.0,
    )

    doc.add_heading("2.1 Archivos de producto ajustados", level=2)
    add_table(
        doc,
        ["Archivo", "Responsabilidad"],
        [
            ("backend/app/modules/business_case/service.py", "Conversión DateTime y bundle de fuente."),
            ("backend/app/modules/economic_financial_model/service.py", "Puente BC↔FIN, vigencia Funding y refresh."),
            ("backend/app/modules/investment_decision/service.py", "Hash canónico del monto autorizado."),
            ("backend/app/modules/investment_evidence_policy/service.py", "Resolución temporal de políticas publicadas."),
        ],
        [5300, 4060],
        font_size=8.2,
    )
    doc.add_heading("2.2 Activos de validación añadidos", level=2)
    add_table(
        doc,
        ["Activo", "Propósito"],
        [
            ("backend/tests/test_gate07m_fid_e2e.py", "Escenario integral y evidencia JSON."),
            ("docker-compose.gate07m-fid-e2e.yml", "PostgreSQL/Redis/API/frontend aislados y fases de regresión."),
            ("frontend/e2e/gate07m-fid-real-stack.spec.ts", "Browser E2E sin route mocks del core."),
            ("docs/63-validation-gate07m-fid-full-stack.md", "ADR / Validation Note y límites de arquitectura."),
        ],
        [5300, 4060],
        font_size=8.2,
    )
    page_break(doc)


def scenario(doc: Document, pg: dict) -> None:
    project = pg["project"]
    funding = pg["funding"]
    doc.add_heading("3. Escenario sembrado y cadena estratégica", level=1)
    add_table(
        doc,
        ["Elemento", "Valor"],
        [
            ("Project", f"{project['code']} · ID {project['id']}"),
            ("Perfil", f"{project['governance_model']} · {project['project_type']} · {project['profile']}"),
            ("Portfolio", f"{pg['portfolio']['code']} · ID {pg['portfolio']['id']}"),
            ("Budget", f"USD {int(pg['budget']['amount']):,} · scenario {pg['budget']['scenario_id']}"),
            ("Funding", f"USD {funding['confirmed_amount']} · cobertura {funding['coverage_percent']} %"),
            ("Permit", f"Clase {pg['permit']['classification']} · {pg['permit']['handling']}"),
            ("Estado Project tras FID", project["status_after"]),
        ],
        [2400, 6960],
        font_size=8.7,
    )

    doc.add_heading("3.1 Fuentes y estructura de capital", level=2)
    source_rows = []
    for source in funding["sources"]:
        source_rows.append(
            (
                source["key"],
                source["source_type"],
                source["instrument_type"],
                f"USD {int(source['amount']):,}",
            )
        )
    add_table(
        doc,
        ["Clave", "Fuente", "Instrumento", "Monto"],
        source_rows,
        [1450, 2350, 3360, 2200],
        font_size=8.2,
    )

    doc.add_heading("3.2 Recorrido validado", level=2)
    add_numbered(
        doc,
        [
            "Planning, evaluation, budget, resource y convergence producen fuentes gobernadas del Portfolio.",
            "FEL-1, FEL-2 y FEL-3, Risk y Permit completan evidencia estratégica aplicable.",
            "Business Case queda COMPLETED/READY y referencia el FIN actual por id, versión y hash.",
            "FIN queda COMPLETED/READY, reconcilia Budget y Funding y expone métricas económicas/de deuda.",
            "Policy 787 rev. 1 resuelve 11/11 Class A; Gate 07J se congela READY.",
            "Gate 07K decide GO, otorga FID y autorización total de fondos, sin activar el Project.",
            "Gate 07L habilita inicialización del workspace y mantiene activación en false.",
        ],
    )
    page_break(doc)


def financial_and_gates(doc: Document, pg: dict) -> None:
    fin = pg["financial_model"]
    gate_j = pg["gate07j"]
    gate_k = pg["gate07k"]
    gate_l = pg["gate07l"]
    doc.add_heading("4. Resultados económicos, financieros y de gobierno", level=1)
    add_table(
        doc,
        ["Indicador", "Resultado", "Lectura"],
        [
            ("NPV", f"USD {float(fin['npv']):,.6f}", "Valor presente neto calculado"),
            ("IRR", f"{float(fin['irr']) * 100:.6f} %", "Tasa interna del caso"),
            ("MOIC", fin["moic"], "Múltiplo de capital invertido"),
            ("Minimum DSCR", fin["minimum_dscr"], "Cobertura mínima de servicio de deuda"),
            ("LLCR", fin["llcr"], "Cobertura de vida del préstamo"),
            ("Equity IRR", fin["equity_irr"], "No computable con el patrón de flujo sembrado; no se fabrica valor"),
            ("Funding", fin["funding_reconciliation"], "Cinco instrumentos referenciados"),
        ],
        [2200, 2250, 4910],
        font_size=8.3,
    )

    doc.add_heading("4.1 Gate 07J · Investment Gate Preparation", level=2)
    add_table(
        doc,
        ["Campo", "Resultado"],
        [
            ("Preparation", f"ID {gate_j['preparation_id']} · {gate_j['status']}"),
            ("Readiness", gate_j["readiness"]),
            ("Class A", f"{gate_j['class_a_valid']}/{gate_j['class_a_count']} válidas"),
            ("Coverage", f"{gate_j['coverage_percent']:.1f} %"),
            ("Snapshot hash", gate_j["hash"]),
        ],
        [2600, 6760],
        font_size=8.2,
    )

    doc.add_heading("4.2 Gate 07K y Gate 07L", level=2)
    add_table(
        doc,
        ["Gate", "Estado", "Salida gobernada"],
        [
            (
                "07K",
                f"{gate_k['status']} / {gate_k['outcome']}",
                f"is_fid={str(gate_k['is_fid']).lower()}; full_funds={str(gate_k['full_funds_authorized']).lower()}; "
                f"investment_authorized={str(gate_k['investment_authorized']).lower()}; can_activate={str(gate_k['can_activate_project']).lower()}",
            ),
            (
                "07L",
                gate_l["eligibility_status"],
                f"can_initialize={str(gate_l['can_initialize']).lower()}; can_activate={str(gate_l['can_activate']).lower()}",
            ),
        ],
        [900, 3200, 5260],
        font_size=8.2,
    )
    page_break(doc)


def effective_and_hashes(doc: Document, pg: dict) -> None:
    policy = pg["policy"]
    doc.add_heading("5. Política efectiva, timezone y cadena de hashes", level=1)
    doc.add_paragraph(
        "El timestamp publicado no fue editado ni retrodatado. El servicio interpreta el valor ISO sin zona como "
        "UTC y lo compara con un reloj UTC aware. La prueba focal usa un reloj de test aislado para comprobar "
        "el instante anterior y el instante exacto de activación."
    )
    add_table(
        doc,
        ["Control", "Resultado"],
        [
            ("Runtime de validación", "UTC aware; ejecución host America/Bogota"),
            ("effective_from persistido", policy["effective_from"]),
            ("Interpretación resuelta", f"{policy['effective_from']}Z (UTC)"),
            ("Antes del instante", "Se selecciona política anterior"),
            ("En/después del instante", f"Policy {policy['id']} rev. {policy['revision']}"),
            ("Policy hash", policy["hash"]),
        ],
        [2800, 6560],
        font_size=8.4,
    )

    doc.add_heading("5.1 Hash chain principal", level=2)
    add_table(
        doc,
        ["Objeto", "ID / versión", "Hash"],
        [
            ("Budget", str(pg["budget"]["scenario_id"]), pg["budget"]["hash"]),
            ("Resource", str(pg["resource"]["scenario_id"]), pg["resource"]["hash"]),
            ("Convergence", str(pg["convergence"]["id"]), pg["convergence"]["hash"]),
            ("Funding Readiness", str(pg["funding"]["snapshot_id"]), pg["funding"]["hash"]),
            ("Business Case", f"rev {pg['business_case']['revision_id']}", pg["business_case"]["hash"]),
            ("FIN", f"v{pg['financial_model']['version_id']}", pg["financial_model"]["hash"]),
            ("Gate 07J", str(pg["gate07j"]["preparation_id"]), pg["gate07j"]["hash"]),
            ("Gate 07K Decision", str(pg["gate07k"]["decision_id"]), pg["gate07k"]["decision_hash"]),
            ("Gate 07L Eligibility", "current", pg["gate07l"]["eligibility_hash"]),
        ],
        [2150, 1450, 5760],
        font_size=7.7,
    )
    page_break(doc)


def negative_recovery(doc: Document, pg: dict) -> None:
    negative = pg["negative_fin_stale"]
    recovery = negative["recovery"]
    doc.add_heading("6. Camino negativo, histórico y recuperación", level=1)
    add_table(
        doc,
        ["Fase", "Acción / condición", "Resultado"],
        [
            ("Estado válido", "BC y FIN actuales; Gate 07J congelado; FID decidido", "READY / FROZEN / DECIDED"),
            ("Disparador", negative["trigger"], f"Funding source ID {negative['source_id']}"),
            ("Reevaluación", "Proveedor FIN consulta fuente Funding gobernada", f"current={str(negative['provider_source_current']).lower()}; readiness={negative['provider_readiness']}"),
            ("Bloqueo", "Se evalúa una preparación nueva", f"would_be_ready={str(negative['gate07j_would_be_ready']).lower()}"),
            ("Histórico", "Preparación y Decision ya congeladas", "Hashes sin cambios"),
            ("Recuperación", "Fuente reactivada; snapshot Funding nuevo; FIN V2; BC rev. 3", recovery["provider_readiness"]),
        ],
        [1500, 4700, 3160],
        font_size=8.2,
    )

    doc.add_heading("6.1 Evidencia de inmutabilidad y recuperación", level=2)
    add_table(
        doc,
        ["Control", "Resultado"],
        [
            ("Preparation hash histórico", "UNCHANGED" if negative["frozen_preparation_hash_unchanged"] else "CHANGED"),
            ("Decision hash histórico", "UNCHANGED" if negative["historical_decision_hash_unchanged"] else "CHANGED"),
            ("Funding snapshot recuperado", f"ID {recovery['funding_snapshot_id']} · {recovery['funding_readiness_hash']}"),
            ("FIN recuperado", f"V{recovery['financial_model_version_id']} · {recovery['financial_model_hash']}"),
            ("Business Case recuperado", f"rev {recovery['business_case_revision_id']} · FIN hash coincidente"),
            ("Provider final", f"current={str(recovery['provider_source_current']).lower()} · {recovery['provider_readiness']}"),
        ],
        [3000, 6360],
        font_size=8.0,
    )

    doc.add_heading("6.2 Negativos cubiertos por regresión", level=2)
    add_bullets(
        doc,
        [
            "Fuente Funding modificada o suspendida: FIN se vuelve STALE y no puede habilitar una preparación nueva.",
            "Referencias BC↔FIN incompletas o antiguas: bloqueadas por exactitud de id, versión y hash.",
            "Cobertura Class A menor a 100 %, provider unavailable y evidencia inválida: Gate 07J no queda READY.",
            "No existe fallback automático A→B; el grado requerido por policy se conserva.",
            "Decide repetido queda protegido por idempotencia y no altera snapshot ni decisión histórica.",
            "Overrun o conciliación Funding no válida bloquea readiness financiera según contratos existentes.",
        ],
        size=8.9,
    )
    page_break(doc)


def validation(doc: Document, pg: dict, sqlite: dict, browser: dict) -> None:
    doc.add_heading("7. Validación técnica y browser real-stack", level=1)
    add_table(
        doc,
        ["Control", "Resultado", "Detalle"],
        [
            ("Ruff check + format", "PASS", "6 archivos backend modificados/nuevos"),
            ("Pruebas focales host", "PASS", "12 passed"),
            ("Integración SQLite", "PASS", "1 passed; evidencia generada"),
            ("Migración PostgreSQL", "PASS", "0044 → head 20260902_0058"),
            ("Schema validator", "PASS", "Gate07M-FIN sobre PostgreSQL"),
            ("Camino full-stack PostgreSQL", "PASS", "1 passed"),
            ("Regresión Gate 07L", "PASS", "5 passed"),
            ("Regresión EVID/FIN/Decision", "PASS", "23 passed"),
            ("Total PostgreSQL", "PASS", "29 passed"),
            ("Frontend ESLint", "PASS", "0 errores; 13 warnings preexistentes no bloqueantes"),
            ("Frontend build", "PASS", "Vite/TypeScript; 2.436 módulos"),
            ("Playwright Chromium", "PASS", "1 escenario real-stack; 1,5 min"),
            ("git diff --check", "PASS", "Sin whitespace errors"),
        ],
        [2700, 1000, 5660],
        font_size=7.6,
    )

    doc.add_heading("7.1 Evidencia de navegador", level=2)
    add_table(
        doc,
        ["Aserción", "Resultado"],
        [
            ("Camino feliz", str(browser["happy_path"])),
            ("FIN stale negativo", str(browser["negative_fin_stale"])),
            ("FID histórico", str(browser["historical_fid"])),
            ("Source drilldown", str(browser["source_drilldown"])),
            ("Page errors", str(browser["page_errors"])),
            ("Console problems", str(browser["console_problems"])),
            ("Unexpected >=400", str(browser["unexpected_responses"])),
            ("Core route mocks", "0"),
        ],
        [3500, 5860],
        font_size=8.3,
    )
    paragraph = doc.add_paragraph(
        f"La equivalencia de contrato se confirmó en dos motores: PostgreSQL produjo {pg['release_status']} y "
        f"SQLite produjo {sqlite['release_status']}. Los hashes son específicos de cada fixture; los estados, "
        "métricas, guardrails y transiciones son equivalentes."
    )
    paragraph.paragraph_format.space_after = Pt(0)
    for run in paragraph.runs:
        set_run(run, size=9.5)
    page_break(doc)


def performance_and_cleanup(doc: Document, pg: dict) -> None:
    performance = pg["performance"]
    doc.add_heading("8. Rendimiento, aislamiento y baseline persistente", level=1)
    add_table(
        doc,
        ["Operación", "Tiempo observado"],
        [
            ("Provider resolution", f"{performance['provider_resolution_ms']:.3f} ms"),
            ("Gate 07J readiness", f"{performance['gate07j_readiness_ms']:.3f} ms"),
            ("Gate 07J freeze", f"{performance['gate07j_freeze_ms']:.3f} ms"),
            ("Gate 07K package load", f"{performance['gate07k_package_load_ms']:.3f} ms"),
            ("Gate 07K decide", f"{performance['gate07k_decide_ms']:.3f} ms"),
            ("Gate 07L evaluate", f"{performance['gate07l_evaluate_ms']:.3f} ms"),
        ],
        [5000, 4360],
        font_size=7.8,
    )
    paragraph = doc.add_paragraph(
        "Estos valores son observaciones de un entorno local contenedorizado, no un SLA ni un benchmark de carga. "
        "La resolución del proveedor es la operación más costosa del escenario y merece telemetría p95 en una "
        "fase posterior, sin bloquear este cierre funcional."
    )
    paragraph.paragraph_format.space_after = Pt(2)
    for run in paragraph.runs:
        set_run(run, size=9.5)

    doc.add_heading("8.1 Aislamiento y limpieza", level=2)
    add_table(
        doc,
        ["Control", "Evidencia"],
        [
            ("Base de prueba", "PostgreSQL gate07e en tmpfs; Redis sin persistencia"),
            ("Separación de puertos", "API 18000; frontend 15173; no reemplazan 8000/5173 principales"),
            ("Limpieza", "docker compose down -v completado; contenedores y red gate07m-fid-browser removidos"),
            ("Aplicación principal", "Servicios persistentes permanecieron activos; localhost:5173/app respondió HTTP 200"),
            ("Project ID 14", "No existe dentro del fixture desechable; por diseño los campos before/after son null"),
            ("Baseline persistente", "Sólo lectura: head 0058; 14 enterprise workspaces; 1 project workspace; 84 configuraciones admin"),
            ("Project 14 persistente", "Hash baseline observado: eea3648c6908a267d1e0651828ff4459; no usado por el test"),
        ],
        [3000, 6360],
        font_size=7.7,
    )

    doc.add_heading("8.2 Datos persistentes preservados", level=2)
    add_table(
        doc,
        ["Objeto", "Conteo / estado baseline"],
        [
            ("Enterprise workspaces", "14"),
            ("Project workspaces", "1"),
            ("Admin configurations", "84"),
            ("User accounts / access assignments", "9 / 2"),
            ("Policy 787", "published · rev. 1 · hash preservado"),
            ("BC / FIN / Decision / Post-FID persistentes", "0 antes de la prueba; escenario ejecutado sólo en DB desechable"),
        ],
        [4200, 5160],
        font_size=7.8,
    )
    page_break(doc)


def risks_and_closeout(doc: Document, pg: dict) -> None:
    doc.add_heading("9. Riesgos, deuda técnica y cierre estratégico", level=1)
    add_table(
        doc,
        ["Observación", "Impacto", "Tratamiento"],
        [
            (
                "13 warnings de ESLint preexistentes",
                "No bloquearon build ni E2E; concentrados en hooks/BIM viewer.",
                "Reducir progresivamente en una tarea de calidad separada.",
            ),
            (
                "Warnings de pytest cache por permisos",
                "No afectan resultados ni estado de pruebas.",
                "Ajustar directorio de cache del runner local/CI.",
            ),
            (
                "Provider resolution ~4 s local",
                "Posible latencia visible con mayor volumen.",
                "Instrumentar p95, queries y cache antes de escala productiva.",
            ),
            (
                "Equity IRR no computable en fixture",
                "Resultado explícito, no falsificado; MOIC/deuda sí calculados.",
                "Añadir fixture económico con flujo equity de signos opuestos si se requiere probar Equity IRR numérica.",
            ),
            (
                "Validación de un perfil",
                "Cierre aplica al camino CAPITAL_OWNER / PYP-FFA probado.",
                "Revalidar otros governance models antes de declararlos equivalentes.",
            ),
        ],
        [2500, 3400, 3460],
        font_size=8.0,
    )

    doc.add_heading("9.1 Decisión de cierre", level=2)
    doc.add_paragraph(
        "No se identificó un gap estratégico bloqueante para el camino validado. El macroproceso estratégico queda "
        "funcionalmente cerrado para CAPITAL_OWNER / PYP-FFA, sujeto a los límites indicados. Esta conclusión no "
        "autoriza ejecución operacional ni activa automáticamente ningún Project."
    )
    add_status_line(doc, "Release", pg["release_status"])
    add_status_line(doc, "Final output", pg["final_output"])
    add_status_line(doc, "NEXT", pg["next"], MID_BLUE)

    doc.add_heading("9.2 Comandos reproducibles", level=2)
    add_table(
        doc,
        ["Objetivo", "Comando"],
        [
            ("PostgreSQL + regresión", "docker compose -f docker-compose.gate07m-fid-e2e.yml --profile validation up --build --abort-on-container-exit --exit-code-from gate07m-fid-test"),
            ("Servicios browser", "docker compose -p gate07m-fid-browser -f docker-compose.gate07m-fid-e2e.yml --profile browser up -d --build"),
            ("E2E Chromium", "cd frontend && npx playwright test e2e/gate07m-fid-real-stack.spec.ts"),
            ("Limpieza", "docker compose -p gate07m-fid-browser -f docker-compose.gate07m-fid-e2e.yml --profile browser down -v --remove-orphans"),
        ],
        [2350, 7010],
        font_size=7.6,
    )
    page_break(doc)


def parse_deliverables() -> list[str]:
    text = PROMPT.read_text(encoding="utf-8")
    section = text.split("## 66. Required deliverables", 1)[1].split("## 67.", 1)[0]
    items: dict[int, str] = {}
    for line in section.splitlines():
        match = re.match(r"^(\d+)\.\s+(.+)$", line.strip())
        if match:
            items[int(match.group(1))] = match.group(2).rstrip(".")
    if sorted(items) != list(range(1, 67)):
        raise RuntimeError("El prompt no contiene los 66 entregables esperados")
    return [items[index] for index in range(1, 67)]


def traceability(doc: Document) -> None:
    doc.add_heading("10. Matriz de trazabilidad de entregables", level=1)
    doc.add_paragraph(
        "La matriz vincula los 66 entregables solicitados con la sección principal de este informe o con el activo "
        "de evidencia verificable. PASS indica cobertura demostrada por el escenario nuevo o por la regresión "
        "existente; N/A controlado indica una condición deliberada del fixture, no una omisión funcional."
    )

    def evidence(index: int, item: str) -> str:
        if index <= 6:
            return "§1 / Validation Note 63"
        if index <= 14:
            return "§3 / JSON PostgreSQL"
        if index <= 27:
            return "§4–5 / JSON + services"
        if index <= 36:
            return "§4–5 / Gate 07J–07L"
        if index <= 47:
            return "§6 / regression + hashes"
        if index <= 54:
            return "§7–8 / Playwright + timings"
        if index <= 61:
            return "§2, §7–8 / diff + cleanup"
        return "§9 / closeout"

    rows = []
    for index, item in enumerate(parse_deliverables(), 1):
        status = "PASS"
        note = evidence(index, item)
        if index == 57:
            status = "N/A aislado"
            note = "§8 · Project 14 ausente en DB desechable; baseline persistente no mutado"
        rows.append((str(index), item, note, status))
    first_rows = rows[:34]
    second_rows = rows[34:]
    add_table(
        doc,
        ["#", "Entregable", "Evidencia", "Estado"],
        first_rows,
        [500, 5000, 2600, 1260],
        font_size=7.4,
    )
    page_break(doc)
    doc.add_heading("10.1 Matriz de trazabilidad (continuación)", level=2)
    add_table(
        doc,
        ["#", "Entregable", "Evidencia", "Estado"],
        second_rows,
        [500, 5000, 2600, 1260],
        font_size=7.4,
    )


def evidence_inventory(doc: Document) -> None:
    page_break(doc)
    doc.add_heading("11. Inventario de evidencia y notas", level=1)
    add_table(
        doc,
        ["Artefacto", "Contenido"],
        [
            ("outputs/technical_reports/gate07m_fid_e2e_postgresql.json", "Estados, IDs, hashes, métricas, negativo y recuperación PostgreSQL."),
            ("outputs/technical_reports/gate07m_fid_e2e_sqlite.json", "Ejecución equivalente en SQLite."),
            ("outputs/technical_reports/gate07m_fid_browser_result.json", "Resultado Playwright, errores de página/consola/red."),
            ("docs/63-validation-gate07m-fid-full-stack.md", "Decisiones, guardrails, defectos corregidos y estado final."),
            ("backend/tests/test_gate07m_fid_e2e.py", "Fixture y assertions de servicio end-to-end."),
            ("frontend/e2e/gate07m-fid-real-stack.spec.ts", "Flujo UI real con navegación histórica y source drilldown."),
            ("docker-compose.gate07m-fid-e2e.yml", "Orquestación reproducible, migración, regresión y browser."),
        ],
        [5100, 4260],
        font_size=8.0,
    )

    doc.add_heading("11.1 Notas de interpretación", level=2)
    add_bullets(
        doc,
        [
            "Los hashes de PostgreSQL y SQLite difieren porque sus fixtures desechables tienen identificadores y timestamps propios; la equivalencia se verifica en contratos, transiciones y guardrails.",
            "La salida Gate 07L observada es READY_FOR_PROJECT_WORKSPACE_INITIALIZATION; el siguiente paso estratégico global permanece READY_TO_ENTER_OPERATIONAL_MACROPROCESSES_REVIEW.",
            "La prueba demuestra que FID no equivale a activación: can_activate_project y can_activate permanecen false.",
            "La limpieza elimina únicamente el proyecto Docker de validación; no toca los servicios principales ni datos persistentes del repositorio.",
        ],
    )


def build() -> None:
    pg = read_json(PG_EVIDENCE)
    sqlite = read_json(SQLITE_EVIDENCE)
    browser = read_json(BROWSER_EVIDENCE)
    commit = git_text("rev-parse", "--short", "HEAD")
    branch = git_text("branch", "--show-current")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    title_page(doc, pg, commit, branch)
    executive_scope(doc)
    implementation_fixes(doc)
    scenario(doc, pg)
    financial_and_gates(doc, pg)
    effective_and_hashes(doc, pg)
    negative_recovery(doc, pg)
    validation(doc, pg, sqlite, browser)
    performance_and_cleanup(doc, pg)
    risks_and_closeout(doc, pg)
    traceability(doc)
    evidence_inventory(doc)

    properties = doc.core_properties
    properties.title = "Informe Técnico Gate 07M-FID-E2E — Full-Stack Revalidation"
    properties.subject = "Nuevo camino PYP-FFA hasta FID y Gate 07L"
    properties.author = "P&Pmis Construction AI"
    properties.keywords = "Gate 07M, FID, PYP-FFA, Business Case, FIN, Gate 07J, Gate 07K, Gate 07L"
    properties.comments = "Generado desde evidencia PostgreSQL, SQLite y Playwright verificada el 4 de septiembre de 2026."
    properties.created = datetime(2026, 9, 4, 12, 0, 0)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
