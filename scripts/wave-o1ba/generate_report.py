"""Generate the requested report from verified Scope evidence and retained styles."""

import hashlib
import importlib.util
import io
import json
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
REFERENCE = (
    ROOT
    / "outputs/wave-o1a-20261019/P&Pmis_Wave_O1A_Operational_Context_and_Period_Technical_Report.docx"
)
REFERENCE_SHA = "e874daacd1fbddca064effd9254941b2b6138e04a881a953835cf8fd9154535f"
OUT = ROOT / "outputs/wave-o1ba-20261020"
QA = ROOT / "outputs/.qa/wave-o1ba-20261020"
DATA = ROOT / "docs/validation/wave-o1ba-20261020/package-data.json"
spec = importlib.util.spec_from_file_location(
    "scope_word_components", ROOT / "scripts/wave-o1a/generate_report.py"
)
components = importlib.util.module_from_spec(spec)
spec.loader.exec_module(components)
components.base.REFERENCE = REFERENCE
p, h, table = components.p, components.h, components.table


def section(document, number, title, *paragraphs):
    h(document, f"{number} {title}")
    for text in paragraphs:
        p(document, text)


def build(data):
    document = Document(REFERENCE)
    logo = deepcopy(document.paragraphs[0]._p)
    section_properties = deepcopy(document._element.body.sectPr)
    for node in list(document._element.body):
        document._element.body.remove(node)
    document._element.body.append(logo)
    document._element.body.append(section_properties)
    p(document, "Wave O1B A\nFundación de Scope y descomposición de ejecución", "Title")
    p(
        document,
        f"P&Pmis Construction AI\nInforme técnico de implementación\n{data['issue_date']}",
        "Subtitle",
    )
    conclusion = (
        "La validación técnica quedó cerrada en el entorno aislado."
        if data["release_validated"]
        else "La liberación técnica permanece pendiente de las comprobaciones indicadas."
    )
    p(
        document,
        "Implementamos Scope como origen operacional canónico, con fuentes IFC y Excel gobernadas, Scope Items, estrategias de ejecución y Plan Items. "
        + conclusion
        + " La base local habitual no fue migrada y no se registró aceptación de negocio.",
    )
    section(
        document,
        1,
        "Resumen ejecutivo",
        "Los cuatro submódulos de Scope Manager comparten identidad canónica, autorización O1A, perfiles publicados, permisos granulares, concurrencia e historial. La importación conserva la fuente; la aprobación convierte una revisión de Scope en un baseline inmutable. La estrategia y Plan Items descomponen la ejecución sin crear actividades, costos ni avance aprobado.",
    )
    section(
        document,
        2,
        "Línea base",
        "El incremento parte de O1A validada sobre main y del commit c75a7da3db3bdcf326b6bed7834182ee126da921, con trabajo previo sin commit preservado. El Proyecto canónico continúa en enterprise_workspaces de tipo PROJECT; projects.id permanece como referencia heredada independiente.",
        "La validación usa docker-compose.wave-o1ba.yml, API en el puerto 8852 y frontend en el 5512, con bases PostgreSQL separadas para aceptación, migración y regresión. Los servicios y datos habituales no se usaron como fixtures. El estado estratégico sigue pendiente de firmas de negocio.",
    )
    section(
        document,
        3,
        "Hallazgos del repositorio",
        "Reutilizamos BimIfcModelViewer, web-ifc y Three.js mediante un proveedor de fuente de solo lectura. Conservamos las cantidades, APU y servicios geometry existentes. IFCOpenShell y los lectores XLSX existentes sustentan el análisis de fuentes; los nuevos controles de identidad, unidad y validación pertenecen al dominio Scope.",
        "No existía un maestro compatible de Scope Items o Plan Items: QuantityTakeoffRun y BimModel representan fuentes legacy. Se incorporó un módulo específico sin ampliar App.tsx, los modelos globales ni los esquemas globales de negocio. El inventario source-capability-reuse.md distingue reutilización, extensión y contratos reservados.",
    )
    section(
        document,
        4,
        "Arquitectura de Scope",
        "Scope Items es la autoridad de qué entregar, su naturaleza y cantidad controlada. El modelo original conserva geometría y propiedades; una estrategia define cómo ejecutar; Plan Items asigna cantidades para ejecución y control. Los vínculos BIM a Scope y Scope a Plan son M:N explícitos, no equivalencias uno a uno.",
    )
    section(
        document,
        5,
        "Modelo 3D",
        "Cada versión identifica el modelo, revisión, nombre de archivo, hash, esquema, unidades, autor, fecha, estado, versión anterior y validación. Los objetos guardan GUID, ExpressID local, clase, propiedades, clasificación, ubicación, materiales, cantidades y referencias geométricas. El código del modelo es su identidad lógica dentro del proyecto; el ID de versión identifica el archivo inmutable.",
    )
    section(
        document,
        6,
        "Importación IFC",
        "El comando carga y analiza la fuente en staging antes de revisión y aceptación independiente. Valida GUID duplicado o ausente, unidades, entidades no soportadas, extracción de cantidades y diferencias respecto a la versión actual. El diff conserva identidad GUID y detecta pérdidas, crecimiento, cambios y roturas de mapeo.",
        "Los bytes originales se guardan con SHA256 y no se reescriben al aceptar una versión ni al crear Plan Items. Una versión rechazada conserva la evidencia y no se convierte en actual. El índice de versión actual y los controles de concurrencia impiden promociones competidoras silenciosas.",
    )
    section(
        document,
        7,
        "Scope Items",
        "El registro canónico incluye código, descripción, jerarquía propia, disciplina, área, zona, sistema, ubicación, fuente, revisión, cantidades, unidad, vigencia y versión. Las jerarquías se validan sin equipararse con WBS, PBS o CBS. Una edición de borrador conserva el linaje y la unidad de origen; cambiar una fuente o un mapeo exige sucesor.",
    )
    section(
        document,
        8,
        "Naturaleza del alcance",
        "INSTALL incorpora un entregable; REMOVE retira un componente; TEMPORARY representa trabajo temporal; MAINTENANCE interviene un componente existente. La naturaleza es un atributo gobernado y queda en revisiones y baselines, disponible para futura planificación y 4D sin activar esas funciones.",
    )
    section(
        document,
        9,
        "Importación Excel",
        "XLSX es una fuente válida sin IFC. La importación verifica plantilla, encabezados y mapeo, unidades, naturaleza, cantidades finitas, códigos repetidos, filas repetidas, padres y ciclos. Conserva archivo, hash, hoja y fila física, valores originales y campos futuros de mapeo. Fórmulas, vínculos externos y contenido activo no se convierten silenciosamente en cantidades.",
        "Staging y preview permiten revisar errores antes de generar candidatos. El comando de candidatos es idempotente y no aprueba Scope. La cantidad fuente permanece separada de la cantidad controlada aun cuando un borrador cambia la cantidad propuesta.",
    )
    section(
        document,
        10,
        "Revisiones de Scope",
        "La preparación congela candidatos, fuentes, altas, retiros, modificaciones, cambios de cantidad y naturaleza e impactos de mapeo. Los comandos DRAFT, SUBMITTED, REVIEWED y APPROVED son explícitos. Si cambia un candidato o el baseline actual después de preparar, la aprobación exige volver a preparar.",
    )
    section(
        document,
        11,
        "Baselines aprobados",
        "La aprobación autorizada crea una instantánea inmutable y conserva la relación con el baseline original y el anterior. No altera la definición histórica en sitio. La fuente Approved Change consume una ChangeRequest existente aprobada y del agregado legacy vinculado; este incremento no implementa otro workflow de Change ni una nueva autoridad de aprobación.",
    )
    section(
        document,
        12,
        "Estrategia y ruta de ejecución",
        "Las estrategias admiten rutas ordenadas, paralelas, condicionales y por área, sistema, fase o paquete. Guardan reglas de segmentación, estrategia temporal, de retiro, mantenimiento, paquetes y frentes. La publicación es independiente del autor y del último editor material, con sucesores en vez de alterar lo publicado.",
        "La segmentación controlada crea Plan Items candidatos por reglas y asignaciones explícitas. Las pruebas cubren área, zona, sistema y rangos de cantidad. Ninguna ruta se transforma automáticamente en actividad ni calcula CPM.",
    )
    section(
        document,
        13,
        "Plan Items",
        "Plan Items vincula Scope aprobado con unidades de planificación y control. Un Plan puede agrupar varios Scope Items y un Scope Item puede repartirse entre varios Plan Items. Las asignaciones conservan identidad, unidad y cantidad; la publicación verifica conciliación según la tolerancia del perfil.",
    )
    section(
        document,
        14,
        "Agrupar dividir y recomponer",
        "GROUP, SPLIT y RECOMPOSE son comandos con versiones esperadas de los orígenes. La operación verifica cardinalidad y conservación de cantidad por Scope y unidad, sustituye los Plan Items de origen y conserva sus enlaces e historia. Las salidas no heredan autoridad de Schedule, Cost o Progress.",
    )
    section(
        document,
        15,
        "Representación operacional derivada",
        "El derivado es una versión lógica inmutable de subconjuntos y grupos con membresía Scope y Plan, GUID, modelo fuente, regla, actor, fecha y hash. Preserva visibilidad y metadatos de sección, además de una referencia de restauración cuando la membresía sigue vigente. No crea fragmentos físicos ni modifica la geometría original.",
    )
    section(
        document,
        16,
        "Autoridad de las fuentes",
        "La separación de autoridades evita que la importación o la planificación modifiquen hechos de otros dominios.",
    )
    table(
        document,
        ["Objeto", "Autoridad"],
        [
            ["Geometría y propiedades IFC", "Fuente original versionada"],
            [
                "Definición naturaleza y cantidad Scope",
                "Scope Items y baseline aprobado",
            ],
            ["Ruta y descomposición de ejecución", "Execution Strategy y Plan Items"],
            [
                "Fechas CPM Cost y Progress",
                "Dominios futuros sin autoridad activa en O1B A",
            ],
        ],
        [3.1, 3.95],
    )
    section(
        document,
        17,
        "ADMIN y USER",
        "ADMIN configura catálogos de naturaleza, tipos, disciplinas, unidades, plantillas y reglas, tolerancias y separación de funciones mediante Scope Profiles. USER importa, prepara, revisa y opera según sus permisos y proyecto. Publicar configuración no aprueba Scope ni autoriza compromisos.",
        "Se evaluaron Project Manager, Project Controls Manager, Scope Manager, BIM Manager y Coordinator, Planner, Cost Controller, Construction Manager, Document Controller, Executive Viewer y Operational Administrator. Los nuevos roles reciben solo la lectura O1A y navegación mínima necesarias, sin elevarlos a administradores de organización.",
    )
    section(
        document,
        18,
        "Seguridad y separación de funciones",
        "Las operaciones materiales requieren autorización O1A vigente y perfiles publicados, además de permisos Scope y asignación al proyecto canónico. Las consultas no aprovisionan roles ni escriben hechos de negocio. Se verificaron aislamiento tenant y proyecto, If-Match, recibos idempotentes, auditabilidad y separación entre preparación y aprobación.",
        "PostgreSQL refuerza identidad canónica, referencias entre proyectos, fuentes originales e historia publicada. Las rutas genéricas de configuración no pueden crear, modificar, clonar o publicar Scope Profiles por fuera de la API dedicada. La aprobación de política y estrategia considera también al último editor material.",
        "La revisión cruzada acotada detectó y corrigió dos casos adicionales: cambio de unidad frente a la fila Excel original y publicación de un Plan preparado contra Scope reemplazado. Ambos quedaron cubiertos por regresiones de API. No quedaron hallazgos S1 o S2 abiertos dentro del alcance revisado; esto no equivale a una auditoría exhaustiva ni a aceptación de negocio.",
    )
    section(
        document,
        19,
        "API",
        f"La API contiene {len(data['api'])} operaciones bajo /api/v1/scope. Las rutas de proyecto reciben project_workspace_id canónico, nunca un ID legacy intercambiable. Incluyen versiones y objetos, descarga autorizada de fuente, staging Excel, Scope Items y jerarquía, revisiones y baselines, estrategias y segmentación, Plan Items, transformaciones y derivados. ADMIN usa su API de perfiles separada.",
        "El inventario completo de método, ruta y operación se entrega en la hoja 12_API. Las escrituras usan Idempotency-Key y los cambios sobre versiones existentes exigen If-Match. El linaje y las fuentes publicadas no admiten un PATCH genérico de estado.",
    )
    section(
        document,
        20,
        "Frontend",
        "La entrada canónica es /workspaces/{workspace_id}/scope-manager. Las cuatro pantallas internas son 3D Model, Scope Items, Strategy & Path of Execution y Plan Items, con CompactModuleHeader. Se conserva la desactivación explícita del módulo por Workspace y no se reescribe App.tsx.",
        "El modelo permite selección de versión, filtros, inspección de propiedades y aislamiento de objetos. Scope muestra tabla o árbol, naturaleza, cantidades, fuente y baseline. Plan muestra asignaciones y preview 3D derivado. WBS, actividades, CBS y AWP aparecen expresamente como integraciones futuras, no activas.",
        "TypeScript y el control de formato pasan. ESLint de los archivos integrados no reporta errores; conserva diez advertencias de efectos y dependencias, nueve en el visor reutilizado y una en ModelWorkspace. La selección del modelo derivado se memoiza antes de preparar su proveedor y representación.",
    )
    section(
        document,
        21,
        "Migración",
        "La revisión única 20261020_0070 sucede a 20261019_0069 y crea doce tablas aditivas. La prueba real verificó upgrade, downgrade vacío seguro y reaplicación, manteniendo las tablas y protecciones O1A. El downgrade rechaza eliminar historia Scope poblada. No se reescribió ninguna migración histórica ni se ejecutó backfill sobre los proyectos habituales.",
    )
    section(
        document,
        22,
        "Validación PostgreSQL",
        "Las siguientes suites son evidencia ejecutada. Los parsers incluidos en la suite Scope no se suman otra vez como pruebas distintas. Los fixtures representan actores y estados canónicos aislados, no aceptación de negocio ni una nueva prueba de creación de proyectos.",
    )
    table(
        document,
        ["Suite", "Casos", "Aprobados", "Fallos"],
        [
            [suite["name"], suite["tests"], suite["passed"], suite["failed"]]
            for suite in data["suites"]
        ],
        [4.65, 0.8, 0.8, 0.8],
    )
    p(
        document,
        f"Total distinto: {data['distinct_test_cases']} casos, con {data['recorded_test_executions']} ejecuciones registradas en las suites finales. Los intentos anteriores y sus errores permanecen separados de la evidencia final.",
    )
    section(
        document,
        23,
        "Validación Chromium",
        f"El recorrido final registra {len(data['browser'].get('checks', []))} comprobaciones con resultado {data['browser'].get('result')}. Usa login real por tenant, sesión emitida legítimamente, API y PostgreSQL reales, sin interceptar respuestas simuladas. Verifica configuración ADMIN, IFC y propiedades, Excel y candidatos, baseline autorizado, estrategia, asignaciones y transformaciones Plan, derivado y fuente original intacta.",
        "Los intentos iniciales detectaron bundle desactualizado y selectores de prueba demasiado estrictos. Se conservaron las fallas y se repitió el recorrido con el bundle reconstruido. Las comprobaciones UI no se presentan como pruebas de carga ni como firma empresarial.",
    )
    section(
        document,
        24,
        "Conservación del Proyecto 14",
        f"El proyecto canónico conserva código PRJ-PPMIS, external_key PRJ-PYP-PPMIS, Record Code 01.03.01.02.01, estado active y versión 1. La comparación de solo lectura antes y después devuelve {'PASS' if data['project14_preserved'] else 'FAIL'}. SHA256 de fila: {data['project14_sha256']}.",
        "No se confunde ese Workspace con projects.id 14, que no existe en la base consultada. Las referencias legacy consultadas permanecen sin cambios. El proyecto no se utilizó para fixtures destructivos ni se le creó una vinculación automática.",
    )
    section(
        document,
        25,
        "Observaciones de rendimiento",
        "Las mediciones corresponden al entorno local y a los archivos pequeños de aceptación. Incluyen parseo y consultas del modelo, objetos, Scope, jerarquía, agrupamiento Plan y derivados. Los tiempos pueden incluir inicialización de motor y contención del equipo; no constituyen SLA ni capacidad productiva.",
    )
    table(
        document,
        ["Operación", "N", "Mín ms", "Mediana ms", "Máx ms"],
        data["performance"],
        [2.65, 0.45, 1.05, 1.45, 1.45],
    )
    section(
        document,
        26,
        "Brechas y límites remanentes",
        "La instalación local habitual no recibió migración ni aprovisionamiento automático. Su despliegue requiere respaldo, upgrade aditivo y asignaciones explícitas. La limitación histórica de inicialización vacía de Alembic se mantiene documentada. Las observaciones locales no acreditan escalabilidad con modelos grandes.",
        "La representación derivada implementada es lógica. Los fragmentos físicos, la línea de tiempo 4D, WBS, P6, Schedule, Cost, Progress, Procurement y Construction completos siguen fuera del alcance. Las firmas estratégicas permanecen pendientes. El detalle de hallazgos y resolución está en la hoja 13_Defects.",
    )
    section(
        document,
        27,
        "Decisión de liberación",
        f"Estado exacto: {data['release_state']}. Salida: {data['final_output']}.",
        "Esta decisión acredita únicamente la fundación Scope O1B A verificada; no declara completo el macroproceso operacional ni Project Controls. Se conserva STRATEGIC_MACROPROCESS_TECHNICALLY_READY_UAT_SIGNOFF_PENDING. No se registra aprobación humana o compromiso presupuestario a partir de estas pruebas.",
    )
    section(
        document,
        28,
        "Siguiente incremento recomendado",
        data["next_recommendation"] + ".",
        "No se inicia automáticamente. La siguiente ejecución requiere autorización explícita y deberá reutilizar Scope aprobado, Plan Items, contexto O1A y las fuentes existentes.",
    )
    for index, paragraph in enumerate(document.paragraphs):
        paragraph.paragraph_format.widow_control = True
        if paragraph.style.name == "Normal":
            paragraph.paragraph_format.keep_together = True
        if paragraph.text in {
            "16 Autoridad de las fuentes",
            "25 Observaciones de rendimiento",
            "27 Decisión de liberación",
        }:
            document.paragraphs[index + 1].paragraph_format.keep_with_next = True
    return document


def main():
    assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_SHA
    data = json.loads(DATA.read_text("utf-8"))
    OUT.mkdir(parents=True, exist_ok=True)
    QA.mkdir(parents=True, exist_ok=True)
    stream = io.BytesIO()
    build(data).save(stream)
    output = OUT / "P&Pmis_Wave_O1B_A_Scope_Foundation_Technical_Report.docx"
    with (
        ZipFile(REFERENCE) as source,
        ZipFile(stream) as generated,
        ZipFile(output, "w") as final,
    ):
        for entry in source.infolist():
            final.writestr(
                entry,
                generated.read("word/document.xml")
                if entry.filename == "word/document.xml"
                else source.read(entry.filename),
            )
    with ZipFile(REFERENCE) as source, ZipFile(output) as final:
        changed = [
            name for name in source.namelist() if source.read(name) != final.read(name)
        ]
        assert changed == ["word/document.xml"]
        namespace = {
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        }
        a, b = [
            etree.fromstring(archive.read("word/document.xml"))
            for archive in (source, final)
        ]
        assert etree.tostring(a.find("w:body/w:sectPr", namespace)) == etree.tostring(
            b.find("w:body/w:sectPr", namespace)
        )
    manifest = {
        "reference_sha256": REFERENCE_SHA,
        "data_sha256": hashlib.sha256(DATA.read_bytes()).hexdigest(),
        "output": str(output),
        "output_sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
        "changed_parts": changed,
        "preserved_parts": len(source.namelist()) - len(changed),
    }
    (QA / "documents-build-manifest.json").write_text(
        json.dumps(manifest, indent=2), encoding="utf-8"
    )
    print(json.dumps(manifest))


if __name__ == "__main__":
    main()
