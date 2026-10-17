"""Collect literal persistent-UAT results and build source-preserving Word successors.

This script never executes application commands, grants approval, starts tests
or edits the retained 15 October artifacts. Authoring markers already ran once
for this logical request and must not be repeated. Root performs render QA.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import sys
from collections import Counter
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile
from zoneinfo import ZoneInfo

import generate_strategic_final_uat_reports as style
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "docs/validation/gate07j-controlled-class-b"
QA = ROOT / "outputs/.qa/gate07j-controlled-class-b-rebuild"
SOURCE = ROOT / "outputs/gate07j-controlled-class-b-20261015"
OUT = ROOT / "outputs/gate07j-controlled-class-b-20261016"
FACTS = DATA / "rebuild-results.json"
NEXT_BLOCKED = "REMEDIATE THE SMALLEST REMAINING STRATEGIC UAT BLOCKER"
NEXT_READY = "AUTHORIZED BUSINESS UAT SIGN-OFF"
NAMES = style.FILES
RETAINED_MINOR_TEXT = (
    "Permanecen OPEN los hallazgos menores previos UAT-S4-001 (rutas no "
    "aplicables), UAT-S4-002 (renderers legacy de Business Case), UAT-S4-003 "
    "(metadatos de Explorer), UAT-S4-004 (desbordamiento visual en Gate) y "
    "UAT-S4-005 (navegación móvil). No se convierten en cerrados por aprobar "
    "este recorrido técnico. El navegador actual reobservó el desbordamiento "
    "de hashes en Strategic Gate, evidencia del mismo UAT-S4-004; no se abre "
    "un duplicado. No se declara una interfaz 100% libre de incidencias de layout."
)


def load(path, default=None):
    return (
        json.loads(Path(path).read_text(encoding="utf-8-sig"))
        if path and Path(path).is_file()
        else default
    )


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def relative(path):
    path = Path(path).resolve()
    return path.relative_to(ROOT).as_posix() if path.is_relative_to(ROOT) else str(path)


def artifact(path):
    path = Path(path)
    return (
        {"file": relative(path), "sha256": digest(path), "bytes": path.stat().st_size}
        if path.is_file()
        else {"file": relative(path), "missing": True}
    )


def junit(path):
    record = artifact(path)
    if record.get("missing"):
        return {**record, "status": "NOT_COMPLETED", "counts": None}
    try:
        tree = ElementTree.parse(path)
    except ElementTree.ParseError as exc:
        return {**record, "status": "INCOMPLETE", "counts": None, "error": str(exc)}
    testcases = list(tree.getroot().iter("testcase"))
    counts, failures = Counter(), []
    for item in testcases:
        state = (
            "ERROR"
            if item.find("error") is not None
            else "FAIL"
            if item.find("failure") is not None
            else "SKIP"
            if item.find("skipped") is not None
            else "PASS"
        )
        counts[state] += 1
        if state in {"FAIL", "ERROR"}:
            error = item.find("error") if state == "ERROR" else item.find("failure")
            failures.append(
                {
                    "test": f"{item.get('classname')}::{item.get('name')}",
                    "message": error.get("message", ""),
                    "phase": "SETUP"
                    if "setup" in error.get("message", "").lower()
                    else "CALL_OR_TEARDOWN",
                    "assertion": next(
                        (
                            line.strip().removeprefix(">").strip()
                            for line in reversed((error.text or "").splitlines())
                            if line.lstrip().startswith(">")
                        ),
                        None,
                    ),
                    "location": next(
                        (
                            line.strip()
                            for line in reversed((error.text or "").splitlines())
                            if line.startswith(("tests/", "backend/"))
                        ),
                        None,
                    ),
                }
            )
    suites = list(tree.getroot().iter("testsuite"))
    return {
        **record,
        "status": "COMPLETED",
        "counts": {k: counts[k] for k in ("PASS", "FAIL", "ERROR", "SKIP")},
        "total": len(testcases),
        "unique_total": len({(i.get("classname"), i.get("name")) for i in testcases}),
        "seconds": sum(float(s.get("time", 0)) for s in suites),
        "failures": failures,
        "error_groups": [
            {
                "phase": phase,
                "location": location,
                "assertion": assertion,
                "case_count": amount,
            }
            for (phase, location, assertion), amount in Counter(
                (row["phase"], row["location"], row["assertion"]) for row in failures
            ).items()
        ],
    }


def junit_identity_multiset(path):
    if not Path(path).is_file():
        return None
    try:
        root = ElementTree.parse(path).getroot()
    except ElementTree.ParseError:
        return None
    return Counter(
        (row.get("classname"), row.get("name")) for row in root.iter("testcase")
    )


def passed(test):
    return (
        test.get("status") == "COMPLETED"
        and test.get("total", 0) > 0
        and test["counts"]["PASS"] > 0
        and not test["counts"]["FAIL"]
        and not test["counts"]["ERROR"]
    )


def operational_snapshot_complete(snapshot):
    if not isinstance(snapshot, dict):
        return False
    workspace = snapshot.get("workspace14", {})
    return (
        snapshot.get("head") == "20261011_0068"
        and workspace.get("id") == 14
        and workspace.get("code") == "PRJ-PPMIS"
        and workspace.get("external_key") == "PRJ-PYP-PPMIS"
        and workspace.get("status") == "active"
        and workspace.get("version") == 1
        and workspace.get("row_md5") == "ef67b78a4d4902d2dc6db3efb017b80b"
        and all(
            isinstance(snapshot.get(key), int) and snapshot[key] >= 0
            for key in ("qra_run_count", "published_count", "workspace_count")
        )
        and bool(snapshot.get("evidence_policies"))
        and all(
            isinstance(row.get("id"), int)
            and row["id"] > 0
            and isinstance(row.get("hash"), str)
            and re.fullmatch(r"[0-9a-fA-F]{64}", row.get("hash", ""))
            for row in snapshot["evidence_policies"]
        )
    )


def persistence_snapshot_complete(snapshot):
    tables = (
        "tenants",
        "admin_configurations",
        "investment_decisions",
        "enterprise_workspaces",
        "investment_gate_preparations",
        "investment_gate_evidence_items",
        "post_fid_project_eligibilities",
        "project_workspace_initializations",
    )
    return (
        isinstance(snapshot, dict)
        and snapshot.get("database") == "strategic_fresh_uat"
        and snapshot.get("head") == "20261011_0068"
        and all(
            isinstance(snapshot.get(table), dict)
            and isinstance(snapshot[table].get("count"), int)
            and snapshot[table]["count"] > 0
            and isinstance(snapshot[table].get("rows_md5"), str)
            and re.fullmatch(r"[0-9a-fA-F]{32}", snapshot[table].get("rows_md5", ""))
            for table in tables
        )
    )


def pick(record, fields):
    return {name: record.get(name) for name in fields if name in record}


def last_stage(chain, name, *, passed_only=True):
    records = [
        item
        for item in chain.get("stages", [])
        if item["stage"] == name and (not passed_only or item.get("status") == "PASS")
    ]
    return records[-1].get("data", {}) if records else {}


def count_text(test):
    c = test.get("counts")
    return (
        f"{c['PASS']} PASS / {c['FAIL']} FAIL / {c['ERROR']} ERROR / {c['SKIP']} SKIP"
        if c
        else f"{test.get('status', 'NOT_COMPLETED')} sin conteo final"
    )


def reconciliation_record(row):
    """Compact records retain null currentness and the literal policy controls."""
    left = (
        f"{row['requirement_key']}\n"
        f"Clase {row['classification']} · Required: {'Sí' if row['required'] else 'No'} "
        f"· Bloqueante: {'Sí' if row['blocking'] else 'No'}\n"
        f"Origen: {row['provider_key'] or 'Referencia'}"
    )
    source_current = row.get("source_current")
    currentness = "No expone (null)" if source_current is None else str(source_current)
    coverage = row.get("coverage_percent")
    required_coverage = row.get("coverage_requirement")
    coverage_text = "No expone (null)" if coverage is None else f"{coverage}%"
    required_text = (
        "No expone (null)" if required_coverage is None else f"{required_coverage}%"
    )
    if row["status"] == "MISSING":
        reason = "No hay referencia válida; " + (
            "requisito bloqueante."
            if row["blocking"]
            else "no bloqueante por política."
        )
    elif row["status"] == "NOT_REQUIRED":
        reason = "Requisito opcional, required=False."
    elif row["status"] == "VALID" and row.get("provider_key"):
        reason = (
            f"Fuente gobernada resuelta VALID; registro {row.get('source_record_id')}."
        )
    elif row["status"] == "VALID":
        reason = f"Referencia VALID por revisor {row.get('validated_by')}."
    else:
        reason = f"Resolución capturada: {row.get('resolution_status', 'No expone')}."
    right = (
        f"Estado: {row['status']}\n"
        f"Currentness: {currentness}\n"
        f"Cobertura: {coverage_text}; exigida: {required_text}\n"
        f"Razón: {reason}\n"
        f"Blockers: {'; '.join(row.get('blockers') or []) or 'ninguno'}\n"
        f"Warnings: {'; '.join(row.get('warnings') or []) or 'ninguno'}"
    )
    return [left, right]


def collect(args):
    chain_path = QA / "fresh-capital-owner-chain.json"
    chain = load(chain_path, {})
    setup = chain.get("setup", {})
    live_path = Path(args.live or QA / "selected-final.json")
    live = load(live_path, {})
    broad_dir = Path(args.broad_dir or QA / "broad-clean")
    broad_manifest = load(broad_dir / "run-manifest.json", {})
    broad_state = broad_manifest.get("status", "NOT_COMPLETED")
    broad_xml = broad_dir / broad_manifest.get(
        "junit", "backend-strategic-regression-clean.xml"
    )
    broad = (
        junit(broad_xml)
        if broad_state in {"COMPLETED", "COMPLETE", "PASS", "FAIL"}
        else {**artifact(broad_xml), "status": broad_state, "counts": None}
    )
    broad["run_manifest"] = artifact(broad_dir / "run-manifest.json")
    broad["module_count"] = broad_manifest.get("module_count")
    broad["same_selection_expected_cases"] = broad_manifest.get(
        "expected_collected_cases"
    )
    broad["focused_added_to_counts"] = False
    broad["interrupted_attempts"] = [
        {
            "artifact": artifact(p),
            "status": record.get("status"),
            "stop_reason": record.get("stop_reason"),
            "result_counts": None,
        }
        for p in sorted(QA.glob("broad*/run-manifest.json"))
        if (record := load(p, {})).get("status", "").startswith("INCOMPLETE")
    ]
    if not broad_manifest and broad["interrupted_attempts"]:
        broad["status"] = "INCOMPLETE_PREVIOUS_ATTEMPT_NO_FINAL_RUN"
    source_hashes = broad_manifest.get("product_sha256_at_start", {})
    broad["product_hashes_match"] = bool(source_hashes) and all(
        (ROOT / p).is_file() and digest(ROOT / p).lower() == h.lower()
        for p, h in source_hashes.items()
    )
    fixture_hashes = broad_manifest.get("fixture_sha256_at_start", {})
    broad["fixture_hashes_match"] = bool(fixture_hashes) and all(
        (ROOT / path).is_file() and digest(ROOT / path).lower() == sha.lower()
        for path, sha in fixture_hashes.items()
    )
    selection_source = broad_manifest.get("selection_source")
    selection_path = ROOT / selection_source if selection_source else None
    baseline_ids = junit_identity_multiset(selection_path) if selection_path else None
    current_ids = junit_identity_multiset(broad_xml)
    broad["selection_baseline"] = artifact(selection_path) if selection_path else None
    broad["identity_multiset_difference_from_baseline"] = (
        sum((current_ids - baseline_ids).values())
        + sum((baseline_ids - current_ids).values())
        if current_ids is not None and baseline_ids is not None
        else None
    )
    broad["identity_multiset_matches_baseline"] = bool(baseline_ids) and (
        current_ids == baseline_ids
    )
    broad["inherited_duplicate_identities"] = [
        {"identity": "::".join(identity), "count": count}
        for identity, count in (baseline_ids or {}).items()
        if count > 1
    ]
    broad["skip_summary"] = broad_manifest.get("skip_summary", {})
    broad["skip_details"] = (
        artifact(broad_dir / broad_manifest["skip_details"])
        if broad_manifest.get("skip_details")
        else None
    )
    bundled_git = Path(
        "C:/Users/Ricardo/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/git/cmd/git.exe"
    )
    git = bundled_git if bundled_git.is_file() else Path("git")
    baseline = {
        "branch": subprocess.check_output(
            [str(git), "branch", "--show-current"], cwd=ROOT, text=True
        ).strip(),
        "commit": subprocess.check_output(
            [str(git), "rev-parse", "HEAD"], cwd=ROOT, text=True
        ).strip(),
        "alembic_head_captured": chain.get("head"),
        "product_hashes_at_collection": {
            p: digest(ROOT / p) for p in source_hashes if (ROOT / p).is_file()
        },
        "fixture_hashes_at_collection": {
            path: digest(ROOT / path)
            for path in fixture_hashes
            if (ROOT / path).is_file()
        },
        "references": [artifact(SOURCE / NAMES[k]) for k in ("closeout", "uat")],
    }
    browser_dir = Path(args.browser_dir or QA / "browser-actual-1")
    browser_data = load(browser_dir / "fresh-browser-result.json", {})
    visual_qa = load(browser_dir / "visual-qa-summary.json", {})
    verified_screenshots = []
    for row in visual_qa.get("screenshots", []):
        image_path = (
            browser_dir
            / visual_qa.get("screenshot_directory", "")
            / row.get("name", "")
        ).resolve()
        verified_screenshots.append(
            {
                **row,
                "artifact": artifact(image_path),
                "sha256_verified": bool(
                    image_path.is_relative_to(browser_dir.resolve())
                    and image_path.is_file()
                    and digest(image_path).lower() == row.get("sha256", "").lower()
                ),
            }
        )
    browser_test = junit(browser_dir / "strategic-capital-current.xml")
    browser_history = []
    for path in sorted(QA.glob("browser-actual-*/fresh-browser-result.json")):
        if path.parent.resolve() == browser_dir.resolve():
            continue
        prior_result = load(path, {})
        prior_visual = load(path.parent / "visual-qa-summary.json", {})
        browser_history.append(
            {
                "artifact": artifact(path),
                "status": prior_result.get("status", "NOT_COMPLETED"),
                "classification": prior_visual.get("status", "NOT_CLASSIFIED"),
                "test": junit(path.parent / "strategic-capital-current.xml"),
                "blocker": prior_result.get("blocker"),
                "failure": prior_result.get("failure"),
                "visual_qa_summary": {
                    "artifact": artifact(path.parent / "visual-qa-summary.json"),
                    "result": prior_visual,
                },
                "current_acceptance_evidence": False,
                "added_to_final_counts": False,
            }
        )
    tests = {
        "class_b": junit(QA / "class-b-focused-final/class-b-all-current.xml"),
        "bridge": junit(QA / "project-initialization-post-fid-bridge.xml"),
        "canonical": junit(QA / "project-initialization-canonical-guards.xml"),
        "frontend": junit(QA / "frontend-presentations-focused.xml"),
        "legacy": junit(QA / "legacy-focused-attempt-2/legacy-focused.xml"),
        "postactivation": junit(
            Path(
                args.postactivation_xml
                or QA / "postactivation-focused-final-3/currentness-focused.xml"
            )
        ),
        "ci_contract_guard": junit(
            QA / "ci-contract-guard-final/ci-contract-guard.xml"
        ),
        "broad": broad,
        "browser": browser_test,
    }
    tests["postactivation"]["previous_attempts"] = [
        {**junit(path), "current_acceptance_evidence": False}
        for path in (
            QA / "postactivation-focused-final/currentness-focused.xml",
            QA / "postactivation-focused-final-2/currentness-focused.xml",
        )
        if path.is_file() and relative(path) != tests["postactivation"]["file"]
    ]
    tests["ci_contract_guard"]["included_in_broad_counts"] = False
    tests["ci_contract_guard"]["run_manifest"] = artifact(
        QA / "ci-contract-guard-final/run-manifest.json"
    )
    tests["ci_contract_guard"]["previous_setup_attempt"] = {
        "artifact": artifact(QA / "ci-contract-guard/run-manifest.json"),
        "current_acceptance_evidence": False,
        "result_counts_added_to_broad": False,
    }
    identity = pick(
        setup,
        (
            "tenant_id",
            "tenant_slug",
            "enterprise_id",
            "business_unit_id",
            "portfolio_id",
            "project_id",
            "capital_cycle_id",
            "capital_demand_id",
            "allocation_plan_id",
            "idea_id",
            "analysis_id",
            "proposal_id",
            "strategic_gate_id",
            "project_creation_request_id",
            "actor_id",
            "reviewer_id",
            "admin_author_id",
            "admin_publisher_id",
            "decision_maker_id",
            "initializer_id",
            "activator_id",
            "capital_endorser_id",
        ),
    )
    recon = last_stage(chain, "A27_ALL_REQUIREMENTS_RECONCILIATION")
    frozen = last_stage(chain, "A29_FROZEN_PACKAGE")
    captured_fid = last_stage(chain, "A30_FID")
    initial_post = last_stage(chain, "A31_POST_FID")
    post_rows = live.get("post_fid", []) or []
    current_post = max(
        post_rows, key=lambda row: row.get("eligibility_round", 0), default={}
    )
    init = live.get("initialization") or {}
    if isinstance(init, list):
        init = init[-1] if init else {}
    project = live.get("project", {})
    references = []
    for key in (
        "permit-readiness",
        "contracting-strategy-readiness",
        "execution-plan-readiness",
        "schedule-maturity",
    ):
        captured = last_stage(chain, f"CLASS_B_VALIDATED_{key}")
        reference, evidence = (
            captured.get("reference", {}),
            captured.get("evidence", {}),
        )
        uri = reference.get("document_uri", "")
        path = (QA / uri.removeprefix("artifact://")).resolve() if uri else None
        document_ok = bool(
            path
            and path.is_relative_to(QA.resolve())
            and path.is_file()
            and digest(path) == reference.get("document_sha256")
        )
        references.append(
            {
                "requirement_key": key,
                "reference": reference,
                "evidence": evidence,
                "document_verified": document_ok,
                "artifact": artifact(path) if path else None,
            }
        )
    class_b_ok = len(references) == 4 and all(
        r["document_verified"]
        and r["evidence"].get("status") == "VALID"
        and r["evidence"].get("classification") == "B"
        and r["evidence"].get("provider_key") is None
        and r["evidence"].get("scope_id") == setup.get("project_id")
        and r["evidence"].get("preparation_id")
        == frozen.get("preparation", {}).get("id")
        and r["reference"].get("project_id") == setup.get("project_id")
        and r["reference"].get("classification_label") == "UAT CONTROLLED EVIDENCE"
        and r["reference"].get("representative_only") is True
        and all(
            r["reference"].get(field)
            for field in (
                "reference_id",
                "title",
                "revision",
                "version",
                "as_of_date",
                "owner",
                "scope_reference_hash",
            )
        )
        and r["reference"].get("document_sha256") == r["reference"].get("evidence_hash")
        and r["evidence"].get("submitted_by") == setup.get("actor_id")
        and r["evidence"].get("validated_by") == setup.get("reviewer_id")
        and isinstance(r["evidence"].get("submitted_by"), int)
        and r["evidence"].get("submitted_by") != r["evidence"].get("validated_by")
        and isinstance(r["evidence"].get("validated_by"), int)
        for r in references
    )
    reconciliation_rows = recon.get("rows", [])
    required_rows = [row for row in reconciliation_rows if row.get("required") is True]
    optional_rows = [row for row in reconciliation_rows if row.get("required") is False]
    reconciliation_ok = (
        len(reconciliation_rows) == 20
        and len({row.get("requirement_key") for row in reconciliation_rows}) == 20
        and len(required_rows) == recon.get("required_row_count") == 19
        and len(optional_rows) == recon.get("optional_row_count") == 1
        and optional_rows[0].get("classification") == "C"
        and optional_rows[0].get("status") == "NOT_REQUIRED"
        and recon.get("readiness", {}).get("required_requirement_count") == 19
        and sum(row.get("status") == "VALID" for row in required_rows)
        == recon.get("readiness", {}).get("valid_requirement_count")
        == 15
        and all(
            row.get("status") == "VALID"
            for row in required_rows
            if row.get("blocking") is True
        )
    )
    fid_ok = (
        captured_fid.get("status") == "DECIDED"
        and captured_fid.get("outcome") == "GO"
        and captured_fid.get("profile_code") == "PYP-FFA"
        and captured_fid.get("tenant_id") == setup.get("tenant_id")
        and captured_fid.get("investment_gate_preparation_id")
        == frozen.get("preparation", {}).get("id")
        and captured_fid.get("source_preparation_hash")
        == frozen.get("preparation", {}).get("preparation_hash")
        and all(
            captured_fid.get(k) is True
            for k in ("is_fid", "investment_authorized", "full_funds_authorized")
        )
        and captured_fid.get("can_activate_project") is False
    )
    live_preparations = [
        row
        for row in live.get("preparations", [])
        if row.get("id") == frozen.get("preparation", {}).get("id")
    ]
    live_decisions = [
        row
        for row in live.get("decisions", [])
        if row.get("id") == captured_fid.get("id")
    ]
    material_hashes_ok = (
        len(live_preparations) == len(live_decisions) == 1
        and bool(frozen.get("preparation", {}).get("preparation_hash"))
        and live_preparations[0].get("status") == "FROZEN"
        and live_preparations[0].get("hash")
        == frozen["preparation"]["preparation_hash"]
        and bool(captured_fid.get("decision_hash"))
        and live_decisions[0].get("decision_hash") == captured_fid["decision_hash"]
        and live_decisions[0].get("status") == "DECIDED"
        and live_decisions[0].get("is_fid") is True
        and live_decisions[0].get("decided_by") == setup.get("decision_maker_id")
    )
    same_identity = (
        chain.get("database") == live.get("database") == "strategic_fresh_uat"
        and live.get("tenant", {}).get("id") == setup.get("tenant_id")
        and live.get("tenant", {}).get("slug") == setup.get("tenant_slug")
        and project.get("id") == setup.get("project_id")
        and project.get("parent_id") == setup.get("portfolio_id")
    )
    activation_ok = (
        same_identity
        and project.get("status") == "active"
        and init.get("state") == "ACTIVATED"
        and init.get("started_by_user_id") == setup.get("initializer_id")
        and init.get("activated_by_user_id") == setup.get("activator_id")
        and setup.get("initializer_id") != setup.get("activator_id")
    )
    live_ok = (
        current_post.get("eligibility_status") == "NO_REACTIVATION_REQUIRED"
        and current_post.get("source_current") is True
        and current_post.get("project_workspace_id") == setup.get("project_id")
        and current_post.get("investment_decision_id") == captured_fid.get("id")
        and current_post.get("fid_decision_hash") == captured_fid.get("decision_hash")
        and not current_post.get("blockers_json", current_post.get("blockers", []))
    )
    browser_identity = browser_data.get("identity", {})
    browser_functional_ok = (
        passed(browser_test)
        and browser_data.get("status") == "PASS"
        and browser_identity.get("tenant_id") == setup.get("tenant_id")
        and browser_identity.get("project_id") == setup.get("project_id")
        and browser_identity.get("portfolio_id") == setup.get("portfolio_id")
        and browser_identity.get("fid_id") == captured_fid.get("id")
        and browser_identity.get("preparation_id")
        == frozen.get("preparation", {}).get("id")
        and all(
            isinstance(browser_data.get(key), list)
            for key in ("runtime_errors", "http_errors", "browser_business_writes")
        )
        and not browser_data.get("runtime_errors")
        and not browser_data.get("http_errors")
        and not browser_data.get("browser_business_writes")
    )
    browser_visual_ok = (
        visual_qa.get("functional_run_status") == "PASS"
        and visual_qa.get("status")
        in {
            "VISUAL_REVIEW_COMPLETED_WITH_OPEN_EXISTING_FINDINGS",
            "VISUAL_REVIEW_COMPLETED",
            "PASS_WITH_OPEN_EXISTING_COSMETIC_FINDINGS",
        }
        and isinstance(visual_qa.get("reviewed_named_screenshots"), int)
        and visual_qa["reviewed_named_screenshots"] == len(verified_screenshots) > 0
        and len({row.get("name") for row in verified_screenshots})
        == len(verified_screenshots)
        and all(row["sha256_verified"] for row in verified_screenshots)
        and visual_qa.get("browser_result_sha256", "").lower()
        == artifact(browser_dir / "fresh-browser-result.json").get("sha256")
        and visual_qa.get("junit_sha256", "").lower()
        == artifact(browser_dir / "strategic-capital-current.xml").get("sha256")
    )
    browser_ok = browser_functional_ok and browser_visual_ok
    before, after = (
        load(QA / "operational-before.json"),
        load(QA / "operational-after.json"),
    )
    preservation_ok = (
        operational_snapshot_complete(before)
        and operational_snapshot_complete(after)
        and before == after
    )
    persistence_before = load(QA / "persistence-before.json")
    persistence_after = load(QA / "persistence-after.json")
    persistence_ok = (
        persistence_snapshot_complete(persistence_before)
        and persistence_snapshot_complete(persistence_after)
        and persistence_before == persistence_after
    )
    focused_ok = all(
        passed(tests[name])
        for name in (
            "class_b",
            "bridge",
            "canonical",
            "frontend",
            "legacy",
            "postactivation",
        )
    )
    frontend_validation = load(DATA / "frontend-build-validation.json", {})
    frontend_ok = frontend_validation.get("result") == "PASS" and all(
        frontend_validation.get(key) == 0
        for key in (
            "exit_code",
            "typescript_no_emit_exit_code",
            "eslint_src_exit_code",
            "eslint_errors",
        )
    )
    broad_ok = (
        passed(broad)
        and broad.get("product_hashes_match")
        and broad.get("fixture_hashes_match")
        and broad.get("module_count") == 61
        and broad.get("total") == broad.get("same_selection_expected_cases") == 914
        and broad.get("identity_multiset_matches_baseline")
        and passed(tests["ci_contract_guard"])
    )
    backend_ok = chain.get("backend_chain") == "PASS" and not chain.get("blocker")
    backend_stages = {
        "A27": class_b_ok
        and reconciliation_ok
        and recon.get("readiness", {}).get("readiness_status") == "READY"
        and recon.get("readiness", {}).get("coverage_percent") == 100,
        "A28": bool(
            last_stage(chain, "A28_CURRENT_SOURCE_CONSISTENCY").get("source_current")
        ),
        "A29": material_hashes_ok
        and frozen.get("preparation", {}).get("status") == "FROZEN"
        and frozen.get("handoff", {}).get("can_enter_investment_decision") is True
        and frozen.get("handoff", {}).get("can_enter_fid") is False,
        "A30": fid_ok and material_hashes_ok,
        "A31": bool(
            initial_post
            and initial_post[0].get("eligibility_status")
            == "READY_FOR_PROJECT_WORKSPACE_INITIALIZATION"
        )
        and activation_ok
        and live_ok,
    }
    complete_proof = (
        backend_ok and browser_ok and live_ok and focused_ok and same_identity
    )
    cases = []
    actuals = {
        "A27": f"Readiness capturada {recon.get('readiness', {}).get('readiness_status', 'NOT_VERIFIED')}; cobertura {recon.get('readiness', {}).get('coverage_percent', 'desconocida')}%; cuatro Class B verificadas {class_b_ok}.",
        "A28": f"Consistencia antes del freeze comprobada {backend_stages['A28']}; situación actual Post-FID {current_post.get('eligibility_status', 'NOT_VERIFIED')}, source_current={current_post.get('source_current')}.",
        "A29": f"Paquete {frozen.get('preparation', {}).get('id')} {frozen.get('preparation', {}).get('status', 'NOT_VERIFIED')}; hash {frozen.get('preparation', {}).get('preparation_hash', 'no capturado')}. Freeze no es FID.",
        "A30": f"Decisión {captured_fid.get('id')} {captured_fid.get('status')} / {captured_fid.get('outcome')}; PYP-FFA, is_fid={captured_fid.get('is_fid')}; sin activación automática.",
        "A31": f"Post-FID inicial {initial_post[0].get('eligibility_status') if initial_post else 'NOT_VERIFIED'}; comando de activación {init.get('state', 'NOT_EXECUTED')} por actor {init.get('activated_by_user_id')}; currentness posterior {current_post.get('eligibility_status', 'NOT_VERIFIED')}, bloqueos {current_post.get('blockers_json', current_post.get('blockers', []))}.",
    }
    common_precondition = (
        f"Reconstrucción autorizada; Tenant {setup.get('tenant_id')}, Portfolio "
        f"{setup.get('portfolio_id')} y Project {setup.get('project_id')} verificados. "
    )
    preconditions = {
        "A27": common_precondition
        + "Perfil y política efectivos publicados; cuatro requisitos Class B bloqueantes sin referencia VALID; documentos representativos UAT con ámbito y hash verificables.",
        "A28": common_precondition
        + "Readiness READY y cobertura del Project 100%; Business Case final y fuentes gobernadas actuales disponibles sin modificar sus resultados para el Gate.",
        "A29": common_precondition
        + f"Preparación {frozen.get('preparation', {}).get('id')} READY, fuentes actuales, referencias revisadas; autor y congelador independientes con permiso y revisión vigente.",
        "A30": common_precondition
        + "Preparación FROZEN con handoff READY_FOR_INVESTMENT_DECISION; Project aún pending; autoridad FID independiente y checklist de decisión disponible.",
        "A31": common_precondition
        + "Decisión PYP-FFA DECIDED / GO y FID real; Project pending; configuración publicada de inicialización y activación y permisos de actores independientes.",
    }
    steps = {
        "A27": f"El autor {setup.get('actor_id')} ejecutó add_evidence para permit-readiness, contracting-strategy-readiness, execution-plan-readiness y schedule-maturity, generando SUBMITTED. El revisor {setup.get('reviewer_id')} ejecutó validate_evidence con la revisión vigente y obtuvo VALID en las cuatro. Se consultaron readiness y project_readiness y se conciliaron 19 filas required más la fila opcional C.",
        "A28": "Se compararon los IDs, versiones y hashes de las referencias proveedor capturadas con Business Case, FIN, Benefits, Commercial, FEL, Risk, Funding, Permit y fuentes de Portfolio. Se conservaron provider_evidence, política capturada, source_current y cross_source_consistency; no se reemplazaron proveedores por referencias manuales.",
        "A29": f"El autor {setup.get('actor_id')} ejecutó ready_for_review y el revisor independiente {setup.get('reviewer_id')} ejecutó freeze, pasando en cada comando la revision_version vigente, equivalente al control If-Match. Se verificaron hash inmutable, status FROZEN y handoff sin autoridad FID.",
        "A30": f"Se creó o reutilizó la misma decisión de la preparación, se asignó al decisor {setup.get('decision_maker_id')} y se ejecutó submit → start_review → ready con checklist. El decisor independiente {setup.get('decision_maker_id')} ejecutó decide GO. Se comprobaron is_fid, fondos y autorización de inversión, hash fuente y que Project seguía pending.",
        "A31": f"Se ejecutó evaluate 07L y se comprobó READY_FOR_PROJECT_WORKSPACE_INITIALIZATION. El inicializador {setup.get('initializer_id')} obtuvo el handoff y ejecutó start y validate sin activar. Tras refresh 07L se verificó READY_FOR_PROJECT_WORKSPACE_ACTIVATION. El activador independiente {setup.get('activator_id')} ejecutó activate con la revisión delegada. La verificación posterior usa refresh 07L, sin ejecutar una segunda activación, y exige NO_REACTIVATION_REQUIRED. La comprobación Chromium de la misma identidad se informa por separado según su evidencia real.",
    }
    for case_id, backend_pass in backend_stages.items():
        cases.append(
            {
                "id": case_id,
                "status": "PASS" if complete_proof and backend_pass else "BLOCKED",
                "backend_status": "PASS" if backend_pass else "BLOCKED",
                "preconditions": preconditions[case_id],
                "steps": steps[case_id],
                "expected": {
                    "A27": "Class B VALID, cobertura 1 de 1 y readiness READY",
                    "A28": "Fuentes y hashes consistentes según el contrato aplicable",
                    "A29": "FROZEN y READY_FOR_INVESTMENT_DECISION sin autoridad FID",
                    "A30": "PYP-FFA GO DECIDED con FID y fondos autorizados sin activar",
                    "A31": "Inicialización y activación separadas; currentness posterior correcto sin reactivación",
                }[case_id],
                "actual": actuals[case_id]
                + (
                    " Prueba formal pendiente de completar."
                    if not complete_proof
                    else ""
                ),
                "evidence": [
                    relative(chain_path),
                    relative(live_path),
                    browser_test["file"],
                ],
            }
        )
    s2_closed = (
        all(row["status"] == "PASS" for row in cases)
        and backend_ok
        and broad_ok
        and preservation_ok
        and persistence_ok
    )
    integration_closed = (
        activation_ok
        and live_ok
        and passed(tests["bridge"])
        and passed(tests["postactivation"])
        and browser_ok
    )
    defects = [
        {
            "id": "UAT-S2-003",
            "status": "CLOSED_SCOPED_TECHNICAL" if s2_closed else "OPEN",
            "actual": "Cuatro Class B, freeze y FID acreditados en la nueva cadena autorizada. "
            + (
                "Cierre integral acreditado."
                if s2_closed
                else "El cierre integral no se acredita mientras existan pruebas o currentness pendientes."
            ),
            "action": "Conservar cadena e historial; resolver únicamente el bloqueo pendiente y repetir las verificaciones afectadas.",
            "evidence": [
                relative(chain_path),
                relative(live_path),
                broad["file"],
                browser_test["file"],
            ],
        },
        {
            "id": "UAT-S3-005",
            "status": "CLOSED_SCOPED_PRESENTATION"
            if browser_ok and passed(tests["frontend"])
            else "OPEN_SCOPED_PRESENTATION",
            "actual": "La proyección visual de Business Case y Funding tiene regresión enfocada; su aceptación visual requiere Chromium actual sin falsa etiqueta STALE.",
            "action": "Mostrar el resultado autoritativo; no modificar Funding ni Business Case para ajustarlos a la pantalla.",
            "evidence": [tests["frontend"]["file"], browser_test["file"]],
        },
        {
            "id": "UAT-S4-006",
            "status": "CLOSED_SCOPED_PRESENTATION"
            if browser_ok and passed(tests["frontend"])
            else "OPEN_SCOPED_LIFECYCLE_LABEL",
            "actual": "Se separan lifecycle histórico de Proposal y elegibilidad actual de reingreso; no se reescribe la aprobación histórica.",
            "action": "Mantener ambos conceptos separados y verificar su presentación en la cadena actual.",
            "evidence": [tests["frontend"]["file"], browser_test["file"]],
        },
        {
            "id": "UAT-ENV-001",
            "status": "CLOSED_AUTHORIZED_RECONSTRUCTION" if same_identity else "OPEN",
            "actual": "La UAT tmpfs original se perdió. El usuario autorizó una reconstrucción gobernada en otra base persistente; no es el antiguo Tenant 4 / Project 16.",
            "action": "Conservar informes del 15 de octubre y respaldos verificables de la nueva base.",
            "evidence": [
                "docs/validation/gate07j-controlled-class-b/rebuild-authorization.md",
                relative(live_path),
            ],
        },
        {
            "id": "UAT-S2-005",
            "isNew": True,
            "classification": "LIFECYCLE INTEGRATION",
            "severity": "S2 CRITICAL",
            "title": "Integración de FID con inicialización y currentness posterior a activación",
            "status": "CLOSED_SCOPED_TECHNICAL" if integration_closed else "OPEN",
            "actual": f"Se reprodujeron dos desacoplamientos: inicialización no reconocía el FID real y la activación canónica provocaba PREPARATION_STALE al reevaluar reglas de ingreso pending-only. El puente de inicialización cuenta con regresión; activación real={activation_ok}. Post-FID actual={current_post.get('eligibility_status', 'NOT_VERIFIED')}; bloqueos={current_post.get('blockers_json', [])}. Focused de currentness: {count_text(tests['postactivation'])}; los errores compartidos de setup se distinguen de fallos de cada caso.",
            "action": "Reconocer la transición canónica autorizada sin omitir cambios genuinos de fuentes, condiciones o configuración.",
            "evidence": [
                tests["bridge"]["file"],
                tests["postactivation"]["file"],
                relative(live_path),
                browser_test["file"],
            ],
        },
    ]
    home_residuals = [
        row
        for row in visual_qa.get("residual_observations", [])
        if row.get("defect_id") == "UAT-S3-006"
    ]
    if home_residuals:
        home = home_residuals[0]
        defects.append(
            {
                "id": "UAT-S3-006",
                "isNew": True,
                "classification": "METADATA PROJECTION / DISCOVERY",
                "severity": "S3 MAJOR",
                "title": "Home omite responsable y plantilla; filtro por responsable incompleto",
                "status": "OPEN",
                "actual": "El mismo Project 4 muestra responsable y plantilla en el Overview canónico, pero Home proyecta Not assigned y No template snapshot. La auditoría fuente identifica impacto en presentación, lista, filtro por responsable y ETag: ese filtro puede omitir el Project. No se encontró consumo en autorización, lifecycle, elegibilidad o selección de Project en el repositorio actual; no se afirma cobertura de clientes externos o futuros. No es UAT-S4-003. La observación inicial sin clasificación se conserva y el triage posterior la clasifica S3 MAJOR OPEN.",
                "action": "Pendiente corrección focalizada del read model y pruebas de proyección/filtro, sin alterar permisos ni datos canónicos. Workaround: Overview canónico, Enterprise Explorer o ruta directa sin filtro por responsable. No se implementa fuera de este alcance.",
                "evidence": [
                    relative(
                        browser_dir
                        / home.get("consumer_audit", "home-context-consumer-audit.md")
                    ),
                    relative(browser_dir / "visual-qa-summary.json"),
                    *[
                        relative(browser_dir / path)
                        for path in home.get("screenshots", [])
                    ],
                ],
            }
        )
    ready = (
        s2_closed
        and integration_closed
        and focused_ok
        and broad_ok
        and browser_ok
        and preservation_ok
        and persistence_ok
        and frontend_ok
    )
    release = {
        "status": "STRATEGIC_MACROPROCESS_TECHNICALLY_READY_UAT_SIGNOFF_PENDING"
        if ready
        else "STRATEGIC_MACROPROCESS_NOT_VALIDATED",
        "decision": "UAT_SIGNOFF_PENDING" if ready else "NO_GO",
        "ready": ready,
        "business_release_authorized": False,
        "production_release_authorized": False,
        "technical_scope": "CONTROLLED_CLASS_B_FID_POSTFID_TARGET_ONLY",
        "full_ui_metadata_validated": False,
        "residual_impact_review_pending": any(
            not row.get("defect_id")
            for row in visual_qa.get("residual_observations", [])
        ),
        "open_metadata_defects": [
            row["id"] for row in defects if row["id"] == "UAT-S3-006"
        ],
        "human_signoff": "PENDING_NOT_ASSUMED",
        "signoffs": [],
        "next_recommendation": NEXT_READY if ready else NEXT_BLOCKED,
        "conditions": {
            "same_identity": same_identity,
            "controlled_references": class_b_ok,
            "reconciliation_19_required_plus_optional": reconciliation_ok,
            "frozen_material_hashes_unchanged": material_hashes_ok,
            "backend_chain": backend_ok,
            "activation_command": activation_ok,
            "post_activation_currentness": live_ok,
            "focused": focused_ok,
            "broad": broad_ok,
            "browser": browser_ok,
            "browser_functional": browser_functional_ok,
            "browser_visual_review": browser_visual_ok,
            "operational_preservation": preservation_ok,
            "named_volume_restart_verified": persistence_ok,
            "frontend_build_typecheck_lint": frontend_ok,
        },
    }
    rc = []
    for criterion_id, ok, actual in (
        (
            "RC02",
            activation_ok,
            "La misma identidad canónica pasa de pending a ACTIVATED por comandos y actores independientes; FID no activa automáticamente.",
        ),
        (
            "RC03",
            passed(tests["class_b"]) and passed(tests["bridge"]),
            "SoD, ámbito y puente gobernado con regresión enfocada; no constituye certificación universal de permisos.",
        ),
        (
            "RC04",
            backend_ok and browser_ok,
            "Recorrido técnico de la cadena actual. No promueve A0 a A26 formalmente solo por completar A27 a A31.",
        ),
        (
            "RC07",
            s2_closed,
            "Cuatro Class B y FID nuevos; cierre condicionado a currentness posterior, navegador y regresión final.",
        ),
        (
            "RC08",
            preservation_ok,
            "Capturas operacionales READ ONLY antes y después iguales, incluidas Project 14 y políticas publicadas.",
        ),
        (
            "RC09",
            chain.get("head") == "20261011_0068"
            and before is not None
            and before.get("head") == "20261011_0068",
            "Cabeza 0068 conservada; no se añade migración ni se atribuye un roundtrip nuevo.",
        ),
        (
            "RC11",
            passed(tests["frontend"]) and browser_ok,
            "Presentación estratégica actual comprobada; no se infiere conservación visual solo por inventario de archivos.",
        ),
        (
            "RC14",
            s2_closed and integration_closed,
            "UAT-S2-003 y UAT-S2-005 deben estar cerrados mediante pruebas actuales.",
        ),
        (
            "RC15",
            False,
            "Quince áreas humanas sin nombre, fecha ni evidencia de firma; ACCEPTED aislado no desbloquea la aceptación.",
        ),
    ):
        rc.append(
            {
                "id": criterion_id,
                "status": "PASS" if ok else "BLOCKED",
                "actual": actual,
                "evidence": [relative(FACTS)],
            }
        )
    return {
        "schema_version": 1,
        "baseline": baseline,
        "collected_at": datetime.now(ZoneInfo("America/Bogota")).isoformat(),
        "authorization": artifact(DATA / "rebuild-authorization.md"),
        "identity": identity,
        "selected": {
            "chain": artifact(chain_path),
            "live": artifact(live_path),
            "project": pick(
                project,
                (
                    "id",
                    "code",
                    "external_key",
                    "record_code",
                    "parent_id",
                    "status",
                    "version",
                ),
            ),
            "attempt_history": chain.get("attempt_history", []),
            "current_blocker": chain.get("blocker"),
            "backend_chain": chain.get("backend_chain", "NOT_COMPLETE"),
        },
        "controlledEvidence": references,
        "reconciliation": recon.get("rows", []),
        "gate": {
            "before": pick(
                next(
                    (
                        row.get("data", {})
                        for row in chain.get("stages", [])
                        if row.get("stage") == "GATE_BEFORE_CONTROLLED_CLASS_B"
                        and row.get("status") == "PASS"
                    ),
                    {},
                ),
                ("id", "status", "readiness_status", "blockers"),
            ),
            "readiness": recon.get("readiness", {}),
            "project_readiness": recon.get("project_readiness", {}),
            "frozen": pick(
                frozen.get("preparation", {}),
                (
                    "id",
                    "status",
                    "revision",
                    "preparation_hash",
                    "investment_gate_profile_id",
                    "profile_code",
                    "profile_revision",
                    "profile_hash",
                    "frozen_by",
                    "frozen_at",
                ),
            ),
            "handoff": frozen.get("handoff", {}),
            "policy": pick(
                frozen.get("preparation", {})
                .get("source_snapshot", {})
                .get("evidence_policy", {}),
                ("policy_id", "policy_revision", "policy_hash"),
            ),
            "source_consistency": last_stage(chain, "A28_CURRENT_SOURCE_CONSISTENCY"),
        },
        "fid": {
            "captured": pick(
                captured_fid,
                (
                    "id",
                    "tenant_id",
                    "portfolio_workspace_id",
                    "investment_gate_preparation_id",
                    "investment_gate_profile_id",
                    "status",
                    "outcome",
                    "profile_code",
                    "decision_round",
                    "decision_maker_user_id",
                    "decided_by",
                    "decided_at",
                    "decision_hash",
                    "source_preparation_hash",
                    "source_package_hash",
                    "source_profile_hash",
                    "source_policy_hash",
                    "configuration_id",
                    "configuration_revision",
                    "configuration_hash",
                    "is_fid",
                    "investment_authorized",
                    "full_funds_authorized",
                    "can_activate_project",
                ),
            ),
            "persisted": live.get("decisions", []),
        },
        "postFid": {
            "initial": initial_post,
            "rounds": post_rows,
            "current": current_post,
        },
        "initialization": pick(
            init,
            (
                "id",
                "state",
                "workspace_id",
                "revision_version",
                "started_by_user_id",
                "validated_by_user_id",
                "activated_by_user_id",
                "started_at",
                "ready_at",
                "activated_at",
                "validation_hash",
                "checklist_hash",
            ),
        ),
        "tests": tests,
        "frontendBuild": {
            "artifact": artifact(DATA / "frontend-build-validation.json"),
            "result": frontend_validation,
        },
        "browser": {
            "artifact": artifact(browser_dir / "fresh-browser-result.json"),
            "status": browser_data.get("status", "NOT_RUN"),
            "identity": browser_identity,
            "visits": browser_data.get("visits", []),
            "blocker": browser_data.get("blocker"),
            "runtime_errors": browser_data.get("runtime_errors", []),
            "http_errors": browser_data.get("http_errors", []),
            "browser_business_writes": browser_data.get("browser_business_writes", []),
            "navigation_metadata_writes": browser_data.get(
                "navigation_metadata_writes", []
            ),
            "navigation_metadata_responses": browser_data.get(
                "navigation_metadata_responses", []
            ),
            "requests": browser_data.get("requests", []),
            "governed_initialization_activation": browser_data.get(
                "governed_initialization_activation", {}
            ),
            "visual_qa_summary": {
                "artifact": artifact(browser_dir / "visual-qa-summary.json"),
                "result": visual_qa,
            },
            "home_consumer_audit": artifact(
                browser_dir / "home-context-consumer-audit.md"
            ),
            "verified_screenshots": verified_screenshots,
            "observations": visual_qa.get("observations", []),
            "residual_observations": visual_qa.get("residual_observations", []),
            "cosmetic_findings": browser_data.get(
                "cosmetic_findings", visual_qa.get("cosmetic_findings", [])
            ),
            "layout_findings": browser_data.get(
                "layout_findings", visual_qa.get("layout_findings", [])
            ),
            "limitations": browser_data.get(
                "limitations", visual_qa.get("limitations", [])
            ),
            "screenshots": browser_data.get("screenshots", []),
            "same_chain_pass": browser_ok,
            "same_chain_functional_pass": browser_functional_ok,
            "visual_review_completed": browser_visual_ok,
            "previous_attempts": browser_history,
        },
        "preservation": {
            "status": "PASS_READ_ONLY_IDENTICAL"
            if preservation_ok
            else "PENDING_OR_DIFFERENT",
            "before": artifact(QA / "operational-before.json"),
            "after": artifact(QA / "operational-after.json"),
            "project14": (before or {}).get("workspace14", {}),
            "persistence_before": artifact(QA / "persistence-before.json"),
            "persistence_after": artifact(QA / "persistence-after.json"),
            "named_volume_restart_verified": persistence_ok,
        },
        "backups": [load(p) for p in sorted((QA / "backups").glob("*.json"))],
        "cases": cases,
        "defects": defects,
        "releaseCriteria": rc,
        "remainingObservations": visual_qa.get("residual_observations", []),
        "release": release,
        "counting_rule": "Cada XML se informa por separado. Los focused y fixtures independientes no se suman a los 914 casos de la única corrida broad limpia final.",
    }


def report_sections(f):
    identity, gate, fid = f["identity"], f["gate"], f["fid"]["captured"]
    release, current = f["release"], f["postFid"]["current"]
    init, tests = f["initialization"], f["tests"]
    case = {r["id"]: r for r in f["cases"]}
    defect = {r["id"]: r for r in f["defects"]}
    frontend = f.get("frontendBuild", {}).get("result", {})
    browser_timing = (
        f["browser"]
        .get("visual_qa_summary", {})
        .get("result", {})
        .get("functional_facts", {})
    )
    browser_run_seconds = browser_timing.get("run_seconds")
    browser_run_text = (
        f"{browser_run_seconds:.3f}"
        if isinstance(browser_run_seconds, (int, float))
        else "no expuesto"
    )
    return [
        (
            "1 Resumen ejecutivo",
            f"El estado técnico es {release['status']} / {release['decision']}. Las cuatro referencias Class B, el freeze y la decisión FID se acreditan con registros de la nueva cadena. La activación registrada es {init.get('state', 'no verificada')}; el resultado actual posterior a ella es {current.get('eligibility_status', 'no verificado')}. Esta conclusión se limita al objetivo técnico Class B/FID/Post-FID: no certifica toda la interfaz ni la metadata de Home. Los hallazgos menores y UAT-S3-006 de proyección y descubrimiento siguen abiertos. Las condiciones pendientes no se sustituyen por resultados esperados.",
        ),
        (
            "2 Línea base",
            f"La reconstrucción fue autorizada expresamente el 16 de octubre de 2026 tras perderse la base temporal anterior. Rama {f['baseline']['branch']}; commit de referencia {f['baseline']['commit']}; cabeza capturada {f['baseline']['alembic_head_captured']}, sin migración nueva. Los informes bloqueados del 15 de octubre permanecen intactos. No se ejecuta commit, push ni despliegue operativo como parte de este cierre.",
        ),
        (
            "3 Identidad de la cadena seleccionada",
            f"La base PostgreSQL strategic_fresh_uat usa almacenamiento persistente dedicado y respaldos custom verificables. La identidad actual es Tenant {identity.get('tenant_id')} / {identity.get('tenant_slug')}, Portfolio {identity.get('portfolio_id')} y Project {identity.get('project_id')}. No corresponde al Tenant 4 / Project 16 perdidos y no mezcla sus pruebas.",
        ),
        (
            "4 Estado y bloqueos de Gate07J",
            f"El paquete antes del freeze alcanzó {gate['readiness'].get('readiness_status', 'no verificado')}, con {gate['readiness'].get('valid_requirement_count')} requisitos válidos de {gate['readiness'].get('required_requirement_count')}. El paquete persistido es {gate['frozen'].get('status', 'no verificado')}. Los bloqueos actuales Post-FID son {current.get('blockers_json', current.get('blockers', []))}; se distinguen de las referencias Class B que faltaban inicialmente.",
        ),
        (
            "5 Arquitectura de evidencias reutilizada",
            "Se reutilizan ManualEvidenceIn, evidence_refs, add_evidence y validate_evidence del mismo Investment Gate Preparation. El alta registra SUBMITTED y la revisión independiente registra VALID; la preparación conserva revision_version, hashes y alcance. No se crean proveedores ni dominios de permisos, contratación, ejecución o madurez del cronograma.",
        ),
        (
            "6 Referencia controlada de permisos",
            "El bloqueo literal MISSING_PERMIT_READINESS correspondía al requirement permit-readiness Class B con provider_key=null; no indicaba que el assessment gobernado Permit estuviera bloqueado: ese assessment ya estaba READY. Tampoco requería un nuevo adapter. La referencia controlada conservó EACH_PROJECT y 180 días, pasó de SUBMITTED a VALID por revisión independiente y eliminó ese bloqueo. El assessment Permit se usa solo como soporte; su aprobación no sustituye la validación de la referencia Gate.",
        ),
        (
            "7 Referencia controlada de contratación",
            "contracting-strategy-readiness conserva Class B y 365 días. El documento representativo referencia el identificador, versión y hash de Commercial Strategy como soporte; la revisión del Gate sigue siendo un acto separado. No se declara un contrato real firmado.",
        ),
        (
            "8 Referencia controlada del plan de ejecución",
            "execution-plan-readiness utiliza el mecanismo de referencia documental existente, con revisión, fecha, Project, propietario y hash. Su documento identifica expresamente el carácter representativo UAT y una ventana de 365 días. No constituye aprobación para ejecutar Project Delivery.",
        ),
        (
            "9 Referencia controlada de madurez del cronograma",
            "schedule-maturity identifica la representación de cronograma, revisión, fecha de datos, dimensiones de evaluación y autoridad de revisión UAT. La ventana es 180 días. No se equipara automáticamente Schedule Source Readiness con esta evaluación ni se emite una certificación P6.",
        ),
        (
            "10 Presentación y validación independiente",
            "Las cuatro referencias conservan los actores y timestamps reales de presentación y validación. El documento inmutable se verifica mediante SHA-256 antes de la revisión. La guarda SoD ahora aplica a los requisitos Class B con manual_validation_required, aunque no figuren en la lista histórica FEL y Funding del perfil.",
        ),
        (
            "11 Vigencia y caducidad",
            "Se reprodujeron y corrigieron dos bypass: fecha futura y validez superior a la ventana capturada. Add y update los rechazan antes de persistir; validate conserva una defensa adicional. La caducidad legítima continúa produciendo EXPIRED y no aporta cobertura. Los paquetes frozen no se reinterpretan históricamente.",
        ),
        (
            "12 Cobertura de cada Project",
            f"La cobertura capturada es {gate['readiness'].get('coverage_percent', 'desconocida')}%, con {gate['readiness'].get('valid_project_count')} de {gate['readiness'].get('required_project_count')} Projects válidos. Las referencias apuntan al Project {identity.get('project_id')} y al hash real del conjunto recomendado. No se promedian omisiones ni se sustituye EACH_PROJECT por evidencia de Portfolio.",
        ),
        ("13 Readiness A27", case["A27"]["actual"]),
        (
            "14 Conciliación de los diecinueve requisitos",
            "La matriz conserva 20 filas: 19 required, de las cuales 15 son VALID y cuatro permanecen MISSING no bloqueantes según la política: environmental-social-readiness, land-access-readiness, stakeholder-readiness y technical-assurance. La fila opcional Class C other-governed-evidence queda NOT_REQUIRED. No se declara 19 de 19 VALID ni se omiten advertencias. Fuente/as-of: checkpoint A27_ALL_REQUIREMENTS_RECONCILIATION del manifest, antes del freeze "
            + str(gate["frozen"].get("frozen_at", "sin fecha expuesta"))
            + " UTC. A27 no expone timestamp propio. No es una lectura viva posterior a la activación; null currentness se muestra como No expone, no como False.",
        ),
        (
            "15 Consistencia entre Business Case y Gate",
            case["A28"]["actual"]
            + " Se conservan los IDs y hashes capturados de Business Case, FIN, Benefits, Commercial, FEL, Risk, Funding, Permit y QRA cuando aplica. Las nuevas referencias controladas no sustituyen esas fuentes.",
        ),
        (
            "16 Paquete frozen A29",
            case["A29"]["actual"]
            + " La instantánea incluye perfil, política, criterios, evidencias, validaciones, ámbito y fuentes. El handoff mantiene can_enter_fid=false; la autoridad formal permanece en Gate07K.",
        ),
        (
            "17 Inmutabilidad del paquete",
            f"La prueba enfocada de sucesor controlado forma parte de la corrida Class B {count_text(tests['class_b'])}. En una fixture mínima aislada se clonó el paquete, actualizó y validó la referencia sucesora por comandos oficiales; el modelo íntegro del paquete frozen original permaneció idéntico. Esta prueba no altera las fuentes de la cadena seleccionada.",
        ),
        (
            "18 Decisión FID A30",
            case["A30"]["actual"]
            + f" La autoridad registrada es el actor {fid.get('decided_by')}, distinto del preparador y del revisor. Se conserva la ronda {fid.get('decision_round')} y la fecha {fid.get('decided_at')} UTC como decisión histórica real.",
        ),
        (
            "19 Indicadores y hashes FID",
            f"is_fid={fid.get('is_fid')}, investment_authorized={fid.get('investment_authorized')}, full_funds_authorized={fid.get('full_funds_authorized')} y can_activate_project={fid.get('can_activate_project')}. El hash de decisión y el hash del paquete fuente se conservan en la tabla de trazabilidad. El registro DECIDED no se borra cuando una evaluación de currentness posterior detecta un bloqueo.",
        ),
        (
            "20 Gobierno posterior a FID A31",
            case["A31"]["actual"]
            + " Se separan las rondas históricas que habilitaron inicialización y activación de la evaluación viva posterior. Un resultado STALE no significa que el comando de activación haya fallado.",
        ),
        (
            "21 Inicialización y activación",
            f"El comando real dejó Project {identity.get('project_id')} {f['selected']['project'].get('status')}, versión {f['selected']['project'].get('version')}, con estado {init.get('state')}. Inicializó el actor {init.get('started_by_user_id')} y activó {init.get('activated_by_user_id')}; sus identidades son distintas. El primer bloqueo de inicialización y el checkpoint prematuro del harness se preservan como intentos anteriores. No se inicia ejecución operativa.",
        ),
        (
            "22 Disposición UAT S2 003",
            f"{defect['UAT-S2-003']['status']}. {defect['UAT-S2-003']['actual']} La autorización de reconstrucción cambia únicamente el prerequisito de la base perdida, no las políticas de evidencia ni las condiciones de cierre.",
        ),
        (
            "23 Reconciliación de la prueba FID heredada",
            f"Clasificación explícita Category A: fixture desactualizada, prueba válida. La fixture positiva fue modernizada para consumir las fuentes actuales de Five Case y referencias controladas, sin relajar la prueba ni la política. La prueba enfocada independiente obtuvo {count_text(tests['legacy'])}. Su resultado no se suma a la corrida broad ni se usa como identidad alternativa de la UAT seleccionada. No se añade xfail para ocultar el fallo anterior.",
        ),
        (
            "24 Presentación de Business Case y Funding",
            f"{defect['UAT-S3-005']['status']}. {defect['UAT-S3-005']['actual']} Regresión frontend: {count_text(tests['frontend'])}. Build de producción {frontend.get('result', 'no verificado')} en {frontend.get('reported_build_seconds', 'tiempo no verificado')} s; TypeScript exit {frontend.get('typescript_no_emit_exit_code', 'no verificado')}; ESLint {frontend.get('eslint_errors', 'no verificado')} errores y {frontend.get('eslint_warnings', 'no verificado')} advertencias. Las advertencias BIM preexistentes quedan fuera de este incremento. La pantalla debe mostrar currentness autoritativa sin alterar el dominio financiero.",
        ),
        (
            "25 Lifecycle histórico de Proposal",
            f"{defect['UAT-S4-006']['status']}. {defect['UAT-S4-006']['actual']} Una restricción de reingreso posterior no cancela la aprobación Strategic Gate histórica.",
        ),
        (
            "26 Regresión estratégica global",
            f"Única corrida final seleccionada: {count_text(tests['broad'])}; duración XML {tests['broad'].get('seconds', 'no completada')} s. Se conservan {tests['broad'].get('total', 'sin conteo')} entradas y {tests['broad'].get('unique_total', 'sin conteo')} identidades únicas, mediante comparación ordinal sensible a mayúsculas: el multiset coincide con la baseline, diferencia {tests['broad'].get('identity_multiset_difference_from_baseline')}. Los parámetros project y PROJECT se consideran distintos; duplicados heredados detectados: {len(tests['broad'].get('inherited_duplicate_identities', []))}. Los 16 SKIP son 13 específicos PostgreSQL, dos perfiles opt-in y un contrato CI no empaquetado en la imagen. Este último se ejecutó aparte desde la raíz read-only: {count_text(tests['ci_contract_guard'])}, {tests['ci_contract_guard'].get('seconds', 'no completado')} s, sin sumarlo a los 914 ni convertir el SKIP original en PASS. Los intentos interrumpidos y el setup inicial del guard se conservan sin aportar conteos finales. {f['counting_rule']}",
        ),
        (
            "27 Verificación PostgreSQL y persistencia",
            f"Los comandos del recorrido actual se ejecutaron en PostgreSQL 16 separado de la base operativa. Los dumps se conservan con tamaño, fecha, SHA-256 y verificación de listado. Persistencia tras reiniciar únicamente los servicios UAT: named_volume_restart_verified={f['preservation']['named_volume_restart_verified']}. El bootstrap de compatibilidad hasta 0067 seguido del upgrade real 0068 no acredita un roundtrip de todas las migraciones históricas.",
        ),
        (
            "28 Verificación de navegador",
            f"Chromium actual: {count_text(tests['browser'])}; resultado {f['browser']['status']}; TOTAL JUnit {browser_run_text} s y testcase {browser_timing.get('testcase_seconds', 'no expuesto')} s, sin sumar la duración raíz a la del caso. La identidad, las vistas visitadas, errores y escrituras se conservan en el JSON de navegador. No se presenta como una sesión totalmente read-only: el GET de WorkspaceContext registra la auditoría normal SecurityEvent; el POST existente de apertura actualiza metadatos RecentWorkspace y auditoría de navegación. Estos metadatos se distinguen de los business writes, que no se permiten desde el navegador; los comandos de negocio pertenecen al recorrido backend. No se reutilizan capturas de otro tenant. "
            + (
                "Intentos previos conservados por separado, sin sumarlos al resultado actual: "
                + "; ".join(
                    f"{count_text(row['test'])} / {row['classification']}"
                    for row in f["browser"].get("previous_attempts", [])
                )
                + "."
                if f["browser"].get("previous_attempts")
                else "No se agrega un intento parcial como PASS."
            ),
        ),
        (
            "29 Seguridad y segregación de funciones",
            f"Class B: {count_text(tests['class_b'])}; puente Post-FID: {count_text(tests['bridge'])}; guardas canónicas: {count_text(tests['canonical'])}; currentness posterior: {count_text(tests['postactivation'])}. Los saltos de ámbito, actor, configuración, hash o vigencia no se aceptan como autorización. Los SKIP se mantienen explícitos y no equivalen a PASS.",
        ),
        (
            "30 Preservación del Project operativo 14",
            f"La comparación actual es {f['preservation']['status']}. La referencia protegida es PRJ-PPMIS, external_key PRJ-PYP-PPMIS, active, versión 1 y MD5 ef67b78a4d4902d2dc6db3efb017b80b. Las observaciones before y after actuales coinciden, incluidas la cabeza, las políticas publicadas y los conteos protegidos; no se sustituye el after por una captura histórica.",
        ),
        (
            "31 Hallazgos restantes",
            "El registro siguiente distingue el bloqueo de entorno resuelto por reconstrucción autorizada, los hallazgos de presentación y la integración de lifecycle. UAT-S2-005 no se cierra únicamente porque activar haya funcionado; exige currentness correcta posterior, pruebas del fix y navegador de la misma identidad. "
            + RETAINED_MINOR_TEXT
            + " Lecturas literales actuales: el header de Gate muestra BLOCKED Governance y su cuerpo muestra READY_FOR_INVESTMENT_DECISION, Source CURRENT, Decision entry READY y can_enter_fid=false. En Post-FID las columnas Initialization Status y Activation Status muestran guiones mientras la API registra initialization_state=ACTIVATED, can_initialize=false y can_activate=false. Se documentan ambas proyecciones sin inventar una explicación de su semántica ni afirmar que la UI sea impecable; los guiones no se convierten en una activación fallida."
            + (
                " Además, UAT-S3-006 permanece OPEN / S3 MAJOR por proyección de responsable y plantilla en Home y posible omisión del Project al filtrar por responsable. Es distinto de UAT-S4-003. El triage posterior y la auditoría de consumidores conservan la observación original, sin cambiarla: no se halló acoplamiento a autorización, lifecycle, elegibilidad o selección de Project en los consumidores actuales revisados. Workaround: Overview canónico y ruta directa sin filtro por responsable. No se afirma que esté corregido ni que se hayan probado clientes externos."
                if f.get("remainingObservations")
                else ""
            ),
        ),
        (
            "32 Decisión técnica de liberación",
            f"{release['status']} / {release['decision']}. Condiciones verificadas: {', '.join(k for k, v in release['conditions'].items() if v)}. Condiciones pendientes: {', '.join(k for k, v in release['conditions'].items() if not v) or 'ninguna técnica del alcance'}. Esta conclusión se limita al objetivo técnico Class B/FID/Post-FID; no certifica toda la metadata de Home, la limpieza completa de UI ni producción. UAT-S3-006 y los menores abiertos se conservan con sus límites de revisión. No se declara STRATEGIC_MACROPROCESS_READY_FOR_BUSINESS_RELEASE.",
        ),
        (
            "33 Aceptación humana",
            "PENDING_NOT_ASSUMED. Las quince áreas de aceptación mantienen vacíos el firmante, la fecha y la evidencia. Una prueba técnica aprobada o una celda ACCEPTED aislada no constituye aceptación de negocio ni autorización de despliegue.",
        ),
        ("34 Recomendación siguiente", release["next_recommendation"]),
    ]


def add_table(doc, headers, rows, widths):
    style.table(doc, headers, rows, widths)
    for cell in doc.tables[-1].rows[0].cells:
        cell.paragraphs[0].paragraph_format.keep_with_next = True


def technical(f):
    doc = Document(SOURCE / NAMES["closeout"])
    body = doc._element.body
    logo = deepcopy(body[0]) if body[0].xpath(".//w:drawing") else None
    for element in list(body):
        if element.tag != qn("w:sectPr"):
            body.remove(element)
    if logo is not None:
        body.insert(0, logo)
    doc.add_paragraph("Informe técnico de evidencias controladas Gate07J", "Title")
    style.p(
        doc,
        "P&Pmis Construction AI. Capital Portfolio & Investment Governance. Corte del 16 de octubre de 2026.",
    )
    style.p(
        doc,
        f"{f['release']['status']} / {f['release']['decision']}. Este informe presenta la reconstrucción autorizada persistente, el cierre de Class B y la comprobación de FID y Post-FID. La aceptación humana permanece pendiente.",
    )
    for heading, paragraph in report_sections(f):
        style.h(doc, heading)
        style.p(doc, paragraph)
        number = int(heading.split()[0])
        if number == 3:
            identity = f["identity"]

            def identities(*keys, values=identity):
                return " / ".join(str(values.get(key, "No capturado")) for key in keys)

            add_table(
                doc,
                ["Identidad", "Valor actual"],
                [
                    ["Tenant y slug", identities("tenant_id", "tenant_slug")],
                    [
                        "Enterprise / BU / Portfolio / Project",
                        identities(
                            "enterprise_id",
                            "business_unit_id",
                            "portfolio_id",
                            "project_id",
                        ),
                    ],
                    [
                        "Capital cycle / demand / allocation",
                        identities(
                            "capital_cycle_id",
                            "capital_demand_id",
                            "allocation_plan_id",
                        ),
                    ],
                    [
                        "Idea / Business Analysis / Proposal",
                        identities("idea_id", "analysis_id", "proposal_id"),
                    ],
                    [
                        "Strategic Gate / Creation Request",
                        identities("strategic_gate_id", "project_creation_request_id"),
                    ],
                    [
                        "Autor / revisor / FID / inicializador / activador",
                        identities(
                            "actor_id",
                            "reviewer_id",
                            "decision_maker_id",
                            "initializer_id",
                            "activator_id",
                        ),
                    ],
                    [
                        "Admin autor / publicador / endorser capital",
                        identities(
                            "admin_author_id",
                            "admin_publisher_id",
                            "capital_endorser_id",
                        ),
                    ],
                ],
                [2.8, 4.26],
            )
        if number == 10:
            add_table(
                doc,
                [
                    "Requisito",
                    "Documento y revisión",
                    "Presentó UTC",
                    "Validó UTC",
                    "Estado",
                ],
                [
                    [
                        r["requirement_key"],
                        r["reference"].get("reference_id", "No verificado") + " / r1",
                        str(r["evidence"].get("submitted_by"))
                        + "\n"
                        + str(r["evidence"].get("submitted_at", ""))[:19].replace(
                            "T", " "
                        ),
                        str(r["evidence"].get("validated_by"))
                        + "\n"
                        + str(r["evidence"].get("validated_at", ""))[:19].replace(
                            "T", " "
                        ),
                        r["evidence"].get("status", "No verificado"),
                    ]
                    for r in f["controlledEvidence"]
                ],
                [1.7, 2.3, 1.0, 1.25, 0.81],
            )
        if number == 11:
            add_table(
                doc,
                ["Referencia UAT", "Fecha y validez UTC", "SHA256 del documento"],
                [
                    [
                        r["reference"].get("reference_id"),
                        "Evidencia "
                        + str(r["evidence"].get("evidence_date", ""))[:19].replace(
                            "T", " "
                        )
                        + "\nLímite "
                        + str(r["evidence"].get("valid_until", ""))[:19].replace(
                            "T", " "
                        ),
                        r["reference"].get("document_sha256"),
                    ]
                    for r in f["controlledEvidence"]
                ],
                [1.8, 2.2, 3.06],
            )
        if number == 4:
            add_table(
                doc,
                ["Referencia publicada", "ID y revisión", "Hash capturado"],
                [
                    [
                        "Política de evidencia",
                        str(f["gate"]["policy"].get("policy_id"))
                        + " / r"
                        + str(f["gate"]["policy"].get("policy_revision")),
                        f["gate"]["policy"].get("policy_hash"),
                    ],
                    [
                        "Perfil PYP FFA",
                        str(f["gate"]["frozen"].get("investment_gate_profile_id"))
                        + " / r"
                        + str(f["gate"]["frozen"].get("profile_revision")),
                        f["gate"]["frozen"].get("profile_hash"),
                    ],
                ],
                [1.8, 1.3, 3.96],
            )
        if number == 14:
            add_table(
                doc,
                ["Requisito y control", "Estado conciliado y cobertura"],
                [reconciliation_record(row) for row in f["reconciliation"]],
                [3.0, 4.06],
            )
        if number == 19:
            add_table(
                doc,
                ["Referencia", "Hash capturado"],
                [
                    [k, f["fid"]["captured"].get(k, "No verificado")]
                    for k in (
                        "decision_hash",
                        "source_preparation_hash",
                        "source_package_hash",
                        "source_profile_hash",
                        "source_policy_hash",
                    )
                ],
                [2.5, 4.56],
            )
        if number == 26:
            add_table(
                doc,
                ["Corrida", "Resultado sin sumar reruns", "XML de evidencia"],
                [
                    [
                        k,
                        count_text(v),
                        v["file"].removeprefix(
                            "outputs/.qa/gate07j-controlled-class-b-rebuild/"
                        ),
                    ]
                    for k, v in f["tests"].items()
                ],
                [1.2, 2.7, 3.16],
            )
        if number == 31:
            add_table(
                doc,
                ["Hallazgo", "Disposición", "Resultado actual"],
                [[r["id"], r["status"], r["actual"]] for r in f["defects"]],
                [1.1, 2.05, 3.91],
            )
    doc.core_properties.title = "Evidencias controladas Gate07J FID y Post FID"
    doc.core_properties.revision = 4
    return doc


def uat(f):
    doc = Document(SOURCE / NAMES["uat"])
    doc.add_page_break()
    style.h(
        doc, "Resultados de la reconstrucción persistente del 16 de octubre de 2026"
    )
    style.p(
        doc,
        f"El estado del incremento es {f['release']['status']} / {f['release']['decision']}. Las 37 páginas anteriores permanecen como historia y no acreditan un PASS nuevo. La reconstrucción fue autorizada; la identidad seleccionada ahora es Tenant {f['identity'].get('tenant_id')} / Project {f['identity'].get('project_id')}.",
    )
    style.p(
        doc,
        "Los siguientes casos distinguen ejecución backend, aceptación formal y currentness posterior. La activación real no se denomina fallida por un refresh posterior STALE. Tampoco se promueven A0 a A26 formalmente solo por completar esta continuación.",
    )
    for case in f["cases"]:
        style.h(doc, f"Caso {case['id']}", 2)
        for label in ("preconditions", "steps", "expected", "actual"):
            style.p(
                doc,
                case[label],
                {
                    "preconditions": "Precondiciones",
                    "steps": "Pasos ejecutados",
                    "expected": "Resultado esperado",
                    "actual": "Resultado real",
                }[label]
                + ".",
            )
        style.p(
            doc,
            f"Estado formal {case['status']}; ejecución backend {case['backend_status']}. Evidencia: "
            + "; ".join(case["evidence"]),
        )
    style.h(doc, "Pruebas del incremento", 2)
    add_table(
        doc,
        ["Corrida independiente", "PASS FAIL ERROR SKIP"],
        [[k, count_text(v)] for k, v in f["tests"].items()],
        [2.4, 4.66],
    )
    style.p(doc, f["counting_rule"])
    style.h(doc, "Disposición de hallazgos y criterios", 2)
    for defect in f["defects"]:
        style.p(doc, f"{defect['id']}: {defect['status']}. {defect['actual']}")
    style.p(doc, RETAINED_MINOR_TEXT)
    for residual in f.get("remainingObservations", []):
        style.p(
            doc,
            f"Residual {residual.get('defect_id', 'sin ID')} / {residual.get('severity', 'sin severidad')} / {residual.get('status', 'sin estado')}: "
            + residual.get("description", "No expone descripción")
            + " La observación histórica y su triage posterior se conservan; no se considera resuelto.",
        )
    style.p(
        doc,
        "Las firmas humanas permanecen vacías. No se autoriza liberación de negocio, despliegue operativo, commit ni push mediante este informe técnico.",
    )
    style.p(doc, f["release"]["next_recommendation"])
    doc.core_properties.revision = 4
    return doc


def save_preserving(source, destination, doc):
    """Only body, descriptive core and the exact footer date may change."""
    xml = etree.tostring(
        doc._element, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    with ZipFile(source) as original, ZipFile(destination, "w") as result:
        core = etree.fromstring(original.read("docProps/core.xml"))
        core.find(
            "{http://purl.org/dc/elements/1.1/}title"
        ).text = doc.core_properties.title
        core.find(
            "{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}revision"
        ).text = str(doc.core_properties.revision)
        replacements = {
            "word/document.xml": xml,
            "docProps/core.xml": etree.tostring(
                core, xml_declaration=True, encoding="UTF-8", standalone=True
            ),
        }
        footer_parts = [
            name
            for name in original.namelist()
            if re.fullmatch(r"word/footer\d+\.xml", name)
        ]
        revised_footers = 0
        for name in footer_parts:
            raw = original.read(name)
            dates = [
                date for date in (b"14 octubre 2026", b"15 octubre 2026") if date in raw
            ]
            if dates:
                assert len(dates) == 1 and raw.count(dates[0]) == 1
                replacements[name] = raw.replace(dates[0], b"16 octubre 2026", 1)
                revised_footers += 1
        assert revised_footers, "Expected retained report footer date not found"
        for info in original.infolist():
            result.writestr(
                info, replacements.get(info.filename, original.read(info.filename))
            )
    with ZipFile(source) as original, ZipFile(destination) as result:
        assert original.namelist() == result.namelist()
        changed_parts = {
            n for n in original.namelist() if original.read(n) != result.read(n)
        }
        assert changed_parts <= {
            "word/document.xml",
            "docProps/core.xml",
            *footer_parts,
        }
        for name in set(footer_parts) & changed_parts:
            assert result.read(name) == replacements[name]


def validate_facts_artifacts(value):
    if isinstance(value, dict):
        if value.get("file") and value.get("sha256"):
            recorded = Path(value["file"])
            path = recorded if recorded.is_absolute() else ROOT / recorded
            assert path.is_file() and digest(path) == value["sha256"], (
                f"Evidence changed since collection: {path}"
            )
        for item in value.values():
            validate_facts_artifacts(item)
    elif isinstance(value, list):
        for item in value:
            validate_facts_artifacts(item)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("--collect", action="store_true")
    parser.add_argument("--documents", action="store_true")
    parser.add_argument("--facts", type=Path, default=FACTS)
    parser.add_argument("--live", type=Path)
    parser.add_argument("--broad-dir", type=Path)
    parser.add_argument("--browser-dir", type=Path)
    parser.add_argument("--postactivation-xml", type=Path)
    parser.add_argument("--private-content", action="store_true")
    args = parser.parse_args()
    if args.collect or not args.documents:
        facts = collect(args)
        args.facts.parent.mkdir(parents=True, exist_ok=True)
        args.facts.write_text(
            json.dumps(facts, indent=2, ensure_ascii=False, default=str),
            encoding="utf-8",
        )
    else:
        facts = load(args.facts)
    if args.private_content:
        private = QA / "report-authoring"
        private.mkdir(parents=True, exist_ok=True)
        (private / "technical-content.json").write_text(
            json.dumps(report_sections(facts), indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
    if args.documents:
        validate_facts_artifacts(facts)
        assert all(
            digest(ROOT / path) == sha
            for path, sha in {
                **facts["baseline"]["product_hashes_at_collection"],
                **facts["baseline"].get("fixture_hashes_at_collection", {}),
            }.items()
        ), "Product changed since evidence collection"
        retained = load(DATA / "document-verification.json")["artifacts"]
        expected = {Path(item["file"]).name: item["sha256"] for item in retained}
        assert len(report_sections(facts)) == 34
        for kind in ("closeout", "uat"):
            assert digest(SOURCE / NAMES[kind]) == expected[NAMES[kind]]
        OUT.mkdir(parents=True, exist_ok=True)
        for kind, doc in (("closeout", technical(facts)), ("uat", uat(facts))):
            save_preserving(SOURCE / NAMES[kind], OUT / NAMES[kind], doc)
            assert digest(SOURCE / NAMES[kind]) == expected[NAMES[kind]]
    print(
        json.dumps(
            {
                "facts": str(args.facts),
                "release": facts["release"],
                "documents_written": args.documents,
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
