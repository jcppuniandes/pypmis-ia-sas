from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(r"C:\Users\Ricardo\Documents\GitHub\pypmis-ia-sas")
OUTPUT = ROOT / "outputs" / "Informe_Tecnico_Navegacion_Estrategica_Frontend_V2_2026-09-10.docx"
LOGO = ROOT / "frontend" / "public" / "pypmis-construction-ai-logo.png"
EVIDENCE = ROOT / "docs" / "validation" / "strategic-navigation-v2"

NAVY = "12212B"
TEAL = "087F80"
LIGHT_TEAL = "EAF5F4"
LIGHT_GRAY = "F3F6F7"
MID_GRAY = "D9E1E5"
TEXT = "17222B"
MUTED = "566874"
WHITE = "FFFFFF"
GREEN = "18794E"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=MID_GRAY, size="5") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def set_run_font(run, size=None, bold=None, color=None, name="Aptos") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc, text: str, level: int = 1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_paragraph(doc, text: str = "", *, bold_lead: str | None = None, style=None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08
    if bold_lead and text.startswith(bold_lead):
        first = paragraph.add_run(bold_lead)
        first.bold = True
        paragraph.add_run(text[len(bold_lead) :])
    else:
        paragraph.add_run(text)
    return paragraph


def add_bullets(doc, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_numbered(doc, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for index, (cell, label) in enumerate(zip(header.cells, headers)):
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            cell.width = Inches(widths[index])
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(label)
        set_run_font(run, 9, True, WHITE)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        prevent_row_split(row)
        for index, (cell, value) in enumerate(zip(row.cells, values)):
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else LIGHT_GRAY)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if widths:
                cell.width = Inches(widths[index])
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(value))
            set_run_font(run, 8.5, False, TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_status_table(doc, rows: list[list[str]]):
    table = add_table(doc, ["Criterio", "Estado", "Evidencia"], rows, [2.45, 0.7, 3.35])
    for row in table.rows[1:]:
        status_cell = row.cells[1]
        set_cell_shading(status_cell, LIGHT_TEAL)
        for run in status_cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(GREEN)
    return table


def add_caption(doc, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    set_run_font(run, 8.5, False, MUTED)
    run.italic = True


def add_picture_fit(run, image_path: Path, max_width: float = 6.2, max_height: float = 8.25) -> None:
    with Image.open(image_path) as source:
        aspect = source.width / source.height
    width = min(max_width, max_height * aspect)
    run.add_picture(str(image_path), width=Inches(width))


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(TEXT)
for name, size, color in (("Title", 28, NAVY), ("Heading 1", 18, NAVY), ("Heading 2", 13, TEAL), ("Heading 3", 11, NAVY)):
    style = styles[name]
    style.font.name = "Aptos Display" if name != "Normal" else "Aptos"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), style.font.name)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_with_next = True

for section_item in doc.sections:
    header = section_item.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("P&Pmis Construction AI  |  Informe técnico")
    set_run_font(run, 8, False, MUTED)
    footer = section_item.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_paragraph.add_run("Navegación estratégica frontend V2  •  Página ")
    set_run_font(run, 8, False, MUTED)
    add_field(footer_paragraph, "PAGE")

# Cover
if LOGO.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run().add_picture(str(LOGO), width=Inches(0.8))
doc.add_paragraph().paragraph_format.space_after = Pt(18)
title = doc.add_paragraph(style="Title")
title.paragraph_format.space_after = Pt(8)
title.add_run("Informe técnico")
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(22)
run = subtitle.add_run("Frontend Strategic Workflow Navigation V2")
set_run_font(run, 20, True, TEAL)
add_paragraph(doc, "P&Pmis Construction AI", bold_lead="P&Pmis Construction AI")
add_paragraph(doc, "Repositorio: pypmis-ia-sas")
add_paragraph(doc, "Rama: main")
add_paragraph(doc, "Línea base: 9cb0969")
add_paragraph(doc, "Fecha de cierre: 10 de septiembre de 2026")
doc.add_paragraph().paragraph_format.space_after = Pt(26)
opening = doc.add_paragraph()
opening.paragraph_format.space_after = Pt(10)
opening.paragraph_format.line_spacing = 1.15
run = opening.add_run(
    "Se implementó una navegación estratégica gobernada que separa la Fundación Empresarial de las transacciones y ordena las fases 1–7 según el flujo de negocio. La solución reutiliza servicios, rutas, permisos y pantallas existentes; no introduce cambios de backend ni migraciones."
)
set_run_font(run, 12, False, TEXT)
conclusion = doc.add_paragraph()
conclusion.paragraph_format.space_before = Pt(14)
run = conclusion.add_run("Conclusión principal: ")
set_run_font(run, 11, True, NAVY)
run = conclusion.add_run(
    "el frontend queda alineado con la Fundación Empresarial, la secuencia estratégica, las ramas paralelas de escenarios y la convergencia de evidencia previa a la decisión de inversión."
)
set_run_font(run, 11, False, TEXT)
doc.add_page_break()

# Contents and executive summary
add_heading(doc, "Contenido", 1)
add_numbered(
    doc,
    [
        "Resumen ejecutivo y alcance",
        "Diagnóstico de la navegación anterior",
        "Arquitectura final Phase 0 / Phase 1+",
        "Jerarquías USER MODE y ADMIN MODE",
        "Blueprint, rutas, permisos y contexto",
        "Implementación y archivos",
        "Verificación y evidencias",
        "Matriz de aceptación",
        "Comparación visual antes/después",
        "Conclusiones y señal de cierre",
    ],
)
add_heading(doc, "1. Resumen ejecutivo y alcance", 1)
add_paragraph(
    doc,
    "El objetivo fue reorganizar Enterprise Strategy Manager para expresar la lógica empresarial real sin reimplementar capacidades ya entregadas. La Fundación Empresarial se presenta como Fase 0 y prerrequisito; Idea & Demand se presenta como Fase 1 y primer punto transaccional. Las fases posteriores respetan propuesta, gate estratégico, planeación, escenarios, definición y gobierno de inversión.",
)
add_bullets(
    doc,
    [
        "Una fuente de verdad tipada genera las jerarquías USER y ADMIN.",
        "Budget Planning y Resource Planning aparecen como ramas paralelas que convergen en una recomendación.",
        "FEL, Business Case, Funding, FIN, Risk y Permit aparecen como aportantes de evidencia hacia Gate07J.",
        "Los enlaces profundos y breadcrumbs se derivan del mismo blueprint.",
        "Los modos CAPITAL_OWNER, CONTRACTOR_DELIVERY, DIRECT_INTERNAL y legacy conservan sus rutas aplicables.",
        "La creación alternativa de proyectos se conserva en Project Workspace Governance, fuera del flujo universal de Ideas.",
    ],
)
add_heading(doc, "Límites respetados", 2)
add_paragraph(
    doc,
    "El cambio es frontend. No se agregaron estados de negocio, políticas de autorización, endpoints, tablas, migraciones ni automatismos de publicación o activación. La elegibilidad y las mutaciones siguen siendo autoridad del backend existente.",
)

# Baseline
doc.add_page_break()
add_heading(doc, "2. Diagnóstico de la navegación anterior", 1)
add_paragraph(
    doc,
    "La línea base agrupaba capacidades por evolución técnica y no por secuencia de negocio. El resultado era funcional, pero dificultaba reconocer los prerrequisitos, el orden de decisiones y las convergencias.",
)
add_table(
    doc,
    ["Área", "Situación anterior", "Riesgo de interpretación"],
    [
        ["USER", "Enterprise Structure aparecía después de Idea y Strategic Gate.", "Podía sugerir que el contexto empresarial se define después de iniciar la Idea."],
        ["Portfolio", "Intake, Evaluation, Analysis, Budget, Resources, Convergence y Gate Prep compartían un módulo amplio.", "No era visible la secuencia ni el paralelismo de escenarios."],
        ["Evidence", "FEL, BC, Funding, FIN, Risk y Permit eran módulos separados.", "La contribución conjunta a Gate07J quedaba implícita."],
        ["ADMIN", "Enterprise Strategy Manager era una lista plana y Enterprise Structure vivía aparte.", "No se distinguía Foundation de configuración transaccional."],
        ["FID", "Las capacidades de decisión y post-FID coexistían sin una fase explícita.", "Podía confundirse FID con activación automática del proyecto."],
    ],
    [1.05, 2.85, 2.6],
)
add_heading(doc, "Inventario funcional reutilizado", 2)
add_paragraph(
    doc,
    "Se conservaron Idea Lifecycle, Project Proposal, Strategic Gate, Portfolio Intake, Evaluation, Analysis, Investor Map, Scorecard, Budget, Resources, Convergence, FEL, Business Case, Funding, Economic & Financial Model, Risk, Permit, Gate Preparation, FID y Post-FID. También se reutilizaron Enterprise Explorer, Enterprise Structure y Workspace Context.",
)

# Architecture
add_heading(doc, "3. Arquitectura final Phase 0 / Phase 1+", 1)
add_table(
    doc,
    ["Fase", "Dominio", "Semántica y capacidades"],
    [
        ["0", "Enterprise Foundation", "Foundation / prerequisite. EWS, Explorer, Workspace Context, Strategic Objectives y estado de Fundación."],
        ["1", "Idea & Demand", "Inicio del flujo estratégico transaccional. Idea Lifecycle."],
        ["2", "Project Proposal", "Transformación de una Idea aceptada en propuesta gobernada."],
        ["3", "Strategic Gate", "Decisión previa a la entrada al Portafolio."],
        ["4", "Portfolio Planning", "Intake & Planning, Evaluation & Prioritization, Analysis, Investor Map y Scorecard."],
        ["5", "Portfolio Scenario Planning", "Budget y Resources en paralelo; convergencia y recomendación."],
        ["6", "Project Definition & Investment Case", "FEL, BC, Funding, FIN, Risk y Permit como evidencia concurrente."],
        ["7", "Investment Governance", "Evidence Policy (ADMIN), Gate Preparation, FID y Post-FID en orden de decisión."],
    ],
    [0.55, 2.05, 3.9],
)
add_heading(doc, "Reglas empresariales visibles", 2)
add_bullets(
    doc,
    [
        "EWS existe antes de la Idea y no se repite por cada Idea.",
        "La selección de un owning Workspace autorizado precede la creación de la Idea.",
        "Strategic Objectives se reutiliza como dependencia de preparación; no se duplica un catálogo.",
        "Business Case se marca como evidencia longitudinal, aunque tenga un acceso principal en Fase 6.",
        "Funding conserva capacidad y asignaciones; FIN conserva el modelamiento económico-financiero.",
        "FID no activa el proyecto automáticamente; Post-FID gobierna el handoff y la elegibilidad posterior.",
    ],
)

# User/Admin
add_heading(doc, "4. Jerarquías USER MODE y ADMIN MODE", 1)
add_heading(doc, "USER MODE", 2)
add_paragraph(
    doc,
    "La cinta lateral muestra Enterprise Strategy Manager como macroproceso y ocho fases desplegables. Cada fase conserva sus pantallas existentes y expone guía contextual en la landing principal.",
)
add_table(
    doc,
    ["Fase", "Accesos USER"],
    [
        ["0", "Foundation Status; Enterprise Explorer; Workspace Context"],
        ["1", "Idea Lifecycle"],
        ["2", "Project Proposals"],
        ["3", "Strategic Gate Decision"],
        ["4", "Portfolio Intake & Planning; Evaluation & Prioritization; Portfolio Analysis → Investor Map / Portfolio Scorecard"],
        ["5", "Budget Planning & Scenarios; Resource Planning & Scenarios; Scenario Convergence & Recommendation"],
        ["6", "Project Definition / FEL; Business Case Manager; Funds & Funding Readiness; Economic & Financial Model; Risk Readiness; Permit Readiness"],
        ["7", "Investment Gate Preparation; Investment Decision / FID; Post-FID Project Governance"],
    ],
    [0.7, 5.8],
)
add_heading(doc, "ADMIN MODE", 2)
add_paragraph(
    doc,
    "ADMIN replica las fases para evitar una taxonomía paralela. La diferencia es operacional: configura y gobierna fuentes ya existentes. Investment Decision Evidence Policy precede a Gate Preparation en Fase 7.",
)
add_table(
    doc,
    ["Zona", "Configuración ADMIN"],
    [
        ["Fase 0", "Foundation Configuration / Preview; Enterprise Structure & Workspace Manager; Workspace Types; Composition Rules; Strategic Objectives; Workspace Defaults & Inheritance; Module Catalog & Activation; Workspace Navigation Profiles."],
        ["Fases 1–6", "Configuraciones equivalentes de Idea, Proposal, Strategic Gate, Portfolio y aportantes de evidencia."],
        ["Fase 7", "Investment Decision Evidence Policy → Investment Gate Preparation → Investment Decision / FID → Post-FID Project Governance."],
        ["Plataforma", "Organization, Security y configuración compartida permanecen en Platform Administration."],
    ],
    [1.15, 5.35],
)

# Blueprint and routes
add_heading(doc, "5. Blueprint, rutas, permisos y contexto", 1)
add_heading(doc, "Fuente única de navegación", 2)
add_paragraph(
    doc,
    "strategicNavigationBlueprint.ts declara id, label, description, phase, order, route, mode, permission, workspaceTypes, governanceModels, children, isFoundation, predecessorIds, parallelGroup y convergesInto. applicationNavigation.ts deriva la cinta, mientras App.tsx se limita a resolver rutas y renderizar ramas recursivas.",
)
add_table(
    doc,
    ["Metadato", "Uso"],
    [
        ["phase / order", "Orden de negocio determinista en USER, ADMIN y Workspace Navigator."],
        ["predecessorIds", "Dependencias visibles sin introducir un motor de workflow."],
        ["parallelGroup", "Identifica Budget/Resources y los aportantes de evidencia concurrentes."],
        ["convergesInto", "Expresa convergencia a Recommendation o Gate07J."],
        ["route", "Enlaces profundos estables y navegación directa."],
        ["permission", "Conserva intención de permisos existentes; el backend sigue siendo autoridad."],
        ["workspaceTypes / governanceModels", "Evita forzar el camino CAPITAL_OWNER sobre rutas alternativas."],
    ],
    [1.8, 4.7],
)
add_heading(doc, "Rutas preservadas y ampliadas", 2)
add_table(
    doc,
    ["Ámbito", "Patrón", "Comportamiento"],
    [
        ["USER", "/app/strategy/<view>", "Resuelve cada view existente sin cambiar su clave interna."],
        ["ADMIN", "/app/admin/strategy/<view>", "Abre la configuración estratégica equivalente."],
        ["Base", "/app", "Permanece válido y conserva Dashboard."],
        ["Control", "/app/schedule-control", "Permanece válido."],
        ["Breadcrumb", "Blueprint → phase → capability → view", "Se genera desde la misma fuente de verdad."],
    ],
    [1.0, 2.15, 3.35],
)
add_heading(doc, "Fundación y autorización", 2)
add_paragraph(
    doc,
    "La landing consulta únicamente servicios existentes de Enterprise Structure e Idea options. Presenta cuatro señales: estructura publicada, owning Workspace elegible, Strategic Objectives e Idea configuration. Start Strategic Workflow aparece solo cuando las cuatro señales son devueltas. La UI no inventa elegibilidad, no crea Workspaces y no publica configuraciones.",
)
add_heading(doc, "Aplicabilidad por Governance Model", 2)
add_table(
    doc,
    ["Modelo", "Tratamiento"],
    [
        ["CAPITAL_OWNER", "Ruta estratégica completa, sujeta a workspaceTypes y permisos existentes."],
        ["CONTRACTOR_DELIVERY", "No se fuerza a fases owner-only; se conserva la creación alternativa de Project."],
        ["DIRECT_INTERNAL", "Mantiene su ruta aplicable y no hereda restricciones owner-only sin señal backend."],
        ["Legacy / unclassified", "Se conserva compatibilidad y orden estable; no se inventan clasificaciones."],
    ],
    [1.65, 4.85],
)

# Implementation
doc.add_page_break()
add_heading(doc, "6. Implementación y archivos", 1)
add_table(
    doc,
    ["Componente", "Responsabilidad"],
    [
        ["features/enterprise-strategy/navigation/strategicNavigationBlueprint.ts", "Blueprint USER/ADMIN, metadatos, filtros, breadcrumbs, rutas, Workspace Navigator y macroproceso derivado."],
        ["features/enterprise-strategy/EnterpriseStrategyLanding.tsx", "Landing de Fundación, secuencia, ramas y convergencia; reutiliza APIs existentes."],
        ["features/enterprise-strategy/enterpriseStrategy.css", "Diseño responsive desktop/tablet/mobile."],
        ["navigation/applicationNavigation.ts", "Integra el macroproceso derivado y preserva macroprocesos no estratégicos."],
        ["App.tsx", "Render recursivo, deep links, breadcrumbs y selección de vistas existentes."],
        ["WorkspaceOperationalPage.tsx", "Ordena entradas backend sin modificar estado, razón, permiso ni ruta."],
        ["PortfolioAnalysisWorkspace.tsx", "Permite abrir Investor Map y Scorecard como vistas hijas."],
        ["AdminEnterpriseStructurePage.tsx", "Permite entrada directa a la pestaña estratégica existente."],
        ["BimIfcModelViewer.tsx", "Corrección accesible mínima: anuncia el bloqueo por cache backend para IFC sobredimensionado."],
        ["scripts/capture-strategic-navigation.mjs", "Captura reproducible de evidencia USER/ADMIN."],
        ["docs/frontend-strategic-workflow-navigation-v2.md", "Documentación técnica actualizada."],
    ],
    [2.85, 3.65],
)
add_heading(doc, "Control de duplicidad", 2)
add_paragraph(
    doc,
    "No se copiaron pantallas ni servicios. Las capacidades nuevas son de composición: blueprint, landing, rutas, breadcrumb y agrupación visual. Las páginas de Idea, Proposal, Portfolio, FEL, BC, Funding, FIN, Risk, Permit, FID y Post-FID siguen siendo las implementaciones desarrolladas en entregas anteriores.",
)

# Validation
add_heading(doc, "7. Verificación y evidencias", 1)
add_table(
    doc,
    ["Verificación", "Resultado", "Evidencia"],
    [
        ["TypeScript", "PASS", "tsc --noEmit; 0 errores."],
        ["Lint", "PASS", "0 errores. 13 warnings históricos en BIM/FIN; 0 warnings nuevos de Enterprise Strategy."],
        ["Build", "PASS", "Vite 8.0.12; 2.439 módulos transformados."],
        ["Pruebas focalizadas", "PASS", "19/19: blueprint, landing y caracterización Enterprise Structure."],
        ["AppShell regression", "PASS", "Flujo completo de navegación principal y Project Workspace Governance."],
        ["BIM regression", "PASS", "Estado accesible de cache backend para IFC >100 MB."],
        ["Browser USER/ADMIN", "PASS", "2/2 Playwright Chromium contra Docker localhost:5173."],
        ["Responsive", "PASS", "Viewports 1440, 820 y 390 px conservan orden y contenido."],
        ["Console", "PASS", "Sin page errors en el recorrido E2E estratégico."],
        ["Infraestructura local", "PASS", "Durante E2E: frontend 5173, API 8000, PostgreSQL, Redis, worker y beat activos."],
    ],
    [1.65, 0.75, 4.1],
)
add_heading(doc, "Cobertura funcional comprobada", 2)
add_bullets(
    doc,
    [
        "Orden de las ocho fases en USER y ADMIN.",
        "Foundation / prerequisite y Start of strategic workflow visibles.",
        "Deep link a Portfolio Scorecard y breadcrumb completo.",
        "Budget/Resources como ramas paralelas y Scenario Convergence como destino.",
        "Seis aportantes de evidencia convergiendo en Investment Gate Preparation.",
        "Evidence Policy antes de Gate Preparation en ADMIN.",
        "ARIA en botones, grupos, regiones, expansión, selección y estado BIM.",
    ],
)

# Acceptance
add_heading(doc, "8. Matriz de aceptación", 1)
criteria = [
    ("Phase 0 Enterprise Foundation explicit", "Fase 0 visible en ambos modos."),
    ("Foundation marked as prerequisite", "Etiqueta Foundation / prerequisite."),
    ("Idea & Demand marked as workflow start", "Etiqueta Start of strategic workflow."),
    ("EWS before Idea conceptually clear", "Landing y orden lateral."),
    ("EWS not repeated per Idea", "Fundación separada del flujo transaccional."),
    ("Owning Workspace selection preserved", "Workspace Context reutilizado."),
    ("Strategic Objectives reused", "Catálogo y tab existentes."),
    ("USER sequence follows business flow", "Blueprint y E2E."),
    ("ADMIN sequence mirrors business flow", "Blueprint y E2E."),
    ("Idea before Proposal", "Fases 1 y 2."),
    ("Proposal before Strategic Gate", "Fases 2 y 3."),
    ("Strategic Gate before Portfolio Intake", "Fases 3 y 4."),
    ("Portfolio Intake before Evaluation", "Orden de hijos en Fase 4."),
    ("Evaluation before Analysis", "Orden de hijos en Fase 4."),
    ("Budget/Resources parallel", "parallelGroup scenario-planning."),
    ("Convergence after scenario branches", "convergesInto recommendation."),
    ("FEL / BC / Funding / FIN / Risk / Permit feed Gate07J", "Grupo de evidencia y landing."),
    ("Business Case longitudinal nature visible", "Metadato y descripción longitudinal."),
    ("Investment Gate before FID", "Orden de Fase 7."),
    ("FID before Post-FID", "Orden de Fase 7."),
    ("No automatic Project Activation implied", "Activación excluida del FID."),
    ("CAPITAL_OWNER route preserved", "Mapa de Governance Model."),
    ("CONTRACTOR_DELIVERY route preserved", "Project Workspace Governance alternativo."),
    ("DIRECT_INTERNAL route preserved", "Filtro de aplicabilidad no owner-only."),
    ("Single navigation blueprint", "Archivo estratégico único."),
    ("No App.tsx monolithic expansion", "Blueprint y landing encapsulados en feature."),
    ("Routes preserved", "Claves existentes y rutas base conservadas."),
    ("Permissions preserved", "Metadato existente; backend autoritativo."),
    ("No backend business change", "Diff limitado a frontend, docs, pruebas y reporte."),
    ("No schema migration", "Sin archivos de migración."),
    ("Browser USER", "Playwright Chromium."),
    ("Browser ADMIN", "Playwright Chromium."),
    ("Responsive", "820 px y 390 px."),
    ("Accessibility", "ARIA y estado accesible."),
    ("Regression", "Suite focalizada, AppShell, BIM y E2E."),
]
add_status_table(doc, [[name, "PASS", evidence] for name, evidence in criteria])

# Deliverables map
add_heading(doc, "Mapa de los 39 entregables", 1)
deliverables = [
    "Inventario de navegación actual", "Jerarquía USER actual", "Jerarquía ADMIN actual", "Inconsistencias de orden", "Arquitectura final Phase 0 / Phase 1+", "Jerarquía USER objetivo", "Jerarquía ADMIN objetivo", "Explicación del prerrequisito EWS", "Comportamiento de owning Workspace", "Dependencia de Strategic Objective", "Mapa de Governance Model", "Blueprint único", "Metadatos de workflow", "Metadatos de ramas paralelas", "Metadatos de convergencia de evidencia", "Especificación de breadcrumb", "Rediseño de landing", "Presentación de readiness de Fundación", "Descripciones empresariales", "Guía de siguiente paso", "Guía de dependencias ADMIN", "Matriz de rutas", "Matriz de permisos", "Navegación por contexto Workspace", "Refactor implementado", "Pruebas de caracterización", "Pruebas unitarias", "Pruebas de integración frontend", "Browser E2E USER", "Browser E2E ADMIN", "Prueba EWS prerequisite", "Prueba Budget/Resource", "Prueba de convergencia de evidencia", "Regresión Governance Model", "Responsive/accesibilidad", "Informe console/network", "Capturas antes/después", "Documentación de navegación", "Informe técnico final"
]
deliverable_rows = []
for index, name in enumerate(deliverables, start=1):
    evidence = "Código y documentación"
    if index in (26, 27, 28):
        evidence = "Vitest"
    elif index in (29, 30, 31, 32, 33, 35, 36):
        evidence = "Playwright Chromium"
    elif index == 37:
        evidence = "docs/validation/strategic-navigation-v2"
    elif index == 38:
        evidence = "docs/frontend-strategic-workflow-navigation-v2.md"
    elif index == 39:
        evidence = OUTPUT.name
    deliverable_rows.append([str(index), name, "ENTREGADO", evidence])
add_table(doc, ["#", "Entregable", "Estado", "Evidencia"], deliverable_rows, [0.35, 3.25, 1.0, 1.9])

# Screenshots
for title_text, file_name, caption in [
    ("9.1 USER MODE — antes", "before/user-mode-desktop.png", "Línea base: módulos estratégicos agrupados por evolución técnica y no por fases."),
    ("9.2 USER MODE — después", "user-mode-desktop.png", "Resultado: Fundación separada, fases 1–7, ramas paralelas y convergencia de evidencia."),
    ("9.3 ADMIN MODE — antes", "before/admin-mode-desktop.png", "Línea base: lista plana de configuraciones estratégicas."),
    ("9.4 ADMIN MODE — después", "admin-mode-desktop.png", "Resultado: configuración espejo por fases y Evidence Policy antes de Gate Preparation."),
]:
    doc.add_page_break()
    add_heading(doc, title_text, 1)
    image_path = EVIDENCE / file_name
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if image_path.exists():
        add_picture_fit(paragraph.add_run(), image_path)
    add_caption(doc, caption)

# Closeout
doc.add_page_break()
add_heading(doc, "10. Conclusiones y señal de cierre", 1)
add_paragraph(
    doc,
    "La solución expresa el flujo estratégico desde la Fundación Empresarial hasta el gobierno post-FID con una jerarquía única, tipada y reutilizable. El usuario puede entender qué es prerrequisito, qué inicia una transacción, qué actividades son paralelas y dónde convergen las decisiones. ADMIN configura la misma arquitectura sin introducir otra taxonomía.",
)
add_paragraph(
    doc,
    "El alcance se detiene en la navegación estratégica solicitada. No se inició la implementación de macroprocesos operacionales. Durante la validación de navegador, frontend y API respondieron en localhost:5173 y localhost:8000; su disponibilidad posterior depende de Docker Desktop.",
)
add_heading(doc, "Señales de validación", 2)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("FRONTEND_STRATEGIC_WORKFLOW_NAVIGATION_V2_VALIDATED")
set_run_font(run, 12, True, TEAL, "Consolas")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("STRATEGIC_FRONTEND_ALIGNED_TO_FOUNDATION_AND_BUSINESS_WORKFLOW")
set_run_font(run, 11, True, NAVY, "Consolas")
add_heading(doc, "Recomendación de continuidad", 2)
add_paragraph(
    doc,
    "El siguiente incremento debe partir de esta jerarquía y del blueprint; no debe duplicar rutas ni reconstruir permisos. Antes de incorporar macroprocesos operacionales se recomienda estabilizar los trece warnings históricos de React Hooks identificados por lint, sin alterar este cierre funcional.",
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.core_properties.title = "Informe técnico Frontend Strategic Workflow Navigation V2"
doc.core_properties.subject = "P&Pmis Construction AI"
doc.core_properties.author = "Codex"
doc.core_properties.comments = "Implementación y validación técnica del PROMPT V2."
doc.save(OUTPUT)
stale_output = ROOT / "outputs" / "Informe_Tecnico_Navegacion_Estrategica_Frontend_V2_2026-09-09.docx"
if stale_output.exists():
    stale_output.unlink()
print(OUTPUT)
