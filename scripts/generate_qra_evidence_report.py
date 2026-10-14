"""Build the contextual QRA evidence report from reviewed release observations."""

from __future__ import annotations

import re
from pathlib import Path

import generate_cpig_gap_review_report as base
import generate_qra_foundation_report as foundation
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "85-qra-evidence-integration-release.md"
OUTPUT = (
    ROOT
    / "outputs"
    / "Informe_Tecnico_QRA_Evidence_Policy_Gate07J_BC_FID_2026-10-14.docx"
)
LOGO = ROOT / "frontend" / "public" / "pypmis-construction-ai-logo.png"


def main():
    content = SOURCE.read_text(encoding="utf-8")
    sections = [
        int(value) for value in re.findall(r"^## (\d+)\. ", content, re.MULTILINE)
    ]
    if sections != list(range(1, 90)):
        raise ValueError("Expected all 89 requested technical report sections")
    status = re.search(r"^\*\*Estado actual:\*\*\s+(.+)$", content, re.MULTILINE).group(
        1
    )
    document = Document()
    base.configure_document(document)
    base.set_language(document)
    document.styles["Normal"].font.size = Pt(11)
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        document.styles[name].font.color.rgb = RGBColor.from_string(base.BLACK)
        props = document.styles[name]._element.find(qn("w:pPr"))
        if props is not None:
            for border in props.findall(qn("w:pBdr")):
                props.remove(border)
    document.core_properties.title = "Informe técnico Integración de evidencia QRA"
    document.core_properties.subject = "Política contextual Gate07J Business Case y FID"
    document.core_properties.author = "P&Pmis Construction AI"
    document.core_properties.comments = (
        "Implementación y validación local de evidencia QRA contextual; "
        "sin publicación operacional automática ni decisión FID automática."
    )
    document.core_properties.keywords = (
        "QRA, Evidence Policy, Gate07J, Business Case, FID"
    )
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(22)
    if LOGO.exists():
        shape = paragraph.add_run().add_picture(str(LOGO), width=Inches(1.15))
        base.set_alt_text(
            shape, "P&Pmis Construction AI", "Logotipo azul de P&Pmis Construction AI"
        )
    paragraph = document.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(18)
    paragraph.add_run("Informe técnico\nIntegración de evidencia QRA")
    paragraph = document.add_paragraph(style="Subtitle")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("Gate07J Business Case y FID")
    for text in (
        "P&Pmis Construction AI",
        "Capital Portfolio and Investment Governance",
        "Corte técnico 14 de octubre de 2026",
        status.replace("_", " "),
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        base.add_inline(paragraph, text, size=10)
    document.add_paragraph(
        "El informe describe la integración del resultado QRA aceptado en las exigencias de evidencia de inversión. "
        "Explica su aplicabilidad, los límites del motor, la conservación del historial y las comprobaciones realmente ejecutadas. "
        "QRA aporta evidencia analítica; Gate07K conserva la decisión formal y los dominios de origen conservan su autoridad."
    )
    document.add_page_break()
    document.add_heading("Control del documento", level=1)
    base.add_table(
        document,
        [
            ["Campo", "Detalle"],
            ["Baseline inicial", "main / 3cadbe0 / working tree limpio al inicio"],
            ["Persistencia", "Sin migración nueva / Alembic 20261011_0068"],
            [
                "Contenido",
                "89 secciones de arquitectura implementación validación y límites",
            ],
            ["Política operacional", "No se publica ni modifica automáticamente"],
            ["Control Git", "Sin commit ni push del agente"],
        ],
    )
    document.add_heading("Guía de lectura", level=1)
    for text in (
        "Secciones 1 a 19 Baseline autoridad aplicabilidad y ciclo de la política",
        "Secciones 20 a 45 Resultados criterios Business Case y decisiones históricas",
        "Secciones 46 a 57 Interfaz seguridad persistencia API y navegación",
        "Secciones 58 a 89 Pruebas rendimiento deuda y decisión de cierre",
    ):
        base.add_inline(document.add_paragraph(style="List Bullet"), text)
    document.add_page_break()
    base.SOURCE, base.DIAGRAMS = SOURCE, []
    source_start = len(document.paragraphs)
    base.add_source_content(
        document, [], page_break_heading_prefixes=(), skip_preface=True
    )
    for index, paragraph in enumerate(document.paragraphs):
        if paragraph.style.name.startswith("Heading"):
            paragraph.text = re.sub(
                r"\s+", " ", re.sub(r"[^\w\s]", " ", paragraph.text).replace("_", " ")
            ).strip()
        elif index >= source_start and paragraph.style.name in {
            "Normal",
            "List Bullet",
            "List Number",
        }:
            for run in paragraph.runs:
                run.font.size = Pt(11)
    foundation.format_tables(document)
    for table in document.tables:
        if len(table.columns) == 2:
            widths = (1.35, 5.15)
        elif table.cell(0, 0).text == "Ubicación principal":
            widths = (2.45, 2.2, 1.85)
            for row in table.rows[1:]:
                cell = row.cells[0]
                if "/" in cell.text:
                    cell.text = cell.text.replace("/", "/\n", 1)
        else:
            widths = (1.5, 2.25, 2.75)
        for column, width in zip(table.columns, widths, strict=True):
            column.width = Inches(width)
        for row in table.rows:
            for cell, width in zip(row.cells, widths, strict=True):
                cell.width = Inches(width)
    for section in document.sections:
        section.different_first_page_header_footer = True
        paragraph = section.header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        base.set_run_font(
            paragraph.add_run("P&Pmis Construction AI | Evidencia QRA"),
            size=7.5,
            color=base.BLACK,
        )
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        base.set_run_font(
            paragraph.add_run("Informe técnico | 2026-10-14 | Página "),
            size=7.5,
            color=base.MUTED,
        )
        base.add_field(paragraph, "PAGE")
        base.set_run_font(paragraph.add_run(" de "), size=7.5, color=base.MUTED)
        base.add_field(paragraph, "NUMPAGES")
    base.update_field_settings(document)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(f"created={OUTPUT}")
    print(f"sections={len(sections)}")
    print(f"status={status}")


if __name__ == "__main__":
    main()
