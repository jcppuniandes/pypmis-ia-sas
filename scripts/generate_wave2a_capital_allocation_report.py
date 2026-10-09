from __future__ import annotations

import re
from pathlib import Path

import generate_cpig_gap_review_report as base
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "73-wave2a-strategic-capital-buckets-allocation-release.md"
QA = ROOT / "outputs" / ".qa" / "wave2a-capital-allocation"
DIAGRAM_DIR = QA / "diagrams"
REPORT_DATE = "2026-10-05"
REPORT_DATE_DISPLAY = "5 de octubre de 2026"
OUTPUT = (
    ROOT
    / "outputs"
    / f"Informe_Tecnico_Wave2A_Strategic_Capital_Buckets_Allocation_Framework_{REPORT_DATE}.docx"
)
LOGO = ROOT / "frontend" / "public" / "pypmis-construction-ai-logo.png"


DIAGRAMS = [
    {
        "name": "01-wave2a-process-chain.png",
        "title": "Cadena conceptual de Wave 2A",
        "rows": [
            ["Strategic Objectives", "Investment Themes", "Planning Cycle"],
            ["Capital Demands", "Allocation Plan", "Capital Buckets"],
            ["Target allocation", "Position comparison", "Readiness"],
        ],
        "edges": [
            (0, 0, 0, 1),
            (0, 1, 0, 2),
            (0, 2, 1, 0),
            (1, 0, 1, 1),
            (1, 1, 1, 2),
            (1, 2, 2, 0),
            (2, 0, 2, 1),
            (2, 1, 2, 2),
        ],
        "alt": (
            "Cadena desde objetivos y temas estratégicos hasta demanda, plan de asignación, "
            "capital buckets, comparación de posiciones y readiness."
        ),
    },
    {
        "name": "02-wave2a-comparison.png",
        "title": "Modelo de comparación de posiciones",
        "rows": [
            ["Target", "Current Demand", "Current Portfolio"],
            ["Recommended", "Approved", "Variance and Readiness"],
        ],
        "edges": [
            (0, 0, 1, 2),
            (0, 1, 1, 2),
            (0, 2, 1, 2),
            (1, 0, 1, 2),
            (1, 1, 1, 2),
        ],
        "alt": (
            "Target, demanda actual, portafolio actual, recomendado y aprobado convergen en "
            "un read model backend de variación y readiness."
        ),
    },
    {
        "name": "03-wave2a-admin-user.png",
        "title": "Separación ADMIN y USER",
        "rows": [
            ["ADMIN configuration", "Published Bucket policy"],
            ["USER Allocation Plan", "Submit", "Independent Review", "Endorse and Freeze"],
        ],
        "edges": [
            (0, 0, 0, 1),
            (0, 1, 1, 0),
            (1, 0, 1, 1),
            (1, 1, 1, 2),
            (1, 2, 1, 3),
        ],
        "alt": (
            "ADMIN publica la política de buckets y USER opera el plan live mediante submit, "
            "revisión independiente y endorsement inmutable."
        ),
    },
]


def read_release_status() -> str:
    source_text = SOURCE.read_text(encoding="utf-8")
    match = re.search(
        r"^\*\*Estado (?:actual|final(?: de release)?):\*\*\s+(.+?)\.?$",
        source_text,
        re.MULTILINE,
    )
    if match is not None:
        return match.group(1).strip().strip("`")
    if "STRATEGIC_CAPITAL_BUCKETS_AND_ALLOCATION_FRAMEWORK_RELEASE_VALIDATED" in source_text:
        return "VALIDACIÓN CONDICIONADA A LA EVIDENCIA REGISTRADA"
    return "EN VALIDACIÓN"


def configure_metadata(document: Document) -> None:
    core = document.core_properties
    core.title = "Informe técnico Wave 2A Strategic Capital Buckets and Allocation Framework"
    core.subject = "Implementación, controles y validación del marco estratégico de asignación de capital"
    core.author = "P&Pmis Construction AI"
    core.keywords = "CPIG, Capital Buckets, Allocation Plan, Portfolio, governance"
    core.comments = "Informe técnico de implementación y evidencia reproducible de Wave 2A."


def add_cover(document: Document, release_status: str) -> None:
    logo = document.add_paragraph()
    logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo.paragraph_format.space_before = Pt(28)
    if LOGO.exists():
        shape = logo.add_run().add_picture(str(LOGO), width=Inches(1.3))
        base.set_alt_text(
            shape,
            "P&Pmis Construction AI",
            "Logotipo azul de P&Pmis Construction AI.",
        )

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(22)
    title.paragraph_format.space_after = Pt(12)
    title.add_run("Informe técnico Wave 2A\nStrategic Capital Buckets and Allocation Framework")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("P&Pmis Construction AI")
    base.set_run_font(run, size=13, color=base.BLACK)

    metadata = document.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_before = Pt(22)
    base.add_inline(
        metadata,
        f"Capital Portfolio and Investment Governance | {REPORT_DATE_DISPLAY}",
        size=10,
    )

    status_label = document.add_paragraph()
    status_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status_label.paragraph_format.space_before = Pt(38)
    run = status_label.add_run("ESTADO DE VALIDACIÓN")
    base.set_run_font(run, size=8, bold=True, color=base.MUTED)

    status = document.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = status.add_run(release_status)
    base.set_run_font(run, size=11, bold=True, color=base.TEAL)
    document.add_page_break()


def add_document_control(document: Document, release_status: str) -> None:
    document.add_heading("Control del documento", level=1)
    base.add_table(
        document,
        [
            ["Campo", "Valor"],
            ["Producto", "P&Pmis Construction AI"],
            ["Macroproceso", "Capital Portfolio and Investment Governance"],
            ["Incremento", "Wave 2A Strategic Capital Buckets and Allocation Framework"],
            ["Rama y commit de baseline", "main / 9cb0969"],
            ["Migración objetivo", "20261001_0063 sobre 20260929_0062"],
            ["Fecha de corte", REPORT_DATE],
            ["Estado al corte", release_status],
            ["Fuente técnica", "Código, migración, pruebas, navegador y guardrails del repositorio"],
        ],
    )

    document.add_heading("Lectura ejecutiva", level=1)
    paragraph = document.add_paragraph()
    base.add_inline(
        paragraph,
        (
            "Wave 2A incorpora Strategic Capital Buckets versionados y Allocation Plans live ligados a "
            "Planning Cycles. La solución compara target, demanda actual, portafolio actual, recomendado y "
            "aprobado sin asumir autoridad de ranking, presupuesto, recursos, funding, FIN o FID."
        ),
    )

    document.add_heading("Contenido", level=1)
    for item in (
        "Baseline, ADR, autoridad y reutilización",
        "Bucket configuration y Allocation Plan gobernado",
        "Clasificación, precedencia, multi-moneda y no doble conteo",
        "Comparación Target, Current, Recommended y Approved",
        "ADMIN, USER, seguridad, persistencia, API y navegación",
        "Pruebas, regresiones, brechas y cierre de release",
    ):
        paragraph = document.add_paragraph(style="List Bullet")
        base.add_inline(paragraph, item)
    document.add_page_break()


def add_headers_and_footers(document: Document) -> None:
    for section in document.sections:
        section.different_first_page_header_footer = True
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = header.add_run("P&Pmis Construction AI  |  Wave 2A Capital Allocation")
        base.set_run_font(run, size=7.5, bold=True, color=base.BLACK)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run(f"Confidencial | {REPORT_DATE} | Página ")
        base.set_run_font(run, size=7.5, color=base.MUTED)
        base.add_field(footer, "PAGE")
        run = footer.add_run(" de ")
        base.set_run_font(run, size=7.5, color=base.MUTED)
        base.add_field(footer, "NUMPAGES")


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    QA.mkdir(parents=True, exist_ok=True)
    DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)

    base.SOURCE = SOURCE
    base.QA = QA
    base.DIAGRAM_DIR = DIAGRAM_DIR
    base.DIAGRAMS = DIAGRAMS
    diagram_paths = [base.render_diagram(spec) for spec in DIAGRAMS]
    release_status = read_release_status()

    document = Document()
    base.configure_document(document)
    configure_metadata(document)
    base.set_language(document)
    add_cover(document, release_status)
    add_document_control(document, release_status)
    base.add_source_content(
        document,
        diagram_paths,
        page_break_heading_prefixes=("4. ", "25. ", "36. ", "47. ", "64. ", "67. "),
    )
    add_headers_and_footers(document)
    base.update_field_settings(document)
    document.save(OUTPUT)
    print(f"created={OUTPUT}")
    print(f"diagrams={len(diagram_paths)}")
    print(f"source={SOURCE}")


if __name__ == "__main__":
    main()
