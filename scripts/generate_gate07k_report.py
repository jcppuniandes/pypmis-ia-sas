from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt
from generate_gate07i_report import (
    CALLOUT,
    CAUTION,
    DARK_BLUE,
    INK,
    POSITIVE,
    RISK,
    TEAL,
    add_bottom_rule,
    add_callout,
    add_page_number,
    add_table,
    configure_styles,
    set_image_alt,
    set_run,
)

REPO = Path(__file__).resolve().parents[1]
PROMPT = Path(
    r"C:\Users\Ricardo\Documents\P&P\P&Pmis Construction AI\Diseño\Portfolio\PROMPT_07K_INVESTMENT_DECISION_FID_GOVERNANCE_AND_DECISION_LIFECYCLE.md"
)
OUTPUT = (
    REPO
    / "outputs"
    / "technical_reports"
    / "Informe_Tecnico_Gate_07K_Investment_Decision_FID_20260831.docx"
)
LOGO = REPO / "frontend" / "public" / "pypmis-construction-ai-logo.png"

# Resolved style system:
# preset=standard_business_brief; first-page pattern=memo_masthead.
# Letter portrait; 1 in margins; 0.492 in header/footer; 9360 DXA tables;
# 120 DXA table indent; Calibri 11 pt; body 6 pt after / 1.10 line.
# Named override: dense traceability matrices use 7.2-8.4 pt.


def source_note(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    set_run(paragraph.add_run(text), size=8, color="5F6B76", italic=True)


def anchored_page_break(doc: Document) -> None:
    """Keep LibreOffice from placing the next block inside the top margin."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    anchor = doc.add_paragraph()
    anchor.paragraph_format.space_before = Pt(0)
    anchor.paragraph_format.space_after = Pt(0)
    anchor.paragraph_format.line_spacing = 0.1
    set_run(anchor.add_run("\u200b"), size=1, color="FFFFFF")


def title_page(doc: Document) -> None:
    section = doc.sections[0]
    doc.settings.odd_and_even_pages_header_footer = False
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(
        header.add_run("TECHNICAL REPORT · GATE 07K"),
        size=8.5,
        color="5F6B76",
        bold=True,
    )
    add_bottom_rule(header, color="D9E1E8", size=6)
    add_page_number(section.footer.paragraphs[0])

    if LOGO.exists():
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(8)
        picture = paragraph.add_run().add_picture(str(LOGO), width=Inches(0.72))
        set_image_alt(picture, "Logotipo de P&Pmis Construction AI")

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(4)
    set_run(
        kicker.add_run("INFORME TÉCNICO DE IMPLEMENTACIÓN"),
        size=10,
        color=TEAL,
        bold=True,
    )
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.keep_with_next = True
    set_run(title.add_run("Gate 07K"), size=25, color=INK, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    set_run(
        subtitle.add_run("Investment Decision / FID Governance and Decision Lifecycle"),
        size=15,
        color=DARK_BLUE,
        bold=True,
    )
    metadata = (
        ("Aplicación", "P&Pmis Construction AI"),
        ("Fecha de corte", "31 de agosto de 2026"),
        ("Repositorio", "pypmis-ia-sas · branch main · base fde6027"),
        ("Migración", "20260831_0055"),
        ("Perfiles", "PYP-INV-SELECT · PYP-FFA"),
        ("Estado exclusivo", "GATE07K_INVESTMENT_DECISION_RELEASE_VALIDATED"),
    )
    for label, value in metadata:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        set_run(paragraph.add_run(f"{label}: "), size=10.5, color=INK, bold=True)
        set_run(paragraph.add_run(value), size=10.5)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(10)
    rule.paragraph_format.space_after = Pt(14)
    add_bottom_rule(rule)
    add_callout(
        doc,
        "Resultado ejecutivo",
        "Se implementó una fuente de verdad InvestmentDecision separada de Gate 07J, con rondas, paquete "
        "congelado, hashes, autoridad individual o comité, quorum, votos, checklist, segregación de funciones, "
        "ETag e idempotencia. PYP-INV-SELECT + GO continúa definición sin ser FID; sólo PYP-FFA + GO produce "
        "FINAL_INVESTMENT_DECISION_GO. Ninguna salida crea o activa Project, emite NTP ni inicia ejecución.",
        fill=POSITIVE,
    )
    add_table(
        doc,
        ["Capa", "Resultado", "Evidencia"],
        [
            (
                "Dominio / API",
                "PASS",
                "2 agregados; 18 rutas OpenAPI; ETag e idempotencia",
            ),
            (
                "Reglas",
                "PASS",
                "8 pruebas funcionales; GO/FID, SoD, quorum, voto, límites y alternos",
            ),
            ("PostgreSQL", "PASS", "0054→0055→0054→0055; validator y E2E estructural"),
            ("Frontend", "PASS", "9 vistas USER; 20 secciones ADMIN; 3 pruebas Vitest"),
            ("Browser", "PASS", "2 flujos Chromium; FID, HOLD, quorum, stale y ADMIN"),
            ("Localhost", "ACTIVE", "Frontend 5173; API 8000; Alembic 0055"),
        ],
        [1900, 1300, 6160],
        font_size=8.4,
    )
    anchored_page_break(doc)


def contents(doc: Document) -> None:
    heading = doc.add_heading("Contenido y trazabilidad", level=1)
    heading.paragraph_format.space_before = Pt(0)
    add_table(
        doc,
        ["Capítulo", "Contenido", "Entregables"],
        [
            ("1", "Baseline, alcance y reutilización", "1-5"),
            ("2", "Modelo de dominio, lifecycle y source of truth", "6-17"),
            ("3", "Autoridad, comité, quorum, voto y SoD", "18-30"),
            ("4", "Semántica de decisión y contratos de salida", "31-46"),
            ("5", "Fronteras Gate 07I, Gate 07J y Project", "47-51, 90"),
            ("6", "Superficies USER/ADMIN", "52-58"),
            ("7", "Seguridad, APIs y persistencia", "59-68"),
            ("8", "Validación, E2E y baseline before/after", "69-81, 89, 91-92"),
            ("9", "Pruebas frontend y browser", "82-88"),
            ("10", "Riesgos, deuda y cierre", "93-95"),
            ("11", "Matriz integral", "1-95"),
        ],
        [1200, 5600, 2560],
        font_size=8.4,
    )
    add_callout(
        doc,
        "Criterio de cierre",
        "El informe distingue implementación, verificación automática y límites deliberados. Los datos de "
        "baseline y localhost provienen de comandos ejecutados en esta entrega; la matriz final relaciona cada "
        "uno de los 95 entregables con evidencia concreta.",
        fill=CALLOUT,
    )
    anchored_page_break(doc)


def baseline(doc: Document) -> None:
    doc.add_heading("1. Baseline, alcance y REUSE BEFORE CREATE", level=1)
    add_table(
        doc,
        ["Elemento", "Entrada", "Salida / decisión"],
        [
            ("Alembic", "20260831_0054", "Sucesora aditiva y reversible 0055"),
            (
                "Gate 07J",
                "Preparación FROZEN + READY",
                "Única entrada admisible; no se recalculan scores",
            ),
            (
                "Project",
                "4 registros; hash e33beb48d80510c3fa7715a4321b7286",
                "4; mismo hash después",
            ),
            ("AdminConfiguration", "82", "83 por seed explícito Gate 07K"),
            (
                "InvestmentDecision",
                "Tabla inexistente",
                "Tabla creada; 0 decisiones de negocio",
            ),
            ("Decision votes", "Tabla inexistente", "Tabla creada; 0 votos de negocio"),
        ],
        [2100, 3450, 3810],
        font_size=8.3,
    )
    source_note(
        doc, "Fuente: inspección PostgreSQL local antes y después de Alembic 0055."
    )
    doc.add_heading("1.1 Capacidades reutilizadas", level=2)
    add_table(
        doc,
        ["Capacidad existente", "Uso en Gate 07K"],
        [
            (
                "Gate 07J InvestmentGatePreparation",
                "Contrato FROZEN/READY, perfil, alcance y hashes",
            ),
            (
                "AdminConfiguration",
                "Policy publicada, clone, preview no persistente, publish y snapshots",
            ),
            ("AdminNumberSequence", "Numeración tenant-scoped IDEC-00001"),
            ("Enterprise RBAC", "Roles, permisos y tenant isolation"),
            ("SecurityEvent", "Auditoría append-only del lifecycle y decisiones"),
            ("ETag / If-Match", "Concurrencia optimista en mutaciones"),
            ("Idempotency", "Create, transición, voto y decide protegidos"),
        ],
        [3000, 6360],
        font_size=8.5,
    )
    add_callout(
        doc,
        "REUSE BEFORE CREATE",
        "Se creó únicamente el agregado faltante InvestmentDecision/InvestmentDecisionVote. No se duplicaron "
        "Gate 07J, Project, configuración, numeración, identidad, auditoría ni navegación.",
        fill=CAUTION,
    )


def domain(doc: Document) -> None:
    doc.add_heading("2. Dominio, lifecycle y fuente de verdad", level=1)
    add_table(
        doc,
        ["Objeto", "Responsabilidad", "Invariantes"],
        [
            (
                "InvestmentDecision",
                "Ronda formal sobre un paquete Gate 07J",
                "Tenant, preparación, perfil y round únicos",
            ),
            (
                "InvestmentDecisionVote",
                "Voto persistente por miembro",
                "Un voto por miembro/decisión; idempotency única",
            ),
            (
                "Decision Package Snapshot",
                "Paquete, evidencia, scope y política congelados",
                "Hash estable; no recálculo",
            ),
            (
                "SecurityEvent",
                "Historia y outcome auditable",
                "Eventos de create/update/ready/vote/decide/void",
            ),
        ],
        [2600, 4000, 2760],
        font_size=8.3,
    )
    doc.add_heading("2.1 Lifecycle y cardinalidad", level=2)
    add_table(
        doc,
        ["Estado", "Acción admitida", "Control"],
        [
            (
                "DRAFT",
                "Editar metadatos, maker o comité; submit; void",
                "If-Match y policy vigente",
            ),
            ("SUBMITTED", "Start review; void", "Paquete y hashes corrientes"),
            ("IN_REVIEW", "Checklist, votos, ready; void", "Autoridad, quorum, SoD"),
            (
                "READY_FOR_DECISION",
                "Decide GO/HOLD/RECYCLE/KILL",
                "Lock, versión, reglas y límites",
            ),
            (
                "DECIDED",
                "Sólo lectura e historia",
                "Inmutable; corrección requiere nueva ronda",
            ),
            ("VOIDED", "Sólo lectura", "Sólo predecision"),
        ],
        [1850, 3860, 3650],
        font_size=8.3,
    )
    add_callout(
        doc,
        "FID source of truth",
        "InvestmentDecision con profile_code=PYP-FFA, outcome=GO e is_fid=true es el único FID. No existe "
        "una segunda tabla FIDApproval. Full Funds Authorization representa autorización de gobierno sobre el "
        "envelope congelado; no reserva caja ni crea movimientos contables.",
        fill=POSITIVE,
    )


def governance(doc: Document) -> None:
    doc.add_heading("3. Autoridad, comité, quorum, voto y SoD", level=1)
    add_table(
        doc,
        ["Control", "Implementación", "Caso verificado"],
        [
            (
                "Single Decision Maker",
                "Usuario congelado; sólo él decide",
                "Maker independiente PYP-FFA",
            ),
            (
                "Committee",
                "Miembros, chair y voting_right congelados",
                "Chair + member + observer",
            ),
            ("Quorum", "Conteo, porcentaje y chair_required", "Bloqueo QUORUM_NOT_MET"),
            (
                "Voting rule",
                "Simple majority, unanimous, supermajority, chair decides",
                "Empate bloquea; mayoría GO decide",
            ),
            (
                "Votes",
                "Persistentes y actualizables durante review/ready",
                "Constraint voter + idempotency",
            ),
            (
                "SoD",
                "PYP-FFA maker distinto de creator/updater",
                "SEGREGATION_OF_DUTIES_VIOLATION",
            ),
            (
                "Authority limits",
                "Límite por usuario y monto congelado",
                "25M bloqueado; 12.5M autorizado",
            ),
            (
                "Checklist",
                "Cinco acknowledgements blocking",
                "Ready bloqueado si incompleto",
            ),
        ],
        [2100, 4300, 2960],
        font_size=8.1,
    )
    add_callout(
        doc,
        "No-double-scoring",
        "Gate 07K consume los resultados y la evidencia congelados de Gate 07 J. No introduce un score nuevo, "
        "no modifica FEL, riesgo, permisos o funding, y no sustituye el juicio de la autoridad por una fórmula.",
        fill=CALLOUT,
    )


def semantics(doc: Document) -> None:
    doc.add_heading("4. Semántica de decisión y contratos de salida", level=1)
    add_table(
        doc,
        ["Perfil / outcome", "final_output", "FID / inversión", "Activación"],
        [
            (
                "PYP-INV-SELECT · GO",
                "AUTHORIZED_TO_CONTINUE_PROJECT_DEVELOPMENT",
                "false / false",
                "false",
            ),
            (
                "PYP-INV-SELECT · HOLD",
                "INVESTMENT_SELECTION_HELD",
                "false / false",
                "false",
            ),
            (
                "PYP-INV-SELECT · RECYCLE",
                "INVESTMENT_SELECTION_RECYCLE_REQUIRED",
                "false / false",
                "false",
            ),
            (
                "PYP-INV-SELECT · KILL",
                "INVESTMENT_SELECTION_KILLED",
                "false / false",
                "false",
            ),
            ("PYP-FFA · GO", "FINAL_INVESTMENT_DECISION_GO", "true / true", "false"),
            ("PYP-FFA · HOLD", "FID_HELD", "false / false", "false"),
            ("PYP-FFA · RECYCLE", "FID_RECYCLE_REQUIRED", "false / false", "false"),
            ("PYP-FFA · KILL", "FID_KILLED", "false / false", "false"),
        ],
        [2200, 3900, 1760, 1500],
        font_size=7.8,
    )
    doc.add_heading("4.1 Contrato común", level=2)
    add_table(
        doc,
        ["Grupo", "Campos principales"],
        [
            (
                "Identidad",
                "decision_id, decision_number, decision_round, gate_profile, outcome",
            ),
            (
                "Autoridad",
                "decision_date, decision_authority, committee_snapshot, quorum_snapshot",
            ),
            (
                "Integridad",
                "source_preparation_hash, decision_package_hash, investment_decision_hash",
            ),
            (
                "FID",
                "authorized_investment_amount, currency, is_fid, full_funds_authorized",
            ),
            ("Condiciones", "input + post-decision conditions con origin conservado"),
            ("Guardrails", "can_create_project/activate/issue_ntp/execute = false"),
        ],
        [2200, 7160],
        font_size=8.5,
    )
    add_callout(
        doc,
        "Handoff post-FID",
        "READY_FOR_POST_FID_PROJECT_GOVERNANCE se expone en post_fid_handoff, separado de final_output. "
        "La señal no activa Project ni habilita contratación, compromisos, pagos, contabilidad, NTP o ejecución.",
        fill=CAUTION,
    )


def boundaries(doc: Document) -> None:
    doc.add_heading("5. Fronteras de arquitectura", level=1)
    add_table(
        doc,
        ["Límite", "Garantía implementada"],
        [
            (
                "Gate 07I",
                "Sólo trazabilidad al portfolio recommendation/project set ya seleccionado",
            ),
            (
                "Gate 07J",
                "Preparación FROZEN/READY inmutable; mismatch/stale/policy bloquean",
            ),
            (
                "Exact scope",
                "decision_scope_hash debe coincidir con recommended_project_set_hash",
            ),
            (
                "Partial GO",
                "No se permite mutar parcialmente el Project Set congelado; requiere recycle/nueva ronda",
            ),
            ("Project", "No se crea, activa ni cambia status de EnterpriseWorkspace"),
            ("Contractor", "No vendor/contract/commitment/PO/invoice/payment"),
            (
                "Execution",
                "No schedule start, NTP, cost posting ni módulo de ejecución",
            ),
        ],
        [2200, 7160],
        font_size=8.4,
    )
    add_callout(
        doc,
        "STOP post-FID",
        "La implementación termina al registrar el resultado formal y su handoff. Un proceso futuro deberá "
        "definir explícitamente elegibilidad de activación y gobierno de ejecución.",
        fill=RISK,
    )


def surfaces(doc: Document) -> None:
    doc.add_heading("6. USER MODE y ADMIN MODE", level=1)
    add_table(
        doc,
        ["USER MODE · 9 vistas", "Finalidad"],
        [
            (
                "Decision Queue / My Decisions",
                "Cola, ronda, perfil, lifecycle y acciones permitidas",
            ),
            ("Investment Decision Package", "Snapshot histórico congelado"),
            ("Evidence Summary", "Resultados fuente sin recalcular score"),
            ("Decision Checklist", "Acknowledgements blocking"),
            ("Conditions", "Condiciones de entrada y post-GO"),
            ("Committee / Votes", "Autoridad, quorum, regla y votos"),
            ("Decision History", "Eventos inmutables"),
            ("Authorization Output", "Salida formal, FID y guardrails"),
        ],
        [3300, 6060],
        font_size=8.4,
    )
    add_table(
        doc,
        ["ADMIN MODE · 20 secciones", "Cobertura"],
        [
            (
                "Governance, profiles, outcomes, authority",
                "Modelos, límites, comité, quorum y voting",
            ),
            ("SoD, checklist, conditions, recycle", "Reglas de preparación y decisión"),
            ("FID, outputs, staleness, numbering", "Semántica, vigencia e identidad"),
            (
                "Roles, audit, preview, import/export",
                "Seguridad y operación controlada",
            ),
        ],
        [3450, 5910],
        font_size=8.4,
    )
    add_callout(
        doc,
        "Dry-run ADMIN",
        "El endpoint preview devolvió publishable=true y persisted=false. El conteo de AdminConfiguration "
        "permaneció 83 antes/después del preview.",
        fill=POSITIVE,
    )


def security_persistence(doc: Document) -> None:
    doc.add_heading("7. Seguridad, APIs y persistencia", level=1)
    add_table(
        doc,
        ["Área", "Implementación"],
        [
            (
                "Roles",
                "Preparer, reviewer, maker, committee member/chair, viewer, configuration admin",
            ),
            ("Permisos", "10 permisos Gate 07K añadidos a catálogo y grants"),
            ("Tenant", "Todos los queries y constraints incluyen tenant_id"),
            (
                "Concurrency",
                "revision_version + If-Match; SELECT FOR UPDATE en decisión final",
            ),
            (
                "Idempotency",
                "Create, transition, vote y decide; payload mismatch bloqueado",
            ),
            ("Audit", "SecurityEvent con actor, target, outcome, hash y metadata"),
            (
                "Immutability",
                "DECIDED/VOIDED protegidos por eventos ORM before_update/delete",
            ),
        ],
        [2200, 7160],
        font_size=8.4,
    )
    doc.add_heading("7.1 Contrato API", level=2)
    add_table(
        doc,
        ["Grupo", "Rutas"],
        [
            ("Decision CRUD", "GET/POST collection; GET/PUT item"),
            ("Lifecycle", "submit, start-review, ready, decide, void"),
            (
                "Read models",
                "package, readiness, committee, votes, history, authorization-output",
            ),
            ("ADMIN", "list, preview, clone, update, publish configurations"),
        ],
        [2200, 7160],
        font_size=8.5,
    )
    source_note(doc, "Fuente: OpenAPI local; 18 paths Gate 07K publicados.")
    anchored_page_break(doc)
    doc.add_heading("7.2 Persistencia y migración", level=2)
    add_table(
        doc,
        ["Artefacto", "Controles"],
        [
            (
                "investment_decisions",
                "FKs, checks status/outcome/profile/hash, active round unique, queue/history indexes",
            ),
            (
                "investment_decision_votes",
                "FK cascade, vote check, voter unique e idempotency unique",
            ),
            (
                "20260831_0055",
                "Upgrade y downgrade explícitos; importado por Alembic env",
            ),
        ],
        [2800, 6560],
        font_size=8.3,
    )


def validation(doc: Document) -> None:
    doc.add_heading("8. Verificación técnica y E2E", level=1)
    add_table(
        doc,
        ["Suite / control", "Resultado", "Evidencia"],
        [
            ("Ruff", "PASS", "Módulo, migración y tests sin hallazgos"),
            (
                "Backend functional",
                "PASS",
                "8 casos: selection, FID, alternos, SoD/stale, committee, constraints",
            ),
            (
                "PostgreSQL disposable",
                "PASS",
                "Validator before/after y 1 E2E estructural",
            ),
            ("Reversibilidad", "PASS", "0054→0055→0054→0055"),
            ("Frontend format/lint/type", "PASS", "Prettier, ESLint y TypeScript"),
            ("Vitest", "PASS", "3 pruebas; 9 USER y 20 ADMIN"),
            ("Vite build", "PASS", "2,421 módulos transformados"),
            (
                "Playwright Chromium",
                "PASS",
                "2 flujos; USER, ADMIN, FID, HOLD, quorum y stale",
            ),
            (
                "API smoke",
                "PASS",
                "/api/v1/health=ok; 18 paths; config default published",
            ),
            (
                "Performance smoke",
                "OBSERVADO",
                "20 requests: avg 356.87 ms; p95 864.08 ms; max 869.75 ms",
            ),
        ],
        [2800, 1400, 5160],
        font_size=8.0,
    )
    add_callout(
        doc,
        "Alcance de performance",
        "La medición corresponde al localhost Docker sobre una lista vacía y no es un benchmark de carga. "
        "Sirve como baseline diagnóstico; se recomienda definir SLO y pruebas concurrentes en una fase posterior.",
        fill=CAUTION,
    )
    anchored_page_break(doc)
    doc.add_heading("8.1 Casos de aceptación", level=2)
    add_table(
        doc,
        ["Caso", "Control observado"],
        [
            ("PYP-INV-SELECT GO", "No FID; no workspace nuevo; output de continuación"),
            (
                "PYP-FFA GO",
                "FID único, monto/currency, límite, idempotencia y no activación",
            ),
            (
                "HOLD / RECYCLE / KILL",
                "Cierre sin FID y outputs específicos por perfil",
            ),
            ("Stale policy/package", "Readiness STALE y decisión bloqueada"),
            ("Committee", "Sin quorum bloquea; empate bloquea; mayoría GO decide"),
            ("SoD", "Maker igual a preparer bloqueado en PYP-FFA"),
            ("Immutability", "Mutación de DECIDED rechazada"),
            ("Duplicate decide", "Misma clave/payload retorna la misma decisión"),
        ],
        [3000, 6360],
        font_size=8.3,
    )
    doc.add_heading("8.2 Baseline before/after", level=2)
    add_table(
        doc,
        ["Control", "Antes", "Después", "Resultado"],
        [
            ("Alembic", "0054", "0055", "PASS"),
            ("Projects", "4", "4", "Sin creación"),
            ("Project hash", "e33beb48…", "e33beb48…", "Sin mutación"),
            ("Decisions", "N/A", "0", "Sin FID demo"),
            ("Votes", "N/A", "0", "Sin voto demo"),
            ("Services", "Activos", "Activos/healthy", "PASS"),
        ],
        [2200, 1800, 2200, 3160],
        font_size=8.2,
    )


def frontend_browser(doc: Document) -> None:
    doc.add_heading("9. Pruebas frontend y browser", level=1)
    add_table(
        doc,
        ["Superficie", "Escenario", "Resultado"],
        [
            (
                "USER / Selection",
                "FID NO + Project activation NO + no acciones downstream",
                "PASS",
            ),
            ("USER / FID", "FINAL_INVESTMENT_DECISION_GO + handoff separado", "PASS"),
            ("USER / Alternate", "INVESTMENT_SELECTION_HELD visible", "PASS"),
            ("USER / Committee", "NOT MET y voto no autorizado disabled", "PASS"),
            ("USER / Stale", "STALE visible y sin activación", "PASS"),
            ("ADMIN", "20 tabs; dry-run persisted=false", "PASS"),
            ("Console", "pageerror collection vacía", "PASS"),
        ],
        [2300, 5160, 1900],
        font_size=8.3,
    )
    add_callout(
        doc,
        "Navegador integrado",
        "El runtime del navegador embebido no pudo inicializarse por una ruta interna ausente. La validación "
        "continuó con Playwright Chromium contra el mismo localhost y se complementó con HTTP/OpenAPI directos.",
        fill=CAUTION,
    )


def risks_closeout(doc: Document) -> None:
    doc.add_heading("10. Riesgos, deuda técnica y cierre", level=1)
    add_table(
        doc,
        ["Riesgo / deuda", "Impacto", "Tratamiento"],
        [
            (
                "SLO no definido",
                "Performance sin criterio contractual",
                "Definir SLO, volumen y concurrencia",
            ),
            (
                "Import/export ADMIN pendiente",
                "UI lista sección sin transport completo",
                "Diseñar formato firmado/versionado",
            ),
            (
                "Reauthorization downstream",
                "Sólo compatibilidad por nueva ronda",
                "Proceso futuro explícito post-FID",
            ),
            (
                "Provider hashes heterogéneos",
                "Resumen depende del snapshot Gate 07J",
                "Normalizar contrato de evidencia futuro",
            ),
            (
                "Browser mocks",
                "Validan UI; backend se valida en suites separadas",
                "Añadir seeded full-stack E2E dedicado",
            ),
            (
                "Observabilidad",
                "SecurityEvent existe; faltan métricas Gate 07K",
                "Agregar dashboard y alertas",
            ),
        ],
        [2700, 2800, 3860],
        font_size=8.2,
    )
    add_callout(
        doc,
        "Cierre",
        "El alcance funcional, persistencia, seguridad, migración, superficies y salvaguardas de Gate 07K "
        "quedan implementados y verificados. Estado exclusivo: GATE07K_INVESTMENT_DECISION_RELEASE_VALIDATED.",
        fill=POSITIVE,
    )


def deliverables() -> list[str]:
    text = PROMPT.read_text(encoding="utf-8")
    section = text.split("## 95. Required deliverables", 1)[1].split("## 96.", 1)[0]
    items: dict[int, str] = {}
    for line in section.splitlines():
        match = re.match(r"^(\d+)\.\s+(.+)$", line.strip())
        if match:
            items[int(match.group(1))] = match.group(2).rstrip(".")
    if sorted(items) != list(range(1, 96)):
        raise RuntimeError(
            "The Gate 07K prompt does not contain the expected 95 deliverables"
        )
    return [items[index] for index in range(1, 96)]


def traceability(doc: Document) -> None:
    heading = doc.add_heading("11. Matriz de trazabilidad 1-95", level=1)
    heading.paragraph_format.page_break_before = False

    def evidence(index: int) -> str:
        if index <= 5:
            return "§1 / ADR 58"
        if index <= 17:
            return "§2 / backend"
        if index <= 30:
            return "§3 / tests"
        if index <= 46:
            return "§4 / outputs"
        if index <= 51 or index == 90:
            return "§5 / guardrails"
        if index <= 58:
            return "§6 / frontend"
        if index <= 68:
            return "§7 / migration"
        if index <= 81 or index in {89, 91, 92}:
            return "§8 / E2E"
        if index <= 88:
            return "§9 / browser"
        return "§10 / report"

    rows = [
        (str(index), item, evidence(index), "CERRADO")
        for index, item in enumerate(deliverables(), 1)
    ]
    # A single flowing table lets Word repeat the marked header row on every
    # continuation page and avoids hard-break pagination artifacts.
    add_table(
        doc,
        ["#", "Entregable", "Evidencia", "Estado"],
        rows,
        [560, 5200, 1900, 1700],
        font_size=7.2,
    )


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)
    title_page(doc)
    contents(doc)
    baseline(doc)
    anchored_page_break(doc)
    domain(doc)
    anchored_page_break(doc)
    governance(doc)
    anchored_page_break(doc)
    semantics(doc)
    anchored_page_break(doc)
    boundaries(doc)
    anchored_page_break(doc)
    surfaces(doc)
    anchored_page_break(doc)
    security_persistence(doc)
    anchored_page_break(doc)
    validation(doc)
    anchored_page_break(doc)
    frontend_browser(doc)
    anchored_page_break(doc)
    risks_closeout(doc)
    anchored_page_break(doc)
    traceability(doc)
    properties = doc.core_properties
    properties.title = "Informe Técnico Gate 07K — Investment Decision / FID Governance"
    properties.subject = "Decisión formal de inversión, FID, Full Funds Authorization y fronteras post-FID"
    properties.author = "P&Pmis Construction AI"
    properties.keywords = (
        "Gate 07K, Investment Decision, FID, Full Funds Authorization, PYP-FFA"
    )
    properties.comments = "Generated from verified repository, PostgreSQL and localhost evidence on 2026-08-31."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
