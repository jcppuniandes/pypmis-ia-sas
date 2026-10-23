"""Cost report from retained native Word components and executed evidence."""

import hashlib
import importlib.util
import io
import json
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
REFERENCE = (
    ROOT
    / "outputs/wave-o1bb1-20261021/P&Pmis_Wave_O1B_B1_AWP_Readiness_4D_Technical_Report.docx"
)
REFERENCE_SHA = "e794fdacd693fa8cc3e765d2b8a73b91522f4af9e0ac431c80d8fb46d7e0b971"
OUT = ROOT / "outputs/wave-o1bc-20261022"
QA = ROOT / "outputs/.qa/wave-o1bc-20261022"
DATA = ROOT / "docs/validation/wave-o1bc-20261022/package-data.json"
spec = importlib.util.spec_from_file_location(
    "retained_components", ROOT / "scripts/wave-o1bb1/generate_report.py"
)
components = importlib.util.module_from_spec(spec)
spec.loader.exec_module(components)
p, h, table = components.p, components.h, components.table


def main():
    assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_SHA
    data = json.loads(DATA.read_text("utf-8"))
    assert len(data["chapters"]) == 36
    document = Document(REFERENCE)
    logo = deepcopy(document.paragraphs[0]._p)
    section = deepcopy(document._element.body.sectPr)
    for node in list(document._element.body):
        document._element.body.remove(node)
    document._element.body.append(logo)
    document._element.body.append(section)
    p(document, "Wave O1B C\nFundamentos de costos operativos", "Title")
    p(
        document,
        "P&Pmis Construction AI\nInforme técnico de implementación\n23 octubre 2026",
        "Subtitle",
    )
    p(document, data["opening"])
    for index, (title, body) in enumerate(data["chapters"], 1):
        h(document, f"{index} {title}")
        for paragraph in body.split("\n\n"):
            p(document, paragraph)
        if index == 30:
            table(
                document,
                ["Suite", "Casos", "Resultado"],
                data["suite_summary"],
                [3.6, 0.8, 1.5],
            )
        if index == 33:
            p(document, data["volume_summary"])
            table(
                document,
                ["Operación", "N", "Mediana ms", "Máximo ms"],
                [
                    [r[0], str(r[1]), f"{r[3]:.2f}", f"{r[4]:.2f}"]
                    for r in data["performance"]
                ],
                [3.15, 0.7, 1.15, 1.15],
            )
    for paragraph in document.paragraphs:
        paragraph.paragraph_format.widow_control = True
        if paragraph.style.name == "Normal":
            paragraph.paragraph_format.keep_together = True
    OUT.mkdir(parents=True, exist_ok=True)
    QA.mkdir(parents=True, exist_ok=True)
    buffer = io.BytesIO()
    document.save(buffer)
    target_path = OUT / "P&Pmis_Wave_O1B_C_Cost_Foundation_Technical_Report.docx"
    with (
        ZipFile(REFERENCE) as source,
        ZipFile(buffer) as generated,
        ZipFile(target_path, "w") as target,
    ):
        for entry in source.infolist():
            payload = source.read(entry.filename)
            if entry.filename == "word/footer1.xml":
                assert payload.count(b"21 octubre 2026") == 1
                payload = payload.replace(b"21 octubre 2026", b"23 octubre 2026")
            target.writestr(
                entry,
                generated.read(entry.filename)
                if entry.filename == "word/document.xml"
                else payload,
            )
    with ZipFile(REFERENCE) as source, ZipFile(target_path) as target:
        changed = [
            name for name in source.namelist() if source.read(name) != target.read(name)
        ]
        assert set(changed) == {"word/document.xml", "word/footer1.xml"}
        assert source.read("word/footer1.xml").replace(
            b"21 octubre 2026", b"23 octubre 2026"
        ) == target.read("word/footer1.xml")
    proof = {
        "reference_sha256": REFERENCE_SHA,
        "output_sha256": hashlib.sha256(target_path.read_bytes()).hexdigest(),
        "changed_parts": changed,
        "chapters": 36,
    }
    (QA / "documents-build-manifest.json").write_text(
        json.dumps(proof, indent=2), encoding="utf-8"
    )
    print(json.dumps(proof))


if __name__ == "__main__":
    main()
