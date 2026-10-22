"""Build the B1 report from retained Word parts and executed release evidence."""

import hashlib
import importlib.util
import io
import json
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
REFERENCE = (
    ROOT
    / "outputs/wave-o1bb-20261020/P&Pmis_Wave_O1B_B_Schedule_Foundation_Technical_Report.docx"
)
REFERENCE_SHA = "d416d007c65b052b42c172a31ce13f4b5d1736f0f1c8872bdaf2df50d7ba6417"
OUT = ROOT / "outputs/wave-o1bb1-20261021"
QA = ROOT / "outputs/.qa/wave-o1bb1-20261021"
DATA = ROOT / "docs/validation/wave-o1bb1-20261021/package-data.json"
spec = importlib.util.spec_from_file_location(
    "word_components", ROOT / "scripts/wave-o1a/generate_report.py"
)
components = importlib.util.module_from_spec(spec)
spec.loader.exec_module(components)
components.base.REFERENCE = REFERENCE
p, h, table = components.p, components.h, components.table


def build(data):
    document = Document(REFERENCE)
    logo = deepcopy(document.paragraphs[0]._p)
    section = deepcopy(document._element.body.sectPr)
    for node in list(document._element.body):
        document._element.body.remove(node)
    document._element.body.append(logo)
    document._element.body.append(section)
    p(document, "Wave O1B B1\nPaquetes AWP readiness y modelo 4D", "Title")
    p(
        document,
        "P&Pmis Construction AI\nInforme técnico de implementación\n21 octubre 2026",
        "Subtitle",
    )
    p(
        document,
        "Añadimos agrupaciones de ejecución AWP, restricciones operacionales y evaluación de preparación al cronograma gobernado. La reproducción 4D consume el IFC conservado y las fechas de versiones aceptadas o baselines aprobados. "
        + data["release_conclusion"],
    )
    chapters = [
        (
            "Resumen ejecutivo",
            "El incremento incorpora CWA, Work Packages, dependencias, vínculos M:N, Execution Constraints, readiness y 4D Model. Scope sigue siendo la autoridad del alcance y Schedule conserva versiones y fechas importadas. Primavera P6 mantiene el cálculo CPM. El incremento no inicia Cost Foundation ni los procesos completos de ingeniería, compras, logística o construcción.",
        ),
        (
            "Línea base validada",
            "Reutilizamos las fundaciones O1A, O1B A y O1B B, incluidos Project Workspace canónico, autorización operacional, comandos idempotentes, perfiles publicados, Scope Baseline, Plan Items, Derived 3D y Schedule. O1B B ya estaba desarrollado y sin commit al iniciar. Conservamos esos cambios. Esta validación utiliza API 8854 y frontend 5514, con bases PostgreSQL propias y volúmenes independientes de la aplicación habitual.",
        ),
        (
            "Hallazgos del repositorio",
            "El WorkPackage legado pertenece al identificador de proyecto histórico e incluye fechas y progreso que no deben convertirse en autoridad de AWP canónico. Sus restricciones tampoco resuelven los vínculos M:N requeridos. Se conservó ese código y se añadió un dominio awp separado. La implementación 4D reutiliza fuentes y modelos de Scope, sin otra tabla maestra de geometría. App.tsx, BimIfcModelViewer.tsx y los modelos y esquemas globales permanecen intactos.",
        ),
        (
            "Autoridad de AWP",
            "AWP gobierna agrupaciones y condiciones de ejecución. No redefine cantidades, alcance aprobado, calendarios ni lógica P6. Un paquete puede contener varios Scope Items y Plan Items y relacionarse con varias Activities. El identificador de cada referencia incluye tenant y Project Workspace. El mapeo Activity conserva además la versión del cronograma para evitar confundir identidades entre fuentes.",
        ),
        (
            "Construction Work Areas",
            "CWA representa un área o contenedor. Conserva código, nombre, descripción, área, zona, localización, padre y estado. La jerarquía rechaza ciclos y padres inactivos. Las revisiones incrementan ETag y guardan un snapshot de historia. Una CWA puede agrupar varios tipos de paquete y no implica una relación padre hijo EWP PWP CWP IWP.",
        ),
        (
            "Tipos de Work Package",
            "Los tipos base son EWP, PWP, CWP, IWP y TWP. La configuración publicada también habilita SUPPORT, LOGISTICS y TEMPORARY_WORKS. Cada paquete conserva identidad, descripción, CWA opcional, responsable, estado, revisión, readiness, release y autores. Sus contratos de entrada rechazan fechas CPM o valores de progreso. Las fechas mostradas proceden de Activities ya existentes.",
        ),
        (
            "Mapeos entre Scope Plan y paquetes",
            "Los vínculos de paquetes con Scope, Plan y Activity se añaden como revisiones de mapping con acciones SET o REMOVE. Scope debe pertenecer a la baseline aprobada y Plan debe estar publicado. Un paquete admite varios objetos y un objeto varios paquetes. Las fases EXECUTE, INSTALL, USE y REMOVE expresan la intención temporal sin deducirla del nombre de una actividad. Los mapeos obsoletos conservan su historia y muestran STALE, BROKEN o REQUIRES_REVIEW.",
        ),
        (
            "Dependencias de paquetes",
            "Cada dependencia identifica predecesor, sucesor, tipo, bloqueo, evidencia y estado. El grafo no impone una cadena fija por tipos AWP. Una dependencia bloqueante activa exige completar o cerrar su predecesor para alcanzar readiness. La validación detecta ciclos sin recurrencia ilimitada. Los cambios quedan en la historia del agregado y no crean relaciones P6.",
        ),
        (
            "Ciclo de vida de paquetes",
            "El ciclo publicado distingue DRAFT, PLANNED, UNDER_REVIEW, READY, RELEASED, IN_PROGRESS, COMPLETED, CLOSED y SUPERSEDED. La evaluación no cambia por sí sola el estado del ciclo de vida. READY exige una evaluación vigente con requisitos satisfechos. RELEASED se emite con una acción autorizada independiente. Completar el paquete no despeja restricciones de manera implícita.",
        ),
        (
            "Execution Constraints",
            "Las restricciones guardan tipo, código, título, descripción, bloqueo, responsable, vencimiento, dominio de origen, referencias, evidencia y autorización de waiver. El catálogo incluye diecisiete tipos desde ENGINEERING hasta OTHER. Los estados OPEN, IN_PROGRESS, CLEARED, WAIVED y CANCELLED se gestionan mediante comandos versionados. El vencimiento no equivale a clearance. Se conservan autor, fecha e historia de los cambios.",
        ),
        (
            "Separación de restricciones P6",
            "ExecutionConstraint pertenece a AWP y se almacena en awp_execution_constraints. Las Schedule Constraints continúan dentro del snapshot P6 inmutable. Despejar o dispensar una restricción de ejecución no modifica lag, calendario, duración, fecha, relación ni restricción importada. Los contratos y las pruebas de integridad verifican la separación de ambas autoridades.",
        ),
        (
            "Evaluación de readiness",
            "PackageReadinessAssessment conserva fecha as of, política publicada, evidencia, restricciones, dependencias, dimensiones y hashes de las fuentes. Los resultados son NOT_ASSESSED, NOT_READY, PARTIALLY_READY, READY y RELEASED. La evaluación es determinista y no calcula CPM. Si cambian fuentes o política, una evaluación antigua deja de servir como autorización de release. Los registros históricos permanecen inmutables.",
        ),
        (
            "Workfront Readiness",
            "La pantalla relaciona CWA, paquetes, estado de mapping Schedule, bloqueos y release. Las dimensiones son Engineering, Material y Procurement, Logistics, Access, Permit, Quality, HSE, Resource y Equipment, y Predecessor. ADMIN determina cuáles son obligatorias por tipo. Una actividad puede estar programada mientras su paquete sigue NOT_READY. El usuario ve la evidencia faltante y los hashes que exigen reevaluar.",
        ),
        (
            "Integración entre paquetes y Activities",
            "Cada mapping referencia una Activity externa de una versión aceptada. No crea Activities ni sobrescribe ScheduleVersion. Las asociaciones son M:N y permiten representar instalación, uso y retiro de un elemento temporal con actividades distintas. Las actualizaciones del cronograma conservan las referencias históricas y señalan identidades que necesitan revisión o han desaparecido.",
        ),
        (
            "Protección del cronograma",
            "La integración escribe únicamente en los agregados AWP y sus recibos de comando. Las versiones aceptadas, sus snapshots, las baselines aprobadas, la fuente IFC y el Scope aprobado no se recalculan. Las pruebas comparan contenido y hashes antes y después de operaciones AWP y consultas 4D. No se añadió sincronización con API P6 ni un motor de ruta crítica.",
        ),
        (
            "Arquitectura del modelo 4D",
            "El servicio construye una proyección de lectura identificada por versión y hashes. Resuelve IFC y objetos de Scope, Plan Items, Derived 3D, paquetes, WBS y Activities. No persiste una segunda geometría ni otro cronograma maestro. El componente de geometría consume los bytes inmutables mediante el proveedor canónico y usa web ifc y Three.js para aplicar visibilidad y colores transitorios. Las consultas GET no generan filas de negocio ni auditoría por cada movimiento temporal.",
        ),
        (
            "Semántica 4D del alcance",
            "INSTALL pasa de futuro no instalado a instalación y después instalado. REMOVE muestra el existente antes del retiro y lo oculta al finalizar. TEMPORARY exige fases explícitas de instalación y retiro, con una etapa visible de uso entre ambas. MAINTENANCE conserva el elemento existente, identifica la intervención y muestra el servicio restaurado. Una relación temporal incompleta o ambigua queda UNKNOWN. El estado proyectado por fecha no certifica avance físico.",
        ),
        (
            "Baseline y Current en 4D",
            "El selector permite Approved Baseline o Accepted Current. Una versión staged o rechazada no puede presentarse como Current. La baseline aprobada mantiene su snapshot aunque su versión de origen quede supersedida. La comparación usa la misma fecha y el mismo registro para mostrar alineación o diferencia de estados. Las desviaciones no producen fechas forecast ni float nuevo.",
        ),
        (
            "Vista Data Date",
            "AS OF DATA DATE utiliza la fecha y hora del cronograma seleccionado, incluida su zona horaria. Se muestran versión, baseline o current y asociación con Project Controls Period cuando existe. La ausencia de asociación se informa, no se fabrica un período. El Data Date no se sustituye por el día del equipo ni por el cutoff. La prueba utiliza fechas distintas entre baseline y current.",
        ),
        (
            "Trazabilidad del modelo",
            "La inspección presenta modelo original, GUID y ExpressID dependiente de versión, Scope Item, Scope Nature, Plan, Derived 3D, CWA, paquetes, Activities, WBS y referencias Schedule. También muestra readiness y restricciones de ejecución. El filtro por WBS incluye descendientes y los demás filtros conservan sus identidades canónicas. La selección en la geometría permite abrir los mismos datos de trazabilidad.",
        ),
        (
            "Separación de ADMIN y USER",
            "ADMIN administra políticas AWP, tipos, requisitos de readiness, workflow de release y perfiles de presentación 4D. El borrador y su publicación tienen permisos distintos y separación de autor y publicador. USER trabaja dentro del Project Workspace autorizado. Las nuevas pantallas reutilizan la navegación existente. El dominio AWP ofrece subpantallas propias y Schedule Manager incorpora 4D Model como su décima pestaña.",
        ),
        (
            "Seguridad y separación de funciones",
            "Los permisos distinguen lectura, configuración, publicación, CWA, edición y revisión de paquetes, release, mapping, dependencias y raise, clear o waive de constraints. El preparador, responsable o evaluador no puede liberar su propio paquete. Quien plantea o posee una restricción no puede despejarla cuando la política lo prohíbe. Los comandos usan reason, If Match, Idempotency Key y auditoría. No se registraron firmas ni aceptación humana de negocio.",
        ),
        (
            "API",
            "Los contratos nuevos se exponen bajo api v1 awp y las rutas 4d del dominio Schedule. Las lecturas separan colecciones, overview, historia, readiness, modelo, comparación y trazabilidad. Las escrituras emplean POST o PUT versionados. Los errores de precondición y permisos no se convierten en resultados correctos. El libro de trazabilidad contiene el inventario de rutas verificadas en la aplicación montada.",
        ),
        (
            "Frontend",
            "Las pantallas de AWP incluyen formularios, tablas, revisión de mapeos, dependencias, evidencia y dimensiones de readiness. 4D ofrece fecha, play, pause, pasos, velocidad, salto a Data Date y hitos cuando existan, selección de fuente, comparación y filtros. El visor informa la ausencia de geometría o correspondencias en vez de mostrar un reemplazo ficticio. Las pruebas de componentes se complementan con la ejecución real de Chromium.",
        ),
        (
            "Migración",
            "La revisión 20261021 0072 sucede a 0071 y añade ocho tablas de AWP y guards PostgreSQL para ámbito e inmutabilidad. No reescribe migraciones históricas ni aplica backfill a proyectos existentes. El bootstrap de validación crea una base ORM descartable anterior a O1A y luego ejecuta las migraciones reales aditivas. Eso no constituye una prueba de migración vacía hasta head. El ciclo upgrade downgrade y reupgrade utiliza otra base aislada.",
        ),
        ("Validación PostgreSQL", data["postgres_summary"]),
        ("Validación Chromium", data["chromium_summary"]),
        ("Preservación del Proyecto 14", data["preservation_summary"]),
        (
            "Observaciones de desempeño",
            "Las mediciones corresponden al entorno local aislado y a los volúmenes descritos en el libro. Se observan listados CWA, paquetes, dependencias y constraints, assessment, mapping, carga inicial 4D, paso temporal, trazabilidad y comparación. Los tiempos de comandos UI incluyen navegación y espera de respuesta. No se declara SLA, validación industrial ni comportamiento a escala productiva.",
        ),
        ("Brechas restantes", data["gaps_summary"]),
        (
            "Decisión de liberación",
            data["release_conclusion"]
            + " Estado exacto: "
            + data["release_state"]
            + ". La aplicación habitual no se migró durante esta validación. No se efectuó commit ni push.",
        ),
        (
            "Siguiente incremento recomendado",
            "WAVE O1B-C — COST FOUNDATION. Su implementación requiere una solicitud posterior. No se implementó en esta entrega.",
        ),
    ]
    assert len(chapters) == 32
    for index, (title, body) in enumerate(chapters, 1):
        h(document, f"{index} {title}")
        p(document, body)
        if index == 26:
            table(
                document,
                ["Suite", "Casos", "Resultado"],
                data["suite_summary"],
                [3.6, 0.8, 1.5],
            )
        if index == 29:
            volume = data["data_volume"]
            p(
                document,
                f"Volumen observado: {volume['packages']} paquetes, {volume['scope_items']} Scope Items, {volume['activities_per_version']} Activities por versión y {volume['ifc_objects']} objetos IFC. Las fuentes incluyen dos versiones aceptadas históricamente y una baseline aprobada.",
            )
            table(
                document,
                ["Operación", "N", "Mediana ms", "Máximo ms"],
                [
                    [row[0], str(row[1]), f"{row[3]:.2f}", f"{row[4]:.2f}"]
                    for row in data["performance"]
                ],
                [3.15, 0.7, 1.15, 1.15],
            )
    for paragraph in document.paragraphs:
        paragraph.paragraph_format.widow_control = True
        if paragraph.style.name == "Normal":
            paragraph.paragraph_format.keep_together = True
    return document


def main():
    assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_SHA
    data = json.loads(DATA.read_text("utf-8"))
    OUT.mkdir(parents=True, exist_ok=True)
    QA.mkdir(parents=True, exist_ok=True)
    doc = build(data)
    buffer = io.BytesIO()
    doc.save(buffer)
    output = OUT / "P&Pmis_Wave_O1B_B1_AWP_Readiness_4D_Technical_Report.docx"
    with (
        ZipFile(REFERENCE) as source,
        ZipFile(buffer) as generated,
        ZipFile(output, "w") as target,
    ):
        for entry in source.infolist():
            target.writestr(
                entry,
                generated.read(entry.filename)
                if entry.filename == "word/document.xml"
                else source.read(entry.filename),
            )
    with ZipFile(REFERENCE) as source, ZipFile(output) as target:
        changed = [
            name for name in source.namelist() if source.read(name) != target.read(name)
        ]
        assert changed == ["word/document.xml"]
    proof = {
        "reference_sha256": REFERENCE_SHA,
        "output_sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
        "changed_parts": changed,
        "chapters": 32,
    }
    (QA / "documents-build-manifest.json").write_text(
        json.dumps(proof, indent=2), encoding="utf-8"
    )
    print(json.dumps(proof))


if __name__ == "__main__":
    main()
