"""Create the Wave 2B technical report using the established report primitives."""

from __future__ import annotations

import re
from pathlib import Path

import generate_cpig_gap_review_report as base
import generate_wave3a_portfolio_review_report as layout
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/77-wave2b-risk-governance-basic-data-release.md"
QA = ROOT / "outputs/.qa/wave2b-risk-basic-data"
DATE = "2026-10-07"
OUTPUT = ROOT / "outputs" / f"Informe_Tecnico_Wave2B_Risk_Governance_Basic_Data_Readiness_{DATE}.docx"
DIAGRAMS = [
    {
        "name": "wave2b-domain-authority.png",
        "title": "Política empresarial y evidencia técnica separadas",
        "rows": [
            ["ADMIN Risk Policy", "Capacity Appetite Thresholds", "Read-only Risk Context"],
            ["ADMIN Basic Data Policy", "USER Technical Evidence", "Immutable Assessment"],
            ["FEL PDRI Context", "Basic Data Provider", "Future Gate Consumers"],
        ],
        "edges": [(0, 0, 0, 1), (0, 1, 0, 2), (1, 0, 1, 1), (1, 1, 1, 2), (2, 0, 1, 2), (1, 2, 2, 1), (2, 1, 2, 2)],
        "alt": "Risk Governance publica contexto empresarial independiente. Basic Data aplica requisitos a evidencia técnica y congela evaluaciones. FEL y PDRI aportan solo contexto; futuros consumidores Gate requieren integración explícita.",
    },
    {
        "name": "wave2b-assessment-history.png",
        "title": "Evaluación versionada y preservación de la historia",
        "rows": [
            ["Items and References", "DRAFT Assessment", "IN REVIEW"],
            ["COMPLETED Snapshot", "Current Freshness Check", "STALE or READY"],
            ["Immutable History", "Refresh Preview", "Successor DRAFT"],
        ],
        "edges": [(0, 0, 0, 1), (0, 1, 0, 2), (0, 2, 1, 0), (1, 0, 1, 1), (1, 1, 1, 2), (1, 0, 2, 0), (1, 2, 2, 1), (2, 1, 2, 2)],
        "alt": "Los datos alimentan una evaluación DRAFT revisada y completada. El snapshot histórico permanece inmutable. La lectura actual puede detectar staleness; Refresh Preview prepara una nueva evaluación sin cambiar la anterior.",
    },
]


def render_diagram(spec: dict) -> Path:
    """Route cross-row evidence links through whitespace, never through nodes."""
    canvas = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(canvas)
    draw.text((80, 46), spec["title"], fill=(0, 0, 0), font=base.font("arialbd.ttf", 42))
    face = base.font("arial.ttf", 32)
    rows = spec["rows"]
    row_height = 655 / len(rows)
    positions = {}
    for row_index, labels in enumerate(rows):
        box_width = min(390, int((1440 - 28 * (len(labels) - 1)) / len(labels)))
        left = (1600 - box_width * len(labels) - 28 * (len(labels) - 1)) // 2
        box_height = min(130, int(row_height * 0.62))
        top = int(165 + row_index * row_height + (row_height - box_height) / 2)
        for column in range(len(labels)):
            x = left + column * (box_width + 28)
            positions[row_index, column] = (x, top, x + box_width, top + box_height)
    for r1, c1, r2, c2 in spec["edges"]:
        a, b = positions[r1, c1], positions[r2, c2]
        if r1 == r2:
            start = (a[2] + 3, (a[1] + a[3]) // 2)
            end = (b[0] - 5, (b[1] + b[3]) // 2)
            base.draw_arrow(draw, start, end)
            continue
        upward = r2 < r1
        offset = 45 if upward else -45
        start = ((a[0] + a[2]) // 2 + offset, a[1] - 3 if upward else a[3] + 3)
        end = ((b[0] + b[2]) // 2 + offset, b[3] + 5 if upward else b[1] - 5)
        if c1 == c2:
            base.draw_arrow(draw, start, end)
        else:
            channel = int((start[1] + end[1]) / 2) + (15 if upward else -15)
            turn1, turn2 = (start[0], channel), (end[0], channel)
            draw.line([start, turn1, turn2], fill=(83, 97, 107), width=5)
            base.draw_arrow(draw, turn2, end)
    fills = [(234, 242, 247), (232, 245, 243), (233, 245, 238)]
    for (row_index, column), box in positions.items():
        draw.rounded_rectangle(box, radius=18, fill=fills[row_index % 3], outline=(23, 50, 77), width=4)
        lines = base.wrap(draw, rows[row_index][column], face, box[2] - box[0] - 34)
        y = box[1] + (box[3] - box[1] - len(lines) * 35) // 2
        for line in lines:
            bounds = draw.textbbox((0, 0), line, font=face)
            x = box[0] + (box[2] - box[0] - bounds[2] + bounds[0]) // 2
            draw.text((x, y), line, fill=(32, 42, 51), font=face)
            y += 35
    path = base.DIAGRAM_DIR / spec["name"]
    canvas.save(path, format="PNG", optimize=True)
    return path


def main() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    status_match = re.search(r"^\*\*Estado actual:\*\*\s+(.+)$", text, re.MULTILINE)
    status = status_match.group(1).strip().strip("`") if status_match else "VALIDACIÓN EN CURSO"
    if status == "RISK_GOVERNANCE_AND_BASIC_DATA_READINESS_RELEASE_VALIDATED":
        status = "RELEASE VALIDADO"
    QA.mkdir(parents=True, exist_ok=True)
    base.SOURCE, base.QA, base.DIAGRAM_DIR = SOURCE, QA, QA / "diagrams"
    base.DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
    paths = [render_diagram(spec) for spec in DIAGRAMS]
    base.DIAGRAMS = []
    doc = Document()
    base.configure_document(doc)
    base.set_language(doc)
    for style_name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        properties = doc.styles[style_name]._element.find(qn("w:pPr"))
        if properties is not None:
            for border in properties.findall(qn("w:pBdr")):
                properties.remove(border)
    doc.core_properties.title = "Informe técnico Wave 2B Risk Governance and Basic Data Readiness"
    doc.core_properties.subject = "Configuración empresarial de riesgo y evidencia técnica gobernada por proyecto"
    doc.core_properties.author = "P&Pmis Construction AI"
    doc.core_properties.keywords = "CPIG, Wave 2B, Risk Governance, Basic Data, Readiness, Evidence"
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(30)
    if layout.LOGO.exists():
        shape = paragraph.add_run().add_picture(str(layout.LOGO), width=Inches(1.3))
        base.set_alt_text(shape, "P&Pmis Construction AI", "Logotipo de P&Pmis Construction AI")
    title = doc.add_paragraph("Informe técnico Wave 2B\nRisk Governance and Basic Data Readiness", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(25)
    for value in ("P&Pmis Construction AI", "Capital Portfolio and Investment Governance", DATE, status):
        paragraph = doc.add_paragraph(value)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(15)
    doc.add_page_break()
    doc.add_heading("Control del documento", level=1)
    base.add_table(doc, [
        ["Campo", "Valor"], ["Producto", "P&Pmis Construction AI"],
        ["Incremento", "Wave 2B Risk Governance and Basic Data Readiness"],
        ["Baseline", "main / 9cb0969 con cambios locales anteriores preservados"],
        ["Migración", "20261005_0064 a 20261008_0065"], ["Fecha", DATE], ["Estado", status],
    ])
    doc.add_heading("Lectura ejecutiva", level=1)
    base.add_inline(doc.add_paragraph(), "La capacidad, el apetito y los umbrales de riesgo son políticas empresariales. Basic Data evalúa la suficiencia, validación, calidad, confianza y actualidad de los datos técnicos de un proyecto. Este informe presenta la implementación, su reutilización, las fronteras de autoridad y las evidencias de validación.")
    doc.add_heading("Contenido", level=1)
    for item in (
        "Baseline y decisiones arquitectónicas", "Política empresarial de riesgo y resolución efectiva",
        "Datos técnicos y evaluación de readiness", "Integraciones de solo lectura y proveedores",
        "Seguridad y persistencia", "Pruebas y cierre del release",
    ):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_page_break()
    start = len(doc.paragraphs)
    base.add_source_content(doc, [], page_break_heading_prefixes=(), skip_preface=True)
    for prefix, path, spec in zip(("4. ", "29. "), paths, DIAGRAMS, strict=True):
        anchor = next(p for p in doc.paragraphs if p.style.name.startswith("Heading") and p.text.startswith(prefix))
        figure = doc.add_paragraph()
        figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
        figure.paragraph_format.keep_together = True
        shape = figure.add_run().add_picture(str(path), width=Inches(6.75))
        base.set_alt_text(shape, spec["title"], spec["alt"])
        anchor._p.addnext(figure._p)
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.style.name.startswith("Heading"):
            paragraph.text = re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", paragraph.text).replace("_", " ")).strip()
        elif index >= start and paragraph.style.name in ("Normal", "List Bullet", "List Number") and "w:drawing" not in paragraph._p.xml:
            for run in paragraph.runs:
                run.font.size = Pt(11)
    layout.format_tables(doc)
    widths_by_headers = {
        ("Dominio", "Autoridad de escritura", "Límite"): (1.65, 2.40, 3.01),
        ("Componente existente", "Uso en Wave 2B", "Decisión"): (1.55, 2.45, 3.06),
        ("Acción", "ADMIN Configuration", "USER operativo"): (2.45, 2.45, 2.16),
        ("Prerrequisito", "Estado arquitectónico", "Brecha para QRA"): (1.60, 2.35, 3.11),
        ("ID", "Criterio de aceptación", "Evidencia verificada", "Estado"): (0.36, 2.80, 3.15, 0.75),
        ("Validación", "Resultado", "Evidencia"): (2.30, 1.30, 3.46),
    }
    for table in doc.tables:
        headers = tuple(cell.text for cell in table.rows[0].cells)
        widths = widths_by_headers.get(headers)
        if widths:
            for column, width in zip(table.columns, widths, strict=True):
                column.width = Inches(width)
            for row in table.rows:
                for index, cell in enumerate(row.cells):
                    cell.width = Inches(widths[index])
    for section in doc.sections:
        section.different_first_page_header_footer = True
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        base.set_run_font(header.add_run("P&Pmis Construction AI | Wave 2B"), size=7.5, color=base.BLACK)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        base.set_run_font(footer.add_run(f"Informe técnico | {DATE} | Página "), size=7.5, color=base.MUTED)
        base.add_field(footer, "PAGE")
        footer.add_run(" de ")
        base.add_field(footer, "NUMPAGES")
    base.update_field_settings(doc)
    doc.save(OUTPUT)
    print(f"created={OUTPUT}")


if __name__ == "__main__":
    main()
