"""Generate the reviewed Wave 4 release report, without editing prior reports."""

from __future__ import annotations

import re
from pathlib import Path

import generate_cpig_gap_review_report as base
import generate_qra_foundation_report as foundation
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "83-qra-engine-release.md"
OUTPUT = ROOT / "outputs" / "Informe_Tecnico_Wave4_QRA_Calculation_Engine_2026-10-12.docx"
LOGO = ROOT / "frontend" / "public" / "pypmis-construction-ai-logo.png"


def math_run(text):
    run = OxmlElement("m:r")
    value = OxmlElement("m:t")
    value.text = text
    run.append(value)
    return run


def math_fraction(numerator, denominator):
    fraction = OxmlElement("m:f")
    for name, parts in (("m:num", numerator), ("m:den", denominator)):
        container = OxmlElement(name)
        for part in parts:
            container.append(part)
        fraction.append(container)
    return fraction


def math_subscript(text, subscript):
    node = OxmlElement("m:sSub")
    for name, value in (("m:e", text), ("m:sub", subscript)):
        container = OxmlElement(name)
        container.append(math_run(value))
        node.append(container)
    return node


def math_sqrt(text):
    node = OxmlElement("m:rad")
    props = OxmlElement("m:radPr")
    hidden = OxmlElement("m:degHide")
    hidden.set(qn("m:val"), "1")
    props.append(hidden)
    node.append(props)
    node.append(OxmlElement("m:deg"))
    element = OxmlElement("m:e")
    element.append(math_run(text))
    node.append(element)
    return node


def add_native_equations(document):
    """Use editable OMML fractions and subscripts, verified in native Word QA."""
    for paragraph in list(document.paragraphs):
        text = paragraph.text
        equations = []
        if text.startswith("Para mínimo a, modo m"):
            paragraph.text = "Para mínimo a, modo m y máximo b se exige que el modo pertenezca al intervalo. La inversa usa el punto de separación y dos ramas cuadráticas; el intervalo degenerado es determinista. Se validan soporte, media teórica y reproducción con el mismo manifiesto."
            equations = [
                [math_subscript("u", "0"), math_run(" = "), math_fraction([math_run("m − a")], [math_run("b − a")])],
                [math_run("μ = "), math_fraction([math_run("a + m + b")], [math_run("3")])],
            ]
        elif text.startswith("Con peso λ explícito"):
            paragraph.text = "Con peso λ explícito, los parámetros α y β y la media siguen las fórmulas siguientes. Se aplica betaincinv y se escala al intervalo entre a y b. Un intervalo degenerado produce una constante. No se introduce un peso oculto ni una beta distinta de la metodología."
            equations = [
                [math_run("α = 1 + λ "), math_fraction([math_run("m − a")], [math_run("b − a")])],
                [math_run("β = 1 + λ "), math_fraction([math_run("b − m")], [math_run("b − a")])],
                [math_run("μ = "), math_fraction([math_run("a + λm + b")], [math_run("λ + 2")])],
            ]
        elif text.startswith("La inversa a+"):
            paragraph.text = "La inversa uniforme preserva los límites y permite un caso constante. Las pruebas contrastan media, desviación teórica y percentiles. Una amplitud negativa o parámetros no finitos impiden cálculo; no se intercambian límites silenciosamente."
            equations = [[math_run("x = a + (b − a)u     μ = "), math_fraction([math_run("a + b")], [math_run("2")])]]
        elif text.startswith("Se admiten FS, SS, FF y SF en un DAG"):
            paragraph.text = "Se admiten FS, SS, FF y SF en un grafo sin ciclos. Para el inicio temprano del sucesor j se aplican las restricciones siguientes para cada predecesor i. Se toma el máximo entre cero y todas las restricciones; el fin temprano suma la duración. Ciclos, referencias externas, tipos desconocidos o unidades de lag distintas se rechazan."
            for label, expression in (
                ("FS", [math_subscript("ES", "i"), math_run(" + "), math_subscript("d", "i"), math_run(" + lag")]),
                ("SS", [math_subscript("ES", "i"), math_run(" + lag")]),
                ("FF", [math_subscript("ES", "i"), math_run(" + "), math_subscript("d", "i"), math_run(" + lag − "), math_subscript("d", "j")]),
                ("SF", [math_subscript("ES", "i"), math_run(" + lag − "), math_subscript("d", "j")]),
            ):
                equations.append([math_run(label + "     "), math_subscript("ES", "j"), math_run(" ≥ "), *expression])
            equations.append([math_subscript("EF", "j"), math_run(" = "), math_subscript("ES", "j"), math_run(" + "), math_subscript("d", "j")])
        elif text.startswith("Estados CONVERGED, NOT_CONVERGED"):
            paragraph.text = text.replace("s/√N", "SE")
            equations = [[math_run("SE = "), math_fraction([math_run("s")], [math_sqrt("N")])]]
        elif text.startswith("Interpolación lineal Hyndman"):
            paragraph.text = text.replace("h=(N−1)p/100", "con posición h definida a continuación")
            equations = [[math_run("h = "), math_fraction([math_run("(N − 1)p")], [math_run("100")])]]
        if equations:
            paragraph.paragraph_format.keep_with_next = True
        cursor = paragraph._p
        for equation_index, parts in enumerate(equations):
            node = OxmlElement("w:p")
            props = OxmlElement("w:pPr")
            spacing = OxmlElement("w:spacing")
            spacing.set(qn("w:before"), "80")
            spacing.set(qn("w:after"), "100")
            props.append(spacing)
            if equation_index < len(equations) - 1:
                props.append(OxmlElement("w:keepNext"))
            node.append(props)
            display = OxmlElement("m:oMathPara")
            equation = OxmlElement("m:oMath")
            for part in parts:
                equation.append(part)
            display.append(equation)
            node.append(display)
            cursor.addnext(node)
            cursor = node


def layout_tables(document):
    for table in document.tables:
        count = len(table.columns)
        if count == 6:
            widths = (1.7, 0.65, 0.65, 1.15, 1.05, 1.3)
        elif count == 3:
            widths = (1.6, 2.35, 2.55)
        elif count == 2:
            widths = (1.45, 5.05)
        else:
            widths = tuple(6.5 / count for _ in range(count))
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        for row in table.rows:
            for index, (cell, width) in enumerate(zip(row.cells, widths)):
                cell.width = Inches(width)
                if count == 6 and index > 0:
                    for paragraph in cell.paragraphs:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def main():
    content = SOURCE.read_text(encoding="utf-8")
    numbers = [int(value) for value in re.findall(r"^## (\d+)\. ", content, re.MULTILINE)]
    if numbers != list(range(1, 94)):
        raise ValueError("Expected all 93 required report sections in order")
    match = re.search(r"^\*\*Estado actual:\*\*\s+(.+)$", content, re.MULTILINE)
    status = match.group(1) if match else "VALIDACIÓN EN CURSO"
    document = Document()
    base.configure_document(document)
    base.set_language(document)
    document.styles["Normal"].font.size = Pt(11)
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        props = document.styles[name]._element.find(qn("w:pPr"))
        if props is not None:
            for border in props.findall(qn("w:pBdr")):
                props.remove(border)
    document.core_properties.title = "Informe técnico Wave 4 QRA Calculation Engine"
    document.core_properties.subject = "Cost and Schedule QRA; governed analytical results"
    document.core_properties.author = "P&Pmis Construction AI"
    document.core_properties.keywords = "QRA, Monte Carlo, Cost, Schedule, governance, Wave 4"
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(28)
    if LOGO.exists():
        shape = paragraph.add_run().add_picture(str(LOGO), width=Inches(1.3))
        base.set_alt_text(shape, "P&Pmis Construction AI", "Logotipo azul de P&Pmis Construction AI")
    paragraph = document.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(22)
    paragraph.add_run("Informe técnico\nWave 4\nQRA Calculation Engine")
    for text, before, size in (
        ("P&Pmis Construction AI", 12, 13),
        ("Capital Portfolio and Investment Governance", 22, 10),
        ("Corte técnico: 12 de octubre de 2026", 6, 10),
        (status.replace("_", " "), 38, 10),
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(before)
        base.add_inline(paragraph, text, size=size)
    document.add_page_break()
    document.add_heading("Control y alcance del documento", level=1)
    base.add_table(document, [
        ["Campo", "Detalle"],
        ["Incremento", "Wave 4: motor QRA de costo y cronograma"],
        ["Baseline", "main / d762d873; cambios previos Foundation preservados"],
        ["Persistencia", "Alembic 20261011_0068 sobre 20261009_0067"],
        ["Resultado", status],
        ["Fuera de alcance", "QRA integrado, FID automático y aprobación de contingencia"],
        ["Control de cambios", "Desarrollo local; sin commit ni push del agente"],
    ])
    paragraph = document.add_paragraph()
    base.add_inline(paragraph, "El documento separa implementación, pruebas realmente ejecutadas, límites metodológicos y deuda técnica. Una corrida completada no constituye aprobación de inversión. La aceptación independiente solo habilita evidencia analítica gobernada.")
    document.add_heading("Guía de lectura", level=1)
    for text in (
        "1–13: baseline, autoridad, dominio, validación y reproducción",
        "14–43: matemática, distribuciones, correlación, costo, red y sensibilidad",
        "44–69: gobierno de resultados, seguridad, UX, rendimiento, migración y API",
        "70–93: pruebas, navegador, performance, limitaciones y decisión de cierre",
    ):
        base.add_inline(document.add_paragraph(style="List Bullet"), text)
    document.add_page_break()
    base.SOURCE, base.DIAGRAMS = SOURCE, []
    source_start = len(document.paragraphs)
    base.add_source_content(document, [], page_break_heading_prefixes=(), skip_preface=True)
    add_native_equations(document)
    for index, paragraph in enumerate(document.paragraphs):
        if paragraph.style.name.startswith("Heading"):
            paragraph.text = re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", paragraph.text).replace("_", " ")).strip()
        elif index >= source_start and paragraph.style.name in {"Normal", "List Bullet", "List Number"}:
            for run in paragraph.runs:
                run.font.size = Pt(11)
    foundation.format_tables(document)
    layout_tables(document)
    for section in document.sections:
        section.different_first_page_header_footer = True
        paragraph = section.header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        base.set_run_font(paragraph.add_run("P&Pmis Construction AI | Wave 4 QRA"), size=7.5, color=base.BLACK)
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        base.set_run_font(paragraph.add_run("Informe técnico | 2026-10-12 | Página "), size=7.5, color=base.MUTED)
        base.add_field(paragraph, "PAGE")
        base.set_run_font(paragraph.add_run(" de "), size=7.5, color=base.MUTED)
        base.add_field(paragraph, "NUMPAGES")
    base.update_field_settings(document)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(f"created={OUTPUT}")
    print(f"sections={len(numbers)}")
    print(f"status={status}")


if __name__ == "__main__":
    main()
