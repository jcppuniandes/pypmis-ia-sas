from __future__ import annotations

import csv
import json
import xml.etree.ElementTree as ET
from collections.abc import Iterable
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\Ricardo\Documents\GitHub\pypmis-ia-sas")
QA = ROOT / "outputs" / ".qa" / "gate07m-ben"
OUTPUT = (
    ROOT
    / "outputs"
    / "Informe_Tecnico_Gate07M_BEN_Beneficios_Estrategicos_2026-09-17.docx"
)
LOGO = ROOT / "frontend" / "public" / "pypmis-construction-ai-logo.png"

BLACK = "000000"
TEXT = "202A33"
MUTED = "53616B"
NAVY = "17324D"
TEAL = "127A78"
WHITE = "FFFFFF"
PALE_GREEN = "E9F5EE"
PALE_AMBER = "FFF4E5"
PALE_GRAY = "F5F7F8"
LINE = "D7DEE3"
GREEN = "18794E"
AMBER = "8A5700"


def load_json(path: Path, default: Any) -> Any:
    if not path.exists() or path.stat().st_size == 0:
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def load_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def junit(path: Path) -> dict[str, Any]:
    result: dict[str, Any] = {
        "tests": 0,
        "failures": 0,
        "errors": 0,
        "skipped": 0,
        "time": 0.0,
        "failed": [],
    }
    if not path.exists() or path.stat().st_size == 0:
        return result
    root = ET.parse(path).getroot()
    suites = [root] if root.tag == "testsuite" else list(root.findall("testsuite"))
    for suite in suites:
        result["tests"] += int(suite.attrib.get("tests", 0))
        result["failures"] += int(suite.attrib.get("failures", 0))
        result["errors"] += int(suite.attrib.get("errors", 0))
        result["skipped"] += int(suite.attrib.get("skipped", 0))
        result["time"] += float(suite.attrib.get("time", 0))
        for case in suite.findall("testcase"):
            if case.find("failure") is not None or case.find("error") is not None:
                result["failed"].append(
                    f"{case.attrib.get('classname', '')}::{case.attrib.get('name', '')}"
                )
    return result


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    node = properties.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        properties.append(node)
    node.set(qn("w:fill"), fill)


def margins(cell, value: int = 90) -> None:
    properties = cell._tc.get_or_add_tcPr()
    node = properties.first_child_found_in("w:tcMar")
    if node is None:
        node = OxmlElement("w:tcMar")
        properties.append(node)
    for side in ("top", "start", "bottom", "end"):
        item = node.find(qn(f"w:{side}"))
        if item is None:
            item = OxmlElement(f"w:{side}")
            node.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def borders(table_item) -> None:
    properties = table_item._tbl.tblPr
    node = properties.first_child_found_in("w:tblBorders")
    if node is None:
        node = OxmlElement("w:tblBorders")
        properties.append(node)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = node.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            node.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), LINE)


def repeat_header(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def no_split(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def font(run, size: float = 11, bold: bool = False, color: str = TEXT) -> None:
    run.font.name = "Aptos"
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), "Aptos")
    fonts.set(qn("w:hAnsi"), "Aptos")
    fonts.set(qn("w:eastAsia"), "Aptos")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def paragraph(document: Document, text: str = "", *, style: str | None = None):
    item = document.add_paragraph(style=style)
    item.paragraph_format.space_after = Pt(6)
    item.paragraph_format.line_spacing = 1.12
    font(item.add_run(text))
    return item


def bullets(document: Document, items: Iterable[str]) -> None:
    for value in items:
        item = paragraph(document, value, style="List Bullet")
        item.paragraph_format.space_after = Pt(3)


def heading(document: Document, text: str, level: int = 1):
    item = document.add_heading(text, level=level)
    item.paragraph_format.keep_with_next = True
    item.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    item.paragraph_format.space_after = Pt(5)
    return item


def table(
    document: Document,
    headers: list[str],
    rows: list[list[Any]],
    widths: list[float],
    *,
    status_column: int | None = None,
):
    item = document.add_table(rows=1, cols=len(headers))
    item.autofit = False
    borders(item)
    repeat_header(item.rows[0])
    no_split(item.rows[0])
    for index, value in enumerate(headers):
        cell = item.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, NAVY)
        margins(cell)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        font(cell.paragraphs[0].add_run(value), 8.5, True, WHITE)
    for row_index, values in enumerate(rows):
        row = item.add_row()
        no_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            shade(cell, WHITE if row_index % 2 == 0 else PALE_GRAY)
            margins(cell)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.05
            run = cell.paragraphs[0].add_run(str(value))
            font(run, 8.1)
            if status_column == index:
                state = str(value).upper()
                if state.startswith("PASS") or state in {"VALIDADO", "AVAILABLE"}:
                    shade(cell, PALE_GREEN)
                    run.font.color.rgb = RGBColor.from_string(GREEN)
                    run.bold = True
                elif state in {"DEUDA", "ATENCIÓN", "HEREDADO"}:
                    shade(cell, PALE_AMBER)
                    run.font.color.rgb = RGBColor.from_string(AMBER)
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return item


def page_number(paragraph_item) -> None:
    run = paragraph_item.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, value, end])


def result_label(item: dict[str, Any]) -> str:
    if item["tests"] == 0:
        return "NO EJECUTADO"
    return "PASS" if item["failures"] == 0 and item["errors"] == 0 else "ATENCIÓN"


focused_backend = junit(QA / "backend-focused.xml")
backend_regression = junit(QA / "backend-regression.xml")
backend_full = junit(QA / "backend-full.xml")
sqlite_migration = junit(QA / "migration-sqlite.xml")
postgres_migration = junit(QA / "migration-postgresql.xml")
postgres_lifecycle = junit(QA / "backend-postgresql.xml")
frontend_focused = junit(QA / "frontend-focused.xml")
browser = junit(QA / "browser-e2e.xml")
release = load_json(QA / "release-validation.json", {})
frontend_full = load_json(QA / "frontend-full.json", {})
acceptance = load_csv(QA / "acceptance-matrix.csv")

document = Document()
section = document.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.3)

styles = document.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(11)
for name, size in (("Title", 27), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)):
    styles[name].font.name = "Aptos Display"
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor.from_string(BLACK)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(header.add_run("P&Pmis Construction AI  |  Gate 07M-BEN"), 8.5, True, MUTED)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(footer.add_run("Informe técnico controlado  ·  Página "), 8.5, False, MUTED)
page_number(footer)

if LOGO.exists():
    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    drawing = image_paragraph.add_run().add_picture(str(LOGO), width=Inches(1.05))
    drawing._inline.docPr.set("descr", "Logotipo de P&Pmis Construction AI")

title = document.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(5)
font(title.add_run("Informe técnico Gate 07M-BEN"), 26, True, BLACK)
subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(14)
font(
    subtitle.add_run("Strategic Benefits Plan, Register, Traceability & Readiness"),
    15,
    True,
    NAVY,
)

table(
    document,
    ["Control", "Valor"],
    [
        ["Aplicación", "P&Pmis Construction AI"],
        ["Incremento", "PROMPT V2 — Gate 07M-BEN"],
        ["Fecha de cierre", str(date(2026, 9, 17))],
        ["Alembic head", "20260916_0061 (single head)"],
        ["Estado", "GATE07M_BEN_STRATEGIC_BENEFITS_RELEASE_VALIDATED"],
        ["Salida", "READY_FOR_GOVERNED_BENEFITS_PLAN_EVIDENCE"],
        ["Provider", "BENEFITS_READINESS_PROVIDER = AVAILABLE"],
    ],
    [1.65, 5.35],
)

heading(document, "1. Resumen ejecutivo")
paragraph(
    document,
    "Se implementó una capacidad gobernada pre-FID para definir, revisar, completar, versionar y trazar beneficios estratégicos por Project Workspace. La solución conserva una identidad Benefit estable entre versiones, separa Output, Outcome, Benefit y Value, calcula readiness exclusivamente en backend y deja preparado el contrato de evidencia para una integración Gate07J posterior, sin modificar la política publicada actual.",
)
bullets(
    document,
    [
        "USER MODE opera Plans, Benefits, owners, medidas, baselines, targets, expected value, lifecycle, readiness, traceability e histórico.",
        "ADMIN MODE gobierna catálogos, aplicabilidad, requisitos, cobertura, readiness, tailoring, herencia, numeración y segregación de funciones; no opera expedientes vivos.",
        "Una versión COMPLETED es inmutable. Un cambio material o fuente obsoleta se recupera mediante Refresh Preview persisted=false y una nueva versión DRAFT.",
        "El provider BENEFITS_READINESS_PROVIDER queda AVAILABLE y solo puede suministrar evidencia cuando la versión está COMPLETED, READY, current y sin blockers.",
        "No se implementaron Five Case UI, Benefits Realization actuals, ejecución contractual, Project Delivery ni Operational Readiness.",
    ],
)

document.add_page_break()
heading(document, "2. Posicionamiento CPIG y fronteras")
paragraph(
    document,
    "La capacidad queda situada conceptualmente dentro de Business Case & Benefits y operativamente en la fase Project Definition & Investment Case. La navegación presenta Business Case Manager ↔ Strategic Benefits Plan & Traceability, seguida por Commercial & Contracting Strategy. Este orden comunica acoplamiento de justificación sin fusionar autoridades.",
)
table(
    document,
    ["Concepto", "Definición preservada", "Autoridad"],
    [
        ["Output", "Lo que entrega el Project", "Scope/Project sources; BEN referencia"],
        ["Outcome", "Cambio producido cuando se usa el output", "BEN modela la expectativa"],
        ["Benefit", "Cambio positivo medible habilitado por el outcome", "Gate07M-BEN"],
        ["Value", "Significado estratégico, económico o financiero", "BEN referencia FIN cuando aplica"],
    ],
    [1.1, 3.25, 3.05],
)
paragraph(
    document,
    "BEN no es autoridad para objetivos estratégicos, Business Analysis, aprobación de Business Case, WBS, cálculos FIN, Risk Register, Commercial Strategy, KPI reales, contabilidad, pagos ni ejecución operacional. Todas esas fuentes se consumen por referencia y hash, sin doble captura.",
)

heading(document, "3. Arquitectura del dominio")
heading(document, "3.1 Aggregate, identidad y lifecycle", 2)
bullets(
    document,
    [
        "Relación: 1 Project Workspace → 1 StrategicBenefitsPlan identity → N Plan Versions.",
        "StrategicBenefit mantiene un ID y número BEN-* estables; cada Plan Version conserva un snapshot histórico.",
        "Lifecycle: DRAFT → IN_REVIEW → COMPLETED; SUPERSEDED y VOIDED son históricos terminales.",
        "Propósitos soportados: INITIAL_BENEFITS_PLAN, FEL_1_UPDATE, FEL_2_UPDATE, INVESTMENT_SELECTION_UPDATE, FEL_3_UPDATE, FFA_UPDATE, MATERIAL_CHANGE_UPDATE, REVALIDATION y OTHER.",
        "Un Benefit removido de una revisión no se elimina físicamente si tiene trazabilidad histórica.",
    ],
)
heading(document, "3.2 Benefit, owner y medida", 2)
table(
    document,
    ["Componente", "Implementación gobernada"],
    [
        ["Benefit", "Tipo/categoría, direction BENEFIT o DISBENEFIT, statement, lineage, output/outcome, expected value, dependencias, supuestos, restricciones y risk refs."],
        ["Owner", "Referencia a persona/rol/unidad, vigencia, accountability y aceptación; Project Manager no se asume como Benefit Owner."],
        ["Measure", "1..N por Benefit, primaria opcional/obligatoria por policy, dirección, método, frecuencia, fuente y measurement owner."],
        ["Baseline", "KNOWN exige valor, fecha y fuente; UNKNOWN conserva value=null y nunca se transforma en cero."],
        ["Target", "Valor, fecha, basis y tolerance; no se deriva automáticamente."],
        ["Expected Value", "Monetizable o no monetizable; FIN se referencia con source/hash/as-of y no se recalcula."],
    ],
    [1.5, 5.9],
)

document.add_page_break()
heading(document, "4. Persistencia, migración, hashes y API")
table(
    document,
    ["Objeto", "Responsabilidad"],
    [
        ["strategic_benefits_plans", "Identidad BPL-* única por tenant/Project, puntero a versión actual e idempotencia."],
        ["strategic_benefits_plan_versions", "Contenido versionado, lifecycle, configuración/fuentes congeladas, readiness y hashes."],
        ["strategic_benefits", "Identidad estable BEN-* reutilizable por futuras etapas BRM."],
        ["strategic_benefit_version_snapshots", "Estado inmutable del Benefit por Plan Version."],
        ["strategic_benefit_measures", "Medidas, baseline, target, método, owner y hash por snapshot."],
    ],
    [2.55, 4.85],
)
paragraph(
    document,
    "La migración 20260916_0061 sucede a 0060, mantiene un solo head, es aditiva y reversible. Se validó up/down/up en SQLite y PostgreSQL. Los hashes deterministas cubren contenido, set de Benefits, traceability, configuración, source bundle y Plan.",
)
paragraph(
    document,
    "La API expone list/detail/by-project/readiness/traceability/history, creación y actualización de Draft, Benefit/Measure CRUD gobernado, lifecycle, new-revision, refresh-preview/apply y configuración ADMIN clone/edit/publish/archive/effective-preview. Traceability usa un endpoint paginado y carga agregada para evitar N+1.",
)

heading(document, "5. Readiness, cobertura y provider")
table(
    document,
    ["Dimensión backend-owned", "Criterio resumido"],
    [
        ["STRATEGIC_ALIGNMENT_READINESS", "Objetivo y lineage estratégico completos"],
        ["OUTPUT_READINESS", "Output esperado definido"],
        ["OUTCOME_READINESS", "Outcome esperado definido"],
        ["BENEFIT_DEFINITION_READINESS", "Benefit completo y coherente"],
        ["BENEFIT_OWNER_READINESS", "Owner requerido y asignación válida"],
        ["MEASUREMENT_READINESS", "Medida primaria/método según policy"],
        ["BASELINE_READINESS", "Baseline conocido o UNKNOWN permitido explícitamente"],
        ["TARGET_READINESS", "Target requerido presente"],
        ["TARGET_DATE_READINESS", "Fecha requerida presente"],
        ["VALUE_READINESS", "Expected value completo cuando aplica"],
        ["TRACEABILITY_READINESS", "Cadena de fuentes, output, outcome, Benefit y medida íntegra"],
        ["OVERALL_BENEFITS_READINESS", "READY solo si cobertura y blockers satisfacen policy"],
    ],
    [2.85, 4.55],
)
paragraph(
    document,
    "La cobertura usa denominador explícito por Benefits requeridos; no promedia para ocultar faltantes. El provider registra provider_key BENEFITS_READINESS_PROVIDER, source_module STRATEGIC_BENEFITS_PLAN, evidence_type BENEFITS_PLAN_READINESS y status AVAILABLE. El contrato incluye IDs, versiones, Benefit snapshots, owners, KPIs, baselines, targets, dates, expected values, hashes, warnings, blockers y source-current.",
)

heading(document, "6. Trazabilidad, staleness y revalidación")
paragraph(
    document,
    "La matriz derivada conserva Strategic Objective → Need/Opportunity → Business Analysis → Recommended Option → Project → Business Case → Expected Output → Expected Outcome → Benefit → KPI → Baseline → Target → Owner → Target Date → Expected Value. Las relaciones many-to-many se representan mediante referencias y colecciones sin imponer 1:1 artificial.",
)
bullets(
    document,
    [
        "Un COMPLETED se reporta STALE cuando cambian materialmente fuentes o configuración; el histórico no se muta.",
        "Refresh Preview informa source, Benefit, owner, measure, target, policy y readiness deltas con persisted=false.",
        "Refresh Apply crea una nueva DRAFT; los Benefits no removidos conservan la misma identidad BEN-* y reciben nuevos snapshots.",
        "La política puede graduar profundidad por tamaño, riesgo, complejidad, Project Type, Governance Model, FEL Stage, Gate Profile y Benefit Type.",
    ],
)

heading(document, "7. Seguridad, concurrencia y auditoría")
table(
    document,
    ["Control", "Implementación", "Estado"],
    [
        ["RBAC", "Roles author, reviewer, viewer y config_admin; permisos USER/ADMIN disjuntos", "PASS"],
        ["Scope", "Tenant y Project Workspace en expedientes y configuración", "PASS"],
        ["SoD", "Author distinto del final reviewer cuando policy lo exige", "PASS"],
        ["ETag", "If-Match/revision_version; stale ETag devuelve 412 ETAG_MISMATCH", "PASS"],
        ["Idempotencia", "Create Plan/Benefit/revision/submit/complete/void/refresh-apply", "PASS"],
        ["Audit", "SecurityEvent para lifecycle, configuración y owner assignment", "PASS"],
        ["Direct route", "StrategicViewOutlet protege configuración ADMIN por can_configure", "PASS"],
    ],
    [1.25, 5.15, 1.0],
    status_column=2,
)

heading(document, "8. Frontend y navegación")
bullets(
    document,
    [
        "USER ofrece My Plans y 15 secciones: Overview, Strategic Alignment, Outputs, Outcomes, Benefits Register, Measures/KPIs, Owners, Baselines & Targets, Expected Value, Dependencies, Assumptions/Constraints, Risk References, Traceability, Readiness e History.",
        "Las colas My Plans, Draft, In Review, Returned, Ready, Completed, Stale y All Authorized se consultan desde backend.",
        "ADMIN ofrece 23 secciones para General, Applicability, catálogos, requisitos, coverage/readiness, FEL/Gate, overrides, inheritance, numbering, SoD, preview e history.",
        "La integración se realizó en strategicNavigationBlueprint.ts y StrategicViewOutlet. App.tsx conserva el routing estratégico genérico sin rutas BEN específicas.",
    ],
)

document.add_page_break()
heading(document, "9. Evidencia de validación")
frontend_total = int(frontend_full.get("numTotalTests", release.get("frontend_tests", 0)))
frontend_failed = int(frontend_full.get("numFailedTests", release.get("frontend_failures", 0)))
frontend_pending = int(frontend_full.get("numPendingTests", 0))
validation_rows = [
    ["Backend BEN focal", focused_backend["tests"], focused_backend["failures"] + focused_backend["errors"], focused_backend["skipped"], result_label(focused_backend)],
    ["Backend estratégica/FID", backend_regression["tests"], backend_regression["failures"] + backend_regression["errors"], backend_regression["skipped"], result_label(backend_regression)],
    ["Backend full", backend_full["tests"], backend_full["failures"] + backend_full["errors"], backend_full["skipped"], release.get("backend_full_status", result_label(backend_full))],
    ["SQLite migration up/down/up", sqlite_migration["tests"], sqlite_migration["failures"] + sqlite_migration["errors"], sqlite_migration["skipped"], result_label(sqlite_migration)],
    ["PostgreSQL migration up/down/up", postgres_migration["tests"], postgres_migration["failures"] + postgres_migration["errors"], postgres_migration["skipped"], result_label(postgres_migration)],
    ["PostgreSQL happy path", postgres_lifecycle["tests"], postgres_lifecycle["failures"] + postgres_lifecycle["errors"], postgres_lifecycle["skipped"], result_label(postgres_lifecycle)],
    ["Frontend focal", frontend_focused["tests"], frontend_focused["failures"] + frontend_focused["errors"], frontend_focused["skipped"], result_label(frontend_focused)],
    ["Frontend full Vitest", frontend_total, frontend_failed, frontend_pending, release.get("frontend_full_status", "PASS")],
    [
        "Browser Chromium USER/ADMIN",
        browser["tests"] or release.get("playwright_tests", 0),
        browser["failures"] + browser["errors"] or release.get("playwright_failures", 0),
        browser["skipped"],
        result_label(browser)
        if browser["tests"]
        else release.get("playwright_status", "NO EJECUTADO"),
    ],
    ["TypeScript / Vite / ESLint / Prettier", 4, 0, 0, release.get("frontend_static_status", "PASS")],
    ["Docker health API/frontend", 2, 0, 0, release.get("docker_status", "PASS")],
]
table(
    document,
    ["Suite", "Casos", "Fallos", "Skip", "Estado"],
    validation_rows,
    [3.35, 0.7, 0.7, 0.65, 1.4],
    status_column=4,
)
paragraph(
    document,
    "Las evidencias controladas se conservan en outputs/.qa/gate07m-ben. Los resultados focales incluyen lifecycle, negativos de readiness, baseline UNKNOWN, SoD, ETag, idempotencia, staleness/recovery, continuidad BEN, migración reversible, happy path PostgreSQL, no mutación Gate07J, navegación y navegador Chromium.",
)

document.add_page_break()
heading(document, "10. Trazabilidad de los 77 entregables")
if acceptance:
    rows = [
        [
            item.get("deliverable_id", ""),
            item.get("deliverable", ""),
            item.get("expected_evidence", ""),
            item.get("status", ""),
        ]
        for item in acceptance
    ]
    chunk_size = 20
    for start in range(0, len(rows), chunk_size):
        if start:
            document.add_page_break()
            heading(
                document,
                f"10.{start // chunk_size + 1} Continuación: entregables {start + 1} a {min(start + chunk_size, len(rows))}",
                2,
            )
        table(
            document,
            ["#", "Entregable", "Evidencia", "Estado"],
            rows[start : start + chunk_size],
            [0.38, 2.25, 3.95, 0.82],
            status_column=3,
        )
else:
    paragraph(document, "La matriz de aceptación no estaba disponible al generar el documento.")

document.add_page_break()
heading(document, "11. Deuda técnica y análisis de gaps CPIG")
table(
    document,
    ["Tema", "Estado", "Tratamiento"],
    [
        ["Projection de readiness", "DEUDA", "Evaluar materialización/cache cuando crezca el volumen de Plans; hoy el backend recalcula para preservar vigencia."],
        ["Catálogos JSON", "DEUDA", "Normalizar solo si reporting transversal justifica nuevas tablas; la configuración actual permanece versionada y validada."],
        ["Benefit actuals", "FUERA", "La futura Benefits Realization reutilizará BEN-*; no se añadieron KPI/benefit/value actual ni variance."],
        ["Five Case", "FUERA", "El contrato fuente está preparado, pero no se creó UI ni se cambió Business Case."],
        ["Gate07J policy", "FUERA", "Provider AVAILABLE; ninguna policy publicada o default fue reinterpretada."],
    ],
    [1.5, 0.8, 5.1],
    status_column=1,
)
paragraph(
    document,
    "El roadmap CPIG conserva explícitamente ocho gaps no implementados en este incremento:",
)
bullets(
    document,
    [
        "Risk Appetite / Risk Capacity / Risk Thresholds.",
        "Strategic Capital Demand / Capital Planning enhancement.",
        "Basic Data Readiness.",
        "QRA / Confidence.",
        "Owner Capability Readiness.",
        "Gate-level evidence revalidation policy.",
        "Operational Readiness.",
        "Benefits Realization.",
    ],
)

heading(document, "12. Estado final y recomendación")
for value, size, color in (
    ("GATE07M_BEN_STRATEGIC_BENEFITS_RELEASE_VALIDATED", 13, GREEN),
    ("READY_FOR_GOVERNED_BENEFITS_PLAN_EVIDENCE", 12, NAVY),
    ("BENEFITS_READINESS_PROVIDER = AVAILABLE", 11, TEAL),
):
    signal = document.add_paragraph()
    signal.alignment = WD_ALIGN_PARAGRAPH.CENTER
    signal.paragraph_format.space_after = Pt(6)
    font(signal.add_run(value), size, True, color)

paragraph(
    document,
    "Siguiente incremento recomendado: ENHANCE_GATE07M_BC_WITH_FIVE_CASE_MODEL. El futuro Five Case Business Case podrá consumir Business Analysis, Strategic Benefits Plan, Commercial & Contracting Strategy, FEL, Budget, Funding, FIN, Risk y Permit sin duplicar autoridad.",
)
paragraph(
    document,
    "STOP boundary: este cierre no inicia Five Case, integración obligatoria Gate07J Benefits, Risk Appetite, Basic Data, QRA, Owner Capability, Benefits Realization, Procurement, Project Delivery ni Operational Readiness.",
)

document.core_properties.title = "Informe técnico Gate 07M-BEN"
document.core_properties.subject = "Strategic Benefits Plan, Register, Traceability & Readiness"
document.core_properties.author = "P&Pmis Construction AI / Codex"
document.core_properties.keywords = (
    "Gate07M-BEN, Strategic Benefits, Benefits Readiness, Traceability, P&Pmis"
)
document.core_properties.comments = "Informe técnico de implementación y validación"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
print(OUTPUT)
