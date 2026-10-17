"""Evidence-led current UAT addenda and the 27-section technical closeout.

Run the document authoring marker before --documents. Prior release artifacts
remain byte-for-byte unchanged; their results are retained as history.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree
from zoneinfo import ZoneInfo

import generate_strategic_final_uat_reports as style
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs/strategic-final-uat-20261014"
OUT = ROOT / "outputs/strategic-uat-remediation-20261015"
DATA = ROOT / "docs/validation/strategic-uat-remediation"
QA = ROOT / "outputs/.qa/strategic-uat-remediation"
NAMES = style.FILES
NEXT = "REMEDIATE THE SMALLEST REMAINING STRATEGIC UAT BLOCKER"


def load(path, default=None):
    return (
        json.loads(path.read_text(encoding="utf-8-sig")) if path.is_file() else default
    )


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def junit(name):
    path = QA / name
    if not path.is_file():
        return {"file": name, "status": "NOT_COMPLETED", "counts": None}
    tree = ElementTree.parse(path)
    cases = list(tree.getroot().iter("testcase"))
    counts = Counter()
    failed = []
    for case in cases:
        state = (
            "ERROR"
            if case.find("error") is not None
            else "FAIL"
            if case.find("failure") is not None
            else "SKIP"
            if case.find("skipped") is not None
            else "PASS"
        )
        counts[state] += 1
        if state in {"FAIL", "ERROR"}:
            failure = case.find("failure") if state == "FAIL" else case.find("error")
            failed.append(
                {
                    "test": f"{case.get('classname')}::{case.get('name')}",
                    "message": failure.get("message", ""),
                }
            )
    return {
        "file": name,
        "status": "COMPLETED",
        "sha256": digest(path),
        "total": len(cases),
        "unique_total": len(
            {(case.get("classname"), case.get("name")) for case in cases}
        ),
        "counts": {key: counts[key] for key in ("PASS", "FAIL", "ERROR", "SKIP")},
        "failures": failed,
    }


def chain_case_notes(chain, setup):
    """Describe recorded stages without promoting composite formal UAT cases."""
    stages = {item["stage"]: item for item in chain.get("stages", [])}
    passed = {name for name, item in stages.items() if item.get("status") == "PASS"}
    blocker = chain.get("blocker") or setup.get("error") or {}
    stop = str(blocker.get("stage") or "sin punto de cierre registrado")
    notes = {}
    if setup.get("status") == "UPSTREAM_COMPLETE":
        notes.update(
            {
                "A0": "Configuraciones y asignaciones reales registradas; el inventario efectivo y todos los negativos ADMIN/USER del caso formal no se acreditan solo por publicar sucesores",
                "A1": "Jerarquía nueva y Project canónico materializados con Record Codes. La visita de Explorer se registra aparte y no sustituye todos los controles de identidad y ámbito",
                "A2": "Ciclo, Demand y Allocation recorridos y utilizados para originar la Idea de esta misma cadena; la auditoría compuesta completa sigue pendiente",
                "A3": "Idea aceptada y BA completado por revisor independiente mediante comandos HTTP reales; una pantalla abierta no equivale a ejecutar allí todo el lifecycle",
                "A4": "Proposal evaluada y handoff al Strategic Gate ejecutado en esta cadena; no se trasladan hallazgos o aprobaciones de otro tenant a este caso formal",
                "A5": "Strategic Gate APPROVE por autoridad independiente ejecutado; no equivale a FID, financiación ni activación",
                "A6": "Solicitud revisada, aprobada y materializada como Project pending. Replay, elegibilidad e inicialización del caso compuesto requieren su propia evidencia",
                "A7": "Intake y membership principal ACTIVE de la misma identidad. No se infieren todas las variaciones de membresía ni cambios de padre",
            }
        )
    if "A8_A12_PORTFOLIO_SOURCES" in passed:
        notes.update(
            {
                "A8": "Etapa técnica completada: evaluación inmutable, historia real y rank contextual de la misma inversión. El caso formal compuesto continúa BLOCKED",
                "A9": "Etapa técnica completada: análisis, moneda y cobertura comprobados, Investor Map, Scorecard y distribuciones consultados. No implica validación visual integral",
                "A10": "Budget Target publicado y escenario Budget FROZEN mediante comandos normales; no constituye inversión aprobada",
                "A11": "Capacity y Demand publicados y escenario Resource FROZEN en esta cadena; la identidad y fechas fuente del Project se conservaron",
                "A12": "Budget y Resource actuales convergieron en alternativa RECOMMENDED y READY_FOR_INVESTMENT_GATE_PREPARATION; todavía no es FID",
            }
        )
    elif setup.get("status") == "UPSTREAM_COMPLETE":
        for case_id in ("A8", "A9", "A10", "A11", "A12"):
            notes[case_id] = (
                f"El grupo técnico A8 a A12 no tiene checkpoint completo. Detención registrada: {stop}; no se acredita el grupo por registros parciales"
            )
    if "A13_PERIODIC_PORTFOLIO_REVIEW" in passed:
        notes["A13"] = (
            "Review real COMPLETED con snapshot histórico congelado, conclusión y acción por actores independientes; Project y membership preservados. No es decisión de inversión"
        )
    if "FEL1_PRELIMINARY_BUSINESS_CASE" in passed:
        notes["A14"] = (
            "Business Case preliminar completado en FEL 1; no sustituye la revisión Five Case FFA final ni su auditoría compuesta"
        )
    if "FEL1_TO_FEL3" in passed:
        evidence = dict(stages["FEL1_TO_FEL3"].get("data", {}).get("evidence", {}))
        notes["A17"] = (
            "Assessments FEL 1 a FEL 3 registrados; evidencia de definición "
            f"{evidence.get('FEL_stage_code', 'sin etapa')} / {evidence.get('readiness_status', 'sin readiness')}. "
            "Assessment READY no equivale por sí solo al cierre de revisión independiente de la etapa FEL exigida por Business Case"
        )
    if "RISK_PERMIT_BASIC_DATA_OWNER" in passed:
        notes.update(
            {
                "A18": "Basic Data completó su lifecycle con validación y revisión independientes; el caso formal completo no se infiere solo de su provider",
                "A20": "Risk Assessment completado en la misma cadena por actores reales. No sustituye la auditoría completa de Risk Governance ni una cuantificación QRA",
                "A21": "Permit completó su lifecycle y assessment mediante comandos normales y política efectiva; no se escribió un provider terminal prefabricado",
                "A22": "Owner Capability completó preparación, revisión y completion reales; el caso formal y sus negativos permanecen sin acreditación integral",
            }
        )
    if "FUNDING_AND_FINANCIAL_MODEL" in passed:
        notes.update(
            {
                "A23": "Funds y asignaciones recorrieron autorización y confirmación con actores independientes; no equivalen a FID ni a ejecución del Project",
                "A24": "Modelo económico financiero y versión completados con fuentes reales. Sus resultados no sustituyen Five Case final, Gate ni decisión de inversión",
            }
        )
    if "COMMERCIAL_BENEFITS_FINAL_FIVE_CASE" in passed:
        notes.update(
            {
                "A14": "Revisión Five Case FFA final completada por revisor independiente y providers actuales comprobados; falta acreditación completa del caso formal y del E2E browser",
                "A15": "Benefits completado y provider source_current comprobado tras el enlace inverso final del Business Case; no se acredita todo el caso compuesto automáticamente",
                "A16": "Commercial completado y provider source_current comprobado tras el Business Case final; no se infiere aprobación de inversión",
                "A17": "Las etapas FEL 1, FEL 2 y FEL 3 recorrieron assessment y revisión independiente antes de las fuentes finales. El Business Case final completado consume la definición vigente; no equivale a aprobación FID",
            }
        )
    elif stop == "COMMERCIAL_BENEFITS_FINAL_FIVE_CASE":
        for case_id in ("A14", "A15", "A16"):
            notes[case_id] = (
                f"El grupo final Commercial, Benefits y Five Case no cerró su checkpoint: {stop}. El bloqueo exacto del corte se conserva sin sustituir fuentes ni etapas"
            )
    if "EFFECTIVE_POLICY_PREVIEW" in passed:
        for case_id in ("A25", "A26"):
            notes[case_id] = (
                "Preview de política efectiva capturado. La exigencia contextual QRA debe leerse de ese resultado; no se infiere una corrida ni se fabrica evidencia cuando no corresponde"
            )
    if "A27_A29_GATE_PREPARATION_FREEZE" in passed:
        notes.update(
            {
                "A27": "Preview efectivo, evidencia requerida y readiness del paquete recorridos por comandos reales; no acreditan el caso formal completo por sí solos",
                "A28": "Fuentes Five Case y paquete de esta identidad enlazados en la etapa técnica; los negativos y la inmutabilidad posterior al FID requieren evidencia separada",
                "A29": "Paquete Gate de la cadena seleccionado quedó FROZEN mediante el lifecycle normal; no equivale a decisión GO ni FID",
            }
        )
    elif (
        stop == "A27_A29_GATE_PREPARATION_FREEZE"
        and "EFFECTIVE_POLICY_PREVIEW" in passed
    ):
        notes.update(
            {
                "A27": "Gate creado como DRAFT con fuentes actuales. La política exige cuatro referencias Class B todavía sin presentación y validación independiente; readiness BLOCKED. No es defecto por provider STALE",
                "A28": "Business Case, FIN, Benefits y Commercial actuales antes del Gate; falta la cobertura de evidencias controladas y la auditoría completa de compatibilidad y negativos",
                "A29": "Freeze no ejecutado: el paquete sigue DRAFT/BLOCKED bajo su política sin rebajar requisitos. No existe handoff frozen nuevo",
            }
        )
    if "A30_FID" in passed:
        notes["A30"] = (
            "Decisión FID registrada por autoridad independiente en la etapa técnica de esta cadena; no convierte automáticamente todos los casos compuestos ni el browser en PASS"
        )
    if "A31_POST_FID" in passed:
        notes["A31"] = (
            "Post FID evaluado sobre la decisión nueva y su Project. Elegibilidad no equivale a inicialización ni activación; el caso formal conserva sus límites de evidencia"
        )
    return notes


def blocker_text(chain, setup, gate_diagnosis=None):
    blocker = chain.get("blocker") or setup.get("error")
    if not blocker:
        if not chain.get("finished_at"):
            return "La ejecución seleccionada no tiene resultado terminal registrado en este corte; no se convierte una ejecución en curso en PASS ni en liberación."
        return "Sin bloqueo serializado en el manifest seleccionado; ello no sustituye la auditoría formal, el browser completo ni las firmas humanas."
    if not isinstance(blocker, dict):
        return (
            "Bloqueo registrado: "
            + str(blocker)
            + ". No se sustituye por evidencia de otro intento."
        )
    diagnosis = gate_diagnosis or {}
    context = diagnosis.get("context", {})
    if (
        blocker.get("stage") == "A27_A29_GATE_PREPARATION_FREEZE"
        and (context.get("tenant_id"), context.get("project_id"))
        == (setup.get("tenant_id"), setup.get("project_id"))
        and diagnosis.get("literal_gate_blockers")
    ):
        return (
            "Punto de detención A27 a A29. Gate "
            + str(diagnosis["preparation"]["id"])
            + " permanece DRAFT / BLOCKED con fuentes actuales. Bloqueadores literales: "
            + ", ".join(diagnosis["literal_gate_blockers"])
            + ". Son cuatro referencias controladas Class B sin presentación y validación independiente; "
            "su detalle se encuentra en la sección 12. La respuesta íntegra, sin truncar, se conserva en "
            "fresh-capital-owner-chain.json y gate-evidence-blocker.json. No se ejecutaron freeze, FID ni Post FID."
        )
    return (
        f"Punto de detención {blocker.get('stage', 'preparación')}. Respuesta "
        + json.dumps(blocker.get("detail", blocker), ensure_ascii=False, default=str)[
            :1800
        ]
        + ". Se conserva este resultado de la cadena seleccionada; no se reemplaza con otro intento ni con una regla debilitada."
    )


def browser_text(browser, validation, setup):
    visits = browser.get("visits", [])
    identity = browser.get("identity") or validation.get("identity") or {}
    matching = (identity.get("tenant_id"), identity.get("project_id")) == (
        setup.get("tenant_id"),
        setup.get("project_id"),
    )
    details = (
        f"Chromium registró {len(visits)} vistas en tenant {identity.get('tenant_id', 'no registrado')} "
        f"y Project {identity.get('project_id', 'no registrado')}. "
        + (
            "Corresponde a la cadena seleccionada. "
            if matching
            else "Es una ejecución distinta de la cadena seleccionada y no se traslada su cobertura. "
        )
    )
    if visits:
        details += (
            "Vistas: "
            + ", ".join(str(item.get("view", "sin nombre")) for item in visits)
            + ". "
        )
    details += (
        "La navegación y selección de registros no equivale a creación browser de todos los lifecycles. "
        f"Estado browser {browser.get('status', 'NOT_COMPLETED')}; bloqueo "
        + str(browser.get("blocker", {}).get("stage", "no serializado"))
        + f". Colecciones registradas: {len(browser.get('http_errors', []))} errores HTTP, "
        + f"{len(browser.get('runtime_errors', []))} errores runtime y {len(browser.get('browser_business_writes', []))} escrituras browser de negocio. "
        + "Los rechazos de autorización y límites de hidratación observados se conservan; no se afirma una red limpia ni E2E global aprobado."
    )
    return details


def collect():
    old = load(ROOT / "docs/validation/strategic-final-uat/release-results.json")
    register = load(
        ROOT / "docs/validation/strategic-final-uat/scenario-coverage-register.json"
    )
    chain = load(QA / "fresh-capital-owner-chain.json", {})
    setup = chain.get("setup") or load(QA / "fresh-setup.json", {})
    # Every current case starts BLOCKED. Evidence of one portion cannot turn a
    # composite formal UAT case into PASS; a complete audit must explicitly say so.
    cases = [
        {
            "id": item["uat_id"],
            "process": item["business_process"],
            "status": "BLOCKED",
            "actual": "BLOCKED NOT EXECUTED como caso UAT completo en esta revisión",
            "evidence": "Sin acreditación completa de todos los controles del caso",
        }
        for item in register["cases"]
        if item["uat_id"].startswith("A")
    ]
    assert len(cases) == 32 and {item["id"] for item in cases} == {
        f"A{number}" for number in range(32)
    }
    # Supplemental audits retain their own identity and historical status; they
    # never overwrite the selected formal matrix or the retained 51-case record.
    explicit = load(DATA / "audited-current-cases.json", {})
    partial = chain_case_notes(chain, setup)
    for item in cases:
        if item["id"] in partial:
            item["actual"] = partial[item["id"]]
            item["evidence"] = (
                f"fresh-capital-owner-chain.json; tenant {setup.get('tenant_id')}; Project {setup.get('project_id')}; estado por etapa, no suma de fixtures"
            )
    suites = {
        "core": junit("core-parent-focused-current.xml"),
        "sources": junit("source-links/domain-regression.xml"),
        "isolation_probe": junit("isolation-probe.xml"),
        "isolation_safety": junit("isolation-safety-linux.xml"),
        "global": junit("backend-strategic-regression-global-resume.xml"),
        "browser": junit("browser/strategic-capital-current.xml"),
        "currency": junit("currency-projection-focused.xml"),
        "shell": junit("strategic-shell-focused.xml"),
        "core_scope_fixture_attempt": junit("core-parent-scope-fixture-104.xml"),
    }
    defects = deepcopy(old.get("defects", []))
    for defect in defects:
        defect["prior_status"] = defect["status"]
        if defect["id"] == "UAT-S3-003" and suites["core"].get("counts") == {
            "PASS": 72,
            "FAIL": 0,
            "ERROR": 0,
            "SKIP": 0,
        }:
            defect.update(
                status="CLOSED_SCOPED_TECHNICAL",
                actual="Ruta explícita de reconciliación de padres existentes, ETag/RBAC y protección concurrente acreditadas con 72 pruebas enfocadas; no equivale a regresión PostgreSQL completa",
                evidence=suites["core"]["file"],
            )
        elif defect["id"] == "UAT-S3-002":
            failures = suites["global"].get("failures", [])
            isolated = (
                suites["global"].get("counts")
                == {"PASS": 898, "FAIL": 1, "ERROR": 0, "SKIP": 15}
                and suites["global"].get("unique_total") == 914
                and len(failures) == 1
                and failures[0]["test"]
                == "tests.test_gate07m_fid_e2e::test_new_pyp_ffa_fid_path_full_stack_revalidation"
                and "BUSINESS_CASE_NOT_READY" in failures[0]["message"]
            )
            defect.update(
                status="CLOSED_ISOLATION_SCOPED"
                if isolated
                else "OPEN_GLOBAL_REVALIDATION",
                actual=(
                    "Aislamiento cerrado en alcance acotado: misma invocación original de 61 módulos y 914 casos únicos, 898 PASS, 1 FAIL, 0 ERROR, 15 SKIP. Cinco antiguos FAIL y tres ERROR por estado compartido ahora PASS; el FID legado sigue fallando con BUSINESS_CASE_NOT_READY. No existe global green"
                    if isolated
                    else "Aislamiento por módulo implementado y protegido; el resultado global disponible no permite aplicar automáticamente el cierre acotado auditado"
                ),
                evidence=suites["global"]["file"],
            )
        elif defect["id"] == "UAT-S2-003":
            defect.update(
                status="OPEN",
                actual="La cadena nueva está registrada por etapas; no se acredita A0 a A31 completo hasta cumplir fuentes actuales, FID, Post FID, regresión y Chromium sobre la misma identidad",
                evidence="fresh-capital-owner-chain.json",
            )
    browser = load(QA / "fresh-browser-result.json", {})
    browser_validation = load(DATA / "browser-current-validation.json", {})
    shell_closed = (
        any(
            item.get("id") == "UAT-S2-004" and item.get("status") == "CLOSED_SCOPED_UI"
            for item in browser_validation.get("closed_defects", [])
        )
        and len(browser.get("visits", [])) >= 6
        and suites["shell"].get("counts")
        == {"PASS": 29, "FAIL": 0, "ERROR": 0, "SKIP": 0}
    )
    if (
        shell_closed
        or browser_validation.get("defect_id") == "UAT-S2-004"
        or browser.get("visits")
    ):
        defects.append(
            {
                "id": "UAT-S2-004",
                "severity": "S2 CRITICAL",
                "classification": "UX/ARCHITECTURE",
                "title": "Las pantallas estratégicas nuevas dependen del dashboard Project heredado",
                "actual": (
                    f"El usuario autorizó expresamente el ajuste mínimo de App.tsx. El fallback USER reutiliza StrategicWorkspaceShell, WorkspaceNavigation y StrategicViewOutlet sin Project operacional ficticio y con canConfigure=false. Shell 29 PASS; Chromium abrió {len(browser.get('visits', []))} vistas de registros reales. No equivale a E2E ni valida freeze/FID/Post-FID"
                    if shell_closed
                    else "La evidencia actual del shell y su navegación conserva un límite pendiente; no se declara su cierre sin pruebas enfocadas y browser real"
                ),
                "status": "CLOSED_SCOPED_UI" if shell_closed else "OPEN",
                "prior_status": "NOT_PREVIOUSLY_RECORDED",
                "owner": "Frontend estratégico",
                "workaround": "No se creó Project operacional heredado ficticio; no se habilita configuración ADMIN desde el fallback USER",
                "evidence": "strategic-shell-focused.xml; browser/strategic-capital-current.xml; browser-current-validation.json",
            }
        )
    if suites["currency"].get("counts") == {
        "PASS": 4,
        "FAIL": 0,
        "ERROR": 0,
        "SKIP": 0,
    }:
        defects.append(
            {
                "id": "UAT-S3-004",
                "severity": "S3 MAJOR",
                "classification": "CODE",
                "title": "La entrada Portfolio omitía la moneda capturada en Proposal",
                "actual": "Snapshots nuevos conservan currency y currency_code de Proposal. Cuatro pruebas cubren USD, ausencia de moneda y ausencia de backfill; el avance de la cadena seleccionada se registra por separado sin extrapolar un Budget anterior",
                "status": "CLOSED_NEW_ENTRIES_SCOPED",
                "prior_status": "NEW_FINDING_FIXED",
                "owner": "Portfolio Planning",
                "workaround": "Los snapshots históricos no se reescriben; la carencia anterior permanece explícita",
                "evidence": "currency-projection-focused.xml; fresh-capital-owner-chain-project4-currency-blocker.json; fresh-capital-owner-chain.json",
            }
        )
    observation_ids = {
        item.get("id")
        for item in browser_validation.get("presentation_observations", [])
    }
    if "OBS-BC-SOURCE-MATRIX-MAPPING" in observation_ids:
        defects.append(
            {
                "id": "UAT-S3-005",
                "severity": "S3 MAJOR",
                "classification": "UX/DATA_PRESENTATION",
                "title": "La tabla de fuentes Five Case no refleja algunos valores serializados",
                "status": "OPEN_SCOPED_PRESENTATION",
                "prior_status": "NEW_BROWSER_FINDING",
                "actual": "Chromium muestra NOT_AVAILABLE y Funding STALE en la tabla informativa, aunque BC4/revisión n.º2 (ID6) es READY/current y los providers gobernados del Gate están vigentes. Es un hallazgo de presentación; no se afirma Funding realmente STALE",
                "owner": "Frontend Business Case",
                "workaround": "Contrastar la revisión capturada y readiness/provider API; no interpretar esa etiqueta de tabla como estado live de Funding",
                "evidence": "browser-current-validation.json; screenshot 06-fresh-completed-five-case-current-sources.png; normalizedSources y _five_case_contract",
            }
        )
    if "OBS-PROPOSAL-AFTER-APPROVE-READINESS" in observation_ids:
        defects.append(
            {
                "id": "UAT-S4-006",
                "severity": "S4 MINOR",
                "classification": "UX/LIFECYCLE_LABEL",
                "title": "La etiqueta de elegibilidad posterior a APPROVE requiere contexto de lifecycle",
                "status": "OPEN_SCOPED_LIFECYCLE_LABEL",
                "prior_status": "NEW_BROWSER_FINDING",
                "actual": "Proposal4 conserva STRATEGIC_GATE_APPROVED e historia DECIDED APPROVE, pero su elegibilidad de entrada posterior muestra GATE07B_REWORK_REQUIRED / PROPOSAL_NOT_EVALUATED. No se borra la aprobación histórica ni se infiere otro bloqueo FID",
                "owner": "Project Proposal y frontend estratégico",
                "workaround": "Consultar el snapshot y la decisión aprobada; distinguir elegibilidad para una nueva entrada de la historia de APPROVE",
                "evidence": "browser-current-validation.json; screenshot 03-fresh-proposal.png; _gate_readiness y proposalReadiness",
            }
        )
    gate_diagnosis = load(DATA / "gate-evidence-blocker.json", {})
    facts = {
        "recorded_at": datetime.now(ZoneInfo("America/Bogota")).isoformat(),
        "baseline_commit": "10a2442bfb69b97dc72f9c78e2419f157b400973",
        "head": "20261011_0068",
        "release_status": "STRATEGIC_MACROPROCESS_NOT_VALIDATED",
        "release_decision": "NO_GO",
        "next": NEXT,
        "human_signoff": "PENDING_NOT_ASSUMED",
        "chain": chain,
        "setup": setup,
        "cases": cases,
        "case_counts": dict(Counter(item["status"] for item in cases)),
        "other_formal_cases": "19 casos anteriores no reejecutados formalmente; historial preservado y estado actual BLOCKED",
        "suites": suites,
        "browser_evidence": browser,
        "browser_validation": browser_validation,
        "current_blocker_summary": blocker_text(chain, setup, gate_diagnosis),
        "current_browser_summary": browser_text(browser, browser_validation, setup),
        "gate_blocker_diagnosis": gate_diagnosis,
        "supplemental_case_audit": explicit,
        "formal_history": {
            "file": "docs/validation/strategic-final-uat/scenario-coverage-register.json",
            "case_count": len(register["cases"]),
            "preserved_separately": True,
        },
        "app_change": {
            "authorization": "EXPLICIT_USER_AUTHORIZED_MINIMAL_APP_CHANGE",
            "scope": "USER fallback for canonical Workspaces without legacy operational Project; existing navigation/outlet reused; canConfigure=false",
            "actual_sha256": digest(ROOT / "frontend/src/App.tsx"),
            "unchanged_claim": False,
            "scoped_ui_defect_closed": shell_closed,
        },
        "core_scope_rerun_boundary": (
            "Se conserva por separado el intento de 104 casos: 103 PASS y un FAIL del positivo HTTP organizacional. "
            "El guard devolvió PROJECT_PARENT_RECONCILIATION_REQUIRED con los dos IDs reales requeridos [5,6], "
            "porque la fixture solo confirmó uno. Se corrigió la confirmación de ambos padres, no el control de ámbito "
            "ni la regla de reconciliación. La ejecución enfocada final tiene su propio current.xml; no se suman ni se oculta el intento fallido."
            if suites["core_scope_fixture_attempt"].get("counts")
            == {"PASS": 103, "FAIL": 1, "ERROR": 0, "SKIP": 0}
            else "El intento diagnóstico CORE, si existe, se conserva con sus propios conteos y fallos; no se suma ni sustituye la ejecución enfocada final."
        ),
        "defects": defects,
        "operational_before": load(DATA / "operational-before.json", {}),
        "operational_after": load(DATA / "operational-after.json", {}),
        "prior_documents": {
            name: {"path": str(SOURCE / filename), "sha256": digest(SOURCE / filename)}
            for name, filename in NAMES.items()
        },
        "protected_app_sha256": digest(ROOT / "frontend/src/App.tsx"),
    }
    DATA.mkdir(parents=True, exist_ok=True)
    (DATA / "release-results.json").write_text(
        json.dumps(facts, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return facts


def para(doc, text):
    style.p(doc, str(text))


def table(doc, headings, rows, widths=None):
    style.table(doc, headings, rows, widths or [7.06 / len(headings)] * len(headings))
    for cell in doc.tables[-1].rows[0].cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True


def matrix(doc, facts):
    table(
        doc,
        ["Caso", "Proceso", "Estado y alcance actual"],
        [
            [case["id"], case["process"], case["status"] + "\n" + case["actual"]]
            for case in facts["cases"]
        ],
        [0.5, 2.3, 4.26],
    )


def new_closeout(facts):
    # Reuse the actual retained document and its page/style/relationship system.
    doc = Document(SOURCE / NAMES["closeout"])
    body = doc._element.body
    logo = deepcopy(body[0]) if body[0].xpath(".//w:drawing") else None
    for element in list(body):
        if element.tag != qn("w:sectPr"):
            body.remove(element)
    if logo is not None:
        body.insert(0, logo)
    doc.add_paragraph("Informe técnico de remediación de bloqueos UAT", "Title")
    para(
        doc,
        "P&Pmis Construction AI. Capital Portfolio & Investment Governance. Revisión del 15 de octubre de 2026.",
    )
    para(
        doc,
        "El incremento incorpora correcciones acotadas de reconciliación CORE, aislamiento de pruebas y referencias cerradas Five Case. La liberación permanece NO GO mientras no exista evidencia suficiente de una sola cadena nueva A0 a A31, su decisión FID y la validación independiente exigida. El informe separa pruebas enfocadas, ejecución parcial y aceptación humana.",
    )
    current = facts["chain"]
    setup = facts["setup"]
    stages = current.get("stages", [])
    dispositions = {item["id"]: item["status"] for item in facts["defects"]}
    sections = [
        (
            "1 Resumen ejecutivo",
            "Se mantienen las reglas de negocio, la segregación de funciones, los controles de versión y la identidad Project canónica. Las pruebas aprobadas que se detallan aquí no constituyen aceptación empresarial ni una suma equivalente a la cadena completa.",
        ),
        (
            "2 Línea base",
            f"La rama de trabajo es main. Commit inicial {facts['baseline_commit']}. Alembic conserva una sola cabeza {facts['head']}. No se añade migración. La liberación anterior era NO GO y NOT VALIDATED; sus documentos permanecen intactos en el directorio del 14 de octubre.",
        ),
        (
            "3 Bloqueos vigentes",
            (
                "UAT S2 003 sigue abierto hasta acreditar la cadena completa. "
                f"La disposición de aislamiento UAT S3 002 es {dispositions.get('UAT-S3-002', 'NOT_RECORDED')}; la repetición global mantiene su fallo legado y no es global verde. "
                f"La disposición CORE UAT S3 003 es {dispositions.get('UAT-S3-003', 'NOT_RECORDED')}, según la evidencia enfocada registrada. "
                f"El bloqueo de shell UAT S2 004 figura {dispositions.get('UAT-S2-004', 'NOT_RECORDED')}; una apertura de vistas no equivale al recorrido browser completo. Los hallazgos menores anteriores se conservan."
            ),
        ),
        (
            "4 Cambios realizados",
            "CORE añade diagnóstico de padres faltantes y un comando explícito que acepta únicamente IDs existentes del tenant. La reconciliación crea o actualiza un DRAFT, preserva las demás ediciones e invalida aprobación y validación si cambia su contenido. Antes de publicar se revalida la identidad exacta del padre reconciliado. SQLite se aísla por módulo sin reemplazar el objeto engine y sin reiniciar bases PostgreSQL explícitas. Benefits y Commercial aceptan únicamente la referencia inversa exacta del Business Case completado que consume la misma versión y hash; cualquier otra variación sigue produciendo STALE.",
        ),
        (
            "5 Entorno PostgreSQL de UAT",
            "docker-compose.strategic-fresh-uat.yml usa PostgreSQL 16 con almacenamiento temporal, Redis privado y API en 127.0.0.1 puerto 8840. No comparte volúmenes ni conexión con el localhost operacional. El esquema utiliza la compatibilidad histórica de metadata hasta 0067 y aplica realmente la migración 0068; esto no demuestra un upgrade íntegro desde un Alembic vacío. AUTO CREATE SCHEMA y SEED DEMO DATA permanecen desactivados para la API aislada.",
        ),
        (
            "6 Conjunto de datos nuevo",
            "El escenario registra una nueva jerarquía y nuevos registros estratégicos. Solo identidades, credenciales de prueba, asignaciones RBAC y un objetivo administrativo explícito se inicializan como referencias. Idea, BA, Proposal, decisiones, solicitudes, Project, membership y fuentes finales se producen mediante comandos normales. Los intentos previos se conservan y no se suman a la cadena seleccionada.",
        ),
        (
            "7 Actores y segregación de funciones",
            "ADMIN autor y publicador son personas distintas. Preparador y revisor de negocio son independientes. También se separan la autoridad FID, la aprobación de creación, materialización, inicialización y activación. Los permisos se resuelven desde las asignaciones reales con bootstrap desactivado en los comandos de dominio. La existencia de una cuenta no sustituye la autorización para el comando.",
        ),
        (
            "8 Inventario de configuración",
            "Los IDs, revisiones y hashes capturados se consultan en el registro de configuración. Los defaults que nacen publicados por bootstrap del producto se identifican como tales; los sucesores UAT siguen clone o DRAFT, validación y publicación por un actor independiente. No se edita contenido de políticas publicadas históricas. El inventario incompleto no se presenta como A0 aprobado.",
        ),
        (
            "9 Matriz de ejecución A0 a A31",
            "La matriz distingue el caso UAT compuesto de una etapa técnica parcial. Un caso sin evidencia suficiente de todos sus controles permanece BLOCKED. Los resultados anteriores se conservan en la trazabilidad y no se reemplazan silenciosamente.",
        ),
        (
            "10 Preparación de fuentes Five Case",
            "La secuencia parte de un Business Case preliminar completado mientras el Project está realmente en FEL 1. La progresión FEL requiere assessment y revisión independiente de cada etapa, incluida FEL 3, antes de capturar las fuentes finales. Assessment READY por sí solo no sustituye complete stage. FIN consume el Business Case preliminar; Commercial y Benefits deben completar sus lifecycles con entradas estabilizadas. La revisión FFA debe consumir IDs, versiones y hashes reales de las fuentes exigidas. Una revisión DRAFT abierta no se elimina ni sustituye para ocultar un snapshot bloqueado; un nuevo escenario mantiene su propia historia.",
        ),
        (
            "11 Evidencia de readiness del Business Case",
            "Solo la ejecución satisfactoria de readiness, revisión independiente y complete puede acreditar el Business Case. No se escribe READY, COMPLETED ni el hash esperado directamente en la base. El registro de etapas contiene la revisión preliminar y, si se alcanzó, la revisión Five Case final; una no equivale a la otra.",
        ),
        (
            "12 Evidencia de Gate07J",
            "Se resuelve la política efectiva PYP FFA y se registra su preview antes de crear el paquete. Una evidencia Class A faltante u obsoleta bloquea el paquete. No se sustituye por evidencia manual ni se fuerza READY. El estado del paquete y su freeze solo se acreditan si el comando correspondiente quedó ejecutado y registrado.",
        ),
        (
            "13 Evidencia de Gate07K y FID",
            "El FID requiere paquete frozen, preparación, revisión, checklist y decisión GO por la autoridad independiente. Solo el resultado de esos comandos permite acreditar is_fid, full_funds_authorized e investment_authorized. Si A30 no se ejecutó, no existe prueba nueva de FID en este incremento.",
        ),
        (
            "14 Evidencia de Post FID",
            "La evaluación Post FID debe referenciar exactamente la decisión nueva y conservar el Project pending. Elegibilidad para inicialización no es activación ni ejecución. Si el comando no se alcanzó, su resultado permanece bloqueado y no se reutiliza una fixture downstream previa.",
        ),
        (
            "15 Disposición de UAT S2 003",
            "Permanece OPEN hasta que una sola cadena completa y la evidencia browser sobre esas mismas identidades satisfagan el criterio. El informe lee las etapas y el bloqueo del manifest seleccionado; no mezcla los intentos anteriores ni confunde un checkpoint técnico con el caso formal compuesto. Si comienza otro escenario realmente fresh, sus identidades y resultados reemplazan solo la selección actual, no su historia.",
        ),
        (
            "16 Disposición de UAT S3 002",
            "CLOSED ISOLATION SCOPED. La base SQLite implícita se separa por módulo con verificación de propiedad, ruta y sidecars antes de eliminar; no se reinician URLs PostgreSQL explícitas ni los pasos ordenados internos del módulo. El primer intento interrumpido sin JUnit se conserva y no se suma. La repetición única de 914 casos produjo 898 PASS, un FAIL legado, cero ERROR y 15 SKIP. Cinco fallos y tres errores anteriores por estado compartido ahora PASS en esa misma invocación. No se promete orden arbitrario de pruebas individuales ni se declara global verde.",
        ),
        (
            "17 Disposición de UAT S3 003",
            (
                f"Disposición registrada: {dispositions.get('UAT-S3-003', 'NOT_RECORDED')}; su evidencia es enfocada, no una aprobación global. "
                "GET parent reconciliation es diagnóstico de solo lectura y también exige permiso de lectura ADMIN y ámbito organizacional, con bootstrap desactivado; un lector acotado a Workspace no recibe el inventario organizacional. POST reconcile live parents exige permisos existentes de crear y editar revisión, ámbito organizacional, IDs exactos y ETag. La implementación aplica rechazos de tenant ajeno, padre inexistente, Project como padre, ancestros incompletos, identidades no relacionadas y claves reemplazadas concurrentemente antes de DML; los resultados reales se consultan en la suite enfocada. El clone ordinario conserva su rechazo preventivo."
            ),
        ),
        (
            "18 Regresión ADMIN y USER",
            "Se conservan las guardas que impiden crear o activar Project mediante configuración genérica y CORE. La ruta canónica USER materializa pending. Con autorización expresa posterior del usuario, App.tsx recibió un ajuste mínimo de fallback USER para tenants sin Project operacional heredado: reutiliza StrategicWorkspaceShell, WorkspaceNavigation y StrategicViewOutlet, y mantiene canConfigure=false. No se afirma App inalterado, no se crea un Project ficticio ni un catálogo de navegación paralelo, y no se conceden capacidades ADMIN desde USER. La suite shell se registra separada de la regresión backend anterior al cambio.",
        ),
        (
            "19 Preservación del Project operacional 14",
            "La comprobación operacional usa transacciones de solo lectura y el MD5 de to_jsonb de la fila completa. Se comparan la identidad, estado, versión, número de Workspaces, políticas publicadas y cabeza de Alembic antes y después. La ausencia de ID14 en la base temporal nunca se presenta como protección del ID14 operacional.",
        ),
        (
            "20 Inmutabilidad histórica",
            "Las pruebas CORE enfocadas protegen el snapshot publicado y las identidades existentes. El puente Five Case no modifica ningún registro ni hash histórico. La prueba negativa completa de fuentes live después de un FID nuevo requiere alcanzar ese FID; mientras no se ejecute, no se declara demostrada la inmutabilidad end to end de Gate, decisión, QRA y Post FID.",
        ),
        (
            "21 Regresión y límites PostgreSQL",
            "Los resultados se desglosan por ejecución, motor y alcance. La suite global terminó y mantiene su FAIL real; sus 914 casos no se suman a reruns o focused tests para inflar cobertura única. Las pruebas estructurales CORE usan SQLite privado; la cadena nueva usa PostgreSQL. La moneda y el shell USER cambiaron después de iniciar el broad y tienen evidencia enfocada propia; no se les atribuye retroactivamente esa cobertura. Los SKIP PostgreSQL u opt in no son PASS ni sustituyen la prueba A0 a A31.",
        ),
        (
            "22 Evidencia Chromium",
            "Chromium real usa la API aislada sin mocks ni estados terminales prefabricados. Tras el fallback USER autorizado se abrieron las vistas estratégicas con identidad explícita; eso cierra solo el bloqueo de shell. El conteo y las identidades exactas se registran a continuación. La prueba sigue fallando si la cadena backend no ha completado las etapas downstream. No se atribuye navegación de un tenant a otro escenario seleccionado. El login usa el endpoint real con tenant_slug y su sesión auténtica; no valida el selector de tenant inexistente en el formulario.",
        ),
        (
            "23 Defectos restantes",
            "Los defectos previos se conservan con su estado anterior y su disposición actual. La falta de ejecución completa es un bloqueo de validación; no se convierte automáticamente en defecto de producto. Los hallazgos de la fixture se distinguen de respuestas legítimas de negocio y de fallos funcionales reproducibles.",
        ),
        (
            "24 Deuda técnica",
            "Persisten el bootstrap de esquema anterior a 0068 y los starters administrativos con semánticas distintas entre módulos. El harness histórico FID no acredita esta release y no se omite ni se marca xfail para aparentar éxito. Cualquier deprecación exige que el reemplazo moderno haya demostrado primero su cobertura. Las credenciales aisladas no deben usarse fuera de UAT.",
        ),
        (
            "25 Decisión técnica de liberación",
            facts["release_status"]
            + ". Decisión NO GO. La presencia de código, un subset verde o el avance parcial no satisfacen los criterios completos de liberación.",
        ),
        (
            "26 Estado de aceptación humana",
            "PENDING. No se infieren firmas de Sponsor, Enterprise, Portfolio, Project Controls, Risk, Permits, Finance, Investment Authority ni representantes de negocio. Las celdas de nombre, fecha y evidencia de firma se mantienen sin completar hasta que intervenga una persona autorizada.",
        ),
        ("27 Siguiente recomendación", NEXT),
    ]
    for heading, body_text in sections:
        style.h(doc, heading)
        para(doc, body_text)
        number = int(heading.split()[0])
        if number == 4:
            para(
                doc,
                "También se corrige la propagación de moneda de Proposal al snapshot nuevo de entrada Portfolio. Se conservan currency y currency_code del origen, sin fallback a una moneda solicitada ni actualización de snapshots o evaluaciones COMPLETED anteriores. Cuatro pruebas específicas aprobaron.",
            )
        elif number == 6:
            table(
                doc,
                ["Referencia", "Valor"],
                [
                    [field, setup.get(field, "No alcanzado")]
                    for field in (
                        "tenant_id",
                        "tenant_slug",
                        "enterprise_id",
                        "portfolio_id",
                        "capital_cycle_id",
                        "capital_demand_id",
                        "allocation_plan_id",
                        "idea_id",
                        "analysis_id",
                        "proposal_id",
                        "strategic_gate_id",
                        "project_creation_request_id",
                        "project_id",
                    )
                ],
                [2.2, 4.86],
            )
        elif number == 7:
            table(
                doc,
                ["Persona UAT", "ID", "Funciones asignadas"],
                [
                    [name, value["id"], ", ".join(value.get("roles", []))]
                    for name, value in setup.get("actors", {}).items()
                ],
                [1.4, 0.5, 5.16],
            )
        elif number == 8:
            configs = setup.get("configurations", {})
            table(
                doc,
                ["Dominio", "Referencia capturada"],
                [
                    [name, json.dumps(value, ensure_ascii=False, default=str)]
                    for name, value in configs.items()
                    if name not in {"project_template"}
                ],
                [2.1, 4.96],
            )
            for stage in stages:
                if stage["stage"] == "CURRENT_SOURCE_CONFIGURATION":
                    table(
                        doc,
                        ["Fuente", "ID", "Revisión", "Hash"],
                        [
                            [
                                name,
                                val.get("id"),
                                val.get("revision"),
                                val.get("content_hash"),
                            ]
                            for name, val in stage["data"].items()
                        ],
                        [1.5, 0.5, 0.7, 4.36],
                    )
        elif number == 9:
            matrix(doc, facts)
        elif number == 12:
            diagnosis = facts["gate_blocker_diagnosis"]
            if diagnosis.get("context", {}).get("project_id") == setup.get(
                "project_id"
            ):
                policy, profile = diagnosis["policy"], diagnosis["profile"]
                counts = diagnosis["counts"]
                para(
                    doc,
                    f"Gate {diagnosis['preparation']['id']} conserva DRAFT / BLOCKED y source_current=true. "
                    f"Política {policy['id']}, revisión {policy['revision']}; perfil {profile['code']} "
                    f"{profile['id']}, revisión {profile['revision']}. Contexto CAPITAL / CAPITAL_OWNER. "
                    f"Hay {counts['valid_requirements']} de {counts['required_requirements']} requisitos válidos, "
                    f"cuatro faltantes bloqueantes y cuatro no bloqueantes; cobertura del Project 0 por ciento.",
                )
                labels = {
                    "PERMIT_READINESS": "Permisos",
                    "CONTRACTING_STRATEGY_READINESS": "Contratación",
                    "EXECUTION_PLAN_READINESS": "Plan de ejecución",
                    "SCHEDULE_MATURITY": "Madurez del cronograma",
                }
                table(
                    doc,
                    ["Referencia controlada pendiente", "Clase", "Regla efectiva"],
                    [
                        [
                            labels.get(item["evidence_type"], item["evidence_type"]),
                            item["evidence_class"],
                            f"Validación independiente; vigencia {item['validity_days']} días; EACH_PROJECT; cobertura 100 por ciento",
                        ]
                        for item in diagnosis["missing_blocking_requirements"]
                    ],
                    [2.5, 0.6, 3.96],
                )
                para(
                    doc,
                    "Los cuatro requisitos tienen provider_key nulo y exigen referencias controladas presentadas y validadas. "
                    "No son providers obsoletos: Commercial está actual; Permit tiene assessment READY y 3 de 3 permisos aprobados. "
                    "El override Permit Class A solo aplica MINING, no CAPITAL. Los registros existentes no sustituyen automáticamente "
                    "la validación de una referencia Class B. No se inventaron documentos de ejecución o cronograma ni se creó evidencia manual para forzar el paquete.",
                )
        elif number == 15:
            para(
                doc,
                facts["current_blocker_summary"],
            )
            table(
                doc,
                ["Etapa técnica del manifest seleccionado", "Resultado registrado"],
                [
                    [item["stage"], item.get("status", "NOT_RECORDED")]
                    for item in stages
                ],
                [5.16, 1.9],
            )
        elif number == 18:
            para(
                doc,
                "Shell enfocado: "
                + json.dumps(facts["suites"]["shell"].get("counts"))
                + ". SHA-256 actual de App.tsx: "
                + facts["app_change"]["actual_sha256"]
                + ". El hash identifica el cambio autorizado; no demuestra igualdad con el baseline.",
            )
        elif number == 17:
            para(doc, facts["core_scope_rerun_boundary"])
        elif number == 19:
            before, after = facts["operational_before"], facts["operational_after"]
            table(
                doc,
                ["Control", "Antes", "Después"],
                [
                    [
                        field,
                        before.get(field),
                        after.get(field, "Pendiente de captura final"),
                    ]
                    for field in (
                        "workspace_count",
                        "workspace14",
                        "published_count",
                        "qra_run_count",
                        "head",
                    )
                ],
                [1.5, 2.78, 2.78],
            )
        elif number == 21:
            table(
                doc,
                ["Ejecución", "Estado", "Resultados"],
                [
                    [
                        name,
                        suite["status"],
                        json.dumps(suite.get("counts"), default=str),
                    ]
                    for name, suite in facts["suites"].items()
                ],
                [2.0, 1.8, 3.26],
            )
            for name, suite in facts["suites"].items():
                for failure in suite.get("failures", []):
                    explanation = {
                        "global": "Único fallo legado: HTTP 409 BUSINESS_CASE_NOT_READY. Su fixture no prepara las fuentes Five Case actuales; no se excluyó ni se deprecó para aparentar éxito.",
                        "browser": "Detención explícita sin FID: el Gate de la misma cadena permanece DRAFT/BLOCKED. Abrir ocho vistas no satisface el recorrido completo.",
                        "core_scope_fixture_attempt": "El guard devolvió 409 y los IDs requeridos [5,6] porque la fixture confirmó solo uno. La ejecución final de 72 casos corrigió esa confirmación sin rebajar la regla.",
                    }.get(name, failure["message"][:180])
                    para(
                        doc,
                        f"{name}: {explanation} Test: {failure['test']}. Respuesta íntegra en {suite['file']}.",
                    )
        elif number == 22:
            para(
                doc,
                "Chromium registró "
                + json.dumps(facts["suites"]["browser"].get("counts"))
                + ". "
                + facts["current_browser_summary"],
            )
        elif number == 23:
            table(
                doc,
                ["ID", "Anterior", "Disposición actual"],
                [
                    [item["id"], item["prior_status"], item["status"]]
                    for item in facts["defects"]
                ],
                [1.25, 2.4, 3.41],
            )
            for finding in facts["defects"]:
                if finding["id"] in {"UAT-S3-005", "UAT-S4-006"}:
                    para(
                        doc,
                        f"{finding['id']} {finding['severity']}. {finding['actual']}. No se aplicó corrección de producto a esta observación en el corte.",
                    )
    doc.core_properties.title = "Informe técnico de remediación de bloqueos UAT"
    doc.core_properties.revision = 2
    # Preserve the footer's existing PAGE field; replace only its literal date.
    for paragraph in doc.sections[0].footer.paragraphs:
        for run in paragraph.runs:
            if "14 octubre 2026" in run.text:
                run.text = run.text.replace("14 octubre 2026", "15 octubre 2026")
    return doc


def updated_addendum(kind, facts):
    doc = Document(SOURCE / NAMES[kind])
    doc.add_page_break()
    style.h(doc, "Revisión de remediación del 15 de octubre de 2026")
    para(
        doc,
        "Esta revisión conserva el documento anterior como historial. Sus resultados del 14 de octubre no se convierten en ejecución de la nueva cadena. La decisión actual continúa NO GO y la aceptación empresarial está pendiente.",
    )
    if kind == "uat":
        para(
            doc,
            "La ejecución nueva se identifica por el manifest fresh capital owner chain y los registros de su tenant aislado. Los casos completos sin todos sus controles acreditados permanecen BLOCKED, aunque una etapa técnica haya aprobado.",
        )
        matrix(doc, facts)
        table(
            doc,
            ["Suite", "Estado", "Conteos"],
            [
                [name, suite["status"], suite.get("counts")]
                for name, suite in facts["suites"].items()
            ],
            [1.8, 1.8, 3.46],
        )
        para(
            doc,
            "Consultar la revisión actual del Final Release Closeout para el bloqueo exacto, la disposición de los hallazgos y la única recomendación siguiente. Las firmas anteriores pendientes siguen sin ser asumidas.",
        )
    else:
        style.h(doc, "Reconciliación explícita de padres CORE", 2)
        para(
            doc,
            "El clone normal sigue devolviendo PROJECT_PARENT_RECONCILIATION_REQUIRED cuando el padre live no está representado. Usar GET /api/v1/admin-configuration/enterprise-structure/enterprise-core-releases/{id}/parent-reconciliation para diagnosticar los IDs exactos existentes sin DML. GET exige admin.enterprise_structure.read y ámbito organizacional con bootstrap desactivado; no entrega inventario global a un lector limitado a Workspace.",
        )
        para(
            doc,
            'Tras revisar las identidades, enviar POST al mismo recurso con /reconcile-live-parents, body {"parent_workspace_ids":[IDs confirmados]} e If-Match con revision_version vigente. Se requieren los permisos existentes de crear y editar revisión, ámbito organizacional y la función structure_editor. No se aceptan nombres ni padres inventados.',
        )
        para(
            doc,
            "Si ya existe DRAFT, el comando sobre la publicación devuelve DRAFT_ALREADY_EXISTS y su ID. Revisar y reconciliar explícitamente ese DRAFT con su propio ETag. Las demás ediciones se preservan; cualquier cambio invalida la validación y aprobación previas. Validar, aprobar y publicar siguen siendo comandos independientes.",
        )
        para(
            doc,
            "No usar esta ruta para crear Project, alterar su lifecycle o copiar una identidad ajena. Antes de publicar o materializar se comprueba de nuevo el ID, external_key y tipo de los padres reconciliados; un reemplazo concurrente bloquea con PARENT_WORKSPACE_IDENTITY_MISMATCH antes de DML.",
        )
        table(
            doc,
            ["Control", "Evidencia"],
            [
                [
                    "Regresión enfocada",
                    facts["suites"]["core"]["status"]
                    + "; "
                    + json.dumps(facts["suites"]["core"].get("counts"))
                    + "; evidencia "
                    + facts["suites"]["core"]["file"],
                ],
                [
                    "Historia",
                    "Snapshot publicado e identidad y lifecycle Project sin reescritura",
                ],
                [
                    "Alcance",
                    "Pruebas estructurales SQLite privadas; no acredita A0 a A31 ni aceptación de negocio",
                ],
                ["Intento previo conservado", facts["core_scope_rerun_boundary"]],
            ],
            [2.1, 4.96],
        )
    doc.core_properties.revision = 2
    return doc


def documents(facts, kinds=("admin", "uat", "closeout")):
    OUT.mkdir(parents=True, exist_ok=True)
    paths = []
    for kind in kinds:
        doc = (
            new_closeout(facts) if kind == "closeout" else updated_addendum(kind, facts)
        )
        path = OUT / NAMES[kind]
        doc.save(path)
        paths.append(path)
    # Retained inputs must still match the exact pre-authoring hashes.
    for value in facts["prior_documents"].values():
        assert digest(Path(value["path"])) == value["sha256"]
    print(
        json.dumps(
            {
                "documents": [str(path) for path in paths],
                "retained_inputs": "UNCHANGED",
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--documents", action="store_true")
    parser.add_argument("--document-kind", choices=("admin", "uat", "closeout"))
    args = parser.parse_args()
    report = collect()
    if args.documents:
        documents(
            report,
            (args.document_kind,)
            if args.document_kind
            else ("admin", "uat", "closeout"),
        )
    else:
        print(
            json.dumps(
                {
                    "status": report["release_status"],
                    "cases": report["case_counts"],
                    "suites": report["suites"],
                },
                indent=2,
            )
        )
