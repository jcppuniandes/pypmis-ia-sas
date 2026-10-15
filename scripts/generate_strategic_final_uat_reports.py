"""Build five strategic UAT documents from inspected sources and explicit results.

Run only after the artifact-operation marker. Builders do not execute UAT,
publish configurations, or convert source/test presence into a passing result.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

import generate_cpig_gap_review_report as base
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "docs/validation/strategic-final-uat"
RESULTS = DATA / "release-results.json"
OUT = ROOT / "outputs/strategic-final-uat-20261014"
FILES = {
    "business": "P&Pmis_Strategic_Macroprocess_Business_Process.docx",
    "user": "P&Pmis_Strategic_Macroprocess_User_Guide.docx",
    "admin": "P&Pmis_Strategic_Macroprocess_Admin_Guide.docx",
    "uat": "P&Pmis_Strategic_Macroprocess_UAT_Plan_and_Results.docx",
    "closeout": "P&Pmis_Strategic_Macroprocess_Final_Release_Closeout.docx",
}
BOUNDARY = (
    "El alcance termina en la gobernanza posterior al FID y en la elegibilidad de "
    "inicialización y activación. No incluye Project Delivery, ejecución de compras, "
    "construcción, pagos, contabilidad, tesorería, operación de activos ni realización "
    "operativa de beneficios. QRA integrado y QRA avanzado no se liberan."
)
CASE_ROUTES = {
    "A1": "enterprise-explorer",
    "A2": "capital-demand-pipeline",
    "A3": "idea-lifecycle",
    "A4": "project-proposal",
    "A5": "strategic-gate-decision",
    "A6": "enterprise-explorer",
    "A7": "strategic-project-planning-entry",
    "A8": "portfolio-evaluation",
    "A9": "portfolio-analysis",
    "A10": "portfolio-budget-planning",
    "A11": "portfolio-resource-planning",
    "A12": "portfolio-scenario-convergence",
    "A13": "portfolio-review-sessions",
    "A14": "business-case-my",
    "A15": "strategic-benefits-my",
    "A16": "commercial-strategy-my",
    "A17": "project-definition-fel",
    "A18": "basic-data-readiness",
    "A19": "risk-readiness-my",
    "A20": "risk-readiness-my",
    "A21": "permit-readiness-my",
    "A22": "owner-capability-readiness",
    "A23": "funds-overview",
    "A24": "financial-model-my-models",
    "A25": "qra-foundation",
    "A26": "qra-foundation",
    "A27": "investment-gate-preparation",
    "A28": "investment-gate-preparation",
    "A29": "investment-gate-preparation",
    "A30": "investment-decision-queue",
    "A31": "post-fid-queue",
    "B1": "enterprise-explorer",
    "C1": "enterprise-explorer",
}
LIMITS = {
    "A1": "El Explorer no crea una segunda identidad Project ni otorga permisos por mostrar un nodo.",
    "A2": "La asignación estratégica de capital no aprueba Budget, Funding ni una inversión.",
    "A3": "Una Idea no es un Workspace. ACCEPTED no omite el Business Analysis requerido por la política.",
    "A4": "La Proposal no autoriza inversión y no sustituye el Strategic Gate.",
    "A5": "APPROVE no es FID ni activa Project. RETURN, REJECT y DEFER conservan efectos distintos.",
    "A6": "La materialización crea el PROJECT Workspace pending. Inicializar no equivale a activar.",
    "A7": "Membership no otorga por sí solo aprobación de fondos ni autorización de ejecución.",
    "A8": "El Portfolio Rank es contextual. No se produce un ranking global independiente del Portfolio.",
    "A9": "El análisis no convierte monedas silenciosamente ni aprueba inversión.",
    "A10": "FROZEN conserva un escenario. No significa aprobado y no modifica Funding.",
    "A11": "La capacidad estratégica no asigna personas ni ejecuta resource leveling de P6 o cambios de fechas Project.",
    "A12": "Una recomendación no es FID. Alternativas incompatibles o inviables no se fuerzan como recomendadas.",
    "A13": "HOLD_FOR_CAPACITY es una solicitud de gobernanza Portfolio, no Gate07K HOLD ni escritura sobre las fuentes.",
    "A14": "Refresh crea una revisión nueva. No reescribe una revisión COMPLETED ni calcula el modelo FIN.",
    "A15": "El plan y registro no ejecutan Benefits Realization.",
    "A16": "La estrategia no ejecuta licitaciones, contratos, compromisos ni compras.",
    "A17": "FEL Stage no equivale a Definition Maturity o PDRI. El puntaje no oculta gaps bloqueantes.",
    "A18": "Basic Data es distinto de FEL y PDRI. Disponibilidad no significa validación ni alta confianza.",
    "A19": "USER consulta el contexto efectivo. Los umbrales reutilizables se editan únicamente en ADMIN.",
    "A20": "Risk Readiness no sustituye la aceptación autorizada ni elimina tratamientos pendientes.",
    "A21": "La cobertura y aplicabilidad no emiten permisos regulatorios ni aprueban condiciones automáticamente.",
    "A22": "Owner Capability no es Resource Planning, Contractor Capability ni Operational Readiness.",
    "A23": "Funding es distinto de Budget. Una fuente o asignación no registra desembolsos ni pagos.",
    "A24": "FIN es determinístico. No altera Treasury o Accounting ni aprueba inversión.",
    "A25": "APPROVED_FOR_ANALYSIS habilita un modelo. No constituye aceptación de un resultado.",
    "A26": "COMPLETED no es ACCEPTED. La disponibilidad del provider no vuelve QRA universalmente obligatorio.",
    "A27": "Gate07J prepara evidencia. No toma la decisión formal Gate07K ni recalcula QRA.",
    "A28": "La comparación de fuentes se exige según la política efectiva. No reinterpreta retrospectivamente paquetes congelados.",
    "A29": "Freeze no es autorización de inversión. Preserva política, providers, fuentes y criterios capturados.",
    "A30": "Solo Gate07K con autoridad configurada decide. PYP-FFA más GO constituye FID; SELECT más GO no lo constituye.",
    "A31": "Elegibilidad posterior al FID no ejecuta Project Delivery ni reactiva un Project ya activo.",
    "B1": "No se imponen por defecto Idea, Portfolio, FEL, Funding, QRA o FID. Un opt-in explícito puede cambiar aplicabilidad.",
    "C1": "No se impone la ruta completa Capital Owner. La autorización interna de la fuente sigue siendo necesaria.",
}
ERRORS = {
    "A3": "Si la Idea aceptada continúa en READY_FOR_BUSINESS_ANALYSIS, completar y revisar BA antes de crear Proposal.",
    "A4": "Distinguir el estado persistido READY_FOR_STRATEGIC_GATE del handoff READY_FOR_STRATEGIC_GATE_DECISION.",
    "A6": "Ante PROJECT_CREATION_PROCESS_REQUIRED usar Project Creation Process, no una API genérica de creación o status.",
    "A9": "Revisar incompatibilidad de moneda y expected_analysis_hash obsoleto antes de repetir el análisis.",
    "A14": "Si una fuente está STALE, actualizar el dominio fuente y crear una nueva revisión antes de volver a revisión independiente.",
    "A25": "Corregir fuentes obsoletas, modelo no aprobado, distribución o correlación inválida. No forzar aprobación.",
    "A26": "Revisar analysis type, scope, capacidad soportada, aceptación independiente, vigencia y convergencia. No adjuntar un resultado manual para suplir Class A.",
    "A27": "NOT_EVALUABLE obligatorio, Class A obsoleta o provider ausente bloquean. Resolver el origen y volver a evaluar.",
    "A28": "Ante CROSS_SOURCE_MISMATCH revisar IDs y hashes capturados en BC y Gate, no editar manualmente el hash esperado.",
    "A30": "Revisar SoD, composición del comité, quorum y autoridad del actor. Evidencia READY no autoriza auto GO.",
    "B1": "Comprobar mobilization_authorized o notice_to_proceed_date y autorización independiente para activar.",
    "C1": "Comprobar authorization_approved, inicialización y actor de activación independiente.",
}

# Business controls are drawn from the inspected route/contracts and the
# required domains, not invented organization-specific policies or defaults.
ADMIN_CONTROLS = {
    "enterprise-structure-configuration": "Tipos permitidos, categorías, composición y revisión de estructura. Ninguna revisión de configuración debe crear o activar PROJECT live fuera del proceso gobernado.",
    "workspace-types": "Workspace Type, categorías, semántica de instancia y reglas de pertenencia. PROJECT conserva su identidad en enterprise_workspaces.",
    "enterprise-workspace-structure": "Relaciones parent/child, tipos admitidos, límites y publicación de Composition Rules. Esta pantalla edita reglas reutilizables, no instancias Project.",
    "foundation-strategic-objectives": "Objetivos estratégicos y trazabilidad de demandas e inversiones. Mantener códigos y referencias históricas al publicar cambios.",
    "workspace-defaults": "Defaults y resolución heredada por contexto. Comparar fuente efectiva y overrides antes de modificar descendientes.",
    "module-catalog-activation": "Catálogo y disponibilidad modular por Workspace. Activación de módulo no es activación de un Project Workspace ni autorización de inversión.",
    "workspace-navigation-profiles": "Perfiles de navegación y visibilidad. La UI ayuda a descubrir acciones, pero no sustituye permisos ni reglas de elegibilidad del backend.",
    "risk-governance-configuration": "Capacidad, apetito y umbrales de riesgo. Son configuración ADMIN; USER consume el contexto efectivo y registra sus evaluaciones de riesgo separadas.",
    "capital-planning-configuration": "Ciclos y taxonomías de demanda, evaluación, handoffs y reglas de planeación. Diferenciar posiciones Target, Current, Recommended y Approved.",
    "strategic-capital-bucket-configuration": "Buckets estratégicos, reglas de asignación y objetivos. Los targets y posiciones no se convierten en Budget o Funding autorizados.",
    "idea-demand-configuration": "Elegibilidad del Workspace propietario, taxonomía Idea y requisitos de evaluación/aceptación. No convertir Idea en Workspace.",
    "business-analysis-configuration": "Aplicabilidad de BA, requisitos de opciones, criterios, recomendación y revisión independiente. Resolver si BA es obligatorio antes del handoff a Proposal.",
    "project-proposal-configuration": "Requisitos de alcance, beneficios, ROM, fechas, riesgos, supuestos y restricciones. Evaluación y handoff no sustituyen Strategic Gate.",
    "strategic-gate-configuration": "Perfil Strategic Gate, criterios y autoridad para APPROVE, RETURN, REJECT y DEFER. Esta autoridad no constituye FID.",
    "portfolio-planning-configuration": "Elegibilidad de intake, membership, contexto Portfolio y gobierno de la entrada de un Project canónico.",
    "portfolio-evaluation-configuration": "Criterios y pesos de evaluación contextual, revisión y priorización. Conservar el Portfolio Rank en su contexto, no un ranking global paralelo.",
    "portfolio-analysis-configuration": "Investor Map y Scorecard, comparabilidad, umbrales y restricciones de moneda. No mezclar monedas silenciosamente.",
    "portfolio-review-configuration": "Ciclo de revisión, snapshot, Project Matrix, conclusiones, acciones y revisión. HOLD_FOR_CAPACITY no equivale a Gate07K HOLD.",
    "portfolio-budget-planning-configuration": "Budget Target, escenarios, selección, fases temporales, Waterline y CAPEX/Expense. FROZEN protege una alternativa, no la aprueba.",
    "portfolio-resource-planning-configuration": "Taxonomía de capacidad y demanda, periodos, escenarios y sobrecargas. No asigna personas ni nivela recursos P6.",
    "portfolio-scenario-convergence-configuration": "Comparabilidad de conjuntos Budget/Resource, viabilidad combinada, conflictos y reglas de recomendación. Recomendación sigue separada de inversión.",
    "project-definition-configuration": "Etapas FEL, requisitos por etapa y evaluación de madurez compatible PDRI. No equiparar etapa FEL, madurez y puntaje ni ocultar gaps bloqueantes.",
    "business-case-configuration": "Casos Strategic, Economic, Commercial, Financial y Management; fuentes y cobertura requeridas, revisión, completion y consistencia con Gate.",
    "benefits-configuration": "Plan, registro, propietarios, medidas, baseline conocida o desconocida y trazabilidad estratégica. La configuración no ejecuta realización operativa.",
    "commercial-strategy-configuration": "Modelo de entrega, packaging, modelo comercial, enfoque contractual y distribución de riesgos. No habilita ejecución de Procurement.",
    "funds-manager-configuration": "Tipos de Funding Source, Fund y Assignment; condiciones y cobertura de Funding Readiness. Mantener la autoridad de fondos separada de Budget y Treasury.",
    "financial-model-configuration": "Periodos, métricas, reconciliación, tasas, estructura de capital, inversionistas y deuda del modelo determinístico. No ejecuta contabilidad ni tesorería.",
    "risk-readiness-configuration": "Aplicabilidad, cobertura, riesgo inherente y residual, tratamientos y aceptación. Conservar fuentes y su vigencia para el provider gobernado.",
    "permit-readiness-configuration": "Requisitos, permisos Project, condiciones, aplicabilidad y cobertura. La política no emite una licencia ni elimina condiciones regulatorias.",
    "basic-data-configuration": "Requisitos por etapa, disponibilidad, validación, calidad, confianza y frescura. Distinguir estas dimensiones de FEL/PDRI.",
    "owner-capability-configuration": "Framework de capacidad, madurez, gaps, mitigación y fuentes INTERNAL, EXTERNAL_CONTRACTED o HYBRID. No confundir con Resource Planning u Operational Readiness.",
    "qra-methodology-source-readiness": "methodology, analysis_types, cost_requirements, schedule_requirements, range_policies, distribution_rules, correlation_rules, model_governance, separation_of_duties y freshness. Solo capacidad Cost/Schedule soportada.",
    "investment-decision-evidence-policy": "gate_profile_code, project contexts, Class A/B/C, providers, cobertura, currentness y Decision Criteria. Preview expone effective_matrix, blockers, warnings y persisted=false. QRA es contextual y no Integrated.",
    "investment-gate-preparation-configuration": "Perfiles Gate, evidencias, revisión y congelación de paquetes. Gate07J no decide la inversión ni ejecuta QRA.",
    "investment-decision-configuration": "Autoridad decisora, comité, miembros, quorum, votación y SoD; outcomes GO, HOLD, RECYCLE y KILL. Solo el contexto PYP-FFA más GO es FID.",
    "post-fid-governance-configuration": "Condiciones posteriores a FID, bloqueos y elegibilidad de inicialización/activación. No activar ni reactivar automáticamente Project.",
}


def load(name: str, default: Any = None) -> Any:
    path = DATA / name
    if not path.exists():
        return default if default is not None else []
    return json.loads(path.read_text(encoding="utf-8-sig"))


def text(value: Any) -> str:
    if value is None:
        return "No registrado"
    if isinstance(value, list):
        return "; ".join(text(v) for v in value)
    if isinstance(value, dict):
        return "; ".join(f"{k}={text(v)}" for k, v in value.items())
    return str(value)


def result_rows(results: dict[str, Any]) -> dict[str, dict[str, Any]]:
    value = results.get("uat_results", results.get("cases", []))
    if isinstance(value, dict):
        return {
            key: item if isinstance(item, dict) else {"status": item}
            for key, item in value.items()
        }
    return {item.get("uat_id", item.get("id", "")): item for item in value}


def effective_cases(
    register: dict[str, Any], results: dict[str, Any]
) -> list[dict[str, Any]]:
    mapped = result_rows(results)
    output = []
    for source in register["cases"]:
        item = dict(source)
        update = mapped.get(source["uat_id"], {})
        for key in (
            "status",
            "actual_result",
            "defect_id",
            "evidence",
            "execution_scope",
        ):
            if key in update:
                item[key] = update[key]
        if item.get("status") not in {"PASS", "FAIL", "BLOCKED"}:
            item["status"] = "BLOCKED"
        output.append(item)
    return output


def p(doc: Document, body: str, lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    if lead:
        run = paragraph.add_run(lead + " ")
        run.bold = True
    paragraph.add_run(body)


def h(doc: Document, title: str, level: int = 1) -> None:
    # Native UI terms retain exact punctuation in prose; plain headings remain black.
    clean = re.sub(r"[^\w\sáéíóúüñÁÉÍÓÚÜÑ]", " ", title)
    doc.add_heading(re.sub(r"\s+", " ", clean).strip(), level)


def table(
    doc: Document, headers: list[str], rows: list[list[Any]], widths: list[float]
) -> None:
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.autofit = False
    base.LINE = "D9D9D9"
    base.set_table_borders(tbl)
    for c, width in zip(tbl.columns, widths, strict=True):
        c.width = Inches(width)
    for i, label in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        cell.width = Inches(widths[i])
        base.set_cell_shading(cell, "17324D")
        base.set_cell_margins(cell, 100)
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = para.add_run(label)
        base.set_run_font(run, 9.5, True, "FFFFFF")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    base.repeat_header(tbl.rows[0])
    for row_index, row in enumerate(rows):
        cells = tbl.add_row().cells
        # Each compact record remains intact across page breaks.
        row_properties = tbl.rows[-1]._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        for i, value in enumerate(row):
            cells[i].width = Inches(widths[i])
            base.set_cell_margins(cells[i], 100)
            if row_index % 2:
                base.set_cell_shading(cells[i], "F5F7F8")
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para = cells[i].paragraphs[0]
            para.paragraph_format.space_after = Pt(1)
            para.paragraph_format.line_spacing = 1.04
            # Allow long stable technical identifiers to wrap without altering their text.
            run = para.add_run(text(value))
            base.set_run_font(run, 9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def make_doc(title: str, intro: str, results: dict[str, Any]) -> Document:
    doc = Document()
    base.configure_document(doc)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.07
    normal.paragraph_format.space_after = Pt(5)
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        for border in list(style._element.xpath("./w:pPr/w:pBdr")):
            border.getparent().remove(border)
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    if base.LOGO.exists():
        shape = (
            doc.add_paragraph().add_run().add_picture(str(base.LOGO), width=Inches(0.6))
        )
        base.set_alt_text(
            shape, "P&Pmis Construction AI", "Logotipo de P&Pmis Construction AI"
        )
    doc.add_paragraph(title, "Title")
    p(
        doc,
        "P&Pmis Construction AI. Capital Portfolio & Investment Governance. Corte del 14 de octubre de 2026.",
    )
    p(doc, intro)
    p(
        doc,
        "La aceptación de negocio requiere firmas de representantes autorizados. Una prueba técnica aprobada o un inventario de código no constituye esa aceptación.",
    )
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("14 octubre 2026  Página ")
    base.set_run_font(run, 8, color="000000")
    base.add_field(footer, "PAGE")
    core = doc.core_properties
    core.title = title
    core.subject = "Validación final del macroproceso estratégico"
    core.author = "P&Pmis Construction AI"
    core.comments = (
        "Validación, guías y cierre estratégico; ninguna firma de negocio es inferida."
    )
    base.set_language(doc)
    base.update_field_settings(doc)
    return doc


def evidence_lines(case: dict[str, Any]) -> list[str]:
    evidence = case.get("evidence", {})
    if isinstance(evidence, dict):
        items = evidence.get("current_increment_execution", [])
        sources = evidence.get("source_tests", [])
        return [
            "Ejecución: " + text(items or "sin ejecución enlazada"),
            "Fuente: " + text(sources),
            "Alcance: " + text(evidence.get("evidence_class", "alcance no registrado")),
        ]
    return [text(evidence)]


def baseline_section(
    doc: Document, results: dict[str, Any], register: dict[str, Any]
) -> None:
    h(doc, "Baseline de validación")
    baseline = results.get("baseline", {})
    fields = [
        (
            "Rama y commit inicial",
            text(baseline.get("branch", "main"))
            + " / "
            + text(
                baseline.get(
                    "commit", register["baseline"].get("parent_reported_commit")
                )
            ),
        ),
        (
            "Working tree",
            baseline.get(
                "working_tree",
                "Conservar todos los cambios locales; consultar captura de baseline del incremento",
            ),
        ),
        ("Alembic requerido", "20261011_0068, una sola cabeza"),
        (
            "Release vinculante",
            "QRA_RESULT_EVIDENCE_POLICY_INTEGRATION_RELEASE_VALIDATED",
        ),
        (
            "Registro de resultados",
            results.get("recorded_at", "Corte del registro local de validación"),
        ),
        ("Estado de release", results.get("release_state", "NOT_VALIDATED")),
        ("Decisión", results.get("release_decision", "NO_GO")),
    ]
    table(doc, ["Control", "Valor registrado"], fields, [1.7, 5.36])
    p(
        doc,
        "No se descarta trabajo local ni se reescribe una migración histórica. La existencia de tablas Project operacionales heredadas no constituye una segunda identidad estratégica: Project Workspace se identifica en enterprise_workspaces con tipo PROJECT.",
    )


def business_doc(
    register: dict[str, Any], results: dict[str, Any], cases: list[dict[str, Any]]
) -> Document:
    doc = make_doc(
        "Proceso de negocio del macroproceso estratégico",
        "Este documento describe cómo se prepara, compara, desarrolla y decide una inversión. Distingue las dependencias obligatorias de las actividades paralelas o condicionales y conserva las autoridades de cada dominio.",
        results,
    )
    h(doc, "Objetivo y principios")
    p(
        doc,
        "ADMIN configura tipos, taxonomías, políticas, umbrales, herencia, numeración, permisos y publicación reutilizables. USER opera instancias Workspace, registros, evaluaciones, escenarios y decisiones con permisos acotados al contexto.",
    )
    p(
        doc,
        "El orden de navegación ayuda a encontrar la siguiente actividad, pero no sustituye la elegibilidad del backend. Una recomendación Portfolio, un escenario congelado, evidencia READY y una decisión FID son objetos de negocio diferentes.",
    )
    h(doc, "Mapa del proceso y condiciones de secuencia")
    phases = [
        (
            "1",
            "Enterprise y Workspace",
            "Configurar primero; crear y operar después. Materializar Project solo en el punto autorizado.",
        ),
        (
            "2",
            "Capital Planning y origination",
            "Demand, Idea y evaluación; BA cuando aplique; Proposal y Strategic Gate.",
        ),
        (
            "3",
            "Portfolio y escenarios",
            "Intake, evaluación y análisis; Budget y Resources pueden elaborarse en paralelo antes de Convergence y Review.",
        ),
        (
            "4",
            "Business Case e inversión",
            "Five Case reúne fuentes de Benefits, Commercial, FEL/PDRI, Basic Data, Risk/Permit, Capability, Funding y FIN.",
        ),
        (
            "5",
            "QRA contextual",
            "Foundation y resultados Cost/Schedule únicamente cuando corresponda por política y capacidad soportada.",
        ),
        (
            "6",
            "Gate y decisión",
            "Preparar evidencia, verificar fuentes/criterios, congelar paquete y someterlo a Gate07K.",
        ),
        (
            "7",
            "Post FID",
            "Evaluar condiciones, bloqueos y elegibilidad. Detenerse antes de ejecución operacional.",
        ),
    ]
    table(doc, ["Fase", "Ámbito", "Relación de negocio"], phases, [0.5, 2.05, 4.51])
    p(
        doc,
        "MANDATORY PREDECESSOR exige una salida previa dentro de la ruta indicada. RECOMMENDED PREDECESSOR orienta sin crear bloqueo. PARALLEL-CAPABLE permite avanzar actividades compatibles. CONDITIONAL depende de política, tipo o contexto. OPTIONAL no se transforma en requisito universal.",
    )
    h(doc, "Capital Owner desde origen hasta FID")
    p(
        doc,
        "La ruta de referencia usa una inversión de expansión de procesamiento de cobre y taxonomía realmente configurada. Idea es un registro estratégico. Después de evaluación y, cuando aplique, Business Analysis independiente, Proposal llega a Strategic Gate. APPROVE habilita el handoff READY_FOR_PORTFOLIO_INTAKE, no constituye FID.",
    )
    p(
        doc,
        "ProjectCreationService.materialize crea la identidad EnterpriseWorkspace de tipo PROJECT después de aprobar la solicitud. La instancia comienza pending. Portfolio Membership, definición, evidencias y autorizaciones no se sustituyen por una activación anticipada.",
    )
    p(
        doc,
        "Budget y Resources producen alternativas estratégicas. Scenario Convergence recomienda una alternativa viable y Periodic Portfolio Review registra conclusiones y solicitudes de gobernanza. La recomendación no decide formalmente la inversión.",
    )
    p(
        doc,
        "Five Case mantiene Strategic, Economic, Commercial, Financial y Management con referencias gobernadas. FEL es etapa; PDRI es madurez. Funding conserva fuentes, fondos y asignaciones separadas del Budget. FIN calcula el caso determinístico sin modificar Treasury/Accounting.",
    )
    p(
        doc,
        "Cuando la política lo requiere, QRA usa un modelo APPROVED_FOR_ANALYSIS y un resultado ACCEPTED actual, compatible con analysis type y scope. Gate07J captura evidencias y criterios. Solo Gate07K con actor autorizado e independiente decide GO, HOLD, RECYCLE o KILL. PYP-FFA más GO establece FID y autorización de inversión/fondos completos, sin activar Project automáticamente.",
    )
    p(
        doc,
        "Esta descripción es el recorrido de uso. La cobertura A0 a A31 debe evaluarse como una ejecución integrada; varias pruebas parciales o históricas no equivalen a una demostración completa fresca.",
    )
    h(doc, "Rutas Contractor Delivery y Direct Internal")
    p(
        doc,
        "CONTRACTOR_DELIVERY parte de una fuente contractual o necesidad de entrega gobernada. La autorización de fuente considera mobilization_authorized o notice_to_proceed_date. Sigue Project Creation, Initialization y Activation independiente. No se fuerza por defecto por Idea, Portfolio, FEL, Funding, QRA o FID.",
    )
    p(
        doc,
        "DIRECT_INTERNAL usa su fuente de autorización interna y authorization_approved. Mantiene creación, inicialización y activación como controles separados. La política explícita puede añadir requisitos contextuales sin convertirlo silenciosamente en Capital Owner.",
    )
    h(doc, "Matriz de trazabilidad y autoridades")
    for phase in dict.fromkeys(item["phase"] for item in register["process_matrix"]):
        h(doc, phase, 2)
        for item in (x for x in register["process_matrix"] if x["phase"] == phase):
            h(doc, f"{item['process_id']} {item['process']}", 3)
            p(
                doc,
                f"{item['input']} → {item['output']}. Prerrequisito ADMIN: {item['admin_prerequisite']}. Transacción: {item['user_transaction']}.",
            )
            p(
                doc,
                f"Autoridad: {item['authority']}. Siguiente: {item['next_process']}. Dependencia: {item['dependency_class']}. {item['dependency_interpretation']}",
            )
    h(doc, "RACI por proceso")
    p(
        doc,
        "R es quien ejecuta, A quien responde por la aprobación, C quien se consulta e I quien se informa. Es una propuesta de responsabilidades derivada del modelo de permisos, no una delegación organizacional ni una firma de aceptación. Las concesiones reales siguen acotadas por tenant, Workspace y vigencia.",
    )
    table(
        doc,
        ["ID", "R", "A", "C e I"],
        [
            [
                x["process_id"],
                x["raci"]["responsible"],
                x["raci"]["accountable"],
                "C: "
                + text(x["raci"]["consulted"])
                + "\nI: "
                + text(x["raci"]["informed"]),
            ]
            for x in register["process_matrix"]
        ],
        [0.6, 1.65, 1.9, 2.91],
    )
    h(doc, "Evidencia gobernada y conservación histórica")
    p(
        doc,
        "Class A requiere la fuente nativa gobernada. Un adjunto manual no sustituye un provider obligatorio. El paquete congelado conserva revisión de política, fuentes, hashes, aplicabilidad, readiness, advertencias y criterios. Una nueva publicación o una fuente actual obsoleta no reescribe el historial congelado.",
    )
    h(doc, "Límite posterior al FID")
    p(doc, BOUNDARY)
    return doc


def user_doc(
    register: dict[str, Any], results: dict[str, Any], cases: list[dict[str, Any]]
) -> Document:
    doc = make_doc(
        "Guía de uso del macroproceso estratégico",
        "Use esta guía para operar registros y decisiones en USER MODE. Cada proceso identifica su propósito, campos, pasos, estados, salida y límites. Las acciones disponibles dependen del permiso del actor y de la política efectiva del Workspace.",
        results,
    )
    h(doc, "Ingreso y contexto de trabajo")
    p(
        doc,
        "Abra /app, seleccione USER MODE y confirme tenant y Workspace visibles antes de operar. Enterprise Explorer y Workspace Context comparten la identidad canónica. Para crear Project use las acciones Create Project, My Project Requests o My Project Workspaces dentro del Explorer. Las políticas reutilizables se configuran en ADMIN, no en estas transacciones.",
    )
    p(
        doc,
        "Si la API responde un conflicto de versión o fuente, vuelva a cargar y compare el cambio antes de repetir. No omita ETag, SoD o validaciones para hacer avanzar un estado. Solo utilice acciones que el backend habilite en allowed_actions.",
    )
    case_map = {x["uat_id"]: x for x in cases}
    user_language = load("user-guide-es.json", {})
    for key in [*(f"A{i}" for i in range(1, 32)), "B1", "C1"]:
        case = case_map[key]
        guidance = user_language.get(key, {})
        related = [x for x in register["process_matrix"] if x["uat_case"] == key]
        h(doc, f"{key} {case['business_process']}")
        p(doc, guidance.get("purpose", case["expected_result"]), "Propósito")
        p(
            doc,
            "Cuando la ruta y política del contexto requieren esta actividad. Usuarios: "
            + text(case["persona"])
            + ".",
            "Cuándo y quién",
        )
        business_preconditions = [
            value
            for value in case["preconditions"]
            if not value.startswith("Disposable isolated UAT")
        ]
        p(doc, text(guidance.get("prerequisites", business_preconditions)) + ".", "Prerrequisitos")
        p(
            doc,
            text([x["input"] for x in related])
            + ". Campos principales: "
            + text(
                case.get(
                    "guide_main_fields",
                    [
                        "governance_model",
                        "source_id",
                        "project_workspace_id",
                        "version",
                        "status",
                    ],
                )
            )
            + ".",
            "Entradas y campos",
        )
        p(doc, "/app/strategy/" + CASE_ROUTES[key] + ".", "Pantalla")
        for index, step in enumerate(guidance.get("steps", case["steps"]), 1):
            p(doc, f"{index}. {step}")
        p(
            doc,
            case.get("guide_statuses")
            or "Solicitud draft → submitted → under_review → approved → materialized. Project pending; Initialization NOT_STARTED → READY_FOR_ACTIVATION → ACTIVATED. Los estados de fuente y aprobación son independientes.",
            "Estados",
        )
        p(doc, LIMITS[key], "Reglas de negocio")
        p(
            doc,
            text([x["output"] for x in related])
            + ". Siguiente proceso: "
            + text([x["next_process"] for x in related])
            + ".",
            "Salida y siguiente paso",
        )
        p(
            doc,
            ERRORS.get(
                key,
                "Revise el Workspace, la revisión de política, campos obligatorios, fuente actual y permisos del actor. Corrija el dominio responsable antes de repetir una transición rechazada.",
            ),
            "Errores frecuentes",
        )
        p(doc, LIMITS[key], "Lo que no hace")
    h(doc, "Ejemplo completo de preparación de una inversión")
    p(
        doc,
        "Ejemplo de uso: Copper Processing Expansion, CAPITAL_OWNER y gate PYP-FFA solo si la política publicada lo determina. Los identificadores se obtienen de los registros creados en el entorno aislado; nunca se reutiliza Project ID14.",
    )
    example = [
        "ADMIN publica las configuraciones aplicables con revisiones, hashes y alcance; un USER confirma Enterprise, Business Unit y Portfolio en Explorer.",
        "USER registra Capital Demand y su ciclo. Registra Idea, la evalúa y completa BA independiente cuando el contexto lo exige.",
        "El originador crea Proposal con alcance, beneficios, costo y fechas ROM, riesgos y restricciones. Un Strategic Gate independiente registra APPROVE y produce el handoff de intake.",
        "La solicitud aprobada se materializa como Project Workspace pending. Se prepara la inicialización, sin activar por conveniencia del ensayo.",
        "Portfolio agrega membership, evalúa y prioriza en contexto. Analiza Investor Map y Scorecard sin mezclar monedas no comparables.",
        "Cost/Finance prepara Budget Target y escenarios; Project Controls prepara demanda/capacidad de recursos. Convergence compara viabilidad y recomienda. Review registra conclusiones y acciones de gobernanza.",
        "El equipo desarrolla Five Case, Benefits, Commercial Strategy, FEL/PDRI y Basic Data. Revisores evalúan Risk, Permit y Owner Capability. Funding y FIN conservan sus autoridades independientes.",
        "Si la política requiere QRA, un revisor independiente aprueba el modelo. El analista ejecuta Cost o Schedule soportado y otro actor acepta el resultado actual. No se exige Integrated QRA.",
        "Gate Preparer evalúa evidencia y criterios, resuelve inconsistencias BC/Gate y solicita revisión. El revisor independiente congela un paquete READY con hashes.",
        "Gate07K recibe el paquete. El actor autorizado o comité aplica sus reglas y registra GO, HOLD, RECYCLE o KILL. GO bajo PYP-FFA establece FID y autorizaciones, no activación automática.",
        "Post-FID evalúa condiciones y bloqueos; deriva al flujo existente de inicialización/activación cuando corresponda. Se detiene antes de Project Delivery.",
    ]
    for index, step in enumerate(example, 1):
        p(doc, f"{index}. {step}")
    p(
        doc,
        "Este ejemplo enseña el uso completo, no afirma que los once pasos o A0 a A31 ya se hayan ejecutado como una sola UAT aceptada. El resultado y alcance de cada evidencia están en el registro UAT.",
    )
    h(doc, "Seguridad y límite de uso")
    p(doc, BOUNDARY)
    return doc


def admin_doc(
    register: dict[str, Any], results: dict[str, Any], routes: list[dict[str, Any]]
) -> Document:
    doc = make_doc(
        "Guía de administración del macroproceso estratégico",
        "ADMIN configura y USER opera. Esta guía cubre la configuración reutilizable previa a las transacciones. Cambiar una política exige revisar alcance, versiones, vigencia, publicación, separación de funciones e impacto histórico.",
        results,
    )
    h(doc, "Modelo de administración y cambio seguro")
    p(
        doc,
        "El propietario de negocio define los criterios y la autoridad configurada. El administrador prepara el draft y un publicador independiente actúa cuando la SoD efectiva lo exige. La asignación de organización no se inventa: use roles y concesiones actuales, acotadas a Workspace y periodo válido.",
    )
    p(
        doc,
        "Abra ADMIN MODE, seleccione el Workspace de configuración, consulte la política efectiva y sus fuentes heredadas, clone la revisión publicada a draft cuando el dominio lo soporte, edite, valide y use preview. Compare el impacto en USER antes de publicar. La disponibilidad de un botón no sustituye autorización de backend.",
    )
    p(
        doc,
        "Una revisión publicada es inmutable. Un sucesor cambia el contexto efectivo futuro según su vigencia; no modifica snapshots COMPLETED, FROZEN o DECIDED. Si cambia una fuente material, el estado actual puede resultar STALE y requerir una nueva revisión de evaluación.",
    )
    p(
        doc,
        "Al solicitar una revisión CORE, el servicio incorpora al borrador las identidades Project ya creadas por USER y su estado live actual. También reconcilia un borrador existente conservando sus demás ediciones; si cambia su contenido, invalida validación y aprobación anteriores. No crea Project ni modifica su lifecycle. Si el padre live no está representado con su identidad exacta, PROJECT_PARENT_RECONCILIATION_REQUIRED exige reconciliación controlada y detiene la operación antes de escribir. No se inventan ancestros ni se modifica el snapshot publicado.",
    )
    admin_routes = [
        x
        for x in routes
        if x.get("mode") == "ADMIN" and x.get("classification") != "READ_ONLY_SHARED"
    ]
    for row in admin_routes:
        label = row.get("navigationTrail", row.get("view", "Configuración")).split(
            " > "
        )[-1]
        h(doc, label)
        p(
            doc,
            "Definir los criterios reutilizables de "
            + label
            + " antes de las transacciones aplicables de USER.",
            "Propósito",
        )
        p(
            doc,
            ADMIN_CONTROLS.get(
                row.get("view", ""),
                "Verificar criterios, taxonomías, permisos y referencias del dominio en la revisión efectiva antes de publicar.",
            ),
            "Controles propios del dominio",
        )
        p(
            doc,
            "Propietario funcional del dominio y administrador con permiso "
            + text(row.get("permission"))
            + ". No equivale a una asignación organizacional ya aprobada.",
            "Propietario de negocio",
        )
        p(doc, row["route"] + ".", "Dónde se configura")
        p(
            doc,
            "Revisar el alcance y las fuentes de política efectiva del Workspace. La herencia se conserva según la resolución implementada por el dominio; no se supone que todas las familias compartan un único esquema ni overrides idénticos.",
            "Herencia",
        )
        p(
            doc,
            "Conservar ID, revision, version y content_hash que el contrato del dominio expone. No editar una publicación ni reutilizar un hash para una configuración distinta.",
            "Versionado",
        )
        p(
            doc,
            "Revisar effective_from y effective_to donde el contrato los admita, así como el contexto as_of. Si la familia no expone fechas, no agregar una vigencia ficticia.",
            "Vigencia",
        )
        p(
            doc,
            "Preparar draft, validar campos y referencias, ejecutar preview sin persistencia cuando esté disponible y publicar solo con la acción autorizada del dominio. Registrar el sucesor y comparar la política efectiva antes y después.",
            "Publicación",
        )
        p(
            doc,
            "Usar autores, revisores y publicadores independientes conforme a las reglas efectivas. No convertir configuration_admin en decisor de inversión ni usar organization_admin como actor ordinario del ensayo.",
            "Separación de funciones",
        )
        p(
            doc,
            "USER debe recibir la política contextual y sus permitted actions. Ningún editor reutilizable se traslada a USER. La publicación no ejecuta automáticamente registros, cálculos QRA, aprobaciones de fondos, FID o activación.",
            "Impacto en USER",
        )
        p(
            doc,
            "Validar alcance y referencias en un entorno aislado, comparar ETag/version, conservar evidencia de preview y revisión, y publicar mediante el flujo existente. No cambiar estructuras o permisos operacionales para acelerar la UAT.",
            "Reglas de cambio seguro",
        )
        p(
            doc,
            "Las revisiones completadas y paquetes congelados conservan política, IDs y hashes capturados. Una incompatibilidad nueva debe reflejarse en vigencia o currentness, no mediante edición retroactiva.",
            "Implicaciones históricas",
        )
    h(doc, "Configuraciones críticas por dominio")
    critical = [
        (
            "Workspace y Project",
            "Tipos, composición, numeración, external_key, templates, Project Type, Governance Model, creación, inicialización y activación. La creación live se hace en USER.",
        ),
        (
            "Capital y origination",
            "Ciclos, demands, buckets, allocation, Idea, BA/options, Proposal y Strategic Gate. Las recomendaciones no adquieren autoridad FID.",
        ),
        (
            "Portfolio",
            "Intake, evaluación, analysis, Budget, Resource, Convergence y Review. Umbrales y criterios son ADMIN; escenarios y revisiones live son USER.",
        ),
        (
            "Investment development",
            "Five Case, Benefits, Commercial, FEL/PDRI, Basic Data, Risk, Permit, Owner Capability, Funding y FIN conservan políticas y fuentes propias.",
        ),
        (
            "QRA",
            "Metodología, source readiness, distribuciones, alcance y capacidad soportada. Cost/Schedule separados; no Integrated QRA.",
        ),
        (
            "Gate y decisión",
            "Evidence Policy, Class A/B/C, providers, Decision Criteria, autoridad, comité, quorum y SoD. Gate07J no adquiere autoridad Gate07K.",
        ),
    ]
    table(doc, ["Familia", "Control requerido"], critical, [1.65, 5.41])
    for cell in doc.tables[-1].rows[0].cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True
    h(doc, "Reglas de aplicabilidad QRA")
    p(
        doc,
        "QRA_RESULT_PROVIDER AVAILABLE permite consumir resultados, pero la política decide la obligación contextual por gobernanza, tipo, stage, valor, riesgo, complejidad e incertidumbre disponibles. Cost y Schedule son calificadores separados. REQUIRED no evaluable falla cerrado; NOT_REQUIRED no genera un bloqueo artificial.",
    )
    p(
        doc,
        "Class A exige fuente nativa aceptada y actual. Revise analysis type, scope, whole/partial project, readiness, convergencia, advertencias, edad y percentiles o target probability soportados. Una exigencia Integrated activa se rechaza. Un preview no publica una política operacional.",
    )
    h(doc, "Cambios prohibidos durante el cierre")
    p(doc, BOUNDARY)
    return doc


def uat_doc(
    register: dict[str, Any],
    results: dict[str, Any],
    cases: list[dict[str, Any]],
    personas: list[dict[str, Any]],
) -> Document:
    doc = make_doc(
        "Plan y resultados de UAT estratégica",
        "El registro separa inspección de fuente, regresión técnica, pruebas HTTP o de navegador y aceptación autorizada de negocio. Los resultados se atribuyen a su entorno y alcance exactos. Las pruebas pendientes se mantienen BLOCKED y no se convierten en PASS por existir código.",
        results,
    )
    baseline_section(doc, results, register)
    h(doc, "Alcance y entorno")
    p(
        doc,
        "Se revisan 44 procesos y 51 casos: A0 a A31 Capital Owner, B1 Contractor Delivery, C1 Direct Internal y 17 negativos o de frontera. La fixture debe ser aislada y representativa, con taxonomía configurada y actores de RBAC reales. Project ID14 y políticas operacionales se preservan.",
    )
    p(
        doc,
        text(
            results.get(
                "environment",
                "La evidencia de cada ejecución identifica su base, backend, navegador y comandos. No se atribuyen resultados de SQLite a PostgreSQL ni de tests DOM a Chromium real.",
            )
        ),
    )
    h(doc, "Personas y permisos")
    table(
        doc,
        ["Persona", "Roles actuales de referencia", "Límite de autoridad"],
        [
            [x["persona"], x["actual_role_codes"], x["sod_and_boundary"]]
            for x in personas
        ],
        [1.65, 2.1, 3.31],
    )
    # A heading row must never be left alone at the foot of the page.
    for cell in doc.tables[-1].rows[0].cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True
    h(doc, "Criterio de resultado y severidad")
    p(
        doc,
        "PASS indica la aserción y alcance ejecutados. FAIL registra un resultado contrario al esperado. BLOCKED registra ausencia de ejecución completa o un prerrequisito no resuelto. Un PASS técnico no equivale a firma humana.",
    )
    p(
        doc,
        "S1 BLOCKER y S2 CRITICAL abiertos impiden liberación. Un S3 MAJOR exige una solución aceptada o decisión de release con workaround aprobado. S4 MINOR y S5 COSMETIC se registran con responsable y disposición. La decisión solo puede ser GO, CONDITIONAL_GO o NO_GO.",
    )
    counts = Counter(x["status"] for x in cases)
    table(
        doc,
        ["Estado", "Casos"],
        [[status, counts[status]] for status in ("PASS", "FAIL", "BLOCKED")],
        [4.9, 2.16],
    )
    # The three-result summary is compact enough to read as a single unit.
    for row in doc.tables[-1].rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    h(doc, "Registro formal de casos")
    for case in cases:
        h(doc, f"{case['uat_id']} {case['business_process']}", 2)
        p(doc, text(case["preconditions"]), "Precondiciones")
        p(
            doc,
            text(case["persona"]) + ". Datos: " + text(case["test_data"]),
            "Persona y datos",
        )
        for index, step in enumerate(case["steps"], 1):
            p(doc, f"{index}. {step}")
        p(doc, case["expected_result"], "Resultado esperado")
        p(doc, case["actual_result"], "Resultado real")
        p(
            doc,
            case["status"]
            + ". Defecto: "
            + text(case.get("defect_id") or "sin defecto de producto inferido")
            + ".",
            "Estado",
        )
        for line in evidence_lines(case):
            p(doc, line)
    findings_section(doc, results, register)
    h(doc, "Aceptación de negocio pendiente")
    p(
        doc,
        "Las firmas no se completan por un agente ni se deducen de resultados técnicos. El workbook contiene campos editables de representante autorizado, decisión, fecha y evidencia para las quince áreas requeridas.",
    )
    p(
        doc,
        text(
            results.get(
                "release_conclusion",
                "NO_GO mientras exista evidencia técnica obligatoria incompleta o un bloqueo material. No se afirma aceptación de negocio sin firma autorizada.",
            )
        ),
    )
    return doc


def findings_section(
    doc: Document, results: dict[str, Any], register: dict[str, Any]
) -> None:
    h(doc, "Defectos y acciones abiertas")
    defects = results.get("defects", [])
    if defects:
        for item in defects:
            h(
                doc,
                text(item.get("id", item.get("defect_id", "Hallazgo")))
                + " "
                + text(item.get("title", item.get("summary", "Hallazgo"))),
                2,
            )
            p(
                doc,
                "Clasificación: "
                + text(item.get("classification"))
                + ". Severidad: "
                + text(item.get("severity"))
                + ". Estado: "
                + text(item.get("status"))
                + ".",
            )
            actual = item.get("actual", item.get("description", item.get("summary")))
            if actual:
                p(doc, text(actual))
            p(
                doc,
                "Acción: "
                + text(
                    item.get("remediation", item.get("action", item.get("workaround")))
                )
                + ". Responsable: "
                + text(item.get("owner"))
                + ".",
            )
            p(doc, "Evidencia: " + text(item.get("evidence")))
    else:
        p(
            doc,
            "El registro de defectos consolidado no contiene resultados en este corte. Esta ausencia no prueba cero defectos ni permite liberar; consultar las brechas de cobertura siguientes.",
        )
    for item in register.get("coverage_findings", []):
        p(doc, text(item["finding_id"]) + ". " + item["summary"] + " " + item["impact"])


def closeout_doc(
    register: dict[str, Any],
    results: dict[str, Any],
    cases: list[dict[str, Any]],
    routes: list[dict[str, Any]],
) -> Document:
    doc = make_doc(
        "Informe técnico y cierre de release estratégico",
        text(
            results.get(
                "release_conclusion",
                "La decisión de liberación permanece NO_GO hasta completar los criterios técnicos obligatorios. La aceptación de negocio exige firmas autorizadas y no se presume a partir de regresiones parciales.",
            )
        ),
        results,
    )
    h(doc, "Conclusión ejecutiva")
    p(
        doc,
        "Decisión: "
        + text(results.get("release_decision", "NO_GO"))
        + ". Estado exacto: "
        + text(results.get("release_state", "NOT_VALIDATED"))
        + ". Próximo incremento recomendado: "
        + text(results.get("next_increment", "B. REMEDIATE STRATEGIC UAT BLOCKERS"))
        + ".",
    )
    baseline_section(doc, results, register)
    h(doc, "Flujo validado y límites de la evidencia")
    p(
        doc,
        text(
            results.get(
                "validated_flow",
                "Se inventarían procesos y se enlazan resultados técnicos por caso. La ruta fresca Capital Owner A0 a A31 no se infiere sumando ejecuciones de dominios, fixtures sembradas ni resultados históricos.",
            )
        ),
    )
    summary_cases = [
        x
        for x in cases
        if x["uat_id"]
        in {"A3", "A4", "A5", "A6", "A14", "A26", "A29", "A30", "A31", "B1", "C1"}
    ]
    table(
        doc,
        ["Caso", "Estado", "Resultado y alcance"],
        [[x["uat_id"], x["status"], x["actual_result"]] for x in summary_cases],
        [0.6, 0.85, 5.61],
    )
    for cell in doc.tables[-1].rows[0].cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True
    h(doc, "Auditoría ADMIN y USER")
    inventory = load("backend-inventory-summary.json", {})
    p(
        doc,
        f"El inventario estático contiene {inventory.get('endpoints', 'n.a.')} endpoints en {inventory.get('modules', 'n.a.')} módulos. Clasificación: {text(inventory.get('mode_counts'))}. La clasificación GET como contexto compartido no prueba ausencia de side effects heredados ni sustituye pruebas HTTP de denegación.",
    )
    p(
        doc,
        text(
            results.get(
                "admin_user",
                "El registro consolidado determina el cierre de defectos de creación/status y de las rutas de configuración. La inspección por sí sola no certifica segregación runtime.",
            )
        ),
    )
    h(doc, "Navegación y reutilización del frontend")
    summary = load("frontend_audit_summary.json", {})
    p(
        doc,
        f"Se revisaron {summary.get('navigationNodes', 'n.a.')} nodos y {summary.get('canonicalRoutes', len(routes))} rutas canónicas. La autoridad de navegación sigue en strategicNavigationBlueprint.ts. Hay {summary.get('reusedLandingAliasNodes', 'n.a.')} alias de landing, que no se consideran pantallas duplicadas por sí solos.",
    )
    p(
        doc,
        text(
            results.get(
                "navigation",
                "La revisión identifica el orden de fases, breadcrumbs, contexto y renderer. Las correcciones y su regresión se documentan en el registro de hallazgos. No se amplía App.tsx para introducir lógica nueva.",
            )
        ),
    )
    h(doc, "Estadísticas de UAT y regresión estratégica")
    counts = Counter(x["status"] for x in cases)
    table(
        doc,
        ["Registro", "Valor"],
        [
            ["Procesos", len(register["process_matrix"])],
            ["Casos", len(cases)],
            ["PASS técnicos por caso", counts["PASS"]],
            ["FAIL", counts["FAIL"]],
            ["BLOCKED", counts["BLOCKED"]],
            [
                "Capital Owner fresco completo",
                results.get("capital_owner_e2e", "NOT_PROVEN"),
            ],
            ["Aceptación humana", results.get("business_signoff", "PENDING")],
        ],
        [4.3, 2.76],
    )
    for item in results.get("executions", []):
        h(doc, text(item.get("name", item.get("id", "Ejecución"))), 2)
        p(
            doc,
            "Estado: "
            + text(item.get("status"))
            + ". Resultado: "
            + text(item.get("counts", item.get("result")))
            + ". Entorno: "
            + text(item.get("environment", item.get("scope")))
            + ".",
        )
        p(
            doc,
            "Evidencia: "
            + text(item.get("evidence"))
            + ". "
            + text(
                item.get(
                    "qualification",
                    "La ejecución solo cubre sus aserciones y datos registrados.",
                )
            ),
        )
    h(doc, "Seguridad e integridad de datos")
    p(
        doc,
        text(
            results.get(
                "security",
                "Los controles de tenant, Workspace, RBAC, SoD, ETag, hashes, providers y referencias se verifican por evidencia ejecutada. El inventario estático no sustituye denegaciones HTTP reales.",
            )
        ),
    )
    p(
        doc,
        text(
            results.get(
                "preservation",
                "Project ID14 requiere captura anterior y posterior de fila, external_key, Record Code, status, activo y hash. Las políticas publicadas y el registro provider requieren comparación read-only. Ningún resultado se presume si falta esa captura.",
            )
        ),
    )
    h(doc, "Datos migraciones y catálogo persistente")
    p(
        doc,
        f"El inventario de fuente contiene {inventory.get('tables', 'n.a.')} tablas: {text(inventory.get('table_classification_counts'))}. ACTIVE, CONFIGURATION, HISTORICAL_REQUIRED y TRANSIENT/COMMAND_RECEIPT describen retención; no autorizan borrar tablas o filas.",
    )
    p(
        doc,
        text(
            results.get(
                "migrations",
                "La validación debe comprobar cabeza única 20261011_0068 y, solo en una base descartable, downgrade y re-upgrade seguros. No se reescribe historial ni se ejecuta ese ciclo sobre la base operacional.",
            )
        ),
    )
    h(doc, "Rendimiento accesibilidad y usabilidad")
    p(
        doc,
        text(
            results.get(
                "performance",
                "Las mediciones locales identifican operación, entorno y duración. No son SLA de producción. Gate overview, resolución QRA, BC overview, preview, historia, navegación, Explorer y Project Matrix requieren evidencia representativa.",
            )
        ),
    )
    p(
        doc,
        text(
            results.get(
                "accessibility",
                "Verificar teclado, labels, tablas, errores de formulario, foco modal, breadcrumbs, header compacto, visibilidad Workspace y warnings. El audit estructural de los documentos no certifica WCAG del frontend.",
            )
        ),
    )
    h(doc, "Resumen de defectos por severidad")
    defects = results.get("defects", [])
    table(
        doc,
        ["Severidad", "Abiertos", "Cerrados"],
        [
            [
                severity,
                sum(item.get("severity", "").startswith(severity) and not item.get("status", "").startswith("CLOSED") for item in defects),
                sum(item.get("severity", "").startswith(severity) and item.get("status", "").startswith("CLOSED") for item in defects),
            ]
            for severity in ("S1", "S2", "S3", "S4", "S5")
        ],
        [3.5, 1.78, 1.78],
    )
    findings_section(doc, results, register)
    h(doc, "Revisión de limpieza y deuda técnica")
    backend = load("backend-code-cleanup.json")
    frontend = load("frontend_cleanup_candidates.json")
    counts_cleanup = Counter(
        x.get("classification", "UNKNOWN") for x in backend + frontend
    )
    p(
        doc,
        f"Se clasifican {len(backend)} candidatos backend y {len(frontend)} frontend: {text(dict(counts_cleanup))}. No encontrar un consumidor directo no demuestra código muerto. Requiere distinguir REQUIRES_DEPRECATION, STILL_REFERENCED, HISTORICAL_COMPATIBILITY_REQUIRED, SAFE_TO_REMOVE y UNKNOWN.",
    )
    p(
        doc,
        text(
            results.get(
                "cleanup",
                "No se eliminan candidatos sin clasificación SAFE_TO_REMOVE y regresión que pruebe seguridad. La compatibilidad histórica, rutas externas y fuentes operacionales se conservan; ninguna tabla histórica se descarta.",
            )
        ),
    )
    h(doc, "Autoridades de negocio conservadas")
    for rule in register["business_rules"]:
        p(doc, rule["rule_id"] + ". " + rule["rule"])
    h(doc, "Documentos y trazabilidad entregados")
    p(
        doc,
        "El paquete contiene Business Process, User Guide, Admin Guide, UAT Plan and Results, UAT Traceability de ocho hojas y este Final Release Closeout. Las guías separan instrucciones de uso de resultados ejecutados y el workbook conserva campos de firma humana sin completar.",
    )
    h(doc, "Criterios de liberación y decisión")
    for item in results.get("checks", []):
        p(
            doc,
            text(item.get("id"))
            + ". "
            + text(item.get("title"))
            + ": "
            + text(item.get("status"))
            + ". "
            + text(item.get("actual", item.get("evidence"))),
        )
    p(
        doc,
        "GO requiere cero S1 y S2 abiertos, ningún bloqueo material de flujo, segregación, autoridad, datos o seguridad, ruta Capital Owner completa, regresión Contractor/Direct y Gate/FID. CONDITIONAL_GO solo permite asuntos no bloqueantes aceptados con responsable y workaround. Sin esas pruebas no se declara estado técnico listo.",
    )
    p(
        doc,
        "Decisión final: "
        + text(results.get("release_decision", "NO_GO"))
        + ". Estado final: "
        + text(results.get("release_state", "NOT_VALIDATED"))
        + ".",
    )
    h(doc, "Siguiente incremento y frontera de parada")
    p(doc, text(results.get("next_increment", "B. REMEDIATE STRATEGIC UAT BLOCKERS")))
    p(doc, BOUNDARY)
    return doc


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--results", type=Path, default=RESULTS)
    parser.add_argument("--output-dir", type=Path, default=OUT)
    parser.add_argument("--only", choices=FILES)
    parser.add_argument("--check-inputs", action="store_true")
    args = parser.parse_args()
    register = load("scenario-coverage-register.json", {})
    if (
        len(register.get("cases", [])) != 51
        or len(register.get("process_matrix", [])) != 44
    ):
        raise ValueError("Expected inspected register with 51 cases and 44 processes")
    results = (
        json.loads(args.results.read_text(encoding="utf-8-sig"))
        if args.results.exists()
        else {}
    )
    cases = effective_cases(register, results)
    if args.check_inputs:
        print(
            json.dumps(
                {
                    "cases": len(cases),
                    "processes": len(register["process_matrix"]),
                    "rules": len(register["business_rules"]),
                    "results_present": args.results.exists(),
                    "status_counts": dict(Counter(x["status"] for x in cases)),
                },
                ensure_ascii=False,
            )
        )
        return
    if not args.results.exists():
        raise FileNotFoundError(
            "Explicit release-results.json required before authoring"
        )
    args.output_dir.mkdir(parents=True, exist_ok=True)
    routes = load("frontend_routes.json")
    builders = {
        "business": lambda: business_doc(register, results, cases),
        "user": lambda: user_doc(register, results, cases),
        "admin": lambda: admin_doc(register, results, routes),
        "uat": lambda: uat_doc(
            register, results, cases, load("backend-personas-rbac.json")
        ),
        "closeout": lambda: closeout_doc(register, results, cases, routes),
    }
    for key in [args.only] if args.only else FILES:
        document = builders[key]()
        path = args.output_dir / FILES[key]
        document.save(path)
        print(
            json.dumps(
                {
                    "path": str(path),
                    "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                    "paragraphs": len(document.paragraphs),
                    "tables": len(document.tables),
                },
                ensure_ascii=False,
            )
        )


if __name__ == "__main__":
    main()
