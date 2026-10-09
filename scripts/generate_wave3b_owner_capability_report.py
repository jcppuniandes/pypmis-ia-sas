"""Compose Wave 3B Word from the factual release report and reused style tools."""

from __future__ import annotations

import re
from pathlib import Path

import generate_cpig_gap_review_report as base
import generate_wave2b_risk_basic_data_report as diagrams
import generate_wave3a_portfolio_review_report as layout
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/79-wave3b-owner-capability-readiness-release.md"
QA = ROOT / "outputs/.qa/wave3b-owner-capability"
DATE = "2026-10-08"
OUTPUT = ROOT / "outputs" / f"Informe_Tecnico_Wave3B_Owner_Capability_Readiness_{DATE}.docx"


def main() -> None:
    source = SOURCE.read_text(encoding="utf-8")
    match = re.search(r"^\*\*Estado actual:\*\*\s+(.+)$", source, re.MULTILINE)
    status = match.group(1).strip() if match else "VALIDACIÓN EN CURSO"
    if status == "OWNER_CAPABILITY_READINESS_RELEASE_VALIDATED":
        status = "RELEASE VALIDADO"
    QA.mkdir(parents=True, exist_ok=True)
    base.SOURCE, base.QA, base.DIAGRAM_DIR = SOURCE, QA, QA / "diagrams"
    base.DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
    spec = {
        "name": "wave3b-assessment-history.png",
        "title": "Evaluación gobernada y lectura de vigencia",
        "rows": [
            ["Policy and Evidence", "DRAFT", "IN REVIEW"],
            ["Independent Review", "COMPLETED Snapshot", "Current READY or STALE"],
            ["Immutable History", "Refresh Preview", "Successor DRAFT"],
        ],
        "edges": [(0, 0, 0, 1), (0, 1, 0, 2), (0, 2, 1, 0), (1, 0, 1, 1), (1, 1, 1, 2), (1, 1, 2, 0), (1, 2, 2, 1), (2, 1, 2, 2)],
        "alt": "La política y evidencia alimentan DRAFT, submit y revisión independiente. COMPLETED congela la historia. La proyección actual puede quedar STALE; refresh preview prepara un sucesor DRAFT sin sobrescribir el cierre anterior.",
    }
    image = diagrams.render_diagram(spec)
    doc = Document()
    base.configure_document(doc)
    base.set_language(doc)
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        properties = doc.styles[name]._element.find(qn("w:pPr"))
        if properties is not None:
            for border in properties.findall(qn("w:pBdr")):
                properties.remove(border)
    doc.core_properties.title = "Informe técnico Wave 3B Owner Capability Readiness"
    doc.core_properties.subject = "Capacidad gobernada del propietario y evidencia de validación"
    doc.core_properties.author = "P&Pmis Construction AI"
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(30)
    if layout.LOGO.exists():
        shape = paragraph.add_run().add_picture(str(layout.LOGO), width=Inches(1.3))
        base.set_alt_text(shape, "P&Pmis Construction AI", "Logotipo de P&Pmis Construction AI")
    title = doc.add_paragraph("Informe técnico Wave 3B\nOwner Capability Readiness", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(25)
    for value in ("P&Pmis Construction AI", "Capital Portfolio and Investment Governance", DATE, status):
        paragraph = doc.add_paragraph(value)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(15)
    doc.add_page_break()
    doc.add_heading("Control del documento", level=1)
    base.add_table(doc, [
        ["Campo", "Valor"], ["Producto", "P&Pmis Construction AI"],
        ["Incremento", "Wave 3B Owner Capability Readiness"],
        ["Baseline", "main y 9cb0969 con cambios previos preservados"],
        ["Migración", "20261008_0065 a 20261008_0066"],
        ["Fecha local", DATE], ["Estado", status],
    ])
    doc.add_heading("Propósito del informe", level=1)
    base.add_inline(doc.add_paragraph(), "Documentamos la capacidad del propietario para gobernar y supervisar una inversión, su política configurable, evaluación por proyecto y evidencia de pruebas. Resource Planning, FEL, PDRI, Commercial, Risk, Basic Data, Gate y FID mantienen su autoridad. El cierre técnico diferencia implementación, resultados observados y trabajo futuro.")
    doc.add_heading("Contenido", level=1)
    for text in ("Baseline arquitectura y autoridad", "Framework madurez requisitos y mitigación", "Integraciones proveedor y vigencia", "ADMIN USER seguridad y persistencia", "Pruebas regresiones y cierre", "Prerrequisitos QRA y siguiente incremento"):
        doc.add_paragraph(text, style="List Bullet")
    doc.add_page_break()
    base.DIAGRAMS = []
    base.add_source_content(doc, [], page_break_heading_prefixes=(), skip_preface=True)
    anchor = next(p for p in doc.paragraphs if p.style.name.startswith("Heading") and p.text.startswith("24. "))
    figure = doc.add_paragraph()
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.paragraph_format.keep_together = True
    shape = figure.add_run().add_picture(str(image), width=Inches(6.75))
    base.set_alt_text(shape, spec["title"], spec["alt"])
    anchor._p.addnext(figure._p)
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            paragraph.text = re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", paragraph.text).replace("_", " ")).strip()
        elif paragraph.style.name in ("Normal", "List Bullet", "List Number"):
            for run in paragraph.runs:
                run.font.size = Pt(11)
    layout.format_tables(doc)
    widths = {
        ("Dominio", "Autoridad de escritura", "Límite"): (1.55, 2.5, 3.01),
        ("Componente existente", "Uso en Wave 3B", "Decisión"): (1.6, 3.1, 2.36),
        ("Acción", "ADMIN Configuration", "USER operativo"): (2.3, 2.45, 2.31),
        ("Prerrequisito", "Estado arquitectónico", "Brecha para QRA"): (1.8, 2.05, 3.21),
        ("Validación", "Resultado", "Evidencia"): (2.45, 1.0, 3.61),
        ("Grupo de API", "Operaciones", "Control"): (2.70, 1.80, 2.56),
    }
    for table in doc.tables:
        selected = widths.get(tuple(c.text for c in table.rows[0].cells))
        if selected:
            for column, width in zip(table.columns, selected, strict=True):
                column.width = Inches(width)
            for row in table.rows:
                for index, cell in enumerate(row.cells):
                    cell.width = Inches(selected[index])
    for section in doc.sections:
        section.different_first_page_header_footer = True
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        base.set_run_font(header.add_run("P&Pmis Construction AI | Wave 3B"), size=7.5, color=base.BLACK)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        base.set_run_font(footer.add_run(f"Informe técnico | {DATE} | Página "), size=7.5, color=base.MUTED)
        base.add_field(footer, "PAGE")
        footer.add_run(" de ")
        base.add_field(footer, "NUMPAGES")
    base.update_field_settings(doc)
    doc.save(OUTPUT)
    print(f"created={OUTPUT}")


if __name__ == "__main__":
    main()
