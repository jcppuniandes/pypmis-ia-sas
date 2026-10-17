"""Read-only scope/native-feature verification of the authorized UAT successor."""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
NAME = "P&Pmis_Strategic_Macroprocess_UAT_Traceability.xlsx"
SOURCE = ROOT / "outputs/gate07j-controlled-class-b-20261015" / NAME
OUTPUT = ROOT / "outputs/gate07j-controlled-class-b-20261016" / NAME
FACTS = ROOT / "docs/validation/gate07j-controlled-class-b/rebuild-results.json"
QA = ROOT / "outputs/.qa/gate07j-controlled-class-b-rebuild/workbook"
CASE_IDS = ("A27", "A28", "A29", "A30", "A31")
DEFECT_IDS = ("UAT-S2-003", "UAT-S3-005", "UAT-S4-006", "UAT-ENV-001")
NEW_DEFECT_ROWS = {"UAT-S2-005": 24, "UAT-S3-006": 25}


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def text(value: object) -> str:
    if isinstance(value, list):
        return "\n".join(str(item) for item in value)
    return "" if value is None else str(value)


def style(cell: object) -> tuple[str, ...]:
    return tuple(
        str(getattr(cell, key))
        for key in (
            "font",
            "fill",
            "border",
            "alignment",
            "number_format",
            "protection",
        )
    )


def row_for(sheet, record_id: str) -> int:
    rows = [
        row
        for row in range(8, sheet.max_row + 1)
        if sheet.cell(row, 1).value == record_id
    ]
    assert len(rows) == 1, (sheet.title, record_id, rows)
    return rows[0]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--visual-reviewed", action="store_true")
    args = parser.parse_args()
    facts = json.loads(FACTS.read_text(encoding="utf-8"))
    manifest = json.loads((QA / "edit-manifest.json").read_text(encoding="utf-8"))
    assert facts["schema_version"] == 1
    assert manifest["source_sha256"] == digest(SOURCE)
    assert manifest["facts_sha256"] == digest(FACTS), (
        "Collector changed after authoring; regenerate the final cut"
    )
    assert manifest["signature_boundary"] == "ACCEPTED alone remains BLOCKED"
    assert (
        facts["release"]["signoffs"] == []
        and facts["release"]["business_release_authorized"] is False
    )
    before = openpyxl.load_workbook(SOURCE)
    before_cached = openpyxl.load_workbook(SOURCE, data_only=True)
    after = openpyxl.load_workbook(OUTPUT)
    cached = openpyxl.load_workbook(OUTPUT, data_only=True)
    assert before.sheetnames == after.sheetnames and len(after.sheetnames) == 8
    cases_by_id = {item["id"]: item for item in facts["cases"]}
    defects_by_id = {item["id"]: item for item in facts["defects"]}
    criteria_by_id = {item["id"]: item for item in facts["releaseCriteria"]}
    assert set(cases_by_id) == set(CASE_IDS)
    assert set(defects_by_id) == {*DEFECT_IDS, *NEW_DEFECT_ROWS}
    allowed: dict[str, set[str]] = {name: set() for name in before.sheetnames}
    allowed_heights: dict[str, set[int]] = {name: set() for name in before.sheetnames}
    allowed_columns = {
        "02_UAT_Cases": set("VWXYZ"),
        "06_Defects": {"R"},
        "08_Release_Criteria": {"M"},
    }
    for record_id in CASE_IDS:
        row = row_for(before["02_UAT_Cases"], record_id)
        allowed["02_UAT_Cases"].update(f"{col}{row}" for col in "CFGHIJVWXYZ")
        allowed_heights["02_UAT_Cases"].add(row)
    allowed["02_UAT_Cases"].update(f"{col}7" for col in "VWXYZ")
    for record_id in DEFECT_IDS:
        row = row_for(before["06_Defects"], record_id)
        allowed["06_Defects"].update(f"{col}{row}" for col in "EFGIR")
        allowed_heights["06_Defects"].add(row)
    allowed["06_Defects"].update({"B5", "R7"})
    for row in NEW_DEFECT_ROWS.values():
        allowed["06_Defects"].update(f"{col}{row}" for col in "ABCDEFGHIR")
        allowed_heights["06_Defects"].add(row)
    for record_id in criteria_by_id:
        row = row_for(before["08_Release_Criteria"], record_id)
        assert before["08_Release_Criteria"].cell(row, 3).value == "TECHNICAL"
        allowed["08_Release_Criteria"].update(f"{col}{row}" for col in "DEFM")
        allowed_heights["08_Release_Criteria"].add(row)
    allowed["08_Release_Criteria"].add("M7")
    changes = {}
    formulas = 0
    tables = 0
    for name in before.sheetnames:
        old, new = before[name], after[name]
        assert old.freeze_panes == new.freeze_panes, name
        assert list(old.tables) == list(new.tables), name
        assert old.sheet_state == new.sheet_state, name
        assert old.sheet_view.showGridLines == new.sheet_view.showGridLines, name
        assert old.merged_cells == new.merged_cells, name
        assert old.protection == new.protection, name
        assert len(old._charts) == len(new._charts), name
        assert len(old._images) == len(new._images), name
        assert len(old._pivots) == len(new._pivots), name
        for table_name in old.tables:
            tables += 1
            expected_ref = (
                "A7:I25" if name == "06_Defects" else old.tables[table_name].ref
            )
            assert new.tables[table_name].ref == expected_ref, (name, table_name)
            assert (
                old.tables[table_name].tableStyleInfo
                == new.tables[table_name].tableStyleInfo
            )
        assert etree.tostring(old.data_validations.to_tree()) == etree.tostring(
            new.data_validations.to_tree()
        ), name
        assert list(old.conditional_formatting) == list(new.conditional_formatting), (
            name
        )
        for region in old.conditional_formatting:
            assert [
                etree.tostring(rule.to_tree())
                for rule in old.conditional_formatting[region]
            ] == [
                etree.tostring(rule.to_tree())
                for rule in new.conditional_formatting[region]
            ], (name, str(region))
        for col, dimension in old.column_dimensions.items():
            if col not in allowed_columns.get(name, set()):
                actual = new.column_dimensions[col]
                assert (dimension.width, dimension.hidden, dimension.outlineLevel) == (
                    actual.width,
                    actual.hidden,
                    actual.outlineLevel,
                ), (name, col, "column dimensions")
        for row, dimension in old.row_dimensions.items():
            if row not in allowed_heights[name]:
                actual = new.row_dimensions[row]
                assert (dimension.height, dimension.hidden, dimension.outlineLevel) == (
                    actual.height,
                    actual.hidden,
                    actual.outlineLevel,
                ), (name, row, "row dimensions")
        changed = []
        for row in range(1, max(old.max_row, new.max_row) + 1):
            for col in range(1, max(old.max_column, new.max_column) + 1):
                a, b = old.cell(row, col), new.cell(row, col)
                if a.value != b.value:
                    assert a.coordinate in allowed[name], (
                        name,
                        a.coordinate,
                        a.value,
                        b.value,
                    )
                    changed.append(a.coordinate)
                if a.data_type == "f":
                    formulas += 1
                    if not (name == "06_Defects" and a.coordinate == "B5"):
                        assert a.value == b.value, (name, a.coordinate, "formula text")
                if a.coordinate not in allowed[name]:
                    assert style(a) == style(b), (name, a.coordinate, "unrelated style")
        changes[name] = changed
    assert tables == 8
    assert list(before.defined_names) == list(after.defined_names)
    for record_id in before.defined_names:
        assert etree.tostring(
            before.defined_names[record_id].to_tree()
        ) == etree.tostring(after.defined_names[record_id].to_tree())
    old, new = before["02_UAT_Cases"], after["02_UAT_Cases"]
    for row in range(8, 59):
        assert new.cell(row, 1).value == old.cell(row, 1).value
        for col in range(12, 22):
            assert new.cell(row, col).value == old.cell(row, col).value, (
                row,
                col,
                "older case history",
            )
    expected_date = datetime.fromisoformat(facts["collected_at"]).astimezone(
        ZoneInfo("America/Bogota")
    )
    expected_date = expected_date.replace(tzinfo=None, microsecond=0)
    for record_id, item in cases_by_id.items():
        row = row_for(old, record_id)
        assert new[f"J{row}"].value == item["status"]
        actual = (
            f"Backend: {item['backend_status']}. {text(item['actual'])}"
            if item.get("backend_status")
            else text(item["actual"])
        )
        assert new[f"H{row}"].value == actual
        assert new[f"I{row}"].value == text(item["evidence"])
        for col, key in ((3, "preconditions"), (6, "steps"), (7, "expected")):
            expected = (
                text(item[key])
                if item.get(key) not in (None, "")
                else old.cell(row, col).value
            )
            assert new.cell(row, col).value == expected
        assert [new.cell(row, col).value for col in (22, 23, 24)] == [
            old.cell(row, col).value for col in (10, 8, 9)
        ]
        assert isinstance(new[f"Y{row}"].value, datetime)
        assert abs((new[f"Y{row}"].value - expected_date).total_seconds()) < 1
        assert new[f"Y{row}"].number_format == "yyyy-mm-dd hh:mm"
        assert new[f"Z{row}"].value == text(item["evidence"])
    old, new = before["06_Defects"], after["06_Defects"]
    assert new["A23"].value == "UAT-ENV-001", "ENV record must retain row 23"
    for record_id in DEFECT_IDS:
        row = row_for(old, record_id)
        item = defects_by_id[record_id]
        assert [new.cell(row, col).value for col in (5, 6, 7, 9)] == [
            text(item["actual"]),
            item["status"],
            text(item["action"]),
            text(item["evidence"]),
        ]
        assert json.loads(new[f"R{row}"].value) == {
            "actual": old[f"E{row}"].value,
            "status": old[f"F{row}"].value,
            "action": old[f"G{row}"].value,
            "evidence": old[f"I{row}"].value,
        }
    for row in range(8, 24):
        for col in (16, 17):
            assert new.cell(row, col).value == old.cell(row, col).value, (
                row,
                col,
                "older defect history",
            )
    for record_id, row in NEW_DEFECT_ROWS.items():
        item = defects_by_id[record_id]
        assert item["isNew"] is True
        assert [new.cell(row, col).value for col in range(1, 10)] == [
            item["id"],
            text(item["classification"]),
            text(item["severity"]),
            text(item["title"]),
            text(item["actual"]),
            item["status"],
            text(item["action"]),
            text(
                item.get("owner")
                or "Investment Decision y Project Workspace Governance"
            ),
            text(item["evidence"]),
        ]
        assert json.loads(new[f"R{row}"].value)["previous_record"] is None
        assert json.loads(new[f"R{row}"].value)["new_defect"] == record_id
    assert (
        new["B5"].value == "=COUNTA(A8:A25)" and cached["06_Defects"]["B5"].value == 18
    )
    old, new = before["08_Release_Criteria"], after["08_Release_Criteria"]
    for row in range(8, 38):
        for col in (11, 12):
            assert new.cell(row, col).value == old.cell(row, col).value, (
                row,
                col,
                "older release history",
            )
    for record_id, item in criteria_by_id.items():
        row = row_for(old, record_id)
        assert [new.cell(row, col).value for col in (4, 5, 6)] == [
            text(item["actual"]),
            item["status"],
            text(item["evidence"]),
        ]
        assert cached["08_Release_Criteria"][f"J{row}"].value == item["status"]
        assert json.loads(new[f"M{row}"].value) == {
            "criterion": record_id,
            "prior_condition": old[f"D{row}"].value,
            "prior_status": old[f"E{row}"].value,
            "prior_evidence": old[f"F{row}"].value,
            "prior_calculated_result": before_cached["08_Release_Criteria"][
                f"J{row}"
            ].value,
        }
    signoffs = 0
    for row in range(8, 38):
        if new[f"C{row}"].value == "BUSINESS_SIGNOFF":
            signoffs += 1
            assert all(new.cell(row, col).value is None for col in range(5, 10))
            assert cached["08_Release_Criteria"][f"J{row}"].value == "BLOCKED"
    assert signoffs == 15
    errors = [
        f"{sheet.title}!{cell.coordinate}:{cell.value}"
        for sheet in cached
        for row in sheet
        for cell in row
        if cell.data_type == "e"
    ]
    assert not errors, errors
    untouched = [name for name in before.sheetnames if not allowed[name]]
    assert len(untouched) == 5 and all(not changes[name] for name in untouched)
    result = {
        "status": "PASS_READ_ONLY_SCOPE_VERIFICATION",
        "source": str(SOURCE),
        "source_sha256": digest(SOURCE),
        "output": str(OUTPUT),
        "output_sha256": digest(OUTPUT),
        "facts_sha256": digest(FACTS),
        "sheets": after.sheetnames,
        "untouched_tabs_values_styles_and_native_features": untouched,
        "native_tables_retained": tables,
        "defect_table_extension_only": "A7:I23 to A7:I25",
        "native_validations_panes_conditional_formats_and_defined_names_preserved": True,
        "formula_texts_preserved_except_direct_defect_count_range": formulas,
        "cached_formula_errors": len(errors),
        "case_ids_and_older_history_rows_preserved": 51,
        "new_history": {"cases": "V:Z", "defects": "R", "release": "M"},
        "human_signoff_areas_blank": signoffs,
        "unsigned_ACCEPTED_remains_BLOCKED": True,
        "updated_cell_values": changes,
        "visual_review": "PASS_MAIN_REVIEW_OF_SAVED_RENDER_PANELS"
        if args.visual_reviewed
        else "PENDING_MAIN_REVIEW_OF_SAVED_RENDER_PANELS",
    }
    QA.mkdir(parents=True, exist_ok=True)
    (QA / "scope-verification.json").write_text(
        json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    print(json.dumps(result, ensure_ascii=True))


if __name__ == "__main__":
    main()
