"""Read-only artifact verification and derived final QA manifest.

This does not author Word/Excel content, execute UAT, or grant business signoff.
Manual visual coverage below is recorded only after actual page inspections.
"""

from __future__ import annotations

import hashlib
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "docs/validation/strategic-final-uat"
QA = ROOT / "outputs/.qa/strategic-final-uat"
OUT = ROOT / "outputs/strategic-final-uat-20261014"
PREFIX = "P&Pmis_Strategic_Macroprocess_"
DOCS = {
    "Business_Process": ("business", 13),
    "User_Guide": ("user", 23),
    "Admin_Guide": ("admin", 29),
    "UAT_Plan_and_Results": ("uat", 32),
    "Final_Release_Closeout": ("closeout", 11),
}
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    results = json.loads((DATA / "release-results.json").read_text(encoding="utf-8-sig"))
    assert results["release_decision"] == "NO_GO"
    assert results["release_state"] == "NOT_VALIDATED"
    assert results["signoffs"] == []
    files = []
    for suffix, (kind, expected_pages) in DOCS.items():
        path = OUT / f"{PREFIX}{suffix}.docx"
        with ZipFile(path) as archive:
            assert archive.testzip() is None
            styles = etree.fromstring(archive.read("word/styles.xml"))
            title = styles.xpath('//w:style[@w:styleId="Title"]', namespaces=NS)[0]
            assert not title.xpath("./w:pPr/w:pBdr", namespaces=NS)
            assert title.xpath("./w:rPr/w:color/@w:val", namespaces=NS) == ["000000"]
        pages = list((QA / "word" / kind).glob("page-*.png"))
        assert len(pages) == expected_pages, (kind, len(pages))
        accessibility = json.loads((QA / "word" / kind / "a11y.json").read_text())
        assert not any(accessibility["counts"].values()), accessibility["counts"]
        files.append({"file": path.name, "bytes": path.stat().st_size, "sha256": sha(path),
                      "rendered_pages": expected_pages, "visual_qa": "PASS all pages",
                      "accessibility_findings": accessibility["counts"]})

    workbook = OUT / f"{PREFIX}UAT_Traceability.xlsx"
    cached = load_workbook(workbook, data_only=True, read_only=True)
    structural = load_workbook(workbook, data_only=False)
    assert len(cached.sheetnames) == 8
    cases = cached["02_UAT_Cases"]
    counts = Counter(cases.cell(row, 10).value for row in range(8, 59))
    assert counts == {"PASS": 5, "BLOCKED": 46}, counts
    release = cached["08_Release_Criteria"]
    signatures = [row for row in range(8, 38) if release.cell(row, 3).value == "BUSINESS_SIGNOFF"]
    assert len(signatures) == 15
    for row in signatures:
        assert all(release.cell(row, col).value in (None, "") for col in range(5, 10))
        assert release.cell(row, 10).value == "BLOCKED"
    assert release["G5"].value == "NO_GO"
    assert release["J5"].value == "NO_GO"
    assert all(release.cell(row, 6).value for row in range(8, 23))
    defects = cached["06_Defects"]
    assert all(defects.cell(row, 5).value for row in range(8, 19))
    errors = [(sheet.title, cell.coordinate, cell.value) for sheet in cached for row in sheet
              for cell in row if cell.data_type == "e"]
    assert errors == [], errors
    for sheet in structural:
        assert sheet.freeze_panes == "C8"
        assert len(sheet.tables) == 1
        assert next(iter(sheet.tables.values())).autoFilter is not None
    files.append({"file": workbook.name, "bytes": workbook.stat().st_size,
                  "sha256": sha(workbook), "sheets": cached.sheetnames,
                  "uat_counts": dict(counts), "unsigned_areas": 15,
                  "cached_formula_errors": errors, "visual_qa": "PASS sheet panels; affected panels rechecked",
                  "formula_behavior": "Artifact Tool recalculation verified unsigned ACCEPTED remains BLOCKED"})
    cached.close()
    structural.close()
    assert {path.name for path in OUT.iterdir() if path.is_file()} == {row["file"] for row in files}
    regression = json.loads((QA / "backend-reconciled-final.json").read_text())
    assert regression["unique_cases"] == 960
    assert regression["latest_counts"] == {"PASS": 944, "FAIL": 1, "SKIP": 15}
    manifest = {
        "verified_at": datetime.now(timezone.utc).isoformat(),
        "release_decision": "NO_GO", "release_state": "NOT_VALIDATED",
        "next_increment": results["next_increment"], "human_signoff": "PENDING",
        "source_results_sha256": sha(DATA / "release-results.json"),
        "backend_reconciliation_sha256": sha(QA / "backend-reconciled-final.json"),
        "backend_regression_counts": regression["latest_counts"],
        "all_word_pages_visually_reviewed": 108,
        "review_method": "Full page PNG inspection; unchanged pages after final narrow edits verified by raster SHA256",
        "artifacts": files,
    }
    destination = DATA / "final-artifact-manifest.json"
    destination.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({"outputs": len(files), "word_pages": 108, "formula_errors": 0,
                      "unsigned_areas": 15, "decision": "NO_GO", "manifest": str(destination)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
