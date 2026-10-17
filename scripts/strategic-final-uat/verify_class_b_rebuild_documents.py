"""Read-only structural, source-fidelity and rendered-prefix verification."""

import argparse
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree
from PIL import Image, ImageChops

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "outputs/gate07j-controlled-class-b-20261015"
OUT = ROOT / "outputs/gate07j-controlled-class-b-20261016"
QA = ROOT / "outputs/.qa/gate07j-controlled-class-b-rebuild"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
PREFIX = "P&Pmis_Strategic_Macroprocess_"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def pixels(path):
    with Image.open(path) as picture:
        return picture.size, hashlib.sha256(picture.tobytes()).hexdigest()


def same_historical_body(before, after):
    with (
        Image.open(before).convert("RGB") as old,
        Image.open(after).convert("RGB") as new,
    ):
        assert old.size == new.size
        difference = ImageChops.difference(old, new).getbbox()
        if difference is not None:
            assert difference[1] >= int(old.height * 0.95), difference
        return True


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--visual-reviewed", action="store_true")
    args = parser.parse_args()
    results = []
    for suffix, kind in (
        ("Final_Release_Closeout", "closeout"),
        ("UAT_Plan_and_Results", "uat"),
    ):
        name = PREFIX + suffix + ".docx"
        source, output = SOURCE / name, OUT / name
        with ZipFile(source) as before, ZipFile(output) as after:
            assert after.testzip() is None
            assert set(before.namelist()) == set(after.namelist())
            changed = [
                part
                for part in before.namelist()
                if before.read(part) != after.read(part)
            ]
            footer_parts = {
                p
                for p in before.namelist()
                if p.startswith("word/footer") and p.endswith(".xml")
            }
            assert (
                set(changed)
                <= {"word/document.xml", "docProps/core.xml"} | footer_parts
            ), changed
            for part in footer_parts:
                expected = (
                    before.read(part)
                    .replace(b"14 octubre 2026", b"16 octubre 2026")
                    .replace(b"15 octubre 2026", b"16 octubre 2026")
                )
                assert after.read(part) == expected, (
                    part,
                    "Only the footer revision date may change",
                )
            old_body = etree.fromstring(before.read("word/document.xml")).find(
                "w:body", NS
            )
            new_body = etree.fromstring(after.read("word/document.xml")).find(
                "w:body", NS
            )
            assert etree.tostring(old_body[-1]) == etree.tostring(new_body[-1])
            assert not new_body.xpath(
                ".//w:ins | .//w:del | .//w:commentRangeStart", namespaces=NS
            )
            if kind == "uat":
                for index, element in enumerate(old_body[:-1]):
                    assert etree.tostring(element) == etree.tostring(new_body[index])
        document = Document(output)
        if kind == "closeout":
            assert (
                len([p for p in document.paragraphs if p.style.name == "Heading 1"])
                == 34
            )
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        assert ":codex" not in text and "turn0search" not in text
        pages = list((QA / f"word-{kind}/pages-final").glob("page-*.png"))
        assert pages, "Canonical renderer page images required"
        if kind == "uat":
            assert len(pages) > 37
            for number in range(1, 38):
                previous = (
                    ROOT
                    / f"outputs/.qa/gate07j-controlled-class-b/word-uat/pages-final/page-{number}.png"
                )
                current = QA / f"word-uat/pages-final/page-{number}.png"
                assert same_historical_body(previous, current), number
        results.append(
            {
                "name": name,
                "path": str(output),
                "source_sha256": sha(source),
                "sha256": sha(output),
                "pages": len(pages),
                "changed_package_parts": changed,
                "historical_prefix_pages_body_pixel_identical": 37
                if kind == "uat"
                else 0,
                "intentional_footer_revision_date": "16 octubre 2026",
                "all_pages_visually_reviewed": args.visual_reviewed,
            }
        )
    destination = QA / "document-verification.json"
    destination.write_text(
        json.dumps({"artifacts": results}, indent=2), encoding="utf-8"
    )
    print(json.dumps({"status": "PASS", "artifacts": results}))


if __name__ == "__main__":
    main()
