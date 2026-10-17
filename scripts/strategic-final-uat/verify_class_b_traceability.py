"""Read-only verification of the narrow blocked-preflight XLSX successor."""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
NAME = "P&Pmis_Strategic_Macroprocess_UAT_Traceability.xlsx"
SOURCE = ROOT / "outputs/strategic-uat-remediation-20261015" / NAME
OUTPUT = ROOT / "outputs/gate07j-controlled-class-b-20261015" / NAME
QA = ROOT / "outputs/.qa/gate07j-controlled-class-b/workbook"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def style(cell: object) -> tuple[str, ...]:
    return tuple(
        str(getattr(cell, name))
        for name in (
            "font",
            "fill",
            "border",
            "alignment",
            "number_format",
            "protection",
        )
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--visual-reviewed", action="store_true")
    args = parser.parse_args()
    before = openpyxl.load_workbook(SOURCE)
    after = openpyxl.load_workbook(OUTPUT)
    cached = openpyxl.load_workbook(OUTPUT, data_only=True)
    assert before.sheetnames == after.sheetnames
    allowed: dict[str, set[str]] = {name: set() for name in before.sheetnames}
    for row in range(35, 40):
        allowed["02_UAT_Cases"].update(f"{col}{row}" for col in "HIJQRSTU")
    allowed["02_UAT_Cases"].update(f"{col}7" for col in "QRSTU")
    allowed["06_Defects"].update(
        [
            "B5",
            "E10",
            "F10",
            "G10",
            "I10",
            "I21",
            "I22",
            "Q7",
            "Q10",
            "Q21",
            "Q22",
            "Q23",
        ]
    )
    allowed["06_Defects"].update(f"{col}23" for col in "ABCDEFGHI")
    for row in (11, 14, 21):
        allowed["08_Release_Criteria"].update(f"{col}{row}" for col in "EFL")
    allowed["08_Release_Criteria"].add("L7")
    changes: dict[str, list[str]] = {}
    formula_count = 0
    for name in before.sheetnames:
        old, new = before[name], after[name]
        assert old.freeze_panes == new.freeze_panes, name
        assert list(old.tables) == list(new.tables), name
        assert old.sheet_view.showGridLines == new.sheet_view.showGridLines, name
        assert old.merged_cells == new.merged_cells, name
        for table_name in old.tables:
            expected = "A7:I23" if name == "06_Defects" else old.tables[table_name].ref
            assert new.tables[table_name].ref == expected, (name, table_name)
            assert (
                old.tables[table_name].tableStyleInfo
                == new.tables[table_name].tableStyleInfo
            )
        changed = []
        for row in range(1, max(old.max_row, new.max_row) + 1):
            for col in range(1, max(old.max_column, new.max_column) + 1):
                a, b = old.cell(row, col), new.cell(row, col)
                if a.value != b.value:
                    assert a.coordinate in allowed[name], (
                        name,
                        a.coordinate,
                        a.value,
                        b.value,
                    )
                    changed.append(a.coordinate)
                if a.data_type == "f":
                    formula_count += 1
                    if not (name == "06_Defects" and a.coordinate == "B5"):
                        assert a.value == b.value, (name, a.coordinate)
                if a.coordinate not in allowed[name]:
                    assert style(a) == style(b), (
                        name,
                        a.coordinate,
                        "unrelated cell style",
                    )
        changes[name] = changed
    # Native validation XML is compared structurally, not object string repr.
    from lxml import etree

    for name in before.sheetnames:
        for feature in ("data_validations", "conditional_formatting"):
            if feature == "data_validations":
                assert etree.tostring(
                    before[name].data_validations.to_tree()
                ) == etree.tostring(after[name].data_validations.to_tree()), (
                    name,
                    feature,
                )
            else:
                old_rules = before[name].conditional_formatting
                new_rules = after[name].conditional_formatting
                assert list(old_rules) == list(new_rules), (name, feature)
                for key in old_rules:
                    assert [
                        etree.tostring(rule.to_tree()) for rule in old_rules[key]
                    ] == [etree.tostring(rule.to_tree()) for rule in new_rules[key]], (
                        name,
                        str(key),
                    )
    cases = after["02_UAT_Cases"]
    assert [cases.cell(row, 1).value for row in range(8, 59)] == [
        before["02_UAT_Cases"].cell(row, 1).value for row in range(8, 59)
    ]
    for row in range(8, 59):
        for col in range(12, 17):
            assert (
                cases.cell(row, col).value
                == before["02_UAT_Cases"].cell(row, col).value
            )
    for row in range(35, 40):
        old = before["02_UAT_Cases"]
        assert cases[f"J{row}"].value == "BLOCKED"
        assert "BLOCKED_UAT_DATABASE_NOT_AVAILABLE" in cases[f"H{row}"].value
        assert [cases.cell(row, col).value for col in (17, 18, 19)] == [
            old.cell(row, col).value for col in (10, 8, 9)
        ]
        assert isinstance(cases[f"T{row}"].value, datetime)
        assert cases[f"T{row}"].number_format == "yyyy-mm-dd hh:mm"
    defects = after["06_Defects"]
    for row in (10, 21, 22):
        prior = json.loads(defects[f"Q{row}"].value)
        old = before["06_Defects"]
        expected_prior = {
            "actual": old[f"E{row}"].value,
            "status": old[f"F{row}"].value,
            "evidence": old[f"I{row}"].value,
        }
        if row == 10:
            expected_prior["workaround"] = old[f"G{row}"].value
        assert prior == expected_prior
        assert defects[f"P{row}"].value == old[f"P{row}"].value
        assert defects[f"F{row}"].value == old[f"F{row}"].value
    assert defects["A23"].value == "UAT-ENV-001"
    assert defects["B23"].value == "VALIDATION_ENVIRONMENT"
    assert "S1/S2" not in defects["C23"].value
    assert defects["B5"].value == "=COUNTA(A8:A23)"
    assert cached["06_Defects"]["B5"].value == 16
    release = after["08_Release_Criteria"]
    signoffs = 0
    for row in range(8, 38):
        assert (
            release[f"K{row}"].value == before["08_Release_Criteria"][f"K{row}"].value
        )
        if release[f"C{row}"].value == "BUSINESS_SIGNOFF":
            signoffs += 1
            assert all(release.cell(row, col).value is None for col in range(5, 10))
            assert cached["08_Release_Criteria"][f"J{row}"].value == "BLOCKED"
    assert signoffs == 15
    for row in (11, 14, 21):
        assert release[f"E{row}"].value == "BLOCKED"
        assert cached["08_Release_Criteria"][f"J{row}"].value == "BLOCKED"
        prior = json.loads(release[f"L{row}"].value)
        old = before["08_Release_Criteria"]
        assert prior["prior_status"] == old[f"E{row}"].value
        assert prior["prior_evidence"] == old[f"F{row}"].value
    errors = [
        f"{sheet.title}!{cell.coordinate}:{cell.value}"
        for sheet in cached
        for row in sheet
        for cell in row
        if cell.data_type == "e"
    ]
    assert not errors, errors
    result = {
        "status": "PASS_READ_ONLY_SCOPE_VERIFICATION",
        "source": str(SOURCE),
        "source_sha256": digest(SOURCE),
        "output": str(OUTPUT),
        "output_sha256": digest(OUTPUT),
        "sheets": after.sheetnames,
        "all_native_tables_retained": 8,
        "native_defects_table_extended_only": "A7:I22 to A7:I23",
        "native_validations_panes_and_conditional_formats_preserved": True,
        "formula_texts_preserved_except_direct_defect_count_range": formula_count,
        "cached_formula_errors": len(errors),
        "case_ids_and_older_history_rows_preserved": 51,
        "human_signoff_areas_blank": signoffs,
        "new_business_tests": "NOT_RUN",
        "new_browser": "NOT_RUN",
        "new_business_state_commands": 0,
        "updated_cell_values": changes,
        "visual_review": (
            "PASS_MAIN_REVIEW_OF_SAVED_RENDER_PANELS"
            if args.visual_reviewed
            else "PENDING_MAIN_REVIEW_OF_SAVED_RENDER_PANELS"
        ),
    }
    QA.mkdir(parents=True, exist_ok=True)
    (QA / "scope-verification.json").write_text(
        json.dumps(result, indent=2), encoding="utf-8"
    )
    print(json.dumps(result, ensure_ascii=True))


if __name__ == "__main__":
    main()
