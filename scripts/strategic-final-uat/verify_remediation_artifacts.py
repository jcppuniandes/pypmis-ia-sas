"""Read-only structural/fidelity checks for the current scoped UAT revision."""

import argparse
import hashlib
import json
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree
from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "outputs/strategic-uat-remediation-20261015"
SOURCE = ROOT / "outputs/strategic-final-uat-20261014"
QA = ROOT / "outputs/.qa/strategic-uat-remediation"
DATA = ROOT / "docs/validation/strategic-uat-remediation"
PREFIX = "P&Pmis_Strategic_Macroprocess_"
DOCS = {
    "Admin_Guide": ("admin", 29),
    "UAT_Plan_and_Results": ("uat", 32),
    "Final_Release_Closeout": ("closeout", 0),
}


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def pixels(path):
    with Image.open(path) as image:
        return image.size, hashlib.sha256(image.tobytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--visual-reviewed",
        action="store_true",
        help="Record root's completed visual review of the final rendered pages and workbook panels",
    )
    args = parser.parse_args()
    facts = json.loads((DATA / "release-results.json").read_text(encoding="utf-8"))
    before, after = facts["operational_before"], facts["operational_after"]
    assert before == after, "Operational preservation check differs"
    assert facts["release_status"] == "STRATEGIC_MACROPROCESS_NOT_VALIDATED"
    artifacts = []
    changed_pages = {}
    for suffix, (kind, prefix_pages) in DOCS.items():
        name = PREFIX + suffix + ".docx"
        path, retained = OUT / name, SOURCE / name
        with ZipFile(path) as new, ZipFile(retained) as old:
            assert new.testzip() is None
            changed_parts = sorted(
                part for part in old.namelist() if old.read(part) != new.read(part)
            )
            allowed = {"word/document.xml", "docProps/core.xml"}
            if suffix == "Final_Release_Closeout":
                allowed.add("word/footer1.xml")
            assert set(changed_parts) <= allowed, (suffix, changed_parts)
            assert set(old.namelist()) == set(new.namelist())
            if prefix_pages:
                ns = {
                    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
                }
                old_body = etree.fromstring(old.read("word/document.xml")).find(
                    "w:body", ns
                )
                new_body = etree.fromstring(new.read("word/document.xml")).find(
                    "w:body", ns
                )
                for index, element in enumerate(old_body[:-1]):
                    assert etree.tostring(element) == etree.tostring(new_body[index]), (
                        suffix,
                        index,
                    )
        render_directory = QA / f"word-{suffix}" / "pages-final-v4"
        if not render_directory.is_dir():
            render_directory = QA / f"word-{suffix}" / "pages-final-v3"
        if not render_directory.is_dir():
            render_directory = QA / f"word-{suffix}" / "pages-final-v2"
        if not render_directory.is_dir():
            render_directory = QA / f"word-{suffix}" / "pages-final"
        if not render_directory.is_dir():
            render_directory = QA / f"word-{suffix}" / "pages"
        pages = sorted(
            render_directory.glob("page-*.png"),
            key=lambda path: int(path.stem.split("-")[-1]),
        )
        assert pages
        changed = []
        reused = []
        for index, page in enumerate(pages, 1):
            old_png = (
                ROOT / f"outputs/.qa/strategic-final-uat/word/{kind}/page-{index}.png"
            )
            if (
                index <= prefix_pages
                and old_png.is_file()
                and pixels(page) == pixels(old_png)
            ):
                reused.append(index)
            else:
                changed.append(index)
        changed_pages[suffix] = {
            "new_or_changed_pages": changed,
            "previously_reviewed_identical_pages": reused,
        }
        a11y = json.loads((QA / f"word-{suffix}" / "a11y.json").read_text())
        assert not any(a11y["counts"].values())
        artifacts.append(
            {
                "name": name,
                "sha256": sha(path),
                "bytes": path.stat().st_size,
                "pages": len(pages),
                "changed_package_parts": changed_parts,
                "a11y": a11y["counts"],
            }
        )
    closeout = Document(OUT / (PREFIX + "Final_Release_Closeout.docx"))
    headings = [p.text for p in closeout.paragraphs if p.style.name == "Heading 1"]
    assert len(headings) == 27
    workbook_name = PREFIX + "UAT_Traceability.xlsx"
    old = load_workbook(SOURCE / workbook_name, data_only=False)
    old_cached = load_workbook(SOURCE / workbook_name, data_only=True)
    new = load_workbook(OUT / workbook_name, data_only=False)
    cached = load_workbook(OUT / workbook_name, data_only=True)
    assert old.sheetnames == new.sheetnames and len(new.sheetnames) == 8
    for name in old.sheetnames:
        assert old[name].freeze_panes == new[name].freeze_panes
        assert len(old[name].tables) == len(new[name].tables) == 1
        assert old[name].data_validations == new[name].data_validations
        for row in old[name]:
            for cell in row:
                if cell.data_type == "f":
                    if name == "06_Defects" and cell.coordinate == "B5":
                        continue  # Directly affected record-count range extension.
                    assert cell.value == new[name][cell.coordinate].value, (
                        name,
                        cell.coordinate,
                    )
    for name in (
        "01_Process_Matrix",
        "03_Admin_User_Matrix",
        "04_Navigation_Matrix",
        "05_Business_Rules",
        "07_Code_Cleanup",
    ):
        assert [[cell.value for cell in row] for row in old[name]] == [
            [cell.value for cell in row] for row in new[name]
        ], name
    for row in range(8, 59):
        assert (
            new["02_UAT_Cases"].cell(row, 12).value
            == old["02_UAT_Cases"].cell(row, 10).value
        )
        assert (
            new["02_UAT_Cases"].cell(row, 13).value
            == old["02_UAT_Cases"].cell(row, 8).value
        )
        assert (
            new["02_UAT_Cases"].cell(row, 14).value
            == old["02_UAT_Cases"].cell(row, 9).value
        )
    errors = [
        (s.title, c.coordinate, c.value)
        for s in cached
        for row in s
        for c in row
        if c.data_type == "e"
    ]
    assert errors == [], errors
    current_defects = {
        cached["06_Defects"].cell(row, 1).value: row for row in range(8, 23)
    }
    for defect in facts["defects"]:
        row = current_defects[defect["id"]]
        assert cached["06_Defects"].cell(row, 5).value == defect["actual"]
        assert cached["06_Defects"].cell(row, 6).value == defect["status"]
        assert cached["06_Defects"].cell(row, 9).value == defect["evidence"]
    release = cached["08_Release_Criteria"]
    prior_release = old_cached["08_Release_Criteria"]
    for row in range(8, 38):
        criterion = release.cell(row, 1).value
        if release.cell(row, 11).value:
            # Adjacent history records the prior calculated state, not its
            # active formula. Native formulas are checked separately above.
            assert json.loads(release.cell(row, 11).value) == [
                prior_release.cell(row, col).value for col in range(1, 11)
            ]
        if criterion == "RC11":
            assert release.cell(row, 2).value == (
                "Blueprint canónico y ajuste USER autorizado"
            )
            assert "autorizado" in release.cell(row, 4).value
            assert "intacto" not in release.cell(row, 4).value
            assert facts["app_change"]["actual_sha256"] in release.cell(row, 6).value
        if criterion == "RC07":
            assert release.cell(row, 5).value == "BLOCKED"
    unsigned = []
    for row in range(8, 38):
        if cached["08_Release_Criteria"].cell(row, 3).value == "BUSINESS_SIGNOFF":
            assert all(
                cached["08_Release_Criteria"].cell(row, col).value in (None, "")
                for col in range(5, 10)
            )
            assert cached["08_Release_Criteria"].cell(row, 10).value == "BLOCKED"
            unsigned.append(row)
    assert len(unsigned) == 15
    artifacts.append(
        {
            "name": workbook_name,
            "sha256": sha(OUT / workbook_name),
            "sheets": cached.sheetnames,
            "uat_current_counts": dict(
                Counter(
                    cached["02_UAT_Cases"].cell(row, 10).value for row in range(8, 59)
                )
            ),
            "history_preserved_cases": 51,
            "unsigned_areas": len(unsigned),
            "cached_formula_errors": errors,
        }
    )
    for value in facts["prior_documents"].values():
        assert sha(Path(value["path"])) == value["sha256"]
    manifest = {
        "release_status": facts["release_status"],
        "operational_preservation": "IDENTICAL_READ_ONLY_SNAPSHOTS",
        "word_page_review_queue": changed_pages,
        "artifacts": artifacts,
        "visual_review": (
            "PASS_ROOT_VISUAL_REVIEW_OF_FINAL_WORD_PAGES_AND_CHANGED_WORKBOOK_PANELS"
            if args.visual_reviewed
            else "PENDING_ROOT_INSPECTION_OF_NEW_OR_CHANGED_PAGES"
        ),
    }
    (DATA / "artifact-structural-verification.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
