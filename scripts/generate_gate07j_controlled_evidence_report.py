"""Versioned reports of the blocked same-chain Gate07J prerequisite.

Run the document edit marker for two DOCX immediately before --documents.
Never starts databases, constructs business records, or changes prior outputs.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from copy import deepcopy
from pathlib import Path

import generate_strategic_final_uat_reports as style
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs/strategic-uat-remediation-20261015"
OUT = ROOT / "outputs/gate07j-controlled-class-b-20261015"
DATA = ROOT / "docs/validation/gate07j-controlled-class-b"
OLD_DATA = ROOT / "docs/validation/strategic-uat-remediation"
OLD_QA = ROOT / "outputs/.qa/strategic-uat-remediation"
NAMES = style.FILES


def load(path):
    return json.loads(path.read_text(encoding="utf-8"))


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def prose(text):
    """Keep compact identifiers readable in authored prose, not source records."""
    text = str(text)
    text = re.sub(
        r"\b(Tenant|Project|Enterprise|Portfolio|Gate|BC|ID|Case|assessment|revisión|Preparación|versión|política|CORE|fuentes)(\d+)\b",
        r"\1 \2",
        text,
    )
    text = re.sub(
        r"\b(son|con|no|del|de|las|los|era|y|tiene|mantienen)(?=\d)", r"\1 ", text
    )
    text = re.sub(r"\bvalidez(180|365)d\b", r"validez de \1 días", text)
    text = re.sub(r"\b(180|365)d\b", r"\1 días", text)
    text = re.sub(r",(?=\d)", ", ", text)
    text = text.replace("EACH_PROJECT16", "EACH_PROJECT para Project 16")
    text = text.replace("PYP-FFA373", "PYP-FFA 373")
    text = text.replace("únicos:898", "únicos: 898")
    return text


def p(doc, text):
    style.p(doc, prose(text))


def table(doc, headers, rows, widths):
    style.table(doc, headers, [[prose(v) for v in row] for row in rows], widths)
    for cell in doc.tables[-1].rows[0].cells:
        cell.paragraphs[0].paragraph_format.keep_with_next = True


def cases():
    return [
        [
            "A27",
            "Gate readiness",
            "BLOCKED",
            "No se verificó la preparación en PostgreSQL. No se presentaron ni validaron referencias.",
        ],
        [
            "A28",
            "Consistencia de fuentes",
            "BLOCKED",
            "No se recalcularon hashes ni vigencia. La captura anterior se conserva como historia.",
        ],
        [
            "A29",
            "Freeze",
            "BLOCKED",
            "No ejecutado. Sin preparación actual verificada no se autoriza congelar.",
        ],
        [
            "A30",
            "FID",
            "BLOCKED",
            "No ejecutado. No hay paquete frozen nuevo ni decisión nueva demostrada.",
        ],
        [
            "A31",
            "Post FID",
            "BLOCKED",
            "No ejecutado. Tampoco inicialización ni activación en este incremento.",
        ],
    ]


def reconciliation():
    diagnostic = load(OLD_QA / "fresh-project16-gate1-missing-evidence-diagnosis.json")
    readiness = load(OLD_QA / "fresh-project16-gate1-readiness-current.json")[
        "readiness"
    ]
    rows = []
    # The diagnostic's top-level list deliberately contains only four blockers.
    # Reconcile the full captured effective matrix, not that filtered view.
    for item in diagnostic["gate"]["source_snapshot"]["evidence_policy"][
        "effective_matrix"
    ]:
        if not item["required"]:
            continue
        coverage = readiness["evidence_coverage"].get(item["evidence_type"], {})
        rows.append(
            {
                "requirement_key": item["requirement_key"],
                "class": item["classification"],
                "required": item["required"],
                "evidence_type": item["evidence_type"],
                "provider_key": item["provider_key"],
                "origin": "PROVIDER" if item["provider_key"] else "MANUAL_CONTROLLED",
                "blocking": item["blocking"],
                "scope": item["scope"],
                "historical_status": coverage.get("status", "MISSING"),
                "historical_coverage_percent": coverage.get("coverage_percent", 0.0),
                "current_reference_status": "NOT_VERIFIABLE",
                "current_validation_status": "NOT_VERIFIABLE",
                "currentness": "NOT_VERIFIABLE",
                "current_coverage": None,
                "reason": "BLOCKED_UAT_DATABASE_NOT_AVAILABLE",
            }
        )
    assert len(rows) == 19
    return rows


def technical(preflight, rows):
    doc = Document(SOURCE / NAMES["closeout"])
    body = doc._element.body
    logo = deepcopy(body[0]) if body[0].xpath(".//w:drawing") else None
    for element in list(body):
        if element.tag != qn("w:sectPr"):
            body.remove(element)
    if logo is not None:
        body.insert(0, logo)
    doc.add_paragraph("Informe técnico de evidencias controladas Gate07J", "Title")
    p(
        doc,
        "P&Pmis Construction AI. Capital Portfolio & Investment Governance. Corte del 15 de octubre de 2026.",
    )
    p(
        doc,
        "La continuación A27 a A31 está bloqueada antes de presentar evidencias: la base PostgreSQL de la UAT seleccionada no está disponible y su almacenamiento era temporal. El estado técnico permanece STRATEGIC_MACROPROCESS_NOT_VALIDATED / NO_GO. El código y los archivos de evidencia anteriores permanecen guardados; no sustituyen la base transaccional requerida para continuar la misma cadena.",
    )
    sections = [
        (
            "1 Resumen ejecutivo",
            "No se ejecutaron comandos de negocio ni se reconstruyó la cadena. Se verificó el estado del entorno, se auditó la arquitectura existente y se registró el bloqueo de prerequisito. La aceptación del incremento y el cierre de UAT-S2-003 siguen pendientes.",
        ),
        (
            "2 Línea base",
            f"Rama {preflight['commands']['git_branch']['stdout']}. Commit {preflight['commands']['git_head']['stdout']}. El árbol conserva modificaciones previas sin commit; no se descartaron ni se publicaron. Alembic heads devuelve {preflight['commands']['alembic_heads']['stdout']}. No se creó migración. El diff previo queda preservado en la evidencia del preflight.",
        ),
        (
            "3 Identidad de la cadena seleccionada",
            "La identidad esperada es Tenant4, fresh-copper-debf223b50, Enterprise13, Portfolio15 y Project16. Cycle, Demand, Allocation, Idea, Analysis, Proposal, Strategic Gate y Creation Request tienen ID4 en la última captura. Business Case4 estaba en revisión n.º2, ID6, COMPLETED/READY. Estos datos son históricos: no se han podido confirmar en una base UAT actual.",
        ),
        (
            "4 Bloqueo actual de Gate07J",
            "El prerequisito falla con BLOCKED_UAT_DATABASE_NOT_AVAILABLE. El contenedor fresh-db está detenido con salida255; PGDATA estaba en tmpfs, sin montajes persistentes. Docker no encuentra PG_VERSION. No se inició el contenedor porque su arranque aprovisionaría una base vacía. No se localizó un respaldo de esta cadena en las rutas revisadas. El dump operativo Wave3A del 7 de octubre no corresponde a esta UAT y no se restauró.",
        ),
        (
            "5 Arquitectura de evidencias reutilizable",
            "El producto ya admite referencias controladas. POST /api/v1/investment-gate-preparations/{id}/evidence registra SUBMITTED; POST /api/v1/investment-gate-evidence/{id}/validate registra VALID o INVALID. Los comandos ready-for-review y freeze usan la misma preparación y el If-Match de su revision_version. Se reutilizarán estos endpoints, sus hashes, el ámbito y sus permisos; no se necesita un subsistema paralelo.",
        ),
        (
            "6 Referencia Class B de permisos",
            "La captura anterior identifica permit-readiness / PERMIT_READINESS, Class B, provider_key=null, EACH_PROJECT16, validez180d y validación independiente. Permit assessment2 estaba READY con3/3 permisos aprobados; eso no valida automáticamente una referencia controlada. No se presentó ni validó una referencia nueva en este incremento.",
        ),
        (
            "7 Referencia Class B de contratación",
            "contracting-strategy-readiness / CONTRACTING_STRATEGY_READINESS conserva Class B, provider_key=null y validez365d en la captura anterior. El snapshot Commercial puede ser material referenciado si se recupera y se comprueba su identidad; su estado COMPLETED no reemplaza la validación independiente de la evidencia.",
        ),
        (
            "8 Referencia Class B del plan de ejecución",
            "execution-plan-readiness / EXECUTION_PLAN_READINESS requiere referencia controlada EACH_PROJECT con validez365d. El nuevo alcance permite un documento representativo rotulado UAT CONTROLLED EVIDENCE; no equivale a un plan real aprobado. Su registro depende de verificar primero la identidad actual del Project y del paquete. No se creó un nuevo dominio de planificación.",
        ),
        (
            "9 Referencia Class B de madurez del cronograma",
            "schedule-maturity / SCHEDULE_MATURITY requiere referencia controlada con validez180d. Se debe identificar la revisión, fecha de datos, alcance, responsable, referencia y hash de una evaluación UAT. Schedule Source Readiness no constituye por sí sola esa evidencia. No se creó un nuevo módulo ni se registró una evidencia en una cadena distinta.",
        ),
        (
            "10 Validación independiente",
            "No se ejecutaron validaciones. El submitter y el validador deberán ser actores autorizados distintos, con permisos de agregar y validar evidencia. Se deben comprobar el perfil publicado y su configuración SoD efectivos. La inspección estática de condiciones configurables no demuestra por sí sola un defecto runtime ni habilita autoaprobación.",
        ),
        (
            "11 Vigencia y caducidad",
            "Las ventanas esperadas son180d para permisos y cronograma,365d para contratación y ejecución. No se asignaron fechas futuras ni validaciones retroactivas. Los negativos de caducidad y currentness no se ejecutaron porque la base seleccionada no está disponible; no se registran como PASS ni como SKIP de una suite inexistente.",
        ),
        (
            "12 Cobertura de cada Project",
            "La última captura exigía EACH_PROJECT y100% de cobertura para un único Project. Su cobertura histórica era0/1. La cobertura actual es desconocida, no0% confirmado. No se usó evidencia de Portfolio ni promedios para satisfacer las cuatro filas.",
        ),
        (
            "13 Readiness A27",
            "A27 permanece BLOCKED como caso formal. No se re-resolvió la preparación ni se recalculó readiness. Gate1 DRAFT/BLOCKED/source_current=true es el último estado capturado, no una observación viva posterior a la interrupción.",
        ),
        (
            "14 Conciliación de los diecinueve requisitos",
            "La tabla siguiente conserva los19 requisitos required de la captura anterior. La política era371/revisión1 y el perfil PYP-FFA373/revisión1. Los estados y porcentajes son exclusivamente históricos. Para todas las filas, referencia, validación, currentness y cobertura actuales son NOT_VERIFIABLE por el mismo bloqueo de base; no se infieren fallos de proveedor.",
        ),
        (
            "15 Consistencia de fuentes A28",
            "La comparación nueva entre Business Case, FIN, Benefits, Commercial, FEL, Risk, Funding y Permit no se ejecutó. Los hashes históricos permanecen archivados. No se reconstruyeron fuentes desde esos JSON ni se inyectaron hashes para aparentar currentness.",
        ),
        (
            "16 Paquete frozen A29",
            "No hay freeze nuevo demostrado. Se requiere preparación READY_FOR_REVIEW con readiness READY y fuentes actuales antes del comando del revisor. El incremento no añadió GO, HOLD, RECYCLE o KILL a Gate07J.",
        ),
        (
            "17 Inmutabilidad del paquete",
            "El negativo de sucesor controlado en un entorno separado queda NOT_RUN. La conservación de archivos históricos no equivale a una prueba de inmutabilidad transaccional de un paquete frozen nuevo.",
        ),
        (
            "18 Decisión FID A30",
            "No se creó Investment Decision. Sin paquete frozen verificado no se puede ejecutar Gate07K. No se materializó un GO ni se cambió un estado de decisión directamente en PostgreSQL.",
        ),
        (
            "19 Indicadores FID",
            "is_fid, investment_authorized, full_funds_authorized y can_activate_project no se obtuvieron de una decisión nueva. No se asignan valores esperados como si fueran resultados. Class B válida, QRA y readiness no autorizan inversión automáticamente.",
        ),
        (
            "20 Gobierno posterior a FID A31",
            "No se creó ni refrescó elegibilidad Post-FID. Su entrada debe ser la decisión PYP-FFA GO nueva y verificada de esta cadena. El resultado de routing no se infiere del estado histórico pending del Project.",
        ),
        (
            "21 Inicialización y activación",
            "No se ejecutaron. Cuando exista FID y la elegibilidad lo permita, inicialización y activación deberán conservar comandos y actores autorizados separados. No se inició Project Delivery operativo.",
        ),
        (
            "22 Disposición UAT S2 003",
            "OPEN. A27 a A31 siguen BLOCKED por la imposibilidad de verificar la misma cadena. No se atribuye cierre al código existente, a la auditoría estática ni al Business Case histórico completado.",
        ),
        (
            "23 Prueba FID histórica",
            "La prueba test_new_pyp_ffa_fid_path_full_stack_revalidation conserva su fallo histórico BUSINESS_CASE_NOT_READY. No se excluyó, no se marcó xfail y no se actualizó la fixture antes de acreditar su reemplazo moderno. La clasificación definitiva se difiere hasta completar A27 a A31.",
        ),
        (
            "24 Disposición UAT S3 005",
            "OPEN_SCOPED_PRESENTATION. La última navegación mostró etiquetas NOT_AVAILABLE y Funding STALE que no correspondían al backend capturado. No se modificaron estados de Funding ni Business Case. La corrección y su regresión de pantalla siguen pendientes de un entorno verificable.",
        ),
        (
            "25 Disposición UAT S4 006",
            "OPEN_SCOPED_LIFECYCLE_LABEL. La aprobación histórica de Proposal no se reescribió. La separación visual entre lifecycle histórico y elegibilidad viva de reingreso no se implementó ni se revalidó en este incremento.",
        ),
        (
            "26 Regresión global",
            "No se lanzó una suite nueva. La última invocación global archivada tiene914 casos únicos:898 PASS,1 FAIL,0 ERROR y15 SKIP. Son resultados anteriores, no conteos de este incremento ni una regresión verde. CORE72 PASS y fuentes76 PASS también se conservan solo como historia.",
        ),
        (
            "27 Verificación PostgreSQL",
            "La verificación de disponibilidad sí se ejecutó; la verificación de Tenant4/Project16 dentro de PostgreSQL no fue posible. La causa es el almacenamiento temporal de la UAT detenida. No se conectó la continuación a la base operativa para eludir este prerequisito.",
        ),
        (
            "28 Verificación de navegador",
            "NOT_RUN. No se ejecutó Chromium contra un backend UAT inexistente ni se reutilizaron capturas de otro tenant como prueba nueva. DOWNSTREAM_CHAIN_PREFLIGHT_NO_FID permanece como hallazgo histórico, no como resultado de una ejecución actual.",
        ),
        (
            "29 Seguridad y segregación de funciones",
            "No hubo bootstrap de permisos ni cambios en políticas. Las guardas de tenant, Project, submitter/validator, preparer/reviewer y autoridad Gate07K deben reejecutarse después de recuperar el entorno. Se auditó su arquitectura en código; esa inspección no reemplaza las pruebas negativas requeridas.",
        ),
        (
            "30 Preservación del Project operativo 14",
            "La captura READ ONLY coincide con la referencia previa: Project14 PRJ-PPMIS, external_key PRJ-PYP-PPMIS, active y versión1; hash de fila ef67b78a4d4902d2dc6db3efb017b80b. Se mantienen31 Workspaces,95 configuraciones publicadas, hashes de las5 políticas de evidencia y0 QRA runs. Se repetirá la captura de cierre sin DML.",
        ),
        (
            "31 Hallazgos restantes",
            "Se registra UAT-ENV-001 como bloqueo de entorno, no como nuevo defecto funcional confirmado. UAT-S2-003 y las observaciones de presentación permanecen abiertos. Las observaciones estáticas de validez, SoD e idempotencia requieren reproducción y no se convierten en defectos S1/S2 demostrados.",
        ),
        (
            "32 Decisión técnica de liberación",
            "STRATEGIC_MACROPROCESS_NOT_VALIDATED / NO_GO. No se alcanza STRATEGIC_MACROPROCESS_TECHNICALLY_READY_UAT_SIGNOFF_PENDING: faltan las pruebas de evidencia, freeze, FID, Post-FID, navegador y regresión nueva.",
        ),
        (
            "33 Aceptación humana",
            "PENDING_NOT_ASSUMED. No se completaron nombres, fechas ni evidencias de firma. Las quince áreas de aceptación del workbook siguen sin firma y BLOCKED.",
        ),
        (
            "34 Recomendación siguiente",
            "REMEDIATE THE SMALLEST REMAINING STRATEGIC UAT BLOCKER. Recuperar la base exacta desde un respaldo verificable. Si no existe, obtener autorización expresa para reconstruir la cadena mediante los ciclos reales en almacenamiento persistente, registrando que ya no sería la continuación de la misma base original.",
        ),
    ]
    assert len(sections) == 34
    for heading, text in sections:
        style.h(doc, heading)
        p(doc, text)
        if heading.startswith("3 "):
            table(
                doc,
                ["Identidad", "Valor esperado"],
                [
                    ["Tenant", "4 / fresh-copper-debf223b50"],
                    ["Portfolio y Project", "15 /16"],
                    ["Gate y Business Case", "Preparación1; BC4 revisión2 ID6"],
                    ["Estado actual", "NOT_VERIFIABLE"],
                ],
                [2.3, 4.76],
            )
        if heading.startswith("4 "):
            table(
                doc,
                ["Comprobación", "Observación actual"],
                [
                    [
                        "Contenedor",
                        preflight["finding"]["container_state"] + "; exit255",
                    ],
                    ["Almacenamiento", "tmpfs; sin volumen persistente"],
                    ["PG_VERSION", "No disponible mediante docker cp"],
                    ["Respaldo seleccionado", "No localizado en las rutas revisadas"],
                    ["Escrituras de negocio", "0"],
                ],
                [2.3, 4.76],
            )
        if heading.startswith("14 "):
            table(
                doc,
                [
                    "Requisito",
                    "Clase",
                    "Origen",
                    "Estado previo",
                    "Cobertura previa",
                    "Bloqueante",
                ],
                [
                    [
                        r["requirement_key"],
                        r["class"],
                        r["provider_key"] or "Manual controlado",
                        r["historical_status"],
                        f"{r['historical_coverage_percent']:g}%",
                        "Sí" if r["blocking"] else "No",
                    ]
                    for r in rows
                ],
                [2.0, 0.55, 1.65, 1.0, 0.96, 0.9],
            )
    doc.core_properties.title = (
        "Evidencias controladas Gate07J y bloqueo de continuidad"
    )
    doc.core_properties.revision = 3
    return doc


def uat(preflight):
    doc = Document(SOURCE / NAMES["uat"])
    doc.add_page_break()
    style.h(doc, "Preflight de continuidad de evidencias Gate07J")
    p(
        doc,
        "Corte del15 de octubre de2026. La base UAT seleccionada no está disponible. Los resultados de las páginas anteriores se conservan como historia; no acreditan la vigencia actual de Tenant4/Project16. No se ejecutaron nuevos casos de negocio ni navegador.",
    )
    p(
        doc,
        "Precondición común: recuperar y verificar la base exacta, el tenant, Project16, perfil373 y política371. Pasos nuevos ejecutados: inspección del contenedor, tmpfs, PG_VERSION, respaldos locales, Git y Alembic. Resultado real: BLOCKED_UAT_DATABASE_NOT_AVAILABLE. Evidencia: preflight.json y baseline-working-tree.diff de este incremento.",
    )
    table(
        doc,
        ["Caso", "Proceso", "Estado", "Resultado actual"],
        cases(),
        [0.5, 1.6, 0.8, 4.16],
    )
    p(
        doc,
        "Resultado esperado al recuperar el entorno: referencias Class B SUBMITTED y VALID por actores distintos, cobertura1/1, readiness real, freeze independiente, decisión PYP-FFA GO/FID y Post-FID de la misma cadena. Ninguno se declara logrado en este corte.",
    )
    p(
        doc,
        "UAT-S2-003 continúa OPEN. UAT-S3-005 y UAT-S4-006 mantienen sus disposiciones abiertas sin corrección nueva. El fallo FID histórico no se ocultó. Nueva regresión y navegador: NOT_RUN; no hay conteos PASS/FAIL/SKIP actuales que sumar.",
    )
    p(
        doc,
        "Las firmas humanas permanecen vacías. El Project operativo14 y sus políticas se preservan mediante lectura sin DML. No se reinicia upstream sin autorización explícita ni se escriben estados terminales desde JSON.",
    )
    doc.core_properties.revision = 3
    return doc


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("--documents", action="store_true")
    args = parser.parse_args()
    preflight = load(DATA / "preflight.json")
    assert preflight["status"] == "BLOCKED_UAT_DATABASE_NOT_AVAILABLE"
    manifest = load(OLD_DATA / "artifact-structural-verification.json")
    hashes = {r["name"]: r["sha256"] for r in manifest["artifacts"]}
    for name in (
        NAMES["closeout"],
        NAMES["uat"],
        "P&Pmis_Strategic_Macroprocess_UAT_Traceability.xlsx",
    ):
        assert sha(SOURCE / name) == hashes[name]
    rows = reconciliation()
    (DATA / "requirement-reconciliation.json").write_text(
        json.dumps(rows, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    if args.documents:
        OUT.mkdir(parents=True, exist_ok=True)
        for kind, doc in (
            ("closeout", technical(preflight, rows)),
            ("uat", uat(preflight)),
        ):
            doc.save(OUT / NAMES[kind])
        for name, expected in hashes.items():
            assert sha(SOURCE / name) == expected
        print(
            json.dumps(
                {
                    "outputs": [str(OUT / NAMES[k]) for k in ("closeout", "uat")],
                    "status": preflight["status"],
                }
            )
        )


if __name__ == "__main__":
    main()
