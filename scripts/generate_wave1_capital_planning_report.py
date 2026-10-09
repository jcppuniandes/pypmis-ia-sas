from __future__ import annotations

import re
from pathlib import Path

import generate_cpig_gap_review_report as base
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "71-wave1-strategic-capital-demand-release.md"
QA = ROOT / "outputs" / ".qa" / "wave1-capital-planning"
DIAGRAM_DIR = QA / "diagrams"
REPORT_DATE = "2026-09-30"
REPORT_DATE_DISPLAY = "30 de septiembre de 2026"
OUTPUT = (
    ROOT
    / "outputs"
    / f"Informe_Tecnico_Wave1_Strategic_Capital_Demand_Capital_Planning_{REPORT_DATE}.docx"
)
LOGO = ROOT / "frontend" / "public" / "pypmis-construction-ai-logo.png"


DIAGRAMS = [
    {
        "name": "01-wave1-lineage.png",
        "title": "Trazabilidad canonica de Wave 1",
        "rows": [
            [
                "Strategic Objective",
                "Investment Theme",
                "Planning Cycle",
                "Capital Demand",
            ],
            ["Project Proposal", "Business Analysis", "Idea", "Originate Idea"],
        ],
        "edges": [
            (0, 0, 0, 1),
            (0, 1, 0, 2),
            (0, 2, 0, 3),
            (0, 3, 1, 3),
            (1, 3, 1, 2),
            (1, 2, 1, 1),
            (1, 1, 1, 0),
        ],
        "alt": (
            "Trazabilidad desde Strategic Objective e Investment Theme hasta Capital Demand, "
            "Idea, Business Analysis y Project Proposal."
        ),
    },
    {
        "name": "02-wave1-boundaries.png",
        "title": "Fronteras de autoridad de Capital Demand",
        "rows": [
            ["Capital Demand"],
            ["Idea", "Project", "Portfolio Membership"],
            ["Budget", "Funding", "FIN"],
        ],
        "edges": [],
        "alt": (
            "Capital Demand como identidad independiente de Idea, Project, Portfolio Membership, "
            "Budget, Funding y FIN."
        ),
    },
    {
        "name": "03-wave1-lifecycles.png",
        "title": "Ciclos de vida gobernados",
        "rows": [
            ["Cycle DRAFT", "Cycle OPEN", "Cycle IN REVIEW", "Cycle CLOSED"],
            ["Demand DRAFT", "Demand SUBMITTED", "Demand IN REVIEW", "Demand ENDORSED"],
        ],
        "edges": [
            (0, 0, 0, 1),
            (0, 1, 0, 2),
            (0, 2, 0, 3),
            (1, 0, 1, 1),
            (1, 1, 1, 2),
            (1, 2, 1, 3),
        ],
        "alt": (
            "Capital Planning Cycle avanza de Draft a Closed y Strategic Capital Demand de Draft "
            "a Endorsed mediante transiciones controladas."
        ),
    },
    {
        "name": "04-wave1-admin-user.png",
        "title": "Separacion entre ADMIN y USER",
        "rows": [
            ["ADMIN configuration", "Published effective policy"],
            ["USER Planning Cycles", "USER Capital Demands", "Originate Idea"],
        ],
        "edges": [
            (0, 0, 0, 1),
            (0, 1, 1, 0),
            (0, 1, 1, 1),
            (1, 1, 1, 2),
        ],
        "alt": (
            "ADMIN publica configuracion efectiva que gobierna los registros USER sin mezclar "
            "permisos de configuracion y transaccion."
        ),
    },
]


def read_release_status() -> str:
    source_text = SOURCE.read_text(encoding="utf-8")
    match = re.search(
        r"^\*\*Estado (?:actual|final(?: de release)?):\*\*\s+(.+?)\.?$",
        source_text,
        re.MULTILINE,
    )
    if match is None:
        raise ValueError(f"Missing release status in {SOURCE}")
    return match.group(1).strip().strip("`")


def add_cover(document: Document, release_status: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(32)
    if LOGO.exists():
        shape = paragraph.add_run().add_picture(str(LOGO), width=Inches(1.3))
        base.set_alt_text(
            shape, "P&Pmis Construction AI", "Logotipo azul de P&Pmis Construction AI."
        )

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(24)
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run(
        "Informe técnico Wave 1 Strategic Capital Demand and Capital Planning"
    )
    base.set_run_font(run, size=22, bold=True, color=base.BLACK)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("P&Pmis Construction AI")
    base.set_run_font(run, size=13, color=base.BLACK)

    metadata = document.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_before = Pt(24)
    base.add_inline(
        metadata,
        f"Capital Portfolio and Investment Governance | {REPORT_DATE_DISPLAY}",
        size=10,
    )

    status_label = document.add_paragraph()
    status_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status_label.paragraph_format.space_before = Pt(44)
    run = status_label.add_run("ESTADO DE VALIDACIÓN")
    base.set_run_font(run, size=8, bold=True, color=base.MUTED)

    status = document.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = status.add_run(release_status)
    base.set_run_font(run, size=11, bold=True, color=base.TEAL)
    document.add_page_break()


def add_document_control(document: Document, release_status: str) -> None:
    document.add_heading("Control del documento", level=1)
    base.add_table(
        document,
        [
            ["Campo", "Valor"],
            ["Producto", "P&Pmis Construction AI"],
            ["Macroproceso", "Capital Portfolio and Investment Governance"],
            [
                "Incremento",
                "Wave 1 Strategic Capital Demand and Capital Planning Foundation",
            ],
            ["Rama y commit de baseline", "main / 9cb0969"],
            ["Migracion objetivo", "20260929_0062 sobre 20260916_0061"],
            ["Fecha de corte", REPORT_DATE],
            ["Estado al corte", release_status],
            [
                "Fuente tecnica",
                "Codigo, migracion, pruebas y validacion del repositorio",
            ],
        ],
    )
    document.add_heading("Lectura ejecutiva", level=1)
    paragraph = document.add_paragraph()
    base.add_inline(
        paragraph,
        (
            "Wave 1 introduce una identidad gobernada para la necesidad de capital antes de seleccionar una "
            "solucion. El incremento conserva Idea, Business Analysis, Project Proposal, Portfolio, Budget, "
            "Funding y FIN como autoridades separadas, y origina Ideas solo mediante una accion explicita, "
            "idempotente y auditable."
        ),
    )
    document.add_heading("Contenido", level=1)
    for item in (
        "Baseline, ADR y posicion del proceso",
        "Diseño de Theme, Planning Cycle y Capital Demand",
        "Contrato Demand to Idea y compatibilidad historica",
        "ADMIN, USER, seguridad, concurrencia y persistencia",
        "API, navegacion, pruebas, brechas y cierre de release",
    ):
        paragraph = document.add_paragraph(style="List Bullet")
        base.add_inline(paragraph, item)
    document.add_page_break()


def configure_metadata(document: Document) -> None:
    core = document.core_properties
    core.title = "Informe técnico Wave 1 Strategic Capital Demand and Capital Planning"
    core.subject = "Implementación y validación de la primera ola CPIG"
    core.author = "P&Pmis Construction AI"
    core.keywords = "CPIG, Capital Demand, Capital Planning, Idea, governance"
    core.comments = "Informe tecnico de implementacion y validacion."


def add_headers_and_footers(document: Document) -> None:
    for section in document.sections:
        section.different_first_page_header_footer = True
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = header.add_run("P&Pmis Construction AI  |  Wave 1 Capital Planning")
        base.set_run_font(run, size=7.5, bold=True, color=base.BLACK)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run(f"Confidencial | {REPORT_DATE} | Página ")
        base.set_run_font(run, size=7.5, color=base.MUTED)
        base.add_field(footer, "PAGE")
        run = footer.add_run(" de ")
        base.set_run_font(run, size=7.5, color=base.MUTED)
        base.add_field(footer, "NUMPAGES")


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    QA.mkdir(parents=True, exist_ok=True)
    DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)

    base.SOURCE = SOURCE
    base.QA = QA
    base.DIAGRAM_DIR = DIAGRAM_DIR
    base.DIAGRAMS = DIAGRAMS
    diagram_paths = [base.render_diagram(spec) for spec in DIAGRAMS]
    release_status = read_release_status()

    document = Document()
    base.configure_document(document)
    configure_metadata(document)
    base.set_language(document)
    add_cover(document, release_status)
    add_document_control(document, release_status)
    base.add_source_content(
        document,
        diagram_paths,
        page_break_heading_prefixes=("4. ",),
    )
    add_headers_and_footers(document)
    base.update_field_settings(document)
    document.save(OUTPUT)
    print(f"created={OUTPUT}")
    print(f"diagrams={len(diagram_paths)}")
    print(f"source={SOURCE}")


if __name__ == "__main__":
    main()
