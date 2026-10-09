from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(r"C:\Users\Ricardo\Documents\GitHub\pypmis-ia-sas")
OUTPUT = ROOT / "outputs" / "Informe_Tecnico_Gate07A_BA_Business_Analysis_2026-09-10.docx"
LOGO = ROOT / "frontend" / "public" / "pypmis-construction-ai-logo.png"
EVIDENCE = ROOT / "docs" / "validation" / "gate07a-ba"

NAVY = "112733"
TEAL = "087F80"
PALE = "EAF5F4"
GRAY = "F3F6F7"
LINE = "D5E0E4"
TEXT = "17242C"
MUTED = "5A6C76"
WHITE = "FFFFFF"
GREEN = "18794E"
AMBER = "A05A00"


def shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), fill)


def margins(cell, value: int = 90) -> None:
    props = cell._tc.get_or_add_tcPr()
    node = props.first_child_found_in("w:tcMar")
    if node is None:
        node = OxmlElement("w:tcMar")
        props.append(node)
    for name in ("top", "start", "bottom", "end"):
        item = node.find(qn(f"w:{name}"))
        if item is None:
            item = OxmlElement(f"w:{name}")
            node.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def repeat_header(row) -> None:
    props = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    props.append(node)


def no_split(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def font(run, size: float = 9.5, bold: bool = False, color: str = TEXT, name: str = "Aptos") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, code, separate, value, end])


def heading(doc, text: str, level: int = 1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def paragraph(doc, text: str = "", *, lead: str | None = None, style: str | None = None):
    item = doc.add_paragraph(style=style)
    item.paragraph_format.space_after = Pt(5)
    item.paragraph_format.line_spacing = 1.06
    if lead and text.startswith(lead):
        first = item.add_run(lead)
        font(first, 10, True)
        rest = item.add_run(text[len(lead) :])
        font(rest, 10)
    else:
        font(item.add_run(text), 10)
    return item


def bullets(doc, values: list[str]) -> None:
    for value in values:
        item = paragraph(doc, value, style="List Bullet")
        item.paragraph_format.space_after = Pt(2.5)


def numbered(doc, values: list[str]) -> None:
    for value in values:
        item = paragraph(doc, value, style="List Number")
        item.paragraph_format.space_after = Pt(2.5)


def table(doc, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    result = doc.add_table(rows=1, cols=len(headers))
    result.style = "Table Grid"
    result.autofit = False
    repeat_header(result.rows[0])
    no_split(result.rows[0])
    for index, label in enumerate(headers):
        cell = result.rows[0].cells[index]
        shade(cell, NAVY)
        margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            cell.width = Inches(widths[index])
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        font(cell.paragraphs[0].add_run(label), 8.5, True, WHITE)
    for row_index, values in enumerate(rows):
        row = result.add_row()
        no_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            shade(cell, WHITE if row_index % 2 == 0 else GRAY)
            margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if widths:
                cell.width = Inches(widths[index])
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            font(cell.paragraphs[0].add_run(str(value)), 8.2)
    paragraph(doc)
    return result


def status_table(doc, rows: list[list[str]]) -> None:
    result = table(doc, ["Puerta", "Resultado", "Evidencia"], rows, [2.0, 0.95, 3.75])
    for row in result.rows[1:]:
        value = row.cells[1].text.strip()
        shade(row.cells[1], PALE if value == "PASS" else "FFF3E3")
        for run in row.cells[1].paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(GREEN if value == "PASS" else AMBER)


def picture(doc, path: Path, caption: str) -> None:
    if not path.exists():
        return
    with Image.open(path) as source:
        aspect = source.width / source.height
    width = min(6.85, 7.0 * aspect)
    item = doc.add_paragraph()
    item.alignment = WD_ALIGN_PARAGRAPH.CENTER
    item.add_run().add_picture(str(path), width=Inches(width))
    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.space_after = Pt(8)
    run = label.add_run(caption)
    font(run, 8.3, False, MUTED)
    run.italic = True


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.68)
section.bottom_margin = Inches(0.62)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Aptos")
normal.font.size = Pt(10)
normal.font.color.rgb = RGBColor.from_string(TEXT)
for name, size, color in (
    ("Title", 27, NAVY),
    ("Heading 1", 17, NAVY),
    ("Heading 2", 12.5, TEAL),
    ("Heading 3", 10.5, NAVY),
):
    style = doc.styles[name]
    style.font.name = "Aptos Display"
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(9)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.keep_with_next = True

for item in doc.sections:
    header = item.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(header.add_run("P&Pmis Construction AI  |  Gate 07A-BA"), 8, False, MUTED)
    footer = item.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run("Informe técnico  •  Página "), 8, False, MUTED)
    field(footer, "PAGE")

# Portada
if LOGO.exists():
    logo = doc.add_paragraph()
    logo.add_run().add_picture(str(LOGO), width=Inches(0.78))
doc.add_paragraph().paragraph_format.space_after = Pt(15)
title = doc.add_paragraph(style="Title")
title.add_run("Informe técnico")
subtitle = doc.add_paragraph()
font(subtitle.add_run("Gate 07A-BA · Business Analysis & Options Governance"), 19, True, TEAL)
subtitle.paragraph_format.space_after = Pt(20)
paragraph(doc, "Producto: P&Pmis Construction AI", lead="Producto:")
paragraph(doc, "Repositorio: pypmis-ia-sas", lead="Repositorio:")
paragraph(doc, "Rama de trabajo: main (cambios locales no confirmados)", lead="Rama de trabajo:")
paragraph(doc, "Línea base: 9cb0969", lead="Línea base:")
paragraph(doc, f"Fecha: {date(2026, 9, 10).strftime('%d/%m/%Y')}", lead="Fecha:")
doc.add_paragraph().paragraph_format.space_after = Pt(20)
intro = doc.add_paragraph()
intro.paragraph_format.line_spacing = 1.16
font(
    intro.add_run(
        "Se implementó un expediente gobernado de Business Analysis entre Idea Lifecycle y Project Proposal, "
        "con alternativas comparables, recomendación trazable, readiness determinístico, segregación de funciones "
        "y configuración versionada por ámbito empresarial."
    ),
    12,
)
doc.add_page_break()

heading(doc, "Contenido", 1)
numbered(
    doc,
    [
        "Resumen ejecutivo y alcance",
        "Arquitectura e integración",
        "Modelo funcional y lifecycle",
        "Persistencia, configuración y migración",
        "API, seguridad y concurrencia",
        "Experiencia USER MODE y ADMIN MODE",
        "Compatibilidad y no duplicación",
        "Validación y evidencia",
        "Inventario de cambios",
        "Riesgos, límites y cierre",
    ],
)
heading(doc, "1. Resumen ejecutivo y alcance", 1)
paragraph(
    doc,
    "Gate 07A-BA transforma una Idea aceptada en un análisis de negocio completo antes de habilitar Project Proposal. "
    "La implementación es aditiva: consume la Idea, su evaluación aceptada, objetivos estratégicos, workspace y "
    "configuración efectiva; no reemplaza ninguna capacidad estratégica existente.",
)
bullets(
    doc,
    [
        "USER MODE administra expedientes vivos, colas, análisis de opciones, recomendación, readiness e historial.",
        "ADMIN MODE administra solamente política reutilizable, catálogos, matrices, mapping, numeración, herencia y SoD.",
        "CAPITAL_OWNER exige BA por defecto; modelos alternos y datos legacy preservan compatibilidad.",
        "Project Proposal sólo se habilita para un BA COMPLETED, vigente y sin blockers.",
        "Los registros completos son inmutables; los cambios materiales crean una revisión vinculada.",
    ],
)
heading(doc, "Resultado de negocio", 2)
paragraph(
    doc,
    "La recomendación deja de ser texto implícito dentro de la propuesta: queda sustentada por necesidad, drivers, "
    "stakeholders, estado actual, valor, requisitos y una matriz de opciones con trazabilidad de fuente y configuración.",
)
heading(doc, "Fuera de alcance respetado", 2)
paragraph(
    doc,
    "No se implementaron Five Case Business Case, modelo financiero completo, selección de portafolio, FID, activación "
    "del Project Workspace, ejecución de beneficios ni automatización por agentes.",
)

doc.add_page_break()
heading(doc, "2. Arquitectura e integración", 1)
table(
    doc,
    ["Capa", "Implementación", "Reutilización"],
    [
        ["Dominio", "BusinessAnalysis + snapshots append-only", "Idea, IdeaEvaluation, Strategic Objective, EnterpriseWorkspace"],
        ["Aplicación", "Servicio de lifecycle, readiness, configuración y mapping", "AdminConfiguration, SecurityEvent, permisos empresariales"],
        ["API", "Rutas USER y ADMIN separadas", "JWT tenant context, ETag, If-Match"],
        ["Frontend", "Workspace transaccional y consola de configuración", "App shell, navegación estratégica, CompactModuleHeader"],
        ["Downstream", "Linaje BA opcional y mapping en Project Proposal", "Lifecycle y snapshot existentes de Proposal"],
    ],
    [1.05, 2.75, 2.9],
)
heading(doc, "Secuencia gobernada", 2)
flow = doc.add_paragraph()
flow.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(flow.add_run("IDEA ACCEPTED  →  BUSINESS ANALYSIS  →  READY_FOR_PROJECT_PROPOSAL  →  PROJECT PROPOSAL"), 10.5, True, TEAL)
paragraph(
    doc,
    "El hash de la Idea y evaluación aceptada se congela al crear el BA. Al completar se incorpora el hash de "
    "configuración y del contenido revisado. Project Proposal conserva las referencias y hashes sin mutar el BA.",
)
heading(doc, "Antes / después", 2)
table(
    doc,
    ["Momento", "Proceso"],
    [
        ["Antes", "Idea ACCEPTED → Project Proposal (recomendación y alternativas implícitas)"],
        [
            "Después",
            "Idea ACCEPTED → Business Analysis DRAFT/REVIEW/COMPLETED → READY_FOR_PROJECT_PROPOSAL → Project Proposal",
        ],
    ],
    [1.15, 5.55],
)
heading(doc, "Autoridad por bounded context", 2)
table(
    doc,
    ["Contexto", "Es autoridad de", "No hace"],
    [
        ["Idea Lifecycle", "Aceptación y evaluación de la Idea", "No evalúa alternativas BA"],
        ["Business Analysis", "Problema, alternativas, recomendación y readiness", "No decide FID ni activa proyectos"],
        ["Project Proposal", "Propuesta derivada y su evaluación", "No reabre ni recalcula el BA"],
        ["Administration", "Política versionada", "No opera expedientes vivos"],
    ],
    [1.35, 2.75, 2.6],
)

doc.add_page_break()
heading(doc, "3. Modelo funcional y lifecycle", 1)
table(
    doc,
    ["Estado", "Entrada", "Salida / control"],
    [
        ["DRAFT", "Idea ACCEPTED elegible", "Edición estructurada; ETag"],
        ["SUBMITTED", "Readiness de contenido sin blockers", "Contenido congelado para revisión"],
        ["IN_REVIEW", "Revisor distinto al autor", "Return o Complete"],
        ["RETURNED", "Observación con motivo", "Vuelve a autor; nueva edición controlada"],
        ["COMPLETED", "Revisión independiente y fuente vigente", "Inmutable; output READY_FOR_PROJECT_PROPOSAL"],
        ["SUPERSEDED / VOIDED", "Gobierno excepcional", "Retención histórica; fuera de colas activas"],
    ],
    [1.55, 2.45, 2.7],
)
heading(doc, "Contenido estructurado", 2)
bullets(
    doc,
    [
        "Need / opportunity: statement, background, contexto, urgencia y consecuencia de no actuar.",
        "Drivers y Strategic Alignment referencian catálogos y objetivos existentes.",
        "Stakeholders, current state, outcomes y value dimensions soportan el caso preliminar.",
        "Assumptions, constraints y business requirements registran los límites del análisis.",
        "Cada opción incluye tipo, descripción, alcance, beneficios, costo/tiempo preliminar, riesgos, dependencias, factibilidad, encaje estratégico y ratings.",
        "La recomendación conserva opción, rationale, trade-offs, condiciones y preocupaciones pendientes.",
    ],
)
heading(doc, "Readiness determinístico", 2)
paragraph(
    doc,
    "Los blockers se calculan en backend. Incluyen fuente o configuración obsoleta, secciones obligatorias ausentes, "
    "opciones insuficientes, tipos exigidos ausentes, criterios incompletos, recomendación no seleccionada y asuntos "
    "bloqueantes no aceptados. La interfaz sólo representa el resultado del contrato.",
)

doc.add_page_break()
heading(doc, "4. Persistencia, configuración y migración", 1)
table(
    doc,
    ["Objeto", "Responsabilidad", "Controles"],
    [
        ["business_analyses", "Estado vigente y contenido normalizado", "Tenant scope, número único, idempotencia, ETag, hashes"],
        ["business_analysis_revisions", "Snapshot por cambio/transición", "Append-only, número único por BA, snapshot hash"],
        ["project_proposals", "Linaje opcional BA", "Id/hash/config revision nullable para legacy"],
        ["admin_configurations", "Política y catálogos BA", "DRAFT/PUBLISHED/ARCHIVED, hash, revisión"],
        ["security_events", "Auditoría de comandos", "Actor, outcome, estados y metadata"],
    ],
    [1.6, 2.65, 2.45],
)
heading(doc, "Resolución de configuración", 2)
paragraph(doc, "Tenant → Enterprise → Business Unit → Portfolio. La configuración aplicable más cercana prevalece y puede heredarse a descendientes. El expediente congela id, revisión y hash efectivos.")
heading(doc, "Migración Alembic", 2)
paragraph(
    doc,
    "La revisión 20260910_0059 sucede a 20260902_0058, crea las dos tablas BA, índices y relaciones, y añade cinco "
    "columnas nullable a Project Proposal. El downgrade elimina primero el linaje y luego las tablas BA. Los ciclos "
    "reales PostgreSQL y SQLite verificaron constraints, índices, nulabilidad y reversibilidad.",
)
heading(doc, "Inmutabilidad", 2)
paragraph(
    doc,
    "Listeners del modelo rechazan update/delete sobre snapshots y rechazan la mutación de un Business Analysis cuyo "
    "estado persistido ya era COMPLETED. Un cambio material se representa mediante supersedes_analysis_id y nueva identidad.",
)

doc.add_page_break()
heading(doc, "5. API, seguridad y concurrencia", 1)
table(
    doc,
    ["Familia", "Operaciones", "Garantía"],
    [
        ["USER query", "options, list, get, readiness, history", "Filtrado tenant/workspace/actor"],
        ["USER command", "create, update, submit, start-review, return, complete, void, new-revision", "RBAC, SoD, ETag, idempotencia"],
        ["ADMIN query", "configurations, preview", "Preview persisted=false"],
        ["ADMIN command", "clone, update, publish, archive", "Publicación inmutable y versionada"],
    ],
    [1.25, 3.2, 2.25],
)
heading(doc, "Permisos", 2)
bullets(
    doc,
    [
        "business_analysis.read / create / edit / submit / review",
        "business_analysis.admin.configure / admin.publish",
        "Roles dedicados author, reviewer, viewer y config_admin; organization_admin mantiene acceso gobernado.",
        "La política SoD por defecto impide que el autor inicie o cierre su propia revisión.",
    ],
)
heading(doc, "Control optimista e idempotencia", 2)
paragraph(
    doc,
    "PUT y transiciones comparan If-Match con revision_version y devuelven 412 ante conflicto. POST de creación y nueva "
    "revisión deduplican por Idempotency-Key tenant-scoped y rechazan la reutilización inconsistente.",
)
heading(doc, "Auditoría", 2)
paragraph(
    doc,
    "Se emiten eventos para created, updated, submitted, review_started, returned, completed, voided, superseded, "
    "revision_created y para "
    "clone/update/publish/archive de configuración. Los historiales de USER no exponen operaciones administrativas como transacciones.",
)

doc.add_page_break()
heading(doc, "6. Experiencia USER MODE y ADMIN MODE", 1)
heading(doc, "USER MODE", 2)
paragraph(
    doc,
    "La nueva pantalla aparece bajo Enterprise Strategy Manager → Idea & Demand después de Idea Lifecycle. Presenta "
    "colas My Analyses, Draft, Submitted, In Review, Returned, Ready for Proposal, Completed y All Authorized; ofrece "
    "búsqueda, detalle compacto, editor estructurado, alternativas, recomendación, readiness e historial.",
)
picture(
    doc,
    EVIDENCE / "gate07a-ba-user-business-analysis.png",
    "Figura 1. Business Analysis & Options en USER MODE.",
)
heading(doc, "ADMIN MODE", 2)
paragraph(
    doc,
    "Business Analysis Configuration aparece en Idea & Demand Configuration. Su navegación cubre veinte grupos de "
    "política. Expone clone/edit/publish/archive y preview no persistente; declara explícitamente que no administra expedientes vivos.",
)
picture(
    doc,
    EVIDENCE / "gate07a-ba-admin-configuration.png",
    "Figura 2. Business Analysis Configuration en ADMIN MODE.",
)
heading(doc, "Navegación y accesibilidad", 2)
paragraph(
    doc,
    "El blueprint tipado ubica BA entre Idea y Proposal, genera rutas y breadcrumbs determinísticos, conserva labels de "
    "modo, usa regiones/nav con nombres accesibles y aplica diseño responsive sin replicar la lógica de autorización.",
)

doc.add_page_break()
heading(doc, "7. Compatibilidad y no duplicación", 1)
table(
    doc,
    ["Escenario", "Comportamiento"],
    [
        ["Nueva Idea CAPITAL_OWNER", "BA obligatorio antes de Project Proposal"],
        ["CONTRACTOR_DELIVERY", "BA no obligatorio por defecto; política configurable"],
        ["DIRECT_INTERNAL", "BA no obligatorio por defecto; política configurable"],
        ["Snapshot legacy sin política BA", "Grandfathered; conserva creación histórica de Proposal"],
        ["Proposal existente", "Sigue válida; columnas BA son nullable"],
        ["BA completo con cambio material", "Nueva revisión; original permanece inmutable"],
    ],
    [2.25, 4.45],
)
heading(doc, "Capacidades reutilizadas", 2)
bullets(
    doc,
    [
        "Strategic Objectives: sólo referencias, sin nuevo catálogo paralelo.",
        "Enterprise Workspaces y contexto tenant: mismo modelo de alcance y herencia.",
        "AdminConfiguration: mismo lifecycle y hashing de configuración.",
        "SecurityEvent y roles: mismos mecanismos transversales.",
        "Project Proposal: extensión nullable y mapping; no se duplicó su lifecycle.",
        "App shell y navegación estratégica: nuevas vistas, sin crear una segunda aplicación.",
    ],
)
heading(doc, "Contrato de linaje futuro para Business Case", 2)
paragraph(
    doc,
    "Gate 07M-BC podrá referenciar source_business_analysis_id, completion_hash, recommended_option_id y el linaje "
    "de configuración congelado. BA no crea el Five Case Model ni mantiene sus revisiones; sólo publica evidencia "
    "estable de necesidad, alternativas y recomendación para que Business Case la consuma sin copiar su lifecycle.",
)

doc.add_page_break()
heading(doc, "8. Validación y evidencia", 1)
status_table(
    doc,
    [
        ["Backend lint", "PASS", "Ruff sobre app y tests: sin hallazgos."],
        ["Gate 07A-BA SQLite", "PASS", "16 pass: 8 BA + 7 Idea/Proposal + 1 migración; 3 PostgreSQL omitidas por dialecto."],
        ["Regresión Idea + Proposal", "PASS", "7 pruebas: nuevo gate y compatibilidad grandfathered."],
        ["TypeScript", "PASS", "tsc --noEmit / compilación tipada sin errores."],
        ["Frontend lint/format", "PASS", "ESLint sin errores; Prettier conforme. Advertencias baseline documentadas."],
        ["Frontend build", "PASS", "Vite produjo bundle de producción."],
        ["Frontend unitario completo", "PASS", "52 archivos, 236 pruebas superadas."],
        ["Browser E2E BA", "PASS", "3/3 en Chromium (42,9 s): stack real + contratos USER y ADMIN; sólo API BA interceptada en lifecycle/config."],
        ["Migración up/down/up", "PASS", "Ciclos reales PostgreSQL y SQLite 0058→0059→0058→0059; FKs, índices y columnas verificados."],
        ["PostgreSQL E2E", "PASS", "18/18 en base desechable PostgreSQL 16: lifecycle, negativos, Idea/Proposal, snapshots y constraints."],
        ["Despliegue localhost", "PASS", "Docker reconstruido; API/frontend HTTP 200, Alembic 0059 head y 18 rutas BA publicadas."],
        ["Backend completo", "PARCIAL", "546 pass / 27 skip; 10 fallos baseline ajenos y 28 errores por acceso al temp de pytest."],
    ],
)
heading(doc, "Cobertura funcional demostrada", 2)
bullets(
    doc,
    [
        "Idea aceptada cambia a READY_FOR_BUSINESS_ANALYSIS y bloquea Proposal.",
        "Creación idempotente y conflicto ETag 412.",
        "Contenido completo, submit y rechazo SoD al mismo actor.",
        "Revisor independiente retorna, retoma y completa; produce READY_FOR_PROJECT_PROPOSAL.",
        "Preview/creación de Proposal conserva id y hash de BA y mapping recomendado.",
        "Mutación posterior de COMPLETED rechazada; una revisión material completa supersede sin borrar el original.",
        "VOIDED queda auditado y libera el único slot activo de la Idea para un reemplazo gobernado.",
        "Cambio de configuración publicada marca BA como STALE y vuelve a bloquear Proposal.",
        "Preview ADMIN no cambia el número de configuraciones; published no se edita.",
    ],
)
heading(doc, "Lectura del barrido global", 2)
paragraph(
    doc,
    "El barrido backend global no es una puerta limpia del repositorio: 28 errores provienen de PermissionError sobre "
    "AppData/Local/Temp/pytest-of-Ricardo y los 10 fallos restantes pertenecen a suites PostgreSQL/Enterprise Structure/"
    "Investment Decision ya fuera del delta BA. Las suites focalizadas nuevas y las regresiones directamente afectadas sí pasan.",
)
paragraph(
    doc,
    "Docker Desktop no estaba en PATH, pero su PostgreSQL 16 local fue localizado y usado con dos bases desechables: "
    "una para el ciclo Alembic real y otra para la suite funcional. Ambas puertas Gate 07A pasan y las bases de prueba "
    "fueron eliminadas al terminar. Este informe aún no "
    "emite GATE07A_BA_BUSINESS_ANALYSIS_RELEASE_VALIDATED porque el barrido backend global del baseline no está limpio.",
)
doc.add_page_break()
heading(doc, "9. Inventario de cambios", 1)
table(
    doc,
    ["Área", "Archivos principales"],
    [
        ["Backend BA", "app/modules/business_analysis/{models,schemas,service,router}.py"],
        ["Integraciones", "idea_demand/service.py; project_proposal/{models,schemas,service}.py"],
        ["API y seguridad", "api/v1/router.py; enterprise_structure/{constants,permissions}.py"],
        ["Workspace", "workspace_context/navigator.py"],
        ["Migración", "alembic/versions/20260910_0059_gate07a_ba_business_analysis.py"],
        ["Frontend BA", "features/business-analysis/{api,types,workspace,admin,css,test}"],
        ["Navegación", "strategicNavigationBlueprint.ts; applicationNavigation.ts; App.tsx"],
        ["Regresión UI", "IdeaLifecycleWorkspace; ProjectProposalWorkspace/types/css; WorkspaceOperationalPage"],
        ["Pruebas", "Gate07A BA/migración SQLite+PostgreSQL; navegación/unitarios; Playwright BA"],
        ["Arquitectura", "ADR 61 y guía gate07a-ba-business-analysis-governance.md"],
    ],
    [1.45, 5.25],
)
heading(doc, "Contratos de salida", 2)
table(
    doc,
    ["Señal", "Condición"],
    [
        ["GATE07A_BA_BUSINESS_\u200bANALYSIS_RELEASE_VALIDATED", "Pendiente sólo de cerrar la regresión backend global del baseline."],
        ["READY_FOR_GOVERNED_PROJECT_PROPOSAL_\u200bFROM_BUSINESS_ANALYSIS", "BA COMPLETED, fuente/config vigentes, sin blockers."],
        ["ENHANCE_GATE07M_BC_WITH_FIVE_CASE_MODEL", "Siguiente incremento; no implementado en este alcance."],
    ],
    [3.25, 3.45],
)

doc.add_page_break()
heading(doc, "10. Riesgos, límites y cierre", 1)
table(
    doc,
    ["Riesgo / deuda", "Impacto", "Tratamiento"],
    [
        ["Suite backend global no limpia", "Impide usarla como señal única de release", "Mantener suites focalizadas; corregir permisos temp y baseline en incremento separado"],
        ["Cadena SQLite histórica 0013", "Rebuild desde cero se detiene antes de Gate 07A", "Portar FK heredada a batch mode; el ciclo inmediato 0058↔0059 sí pasa"],
        ["Escala de colas", "Listados grandes pueden requerir paginación", "Añadir paginación/cursor e índices de consulta en hardening"],
        ["Editor JSON ADMIN", "Potente pero técnico", "Evolucionar a formularios por sección sin cambiar el contrato"],
        ["Matrix ratings", "Requiere validación de escala más estricta", "Agregar límites/rangos publicados en hardening"],
        ["Seed concurrente", "Procesos separados aún pueden competir en bootstrap", "RLock mitiga concurrencia en proceso; PostgreSQL/CI debe validar multiproceso"],
        ["Modelos alternos", "Política inicial no obliga BA", "Activar sólo mediante configuración publicada y pruebas de contexto"],
    ],
    [2.0, 2.05, 2.65],
)
heading(doc, "Conclusión", 2)
paragraph(
    doc,
    "Gate 07A-BA queda implementado en la arquitectura actual: dominio, API, seguridad, "
    "configuración, navegación, UI, integración con Idea/Proposal, migración, pruebas y documentación. La solución "
    "preserva compatibilidad y evita duplicar capacidades existentes. El stack Docker local quedó actualizado y sano.",
)
paragraph(
    doc,
    "Las puertas focalizadas, PostgreSQL y frontend están validadas. La liberación formal permanece condicionada a "
    "sanear el barrido backend global, cuyas incidencias baseline explícitas no deben confundirse con defectos del módulo BA.",
)
heading(doc, "Siguiente paso recomendado", 2)
paragraph(
    doc,
    "Ejecutar el hardening de Five Case Model dentro de Gate 07M-BC consumiendo el BA completo como evidencia de "
    "necesidad y recomendación, sin trasladar a BA cálculos financieros o decisiones FID.",
)

doc.add_page_break()
heading(doc, "Anexo A. Trazabilidad de entregables", 1)
table(
    doc,
    ["#", "Entregable", "Evidencia"],
    [
        ["1", "Baseline assessment", "Resumen, arquitectura y lectura del barrido global"],
        ["2", "ADR de inserción", "docs/61-adr-gate07a-ba-business-analysis-options.md"],
        ["3", "Compatibilidad histórica", "Política grandfathered y pruebas Gate07B"],
        ["4", "Matriz USER/ADMIN", "Secciones 2, 5 y 6"],
        ["5", "Reuso de dominio", "Secciones 2 y 7"],
        ["6", "Diseño del agregado", "Secciones 2 y 4"],
        ["7", "Lifecycle", "Sección 3"],
        ["8", "Eligibility/readiness", "Servicio, API y sección 3"],
        ["9", "Need/Problem/Opportunity", "BusinessAnalysisPayload + UI"],
        ["10", "Business Drivers", "Catálogo configurable + payload"],
        ["11", "Strategic Alignment", "Referencias a objetivos existentes"],
        ["12", "Stakeholders", "Payload y pantalla de registro"],
        ["13", "Current State", "Payload y readiness"],
        ["14", "Desired Outcomes", "Payload y readiness"],
        ["15", "Expected Value", "Dimensiones configurables"],
        ["16", "Assumptions", "Payload, mapping y readiness"],
        ["17", "Constraints", "Payload, mapping y readiness"],
        ["18", "High-Level Business Requirements", "Payload y readiness"],
        ["19", "Options model", "BusinessAnalysisOptionIn"],
        ["20", "Options Evaluation", "Matriz publicada + ratings"],
        ["21", "Recommended Option", "Selección, rationale, trade-offs y condiciones"],
        ["22", "Proposal mapping", "ProjectProposalService + snapshot"],
        ["23", "Idea lineage", "IDs, revisión y source hash"],
        ["24", "Future Business Case lineage", "Contrato descrito en sección 7"],
        ["25", "ADMIN configuration model", "AdminConfiguration reutilizada"],
        ["26", "ADMIN lifecycle", "Clone/update/publish/archive"],
        ["27", "Inheritance", "Tenant→Enterprise→BU→Portfolio"],
        ["28", "Roles/permissions/SoD", "Permisos dedicados + revisor independiente"],
        ["29", "ETag/idempotency/audit", "If-Match, Idempotency-Key, SecurityEvent"],
        ["30", "USER navigation", "Blueprint + BusinessAnalysisWorkspace"],
        ["31", "ADMIN navigation", "Blueprint + BusinessAnalysisAdminView"],
        ["32", "PostgreSQL happy path", "Suite postgres Gate07A-BA"],
        ["33", "PostgreSQL negative", "Suite postgres Gate07A-BA"],
        ["34", "Historical compatibility", "test_project_proposal_gate07b"],
        ["35", "Governance Model tests", "CAPITAL_OWNER / alternos / legacy"],
        ["36", "Browser USER E2E", "business-analysis-governance.spec.ts"],
        ["37", "Browser ADMIN E2E", "business-analysis-governance.spec.ts"],
        ["38", "Contamination proof", "Unit + Playwright"],
        ["39", "Migration proof", "Alembic 0059 up/down/up"],
        ["40", "Regression summary", "Sección 8"],
        ["41", "Before/after diagram", "Sección 2"],
        ["42", "Updated USER/ADMIN docs", "gate07a-ba-business-analysis-governance.md"],
        ["43", "Final technical report", "Este documento"],
        ["44", "Final gap analysis", "Sección 10"],
        ["45", "Next improvement", "ENHANCE_GATE07M_BC_WITH_FIVE_CASE_MODEL"],
    ],
    [0.35, 2.5, 3.85],
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
