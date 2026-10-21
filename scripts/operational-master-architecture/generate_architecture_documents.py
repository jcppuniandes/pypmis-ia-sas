"""Two source-style Word reports, architecture only; no product imports."""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import io
import json
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.shared import Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "docs/architecture/operational/architecture-data.json"
OUT = ROOT / "outputs/operational-master-architecture-20261019"
QA = ROOT / "outputs/.qa/operational-master-architecture-20261019"
REFERENCE = (
    ROOT
    / "outputs/business-uat-signoff-20261018/P&Pmis_Informe_Tecnico_Business_UAT_Signoff_20261018.docx"
)
REFERENCE_SHA = "475c1674cf398f005e3239df01cf84c3e649951d8a45b3056d9f0fddf6c7f3a7"
NAMES = [
    "P&Pmis_Operational_Macroprocess_Master_Architecture.docx",
    "P&Pmis_Operational_Macroprocess_Implementation_Roadmap.docx",
]
spec = importlib.util.spec_from_file_location(
    "retained_document_components",
    ROOT / "scripts/business-uat-signoff/generate_signoff_documents.py",
)
assert spec and spec.loader
components = importlib.util.module_from_spec(spec)
spec.loader.exec_module(components)
components.REFERENCE = REFERENCE
components.prose = (
    str  # Preserve technical identifiers and path:line references exactly.
)
p, h, t = components.paragraph, components.heading, components.table


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def lines(value):
    return "; ".join(str(item) for item in value)


def start(title, lead):
    doc = Document(REFERENCE)
    logo = deepcopy(doc.paragraphs[0]._p)
    sect = deepcopy(doc._element.body.sectPr)
    for element in list(doc._element.body):
        doc._element.body.remove(element)
    doc._element.body.append(logo)
    doc._element.body.append(sect)
    p(doc, title, "Title")
    p(
        doc,
        "P&Pmis Construction AI · Operational Macroprocess\nInforme técnico de arquitectura · 19 de octubre de 2026",
        "Subtitle",
    )
    p(doc, lead)
    p(
        doc,
        "Estado: PROPOSED_DESIGN. Arquitectura documental; no implementación de módulos, migración, liberación ni aceptación de negocio.",
        bold=True,
    )
    return doc


def evidence(doc, *refs):
    para = p(doc, "Evidencia inspeccionada: " + "; ".join(refs) + ".")
    for run in para.runs:
        run.font.size = Pt(9)


def pairs(doc, rows):
    return t(doc, ["Elemento", "Definición / control propuesto"], rows, [1.7, 5.36])


def section(doc, number, title, paragraphs):
    h(doc, f"{number}. {title}")
    for text in paragraphs:
        p(doc, text)


def master(data):
    c = data["contracts"]
    doc = start(
        "Operational Macroprocess\nMaster Architecture",
        "Arquitectura maestra y mapa de dominios para evolucionar el control operacional reutilizando las capacidades de P&Pmis, sin duplicar sus autoridades estratégicas ni operativas.",
    )
    section(
        doc,
        1,
        "Resumen ejecutivo",
        [
            "El repositorio contiene una base operacional real pero parcial: cantidades BIM/Excel y APU, importación de cronogramas, CBS/control accounts, costo, progreso, cambios, RFQ, contratos/PO, documentos, AWP y análisis de reclamaciones. Su presencia en código no acredita un ciclo integral gobernado ni una pantalla accesible para cada entrada del menú.",
            "Se definieron 25 candidatos, no 25 módulos nuevos: 12 dominios, 4 subdominios, 3 capacidades transversales, 4 workflows, 1 área de configuración y 1 dominio futuro. Las acciones propuestas son 12 EXTEND_EXISTING, 9 BUILD_LATER, 3 CROSS_CUTTING y 1 MERGE_WITH_ANOTHER_DOMAIN. BUILD_NOW es cero en este incremento.",
            "La dependencia crítica es el enlace entre enterprise_workspaces (identidad canónica PROJECT) y projects.id (agregado operacional heredado). La creación canónica intencionalmente no crea una fila legacy. Un adaptador explícito es indispensable antes de ampliar la operación; no se permite igualar IDs ni crear otro Project master.",
            "La recomendación única es WAVE O1 — PROJECT CONTROLS FOUNDATION, comenzando por O1A: contexto/binding, autorización, período y políticas comunes. O1B integra baselines y fuentes; O1C agrega forecast/reporting gobernados. Este documento se detiene antes de ejecutar O1.",
            "Corte de evidencia: HEAD c75a7da3db3bdcf326b6bed7834182ee126da921. La revisión fue estática de archivos, rutas, servicios, modelos y definiciones de tests. No se ejecutaron tests de producto, UAT de navegador ni comandos de base de datos; los riesgos son hallazgos de diseño, no incidentes de importes de producción demostrados.",
        ],
    )
    pairs(
        doc,
        [
            [
                "Cobertura documental",
                "25 candidatos; 15 handoffs; 26 objetos/facetas; 26 integraciones; 13 eventos conceptuales; 12 ADR propuestos; 18 brechas.",
            ],
            [
                "Inventario",
                "15 unidades de evidencia backend, 25 frontend y 16 procurement/handoff. Son 56 filas de diferente granularidad, no 56 funciones distintas.",
            ],
            [
                "Protección",
                "860 archivos de producto bajo huella SHA; no cambios en lógica estratégica, navegación, esquema ni datos. SAFE_TO_REMOVE=0.",
            ],
            ["Estado estratégico", data["strategic_state_preserved"]],
        ],
    )
    section(
        doc,
        2,
        "Frontera Strategic → Operational",
        [
            "Estrategia conserva demanda, propuesta, portfolio, Business Case, FIN, beneficios, FEL, readiness, QRA, Funding y decisiones de inversión. Operación consume referencias versionadas, restricciones e insumos; no recalcula FID, no aprueba inversión y no convierte cobertura de fondos en presupuesto o gasto operacional.",
            "El handoff identifica tenant, Workspace PROJECT, governance model, Project Type, autorización, política/revisión/hash, activación y fuentes. Cada insumo incluye ID, revisión, hash, as-of/validez, clasificación, condiciones y dueño de actualización. Estado missing/stale/bloqueante permanece visible.",
            "CAPITAL_OWNER consulta la ruta estratégica aplicable y las condiciones Gate/FID/Post-FID actuales. CONTRACTOR_DELIVERY y DIRECT_INTERNAL conservan sus autorizaciones legítimas sin forzar el ciclo completo de inversión Owner. ACTIVE no elimina condiciones commitment-only ni autoriza por sí solo una orden de compra.",
        ],
    )
    handoffs = c["strategic_operational_handoff"]["items"]
    t(
        doc,
        ["ID / insumo", "Tratamiento", "Límite de transferencia"],
        [
            [r["id"] + " " + r["artifact"], r["classification"], r["acceptance_rule"]]
            for r in handoffs
        ],
        [1.75, 1.32, 3.99],
    )
    p(
        doc,
        "Las autoridades NOT TRANSFERRED incluyen inversión, política estratégica, Funding, BC/FIN, beneficios, QRA estratégico, CPM P6 y contabilización/pago. No se cambia el estado histórico ni se reanuda el signoff de negocio en esta tarea.",
    )
    evidence(
        doc,
        "backend/app/modules/post_fid_governance/service.py:492",
        "backend/tests/test_project_creation_process.py:642",
        "docs/architecture/operational/control-contracts.json",
    )
    section(
        doc,
        3,
        "Principios de arquitectura",
        [
            "Monolito modular con límites de dominio explícitos. Cada agregado/faceta tiene un único propietario y API/comando interno validado; consumir una referencia o snapshot no autoriza escribir las tablas del origen. Integración no significa convertir Project Controls en dueño de todos los datos.",
            "Reutilizar antes de extender; extender antes de crear. Configuración ADMIN publicada y versionada se separa de transacciones USER. Todas las operaciones verifican tenant/Workspace, membresía, permiso, estado, revisión y condiciones aplicables. Historial aprobado inmutable, optimistic concurrency e idempotencia en comandos materiales.",
            "El catálogo de módulos no equivale a un dominio implementado. Los placeholders permanecen identificados como reservados. No se propone microservicios, event bus, reescritura de navegación, migración amplia ni integración ERP/P6 directa.",
        ],
    )
    section(
        doc,
        4,
        "Lifecycle operacional",
        [
            "El proceso es una secuencia conceptual con ramas y ciclos, no una lista universal obligatoria. Preparación, aprobación, ejecución, control, forecast y cambio se repiten por período. Lo no aplicable se declara con razón y autoridad según Project Type; no se fabrican commissioning, claims o compras para completar una lista."
        ],
    )
    t(
        doc,
        ["Paso", "Finalidad", "Dominios"],
        [
            [r["id"] + " " + r["name"], r["description"], lines(r["domain_ids"])]
            for r in data["process_map"]
        ],
        [2.0, 4.26, 0.8],
    )
    section(
        doc,
        5,
        "Mapa de dominios recomendado",
        [
            "D01–D25 son identificadores de diseño estables. La clasificación evita crear módulos por cada etiqueta. Forecast y Claims se integran como subdominios; Documents, Interface y Reporting son transversales; Handover/Readiness/Closeout/Asset Transition son workflows diferenciados.",
            "El detalle completo de propósito, objetos, fuentes, entradas, salidas, upstream/downstream, ADMIN/USER, aplicabilidad, handoff y P6 está en las hojas 01 a 06 del workbook y en architecture-data.json. Esta síntesis no sustituye esas matrices.",
        ],
    )
    t(
        doc,
        ["ID / capacidad", "Clasificación", "Acción / ola"],
        [
            [
                d["id"] + " " + d["name"],
                d["classification"],
                d["planned_action"].replace(
                    "MERGE_WITH_ANOTHER_DOMAIN", "MERGE_WITH_ANOTHER_\nDOMAIN"
                )
                + " / "
                + d["wave"],
            ]
            for d in data["domains"]
        ],
        [3.28, 1.77, 2.01],
    )
    section(
        doc,
        6,
        "Project Controls como núcleo de integración",
        [
            "Scope, Schedule, Cost, Resource, Progress, Change, Risk e Issue alimentan un núcleo de coordinación de período, mappings y lineage. Forecast utiliza esas fuentes para emitir una revisión de resultados; Performance/Reporting presenta snapshots reproducibles. Cada origen mantiene sus reglas y aprobaciones.",
            "El núcleo posee el registro de período y los bundles/reportes emitidos, no los contratos, fechas CPM o mediciones fuente. Un reporte declara baseline set, data date, cutoff, fuente, método, versión, excepciones y permiso de acceso. Abrir el dashboard no debe ejecutar comandos de negocio ni fabricar un forecast.",
            "BIM/4D es una capacidad transversal: referencia versión de modelo, GUID/ExpressID, cantidades y actividades controladas. No es otro motor de programación. El visor BIM y APU se reutilizan; el mapping IFC↔actividad y la simulación 4D persistente son alcance futuro, no capacidad completa demostrada.",
        ],
    )
    pairs(
        doc,
        [
            [
                "Fuentes",
                "Scope / Schedule / Cost / Resource / Progress / Change / Risk / Issue conservan autoridad.",
            ],
            [
                "Integración",
                "Project Workspace + período común + mappings versionados + validación de completitud + snapshots.",
            ],
            [
                "Resultados",
                "Forecast / performance / reportes / look-ahead consumen esas fuentes sin escrituras ocultas.",
            ],
        ],
    )
    evidence(
        doc,
        "backend/app/services/control_core.py:27",
        "frontend/src/components/BimIfcModelViewer.tsx:1",
    )
    section(
        doc,
        7,
        "Descripciones y límites por dominio",
        [
            "Las siguientes fichas distinguen capacidad observada de autoridad objetivo. Los modelos conceptuales no son nuevas tablas ni endpoints publicados. Los nombres son etiquetas de diseño, no una modificación del menú vigente."
        ],
    )
    for d in data["domains"]:
        h(doc, d["id"] + " · " + d["name"], 2)
        p(doc, d["purpose"] + " " + d["source_authority"])
        p(doc, "Objetos: " + lines(d["objects"]) + ".")
        p(doc, "Estado observado: " + d["current_status"])
        p(
            doc,
            "Evolución propuesta: "
            + d["planned_action"]
            + " · "
            + d["wave"]
            + " · complejidad "
            + d["complexity"]
            + ". "
            + d["rationale"],
        )
    section(
        doc,
        8,
        "Aplicabilidad Owner, Contractor y Direct Internal",
        [
            "La matriz completa asigna MANDATORY, OPTIONAL, CONDITIONAL o NOT_APPLICABLE a los 25 candidatos para cada governance model. Es política objetivo para revisión, no configuración runtime ya aplicada. Project Type, materialidad, estrategia y alcance determinan profundidad y controles.",
            "CAPITAL_OWNER controla inversión, entrega, interfaces y riesgo del activo, con referencias vigentes FID/Post-FID cuando correspondan. Puede supervisar contratistas sin gestionar directamente crews. CONTRACTOR_DELIVERY controla alcance contractual, costo propio, productividad, ingresos/certificados, avisos, variaciones y obligaciones de entrega; contrato cliente de ingreso no es un compromiso de compra.",
            "DIRECT_INTERNAL parte de autorización interna/sponsor y aplica controles proporcionales. No exige automáticamente FEL/FID completo, external procurement, commissioning ni AWP. Evidencia/contexto mínimo y seguridad sí se preservan en todas las rutas.",
        ],
    )
    pairs(
        doc,
        [
            [
                "Owner",
                "Autorización de inversión y condiciones; oversight contractual y receptor del activo.",
            ],
            [
                "Contractor",
                "Award/contrato de cliente y movilización/NTP según política; ejecución de alcance y subcontratos diferenciados.",
            ],
            [
                "Direct Internal",
                "Referencia de autorización interna/sponsor y controles adaptados al trabajo.",
            ],
            [
                "No aplicable",
                "Razón y autoridad trazables; jamás objetos ficticios ni gates completados por defecto.",
            ],
        ],
    )
    section(
        doc,
        9,
        "ADMIN MODE y USER MODE",
        [
            "ADMIN configura/publica plantillas de Project Type, aplicabilidad, estados, roles/SoD, umbrales, mappings/catálogos, calendario, métodos de progreso, reglas de baseline/change/forecast, paquetes AWP y referencias de evidencia. Publicar configuración no aprueba ninguna transacción del proyecto.",
            "USER prepara/importa, revisa, captura, solicita, aprueba bajo delegación y ejecuta comandos del dueño del dominio. Draft simple puede ser directo; las decisiones materiales tienen la separación requerida. Las consultas comunes muestran historia, configuración aplicable y lineage sin permitir escrituras por acceso al reporte.",
            "Se reutilizan el catálogo de navegación y WorkspaceContext. El diseño no agrega otro router ni cambia el menú actual. La hoja 03 contiene la matriz de los 25 candidatos; la navegación ADMIN operacional específica es una brecha para una fase autorizada.",
        ],
    )
    section(
        doc,
        10,
        "Data ownership y modelos conceptuales",
        [
            "La hoja 05 define 26 objetos/facetas con propietario único. Quantity de alcance pertenece a Scope; hechos instalados a Field; aprobación/consolidación de medición por método a Progress. Work Package definition/baseline pertenece a Scope; release y hechos de ejecución son registros de Field que la referencian.",
            "Objeto conceptual mínimo: identidad canónica, ID del agregado dentro del tenant/Workspace, revisión/configuración, estado, unidad/base temporal cuando aplique, origen/hash, actor y auditoría. IDs externos se guardan con namespace y versión; no se usan nombres para fusionar proyectos.",
            "Las extensiones futuras deben evaluar las tablas existentes y su compatibilidad antes de proponer schema. No se crea un segundo maestro Project, Supplier, Contract/PO, baseline ni documento. Las tablas projects y sus FKs se conservan hasta un plan de transición acotado y autorizado.",
        ],
    )
    t(
        doc,
        ["Grupo", "Autoridad y consumo"],
        [
            [
                "Identidad / autorización",
                "Enterprise Workspace / ruta de autorización vigente; operación consume referencia.",
            ],
            [
                "Scope / WBS / cantidades",
                "Scope controla definición; mappings a PBS/WBS/CBS y actividades son versionados M:N.",
            ],
            [
                "Schedule / calendario",
                "Registro de versión en Schedule; fechas CPM/float/critical de P6 importado.",
            ],
            [
                "Budget / actual",
                "Cost; obligación legal fuente en Contract Administration y posting/pago fuera de este incremento.",
            ],
            [
                "Progress / Forecast / Change",
                "Aprobación propia de cada agregado, con fuentes congeladas y receipts de implementación.",
            ],
            [
                "Document / paquete / claim",
                "Document Controller / Scope definición y Field release / subdominio Claims de Contracts.",
            ],
        ],
        [1.8, 5.26],
    )
    section(
        doc,
        11,
        "Integraciones y eventos",
        [
            "La hoja 06 contiene 26 enlaces de datos/eventos con fuente, destino, dirección, referencia/snapshot/comando, autoridad y prerrequisitos. Ninguno autoriza escritura directa sobre tablas ajenas. Los ciclos de negocio no son ciclos del DAG de implementación.",
            "El envelope de evento conceptual incluye event_id/type/schema, tenant/Workspace, aggregate/version, occurred_at UTC, source owner, actor/action, idempotency/correlation/causation, source hashes, policy y audit reference. Se emite tras una transacción autorizada exitosa; replays son idempotentes y no inventan aprobación.",
            "Comenzar con registros de dominio/auditoría y llamadas de servicios dentro del monolito modular. No se crea infraestructura event bus. Los 13 nombres siguientes son contratos propuestos, no eventos emitidos en esta tarea.",
        ],
    )
    t(
        doc,
        ["Evento", "Dueño / límite"],
        [
            [r["name"], r["source_owner"] + ". " + r["authority_limit"]]
            for r in c["event_model"]["events"]
        ],
        [2.65, 4.41],
    )
    section(
        doc,
        12,
        "Frontera con Primavera P6",
        [
            "P6 permanece como motor CPM operacional externo: calcula fechas, ruta crítica/near-critical y float. P&Pmis controla identidad, versiones, mapping, validación, aceptación de cortes y reporting. No propone CPM principal nativo, API directa P6, OPC ni sincronización automática.",
            "El flujo objetivo es intercambio manual XML/XER: seleccionar versión/manifest → actualizar y calcular en P6 → importar a staging → validar/diff → revisar/aceptar nueva versión current o solicitar baseline → conservar historia/rollback. El exportador de roundtrip completo es futuro y debe declarar subset soportado; no se presenta un CSV/JSON como exportador P6.",
            "Oracle documenta una asimetría relevante: XML transporta baselines, mientras XER no las intercambia de ese modo. Para baseline usar XML compatible o un artefacto aprobado separado y explícitamente vinculado al XER current. Esta restricción del formato no demuestra compatibilidad completa del importador local.",
            "Validar IDs externos/namespace, WBS, relaciones FS/SS/FF/SF y lag, vínculos externos no resueltos, calendarios/turnos/excepciones, unidades, restricciones, data date, timezone y fechas planned/actual/remaining. No asumir 8 h/día; unknown bloquea conversiones dependientes. Critical/float importados no se sustituyen silenciosamente por una heurística local.",
        ],
    )
    pairs(
        doc,
        [
            ["Baseline", "Versión aprobada inmutable, independiente del update."],
            [
                "Current",
                "Último update aceptado con Data Date explícito; importar no aprueba baseline.",
            ],
            [
                "Forecast",
                "Escenario/proyección explícito; no se declara resultado CPM P6 si no fue calculado/importado allí.",
            ],
            [
                "Rollback",
                "Reversión controlada de pointer o nueva revisión; no borrar archivos/aprobaciones históricas.",
            ],
        ],
    )
    p(
        doc,
        "Fuente primaria consultada el 19-10-2026: Oracle, Import/Export File Formats. "
        + c["p6_manual_schedule_contract"]["primary_source"]["url"],
    )
    evidence(
        doc,
        "backend/app/services/schedule_ingestion.py:402",
        "backend/app/services/schedule_ingestion.py:1612",
        "backend/app/services/schedule_ingestion.py:1731",
    )
    section(
        doc,
        13,
        "Gobernanza de baselines y presupuesto",
        [
            "Un Control Baseline Set referencia versiones aprobadas separadas de Scope, Schedule, Cost y Resource opcional, junto con mapping/configuración. Aprobación parcial es visible; no se transforma en aprobación global. Original, approved revision, current control, forecast, actual e historia no son sinónimos.",
            "Original Budget es la primera versión aprobada. Current Budget deriva de cambios/transferencias aprobados e implementados con lineage. BAC identifica el presupuesto aplicable a medición del desempeño, no es automáticamente capital FID, strategic envelope, reserva o saldo bancario.",
            "Scope Item/cantidad/unidad/método y revisiones se alinean M:N con PBS/WBS/CBS/control accounts/actividades/contratos, sin asumir jerarquías idénticas. Scope baseline no publica Cost automáticamente. Cambios de moneda/FX, redondeo, inclusión de contingencia y reservas requieren política, fecha/fuente y reconciliación.",
            "La importación legacy puede actualizar baseline_start y materializar Budget; esa conducta se identifica como brecha, no arquitectura objetivo. El futuro endurecimiento preservará historia y canalizará propuestas al dueño autorizado en lugar de sobrescribir aprobaciones.",
        ],
    )
    section(
        doc,
        14,
        "Change Management transversal",
        [
            "Un solo registro Change enlaza origen, causa, alcance, dominios afectados, evaluaciones, evidencia, versión de política y decisión. Assessment de costo/plazo/riesgo/contrato no implica aprobación. Change aprobado tampoco implica implementación completada.",
            "El dueño de cada fuente aplica comandos idempotentes autorizados, con revisión esperada y receipt de nueva versión. La coordinación registra IMPLEMENTED/PARTIAL/FAILED según respuestas, no falsifica atomicidad interdominio. Un fallo requiere remediación/compensación gobernada, sin alterar baseline histórica ni aprobar automáticamente variación legal.",
            "Variaciones de contrato pertenecen a Contract Administration; evaluación consolidada a Change. Rebaselining solo por decisión explícita de baseline y trazabilidad de antes/después. Issues, NCR o claims pueden originar un cambio, pero no se fusionan sus autoridades.",
        ],
    )
    section(
        doc,
        15,
        "Procurement, Commitments y Contract Administration",
        [
            "Resolución conceptual: Procurement posee requisición, sourcing, RFQ/bids, evaluación y recomendación de adjudicación. Contract Administration & Commitments es un dominio cohesivo distinto que reutiliza Contract, PO/subcontract, SOV, notices, certificaciones/recepciones y obligación legal. Cost consume referencias reconciliadas, no emite otra PO.",
            "Vendors/Partners se tratan como shared party/reference data futura reutilizable, no nuevos maestros porque existan dos menús. Strategic Commercial Strategy y Funds permanecen estratégicos; su handoff aporta contexto/constraints, no transacciones de compra ni gasto. Una recomendación con mejor score RFQ no constituye award aprobado.",
            "Presupuesto aprobado, compromiso bruto, compromiso neto pendiente, actual reconocido, accrual y Funding son medidas distintas. PO incluido o call-off consume exposición del padre sin volver a sumarla; PO ADDITIVE añade solo obligación incremental. Cancelación, cumplimiento y liquidación tienen una política explícita.",
            "Certificado, recepción, factura de referencia y accrual pueden representar el mismo hecho económico. Se requiere clave de origen/proyecto/moneda/línea y reglas de matching/reversión para contabilizar una vez. No se deduce un actual dos veces del compromiso. Certificación no es pago bancario; ERP/Accounting/Treasury quedan fuera de esta fase.",
            "La agregación actual localizada suma Contract.value y PO.committed_amount, y maneja committed/executed sin semántica uniforme de neteo. Esto prueba exposición de diseño, no que se hayan duplicado valores reales de producción. Antes de ampliar el uso vivo deben definirse y probarse casos de call-off, included, additive, cancelación, accrual/reversal y monedas.",
        ],
    )
    pairs(
        doc,
        [
            [
                "KEEP STRATEGIC",
                "Commercial strategy, Funds/readiness, FID y restricciones upstream.",
            ],
            [
                "REUSE OPERATIONAL",
                "RFQ, Contract, PO, SOV, notices, certificados y recepciones existentes.",
            ],
            [
                "SHARED MASTER DATA",
                "Identidad de parties/Vendors/Partners, con contexto y permisos; no string duplicado como maestro nuevo.",
            ],
            [
                "DEPRECATE / FUTURE",
                "Solo tras revisar consumidores y paridad. Ningún componente se elimina aquí.",
            ],
        ],
    )
    evidence(
        doc,
        "backend/app/services/integrated_control.py:392",
        "backend/app/modules/post_fid_governance/service.py:499",
        "docs/validation/operational-master-architecture-20261019/procurement-handoff-inventory.json",
    )
    section(
        doc,
        16,
        "AWP opcional y ejecución por paquetes",
        [
            "AWP se habilita por Project Type/estrategia de ejecución y configuración publicada. CWA agrupa CWP; EWP/PWP aportan dependencias de ingeniería/materiales; IWP representa alcance ejecutable. Las relaciones son tipadas y pueden ser M:N o paralelas, no una cadena lineal obligatoria.",
            "TWP conserva significado configurado de test/turnover y puede abarcar varios sistemas/paquetes; no se afirma que sea un nivel universal de CII ni hijo obligatorio de IWP. La referencia primaria CII reconoce, entre otros, CWP/EWP/PWP/SWP/IWP. El diseño de P&Pmis debe adaptar estos conceptos a su alcance real.",
            "Scope conserva definición/baseline del paquete; Field posee release y hechos. Liberar requiere evidencia vigente de alcance, cronograma/look-ahead, materiales, revisiones técnicas, permisos, QA/HSE e issues aplicables. Field no resuelve restricciones de otros dueños localmente. AWP actual es parcial y no demuestra toda esta arquitectura.",
        ],
    )
    p(
        doc,
        "Fuente primaria consultada el 19-10-2026: Construction Industry Institute, AWP Overview. "
        + c["awp_optional_contract"]["primary_source"]["url"],
    )
    section(
        doc,
        17,
        "BIM, cantidades, APU y 4D",
        [
            "Conservar el viewer web-ifc/Three.js, explorador/propiedades, selección GUID/ExpressID, hashes de archivo, tabla de cantidades, asignación de códigos y APU. Geometría/modelo fuente no reemplaza ítem de alcance aprobado ni presupuesto Cost. raw_data.budget_item_assignment es evidencia de asignación, no publicación financiera automática.",
            "El mapping modelo-versión-objeto ↔ Scope Item ↔ actividad-versión debe ser persistente, auditable, M:N y tolerar elementos sin asociación o revisiones removidas. Una simulación 4D consume fechas aceptadas o escenario declarado y no cambia el cronograma actual.",
            "Las cantidades de diseño, ejecutadas, aceptadas, certificadas y usadas para earned value son facetas diferentes. Unidad, método, tolerancia, fuente y aprobación evitan intercambiar m², m³, horas y porcentajes. No se presenta .xls binario ni .mpp como parser validado solo porque el detector admita su extensión.",
        ],
    )
    evidence(
        doc,
        "backend/app/api/v1/routers/projects.py:1165",
        "frontend/src/lib/bimBudget.ts:115",
    )
    section(
        doc,
        18,
        "Seguridad, roles y segregación de funciones",
        [
            "Reutilizar tenant/Workspace-scoped RBAC, membresías, permisos de acción/módulo y seguridad de documentos. Listas, exportaciones y reportes validan el mismo alcance; un reporte no debe filtrar datos de otra organización ni saltar controles de fuente.",
            "Los comandos verifican estado canónico activo, ruta/autorización, condiciones vigentes, policy/config, ETag/revisión e idempotencia. Una operación staged o de larga duración revalida al promover. El acceso ADMIN no equivale a autoridad para adjudicar o aceptar una medición.",
            "SoD se aplica proporcionalmente: entrada de draft simple no exige cuatro ojos; aprobación material sí requiere actor independiente y umbrales publicados. Los roles son responsabilidades scoped propuestas, no nuevos superroles globales.",
        ],
    )
    t(
        doc,
        ["Decisión", "Prepara", "Aprueba / regla"],
        [
            [r["case"], r["prepare"], r["approve"] + ". " + r["rule"]]
            for r in c["security_and_sod"]["sod"]
        ],
        [1.05, 2.15, 3.86],
    )
    section(
        doc,
        19,
        "Reporting, período común y forecast",
        [
            "Un calendario común identifica período, inicio/cutoff, timezone, data date schedule y revisión. OPEN permite capturas; REVIEW congela candidatos; CLOSED fija fuentes aprobadas/reconciliadas; PUBLISHED emite bundle/reporte inmutable. Datos tardíos usan ajustes o reapertura autorizada con nueva revisión; el publicado anterior se conserva.",
            "Data Date P6, fecha del hecho, recepción, contabilización y cutoff no se fuerzan iguales. El forecast declara fuentes abiertas/cerradas y excepciones de corte. Mezclas no reconciliadas son visibles, no porcentajes aparentemente consistentes.",
            "Physical progress, schedule percent, cost percent, earned progress, quantity-based, weighted y milestone mantienen método/denominador/unidad/aprobación distintos. EV deriva de una baseline de medición y regla aprobadas; no de cualquier porcentaje disponible. PV usa distribución temporal aprobada y AC una base de costo declarada.",
            "ETC explicita alcance remanente, precios/productividad, cambios/riesgos y método bottom-up/estadístico. EAC=actual reconocido+ETC aprobado bajo la misma base. Compromiso neto no se vuelve a sumar si ya está dentro del ETC; contingencia/reserva se incluye una sola vez. Forecast time requiere supuestos o resultado P6 importado; no convierte una heurística CPI/SPI en nueva baseline.",
            "Look-ahead predeterminado de cuatro semanas, configurable, consume current schedule y restricciones vigentes de workfront/materiales/ingeniería/permisos/issues. Indicadores de costo/plazo, curvas, dashboards y exportaciones se reproducen por snapshot y versión de fórmula. Publication no equivale a aprobación de inversión.",
        ],
    )
    section(
        doc,
        20,
        "Checkpoints operacionales",
        [
            "Se definen seis checkpoints de ejecución, distintos de Strategic Gates/FID. Cada decisión consume aprobación de sus fuentes, mantiene bloqueos explícitos y aplica tailoring. No se ejecutó ninguno en este incremento."
        ],
    )
    for row in data["checkpoints"]:
        h(doc, row["id"] + " · " + row["name"], 2)
        p(
            doc,
            "Entrada: " + lines(row["entry"]) + ". Salida: " + lines(row["exit"]) + ".",
        )
        p(
            doc,
            "Bloquea: "
            + lines(row["blocked_when"])
            + ". Autoridad propuesta: "
            + row["decision_owner"],
        )
    section(
        doc,
        21,
        "Inventario de capacidad actual",
        [
            "El inventario distingue REUSE, EXTEND, REFACTOR, RESERVED y NOT_RELEVANT, además de madurez frontend. Hay 15 fichas backend (4 REUSE, 6 EXTEND, 3 REFACTOR, 1 RESERVED, 1 NOT_RELEVANT), 25 fichas frontend y 16 de procurement/handoff. Su granularidad se superpone: no deben sumarse como funciones implementadas.",
            "USER Project Control Manager contiene diez módulos en el catálogo: Project, Scope, Schedule, Cost, Risk, Procurement, Progress/Performance, Resource, Claim y Document. ADMIN contiene enterprise/strategy/platform; no se encontró un área operacional completa configurada para todos los dominios.",
            "Las fichas documentan modelo, API/servicio, rutas, comportamiento, límites y tests definidos con path:line. Tests inspeccionados no significa tests ejecutados hoy. Baseline/Evidence tiene JSX, pero accesibilidad por el flujo/nav vigente requiere prueba futura de navegador; se clasifica UNKNOWN en lugar de declarar pantalla completa.",
            "No se llamó GET dashboard porque el código puede generar registros si faltan. No se iniciaron servicios ni se consultó una base operacional. La auditoría no demuestra disponibilidad localhost, performance, datos reales ni conformidad legal.",
        ],
    )
    section(
        doc,
        22,
        "Gap analysis y riesgo crítico",
        [
            "La hoja 08 registra 18 brechas con severidad de diseño, objetivo, fase, evidencia y guardrail. P0 señala un prerrequisito de diseño crítico, no una clasificación de incidente real. Se destacan:"
        ],
    )
    t(
        doc,
        ["Brecha", "Hallazgo", "Remediación futura"],
        [
            [
                r["id"] + " · " + r["severity"],
                r["topic"] + ". " + r["current"],
                r["target"].replace(
                    "commissioning/handover/readiness/asset/closeout",
                    "commissioning, handover, readiness, asset transition y closeout",
                )
                + " ["
                + r["wave"]
                + "]",
            ]
            for r in data["gaps"]
        ],
        [1.0, 3.0, 3.06],
    )
    section(
        doc,
        23,
        "Deuda técnica y candidatos legacy",
        [
            "Se registran 28 unidades de revisión legacy a partir de tres vistas de evidencia, no 28 archivos muertos únicos. Clasificaciones KEEP, REQUIRES_DEPRECATION y UNKNOWN evitan confundir falta de import local con autorización de borrar. SAFE_TO_REMOVE es cero y no se elimina ningún archivo.",
            "schedule_parser.py tiene tests propios y contrato diferente de la ingestión canónica: requiere compatibilidad/consumidores externos/deprecación. Las vistas antiguas, Sidebar/AppShellCtx y tipos acoplados exigen revisión de paridad y rutas. App.tsx y el viewer grandes son capacidades activas, no basura por tamaño.",
            "projects legacy, Budget/CostCode/control accounts, APU/BIM y recursos se conservan porque tienen datos/consumidores. La transición debe ser incremental, con adaptadores, pruebas de caracterización, instrumentación de uso y rollback. No se propone refactor masivo ni borrar placeholders del catálogo aprobado.",
        ],
    )
    section(
        doc,
        24,
        "Arquitectura recomendada y cierre del incremento",
        [
            "Mantener identidad Enterprise, seguridad y autoridades estratégicas; resolver contexto/binding canónico primero. Construir Project Controls como integración de fuentes dueñas y período común, no otro megaagregado. Separar procurement de obligación legal y aislar análisis forensic de operación viva.",
            "Las olas propuestas son O1 Project Controls Foundation; O2 Procurement/Contracts/Engineering; O3 Construction Execution; O4 Claims/Commissioning/Handover/Closeout. D14 evidencia mínima se adelanta a O1; D05 rates/horas mínimas se reutilizan en O1, productividad completa en O3. Claims existente continúa disponible: O4 endurece integración, no elimina esa capacidad hasta entonces.",
            "Los ADR 86–97 registran los doce mínimos usando la numeración global del repositorio y aliases ADR-O-001–012. Permanecen PROPOSED_DESIGN. La matriz de 32 criterios declara cobertura arquitectónica, no PASS de tests de producto ni aceptación de negocio.",
            "Siguiente implementación única recomendada: WAVE O1 — PROJECT CONTROLS FOUNDATION, empezando por O1A contexto/binding canónico y período/políticas comunes; después O1B baselines/fuentes y O1C forecast/reporting. La autorización de este diseño no libera operación viva ni sustituye el signoff estratégico pendiente. Detenerse aquí y esperar autorización expresa.",
        ],
    )
    p(
        doc,
        "Paquete complementario: Implementation Roadmap.docx y Domain Map and Roadmap.xlsx. Fuente reproducible: docs/architecture/operational/architecture-data.json; inventarios y auditorías: docs/validation/operational-master-architecture-20261019/.",
    )
    return doc


def roadmap(data):
    c = data["contracts"]
    doc = start(
        "Operational Macroprocess\nImplementation Roadmap",
        "Secuencia de implementación propuesta, dependencias y criterios verificables para evolucionar la operación sin duplicar funcionalidades ni autoridades de P&Pmis.",
    )
    section(
        doc,
        1,
        "Decisión de secuencia y límites",
        [
            "El incremento presente termina con arquitectura y documentos. No ejecuta O1, no modifica módulos, no migra datos y no registra aceptación de negocio. La secuencia depende de autoridad de fuentes e identidad, no del orden del menú ni del ranking numérico.",
            "Única siguiente implementación recomendada: WAVE O1 — PROJECT CONTROLS FOUNDATION. O1A resuelve primero el enlace canónico/legacy, autorización, período común y configuración. O1B establece baselines y adapters; O1C forecast/reporting. Evitar iniciar simultáneamente 25 módulos.",
            "Aprobación del diseño habilita preparación/implementación acotada cuando se autorice, no operación productiva. Cada salida a uso vivo requiere aceptación/liberación autorizada y autorización aplicable del Project. El signoff estratégico de negocio sigue pendiente.",
        ],
    )
    t(
        doc,
        ["Ola", "Resultado objetivo", "Nivel"],
        [
            [w["id"] + " · " + w["name"], lines(w["scope"]), w["relative_complexity"]]
            for w in c["roadmap"]["waves"]
        ],
        [1.9, 4.26, 0.9],
    )
    section(
        doc,
        2,
        "Mínimo viable y prerequisitos",
        [
            "Core O1: D01 contexto; D02 Scope; D03 Schedule; D04 Cost; D06 Progress; D07 Forecast; D08 Change; D09 Risk; D10 Issue; D14 evidencia mínima; D19 Project Controls/Reporting. D05 tarifas/horas y D12 compromiso se consumen bajo contrato mínimo de lectura; su suite completa no se adelanta artificialmente.",
            "No incluye suite Field/QA/HSE, Engineering completa, sourcing/award/contratos completos, commissioning/handover, ERP/pagos ni DMS corporativo. No CPM nativo ni API/direct sync P6. Los registros operacionales existentes se preservan y exponen solo tras validar su vínculo y permisos.",
            "O1A requiere resolver proyectos sin binding, colisiones de uno/múltiples legacy, pertenencia tenant, memberships y revocación/estado. Contractor/Direct no se fuerzan al Owner FID; Owner no puede saltar condiciones de compromiso porque esté ACTIVE. El contrato no cambia retroactivamente la autorización histórica.",
        ],
    )
    for index, wave in enumerate(c["roadmap"]["waves"], 3):
        section(
            doc,
            index,
            wave["id"] + " · " + wave["name"],
            [
                "Alcance: " + lines(wave["scope"]) + ".",
                "Dependencias: " + lines(wave["dependencies"]) + ".",
            ],
        )
        pairs(
            doc,
            [
                ["Entrada", lines(wave["entry_criteria"])],
                ["Salida / evidencia futura", lines(wave["exit_criteria"])],
                ["Riesgos", lines(wave["risks"])],
                ["Reutilización", lines(wave["reuse"])],
                [
                    "Complejidad relativa",
                    wave["relative_complexity"]
                    + ". Sin fechas, horas, costo o dotación comprometidos.",
                ],
            ],
        )
        if wave["id"] == "O1":
            for subwave in wave["subwaves"]:
                h(doc, subwave["id"] + " · " + subwave["complexity"], 2)
                p(doc, subwave["purpose"] + " Salida requerida: " + subwave["exit"])
            p(
                doc,
                "Evidencia mínima de Document/version/hash se reutiliza antes de aprobar baseline en O1A/O1B. Reviews/distribución y DMS se amplían en O2. El contrato de compromiso/neteos en lectura es requisito del costo O1, aunque el hardening transaccional completo de D12 sea O2.",
            )
    section(
        doc,
        7,
        "DAG de dependencias y ruta crítica de implementación",
        [
            "El grafo contiene 25 nodos y es acíclico. Representa contratos mínimos de implementación, no prohibición de realimentación transaccional. El grafo de operación sí permite Change↔Forecast y fuente↔reporting por revisiones controladas.",
            "Cadena inicial: D01 → D14 mínimo → D02 → D03/D04 → D06/Change/Risk/Issue → D07 → D19. O2 consume O1; O3 requiere fuentes técnicas/materiales y QA/HSE aplicables antes de liberar Field; O4 requiere historia, aceptación técnica y condiciones residuales explícitas.",
            "D04 no depende del hardening completo D12 para empezar: consume la autoridad Contract/PO existente en lectura y un contrato mínimo de semántica económica. D12 transaccional O2 depende de fuentes O1. Esta separación evita el ciclo falso Cost↔Contracts.",
        ],
    )
    t(
        doc,
        ["ID / candidato", "Prerrequisitos mínimos", "Ola"],
        [
            [
                d["id"] + " " + d["name"],
                lines(data["dependency_graph"]["prerequisites"][d["id"]])
                or "Raíz: identidad/autorización existente",
                d["wave"],
            ]
            for d in data["domains"]
        ],
        [3.8, 2.71, 0.55],
    )
    section(
        doc,
        8,
        "Priorización transparente y complejidad",
        [
            "Se valoran 25 candidatos × 11 criterios en escala entera 1–5. Son juicios arquitectónicos sujetos a discusión, no prioridades aprobadas por negocio ni ROI calculado. Urgencia=3 neutral en todos: no se inventó una preferencia del usuario.",
            "Beneficios: valor, centralidad, reutilización, integración estratégica, aplicabilidad Owner/Contractor, reporting y urgencia. Penalizaciones: complejidad, riesgo de modelo y acoplamiento. Pesos propuestos suman 1; se muestran en la hoja 09 junto con 275 justificaciones y referencias.",
            "Sea B la suma ponderada de beneficios y P la de penalizaciones; WB/WP son las sumas de sus pesos. Score=100×[(B−P)−(WB−5WP)]/[4(WB+WP)]. Con pesos válidos e inputs 1–5 el resultado va de 0 a 100. Vacíos, texto, decimales o fuera de rango se marcan inválidos y no generan score/rank.",
            "Rank descendente admite empates sin fabricar desempate comercial. El score no salta prerequisites: identidad, autorización, SoD, fuente y período pueden imponer O1A aunque tenga alta complejidad. LOW/MEDIUM/HIGH/VERY_HIGH son complejidades relativas; no se traducen aquí a cronograma o presupuesto.",
        ],
    )
    t(
        doc,
        ["Criterio", "Peso", "Sentido"],
        [
            [
                r["label"],
                f"{r['weight']:.0%}",
                "Penalización" if r["direction"] == "penalty" else "Beneficio",
            ]
            for r in data["prioritization"]["criteria"]
        ],
        [4.25, 0.75, 2.06],
    )
    section(
        doc,
        9,
        "Matriz de riesgos y verificación futura",
        [
            "La implementación debe producir tests de caracterización y de contrato antes de alterar fuentes legacy. No se presentan pruebas ejecutadas en esta revisión estática. Los escenarios siguientes son criterios de salida futuros; no evidencia de funcionamiento actual."
        ],
    )
    t(
        doc,
        ["Riesgo", "Escenarios de salida requeridos"],
        [
            [
                "Identidad / acceso",
                "Tenant cruzado, mapping ausente/ambiguo, Workspace no activo, permiso revocado, concurrencia y replay; sin duplicar Project.",
            ],
            [
                "Autoridad estratégica",
                "Owner con condición de compromiso vigente, fuente stale/revocada; Contractor/Direct autorizados sin requisito Owner indebido.",
            ],
            [
                "P6 / baseline",
                "XML/XER subset, calendarios/unidades desconocidos, external links, Data Date, IDs reutilizados, diff/rechazo y rollback sin overwrite histórico.",
            ],
            [
                "Costo",
                "Contract/PO included/call-off/additive, cancelación, certificado/recepción mismo origen, accrual/reversal, monedas y reserva una vez.",
            ],
            [
                "Progress / período",
                "Métodos/denominadores distintos, aprobación independiente, corte/data date, cierre, late adjustment/reapertura y snapshot reproducible.",
            ],
            [
                "Change / Forecast",
                "Aprobar ≠ implementar, receipts parciales/fallo, idempotencia, no baseline write; EAC/ETC bajo la misma base.",
            ],
            [
                "Reporting / documentos",
                "GET sin comando oculto, mismo scope de fuentes, origen/version/hash y evidencia inaccesible no filtrada por export.",
            ],
            [
                "Compatibilidad",
                "Rutas actuales y fuentes estratégicas sin regresión; formatos binarios no soportados rechazados explícitamente; no limpieza sin deprecación.",
            ],
        ],
        [1.6, 5.46],
    )
    section(
        doc,
        10,
        "Decisiones arquitectónicas y criterios de aceptación",
        [
            "ADR86–97 están PROPOSED_DESIGN, con aliases del prompt ADR-O-001–012 y evidencia path:line. No existe aprobación de negocio registrada para estas decisiones. Workbook 12_Decisions añade la clasificación de las 16 capacidades procurement/handoff.",
            "Los 32 criterios siguientes se cubren como diseño documentado. PASS significa arquitectura definida y revisada en matrices/documentos, no software implementado, test de producto aprobado, validación jurídica ni signoff de negocio. Los tres criterios de protección se verifican además mediante huellas del producto y ausencia de migraciones.",
        ],
    )
    t(
        doc,
        ["ID / criterio", "Evidencia de diseño", "Estado"],
        [
            [r["id"] + " " + r["criterion"], r["evidence"], "PASS de diseño"]
            for r in data["acceptance"]
        ],
        [3.5, 2.51, 1.05],
    )
    doc.add_page_break()
    section(
        doc,
        11,
        "Único siguiente prompt recomendado",
        [
            "WAVE O1 — PROJECT CONTROLS FOUNDATION.",
            "Implementar O1 comenzando exclusivamente por O1A: reutilizar Enterprise Workspace PROJECT y sus lifecycle/RBAC; diseñar y validar binding explícito con el agregado operacional legacy, sin crear otro Project master ni asumir IDs iguales; mantener las tres rutas de autorización y condiciones actuales. Reutilizar evidencia/versiones y establecer período/políticas comunes ADMIN/USER, con SoD, idempotencia, concurrencia y auditoría. Preservar toda lógica estratégica e historia aprobada.",
            "Cerrar O1A con pruebas de identidad, acceso, colisiones, estado, currentness, condiciones y ciclo de período; solo entonces habilitar O1B baselines/adapters de Scope/Schedule/Cost/Progress/Change/Risk/Issue y O1C Forecast/Reporting bajo alcance autorizado. P6 sigue externo manual XML/XER, sin CPM/API/sync. No ampliar a O2–O4. Entregar evidencia, inventario de reutilización y reporte técnico; no declarar liberación ni signoff sin aprobación autorizada.",
            "STOP: la tarea actual termina en este diseño. La siguiente implementación requiere instrucción expresa.",
        ],
    )
    p(
        doc,
        "Referencias principales: Master Architecture.docx; Domain Map and Roadmap.xlsx; docs/architecture/operational/architecture-data.json; inventarios y auditorías del corte 20261019. Oracle P6 y CII se citan en el informe maestro para las fronteras XML/XER y AWP.",
    )
    return doc


def save(doc, output):
    stream = io.BytesIO()
    doc.save(stream)
    with ZipFile(REFERENCE) as src, ZipFile(stream) as generated:
        body = generated.read("word/document.xml")
        footer = src.read("word/footer1.xml").replace(
            b"18 octubre 2026", b"19 octubre 2026"
        )
        assert footer != src.read("word/footer1.xml")
        with ZipFile(output, "w") as dst:
            for info in src.infolist():
                dst.writestr(
                    info,
                    body
                    if info.filename == "word/document.xml"
                    else footer
                    if info.filename == "word/footer1.xml"
                    else src.read(info.filename),
                )
    with ZipFile(REFERENCE) as src, ZipFile(output) as dst:
        assert src.namelist() == dst.namelist()
        changed = [name for name in src.namelist() if src.read(name) != dst.read(name)]
        assert set(changed) == {"word/document.xml", "word/footer1.xml"}
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        a, b = [etree.fromstring(z.read("word/document.xml")) for z in (src, dst)]
        for xpath in ("w:body/w:sectPr", "w:body/w:p"):
            assert etree.tostring(a.find(xpath, ns)) == etree.tostring(
                b.find(xpath, ns)
            )
    return {
        "path": str(output),
        "sha256": digest(output),
        "bytes": output.stat().st_size,
        "parts_preserved": 17,
        "changed_parts": changed,
        "visual_review": "PENDING",
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--authoring-marker-confirmed", action="store_true")
    parser.add_argument("--validate-only", action="store_true")
    args = parser.parse_args()
    data = json.loads(DATA.read_text(encoding="utf-8"))
    assert digest(REFERENCE) == REFERENCE_SHA
    assert len(data["domains"]) == 25 and len(data["acceptance"]) == 32
    assert len(data["workbook_sheets"]) == 12 and data["safe_to_remove_count"] == 0
    # Verify expected input structure before authoring marker is requested.
    assert len(data["contracts"]["strategic_operational_handoff"]["items"]) == 15
    if args.validate_only:
        print("ARCHITECTURE_DATA_AND_TEMPLATE_PASS")
        return
    assert args.authoring_marker_confirmed, (
        "Root must confirm create marker DOCX count=2"
    )
    OUT.mkdir(parents=True, exist_ok=True)
    QA.mkdir(parents=True, exist_ok=True)
    results = [
        save(builder(data), OUT / name)
        for builder, name in zip((master, roadmap), NAMES)
    ]
    assert digest(REFERENCE) == REFERENCE_SHA
    manifest = {
        "source": str(REFERENCE),
        "source_sha256": REFERENCE_SHA,
        "data_sha256": digest(DATA),
        "outputs": results,
        "production_changes": 0,
        "business_approvals": 0,
    }
    (QA / "documents-build-manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(results, ensure_ascii=False))


if __name__ == "__main__":
    main()
