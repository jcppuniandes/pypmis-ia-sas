"""Read-only native XLSX verification. This module never saves a workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from pathlib import Path
from zipfile import ZipFile

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "docs/architecture/operational/architecture-data.json"
OUTPUT = (
    ROOT
    / "outputs/operational-master-architecture-20261019"
    / "P&Pmis_Operational_Macroprocess_Domain_Map_and_Roadmap.xlsx"
)
REQUIRED_SHEETS = [
    "01_Domain_Map",
    "02_Process_Map",
    "03_Admin_User_Matrix",
    "04_Owner_Contractor_Matrix",
    "05_Data_Ownership",
    "06_Integration_Matrix",
    "07_Current_Capabilities",
    "08_Gap_Analysis",
    "09_Wave_Prioritization",
    "10_Dependency_Map",
    "11_Legacy_Candidates",
    "12_Decisions",
]
ERROR_PATTERN = re.compile(
    r"^#(?:REF!|DIV/0!|VALUE!|NAME\?|N/A|NUM!|NULL!|SPILL!|CALC!)$"
)


def presentation_value(value, compact_references=False):
    if not isinstance(value, str) or "\n" not in value:
        return value
    lines = [line.strip() for line in value.split("\n")]
    if compact_references and all(
        re.fullmatch(r"(?:backend|frontend|scripts|docs)/.+:\d+", line)
        for line in lines
    ):
        return "; ".join(lines)
    return (
        ", ".join(lines)
        if all(re.fullmatch(r"[A-Z][A-Z0-9_-]*\d[A-Z0-9_-]*", line) for line in lines)
        else value
    )


def score(row: dict, criteria: list[dict]) -> float | None:
    weights = [item["weight"] for item in criteria]
    if not all(isinstance(weight, (int, float)) and weight > 0 for weight in weights):
        return None
    if not math.isclose(sum(weights), 1, abs_tol=1e-9):
        return None
    scores = row["scores"]
    for item in criteria:
        value = scores.get(item["key"])
        if not isinstance(value, int) or isinstance(value, bool) or not 1 <= value <= 5:
            return None
    total = sum(
        scores[item["key"]]
        * item["weight"]
        * (1 if item["direction"] == "benefit" else -1)
        for item in criteria
    )
    minimum = sum(
        item["weight"] * (1 if item["direction"] == "benefit" else -5)
        for item in criteria
    )
    return 100 * (total - minimum) / (4 * sum(weights))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--visual-reviewed", action="store_true")
    args = parser.parse_args()
    data = json.loads(DATA.read_text(encoding="utf-8"))
    formulas = openpyxl.load_workbook(OUTPUT, data_only=False)
    values = openpyxl.load_workbook(OUTPUT, data_only=True)
    checks: list[dict] = []

    def check(name: str, passed: bool, detail: str = "") -> None:
        checks.append({"name": name, "pass": bool(passed), "detail": detail})

    check("exact_12_sheet_names_and_order", formulas.sheetnames == REQUIRED_SHEETS)
    check(
        "all_sheets_gridlines_hidden",
        all(sheet.sheet_view.showGridLines is False for sheet in formulas),
    )
    check("12_native_tables", sum(len(sheet.tables) for sheet in formulas) == 12)
    check(
        "no_merged_working_cells",
        all(not sheet.merged_cells.ranges for sheet in formulas),
    )
    check(
        "row_heights_within_excel_limit",
        all(
            (dimension.height or 0) <= 409.5
            for sheet in formulas
            for dimension in sheet.row_dimensions.values()
        ),
    )
    mismatches = []
    for spec in data["workbook_sheets"]:
        if spec["name"] == REQUIRED_SHEETS[8]:
            continue
        sheet = values[spec["name"]]
        header_row = spec.get("header_row", 7)
        for col, expected in enumerate(spec["headers"], start=1):
            if sheet.cell(header_row, col).value != expected:
                mismatches.append(
                    f"{sheet.title}!{sheet.cell(header_row, col).coordinate}: header mismatch"
                )
        for row_index, row in enumerate(spec["rows"], start=header_row + 1):
            for col, expected in enumerate(row, start=1):
                expected = presentation_value(
                    expected, spec["name"] == "07_Current_Capabilities" and col == 9
                )
                actual = sheet.cell(row_index, col).value
                if actual != expected and not (actual is None and expected == ""):
                    mismatches.append(
                        f"{sheet.title}!{sheet.cell(row_index, col).coordinate}: value mismatch"
                    )
    check(
        "generic_source_rows_and_headers_preserved",
        not mismatches,
        "; ".join(mismatches[:10]),
    )

    unexpected_errors = []
    formula_count = 0
    for source in formulas:
        cached = values[source.title]
        for row in source:
            for cell in row:
                formula_count += int(cell.data_type == "f")
                cached_value = cached[cell.coordinate].value
                if isinstance(cached_value, str) and ERROR_PATTERN.fullmatch(
                    cached_value
                ):
                    unexpected_errors.append(
                        f"{source.title}!{cell.coordinate}:{cached_value}"
                    )
    check(
        "no_cached_formula_errors",
        not unexpected_errors,
        "; ".join(unexpected_errors[:10]),
    )

    p = data["prioritization"]
    source = formulas[REQUIRED_SHEETS[8]]
    cached = values[source.title]
    weights = [cached.cell(5, col).value for col in range(3, 14)]
    check(
        "11_numeric_positive_weights_sum_1",
        all(isinstance(v, (int, float)) and v > 0 for v in weights)
        and math.isclose(sum(weights), 1, abs_tol=1e-9),
    )
    check(
        "positive_weights_control_formula",
        source["B7"].data_type == "f" and "SUM(C5:M5)" in source["B7"].value,
    )
    check("weights_control_cached_valid", cached["B7"].value == "VALID")
    expected_scores = [score(row, p["criteria"]) for row in p["rows"]]
    score_mismatches = []
    rank_mismatches = []
    raw_mismatches = []
    rationale_mismatches = []
    for index, row in enumerate(p["rows"], start=10):
        if (
            cached.cell(index, 1).value != row["domain_id"]
            or cached.cell(index, 2).value != row["domain_name"]
        ):
            raw_mismatches.append(f"row {index}: identity")
        domain = next(
            item for item in data["domains"] if item["id"] == row["domain_id"]
        )
        if (
            cached.cell(index, 20).value
            != domain["rationale"] + "\nDetalle íntegro en nota."
        ):
            rationale_mismatches.append(f"row {index}: concise summary mismatch")
        note = source.cell(index, 20).comment
        if note is None or note.text != row["rationale"]:
            rationale_mismatches.append(
                f"row {index}: complete 11-criterion rationale note mismatch"
            )
        for col, criterion in enumerate(p["criteria"], start=3):
            if cached.cell(index, col).value != row["scores"].get(criterion["key"]):
                raw_mismatches.append(f"row {index}, column {col}: raw score")
        expected = expected_scores[index - 10]
        actual = cached.cell(index, 17).value
        if expected is None:
            if actual not in {None, ""}:
                score_mismatches.append(
                    f"row {index}: missing/invalid input should not calculate a score"
                )
        elif not isinstance(actual, (int, float)) or not math.isclose(
            actual, expected, abs_tol=1e-8
        ):
            score_mismatches.append(f"row {index}: expected {expected}, got {actual}")
        if expected is not None:
            expected_rank = 1 + sum(
                other is not None and other > expected + 1e-8
                for other in expected_scores
            )
            if cached.cell(index, 18).value != expected_rank:
                rank_mismatches.append(
                    f"row {index}: expected rank {expected_rank}, got {cached.cell(index, 18).value}"
                )
        for col in range(14, 19):
            if source.cell(index, col).data_type != "f":
                score_mismatches.append(f"row {index}, column {col}: formula missing")
    check(
        "raw_scores_exact_and_integer",
        not raw_mismatches,
        "; ".join(raw_mismatches[:10]),
    )
    check(
        "all_25_detailed_rationale_notes_exact",
        not rationale_mismatches,
        "; ".join(rationale_mismatches[:10]),
    )
    check(
        "scores_match_independent_python",
        not score_mismatches,
        "; ".join(score_mismatches[:10]),
    )
    check(
        "ranks_match_descending_shared_ties",
        not rank_mismatches,
        "; ".join(rank_mismatches[:10]),
    )
    check(
        "score_table_frozen_at_C10",
        source.freeze_panes == "C10",
        str(source.freeze_panes),
    )
    validations = list(source.data_validations.dataValidation)
    check(
        "whole_number_1_to_5_validation",
        any(
            v.type == "whole"
            and str(v.formula1) == "1"
            and str(v.formula2) == "5"
            and str(v.sqref) == f"C10:M{9 + len(p['rows'])}"
            for v in validations
        ),
    )
    check(
        "dynamic_invalid_input_conditional_formats",
        len(source.conditional_formatting) >= 2,
    )
    with ZipFile(OUTPUT) as archive:
        names = archive.namelist()
        check(
            "no_external_workbook_links",
            not any(name.startswith("xl/externalLinks/") for name in names),
        )
        check(
            "no_macros_or_embedded_executables",
            not any(name.endswith(("vbaProject.bin", ".exe")) for name in names),
        )
    visual_baseline = (
        ROOT
        / "outputs/.qa/operational-master-architecture-20261019/workbook/layout-before-reference-widths.xlsx"
    )
    if visual_baseline.exists():
        previous = openpyxl.load_workbook(visual_baseline, data_only=False)
        preserved = []
        for sheet in formulas:
            if sheet.title == "07_Current_Capabilities":
                continue
            before = previous[sheet.title]

            def signature(worksheet):
                return {
                    "cells": [
                        (
                            cell.coordinate,
                            cell.value,
                            cell.data_type,
                            cell.style_id,
                            cell.comment.text if cell.comment else None,
                        )
                        for row in worksheet
                        for cell in row
                        if cell.value is not None or cell.comment
                    ],
                    "rows": [
                        (key, value.height, value.hidden)
                        for key, value in worksheet.row_dimensions.items()
                    ],
                    "columns": [
                        (key, value.width, value.hidden)
                        for key, value in worksheet.column_dimensions.items()
                    ],
                    "panes": worksheet.freeze_panes,
                    "gridlines": worksheet.sheet_view.showGridLines,
                    "tables": [
                        (table.name, table.ref) for table in worksheet.tables.values()
                    ],
                }

            if signature(sheet) != signature(before):
                preserved.append(sheet.title)
        check(
            "11_unchanged_sheets_render_reuse_proven",
            not preserved,
            ", ".join(preserved),
        )
        with ZipFile(visual_baseline) as before, ZipFile(OUTPUT) as after:
            check(
                "style_definitions_unchanged_for_reused_renders",
                before.read("xl/styles.xml") == after.read("xl/styles.xml"),
            )
    check(
        "every_sheet_visually_reviewed",
        args.visual_reviewed,
        "Flag may only be supplied after human visual inspection of saved renders.",
    )
    result = {
        "operation": "READ_ONLY_NATIVE_XLSX_VERIFICATION_NO_SAVE",
        "output_sha256": hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
        "data_sha256": hashlib.sha256(DATA.read_bytes()).hexdigest(),
        "checks": checks,
        "passed": sum(item["pass"] for item in checks),
        "failed": sum(not item["pass"] for item in checks),
        "formula_cells_observed": formula_count,
        "engine_limitation": "Artifact Tool recalculation plus saved OOXML/cache inspection. No native Microsoft Excel application recalculation was performed.",
    }
    proof = (
        ROOT
        / "outputs/.qa/operational-master-architecture-20261019/workbook/native-verification.json"
    )
    proof.parent.mkdir(parents=True, exist_ok=True)
    proof.write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if result["failed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
