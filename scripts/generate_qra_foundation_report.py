"""Create the QRA Foundation technical report from its reviewed Markdown source."""

from __future__ import annotations

import re
from pathlib import Path

import generate_cpig_gap_review_report as base
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "81-qra-foundation-release.md"
QA = ROOT / "outputs" / ".qa" / "qra-foundation"
OUTPUT = ROOT / "outputs" / "Informe_Tecnico_QRA_Foundation_2026-10-09.docx"
LOGO = ROOT / "frontend" / "public" / "pypmis-construction-ai-logo.png"
REPORT_DATE = "2026-10-09"


def release_status() -> str:
    text = SOURCE.read_text(encoding="utf-8")
    match = re.search(r"^\*\*Estado actual:\*\*\s+(.+?)$", text, re.MULTILINE)
    value = match.group(1).strip() if match else "VALIDACIÓN EN CURSO"
    return "RELEASE VALIDADO" if value == "QRA_SOURCE_READINESS_AND_METHODOLOGY_FOUNDATION_RELEASE_VALIDATED" else value


def metadata(document: Document) -> None:
    core = document.core_properties
    core.title = "Informe técnico QRA Foundation"
    core.subject = "QRA Source Readiness and Methodology Foundation"
    core.author = "P&Pmis Construction AI"
    core.keywords = "QRA, source readiness, methodology, cost, schedule, governance"
    core.comments = "Preparación gobernada; no contiene simulación probabilística."


def cover(document: Document, status: str) -> None:
    logo = document.add_paragraph()
    logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo.paragraph_format.space_before = Pt(28)
    if LOGO.exists():
        shape = logo.add_run().add_picture(str(LOGO), width=Inches(1.3))
        base.set_alt_text(shape, "P&Pmis Construction AI", "Logotipo azul de P&Pmis Construction AI.")
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(22)
    title.add_run("Informe técnico\nQRA Source Readiness and Methodology Foundation")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.set_run_font(subtitle.add_run("P&Pmis Construction AI"), size=13, color=base.BLACK)
    details = document.add_paragraph()
    details.alignment = WD_ALIGN_PARAGRAPH.CENTER
    details.paragraph_format.space_before = Pt(22)
    base.add_inline(details, f"Capital Portfolio and Investment Governance | {REPORT_DATE}", size=10)
    label = document.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.space_before = Pt(38)
    base.set_run_font(label.add_run("ESTADO DE VALIDACIÓN"), size=8, bold=True, color=base.MUTED)
    value = document.add_paragraph()
    value.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.set_run_font(value.add_run(status.replace("_", " ")), size=11, bold=True, color=base.BLACK)
    document.add_page_break()


def document_control(document: Document, status: str) -> None:
    document.add_heading("Control del documento", level=1)
    base.add_table(document, [
        ["Campo", "Valor"],
        ["Producto", "P&Pmis Construction AI"],
        ["Macroproceso", "Capital Portfolio and Investment Governance"],
        ["Incremento", "QRA Source Readiness and Methodology Foundation"],
        ["Rama y commit local", "main / 8f646ba"],
        ["Baseline de referencia", "main / 9cb0969; cambios previos preservados"],
        ["Migración", "20261009_0067 sobre 20261008_0066"],
        ["Fecha de corte", REPORT_DATE],
        ["Estado", status],
    ])
    document.add_heading("Resumen para decisión", level=1)
    paragraph = document.add_paragraph()
    base.add_inline(paragraph, "Este informe documenta la preparación gobernada de fuentes de costo y cronograma, metodología, rangos y definición de modelo. El alcance no contiene Monte Carlo, QraRun, percentiles, contingencias ni fechas probabilísticas. Las comprobaciones pendientes se mantienen visibles como validación en curso.")
    document.add_heading("Contenido", level=1)
    for item in (
        "Baseline, ADR, autoridad y reutilización",
        "Cost Source Readiness y Schedule Source Readiness",
        "Rangos, distribuciones, correlaciones y metodología",
        "QraModelDefinition, SoD, staleness y providers",
        "ADMIN, USER, seguridad, persistencia, API y navegación",
        "Pruebas, performance, deuda técnica, contrato futuro y cierre",
    ):
        paragraph = document.add_paragraph(style="List Bullet")
        base.add_inline(paragraph, item)
    document.add_page_break()


def headers_and_footers(document: Document) -> None:
    for section in document.sections:
        section.different_first_page_header_footer = True
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        base.set_run_font(header.add_run("P&Pmis Construction AI  |  QRA Foundation"), size=7.5, bold=True, color=base.BLACK)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        base.set_run_font(footer.add_run(f"Informe técnico | {REPORT_DATE} | Página "), size=7.5, color=base.MUTED)
        base.add_field(footer, "PAGE")
        base.set_run_font(footer.add_run(" de "), size=7.5, color=base.MUTED)
        base.add_field(footer, "NUMPAGES")


def format_tables(document: Document) -> None:
    for table in document.tables:
        table.autofit = False
        borders = table._tbl.tblPr.find(qn("w:tblBorders"))
        if borders is not None:
            for border in borders:
                border.set(qn("w:color"), "D9D9D9")
        for row_index, row in enumerate(table.rows):
            props = row._tr.get_or_add_trPr()
            if props.find(qn("w:cantSplit")) is None:
                props.append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.05
                    if row_index == 0:
                        paragraph.paragraph_format.keep_with_next = True
                    for run in paragraph.runs:
                        run.font.size = Pt(8.5)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    QA.mkdir(parents=True, exist_ok=True)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    base.SOURCE = SOURCE
    base.DIAGRAMS = []
    document = Document()
    base.configure_document(document)
    base.set_language(document)
    for style_name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        properties = document.styles[style_name]._element.find(qn("w:pPr"))
        if properties is not None:
            for border in properties.findall(qn("w:pBdr")):
                properties.remove(border)
    metadata(document)
    status = release_status()
    cover(document, status)
    document_control(document, status)
    source_start = len(document.paragraphs)
    base.add_source_content(document, [], page_break_heading_prefixes=(), skip_preface=True)
    for index, paragraph in enumerate(document.paragraphs):
        if paragraph.style.name.startswith("Heading"):
            cleaned = re.sub(r"[^\w\s]", " ", paragraph.text).replace("_", " ")
            paragraph.text = re.sub(r"\s+", " ", cleaned).strip()
        elif index >= source_start and paragraph.style.name in {"Normal", "List Bullet", "List Number"}:
            for run in paragraph.runs:
                run.font.size = Pt(10.5)
    format_tables(document)
    headers_and_footers(document)
    base.update_field_settings(document)
    document.save(OUTPUT)
    print(f"created={OUTPUT}")
    section_count = len(re.findall(r"^## \d+\. ", SOURCE.read_text(encoding="utf-8"), re.MULTILINE))
    print(f"sections={section_count}")
    print(f"status={status}")


if __name__ == "__main__":
    main()
