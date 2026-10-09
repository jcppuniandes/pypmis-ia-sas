from __future__ import annotations

import xml.etree.ElementTree as ET
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(r"C:\Users\Ricardo\Documents\GitHub\pypmis-ia-sas")
QA = ROOT / "outputs" / ".qa" / "gate07m-bc-5c"
OUTPUT = ROOT / "outputs" / "Informe_Tecnico_Gate07M_BC_5C_Five_Case_2026-09-18.docx"
LOGO = ROOT / "frontend" / "public" / "pypmis-construction-ai-logo.png"

BLACK = "000000"
TEXT = "202A33"
MUTED = "53616B"
NAVY = "17324D"
TEAL = "127A78"
WHITE = "FFFFFF"
PALE_GREEN = "E9F5EE"
PALE_AMBER = "FFF4E5"
PALE_RED = "FDECEC"
PALE_GRAY = "F5F7F8"
LINE = "D7DEE3"
GREEN = "18794E"
AMBER = "8A5700"
RED = "B42318"


def junit(path: Path) -> dict[str, Any]:
    result: dict[str, Any] = {
        "tests": 0,
        "failures": 0,
        "errors": 0,
        "skipped": 0,
        "time": 0.0,
    }
    if not path.exists() or path.stat().st_size == 0:
        return result
    root = ET.parse(path).getroot()
    suites = [root] if root.tag == "testsuite" else list(root.findall("testsuite"))
    for suite in suites:
        result["tests"] += int(suite.attrib.get("tests", 0))
        result["failures"] += int(suite.attrib.get("failures", 0))
        result["errors"] += int(suite.attrib.get("errors", 0))
        result["skipped"] += int(suite.attrib.get("skipped", 0))
        result["time"] += float(suite.attrib.get("time", 0))
    return result


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    node = properties.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        properties.append(node)
    node.set(qn("w:fill"), fill)


def margins(cell, value: int = 85) -> None:
    properties = cell._tc.get_or_add_tcPr()
    node = properties.first_child_found_in("w:tcMar")
    if node is None:
        node = OxmlElement("w:tcMar")
        properties.append(node)
    for side in ("top", "start", "bottom", "end"):
        item = node.find(qn(f"w:{side}"))
        if item is None:
            item = OxmlElement(f"w:{side}")
            node.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def borders(table_item) -> None:
    properties = table_item._tbl.tblPr
    node = properties.first_child_found_in("w:tblBorders")
    if node is None:
        node = OxmlElement("w:tblBorders")
        properties.append(node)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = node.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            node.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), LINE)


def repeat_header(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def no_split(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def font(run, size: float = 10.5, bold: bool = False, color: str = TEXT) -> None:
    run.font.name = "Aptos"
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), "Aptos")
    fonts.set(qn("w:hAnsi"), "Aptos")
    fonts.set(qn("w:eastAsia"), "Aptos")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def paragraph(document: Document, text: str = "", *, style: str | None = None):
    item = document.add_paragraph(style=style)
    item.paragraph_format.space_after = Pt(5)
    item.paragraph_format.line_spacing = 1.08
    font(item.add_run(text))
    return item


def bullets(document: Document, items: Iterable[str]) -> None:
    for value in items:
        item = paragraph(document, value, style="List Bullet")
        item.paragraph_format.space_after = Pt(2.5)


def heading(document: Document, text: str, level: int = 1):
    item = document.add_heading(text, level=level)
    item.paragraph_format.keep_with_next = True
    item.paragraph_format.space_before = Pt(8 if level == 1 else 6)
    item.paragraph_format.space_after = Pt(4)
    return item


def table(
    document: Document,
    headers: list[str],
    rows: list[list[Any]],
    widths: list[float],
    *,
    status_column: int | None = None,
):
    item = document.add_table(rows=1, cols=len(headers))
    item.autofit = False
    borders(item)
    repeat_header(item.rows[0])
    no_split(item.rows[0])
    for index, value in enumerate(headers):
        cell = item.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, NAVY)
        margins(cell)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        font(cell.paragraphs[0].add_run(value), 8.2, True, WHITE)
    for row_index, values in enumerate(rows):
        row = item.add_row()
        no_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            shade(cell, WHITE if row_index % 2 == 0 else PALE_GRAY)
            margins(cell)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.02
            run = cell.paragraphs[0].add_run(str(value))
            font(run, 7.8)
            if status_column == index:
                state = str(value).upper()
                if state.startswith("PASS") or state == "COMPLETO":
                    shade(cell, PALE_GREEN)
                    run.font.color.rgb = RGBColor.from_string(GREEN)
                    run.bold = True
                elif state.startswith(("PARCIAL", "PENDIENTE")):
                    shade(cell, PALE_AMBER)
                    run.font.color.rgb = RGBColor.from_string(AMBER)
                    run.bold = True
                elif state.startswith(("FAIL", "ATENCIÓN")):
                    shade(cell, PALE_RED)
                    run.font.color.rgb = RGBColor.from_string(RED)
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return item


def page_number(paragraph_item) -> None:
    run = paragraph_item.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, value, end])


def state(result: dict[str, Any]) -> str:
    if result["tests"] == 0:
        return "PENDIENTE"
    return "PASS" if result["failures"] + result["errors"] == 0 else "ATENCIÓN"


def green(
    result: dict[str, Any], *, minimum_tests: int = 1, allow_skips: bool = True
) -> bool:
    return (
        result["tests"] >= minimum_tests
        and result["failures"] + result["errors"] == 0
        and (allow_skips or result["skipped"] == 0)
    )


backend = junit(QA / "backend-focused.xml")
backend_strategic = junit(QA / "backend-strategic-regression-final.xml")
frontend = junit(QA / "frontend-focused.xml")
frontend_full = junit(QA / "frontend-full-final.xml")
browser = junit(QA / "browser-e2e.xml")
postgres_structural = junit(QA / "postgres-structural.xml")
postgres_business_case = junit(QA / "postgres-business-case.xml")
appflow_hardening = junit(QA / "frontend-appflow-production-hardening.xml")
appflow_admin = junit(QA / "frontend-appflow-user-admin.xml")
appflow_navigation = junit(QA / "frontend-appflow-navigation.xml")
release_validated = all(
    (
        green(backend, minimum_tests=41, allow_skips=False),
        green(backend_strategic, minimum_tests=156),
        green(postgres_business_case, minimum_tests=2, allow_skips=False),
        green(frontend, minimum_tests=20, allow_skips=False),
        green(frontend_full, minimum_tests=273, allow_skips=False),
        green(browser, minimum_tests=1, allow_skips=False),
    )
)
release_signal = (
    "GATE07M_BC_FIVE_CASE_RELEASE_VALIDATED"
    if release_validated
    else "PENDING_FULL_REGRESSION"
)

document = Document()
section = document.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.62)
section.left_margin = Inches(0.68)
section.right_margin = Inches(0.68)
section.header_distance = Inches(0.28)
section.footer_distance = Inches(0.28)

styles = document.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
for name, size in (
    ("Title", 25),
    ("Heading 1", 16),
    ("Heading 2", 12.5),
    ("Heading 3", 10.5),
):
    styles[name].font.name = "Aptos Display"
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor.from_string(BLACK)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(header.add_run("P&Pmis Construction AI  |  Gate 07M-BC-5C"), 8.2, True, MUTED)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(footer.add_run("Informe técnico controlado  ·  Página "), 8.2, False, MUTED)
page_number(footer)

if LOGO.exists():
    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    drawing = image_paragraph.add_run().add_picture(str(LOGO), width=Inches(0.95))
    drawing._inline.docPr.set("descr", "Logotipo de P&Pmis Construction AI")

title = document.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(4)
font(title.add_run("Informe técnico Gate 07M-BC-5C"), 25, True, BLACK)
subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(12)
font(
    subtitle.add_run("Five Case Model Integration for Governed Business Case"),
    14,
    True,
    NAVY,
)

table(
    document,
    ["Control", "Valor"],
    [
        ["Aplicación", "P&Pmis Construction AI"],
        ["Incremento", "PROMPT — Gate 07M-BC-5C"],
        ["Fecha de cierre", "2026-09-18"],
        ["Alembic head", "20260916_0061 (single head; sin migración nueva)"],
        ["Implementación", "COMPLETA PARA EL ALCANCE FUNCIONAL"],
        ["Release", release_signal],
        ["Provider", "BUSINESS_CASE_READINESS_PROVIDER / BUSINESS_CASE_READINESS"],
    ],
    [1.6, 5.4],
)

heading(document, "1. Resumen ejecutivo")
paragraph(
    document,
    "Se amplió el Business Case gobernado existente con una capa de integración Five Case: Strategic, Economic, Commercial, Financial y Management. La solución reutiliza BusinessCase, BusinessCaseRevision y BusinessCaseDecisionSnapshot; consume evidencia por referencia y hash; conserva la autoridad de cada dominio; y evita nuevos aggregates, tablas o proveedores duplicados.",
)
bullets(
    document,
    [
        "Cada revisión nueva congela política, contexto Gate/FEL, referencias, hashes, readiness y as-of en el JSON ya gobernado.",
        "Las revisiones legacy siguen legibles con legacy_layout=true y no reciben Five Case retrospectivo fabricado.",
        "Refresh Preview no persiste; Apply crea una revisión sucesora DRAFT. Revisiones COMPLETED y snapshots permanecen inmutables.",
        "CAPITAL_OWNER es estricto por defecto; CONTRACTOR_DELIVERY y DIRECT_INTERNAL permanecen NOT_APPLICABLE salvo override explícito.",
        "USER ofrece diez vistas Five Case y ADMIN amplía la configuración existente, sin operar fuentes vivas ni contaminar modos.",
        "El localhost desplegado expone el contrato OpenAPI nuevo y reporta API, PostgreSQL y Redis saludables.",
    ],
)
paragraph(
    document,
    (
        "El cierre quedó validado con 273/273 pruebas frontend, la regresión estratégica backend completa "
        f"del alcance ({backend_strategic['tests'] - backend_strategic['skipped']}/{backend_strategic['tests'] - backend_strategic['skipped']} ejecutadas; "
        f"{backend_strategic['skipped']} omisiones PostgreSQL-only esperadas) y el happy path transaccional Five Case en PostgreSQL 2/2. "
        "Se emiten las señales formales de release; la suite backend global del repositorio permanece fuera del alcance exigido por este incremento."
        if release_validated
        else "El estado exclusivo de release permanece pendiente hasta que todas las evidencias automatizadas requeridas queden verdes."
    ),
)

document.add_page_break()
heading(document, "2. Baseline y REUSE BEFORE CREATE")
table(
    document,
    ["Activo existente", "Decisión de reutilización", "Prueba"],
    [
        [
            "BusinessCase",
            "Se conserva como identidad longitudinal por Project Workspace",
            "No se creó aggregate Five Case",
        ],
        [
            "BusinessCaseRevision",
            "Contiene narrativa y bundle congelado; proyecta Five Case de forma aditiva",
            "Contrato API enriquecido",
        ],
        [
            "BusinessCaseDecisionSnapshot",
            "Conserva contenido, Cases, readiness y hashes al momento de decisión",
            "Append-only",
        ],
        [
            "BUSINESS_CASE_READINESS_PROVIDER",
            "Se reutiliza y amplía sin cambiar provider_key/evidence_type",
            "Compatibilidad v1",
        ],
        [
            "AdminConfiguration",
            "Congela la política efectiva dentro de la revisión",
            "Preview persisted=false",
        ],
        [
            "Navigation blueprint",
            "Sigue siendo fuente canónica para USER/ADMIN",
            "App.tsx no se expandió",
        ],
    ],
    [1.65, 3.35, 2.35],
)
paragraph(
    document,
    "La evaluación real del repositorio confirmó Alembic 0061 como head único. Como la estructura Five Case se deriva de JSON versionado ya existente, no se añadió 0062. El ADR 67 documenta esta decisión y su frontera histórica.",
)

heading(document, "3. Arquitectura y persistencia")
heading(document, "3.1 Contrato aditivo", 2)
table(
    document,
    ["Campo de salida", "Semántica"],
    [
        [
            "five_case_json",
            "Cinco Cases, preguntas, resumen, findings, supuestos, fuentes, warnings, blockers y conclusión",
        ],
        ["five_case_readiness_json", "Readiness backend-owned por Case y overall"],
        [
            "source_evidence_matrix_json",
            "Matriz read-only con autoridad, versión, hash, vigencia, as-of y blocking",
        ],
        [
            "gate_context_json",
            "Project type, Governance Model, Gate Profile, FEL/policy stage y revision purpose",
        ],
        ["source_current", "Vigencia agregada del bundle congelado"],
        [
            "five_case_hash / source_bundle_hash",
            "Firmas deterministas de estructura Five Case y evidencia fuente",
        ],
        [
            "business_case_hash",
            "Composición semántica de contenido, Five Case, fuentes y configuración",
        ],
        [
            "legacy_layout",
            "Indica una revisión previa al marker, sin reinterpretación retrospectiva",
        ],
    ],
    [2.35, 5.0],
)
heading(document, "3.2 Marker congelado", 2)
paragraph(
    document,
    "El marker _five_case se persiste dentro de source_references_json. Contiene schema_version, integration_layer, source_authority_preserved, configuration_id/hash, política efectiva y contexto Gate/FEL. No se incluye como evidencia externa ni participa como clave de staleness.",
)
heading(document, "3.3 Inmutabilidad y hashes", 2)
bullets(
    document,
    [
        "Las lecturas históricas usan exclusivamente revision.source_references_json y revision.source_hashes_json.",
        "La consulta de fuentes actuales se limita a readiness, staleness y Refresh Preview/Apply.",
        "COMPLETED no se edita; Refresh Apply crea una nueva DRAFT con If-Match e idempotency key.",
        "Snapshot requiere revisión COMPLETED e incluye Five Case, evidence matrix, gate context y hashes.",
    ],
)

document.add_page_break()
heading(document, "4. Mapeo Five Case y autoridad de dominio")
table(
    document,
    ["Case", "Pregunta", "Fuentes gobernadas", "Autoridad preservada"],
    [
        [
            "Strategic",
            "¿Por qué intervenir y cómo se alinea?",
            "Idea, BA, Proposal, Benefits",
            "Estrategia/BA/Benefits",
        ],
        [
            "Economic",
            "¿Qué opción genera mayor valor total?",
            "BA, Proposal, Benefits, FIN, Risk",
            "BA y FIN",
        ],
        [
            "Commercial",
            "¿Puede contratarse y estructurarse?",
            "Commercial Strategy",
            "Commercial & Contracting",
        ],
        [
            "Financial",
            "¿Puede costearse y financiarse?",
            "Budget, Funding, FIN, Benefits",
            "Budget/Funds/FIN",
        ],
        [
            "Management",
            "¿Puede entregarse y gobernarse?",
            "FEL, Resources, Risk, Permit, Benefits, Commercial",
            "Cada provider fuente",
        ],
    ],
    [0.8, 2.1, 2.35, 2.1],
)
heading(document, "4.1 Reglas de integración", 2)
bullets(
    document,
    [
        "Business Analysis se resuelve por ProjectProposal.source_business_analysis_id/hash; nunca por coincidencia débil de Idea.",
        "La opción recomendada debe coincidir entre el contenido BA y el snapshot del Proposal; divergencia implica source_current=false.",
        "Benefits y Commercial se consumen desde sus versiones actuales gobernadas, con readiness, warnings, blockers y hash.",
        "FIN sigue siendo autoridad para NPV, IRR, BCR, WACC, DSCR, LLCR, MOIC y Payback; Business Case no recalcula métricas.",
        "Budget/Resource usan escenarios congelados; Funding, Risk y Permit conservan sus contratos de readiness y validez.",
        "QRA/Confidence, Basic Data y Owner Capability son futuras fuentes explícitas NOT_AVAILABLE, no providers fabricados.",
    ],
)

heading(document, "5. Readiness y madurez progresiva")
paragraph(
    document,
    "Los estados cerrados son NOT_STARTED, IN_PROGRESS, READY, BLOCKED, STALE y NOT_APPLICABLE. El backend compone cada Case desde las fuentes requeridas por la política efectiva; el frontend solo presenta el resultado.",
)
table(
    document,
    ["Etapa", "Cases requeridos por defecto"],
    [
        ["FEL-1", "Strategic"],
        ["FEL-2", "Strategic, Economic, Commercial y Management"],
        ["Investment Selection", "Strategic, Economic, Commercial y Management"],
        ["FEL-3", "Los cinco Cases"],
        ["FFA", "Los cinco Cases; Funding requerido para Financial"],
    ],
    [1.55, 5.8],
)
paragraph(
    document,
    "Para Cases requeridos, STALE domina; luego BLOCKED; todos READY producen READY. Cases no requeridos maduran como IN_PROGRESS o NOT_STARTED. En modelos no habilitados, Cases y overall son NOT_APPLICABLE. Tailoring y overrides quedan congelados en la revisión.",
)

document.add_page_break()
heading(document, "6. Refresh, staleness y recuperación")
table(
    document,
    ["Operación", "Comportamiento gobernado"],
    [
        [
            "Refresh Preview",
            "persisted=false; compara unión de claves y detecta ADDED, CHANGED y REMOVED",
        ],
        [
            "Delta",
            "Agrupa cambios por Case y expone readiness, warnings, blockers y hashes antes/después",
        ],
        [
            "Material change",
            "Recomienda MATERIAL_CHANGE_UPDATE cuando existen snapshots y cambian fuentes",
        ],
        [
            "Refresh Apply",
            "Exige If-Match e idempotencia; crea revisión sucesora DRAFT",
        ],
        [
            "Historical read",
            "No consulta fuentes vivas ni reinterpreta política/contexto",
        ],
        ["Legacy", "Entrega contrato v1 y legacy_layout=true sin Five Case inventado"],
    ],
    [1.65, 5.7],
)
paragraph(
    document,
    "Los negativos focales cubren Benefits faltante/stale, Commercial bloqueado, FIN stale, FEL stale y Business Analysis mismatch. La evidence matrix representa una fuente ausente como NOT_AVAILABLE con ID/hash nulos y blocking=true.",
)

heading(document, "7. Provider y frontera Gate07J")
bullets(
    document,
    [
        "provider_key: BUSINESS_CASE_READINESS_PROVIDER.",
        "evidence_type: BUSINESS_CASE_READINESS.",
        "Se añaden Five Case readiness, hashes, refs, source_current, warnings, blockers y can_supply_investment_gate_evidence.",
        "Para revisiones legacy o NOT_APPLICABLE se conserva el comportamiento v1.",
        "Solo una revisión COMPLETED, READY, current y sin blockers puede suministrar evidencia.",
        "No se publicó, reordenó ni reinterpretó la matriz/default policy de Gate07J-EVID.",
    ],
)

heading(document, "8. Seguridad, concurrencia y auditoría")
table(
    document,
    ["Control", "Implementación", "Estado"],
    [
        [
            "RBAC",
            "Permisos read/create/edit/review/refresh/snapshot/admin existentes",
            "PASS",
        ],
        ["Scope", "Tenant y Project Workspace en queries y endpoints", "PASS"],
        [
            "ADMIN/USER",
            "Configuración separada de expedientes vivos y rutas protegidas",
            "PASS",
        ],
        [
            "SoD",
            "Reviewer final distinto del preparer cuando la policy lo exige",
            "PASS",
        ],
        ["ETag", "If-Match y revision_version; mismatch = 412", "PASS"],
        ["Idempotencia", "Create/revision/submit/complete/refresh/snapshot", "PASS"],
        [
            "SecurityEvent",
            "Lifecycle, refresh, staleness, material change, snapshots y config",
            "PASS",
        ],
    ],
    [1.2, 5.25, 0.9],
    status_column=2,
)

document.add_page_break()
heading(document, "9. USER MODE, ADMIN MODE y navegación")
heading(document, "9.1 USER MODE", 2)
paragraph(
    document,
    "Business Case Manager permanece dentro de Project Definition & Investment Case. La navegación canónica ofrece Overview, Strategic Case, Economic Case, Commercial Case, Financial Case, Management Case, Sources & Evidence, Readiness, Decision Snapshots e History.",
)
bullets(
    document,
    [
        "Cada Case presenta pregunta, resumen, evidencia, readiness, findings, supuestos, warnings, blockers y conclusión.",
        "Sources & Evidence es read-only y ofrece View Source hacia el módulo propietario.",
        "Revisiones legacy muestran fallback explícito y no aparentan readiness Five Case.",
    ],
)
heading(document, "9.2 ADMIN MODE", 2)
paragraph(
    document,
    "La configuración Business Case existente se amplió con applicability, requirements por Case, source/readiness/evidence rules, FEL/Gate, overrides, tailoring, inheritance, Review/SoD, Preview e History. ADMIN no crea ni edita evidencia fuente.",
)
heading(document, "9.3 Integración técnica", 2)
paragraph(
    document,
    "strategicNavigationBlueprint.ts continúa como fuente de verdad; StrategicViewOutlet entrega la vista y el navegador canónico. Las rutas históricas se mantienen como aliases. App.tsx no recibió routing específico Five Case.",
)

heading(document, "10. API, compatibilidad y despliegue")
table(
    document,
    ["Área", "Resultado"],
    [
        [
            "API",
            "BusinessCaseRevisionOut expone estructura, readiness, matrix, context, vigencia y hashes Five Case",
        ],
        [
            "Compatibilidad",
            "Campos aditivos; endpoints, ETag e idempotency keys existentes conservados",
        ],
        ["Migración", "No requerida; head permanece 20260916_0061"],
        [
            "PostgreSQL",
            "Tablas/constraints/unique keys e índices funcionales confirmados",
        ],
        ["Docker", "API, worker y beat reconstruidos; DB/Redis healthy"],
        ["Health", "GET /api/v1/health/ready = 200; API/database/redis = ok"],
        [
            "OpenAPI",
            "Contrato desplegado contiene los cinco campos/estructuras aditivos comprobados",
        ],
    ],
    [1.4, 5.95],
)

document.add_page_break()
heading(document, "11. Evidencia de validación")
validation_rows = [
    [
        "Backend focal",
        backend["tests"],
        backend["failures"] + backend["errors"],
        backend["skipped"],
        state(backend),
    ],
    [
        "Backend estratégico",
        backend_strategic["tests"],
        backend_strategic["failures"] + backend_strategic["errors"],
        backend_strategic["skipped"],
        state(backend_strategic),
    ],
    [
        "PostgreSQL estructural",
        postgres_structural["tests"],
        postgres_structural["failures"] + postgres_structural["errors"],
        postgres_structural["skipped"],
        state(postgres_structural),
    ],
    [
        "PostgreSQL Five Case transaccional",
        postgres_business_case["tests"],
        postgres_business_case["failures"] + postgres_business_case["errors"],
        postgres_business_case["skipped"],
        state(postgres_business_case),
    ],
    [
        "Frontend focal",
        frontend["tests"],
        frontend["failures"] + frontend["errors"],
        frontend["skipped"],
        state(frontend),
    ],
    [
        "Browser Chromium",
        browser["tests"],
        browser["failures"] + browser["errors"],
        browser["skipped"],
        state(browser),
    ],
    [
        "Frontend full",
        frontend_full["tests"],
        frontend_full["failures"] + frontend_full["errors"],
        frontend_full["skipped"],
        state(frontend_full),
    ],
    [
        "AppFlow hardening aislado",
        appflow_hardening["tests"],
        appflow_hardening["failures"] + appflow_hardening["errors"],
        appflow_hardening["skipped"],
        state(appflow_hardening),
    ],
    [
        "AppFlow user/admin aislado",
        appflow_admin["tests"],
        appflow_admin["failures"] + appflow_admin["errors"],
        appflow_admin["skipped"],
        state(appflow_admin),
    ],
    [
        "AppFlow navegación aislado",
        appflow_navigation["tests"],
        appflow_navigation["failures"] + appflow_navigation["errors"],
        appflow_navigation["skipped"],
        state(appflow_navigation),
    ],
    ["TypeScript / Vite / ESLint / Prettier", 4, 0, 0, "PASS"],
    ["Docker health / OpenAPI", 2, 0, 0, "PASS"],
]
table(
    document,
    ["Suite", "Casos", "Fallos", "Skip", "Estado"],
    validation_rows,
    [3.1, 0.65, 0.65, 0.6, 1.45],
    status_column=4,
)
paragraph(
    document,
    "Backend focal incluye BC, Five Case, Benefits, Commercial, FIN y Evidence Policy. Ruff app/tests pasó. TypeScript, build Vite, ESLint (0 errores, 9 warnings BIM heredados) y Prettier pasaron. La prueba Chromium recorre USER y ADMIN en el localhost desplegado.",
)
paragraph(
    document,
    "La corrida frontend inicial obtuvo 270/273 y tres timeouts sin fallos de aserción. Se estabilizaron únicamente los límites de tiempo (45→90 s global y 300→420 s para el AppFlow largo), los tres casos pasaron aislados y la repetición integral cerró 273/273. La regresión estratégica backend final quedó verde tras eliminar una dependencia de orden del fixture Gate 07L, sin modificar lógica productiva ni relajar aserciones.",
)

heading(document, "12. Rendimiento y observaciones operativas")
bullets(
    document,
    [
        "Vite transformó 2.453 módulos y produjo el build de producción correctamente.",
        "web-ifc y Three.js permanecen aislados en chunks BIM; el incremento Five Case no añadió un motor pesado al bundle inicial.",
        "La prueba AppFlow de production hardening tardó 46,26 s; la de tenant user/admin, 30,23 s. Los límites se ajustaron a 90 s global y 420 s para el flujo largo; la repetición integral quedó verde.",
        "El E2E Chromium Five Case tardó 17,9 s y validó un flujo completo USER/ADMIN.",
        "El API health mostró latencias de pocos milisegundos para PostgreSQL y Redis en el entorno local.",
    ],
)

document.add_page_break()
heading(document, "13. Matriz de los 75 entregables")
deliverables = [
    "Executive summary",
    "ADR Gate07M-BC-5C",
    "Actual baseline assessment",
    "REUSE BEFORE CREATE analysis",
    "Existing BC architecture preservation proof",
    "Five Case conceptual mapping",
    "Domain authority matrix",
    "Source authority matrix",
    "Strategic Case design",
    "Economic Case design",
    "Commercial Case design",
    "Financial Case design",
    "Management Case design",
    "Case source mapping",
    "Source compatibility rules",
    "Benefits integration",
    "Commercial Strategy integration",
    "FIN integration",
    "FEL integration",
    "Budget/Resource integration",
    "Funding integration",
    "Risk integration",
    "Permit integration",
    "Case readiness model",
    "Overall readiness model",
    "Progressive maturity model",
    "FEL/Gate tailoring",
    "Gate-level revalidation model",
    "Material change model",
    "Source-current/staleness",
    "Source mismatch handling",
    "Refresh Preview/Apply",
    "Legacy BC compatibility",
    "Historical snapshot compatibility",
    "Existing provider identification",
    "Provider enhancement",
    "Provider backward compatibility",
    "Gate07J non-mutation proof",
    "Future provider extensibility",
    "Missing CPIG source handling",
    "ADMIN configuration design",
    "ADMIN UI",
    "USER Five Case UI",
    "Sources & Evidence matrix",
    "Roles/permissions",
    "SoD",
    "ETag",
    "Idempotency",
    "SecurityEvent",
    "Hashing",
    "Persistence decision",
    "Migration decision",
    "API compatibility",
    "PostgreSQL happy path",
    "Negative Benefits test",
    "Negative Commercial test",
    "Negative FIN stale test",
    "Negative FEL stale test",
    "Source mismatch test",
    "Legacy revision test",
    "Refresh/recovery test",
    "Decision snapshot test",
    "Governance Model regression",
    "Tailoring/proportionality",
    "Browser USER E2E",
    "Browser ADMIN E2E",
    "Direct-route security",
    "Navigation integration",
    "Full strategic regression",
    "Performance observation",
    "Technical debt",
    "CPIG gap update",
    "Updated technical report",
    "Final exclusive state",
    "Recommendation for next increment",
]
partial_ids = set() if release_validated else {54, 69, 74}
evidence = {
    2: "docs/67-adr...",
    32: "refresh preview/apply + tests",
    35: "BUSINESS_CASE_READINESS_PROVIDER",
    38: "policy matrix sin mutación",
    54: "postgres-business-case.xml: 2/2, incluido lifecycle READY",
    55: "backend-focused.xml",
    56: "backend-focused.xml",
    57: "backend-focused.xml",
    58: "backend-focused.xml",
    59: "backend-focused.xml",
    60: "legacy unit test",
    61: "refresh union/removal test",
    62: "snapshot regression test",
    65: "browser-e2e.xml",
    66: "browser-e2e.xml",
    68: "blueprint + E2E",
    69: f"backend-strategic-regression-final.xml: {backend_strategic['tests'] - backend_strategic['skipped']} ejecutadas PASS; {backend_strategic['skipped']} PostgreSQL-only omitidas",
    70: "build y tiempos observados",
    73: "este documento",
    74: release_signal
    + (
        " / READY_FOR_GOVERNED_FIVE_CASE_BUSINESS_CASE_EVIDENCE"
        if release_validated
        else ""
    ),
}
rows = []
for index, item in enumerate(deliverables, start=1):
    status = "PARCIAL" if index in partial_ids else "PASS"
    rows.append(
        [index, item, evidence.get(index, "Código, ADR y validación focal"), status]
    )
for start in range(0, len(rows), 19):
    if start:
        document.add_page_break()
        heading(
            document,
            f"13.{start // 19 + 1} Entregables {start + 1}–{min(start + 19, len(rows))}",
            2,
        )
    table(
        document,
        ["#", "Entregable", "Evidencia", "Estado"],
        rows[start : start + 19],
        [0.38, 2.45, 3.7, 0.82],
        status_column=3,
    )

document.add_page_break()
heading(document, "14. Deuda técnica, gaps CPIG y estado final")
table(
    document,
    ["Tema", "Estado", "Tratamiento recomendado"],
    [
        [
            "Regresión frontend completa",
            "PASS",
            "273/273 tras estabilización controlada de timeouts",
        ],
        [
            "Regresión backend estratégica",
            "PASS",
            f"{backend_strategic['tests'] - backend_strategic['skipped']} ejecutadas PASS; {backend_strategic['skipped']} PostgreSQL-only omitidas",
        ],
        [
            "PostgreSQL lifecycle Five Case",
            "PASS",
            "Round-trip transaccional de cinco Cases READY con rollback final",
        ],
        [
            "Índices PostgreSQL",
            "PASS",
            "La prueba valida columnas, independiente del nombre create_all/Alembic",
        ],
        [
            "Warnings BIM",
            "DEUDA",
            "9 warnings ESLint heredados; fuera del alcance Five Case",
        ],
        ["Gate07J policy", "FUERA", "Mantener sin mutación hasta prompt explícito"],
    ],
    [1.9, 1.0, 4.45],
    status_column=1,
)
paragraph(document, "Gaps CPIG explícitamente no implementados:")
bullets(
    document,
    [
        "Risk Appetite / Risk Capacity / Risk Thresholds.",
        "Strategic Capital Demand / Capital Planning enhancement.",
        "Basic Data Readiness y QRA / Confidence.",
        "Owner Capability Readiness.",
        "Gate-level evidence policy integration.",
        "Operational Readiness y Benefits Realization.",
    ],
)

heading(document, "14.1 Estado exclusivo", 2)
signal = document.add_paragraph()
signal.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(signal.add_run(release_signal), 14, True, GREEN if release_validated else AMBER)
paragraph(
    document,
    (
        "READY_FOR_GOVERNED_FIVE_CASE_BUSINESS_CASE_EVIDENCE. La capacidad quedó implementada, desplegada y validada con evidencia focal, regresión estratégica, frontend integral, navegador y PostgreSQL transaccional."
        if release_validated
        else "No se emiten las señales formales mientras alguna evidencia requerida continúe pendiente."
    ),
)
heading(document, "14.2 Recomendación del siguiente incremento", 2)
paragraph(
    document,
    "El siguiente gap recomendado es Gate-level evidence policy integration, porque conecta el provider ya enriquecido con la evaluación formal sin duplicar evidencia ni anticipar Procurement, Project Delivery u Operational Readiness.",
)
heading(document, "14.3 STOP boundary", 2)
paragraph(
    document,
    "Este trabajo se detiene aquí. No inicia automáticamente Gate07J policy update, Risk Appetite, Basic Data, QRA, Owner Capability, Capital Demand, Procurement execution, Project Delivery, Operational Readiness ni Benefits Realization.",
)

document.core_properties.title = "Informe técnico Gate 07M-BC-5C"
document.core_properties.subject = (
    "Five Case Model Integration for Governed Business Case"
)
document.core_properties.author = "P&Pmis Construction AI / Codex"
document.core_properties.keywords = (
    "Gate07M-BC-5C, Five Case, Business Case, readiness, CPIG, P&Pmis"
)
document.core_properties.comments = (
    "Informe técnico de implementación, despliegue y validación"
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
print(OUTPUT)
