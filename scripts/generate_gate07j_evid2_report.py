from __future__ import annotations

import json
import xml.etree.ElementTree as ET
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
QA = ROOT / "outputs" / ".qa" / "gate07j-evid2"
OUTPUT = (
    ROOT
    / "outputs"
    / "Informe_Tecnico_Gate07J_EVID_2_Criterios_Decision_2026-09-20.docx"
)
LOGO = ROOT / "frontend" / "public" / "pypmis-construction-ai-logo.png"
ADR = ROOT / "docs" / "68-adr-gate07j-evid2-decision-criteria.md"
VALIDATION_RECORD = (
    ROOT / "docs" / "validation" / "gate07j-evid2" / "release-validation.md"
)

BLACK = "000000"
TEXT = "202A33"
MUTED = "53616B"
NAVY = "17324D"
WHITE = "FFFFFF"
PALE_GREEN = "E9F5EE"
PALE_BLUE = "EEF4F8"
PALE_AMBER = "FFF4E5"
PALE_RED = "FDECEC"
PALE_GRAY = "F5F7F8"
LINE = "D9D9D9"
GREEN = "18794E"
AMBER = "8A5700"
RED = "B42318"


@dataclass(frozen=True)
class Check:
    key: str
    label: str
    status: str
    detail: str
    source: str
    tests: int = 0
    failures: int = 0
    errors: int = 0
    skipped: int = 0
    seconds: float = 0.0

    @property
    def passed(self) -> bool:
        return self.status == "PASS"


def first_existing(*names: str) -> Path | None:
    for name in names:
        candidate = QA / name
        if candidate.exists() and candidate.stat().st_size:
            return candidate
    return None


def read_json(path: Path | None) -> dict[str, Any]:
    if path is None or not path.exists() or not path.stat().st_size:
        return {}
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError):
        return {}
    return value if isinstance(value, dict) else {}


def junit_check(key: str, label: str, *names: str) -> Check:
    path = first_existing(*names)
    if path is None:
        return Check(key, label, "PENDIENTE", "Artefacto JUnit no encontrado", "-")

    values = {"tests": 0, "failures": 0, "errors": 0, "skipped": 0}
    seconds = 0.0
    try:
        root = ET.parse(path).getroot()
        suites = [root] if root.tag == "testsuite" else list(root.findall("testsuite"))
        for suite in suites:
            for item in values:
                values[item] += int(suite.attrib.get(item, 0))
            seconds += float(suite.attrib.get("time", 0) or 0)
    except (ET.ParseError, OSError, TypeError, ValueError) as error:
        return Check(
            key,
            label,
            "ATENCION",
            f"JUnit ilegible: {type(error).__name__}",
            path.name,
        )

    executed = values["tests"] - values["skipped"]
    failures = values["failures"] + values["errors"]
    if values["tests"] == 0 or executed == 0:
        status = "PENDIENTE"
    else:
        status = "PASS" if failures == 0 else "FAIL"
    detail = (
        f"{executed}/{values['tests']} ejecutadas; "
        f"{failures} fallos; {values['skipped']} omitidas; {seconds:.2f} s"
    )
    return Check(
        key,
        label,
        status,
        detail,
        path.name,
        values["tests"],
        values["failures"],
        values["errors"],
        values["skipped"],
        seconds,
    )


def normalized_status(value: Any) -> tuple[str, str]:
    if isinstance(value, bool):
        return ("PASS", "Validado") if value else ("FAIL", "Reportado como fallido")
    if isinstance(value, str):
        candidate = value.strip()
        upper = candidate.upper()
        if upper in {"PASS", "PASSED", "VALIDATED", "OK", "GREEN"}:
            return "PASS", candidate
        if upper in {"FAIL", "FAILED", "ERROR", "RED"}:
            return "FAIL", candidate
        if upper in {"PENDING", "PENDIENTE", "SKIPPED", "NOT_RUN"}:
            return "PENDIENTE", candidate
        return "DOCUMENTADO", candidate
    if isinstance(value, dict):
        raw = value.get("status", value.get("passed", value.get("result")))
        status, default_detail = normalized_status(raw)
        detail = str(value.get("detail", value.get("summary", default_detail)))
        return status, detail
    return "PENDIENTE", "Sin resultado estructurado"


def summary_check(
    summary: dict[str, Any], key: str, label: str, *fallback_names: str
) -> Check:
    checks = summary.get("checks", {})
    if isinstance(checks, dict) and key in checks:
        status, detail = normalized_status(checks[key])
        return Check(key, label, status, detail, "validation-summary.json")
    path = first_existing(*fallback_names)
    if path is None:
        return Check(key, label, "PENDIENTE", "Resultado no registrado", "-")
    text = path.read_text(encoding="utf-8", errors="replace")
    upper = text.upper()
    status = "PASS" if any(token in upper for token in ("RESULT=PASS", "STATUS=PASS", "EXIT=0")) else "DOCUMENTADO"
    return Check(key, label, status, text.strip()[:180] or "Log disponible", path.name)


def collect_checks() -> tuple[dict[str, Check], dict[str, Any]]:
    summary = read_json(first_existing("validation-summary.json", "summary.json"))
    checks = {
        item.key: item
        for item in (
            junit_check(
                "backend_focused",
                "Backend focal Gate07J EVID 2",
                "backend-focused.xml",
                "backend-evid2.xml",
            ),
            junit_check(
                "backend_strategic",
                "Regresion estrategica backend",
                "backend-strategic-regression.xml",
                "backend-strategic-regression-final.xml",
                "backend-strategic.xml",
            ),
            junit_check(
                "postgres_ffa",
                "PostgreSQL PYP FFA",
                "postgres-ffa.xml",
                "postgres-pyp-ffa.xml",
                "postgres-evid2-ffa.xml",
            ),
            junit_check(
                "postgres_select",
                "PostgreSQL PYP INV SELECT",
                "postgres-select.xml",
                "postgres-pyp-inv-select.xml",
                "postgres-evid2-select.xml",
            ),
            junit_check(
                "frontend_focused",
                "Frontend focal",
                "frontend-focused.xml",
                "frontend-evid2.xml",
            ),
            junit_check(
                "frontend_full",
                "Regresion frontend completa",
                "frontend-full.xml",
                "frontend-full-final.xml",
            ),
            junit_check(
                "browser_gate07j",
                "Browser USER Gate07J",
                "browser-gate07j.xml",
                "browser-user-gate07j.xml",
            ),
            junit_check(
                "browser_admin",
                "Browser ADMIN policy",
                "browser-admin.xml",
                "browser-admin-policy.xml",
            ),
            junit_check(
                "browser_gate07k",
                "Browser USER Gate07K",
                "browser-gate07k.xml",
                "browser-investment-decision.xml",
            ),
        )
    }
    for key, label, names in (
        ("ruff", "Ruff backend", ("ruff.txt", "backend-ruff.txt")),
        ("typecheck", "TypeScript", ("typescript.txt", "typecheck.txt")),
        ("build", "Vite production build", ("vite-build.txt", "build.txt")),
        ("eslint", "ESLint", ("eslint.txt",)),
        ("prettier", "Prettier", ("prettier.txt", "format.txt")),
        ("alembic", "Alembic head", ("alembic.txt", "alembic-head.txt")),
        ("docker_health", "Docker y health", ("docker-health.txt", "health.txt")),
        ("openapi", "OpenAPI desplegado", ("openapi.txt", "openapi-contract.txt")),
        ("performance", "Rendimiento", ("performance.txt", "performance.json")),
    ):
        checks[key] = summary_check(summary, key, label, *names)
    return checks, summary


def shade(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    node = properties.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        properties.append(node)
    node.set(qn("w:fill"), fill)


def cell_margins(cell: Any, value: int = 90) -> None:
    properties = cell._tc.get_or_add_tcPr()
    node = properties.first_child_found_in("w:tcMar")
    if node is None:
        node = OxmlElement("w:tcMar")
        properties.append(node)
    for side in ("top", "start", "bottom", "end"):
        item = node.find(qn(f"w:{side}"))
        if item is None:
            item = OxmlElement(f"w:{side}")
            node.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def table_borders(table_item: Any) -> None:
    properties = table_item._tbl.tblPr
    node = properties.first_child_found_in("w:tblBorders")
    if node is None:
        node = OxmlElement("w:tblBorders")
        properties.append(node)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = node.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            node.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), LINE)


def repeat_header(row: Any) -> None:
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def no_split(row: Any) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def set_font(
    run: Any,
    size: float = 10.5,
    bold: bool = False,
    color: str = TEXT,
) -> None:
    run.font.name = "Aptos"
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), "Aptos")
    fonts.set(qn("w:hAnsi"), "Aptos")
    fonts.set(qn("w:eastAsia"), "Aptos")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_paragraph(
    document: Document,
    text: str = "",
    *,
    style: str | None = None,
    keep_with_next: bool = False,
) -> Any:
    item = document.add_paragraph(style=style)
    item.paragraph_format.space_after = Pt(5)
    item.paragraph_format.line_spacing = 1.08
    item.paragraph_format.keep_with_next = keep_with_next
    set_font(item.add_run(text))
    return item


def add_bullets(document: Document, items: Iterable[str]) -> None:
    for value in items:
        item = add_paragraph(document, value, style="List Bullet")
        item.paragraph_format.space_after = Pt(2.5)


def add_heading(document: Document, text: str, level: int = 1) -> Any:
    item = document.add_heading(text, level=level)
    item.paragraph_format.keep_with_next = True
    item.paragraph_format.space_before = Pt(8 if level == 1 else 6)
    item.paragraph_format.space_after = Pt(4)
    return item


def add_table(
    document: Document,
    headers: list[str],
    rows: list[list[Any]],
    widths: list[float],
    *,
    status_column: int | None = None,
    font_size: float = 7.8,
) -> Any:
    table_item = document.add_table(rows=1, cols=len(headers))
    table_item.autofit = False
    table_borders(table_item)
    repeat_header(table_item.rows[0])
    no_split(table_item.rows[0])
    for index, value in enumerate(headers):
        cell = table_item.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, NAVY)
        cell_margins(cell)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        cell.paragraphs[0].paragraph_format.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER if index == status_column else WD_ALIGN_PARAGRAPH.LEFT
        )
        set_font(cell.paragraphs[0].add_run(str(value)), 8.1, True, WHITE)
    for row_index, values in enumerate(rows):
        row = table_item.add_row()
        no_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(cell, WHITE if row_index % 2 == 0 else PALE_GRAY)
            cell_margins(cell)
            paragraph_item = cell.paragraphs[0]
            paragraph_item.paragraph_format.space_after = Pt(0)
            paragraph_item.paragraph_format.line_spacing = 1.04
            if index == status_column:
                paragraph_item.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph_item.add_run(str(value))
            set_font(run, font_size)
            if status_column == index:
                status = str(value).upper()
                if status.startswith("PASS"):
                    shade(cell, PALE_GREEN)
                    run.font.color.rgb = RGBColor.from_string(GREEN)
                    run.bold = True
                elif status.startswith(("PENDIENTE", "DOCUMENTADO", "IMPLEMENTADO")):
                    shade(cell, PALE_AMBER if status.startswith("PENDIENTE") else PALE_BLUE)
                    run.font.color.rgb = RGBColor.from_string(AMBER if status.startswith("PENDIENTE") else NAVY)
                    run.bold = True
                elif status.startswith(("FAIL", "ATENCION")):
                    shade(cell, PALE_RED)
                    run.font.color.rgb = RGBColor.from_string(RED)
                    run.bold = True
                elif "STOP" in status:
                    shade(cell, PALE_GRAY)
                    run.font.color.rgb = RGBColor.from_string(MUTED)
                    run.bold = True
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return table_item


def add_page_number(paragraph_item: Any) -> None:
    run = paragraph_item.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, value, end])


def repository_presence(*paths: str) -> bool:
    return all((ROOT / path).exists() for path in paths)


def evidence_status(present: bool, validated: bool = False) -> str:
    if validated:
        return "PASS"
    return "IMPLEMENTADO" if present else "PENDIENTE"


def make_document(checks: dict[str, Check], summary: dict[str, Any]) -> Document:
    required_release_checks = (
        "backend_focused",
        "backend_strategic",
        "postgres_ffa",
        "postgres_select",
        "frontend_focused",
        "frontend_full",
        "browser_gate07j",
        "browser_admin",
        "browser_gate07k",
        "ruff",
        "typecheck",
        "build",
        "eslint",
        "prettier",
        "alembic",
        "docker_health",
        "openapi",
    )
    computed_release = all(checks[key].passed for key in required_release_checks)
    if "release_validated" in summary:
        release_validated = bool(summary["release_validated"]) and computed_release
    else:
        release_validated = computed_release
    release_signal = (
        "GATE07J_EVID_2_DECISION_CRITERIA_POLICY_RELEASE_VALIDATED"
        if release_validated
        else "PENDING_RELEASE_VALIDATION"
    )
    final_output = (
        "READY_FOR_GOVERNED_GATE_DECISION_WITH_INTEGRATED_STRATEGIC_EVIDENCE"
        if release_validated
        else "NOT_EMITTED"
    )

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    for name, size in (
        ("Title", 25),
        ("Heading 1", 16),
        ("Heading 2", 12.5),
        ("Heading 3", 10.5),
    ):
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = RGBColor.from_string(BLACK)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("P&Pmis Construction AI  Gate 07J EVID 2"), 8.2, True, BLACK)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("Informe tecnico controlado  Pagina "), 8.2, False, BLACK)
    add_page_number(footer)

    if LOGO.exists():
        image_paragraph = document.add_paragraph()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        drawing = image_paragraph.add_run().add_picture(str(LOGO), width=Inches(0.95))
        drawing._inline.docPr.set("descr", "Logotipo de P&Pmis Construction AI")

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    set_font(title.add_run("Informe tecnico Gate 07J EVID 2"), 25, True, BLACK)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    set_font(
        subtitle.add_run(
            "Politica de evidencia criterios de decision y revalidacion estrategica"
        ),
        14,
        True,
        BLACK,
    )

    alembic_detail = checks["alembic"].detail
    add_table(
        document,
        ["Control", "Valor"],
        [
            ["Aplicacion", "P&Pmis Construction AI"],
            ["Incremento", "PROMPT Gate 07J EVID 2"],
            ["Fecha de cierre", "20 de septiembre de 2026"],
            ["Alembic", alembic_detail],
            ["Persistencia", "JSON gobernado existente sin migracion nueva"],
            ["Release", release_signal],
            ["Salida final", final_output],
        ],
        [1.55, 5.55],
    )

    add_heading(document, "1 Resumen ejecutivo")
    add_paragraph(
        document,
        "Se integro una politica gobernada de criterios de decision sobre la politica de evidencia existente. "
        "Gate07J prepara, revalida y congela el paquete; Gate07K conserva la autoridad exclusiva para GO, "
        "HOLD, RECYCLE y KILL. La implementacion mantiene separadas la clasificacion de evidencia A B C, "
        "la evaluacion de negocio, el contexto de Portfolio y la decision formal.",
    )
    add_bullets(
        document,
        [
            "La taxonomia incluye READINESS_CHECK, MUST_MEET, DECISION_HURDLE, PRIORITIZATION y SUCCESS_CRITERION.",
            "Los criterios usan operadores controlados y consumen valores de providers gobernados; no ejecutan formulas libres ni recalculan FIN, Five Case o ranking de Portfolio.",
            "Los paquetes congelan politica, evidencias, resultados de criterios, contexto de Portfolio y candidatos de exito con hashes y versionado historico.",
            "Gate07K impide GO cuando el paquete no esta listo, falta evidencia Class A o no se satisface un MUST_MEET o DECISION_HURDLE requerido.",
            "El snapshot de criterios de exito conserva referencias BEN y valores de decision sin duplicar el Benefits Register.",
        ],
    )
    add_paragraph(
        document,
        (
            "La evidencia automatizada obligatoria esta completa. Se emiten las dos señales formales de cierre."
            if release_validated
            else "La implementacion se documenta con su evidencia disponible. Las señales formales permanecen retenidas hasta completar todos los controles obligatorios."
        ),
    )

    document.add_page_break()
    add_heading(document, "2 Baseline y reutilizacion")
    add_table(
        document,
        ["Activo existente", "Decision de reutilizacion", "Resultado"],
        [
            ["AdminConfiguration", "Versiona evidencia y criterios en content_json", "Sin tabla nueva"],
            ["Gate07J Preparation", "Amplia source_snapshot_json y evidence_summary_json", "Paquete v2"],
            ["Gate07K Decision", "Amplia decision_package_snapshot_json", "Decision humana"],
            ["Provider registry", "Reutiliza keys y contratos vigentes", "Sin providers duplicados"],
            ["Benefits", "Referencia BEN estable y snapshot de valores", "Sin Benefit KPI duplicado"],
            ["Portfolio", "Consume rank, recomendacion y restricciones", "Sin segundo ranking"],
            ["SecurityEvent", "Reutiliza auditoria existente", "Eventos aditivos"],
        ],
        [1.55, 3.4, 2.15],
    )
    add_paragraph(
        document,
        "El baseline esperado conserva Alembic 20260916_0061 como head unico. La persistencia JSON existente soporta "
        "version de politica, evaluaciones, lineage y snapshots inmutables; por ello el incremento no requiere migracion.",
    )
    add_heading(document, "2 1 Frontera de autoridad", 2)
    add_table(
        document,
        ["Capa", "Responsabilidad", "No hace"],
        [
            ["Evidence policy", "Define fuente, cobertura, vigencia y Class A B C", "No decide el Gate"],
            ["Gate07J", "Revalida, evalua criterios y congela el paquete", "No crea outcome formal"],
            ["Portfolio", "Aporta mix, rank, escenario y restricciones", "No es recalculado por Gate"],
            ["Gate07K", "Registra GO HOLD RECYCLE o KILL", "No edita evidencia fuente"],
            ["Gate07L", "Consume decision historica para post FID", "No se modifica en este incremento"],
        ],
        [1.35, 3.55, 2.2],
    )

    add_heading(document, "3 Taxonomias y modelo de evaluacion")
    add_heading(document, "3 1 Clasificacion de evidencia", 2)
    add_table(
        document,
        ["Clase", "Significado", "Control"],
        [
            ["A", "GOVERNED_SOURCE_REQUIRED", "Provider gobernado obligatorio; sin fallback manual"],
            ["B", "GOVERNED_EVIDENCE_SUFFICIENT", "Fuente controlada o evidencia validada segun politica"],
            ["C", "NOT_REQUIRED", "Sin requisito efectivo ni blocker"],
        ],
        [0.7, 2.35, 4.05],
    )
    add_heading(document, "3 2 Tipos de criterio", 2)
    add_table(
        document,
        ["Tipo", "Pregunta", "Efecto gobernado"],
        [
            ["READINESS_CHECK", "Esta listo el paquete", "Controla elegibilidad para evaluar"],
            ["MUST_MEET", "Se cumple una condicion obligatoria", "Puede bloquear el paquete"],
            ["DECISION_HURDLE", "Se alcanza el umbral configurado", "Consume metrica fuente"],
            ["PRIORITIZATION", "Cual es el contexto relativo", "Referencia Portfolio sin re rank"],
            ["SUCCESS_CRITERION", "Que compromiso se congela", "Candidato para snapshot FID"],
        ],
        [1.6, 2.35, 3.15],
    )
    add_heading(document, "3 3 Estados y operadores", 2)
    add_paragraph(
        document,
        "Cada evaluacion conserva criterion_id, code, type, source reference, source value, operator, expected value, "
        "result, blocking, evaluated_at y policy revision hash. Los resultados son SATISFIED, NOT_SATISFIED, "
        "NOT_APPLICABLE, NOT_EVALUABLE y STALE.",
    )
    add_table(
        document,
        ["Familia", "Operadores permitidos"],
        [
            ["Comparacion", "EQUALS, NOT_EQUALS, GREATER_THAN, GREATER_THAN_OR_EQUAL"],
            ["Comparacion adicional", "LESS_THAN, LESS_THAN_OR_EQUAL, IN, NOT_IN"],
            ["Booleano", "IS_TRUE, IS_FALSE"],
            ["Gobierno", "IS_READY, IS_CURRENT, COVERAGE_AT_LEAST"],
        ],
        [1.65, 5.45],
    )

    add_heading(document, "4 Fuentes gobernadas y autoridad de dominio")
    add_table(
        document,
        ["Dominio", "Provider reutilizado", "Datos consumidos", "Recalculo Gate"],
        [
            ["Business Case", "BUSINESS_CASE_READINESS_PROVIDER", "Revision, Five Case, hash, current, blockers", "No"],
            ["FIN", "ECONOMIC_FINANCIAL_READINESS_PROVIDER", "Readiness, metric summary, model hash", "No"],
            ["Commercial", "COMMERCIAL_READINESS_PROVIDER", "Strategy, packaging, market, risk allocation", "No"],
            ["Benefits", "BENEFITS_READINESS_PROVIDER", "Coverage, readiness, Benefit IDs, plan hash", "No"],
            ["FEL", "GATE07J_FEL_PROJECT_DEFINITION", "FEL stage, maturity, assessment, gaps", "No"],
            ["Funding", "GATE07J_FUND_FUNDING_READINESS", "Coverage, sources, compatibility, blockers", "No"],
            ["Risk", "RISK_READINESS_PROVIDER", "Readiness y condiciones gobernadas", "No QRA"],
            ["Permit", "PERMIT_READINESS_PROVIDER", "Readiness segun contexto regulado", "No"],
            ["Portfolio", "GATE07I PORTFOLIO y Gate07G Gate07H", "Rank, scenario, budget y resource context", "No rank"],
        ],
        [1.05, 2.05, 3.25, 0.75],
        font_size=7.4,
    )
    add_heading(document, "4 1 Evidencia calidad y confianza", 2)
    add_paragraph(
        document,
        "Readiness expresa suficiencia bajo la politica; quality y confidence describen la solidez de la fuente. "
        "El backend o el provider gobierna VALIDATED, PARTIALLY_VALIDATED, PRELIMINARY, LOW_CONFIDENCE, "
        "NOT_VALIDATED y NOT_APPLICABLE. Una estimacion PRELIMINARY puede ser READY en FEL 2 y no ser suficiente en FFA.",
    )
    add_heading(document, "4 2 Consistencia entre fuentes", 2)
    add_bullets(
        document,
        [
            "El contexto estricto FFA compara hashes y versiones de Business Case con FIN, Benefits, Commercial, FEL, Funding, Risk y Permit.",
            "Una referencia historica distinta de la fuente vigente produce mismatch o stale; Gate07J no corrige ni recalcula la fuente.",
            "Project Set usa denominador exacto. Nueve providers ready de diez Projects equivalen a 90 por ciento y bloquean una exigencia de 100 por ciento.",
        ],
    )

    add_heading(document, "5 Madurez progresiva y tailoring")
    add_table(
        document,
        ["Fuente", "PYP INV SELECT", "PYP FFA"],
        [
            ["Project Definition y FEL", "Class A minimo FEL 2", "Class A minimo FEL 3"],
            ["Business Case Five Case", "Class A con madurez de etapa", "Class A con Cases requeridos READY"],
            ["FIN", "Class B configurable", "Class A"],
            ["Commercial", "Class B configurable", "Class A"],
            ["Benefits", "Class B configurable", "Class A"],
            ["Funding", "Class B o NOT_REQUIRED", "Class A"],
            ["Risk", "Class B", "Class A"],
            ["Permit", "Class B", "Class B o Class A por contexto regulado"],
        ],
        [2.25, 2.35, 2.5],
    )
    add_paragraph(
        document,
        "La aplicabilidad combina Gate Profile, FEL Stage, Governance Model, Project Type, tamano, riesgo, "
        "complejidad e incertidumbre. CAPITAL_OWNER recibe la politica inicial; CONTRACTOR_DELIVERY y "
        "DIRECT_INTERNAL quedan fuera por defecto salvo opt in explicito. Los proyectos publicos o no revenue "
        "usan las metricas economicas que FIN declara aplicables, sin exigir IRR comercial cuando no corresponde.",
    )

    add_heading(document, "6 Paquete congelado y ciclo de politica")
    add_table(
        document,
        ["Grupo", "Contenido congelado"],
        [
            ["Politica", "ID, revision, effective date, hash y matriz de evidencia"],
            ["Fuentes", "Provider, type, source ID, version, hash, readiness, quality, current, coverage y as of"],
            ["Criterios", "Policy ID hash, resultados, actual, expected, result y blocking"],
            ["Portfolio", "Recommendation, rank authority, scenario y restricciones de capital o recursos"],
            ["Exito", "Candidatos con Benefit ID, metric, baseline, target, unit, date y source hash"],
            ["Control", "Warnings, blockers, package state, freeze hash y timestamps"],
        ],
        [1.35, 5.75],
    )
    add_bullets(
        document,
        [
            "Una preparacion congelada no cambia cuando se publica la revision N mas 1; conserva N y su hash.",
            "Un nuevo round posterior a HOLD, RECYCLE o revalidacion explicita resuelve la politica efectiva vigente y exige una nueva preparacion.",
            "Las politicas publicadas historicas sin decision_criterion_policy conservan su comportamiento; el sistema no fabrica criterios retrospectivos.",
            "Clone crea un DRAFT; Preview no persiste; Publish inmoviliza la revision y respeta ETag e idempotencia.",
        ],
    )

    add_heading(document, "7 Contrato Gate07K y resultados formales")
    add_table(
        document,
        ["Outcome", "Semantica", "Regla"],
        [
            ["GO", "Autoriza avance formal", "Package READY y MUST_MEET HURDLE requeridos satisfechos"],
            ["HOLD", "Aplaza por contexto o restriccion", "No marca fallida la evidencia fuente"],
            ["RECYCLE", "Solicita rework o nueva evidencia", "Nuevo round requiere nueva preparacion"],
            ["KILL", "Termina por decision autorizada", "No se infiere por deliverables faltantes"],
        ],
        [0.85, 2.5, 3.75],
    )
    add_paragraph(
        document,
        "Gate07K consume el paquete Gate07J congelado, evidencia, criterios obligatorios, decision hurdles, contexto "
        "de Portfolio, condiciones y candidatos de exito. La autoridad sigue siendo una persona o comite autorizado "
        "con quorum, separacion de funciones y limites vigentes. READY significa elegible para decision, no GO automatico.",
    )
    add_heading(document, "7 1 Snapshot de criterios de exito", 2)
    add_paragraph(
        document,
        "Cuando Gate07K registra GO, congela decision_id, Business Case revision y hash, Benefits Plan version y hash, "
        "identidades BEN, metricas, baseline, target, target date, unit, condition, source, source version, source hash y as_of. "
        "Una revision posterior del Benefits Plan no modifica el snapshot historico.",
    )

    add_heading(document, "8 Interfaces seguridad y persistencia")
    add_table(
        document,
        ["Superficie", "Capacidad", "Limite"],
        [
            ["ADMIN Evidence and Decision Policy", "Clone, edit DRAFT, criteria, tailoring, preview, publish, history", "No transacciones Gate"],
            ["USER Gate07J", "Matrix, groups, source links, statuses, context, candidates y freeze", "Sin controles ADMIN"],
            ["USER Gate07K", "Paquete frozen, criterios, contexto, conditions y decision", "Sin editar fuentes"],
        ],
        [1.75, 3.45, 1.9],
    )
    add_table(
        document,
        ["Control", "Implementacion"],
        [
            ["SoD", "Publisher, preparer y decision maker respetan roles y no bypass de Class A"],
            ["ETag", "If Match y respuesta 412 ETAG_MISMATCH para escrituras stale"],
            ["Idempotencia", "Clone, publish, refresh, freeze, create round, decide y success snapshot"],
            ["SecurityEvent", "Policy clone publish, criterion evaluation, preparation freeze y decision snapshot"],
            ["Persistencia", "AdminConfiguration y snapshots JSON existentes"],
            ["Migracion", "No requerida si Alembic confirma head 20260916_0061"],
            ["API", "Contratos aditivos para policy preview, Gate07J handoff y Gate07K package"],
        ],
        [1.45, 5.65],
    )

    add_heading(document, "9 Evidencia de validacion")
    validation_rows = [
        [item.label, item.detail, item.source, item.status]
        for item in checks.values()
    ]
    add_table(
        document,
        ["Validacion", "Resultado", "Artefacto", "Estado"],
        validation_rows,
        [2.0, 3.2, 1.15, 0.75],
        status_column=3,
        font_size=7.2,
    )
    add_heading(document, "9 1 Casos funcionales obligatorios", 2)
    add_table(
        document,
        ["Caso", "Resultado esperado", "Evidencia"],
        [
            ["PYP FFA happy path", "Evidence, MUST_MEET y HURDLE satisfechos; package READY", checks["postgres_ffa"].status],
            ["PYP INV SELECT", "Madurez FEL 2 sin overreach FFA", checks["postgres_select"].status],
            ["BC faltante", "Class A bloquea; no decision", checks["backend_focused"].status],
            ["FIN stale", "Criterio STALE o NOT_EVALUABLE", checks["backend_focused"].status],
            ["Benefits 80 de 100", "MUST_MEET NOT_SATISFIED", checks["backend_focused"].status],
            ["Commercial blocked", "Package no READY en FFA", checks["backend_focused"].status],
            ["Funding mismatch", "Consume blocker de FIN sin recalcular", checks["backend_focused"].status],
            ["Permit contextual", "Class A regulado bloquea; Class B sigue policy", checks["backend_focused"].status],
            ["Resource constrained", "Package puede estar READY y Gate07K puede HOLD", checks["backend_focused"].status],
            ["Readiness failure", "No crea KILL automatico", checks["backend_focused"].status],
            ["Policy N y N mas 1", "Historico queda en N", checks["backend_focused"].status],
            ["New round", "Resuelve policy vigente con nueva preparation", checks["backend_focused"].status],
            ["Success snapshot", "BEN y targets inmutables", checks["backend_focused"].status],
            ["Contractor Direct", "Sin owner FID forzado", checks["backend_strategic"].status],
            ["Project ID14", "Sin FID o hash fabricado", checks["backend_strategic"].status],
        ],
        [1.8, 4.35, 0.95],
        status_column=2,
        font_size=7.4,
    )

    add_heading(document, "10 Rendimiento regresion y deuda tecnica")
    add_paragraph(
        document,
        "El evaluador opera sobre la matriz ya resuelta y evita consultas por criterio. La API combina evidencia, "
        "evaluaciones, contexto y candidatos en el paquete; la vigencia de fuentes no se sacrifica por cache. "
        f"La evidencia de rendimiento reporta: {checks['performance'].detail}.",
    )
    add_table(
        document,
        ["Area", "Cobertura esperada", "Estado"],
        [
            ["Backend estrategico", "Enterprise, Idea, BA, Proposal, Portfolio, FEL, BC, BEN, COMM, FUND, FIN, RISK, PERMIT, 07J, 07K y 07L", checks["backend_strategic"].status],
            ["Frontend", "USER ADMIN, navegación, vistas historicas y modos", checks["frontend_full"].status],
            ["Browser", "Gate07J USER, policy ADMIN y Gate07K USER", "PASS" if all(checks[key].passed for key in ("browser_gate07j", "browser_admin", "browser_gate07k")) else "PENDIENTE"],
            ["Deuda tecnica", "Warnings heredados o mediciones pendientes se conservan en validation summary", "DOCUMENTADO"],
        ],
        [1.5, 4.75, 0.85],
        status_column=2,
    )

    add_heading(document, "11 Frontera de cierre")
    add_table(
        document,
        ["Capacidad", "Tratamiento"],
        [
            ["Strategic Capital Buckets", "N A STOP; referencia futura opcional"],
            ["Periodic Portfolio Review", "N A STOP; contexto futuro opcional"],
            ["Capital Demand y Capital Planning", "N A STOP"],
            ["Basic Data Readiness", "N A STOP; provider ausente es NOT_AVAILABLE"],
            ["QRA y Confidence engine", "N A STOP; Risk Readiness no equivale a QRA"],
            ["Owner Capability Readiness", "N A STOP"],
            ["Operational Readiness y Benefits Realization", "N A STOP"],
            ["Procurement execution y Project Delivery", "N A STOP"],
        ],
        [2.85, 4.25],
    )

    add_heading(document, "12 Matriz de los 89 entregables")
    architecture_present = repository_presence(
        "backend/app/modules/investment_evidence_policy/decision_criteria.py",
        "backend/app/modules/investment_gate_preparation/service.py",
        "backend/app/modules/investment_decision/service.py",
    )
    admin_ui_present = repository_presence(
        "frontend/src/features/investment-evidence-policy/InvestmentEvidencePolicyAdminView.tsx"
    )
    gate07j_ui_present = repository_presence(
        "frontend/src/features/investment-gate-preparation/InvestmentGatePreparationWorkspace.tsx"
    )
    gate07k_ui_present = repository_presence(
        "frontend/src/features/investment-decision/InvestmentDecisionWorkspace.tsx"
    )
    static_status = evidence_status(architecture_present, checks["backend_focused"].passed)
    api_status = evidence_status(architecture_present, checks["openapi"].passed)
    admin_status = evidence_status(admin_ui_present, checks["browser_admin"].passed)
    gate07j_status = evidence_status(gate07j_ui_present, checks["browser_gate07j"].passed)
    gate07k_status = evidence_status(gate07k_ui_present, checks["browser_gate07k"].passed)
    postgres_ffa_status = checks["postgres_ffa"].status
    postgres_select_status = checks["postgres_select"].status
    backend_status = checks["backend_focused"].status
    regression_status = checks["backend_strategic"].status
    deliverables = [
        ("Executive summary", "Seccion 1", "DOCUMENTADO"),
        ("ADR Gate07J EVID 2", "docs/68-adr-gate07j-evid2-decision-criteria.md", evidence_status(ADR.exists())),
        ("Actual baseline assessment", "Seccion 2 y evidencia Alembic", checks["alembic"].status),
        ("REUSE BEFORE CREATE analysis", "Seccion 2", "DOCUMENTADO"),
        ("Existing 07J EVID preservation proof", "Taxonomia A B C y legacy", backend_status),
        ("Existing 07K authority preservation proof", "Formal authority Gate07K", backend_status),
        ("Evidence Class taxonomy", "Seccion 3 1", static_status),
        ("Decision Criterion taxonomy", "Seccion 3 2", static_status),
        ("Readiness Business Evaluation Portfolio Context model", "Separacion de cuatro estados", static_status),
        ("Authority matrix", "Secciones 2 1 y 4", "DOCUMENTADO"),
        ("Provider registry assessment", "Keys canónicas", backend_status),
        ("Business Case provider integration", "Provider y lineage Five Case", backend_status),
        ("FIN provider integration", "Provider ECONOMIC_FINANCIAL", backend_status),
        ("Commercial provider integration", "Provider COMMERCIAL", backend_status),
        ("Benefits provider integration", "Provider BENEFITS y BEN continuity", backend_status),
        ("FEL provider integration", "Provider GATE07J FEL", backend_status),
        ("Funding provider integration", "Provider GATE07J FUND", backend_status),
        ("Risk provider integration", "RISK_READINESS_PROVIDER", backend_status),
        ("Permit provider integration", "PERMIT_READINESS_PROVIDER", backend_status),
        ("Evidence quality confidence design", "Readiness distinta de quality", backend_status),
        ("Progressive maturity model", "FEL y Gate applicability", backend_status),
        ("PYP INV SELECT policy design", "Seccion 5", postgres_select_status),
        ("PYP FFA policy design", "Seccion 5", postgres_ffa_status),
        ("Project Type tailoring", "Criterion filters", backend_status),
        ("Governance Model tailoring", "Capital owner opt in", regression_status),
        ("Uncertainty complexity tailoring", "Applicability filters", backend_status),
        ("Decision Criteria data model", "AdminConfiguration content_json", static_status),
        ("Controlled operators", "Seccion 3 3", backend_status),
        ("Readiness check design", "READINESS_CHECK", backend_status),
        ("Must meet design", "MUST_MEET", backend_status),
        ("Decision hurdle design", "DECISION_HURDLE", backend_status),
        ("Prioritization design", "Portfolio reference only", backend_status),
        ("Success Criteria design", "Candidates BEN BC FIN FEL", backend_status),
        ("Portfolio authority preservation", "No rank calculation", backend_status),
        ("No double ranking proof", "NO_GATE_DOUBLE_SCORING", backend_status),
        ("Gate07J frozen package design", "PPMIS GATE07J DECISION PREPARATION v2", backend_status),
        ("Criterion evaluation snapshot", "Values operators result and hash", backend_status),
        ("Policy versioning", "Clone publish revision hash", backend_status),
        ("Effective dating", "Current policy resolution", backend_status),
        ("Historical package compatibility", "Legacy no reinterpretation", regression_status),
        ("New round semantics", "New preparation after closed round", backend_status),
        ("Gate07K consumption contract", "Decision package v2", backend_status),
        ("GO guardrails", "Package Class A MUST_MEET HURDLE", backend_status),
        ("HOLD semantics", "Context may constrain advancement", backend_status),
        ("RECYCLE semantics", "New round and new preparation", backend_status),
        ("KILL semantics", "Human decision only", backend_status),
        ("FID Success Criteria snapshot", "Immutable snapshot and hash", backend_status),
        ("Benefit identity continuity", "Stable BEN references", backend_status),
        ("Project Set coverage", "Exact denominator", backend_status),
        ("Cross source consistency", "BC against governed sources", backend_status),
        ("ADMIN configuration", "Existing policy lifecycle", admin_status),
        ("ADMIN UI", "Evidence and Decision Policy", admin_status),
        ("USER Gate07J UI", "Matrix groups context and freeze", gate07j_status),
        ("Gate07K decision UI", "Frozen criteria context and decision", gate07k_status),
        ("SoD", "Roles and authority", regression_status),
        ("ETag", "If Match and 412", backend_status),
        ("Idempotency", "Mutating lifecycle actions", backend_status),
        ("SecurityEvent", "Policy criterion package and decision events", backend_status),
        ("Persistence decision", "Existing JSON snapshots", static_status),
        ("Migration decision", "No migration and head 0061", checks["alembic"].status),
        ("API changes", "Additive schemas and routes", api_status),
        ("PostgreSQL PYP FFA happy path", checks["postgres_ffa"].detail, postgres_ffa_status),
        ("PostgreSQL PYP INV SELECT happy path", checks["postgres_select"].detail, postgres_select_status),
        ("Missing BC test", "Backend focal", backend_status),
        ("Stale FIN test", "Backend focal", backend_status),
        ("Benefits coverage test", "Backend focal", backend_status),
        ("Commercial readiness test", "Backend focal", backend_status),
        ("Funding mismatch test", "Backend focal", backend_status),
        ("Permit context test", "Backend focal", backend_status),
        ("Early financial overreach test", "Backend focal Select", backend_status),
        ("Portfolio resource HOLD context test", "Backend focal", backend_status),
        ("Readiness not KILL test", "Backend focal", backend_status),
        ("Historical policy test", "Backend focal and regression", regression_status),
        ("New round policy test", "Backend focal", backend_status),
        ("Success Criteria snapshot test", "Backend focal", backend_status),
        ("Contractor Direct regression", "Strategic regression", regression_status),
        ("Project ID14 regression", "Strategic regression", regression_status),
        ("Browser USER E2E", checks["browser_gate07j"].detail, checks["browser_gate07j"].status),
        ("Browser ADMIN E2E", checks["browser_admin"].detail, checks["browser_admin"].status),
        ("Browser Gate07K E2E", checks["browser_gate07k"].detail, checks["browser_gate07k"].status),
        ("Performance assessment", checks["performance"].detail, checks["performance"].status),
        ("Full strategic regression", checks["backend_strategic"].detail, regression_status),
        ("Technical debt", "Validation summary y Seccion 10", "DOCUMENTADO"),
        ("Strategic Capital Buckets gap note", "N A STOP", "N A STOP"),
        ("Periodic Portfolio Review gap note", "N A STOP", "N A STOP"),
        ("QRA Basic Data Owner Capability gap note", "N A STOP", "N A STOP"),
        ("Updated technical report", "Este documento", "DOCUMENTADO"),
        ("Final exclusive state", release_signal, "PASS" if release_validated else "PENDIENTE"),
        ("Recommendation for next increment", "Strategic gap review without auto implementation", "DOCUMENTADO"),
    ]
    assert len(deliverables) == 89
    rows = [
        [index, name, evidence, status]
        for index, (name, evidence, status) in enumerate(deliverables, start=1)
    ]
    for start in range(0, len(rows), 18):
        if start:
            document.add_page_break()
            add_heading(
                document,
                f"12 {start // 18 + 1} Entregables {start + 1} a {min(start + 18, len(rows))}",
                2,
            )
        add_table(
            document,
            ["Numero", "Entregable", "Evidencia", "Estado"],
            rows[start : start + 18],
            [0.5, 2.35, 3.35, 0.9],
            status_column=3,
            font_size=7.2,
        )

    document.add_page_break()
    add_heading(document, "13 Estado final y recomendacion")
    add_table(
        document,
        ["Señal", "Valor", "Estado"],
        [
            ["Release", release_signal, "PASS" if release_validated else "PENDIENTE"],
            ["Final output", final_output, "PASS" if release_validated else "PENDIENTE"],
        ],
        [1.15, 5.05, 0.9],
        status_column=2,
    )
    add_paragraph(
        document,
        (
            "Gate07J puede preparar un paquete gobernado completo y Gate07K puede decidir formalmente usando ese paquete. Este estado no significa FID automatico."
            if release_validated
            else "La arquitectura y el avance quedan documentados, pero el cierre no se declara hasta que todos los checks obligatorios figuren PASS en los artefactos de QA."
        ),
    )
    add_heading(document, "13 1 Recomendacion posterior", 2)
    add_paragraph(
        document,
        "Realizar una revision de brechas estrategicas para Strategic Capital Buckets, Periodic Portfolio Review, "
        "Capital Demand, Basic Data Readiness, QRA Confidence y Owner Capability Readiness. El siguiente incremento "
        "debe elegirse despues de revisar este cierre; ninguna de esas capacidades se implementa automaticamente.",
    )
    add_heading(document, "13 2 Stop boundary", 2)
    add_paragraph(
        document,
        "El trabajo termina en Gate07J EVID 2. No inicia Procurement execution, Project Delivery, Operational "
        "Readiness, Benefits Realization, Strategic Capital Buckets, Portfolio Review, Capital Demand, Basic Data, "
        "QRA ni Owner Capability.",
    )

    document.core_properties.title = "Informe tecnico Gate 07J EVID 2"
    document.core_properties.subject = (
        "Politica de evidencia criterios de decision y revalidacion estrategica"
    )
    document.core_properties.author = "P&Pmis Construction AI"
    document.core_properties.keywords = (
        "Gate07J EVID 2, Gate07K, decision criteria, evidence, Portfolio, FID"
    )
    document.core_properties.comments = (
        "Informe tecnico de implementacion validacion y cierre controlado"
    )
    document.core_properties.created = datetime(2026, 9, 20, tzinfo=UTC)
    document.core_properties.modified = datetime(2026, 9, 20, tzinfo=UTC)
    return document


def main() -> None:
    checks, summary = collect_checks()
    document = make_document(checks, summary)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
