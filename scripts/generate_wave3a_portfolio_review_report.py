from __future__ import annotations

import re
from pathlib import Path

import generate_cpig_gap_review_report as base
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "75-wave3a-periodic-portfolio-review-release.md"
QA = ROOT / "outputs" / ".qa" / "wave3a-portfolio-review"
DIAGRAM_DIR = QA / "diagrams"
REPORT_DATE = "2026-10-07"
REPORT_DATE_DISPLAY = "7 de octubre de 2026"
OUTPUT = (
    ROOT
    / "outputs"
    / f"Informe_Tecnico_Wave3A_Periodic_Portfolio_Review_Governance_Cycle_{REPORT_DATE}.docx"
)
LOGO = ROOT / "frontend" / "public" / "pypmis-construction-ai-logo.png"

DIAGRAMS = [
    {
        "name": "01-wave3a-governance-loop.png",
        "title": "Ciclo recurrente de gobierno de portafolio",
        "rows": [
            ["Capital Planning", "Allocation", "Portfolio Planning"],
            ["Scenarios", "Recommendation", "Periodic Review"],
            ["Conclusions", "Governance Actions", "Follow-up / Next Review"],
        ],
        "edges": [
            (0, 0, 0, 1),
            (0, 1, 0, 2),
            (0, 2, 1, 0),
            (1, 0, 1, 1),
            (1, 1, 1, 2),
            (1, 2, 2, 0),
            (2, 0, 2, 1),
            (2, 1, 2, 2),
            (2, 2, 1, 2),
        ],
        "alt": (
            "Flujo desde planificación y asignación de capital hacia escenarios, recomendación, "
            "revisión periódica, conclusiones, acciones y la siguiente revisión."
        ),
    },
    {
        "name": "02-wave3a-authority-boundary.png",
        "title": "Frontera de autoridad y adaptadores de lectura",
        "rows": [
            ["Membership / Rank", "Wave 2A Allocation", "Budget / Resources"],
            ["Gate07I", "Read-only adapters", "Gate07J / Gate07K"],
            ["Frozen Review Snapshot", "Conclusions", "Request-only Actions"],
        ],
        "edges": [
            (0, 0, 1, 1),
            (0, 1, 1, 1),
            (0, 2, 1, 1),
            (1, 0, 1, 1),
            (1, 2, 1, 1),
            (1, 1, 2, 0),
            (2, 0, 2, 1),
            (2, 1, 2, 2),
        ],
        "alt": (
            "Fuentes de Membership, Rank, Allocation, Budget, Resources y Gates convergen mediante "
            "adaptadores read-only a un snapshot, conclusiones y acciones sin mutación silenciosa."
        ),
    },
    {
        "name": "03-wave3a-lifecycle.png",
        "title": "Lifecycle gobernado de Review Session",
        "rows": [
            ["DRAFT", "PREPARING", "READY FOR REVIEW"],
            ["IN REVIEW", "COMPLETED / immutable", "Open follow-up actions"],
        ],
        "edges": [
            (0, 0, 0, 1),
            (0, 1, 0, 2),
            (0, 2, 1, 0),
            (1, 0, 1, 1),
            (1, 1, 1, 2),
        ],
        "alt": (
            "Lifecycle desde borrador y preparación hasta revisión, cierre inmutable y acciones de "
            "seguimiento que pueden permanecer abiertas."
        ),
    },
    {
        "name": "04-wave3a-admin-user.png",
        "title": "Separación estricta ADMIN y USER",
        "rows": [
            ["ADMIN policy", "Preview persisted=false", "Publish / Archive"],
            ["USER Session", "Snapshot / Review", "Complete / History"],
        ],
        "edges": [
            (0, 0, 0, 1),
            (0, 1, 0, 2),
            (0, 2, 1, 0),
            (1, 0, 1, 1),
            (1, 1, 1, 2),
        ],
        "alt": (
            "ADMIN configura y publica política reusable; USER opera sesiones live, snapshots, "
            "conclusiones, acciones e historia."
        ),
    },
]


def read_release_status() -> str:
    source_text = SOURCE.read_text(encoding="utf-8")
    match = re.search(r"^\*\*Estado actual:\*\*\s+(.+?)$", source_text, re.MULTILINE)
    if match is not None:
        status = match.group(1).strip().strip("`")
        return (
            "RELEASE VALIDADO"
            if status == "PERIODIC_PORTFOLIO_REVIEW_AND_GOVERNANCE_CYCLE_RELEASE_VALIDATED"
            else status
        )
    if "PERIODIC_PORTFOLIO_REVIEW_AND_GOVERNANCE_CYCLE_RELEASE_VALIDATED" in source_text:
        return "RELEASE VALIDADO"
    return "VALIDACIÓN EN CURSO"


def configure_metadata(document: Document) -> None:
    core = document.core_properties
    core.title = "Informe técnico Wave 3A Periodic Portfolio Review and Governance Cycle"
    core.subject = "Implementación y validación del ciclo periódico de gobierno de portafolio"
    core.author = "P&Pmis Construction AI"
    core.keywords = "CPIG, Portfolio Review, Governance Cycle, Snapshot, Portfolio Actions"
    core.comments = "Informe técnico reproducible de implementación y release Wave 3A."


def add_cover(document: Document, release_status: str) -> None:
    logo = document.add_paragraph()
    logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo.paragraph_format.space_before = Pt(28)
    if LOGO.exists():
        shape = logo.add_run().add_picture(str(LOGO), width=Inches(1.3))
        base.set_alt_text(
            shape,
            "P&Pmis Construction AI",
            "Logotipo azul de P&Pmis Construction AI.",
        )

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(22)
    title.paragraph_format.space_after = Pt(12)
    title.add_run("Informe técnico Wave 3A\nPeriodic Portfolio Review and Governance Cycle")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.set_run_font(subtitle.add_run("P&Pmis Construction AI"), size=13, color=base.BLACK)

    metadata = document.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_before = Pt(22)
    base.add_inline(
        metadata,
        f"Capital Portfolio and Investment Governance | {REPORT_DATE_DISPLAY}",
        size=10,
    )

    status_label = document.add_paragraph()
    status_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status_label.paragraph_format.space_before = Pt(38)
    base.set_run_font(status_label.add_run("ESTADO DE VALIDACIÓN"), size=8, bold=True, color=base.MUTED)

    status = document.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.set_run_font(status.add_run(release_status.replace("_", " ")), size=11, bold=True, color=base.BLACK)
    document.add_page_break()


def add_document_control(document: Document, release_status: str) -> None:
    document.add_heading("Control del documento", level=1)
    base.add_table(
        document,
        [
            ["Campo", "Valor"],
            ["Producto", "P&Pmis Construction AI"],
            ["Macroproceso", "Capital Portfolio and Investment Governance"],
            ["Incremento", "Wave 3A Periodic Portfolio Review and Governance Cycle"],
            ["Rama y commit baseline", "main / 9cb0969"],
            ["Migración objetivo", "20261005_0064 sobre 20261001_0063"],
            ["Fecha de corte", REPORT_DATE],
            ["Estado al corte", release_status],
            ["Fuente", "Código, migración, pruebas, PostgreSQL, Chromium y guardrails"],
        ],
    )

    document.add_heading("Lectura ejecutiva", level=1)
    paragraph = document.add_paragraph()
    base.add_inline(
        paragraph,
        (
            "Wave 3A implementa una revisión periódica del portafolio basada en snapshots inmutables, "
            "conclusiones y acciones gobernadas. Consume Allocation, Budget, Resource, Rank y Gate/FID "
            "sin asumir su autoridad ni generar mutaciones cruzadas."
        ),
    )

    document.add_heading("Contenido", level=1)
    for item in (
        "Baseline, ADR, autoridad y reutilización",
        "Review Session, lifecycle, snapshot y staleness",
        "Conclusiones, acciones y límites de dispatch",
        "Integraciones read-only y Project Matrix",
        "ADMIN, USER, seguridad, persistencia, API y navegación",
        "PostgreSQL, Chromium, regresiones, brechas y cierre",
    ):
        paragraph = document.add_paragraph(style="List Bullet")
        base.add_inline(paragraph, item)
    document.add_page_break()


def add_headers_and_footers(document: Document) -> None:
    for section in document.sections:
        section.different_first_page_header_footer = True
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        base.set_run_font(
            header.add_run("P&Pmis Construction AI  |  Wave 3A Portfolio Review"),
            size=7.5,
            bold=True,
            color=base.BLACK,
        )
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        base.set_run_font(
            footer.add_run(f"Informe técnico | {REPORT_DATE} | Página "),
            size=7.5,
            color=base.MUTED,
        )
        base.add_field(footer, "PAGE")
        base.set_run_font(footer.add_run(" de "), size=7.5, color=base.MUTED)
        base.add_field(footer, "NUMPAGES")


def format_tables(document: Document) -> None:
    """Give evidence columns room and keep short status values compact."""
    widths_by_headers = {
        ("Campo", "Valor"): (2.10, 4.96),
        ("Evidencia", "Baseline"): (2.10, 4.96),
        ("Dimensión", "Portfolio Review", "Project Gate"): (1.20, 2.93, 2.93),
        ("Dominio", "Autoridad", "Uso Wave 3A", "Mutación Wave 3A"): (1.25, 1.75, 2.80, 1.26),
        ("Capacidad", "ADMIN", "USER"): (3.50, 1.55, 2.01),
        ("Incremento candidato", "Dependencia/valor", "Reuso", "Complejidad y riesgo de duplicación"): (
            1.75, 2.05, 1.40, 1.86,
        ),
        ("Ejecución", "Fecha Bogotá", "Resultado observado", "Alcance"): (1.55, 1.20, 1.70, 2.61),
        ("Criterio", "Estado", "Evidencia"): (2.30, 0.75, 4.01),
    }
    for table in document.tables:
        headers = tuple(cell.text for cell in table.rows[0].cells)
        widths = widths_by_headers.get(headers)
        table.autofit = False
        if widths is not None:
            for column, width in zip(table.columns, widths, strict=True):
                column.width = Inches(width)
        borders = table._tbl.tblPr.find(qn("w:tblBorders"))
        if borders is not None:
            for border in borders:
                border.set(qn("w:color"), "D9D9D9")
        for row_index, row in enumerate(table.rows):
            row_properties = row._tr.get_or_add_trPr()
            if row_properties.find(qn("w:cantSplit")) is None:
                row_properties.append(OxmlElement("w:cantSplit"))
            for index, cell in enumerate(row.cells):
                if widths is not None:
                    cell.width = Inches(widths[index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.05
                    if row_index == 0:
                        paragraph.paragraph_format.keep_with_next = True
                    if headers[index] in {"Estado", "ADMIN", "USER"}:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.size = Pt(9.5)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    QA.mkdir(parents=True, exist_ok=True)
    DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)

    base.SOURCE = SOURCE
    base.QA = QA
    base.DIAGRAM_DIR = DIAGRAM_DIR
    base.DIAGRAMS = DIAGRAMS
    diagram_paths = [base.render_diagram({**spec, "title_color": (0, 0, 0), "title_rule": False}) for spec in DIAGRAMS]
    release_status = read_release_status()

    document = Document()
    base.configure_document(document)
    for style_name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        properties = document.styles[style_name]._element.find(qn("w:pPr"))
        if properties is not None:
            for border in properties.findall(qn("w:pBdr")):
                properties.remove(border)
    configure_metadata(document)
    base.set_language(document)
    add_cover(document, release_status)
    add_document_control(document, release_status)
    source_start = len(document.paragraphs)
    embedded_diagrams = SOURCE.read_text(encoding="utf-8").count("```mermaid")
    if embedded_diagrams == 0:
        base.DIAGRAMS = []
    base.add_source_content(
        document,
        diagram_paths,
        page_break_heading_prefixes=(),
        skip_preface=True,
    )
    if embedded_diagrams == 0:
        for heading_number, path, spec in zip(("4. ", "6. ", "9. ", "43. "), diagram_paths, DIAGRAMS, strict=True):
            anchor = next(paragraph for paragraph in document.paragraphs
                          if paragraph.style.name.startswith("Heading") and paragraph.text.startswith(heading_number))
            figure = document.add_paragraph()
            figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
            figure.paragraph_format.keep_together = True
            shape = figure.add_run().add_picture(str(path), width=Inches(6.75))
            base.set_alt_text(shape, spec["title"], spec["alt"])
            figure.add_run("\n")
            base.set_run_font(figure.add_run(spec["title"]), size=9, color=base.BLACK)
            anchor._p.addnext(figure._p)
    # Keep the required numbered outline while removing decorative punctuation.
    for index, paragraph in enumerate(document.paragraphs):
        if paragraph.style.name.startswith("Heading"):
            text = re.sub(r"[^\w\s]", " ", paragraph.text).replace("_", " ")
            paragraph.text = re.sub(r"\s+", " ", text).strip()
        elif index >= source_start and paragraph.style.name in {"Normal", "List Bullet", "List Number"} and (
            "w:drawing" not in paragraph._p.xml
        ):
            for run in paragraph.runs:
                run.font.size = Pt(11)
    format_tables(document)
    add_headers_and_footers(document)
    base.update_field_settings(document)
    document.save(OUTPUT)
    print(f"created={OUTPUT}")
    print(f"diagrams={len(diagram_paths)}")
    print(f"source={SOURCE}")


if __name__ == "__main__":
    main()
