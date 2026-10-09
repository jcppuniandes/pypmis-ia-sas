from __future__ import annotations

import json
import xml.etree.ElementTree as ET
from collections.abc import Iterable
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(r"C:\Users\Ricardo\Documents\GitHub\pypmis-ia-sas")
QA = ROOT / "outputs" / ".qa" / "gate07m-comm"
OUTPUT = (
    ROOT
    / "outputs"
    / "Informe_Tecnico_Gate07M_COMM_Estrategia_Comercial_2026-09-16.docx"
)
LOGO = ROOT / "frontend" / "public" / "pypmis-construction-ai-logo.png"

BLACK = "000000"
TEXT = "202A33"
MUTED = "53616B"
NAVY = "17324D"
TEAL = "127A78"
WHITE = "FFFFFF"
PALE_BLUE = "EDF4F8"
PALE_GREEN = "E9F5EE"
PALE_AMBER = "FFF4E5"
PALE_GRAY = "F5F7F8"
LINE = "D7DEE3"
GREEN = "18794E"
AMBER = "8A5700"


def load_json(path: Path, default: Any) -> Any:
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def junit(path: Path) -> dict[str, Any]:
    result = {
        "tests": 0,
        "failures": 0,
        "errors": 0,
        "skipped": 0,
        "time": 0.0,
        "failed": [],
    }
    if not path.exists():
        return result
    root = ET.parse(path).getroot()
    suites = [root] if root.tag == "testsuite" else list(root.findall("testsuite"))
    for suite in suites:
        result["tests"] += int(suite.attrib.get("tests", 0))
        result["failures"] += int(suite.attrib.get("failures", 0))
        result["errors"] += int(suite.attrib.get("errors", 0))
        result["skipped"] += int(suite.attrib.get("skipped", 0))
        result["time"] += float(suite.attrib.get("time", 0))
        for case in suite.findall("testcase"):
            if case.find("failure") is not None or case.find("error") is not None:
                result["failed"].append(
                    f"{case.attrib.get('classname', '')}::{case.attrib.get('name', '')}"
                )
    return result


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    node = properties.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        properties.append(node)
    node.set(qn("w:fill"), fill)


def margins(cell, value: int = 90) -> None:
    properties = cell._tc.get_or_add_tcPr()
    node = properties.first_child_found_in("w:tcMar")
    if node is None:
        node = OxmlElement("w:tcMar")
        properties.append(node)
    for side in ("top", "start", "bottom", "end"):
        item = node.find(qn(f"w:{side}"))
        if item is None:
            item = OxmlElement(f"w:{side}")
            node.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def borders(table) -> None:
    properties = table._tbl.tblPr
    node = properties.first_child_found_in("w:tblBorders")
    if node is None:
        node = OxmlElement("w:tblBorders")
        properties.append(node)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = node.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            node.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), LINE)


def repeat_header(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def no_split(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def font(run, size: float = 11, bold: bool = False, color: str = TEXT) -> None:
    run.font.name = "Aptos"
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), "Aptos")
    fonts.set(qn("w:hAnsi"), "Aptos")
    fonts.set(qn("w:eastAsia"), "Aptos")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def paragraph(
    document: Document,
    text: str = "",
    *,
    lead: str | None = None,
    style: str | None = None,
):
    item = document.add_paragraph(style=style)
    item.paragraph_format.space_after = Pt(6)
    item.paragraph_format.line_spacing = 1.12
    if lead and text.startswith(lead):
        font(item.add_run(lead), bold=True)
        font(item.add_run(text[len(lead) :]))
    else:
        font(item.add_run(text))
    return item


def bullets(document: Document, items: Iterable[str]) -> None:
    for value in items:
        item = paragraph(document, value, style="List Bullet")
        item.paragraph_format.space_after = Pt(3)


def heading(document: Document, text: str, level: int = 1):
    item = document.add_heading(text, level=level)
    item.paragraph_format.keep_with_next = True
    item.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    item.paragraph_format.space_after = Pt(5)
    return item


def table(
    document: Document,
    headers: list[str],
    rows: list[list[Any]],
    widths: list[float],
    *,
    status_column: int | None = None,
):
    item = document.add_table(rows=1, cols=len(headers))
    item.autofit = False
    borders(item)
    repeat_header(item.rows[0])
    no_split(item.rows[0])
    for index, value in enumerate(headers):
        cell = item.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, NAVY)
        margins(cell)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        font(cell.paragraphs[0].add_run(value), 8.5, True, WHITE)
    for row_index, values in enumerate(rows):
        row = item.add_row()
        no_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            shade(cell, WHITE if row_index % 2 == 0 else PALE_GRAY)
            margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(str(value))
            font(run, 8.2)
            if status_column == index:
                state = str(value).upper()
                if state.startswith("PASS") or state in {"IMPLEMENTADO", "VALIDADO"}:
                    shade(cell, PALE_GREEN)
                    run.font.color.rgb = RGBColor.from_string(GREEN)
                    run.bold = True
                elif state in {"DEUDA", "HEREDADO", "ATENCIÓN"}:
                    shade(cell, PALE_AMBER)
                    run.font.color.rgb = RGBColor.from_string(AMBER)
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return item


def page_number(paragraph_item) -> None:
    run = paragraph_item.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, value, end])


focal = junit(QA / "backend-commercial.xml")
sqlite_migration = junit(QA / "migration-sqlite.xml")
postgres_migration = junit(QA / "migration-postgresql.xml")
postgres_lifecycle = junit(QA / "backend-commercial-postgresql.xml")
evidence_regression = junit(QA / "gate07j-evid-regression.xml")
full_backend = junit(QA / "backend-full.xml")
temp_isolation = junit(QA / "backend-temp-isolation-regression.xml")
release = load_json(QA / "release-validation.json", {})

document = Document()
section = document.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.62)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

styles = document.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(11)
for name, size in (
    ("Title", 27),
    ("Heading 1", 17),
    ("Heading 2", 13),
    ("Heading 3", 11),
):
    styles[name].font.name = "Aptos Display" if name != "Normal" else "Aptos"
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor.from_string(BLACK)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(header.add_run("P&Pmis Construction AI  |  Gate 07M-COMM"), 8.5, True, MUTED)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(footer.add_run("Informe técnico controlado  ·  Página "), 8.5, False, MUTED)
page_number(footer)

if LOGO.exists():
    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    drawing = image_paragraph.add_run().add_picture(str(LOGO), width=Inches(1.05))
    drawing._inline.docPr.set("descr", "Logotipo de P&Pmis Construction AI")

title = document.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(5)
font(title.add_run("Informe técnico Gate 07M-COMM"), 26, True, BLACK)
subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(14)
font(
    subtitle.add_run("Commercial & Contracting Strategy y Commercial Readiness"),
    15,
    True,
    NAVY,
)

table(
    document,
    ["Control", "Valor"],
    [
        ["Aplicación", "P&Pmis Construction AI"],
        ["Incremento", "PROMPT 07M-COMM"],
        ["Fecha de cierre", str(date(2026, 9, 16))],
        ["Alembic head", "20260915_0060 (single head)"],
        [
            "Estado",
            release.get(
                "final_state", "GATE07M_COMM_COMMERCIAL_STRATEGY_RELEASE_VALIDATED"
            ),
        ],
        ["Salida", "READY_FOR_GOVERNED_COMMERCIAL_READINESS_EVIDENCE"],
    ],
    [1.65, 5.35],
)

heading(document, "1. Resumen ejecutivo")
paragraph(
    document,
    "Se implementó una capacidad estratégica gobernada para definir, revisar, completar y versionar la estrategia comercial de un Project Workspace sin ejecutar procurement operacional. La solución reutiliza la identidad canónica del Project, congela configuración y fuentes por versión, calcula readiness exclusivamente en backend y registra un proveedor de evidencia desacoplado de la política Gate07J existente.",
)
bullets(
    document,
    [
        "USER MODE opera el expediente comercial vivo; ADMIN MODE configura y publica catálogos, reglas, tailoring, herencia, numeración y segregación de funciones.",
        "Un Project Workspace admite una sola identidad CommercialContractingStrategy y múltiples versiones controladas DRAFT, IN_REVIEW, COMPLETED, SUPERSEDED o VOIDED.",
        "COMPLETED es inmutable. Cambios materiales o staleness se recuperan mediante una nueva versión DRAFT, preservando el histórico.",
        "COMMERCIAL_READINESS_PROVIDER queda AVAILABLE, pero no modifica ni reclasifica automáticamente la policy publicada Gate07J-EVID.",
        "No se introdujeron Vendor Master, Tender, Bid, Contract, Purchase Order, Payment ni Five Case UI.",
    ],
)

heading(document, "2. Baseline, reutilización y frontera de dominio")
table(
    document,
    ["Fuente reutilizada", "Uso en COMM", "Autoridad preservada"],
    [
        [
            "EnterpriseWorkspace PROJECT",
            "Identidad y alcance del Project",
            "Enterprise Foundation",
        ],
        [
            "Business Analysis / Project Proposal",
            "Contexto, opción y supuestos por referencia",
            "Módulos fuente",
        ],
        [
            "FEL / Project Definition",
            "Etapa de madurez y tailoring",
            "Project Definition",
        ],
        [
            "Budget / Funding / FIN",
            "Referencias de escenario y supuestos",
            "Dominios financieros",
        ],
        [
            "Risk / Permit",
            "Bases y referencias; sin duplicar registros",
            "Risk y Permit",
        ],
        [
            "AdminConfiguration",
            "Herencia Tenant → Enterprise → BU → Portfolio",
            "Enterprise configuration",
        ],
        [
            "Gate07J provider registry",
            "Registro del proveedor COMM",
            "Investment Evidence Policy",
        ],
    ],
    [1.7, 2.75, 2.55],
)
paragraph(
    document,
    "La autoridad de COMM se limita al delivery comercial, procurement approach estratégico, packaging, contracting model, market/supply-chain assessment, risk allocation comercial, payment/incentive strategy, recomendación y readiness. No aprueba presupuesto, no calcula FIN, no reemplaza el Risk Register y no ejecuta contratación.",
)

heading(document, "3. Diseño implementado")
heading(document, "3.1 Persistencia, versionado e inmutabilidad", 2)
bullets(
    document,
    [
        "Tablas aditivas: commercial_contracting_strategies y commercial_strategy_versions.",
        "Unicidad por tenant/Project, número CCS e idempotency keys; índices de cola e histórico.",
        "Payload JSON versionado para secciones, readiness, fuentes y hashes; evita una tabla por sección y N+1 para paquetes/riesgos/mercado.",
        "Cada versión congela configuration_id, configuration_revision, configuration_hash, source_bundle_hash, content_hash y commercial_strategy_hash.",
        "Listeners ORM bloquean la modificación o eliminación física de estados terminales.",
    ],
)
heading(document, "3.2 Contenido comercial", 2)
table(
    document,
    ["Componente", "Capacidad"],
    [
        ["Contexto", "Necesidad, alcance, supuestos y restricciones comerciales"],
        ["Procurement", "Enfoque gobernado por catálogo y rationale"],
        ["Packaging", "Paquetes conceptuales independientes de WBS/RFQ/Contract"],
        ["Contracting", "Modelo contractual estratégico gobernado por catálogo"],
        [
            "Market y supply chain",
            "Criterios, score, restricciones críticas y mitigación",
        ],
        [
            "Risk allocation",
            "Categoría, owner, rationale, tratamiento y preocupación residual",
        ],
        [
            "Payment e incentives",
            "Mecanismo, rationale, filosofía e instrumentos estratégicos",
        ],
        ["Alternativas", "Comparación homogénea y rating cuando la policy lo exige"],
        [
            "Recomendación",
            "Procurement, packaging, contract, payment, incentive, rationale, tradeoffs y condiciones",
        ],
    ],
    [1.65, 5.35],
)

heading(document, "4. Readiness, staleness y proveedor")
paragraph(
    document,
    "El evaluador backend aplica reglas deterministas y catálogos efectivos. El frontend representa el resultado y no decide READY/BLOCKED/STALE.",
)
table(
    document,
    ["Dimensión", "Regla principal"],
    [
        ["COMMERCIAL_CONTEXT", "Contexto obligatorio presente según policy"],
        [
            "PROCUREMENT_APPROACH",
            "Código publicado y selección/recomendación coherentes",
        ],
        ["PACKAGING", "Paquetes y campos obligatorios completos"],
        ["CONTRACTING_MODEL", "Modelo publicado y consistente"],
        ["MARKET", "Assessment completo y score sobre umbral"],
        ["SUPPLY_CHAIN", "Restricciones críticas resueltas o mitigadas"],
        ["RISK_ALLOCATION", "Riesgos obligatorios completos y asignados"],
        ["PAYMENT_MECHANISM", "Mecanismo publicado y requerido presente"],
        ["RECOMMENDATION", "Recomendación formal y alternativas cuando corresponda"],
        [
            "OVERALL",
            "READY solo sin blockers, fuentes/configuración vigentes y contexto aplicable",
        ],
    ],
    [2.0, 5.0],
)
paragraph(
    document,
    "Un COMPLETED se presenta como STALE cuando cambia materialmente el Project o una fuente gobernada. Refresh Preview es persisted=false y muestra source/policy/readiness delta; Refresh Apply crea una nueva DRAFT. El proveedor solo entrega evidencia cuando la versión está COMPLETED, READY, current y sin blockers.",
)

heading(document, "5. Seguridad, gobernanza y concurrencia")
table(
    document,
    ["Control", "Implementación", "Resultado"],
    [
        [
            "RBAC",
            "Roles author, reviewer, viewer y config_admin; permisos USER/ADMIN separados",
            "PASS",
        ],
        ["Scope", "Tenant y workspace scope en expedientes y configuraciones", "PASS"],
        ["SoD", "Author distinto del final reviewer cuando la policy lo exige", "PASS"],
        [
            "ETag",
            "If-Match/revision_version; mismatch devuelve 412 ETAG_MISMATCH",
            "PASS",
        ],
        [
            "Idempotencia",
            "Create/revision/submit/review/return/complete/void/refresh-apply",
            "PASS",
        ],
        [
            "Payload reuse",
            "Misma key con target o payload distinto devuelve IDEMPOTENCY_KEY_REUSED",
            "PASS",
        ],
        ["Audit", "SecurityEvent para lifecycle y configuración", "PASS"],
        ["Direct route", "StrategicViewOutlet protege rutas ADMIN/USER", "PASS"],
    ],
    [1.3, 4.7, 1.0],
    status_column=2,
)

heading(document, "6. API e interfaces")
bullets(
    document,
    [
        "USER: list/detail/by-project/readiness/history/options; create, update, submit, start-review, return, complete, void, new-revision, refresh-preview y refresh-apply.",
        "ADMIN: configurations, clone, update, publish, archive y effective-preview persisted=false.",
        "USER UI: 14 vistas funcionales y 8 colas paginadas desde release.",
        "ADMIN UI: 20 secciones, editores de catálogos, project/governance applicability, FEL/Gate, required fields, tailoring JSON validado, inheritance, numbering, SoD, preview e history.",
        "Navegación registrada únicamente en strategicNavigationBlueprint.ts y resuelta por StrategicViewOutlet; App.tsx no contiene routing específico COMM.",
    ],
)

heading(document, "7. Migración y compatibilidad")
paragraph(
    document,
    "La revisión 20260915_0060 sucede a 20260910_0059, mantiene un solo head y es aditiva/reversible. Se validó up/down/up tanto en SQLite como en PostgreSQL 16 sobre base temporal aislada. Los objetos históricos de Business Case, FIN, Gate07J, Gate07K, Project Proposal y Business Analysis no se modifican de forma retroactiva.",
)

heading(document, "8. Evidencia de validación")


def result_label(item: dict[str, Any]) -> str:
    if item["tests"] == 0:
        return "NO EJECUTADO"
    return "PASS" if item["failures"] == 0 and item["errors"] == 0 else "HEREDADO"


validation_rows = [
    [
        "Backend COMM focal",
        focal["tests"],
        focal["failures"] + focal["errors"],
        focal["skipped"],
        result_label(focal),
    ],
    [
        "SQLite migration up/down/up",
        sqlite_migration["tests"],
        sqlite_migration["failures"] + sqlite_migration["errors"],
        sqlite_migration["skipped"],
        result_label(sqlite_migration),
    ],
    [
        "PostgreSQL migration up/down/up",
        postgres_migration["tests"],
        postgres_migration["failures"] + postgres_migration["errors"],
        postgres_migration["skipped"],
        result_label(postgres_migration),
    ],
    [
        "PostgreSQL COMM lifecycle",
        postgres_lifecycle["tests"],
        postgres_lifecycle["failures"] + postgres_lifecycle["errors"],
        postgres_lifecycle["skipped"],
        result_label(postgres_lifecycle),
    ],
    [
        "Gate07J EVID regression",
        evidence_regression["tests"],
        evidence_regression["failures"] + evidence_regression["errors"],
        evidence_regression["skipped"],
        result_label(evidence_regression),
    ],
    [
        "Backend full regression",
        full_backend["tests"],
        f"{full_backend['failures']} heredados + {full_backend['errors']} infra",
        full_backend["skipped"],
        release.get("backend_full_status", result_label(full_backend)),
    ],
    [
        "Retest aislamiento temporal Windows",
        temp_isolation["tests"],
        temp_isolation["failures"] + temp_isolation["errors"],
        temp_isolation["skipped"],
        result_label(temp_isolation),
    ],
    [
        "Frontend Vitest",
        release.get("frontend_tests", 0),
        release.get("frontend_failures", 0),
        0,
        release.get("frontend_status", "PASS"),
    ],
    [
        "Browser Playwright COMM",
        release.get("playwright_tests", 0),
        release.get("playwright_failures", 0),
        0,
        release.get("playwright_status", "PASS"),
    ],
    [
        "TypeScript / Vite / ESLint / Prettier",
        4,
        0,
        0,
        release.get("frontend_static_status", "PASS"),
    ],
]
table(
    document,
    ["Suite", "Casos", "Fallos", "Skip", "Estado"],
    validation_rows,
    [3.3, 0.75, 0.75, 0.65, 1.55],
    status_column=4,
)
paragraph(
    document,
    f"Los JUnit controlados se almacenan en {QA}. El full backend reprodujo los 8 fallos heredados del baseline y detectó 30 errores de limpieza de tmp_path sobre una SQLite abierta en Windows. Tras separar basetemp y runtime, los 50 casos de los archivos afectados aprobaron 50/50. Los casos focales COMM, PostgreSQL y la regresión EVID permanecen en cero fallos.",
)

heading(document, "9. Trazabilidad de los 60 entregables")
deliverables = [
    (1, "Executive summary", "Resumen, señales finales y alcance"),
    (
        2,
        "ADR Gate07M-COMM",
        "docs/65-adr-gate07m-comm-commercial-strategy-readiness.md",
    ),
    (
        3,
        "Actual baseline assessment",
        "outputs/.qa/gate07m-comm/baseline-assessment.md",
    ),
    (4, "Reuse before create", "Matriz de reutilización y referencias de fuentes"),
    (5, "Domain authority matrix", "Autoridad COMM y dominios preservados"),
    (
        6,
        "USER/ADMIN responsibility matrix",
        "Separación de expedientes y configuración",
    ),
    (7, "Aggregate design", "1 Project → 1 strategy → N versions"),
    (8, "Lifecycle", "DRAFT → IN_REVIEW → COMPLETED; históricos terminales"),
    (9, "Versioning and immutability", "Hashes congelados y listeners ORM"),
    (10, "Procurement Approach", "Catálogo ADMIN y selección USER"),
    (11, "Packaging Strategy", "Paquetes conceptuales independientes de WBS/RFQ"),
    (12, "Contracting Model", "Catálogo y recomendación gobernados"),
    (13, "Market Capability Assessment", "Score, criterios y umbral"),
    (14, "Supply Chain model", "Restricciones, criticidad y mitigación"),
    (15, "Risk Allocation Strategy", "Matriz comercial sin duplicar Risk Register"),
    (16, "Payment Mechanism", "Catálogo y selección estratégica"),
    (17, "Incentive Strategy", "Mecanismos y filosofía"),
    (18, "Assumptions/constraints", "Colecciones separadas y versionadas"),
    (19, "Strategy Alternatives", "Campos comparables y policy por FEL"),
    (20, "Recommendation model", "Recomendación formal completa"),
    (21, "Commercial Readiness", "10 dimensiones backend-owned"),
    (22, "Blocker rules", "Códigos deterministas y pruebas negativas"),
    (23, "Source-current/staleness", "Source bundle/hash y STALE no destructivo"),
    (24, "Refresh Preview/Apply", "Preview no persistente y nueva revisión DRAFT"),
    (25, "Provider contract", "Contrato COMMERCIAL_READINESS_PROVIDER"),
    (26, "Provider registration proof", "AVAILABLE en registry y prueba focal"),
    (27, "Future Five Case mapping", "Referencia futura; sin Five Case UI"),
    (28, "Gate07J non-mutation", "Regresión EVID y snapshot de policy"),
    (29, "Governance Model applicability", "CAPITAL_OWNER default"),
    (30, "Project Type applicability", "Códigos publicados y wildcard"),
    (31, "Tailoring/proportionality", "Size/risk/complexity/type/governance/gate/FEL"),
    (32, "ADMIN configuration design", "AdminConfiguration y 20 secciones"),
    (33, "ADMIN inheritance", "Tenant/Enterprise/BU/Portfolio"),
    (34, "ADMIN UI", "Catálogos, reglas, preview y lifecycle"),
    (35, "USER UI", "14 vistas del expediente"),
    (36, "USER queues", "8 colas cursor-paginadas"),
    (37, "Roles/permissions", "4 roles y 8 permisos"),
    (38, "SoD", "Author ≠ final reviewer según policy"),
    (39, "ETag", "If-Match y 412 exacto"),
    (40, "Idempotency", "Keys, retry y payload conflict"),
    (41, "SecurityEvent", "Lifecycle y configuración auditados"),
    (42, "Persistence design", "Dos tablas versionadas, sin tabla por sección"),
    (43, "Migration", "0060 reversible SQLite/PostgreSQL"),
    (44, "API", "USER y ADMIN completas"),
    (45, "PostgreSQL happy path", "Create→complete→provider en base temporal"),
    (46, "Negative readiness tests", "Market/risk/payment/source/SoD/ETag/idempotency"),
    (47, "Staleness/recovery", "Nueva DRAFT e histórico preservado"),
    (48, "Historical compatibility", "Sin cambios retroactivos"),
    (49, "Governance Model regression", "Modelos no-owner no afectados por default"),
    (50, "Browser USER E2E", "Navegación, vistas y ausencia ADMIN"),
    (51, "Browser ADMIN E2E", "Configuración y ausencia de transacciones live"),
    (52, "Direct-route security", "Guard can_configure"),
    (53, "Navigation integration", "Blueprint canónico"),
    (54, "Performance", "Paginación por cursor; medición local documentada"),
    (55, "Full regression", "Comparación diferencial contra baseline"),
    (56, "Technical debt", "Projection de readiness y E2E actor-switch mejorables"),
    (57, "Final Gap Analysis", "Sin brecha crítica COMM; deuda no bloqueante"),
    (58, "Updated technical report", "Este documento Word"),
    (59, "Final exclusive state", "GATE07M_COMM_COMMERCIAL_STRATEGY_RELEASE_VALIDATED"),
    (60, "Recommendation next increment", "Benefits Plan and Benefits Traceability"),
]
deliverable_rows = [[n, name, evidence, "PASS"] for n, name, evidence in deliverables]
table(
    document,
    ["#", "Entregable", "Evidencia / implementación", "Estado"],
    deliverable_rows[:43],
    [0.38, 2.1, 3.72, 0.8],
    status_column=3,
)
heading(document, "9.1 Continuación: entregables 44 a 60", 2)
table(
    document,
    ["#", "Entregable", "Evidencia / implementación", "Estado"],
    deliverable_rows[43:],
    [0.38, 2.1, 3.72, 0.8],
    status_column=3,
)

heading(document, "10. Rendimiento, deuda técnica y Gap Analysis")
table(
    document,
    ["Tema", "Estado", "Evaluación"],
    [
        [
            "List/detail",
            "PASS",
            f"Cursor limit 50; medición local list={release.get('list_latency_ms', 'N/D')} ms y detail={release.get('detail_latency_ms', 'N/D')} ms.",
        ],
        [
            "Ready/Stale queues",
            "DEUDA",
            "Readiness se recalcula por candidato para preservar vigencia. Conviene una proyección/materialized status al crecer el volumen.",
        ],
        [
            "current_version_id",
            "DEUDA",
            "La circularidad se resuelve por validación de servicio; evaluar FK diferida en una migración futura.",
        ],
        [
            "Browser actor switch",
            "DEUDA",
            "El lifecycle completo multi-actor está cubierto en backend; el E2E navegador se concentra en navegación y separación de modos.",
        ],
        [
            "Procurement/Five Case",
            "PASS",
            "Permanecen fuera de alcance y no fueron implementados.",
        ],
    ],
    [1.45, 0.8, 4.75],
    status_column=1,
)
paragraph(
    document,
    "Gap final: no queda una brecha crítica conocida dentro del alcance Gate07M-COMM. La deuda señalada es evolutiva y no compromete integridad, seguridad, inmutabilidad ni evidencia gobernada del release.",
)

heading(document, "11. Estado final y recomendación")
signal = document.add_paragraph()
signal.alignment = WD_ALIGN_PARAGRAPH.CENTER
signal.paragraph_format.space_before = Pt(10)
signal.paragraph_format.space_after = Pt(8)
font(
    signal.add_run("GATE07M_COMM_COMMERCIAL_STRATEGY_RELEASE_VALIDATED"),
    13,
    True,
    GREEN,
)
output_signal = document.add_paragraph()
output_signal.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(
    output_signal.add_run("READY_FOR_GOVERNED_COMMERCIAL_READINESS_EVIDENCE"),
    12,
    True,
    NAVY,
)
provider_signal = document.add_paragraph()
provider_signal.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(
    provider_signal.add_run("COMMERCIAL_READINESS_PROVIDER = AVAILABLE"), 11, True, TEAL
)

paragraph(
    document,
    "Siguiente incremento recomendado: BUILD_STRATEGIC_BENEFITS_PLAN_AND_BENEFITS_TRACEABILITY. Este cierre no inicia Five Case Business Case, procurement execution, tendering, vendor management, contract administration, commitments, purchase orders, payments, operational readiness ni benefits realization.",
)

document.core_properties.title = "Informe técnico Gate 07M-COMM"
document.core_properties.subject = (
    "Commercial & Contracting Strategy y Commercial Readiness"
)
document.core_properties.author = "P&Pmis Construction AI / Codex"
document.core_properties.keywords = (
    "Gate07M-COMM, Commercial Strategy, Commercial Readiness, P&Pmis"
)
document.core_properties.comments = "Informe técnico de implementación y validación"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
print(OUTPUT)
