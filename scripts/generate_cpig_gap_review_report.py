from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "69-cpig-gap-review-target-architecture-pdri.md"
QA = ROOT / "outputs" / ".qa" / "cpig-gap-review"
DIAGRAM_DIR = QA / "diagrams"
OUTPUT = (
    ROOT
    / "outputs"
    / "Informe_Tecnico_CPIG_Gap_Review_Arquitectura_Objetivo_PDRI_2026-09-21.docx"
)
LOGO = ROOT / "frontend" / "public" / "pypmis-construction-ai-logo.png"

BLACK = "000000"
TEXT = "202A33"
MUTED = "53616B"
NAVY = "17324D"
TEAL = "117D82"
WHITE = "FFFFFF"
PALE_BLUE = "EAF2F7"
PALE_TEAL = "E8F5F3"
PALE_GREEN = "E9F5EE"
PALE_AMBER = "FFF4E5"
PALE_GRAY = "F5F7F8"
LINE = "CCD5DC"


DIAGRAMS: list[dict[str, Any]] = [
    {
        "name": "01-current-strategic-map.png",
        "title": "Mapa estratégico actual",
        "rows": [
            ["0 Foundation", "1 Idea & Demand", "2 Proposal", "3 Strategic Gate"],
            ["7 Governance", "6 Definition & Case", "5 Scenarios", "4 Portfolio"],
        ],
        "edges": [(0, i, 0, i + 1) for i in range(3)] + [(0, 3, 1, 3)] + [(1, i, 1, i - 1) for i in range(3, 0, -1)],
        "alt": "Flujo actual de ocho fases estratégicas desde Enterprise Strategy Foundation hasta Investment Governance.",
    },
    {
        "name": "02-target-strategic-map.png",
        "title": "Arquitectura estratégica objetivo",
        "rows": [
            ["0 Foundation + Appetite", "1 Capital Planning & Demand", "2 Origination & Shaping", "3 Strategic Gate"],
            ["7 Investment Governance", "6 Investment Development", "5 Scenario Planning", "4 Portfolio Planning & Review"],
        ],
        "edges": [(0, i, 0, i + 1) for i in range(3)] + [(0, 3, 1, 3)] + [(1, i, 1, i - 1) for i in range(3, 0, -1)],
        "alt": "Arquitectura objetivo de ocho fases con Capital Planning, Portfolio Review e Investment Development.",
    },
    {
        "name": "03-gap-overlay.png",
        "title": "Superposición de brechas y reutilización PDRI",
        "rows": [
            ["Capacidades vigentes", "PDRI vigente"],
            ["Capital Demand", "Strategic Buckets", "Portfolio Review"],
            ["Basic Data", "QRA / Confidence", "Owner Capability"],
        ],
        "edges": [
            (0, 0, 1, 0), (0, 0, 1, 1), (0, 0, 1, 2),
            (0, 1, 2, 0), (0, 1, 2, 1), (0, 0, 2, 2),
        ],
        "alt": "Mapa de seis gaps y relación de reutilización con las capacidades y PDRI vigentes.",
    },
    {
        "name": "04-demand-chain.png",
        "title": "Cadena Strategic Capital Demand",
        "rows": [
            ["Objectives", "Investment Themes", "Planning Cycle", "Capital Demand"],
            ["Proposal", "BA / Options", "Idea", "Originate Idea"],
        ],
        "edges": [(0, i, 0, i + 1) for i in range(3)] + [(0, 3, 1, 3)] + [(1, i, 1, i - 1) for i in range(3, 0, -1)],
        "alt": "Trazabilidad desde Strategic Objectives hasta Project Proposal mediante Capital Demand e Idea.",
    },
    {
        "name": "05-basic-data-pdri-qra.png",
        "title": "Relación Basic Data, PDRI y QRA",
        "rows": [
            ["Basic Data", "Project Definition", "PDRI / FEL"],
            ["QRA / Confidence", "Five Case Business Case"],
        ],
        "edges": [(0, 0, 0, 1), (0, 1, 0, 2), (0, 0, 1, 0), (0, 2, 1, 0), (1, 0, 1, 1)],
        "alt": "Basic Data aporta contexto a Project Definition y QRA; PDRI se conserva separado y ambos alimentan el Business Case.",
    },
    {
        "name": "06-domain-authority.png",
        "title": "Mapa de autoridades de dominio",
        "rows": [
            ["Enterprise Strategy", "Capital Planning", "Idea / BA / Proposal"],
            ["Portfolio", "Scenarios", "Investment Development"],
            ["Project Definition / PDRI", "Risk / QRA", "Five Case / Gate / FID"],
        ],
        "edges": [
            (0, 0, 0, 1), (0, 1, 0, 2), (0, 2, 1, 0), (1, 0, 1, 1),
            (1, 1, 1, 2), (1, 2, 2, 0), (1, 2, 2, 1), (2, 0, 2, 2), (2, 1, 2, 2),
        ],
        "alt": "Autoridades únicas desde Enterprise Strategy hasta FID, con Investment Development como capa de preparación.",
    },
    {
        "name": "07-provider-architecture.png",
        "title": "Arquitectura de proveedores Gate07J",
        "rows": [
            ["Project Definition / PDRI", "Basic Data futuro", "QRA futuro", "Owner Capability futuro"],
            ["Provider PDRI existente", "Provider Basic Data", "Provider QRA", "Provider Owner Capability"],
            ["Gate07J Registry", "Policy clone / preview / publish", "Evidence snapshot"],
        ],
        "edges": [
            (0, 0, 1, 0), (0, 1, 1, 1), (0, 2, 1, 2), (0, 3, 1, 3),
            (1, 0, 2, 0), (1, 1, 2, 0), (1, 2, 2, 0), (1, 3, 2, 0),
            (2, 0, 2, 1), (2, 1, 2, 2),
        ],
        "alt": "Provider PDRI existente y tres providers futuros convergen en Gate07J y una publicación controlada de policy.",
    },
    {
        "name": "08-dependency-graph.png",
        "title": "Dependencias objetivo",
        "rows": [
            ["Objectives / Themes", "Capital Planning", "Buckets / Review"],
            ["Basic Data", "PDRI / FEL", "Risk + Appetite", "Owner Capability"],
            ["QRA", "Five Case", "Gate07J futuro"],
        ],
        "edges": [
            (0, 0, 0, 1), (0, 1, 0, 2), (1, 0, 1, 1), (1, 0, 2, 0),
            (1, 1, 2, 0), (1, 2, 2, 0), (1, 3, 2, 1), (2, 0, 2, 1),
            (2, 1, 2, 2), (0, 2, 2, 2),
        ],
        "alt": "Grafo de dependencias desde Objectives y Basic Data hasta Five Case y una futura policy Gate07J.",
    },
    {
        "name": "09-roadmap.png",
        "title": "Roadmap por olas",
        "rows": [
            ["Ola 1\nCapital Demand & Planning"],
            ["Ola 2A\nStrategic Buckets", "Ola 2B\nAppetite + Basic Data"],
            ["Ola 3A\nPortfolio Review", "Ola 3B\nOwner Capability"],
            ["Ola 4\nQRA / Confidence"],
            ["Ola 5\nFive Case + Gate07J"],
        ],
        "edges": [
            (0, 0, 1, 0), (0, 0, 1, 1), (1, 0, 2, 0), (1, 1, 2, 1),
            (3, 0, 4, 0),
        ],
        "alt": "Roadmap de cinco olas que inicia con Capital Demand y termina con integración controlada Five Case y Gate07J.",
    },
]


def font(path: str, size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidate = Path("C:/Windows/Fonts") / path
    try:
        return ImageFont.truetype(str(candidate), size=size)
    except OSError:
        return ImageFont.load_default()


def wrap(draw: ImageDraw.ImageDraw, text: str, face: ImageFont.ImageFont, width: int) -> list[str]:
    result: list[str] = []
    for raw in text.split("\n"):
        words = raw.split()
        line = ""
        for word in words:
            trial = f"{line} {word}".strip()
            if draw.textbbox((0, 0), trial, font=face)[2] <= width or not line:
                line = trial
            else:
                result.append(line)
                line = word
        result.append(line)
    return result


def draw_arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int]) -> None:
    color = (83, 97, 107)
    draw.line([start, end], fill=color, width=5)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex > sx else -1
        points = [(ex, ey), (ex - direction * 18, ey - 10), (ex - direction * 18, ey + 10)]
    else:
        direction = 1 if ey > sy else -1
        points = [(ex, ey), (ex - 10, ey - direction * 18), (ex + 10, ey - direction * 18)]
    draw.polygon(points, fill=color)


def render_diagram(spec: dict[str, Any]) -> Path:
    width, height = 1600, 900
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_face = font("arialbd.ttf", 42)
    node_size = 26 if len(spec["rows"]) >= 5 else 32
    node_face = font("arial.ttf", node_size)
    draw.text((80, 46), spec["title"], fill=spec.get("title_color", (23, 50, 77)), font=title_face)
    if spec.get("title_rule", True):
        draw.line((80, 112, width - 80, 112), fill=(17, 125, 130), width=5)

    rows: list[list[str]] = spec["rows"]
    top, bottom = 165, height - 80
    row_height = (bottom - top) / max(len(rows), 1)
    positions: dict[tuple[int, int], tuple[int, int, int, int]] = {}
    for row_index, row in enumerate(rows):
        count = len(row)
        gap = 28
        available = width - 160 - gap * (count - 1)
        box_width = min(390, int(available / max(count, 1)))
        total_width = box_width * count + gap * (count - 1)
        left = (width - total_width) // 2
        box_height = min(130, int(row_height * (0.82 if len(rows) >= 5 else 0.62)))
        y = int(top + row_index * row_height + (row_height - box_height) / 2)
        for col_index, label in enumerate(row):
            x = left + col_index * (box_width + gap)
            positions[(row_index, col_index)] = (x, y, x + box_width, y + box_height)

    for r1, c1, r2, c2 in spec["edges"]:
        a = positions[(r1, c1)]
        b = positions[(r2, c2)]
        if r1 == r2:
            if c2 > c1:
                start = (a[2] + 3, (a[1] + a[3]) // 2)
                end = (b[0] - 5, (b[1] + b[3]) // 2)
            else:
                start = (a[0] - 3, (a[1] + a[3]) // 2)
                end = (b[2] + 5, (b[1] + b[3]) // 2)
        else:
            start = ((a[0] + a[2]) // 2, a[3] + 3)
            end = ((b[0] + b[2]) // 2, b[1] - 5)
        draw_arrow(draw, start, end)

    fills = [(234, 242, 247), (232, 245, 243), (233, 245, 238), (255, 244, 229)]
    for (row_index, col_index), box in positions.items():
        draw.rounded_rectangle(box, radius=18, fill=fills[row_index % len(fills)], outline=(23, 50, 77), width=4)
        lines = wrap(draw, rows[row_index][col_index], node_face, box[2] - box[0] - 34)
        line_height = 30 if len(rows) >= 5 else 35
        text_height = len(lines) * line_height
        y = box[1] + (box[3] - box[1] - text_height) // 2
        for line in lines:
            bounds = draw.textbbox((0, 0), line, font=node_face)
            x = box[0] + (box[2] - box[0] - (bounds[2] - bounds[0])) // 2
            draw.text((x, y), line, fill=(32, 42, 51), font=node_face)
            y += line_height

    path = DIAGRAM_DIR / spec["name"]
    if spec["name"] == "09-roadmap.png":
        image = image.crop((200, 0, 1400, height))
    image.save(path, format="PNG", optimize=True)
    return path


def set_cell_shading(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell: Any, value: int = 85) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side in ("top", "start", "bottom", "end"):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table: Any) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "5")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), LINE)


def repeat_header(row: Any) -> None:
    properties = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    properties.append(flag)


def set_alt_text(inline_shape: Any, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def set_run_font(run: Any, size: float | None = None, bold: bool | None = None, color: str = TEXT) -> None:
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Aptos")
    run.font.color.rgb = RGBColor.from_string(color)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_inline(paragraph: Any, text: str, *, size: float | None = None) -> None:
    pattern = re.compile(r"(\*\*.*?\*\*|`.*?`)")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run_font(run, size=size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, color=NAVY)
            run.font.name = "Aptos Mono"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Aptos Mono")
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=size)


def clean_cell(value: str) -> str:
    value = value.strip()
    value = re.sub(r"\*\*(.*?)\*\*", r"\1", value)
    value = re.sub(r"`(.*?)`", r"\1", value)
    return value.replace("<br>", "\n")


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    headers = rows[0]
    bodies = rows[1:]
    groups: list[list[int]]
    if len(headers) >= 7:
        groups = []
        start = 1
        while start < len(headers):
            groups.append([0, *range(start, min(start + 5, len(headers)))])
            start += 5
    else:
        groups = [list(range(len(headers)))]

    for part, columns in enumerate(groups, start=1):
        if len(groups) > 1:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(3)
            run = paragraph.add_run(f"Matriz integral — parte {part} de {len(groups)}")
            set_run_font(run, size=8.5, bold=True, color=MUTED)
        table = document.add_table(rows=1, cols=len(columns))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        set_table_borders(table)
        for index, source_index in enumerate(columns):
            cell = table.rows[0].cells[index]
            cell.text = clean_cell(headers[source_index])
            set_cell_shading(cell, NAVY)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, size=7.0 if len(columns) > 5 else 8.0, bold=True, color=WHITE)
        repeat_header(table.rows[0])
        for row_index, source_row in enumerate(bodies):
            row = table.add_row()
            for index, source_index in enumerate(columns):
                value = source_row[source_index] if source_index < len(source_row) else ""
                cell = row.cells[index]
                cell.text = clean_cell(value)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                set_cell_margins(cell)
                if row_index % 2:
                    set_cell_shading(cell, PALE_GRAY)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(1)
                    for run in paragraph.runs:
                        set_run_font(run, size=6.7 if len(columns) > 5 else 8.0)
        document.add_paragraph().paragraph_format.space_after = Pt(1)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.35)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size in (("Title", 24), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10.5)):
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)

    core = document.core_properties
    core.title = "Informe técnico de revisión de brechas CPIG y arquitectura objetivo"
    core.subject = "Capital Portfolio & Investment Governance con reutilización PDRI"
    core.author = "P&Pmis Construction AI"
    core.keywords = "CPIG, arquitectura, PDRI, Capital Planning, QRA, Gate07J"
    core.comments = "Artefacto de arquitectura; no implementa cambios productivos."


def add_field(paragraph: Any, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, separate, end])


def add_headers_and_footers(document: Document) -> None:
    for section in document.sections:
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = header.add_run("P&Pmis Construction AI  |  CPIG Architecture Review")
        set_run_font(run, size=7.5, bold=True, color=MUTED)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Confidencial — 2026-09-21  |  Página ")
        set_run_font(run, size=7.5, color=MUTED)
        add_field(footer, "PAGE")
        run = footer.add_run(" de ")
        set_run_font(run, size=7.5, color=MUTED)
        add_field(footer, "NUMPAGES")


def add_cover(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(36)
    if LOGO.exists():
        shape = paragraph.add_run().add_picture(str(LOGO), width=Inches(1.35))
        set_alt_text(shape, "P&Pmis Construction AI", "Logotipo azul de P&Pmis Construction AI.")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(24)
    run = paragraph.add_run("Informe técnico de revisión de brechas CPIG y arquitectura objetivo")
    set_run_font(run, size=23, bold=True, color=BLACK)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Capital Portfolio & Investment Governance con reutilización exacta de PDRI")
    set_run_font(run, size=12.5, color=NAVY)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(24)
    add_inline(paragraph, "Arquitectura y gap review · 21 de septiembre de 2026", size=10)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(42)
    run = paragraph.add_run("RESULTADO")
    set_run_font(run, size=8, bold=True, color=MUTED)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("ARQUITECTURA OBJETIVO VALIDADA")
    set_run_font(run, size=14, bold=True, color=TEAL)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(18)
    run = paragraph.add_run("Límite: documentación y decisiones arquitectónicas; cero cambios productivos")
    set_run_font(run, size=9, color=MUTED)
    document.add_page_break()


def add_document_control(document: Document) -> None:
    document.add_heading("Control del documento", level=1)
    add_table(document, [
        ["Campo", "Valor"],
        ["Producto", "P&Pmis Construction AI"],
        ["Macroproceso", "Capital Portfolio & Investment Governance (CPIG)"],
        ["Tipo", "Revisión de brechas y arquitectura objetivo"],
        ["Rama / commit", "main / 9cb0969"],
        ["Alembic head en código", "20260916_0061"],
        ["Fecha de corte", "2026-09-21"],
        ["Estado", "VALIDADO PARA PLANIFICACIÓN; SIN CAMBIOS PRODUCTIVOS"],
    ])
    document.add_heading("Resumen para decisión", level=1)
    paragraph = document.add_paragraph()
    add_inline(paragraph, (
        "Se confirmaron seis brechas reales. PDRI ya está implementado y debe reutilizarse sin crear un segundo "
        "dominio o provider. La primera ola recomendada es Strategic Capital Demand & Capital Planning Foundation; "
        "las capacidades restantes se secuencian por dependencias y las integraciones Gate07J se reservan para una "
        "policy futura clone/preview/publish."
    ))
    document.add_heading("Contenido", level=1)
    for item in (
        "Baseline y mapas actual/objetivo",
        "Matriz integral de brechas y current versus target",
        "Diseño de los seis gaps y reutilización PDRI",
        "Autoridades, providers, ADMIN/USER y dependencias",
        "Historia, rendimiento, roadmap, riesgos y cierre",
    ):
        paragraph = document.add_paragraph(style="List Bullet")
        add_inline(paragraph, item)
    document.add_page_break()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    raw: list[str] = []
    index = start
    while index < len(lines) and lines[index].lstrip().startswith("|"):
        raw.append(lines[index].strip())
        index += 1
    rows = [[cell.strip() for cell in line.strip("|").split("|")] for line in raw]
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in rows[1]):
        rows.pop(1)
    return rows, index


def add_source_content(
    document: Document,
    diagram_paths: list[Path],
    *,
    page_break_heading_prefixes: tuple[str, ...] = (
        "4. ",
        "7. ",
        "9. ",
        "13. ",
        "18. ",
    ),
    skip_preface: bool = False,
) -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    if skip_preface:
        first_section = next((index for index, line in enumerate(lines) if line.startswith("## ")), len(lines))
        lines = lines[first_section:]
    index = 0
    diagram_index = 0
    in_code = False
    code_language = ""
    code_lines: list[str] = []
    first_h1_skipped = False

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if in_code:
            if stripped == "```":
                if code_language == "mermaid":
                    spec = DIAGRAMS[diagram_index]
                    path = diagram_paths[diagram_index]
                    paragraph = document.add_paragraph()
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.keep_together = True
                    shape = paragraph.add_run().add_picture(str(path), width=Inches(6.75))
                    set_alt_text(shape, spec["title"], spec["alt"])
                    paragraph.add_run("\n")
                    run = paragraph.add_run(
                        f"Figura {diagram_index + 1}. {spec['title']}"
                    )
                    set_run_font(run, size=8, color=MUTED)
                    diagram_index += 1
                else:
                    table = document.add_table(rows=1, cols=1)
                    set_table_borders(table)
                    repeat_header(table.rows[0])
                    cell = table.cell(0, 0)
                    set_cell_shading(cell, PALE_GRAY)
                    cell.text = "\n".join(code_lines)
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.name = "Aptos Mono"
                            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Aptos Mono")
                            run.font.size = Pt(8)
                in_code = False
                code_language = ""
                code_lines = []
            else:
                code_lines.append(raw)
            index += 1
            continue

        if stripped.startswith("```"):
            in_code = True
            code_language = stripped[3:].strip().lower()
            index += 1
            continue
        if not stripped:
            index += 1
            continue
        if stripped == "<!-- PAGE_BREAK -->":
            document.add_page_break()
            index += 1
            continue
        if stripped.startswith("# "):
            if not first_h1_skipped:
                first_h1_skipped = True
            index += 1
            continue
        if stripped.startswith("### "):
            document.add_heading(stripped[4:], level=2)
            index += 1
            continue
        if stripped.startswith("## "):
            heading = stripped[3:]
            if heading.startswith(page_break_heading_prefixes):
                document.add_page_break()
            document.add_heading(heading, level=1)
            index += 1
            continue
        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(document, rows)
            continue
        if stripped.startswith("- [x] "):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(1)
            add_inline(paragraph, f"✓ {stripped[6:]}")
            index += 1
            continue
        if stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:])
            index += 1
            continue
        if re.match(r"^\d+\. ", stripped):
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\. ", "", stripped))
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or candidate.startswith(("#", "|", "- ", "- [", "```")) or re.match(r"^\d+\. ", candidate):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = document.add_paragraph()
        add_inline(paragraph, " ".join(paragraph_lines))

    if diagram_index != len(DIAGRAMS):
        raise RuntimeError(f"Expected {len(DIAGRAMS)} Mermaid diagrams, rendered {diagram_index}")


def set_language(document: Document) -> None:
    for style in document.styles:
        if style.type != 1:
            continue
        rpr = style._element.get_or_add_rPr()
        lang = rpr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            rpr.append(lang)
        lang.set(qn("w:val"), "es-CO")


def update_field_settings(document: Document) -> None:
    settings = document.settings._element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    QA.mkdir(parents=True, exist_ok=True)
    DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    diagram_paths = [render_diagram(spec) for spec in DIAGRAMS]

    document = Document()
    configure_document(document)
    set_language(document)
    add_cover(document)
    add_document_control(document)
    add_source_content(document, diagram_paths)
    add_headers_and_footers(document)
    update_field_settings(document)
    document.save(OUTPUT)
    print(f"created={OUTPUT}")
    print(f"diagrams={len(diagram_paths)}")
    print(f"source={SOURCE}")


if __name__ == "__main__":
    main()
