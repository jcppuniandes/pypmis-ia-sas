"""Preset and structural audit for the final Gate 07J-FEL Word report."""

from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

REPO = Path(__file__).resolve().parents[1]
DOCX = REPO / "outputs" / "technical_reports" / "Informe_Tecnico_Gate_07J_FEL_Project_Definition_20260826.docx"
OUTPUT = REPO / "outputs" / "technical_reports" / "gate07j_fel_preset_audit.json"


def close(actual: int | None, expected: int, tolerance: int = 1000) -> bool:
    return actual is not None and abs(actual - expected) <= tolerance


def run() -> None:
    doc = Document(DOCX)
    section = doc.sections[0]
    body_text = [p.text for p in doc.paragraphs]
    table_text = [p.text for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]
    text = " ".join(body_text + table_text)
    checks: dict[str, bool] = {
        "single_portrait_section": len(doc.sections) == 1 and section.page_height > section.page_width,
        "letter_page_size": close(section.page_width, Inches(8.5)) and close(section.page_height, Inches(11)),
        "one_inch_margins": all(
            close(value, Inches(1))
            for value in (section.left_margin, section.right_margin, section.top_margin, section.bottom_margin)
        ),
        "header_footer_distance": close(section.header_distance, Inches(0.492))
        and close(section.footer_distance, Inches(0.492)),
        "normal_11pt": close(doc.styles["Normal"].font.size, Pt(11), 100),
        "heading1_16pt": close(doc.styles["Heading 1"].font.size, Pt(16), 100),
        "heading2_13pt": close(doc.styles["Heading 2"].font.size, Pt(13), 100),
        "heading1_blue": str(doc.styles["Heading 1"].font.color.rgb) == "2E74B5",
        "heading2_blue": str(doc.styles["Heading 2"].font.color.rgb) == "2E74B5",
        "heading_hierarchy_present": sum(p.style.name == "Heading 1" for p in doc.paragraphs) == 10
        and sum(p.style.name == "Heading 2" for p in doc.paragraphs) == 54,
        "running_header_present": "TECHNICAL REPORT - GATE 07J-FEL" in section.header.paragraphs[0].text,
        "even_running_header_present": "TECHNICAL REPORT - GATE 07J-FEL"
        in section.even_page_header.paragraphs[0].text,
        "running_footer_present": "P&Pmis Construction AI" in section.footer.paragraphs[0].text,
        "even_running_footer_present": "P&Pmis Construction AI" in section.even_page_footer.paragraphs[0].text,
        "image_alt_present": all(shape._inline.docPr.get("descr") for shape in doc.inline_shapes),
        "no_fake_bullets": not any(p.text.strip().startswith(("•", "- ")) for p in doc.paragraphs),
        "no_placeholder_text": not re.search(
            r"\b(?:todo|lorem ipsum|placeholder|codex-file-citation)\b", text, flags=re.IGNORECASE
        ),
        "exclusive_final_state_present": "GATE07J_FEL_RELEASE_VALIDATED" in text
        and "READY_FOR_GOVERNED_PROJECT_DEFINITION_EVIDENCE" not in text,
        "three_consecutive_browser_runs_present": all(token in text for token in ("Run 1", "Run 2", "Run 3", "3/3")),
        "baseline_protection_present": all(
            token in text for token in ("Project 14 defaults MD5", "eea3648c6908a267d1e0651828ff4459", "MATCH")
        ),
        "stop_boundary_present": all(
            token in text
            for token in ("Investment Decision/FID", "Project Initialization", "Activation", "Execution", "CPM", "P6")
        ),
    }
    table_geometry = True
    row_header_flags = True
    cell_margins = True
    no_split_rows = True
    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        table_geometry &= tbl_pr.find(qn("w:tblW")) is not None and tbl_pr.find(qn("w:tblInd")) is not None
        row_header_flags &= table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None
        for row in table.rows:
            no_split_rows &= row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
            for cell in row.cells:
                cell_margins &= cell._tc.get_or_add_tcPr().find(qn("w:tcMar")) is not None
    checks.update(
        {
            "table_geometry_complete": table_geometry,
            "table_header_flags_complete": row_header_flags,
            "cell_margins_complete": cell_margins,
            "rows_prevent_split": no_split_rows,
        }
    )
    report = {
        "preset": "standard_business_brief",
        "first_page_pattern": "memo_masthead",
        "file": str(DOCX),
        "page_geometry": "Letter portrait, 1-inch margins",
        "tables": len(doc.tables),
        "inline_images": len(doc.inline_shapes),
        "checks": checks,
        "result": "PASS" if all(checks.values()) else "FAIL",
        "failed": [key for key, value in checks.items() if not value],
        "direct_formatting_exceptions": "Masthead, metadata, table cells, callouts, headers and footers only.",
    }
    OUTPUT.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(report, indent=2, ensure_ascii=False))
    if report["result"] != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    run()
