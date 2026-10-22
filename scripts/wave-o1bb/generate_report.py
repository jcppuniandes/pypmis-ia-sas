"""Author the 34-chapter report using actual release evidence and retained Word parts."""

import hashlib
import importlib.util
import io
import json
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
REFERENCE = ROOT / "outputs/wave-o1ba-20261020/P&Pmis_Wave_O1B_A_Scope_Foundation_Technical_Report.docx"
REFERENCE_SHA = "c14b3547fa3b96ce3c08cc30494c6760fc83f33df19dff0cda127f8d21ddc9d1"
OUT = ROOT / "outputs/wave-o1bb-20261020"
QA = ROOT / "outputs/.qa/wave-o1bb-20261020"
DATA = ROOT / "docs/validation/wave-o1bb-20261020/package-data.json"
spec = importlib.util.spec_from_file_location("schedule_word_components", ROOT / "scripts/wave-o1a/generate_report.py")
components = importlib.util.module_from_spec(spec)
spec.loader.exec_module(components)
components.base.REFERENCE = REFERENCE
p, h, table = components.p, components.h, components.table


def build(data):
    document = Document(REFERENCE)
    logo = deepcopy(document.paragraphs[0]._p)
    section_properties = deepcopy(document._element.body.sectPr)
    for node in list(document._element.body):
        document._element.body.remove(node)
    document._element.body.append(logo)
    document._element.body.append(section_properties)
    p(document, "Wave O1B B\nFundación de cronogramas y gobierno P6", "Title")
    p(document, "P&Pmis Construction AI\nInforme técnico de implementación\n21 octubre 2026", "Subtitle")
    p(document, "Implementamos Schedule Manager como dominio dependiente del Scope aprobado y de los Plan Items. "
      + ("La validación técnica del entorno aislado quedó cerrada." if data["release_validated"] else "La liberación permanece pendiente de los controles señalados en este informe.")
      + " La base habitual no fue migrada y la aceptación de negocio continúa sin firmas.")
    chapters = [
        ("Resumen ejecutivo", "Schedule administra registros, versiones, WBS, actividades, hitos, calendarios y relaciones. El intercambio manual XML/XER pasa por staging, validación, reconciliación y aceptación independiente. Current y Baseline conservan punteros separados. Las cantidades y la descomposición de ejecución siguen bajo Scope; Primavera P6 conserva los cálculos CPM. El libro adjunto reúne requisitos, pruebas y API sin sustituir la evidencia ejecutada."),
        ("Línea base", f"Partimos de main en {data['baseline_head']}, con el repositorio inicialmente limpio y O1A y O1B A previamente validadas. Reutilizamos Project Workspace canónico, binding legado, autorización operacional, perfiles y Project Controls Period. La validación nueva usa PostgreSQL aislado, API 8853 y frontend 5513. No se realizó commit ni push en este incremento."),
        ("Hallazgos del repositorio", "El parser histórico XER reconoce marcadores de un adaptador legado, no todas las tablas XER nativas. La ingesta histórica también mezcla efectos de Cost y valores de float predeterminados. Por ello conserva compatibilidad detrás de sus rutas existentes y no se usa como autoridad del dominio canónico. El inventario clasifica reutilización, extensión y adaptadores; no se eliminó código legado sin evidencia de paridad."),
        ("Modelo de autoridad de Schedule", "Enterprise Structure conserva la identidad de proyecto. Scope conserva Scope Items y Plan Items. Schedule gobierna WBS, actividades, versiones, validación, mapeos, manifiestos, Current y Baselines. P6 es la fuente de fechas CPM, float y criticidad importados. No se implementaron cálculos de ruta crítica, nivelación ni recálculo de fechas. Los campos importados de avance se conservan solo como referencia de origen."),
        ("Relación entre Scope Plan y Schedule", "El servicio exige contexto O1A válido, perfil Scope publicado, Scope Baseline aprobado y perfil Schedule vigente. Los mapeos resuelven Activity o WBS hacia varios Plan Items y, a través de sus asignaciones, hacia varios Scope Items. El vínculo directo con Scope depende de política y no puede contradecir el Plan. Ninguna importación crea ni modifica entregables o cantidades Scope."),
        ("Schedule Register", "Cada registro pertenece al Project Workspace canónico y conserva código, nombre, tipo, propósito, fuente, calendario, zona horaria, autor, estado y ETag. Los tipos se publican en ADMIN y pueden habilitar varias corrientes. Crear un registro exige asignación al proyecto y permiso granular. La misma Idempotency Key repite el resultado; un payload distinto con la misma clave se rechaza."),
        ("Versiones del cronograma", "Los cambios materiales crean nuevas versiones con snapshot inmutable, revisión, predecesor y hashes de configuración, contenido y fuente. El ciclo es DRAFT o STAGED, VALIDATED, UNDER_REVIEW y ACCEPTED_CURRENT, con rechazo o supersesión explícitos. El estado Baseline se gobierna en un agregado separado. No se permite sobrescribir la fuente, ni siquiera mientras está en staging. Una corrección manual crea una sucesora."),
        ("WBS", "El árbol WBS conserva identidad externa, código, nombre, padre, ruta, orden y estado. No se obliga a que coincida con la jerarquía Scope. La interfaz presenta nodos jerárquicos e inspección de mapeos. Mover un nodo exige una nueva versión y aparece en el diff. La validación bloquea ciclos, padres inexistentes y referencias Activity a WBS no resueltas."),
        ("Actividades e hitos", "Las actividades conservan código, nombre, tipo, estado, WBS, calendario, duraciones originales y restantes, fechas actuales y planificadas, restricciones y porcentajes importados. Los hitos de inicio y fin son actividades de tipo explícito, no Scope Items. El mapeo permite varios Plan Items por actividad y varias actividades por Plan Item. La creación manual no puede inventar resultados CPM reservados a una fuente importada."),
        ("Calendarios", "El subconjunto importado conserva referencias, horas por día, semana, turnos, excepciones, festivos, zona y campos originales. Las pruebas usan calendarios de diez y seis horas. No se supone una jornada universal de ocho horas. Las bases o conversiones no disponibles quedan UNKNOWN o bloquean la operación. Los calendarios XER opacos se preservan, pero bloquean una exportación XML que no pueda reproducirlos sin pérdida."),
        ("Relaciones", "Se conservan FS, SS, FF y SF, predecesor y sucesor, lags con signo y unidad y referencias externas. El lag negativo no se borra; se evalúa según política. Las relaciones entre proyectos o hacia actividades no disponibles quedan identificadas como externas o no resueltas. No se fabrican actividades internas para cerrar enlaces. La política por defecto bloquea enlaces no resueltos; la UAT usa una excepción WARN publicada explícitamente."),
        ("Restricciones y códigos", "Las restricciones primarias y secundarias conservan tipo y fecha sin convertir Mandatory Start en otra clase. ADMIN define tipos permitidos y tratamiento de restricciones duras. Los códigos de actividad y sus asignaciones se mantienen separados de WBS y de las clasificaciones Scope. Los campos y objetos no soportados se declaran en el manifiesto en lugar de prometer compatibilidad total."),
        ("Data Date", "Cada versión conserva Data Date explícito del archivo o de la definición manual. La fecha no se reemplaza por el cutoff ni por la hora de recepción. Una asociación con Project Controls Period requiere un comando autorizado y registra ambos valores y la versión del período. La prueba usa Data Date distinto del cutoff; no crea costos, avance aprobado ni una nueva autoridad de período."),
        ("Límite con Primavera P6", "El flujo admite archivos manuales y no usa API P6 ni sincronización directa. P6 continúa calculando CPM, float, fechas early y late y nivelación. P&Pmis almacena resultados importados e indicadores de calidad no CPM. No se ha ejecutado una importación del XML generado dentro de una instalación Primavera ni se ha validado contra p6apibo.xsd de una instalación licenciada."),
        ("Importación XML y XER", "El importador conserva bytes, nombre, SHA256, identidad P6, Data Date, formato, subconjunto, campos no soportados y validación. Admite XML APIBusinessObjects y tablas XER nativas delimitadas. El adaptador XER histórico se caracteriza por separado. Los archivos inválidos, entidades XML, números no finitos, límites excedidos y snapshots malformados se rechazan. Importar no acepta Current ni aprueba Baseline automáticamente."),
        ("Exportación XML y XER", "Se genera XML real del subconjunto permitido, junto con manifiesto, hash y evento de exportación. El objetivo CURRENT exige una versión aceptada y linaje vigente; BASELINE exige referencia aprobada explícita; DRAFT_FOR_P6 conserva su carácter de borrador. No existe exportación XER. La referencia a Baseline se gobierna en el manifiesto; no se afirma transporte nativo BaselineProject ni equivalencia con baselines XER."),
        ("Diff y reconciliación", "La comparación usa identidades externas y detecta altas, bajas y cambios en WBS, actividades, calendario, relaciones, restricciones, duraciones, fechas actuales, Data Date y códigos. Los movimientos o cambios materiales marcan mapeos heredados para revisión; una identidad desaparecida queda BROKEN. Los vínculos no se borran silenciosamente. El operador revisa diff y validación antes de la aceptación."),
        ("Gobierno del Baseline", "Una versión aceptada origina BASELINE_CANDIDATE, pasa por SUBMITTED y UNDER_REVIEW y requiere aprobación independiente para APPROVED_BASELINE. El snapshot aprobado contiene cronograma, mapeos y hashes y no admite actualización ni eliminación. Una sucesora exige referencia explícita a la baseline anterior y comprueba que siga vigente al aprobar. El workflow completo de Change permanece fuera del alcance."),
        ("Cronograma Current", "Aceptar una actualización cambia solo el puntero Current y marca la versión anterior SUPERSEDED. La baseline aprobada conserva su versión, contenido y hash. La aceptación compara el Current esperado al importar para evitar aceptar una rama sobre un Current que ya cambió. El historial diferencia versiones rechazadas, supersedidas, iniciales, actualizaciones, correcciones y escenarios."),
        ("Float y criticidad", "Total Float, Free Float, criticidad y fechas early y late se conservan solo cuando provienen de la fuente P6 soportada. La ausencia se muestra UNKNOWN. El XER no inventa un campo IsCritical a partir de float ni confunde Longest Path con criticidad. Near critical compara float importado con un umbral ADMIN en horas, sin calcular ruta crítica. Las unidades incompatibles mantienen UNKNOWN."),
        ("Schedule Health", "Las comprobaciones detectan extremos abiertos, calendarios faltantes, mappings incompletos, relaciones externas, lags excesivos o negativos, restricciones duras, duración larga y fechas inválidas. ADMIN configura umbrales y tratamiento de bloqueo o advertencia. Un resultado de calidad no se presenta como cálculo CPM. Look ahead ofrece un contrato de consulta de veintiocho días configurable, sin implementar el workflow completo de Project Controls."),
        ("Integridad de los mapeos", "Los mapeos son revisiones anexadas con identidad externa, Plan Items, Scope directo, linaje y motivo. VALID, STALE, BROKEN y REQUIRES_REVIEW distinguen casos que requieren reconciliación. REMOVE conserva historia en vez de borrar filas. Un Plan supersedido o un Scope Baseline cambiado invalida la promoción. La aceptación congela el hash de los vínculos; un cambio posterior requiere una nueva versión."),
        ("Separación ADMIN y USER", "ADMIN publica Schedule Profiles con tipos, umbrales, políticas de calendario, restricciones, mapeo y separación de funciones. USER opera el cronograma del proyecto. Las configuraciones publicadas son inmutables y sus cambios crean sucesoras. Las rutas genéricas ADMIN no pueden crear, editar, publicar o clonar un Schedule Profile eludiendo la API dedicada. La configuración publicada no acepta versiones ni aprueba baselines."),
        ("Seguridad y separación de funciones", "Los permisos granulares distinguen lectura, edición de registro y versión, importación, validación, revisión, aceptación, mapeo, baseline, exportación, período y configuración. La autorización combina tenant, proyecto canónico y asignación operacional. Importador y último editor material no pueden aceptar su trabajo cuando SoD está activo. Lo mismo aplica al preparador de baseline y al autor o último editor de configuración."),
        ("API", f"La API dedicada expone {len(data['api'])} rutas bajo /api/v1/schedule. Las familias incluyen perfiles, overview, registros, versiones, importación, transiciones, mappings, baselines, exportación y descarga, asociación a período, consultas temáticas y look ahead. Las escrituras usan Idempotency Key, motivo y If Match cuando corresponde. La matriz adjunta conserva método, ruta y permiso específico. Las lecturas no aprovisionan permisos ni escriben datos de negocio."),
        ("Frontend", "Schedule Manager usa CompactModuleHeader y nueve vistas: Schedule Register, WBS, Activities, Calendars, Relationships, Codes and Constraints, Baselines, Import and Export y Schedule Health. La ruta se dirige al Project Workspace canónico y conserva guards de autorización. Los controles muestran fuente, hashes, Data Date, validación, diff, mapeos e historia. La implementación está separada por feature y no modifica App.tsx."),
        ("Migración", "La única sucesora es 20261021_0071 sobre 20261020_0070. Añade seis tablas: registros, versiones, mapeos, baselines, manifiestos de exportación y asociaciones a período. WBS, actividades y referencias quedan dentro de snapshots versionados, sin duplicar la autoridad Scope. PostgreSQL añade guards de tenant y proyecto e inmutabilidad. Se prueba upgrade, downgrade vacío seguro y re upgrade; no se declara compatible toda la historia desde una base vacía."),
        ("Validación PostgreSQL", f"El paquete registra {data['backend_count']} casos backend únicos en sus suites de Schedule e intercambio, con resultado {data['backend_result']}. La regresión de O1A, Scope y rutas heredadas registra {data['regression_count']} casos con resultado {data['regression_result']}. Se distinguen las pruebas de archivos puros de las que ejecutan HTTP y PostgreSQL real. Los datos son tenants aislados, sin fixtures destructivos sobre el Proyecto 14."),
        ("Validación Chromium", f"La UAT registra {len(data['browser'].get('checks', []))} comprobaciones y resultado {data['browser'].get('result', 'PENDING')}. Usa login real, sesiones emitidas por la API, PostgreSQL y frontend compilado sin mocks de respuestas. El recorrido cubre ADMIN, registro, carga XML y XER, mapeos, aceptación, baseline con actor independiente, exportación y nueva importación. Un parser loop XML no equivale a roundtrip ejecutado en Primavera."),
        ("Preservación del Proyecto 14", f"La comparación antes y después tiene resultado {data['preservation_result']}. Conserva hash de la fila Project Workspace canónica 14 y referencias heredadas. En la base habitual projects.id 14 está ausente; no se confunde esa identidad con el Project Workspace 14. La base de validación tiene sus propios identificadores. No se aplicó la migración 0071 a los datos habituales."),
        ("Observaciones de rendimiento", "Las mediciones corresponden a archivos pequeños y a consultas locales del entorno aislado. Se registran carga de registros, WBS, actividades, parseo XML y XER, diff, mappings y Health. Los tiempos no representan capacidad productiva ni SLA. El sistema comparte recursos con otras aplicaciones; las mediciones permiten reproducir este caso, no extrapolar a cronogramas masivos."),
        ("Brechas pendientes", "La compatibilidad P6 se limita al subconjunto declarado. Quedan fuera export XER, transporte nativo BaselineProject, validación XSD de una instalación Primavera, roundtrip ejecutado en P6, sincronización API y escalabilidad productiva. Cost, Progress aprobado y reproducción 4D no se implementaron. El despliegue sobre la base habitual y la aceptación de negocio siguen pendientes; la validación técnica no los sustituye."),
        ("Decisión de liberación", data["release_state"] + ". " + data["release_explanation"] + " Resultado: " + data["final_output"] + "."),
        ("Única recomendación siguiente", "WAVE O1B-C — COST FOUNDATION. Solo debe iniciarse después del cierre validado de O1B B y de autorización explícita. Deberá consumir Scope y Plan Items vigentes y mantener la identidad canónica y los límites de Schedule. No se implementa automáticamente en este trabajo."),
    ]
    for index, (title, text) in enumerate(chapters, 1):
        h(document, f"{index} {title}")
        p(document, text)
        if index == 31 and data.get("performance"):
            table(document, ["Operación", "N", "Mediana ms", "Máximo ms"],
                  [[row[0], str(row[1]), f"{row[3]:.2f}", f"{row[4]:.2f}"] for row in data["performance"]], [3.15, .7, 1.15, 1.15])
    for paragraph in document.paragraphs:
        paragraph.paragraph_format.widow_control = True
        if paragraph.style.name == "Normal":
            paragraph.paragraph_format.keep_together = True
    return document


def main():
    assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_SHA
    data = json.loads(DATA.read_text("utf-8"))
    OUT.mkdir(parents=True, exist_ok=True)
    QA.mkdir(parents=True, exist_ok=True)
    buffer = io.BytesIO()
    build(data).save(buffer)
    output = OUT / "P&Pmis_Wave_O1B_B_Schedule_Foundation_Technical_Report.docx"
    with ZipFile(REFERENCE) as source, ZipFile(buffer) as generated, ZipFile(output, "w") as final:
        for entry in source.infolist():
            content = generated.read(entry.filename) if entry.filename == "word/document.xml" else source.read(entry.filename)
            if entry.filename == "word/footer1.xml":
                content = content.replace(b"20 octubre 2026", b"21 octubre 2026")
            final.writestr(entry, content)
    with ZipFile(REFERENCE) as source, ZipFile(output) as final:
        changed = [name for name in source.namelist() if source.read(name) != final.read(name)]
        assert set(changed) <= {"word/document.xml", "word/footer1.xml"}
    proof = {"reference_sha256": REFERENCE_SHA, "output_sha256": hashlib.sha256(output.read_bytes()).hexdigest(), "changed_parts": changed, "chapters": 34}
    (QA / "documents-build-manifest.json").write_text(json.dumps(proof, indent=2), encoding="utf-8")
    print(json.dumps(proof))


if __name__ == "__main__":
    main()
